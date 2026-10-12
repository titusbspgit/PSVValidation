#!/usr/bin/env python3
"""
TestPlan Excel Generator - Agent 7
Generates a real .xlsx workbook using openpyxl
"""
import json
import os
import sys
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"testplan_{timestamp}.xlsx"

# Input JSON data
json_data = [{"Index":"1","SS / Module":"MIPI CSI2 Subsystem","Test Case Name":"mipi_csi2_test_pattern_generator","Feature":"NA","Meta Headers":"#include <stdio.h>\n#include <stdlib.h>\n#include \"test_common.h\"\n#include \"mipi_csi2.h\"","Meta Macros":"#define GDMA_CSI2_DATA_DEST_ADDR2 0xE6040080\n#define GDMA_CTRL_DATA_DEST_ADDR2 0xE6040000","Meta Arrays":"NA","Speed":"NA","Mode":"NA","Memory Start Offset":"NA","Memory End Offset":"NA","Meta Test Description":"This testcase configures and enables the MIPI CSI-2 internal test pattern generator (PG) to generate a known image pattern, routes the generated data through the CSI-2 subsystem via a selected virtual channel and GDMA path, programs the DMA engine to transfer the received pixel data from the CSI-2 receive buffer to a destination memory address, waits for the DMA transfer completion interrupt, and then finishes. The test exercises the CSI-2 host controller pattern generator registers (PPI_PG_PATTERN_VRES, PPI_PG_PATTERN_HRES, PPI_PG_CONFIG, PPI_PG_ENABLE), the CSI-2 subsystem wrapper registers (virtual_channel, control_data), the D-PHY initialization and stop-state polling (PHY_STOPSTATE), the DMA higher-order address configuration registers (dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, dma_m0_addr_aw_ch0_Instruction), interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED), and the main interrupt status register (INT_ST_MAIN). The pattern generator is configured for a vertical resolution of 16 lines and horizontal resolution of 320 pixels with 24 valid bits per pixel. The DMA transfer size is computed based on these parameters with 8-byte alignment.","Test Description":"Configure and enable the MIPI CSI-2 internal test pattern generator to produce a known image pattern (320x16 pixels, 24 bits per pixel). Route the generated data through a selected virtual channel and GDMA path. Program the DMA engine to transfer received pixel data to destination memory. Enable all relevant CSI-2 and DMA interrupts. Initialize the D-PHY and wait for stop-state. Start the pattern generator, wait, then disable it. Poll for DMA transfer completion interrupt. Wait and finish the test.","Meta Test Steps / Procedure":"1. Global variable declarations:\n int data_rd, data_wr;\n int def_fail_cnt = 0, wr_fail_cnt = 0;\n int vcid_csi2_wrap_reg;\n int gdma_path;\n long int csi_ctrl_data;\n int gdma_int_rsts, gdma_ch0_rsts, gdma_ch1_rsts;\n extern int_pend;\n unsigned int tx_trnsfr_size;\n\n2. Function void csi2_ctrlr_pg_enable():\n 2a. write_reg(MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES, 0x10);\n 2b. write_reg(MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES, 0x70140);\n 2c. write_reg(MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG, 0xe401);\n 2d. write_reg(MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE, 1);\n\n3. Function int test_case():\n 3a-3ai. [Full test case procedure]","Test Steps / Procedure":"1. Set interrupt pending flag.\n2. Set virtual channel ID (vcid=3) and compute the unselected path value.\n3. Based on the selected GDMA path (GDMA0/1/2/3), compute the virtual channel wrapper register value and set the GDMA path index.\n4. Compute gdma_reg_base address from the GDMA path index.\n5. Write the computed virtual channel value to the virtual_channel register.\n6. Disable control data transfer by writing 0 to the control_data register.\n7. Enable CSI-2 and DMA interrupts by calling the interrupt enable routine, which reads INT_ST_MAIN to clear pending interrupts and writes interrupt mask values to INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED.\n8. Initialize the SNPS D-PHY.\n9. Poll PHY_STOPSTATE register until it reads 0x1000f (PHY in stop state).\n10. Set image parameters: hres=320, vres=16, valid_bits_per_pixel=24.\n11. Compute DMA transfer size with 8-byte alignment.\n12. Program DMA higher-order address registers.\n13. Enable fracdiv output to CSI-2 subsystem.\n14. Set DMA channel 0 program counter and instruction preload address.\n15. Configure DMA transfer instruction preload.\n16. Issue DMAGO command for channel 0.\n17. Enable the pattern generator.\n18. Wait 100 cycles.\n19. Disable the pattern generator.\n20. Poll DMA interrupt masked status register until bit 0 is set.\n21. Wait 10000 cycles.\n22. Call finish(0) to end the test.","Meta Impacted Registers":"PPI_PG_PATTERN_VRES, PPI_PG_PATTERN_HRES, PPI_PG_CONFIG, PPI_PG_ENABLE, virtual_channel, control_data, PHY_STOPSTATE, dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, dma_m0_addr_aw_ch0_Instruction, INT_ST_MAIN, INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED","Impacted Registers":"PPI_PG_PATTERN_VRES, PPI_PG_PATTERN_HRES, PPI_PG_CONFIG, PPI_PG_ENABLE, virtual_channel, control_data, PHY_STOPSTATE, dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, dma_m0_addr_aw_ch0_Instruction, INT_ST_MAIN, INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED","Meta Validation / Acceptance Criteria":"1. PHY stop-state polling: read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) must return 0x1000f before proceeding.\n2. DMA transfer completion polling: read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) must have bit 0 set.\n3. Test completes by calling finish(0), indicating successful completion with return code 0.","Validation / Acceptance Criteria":"1. The PHY_STOPSTATE register must read 0x1000f indicating the D-PHY has entered stop state on all lanes before proceeding with DMA and pattern generator configuration.\n2. The DMA interrupt masked status register bit 0 must be set, indicating the DMA transfer has completed successfully.\n3. The test finishes with finish(0) indicating successful test completion.","Remarks":"The testcase uses conditional compilation (#ifdef GDMA3_PATH, GDMA2_PATH, GDMA1_PATH, else GDMA0_PATH) to select which GDMA path to use. An additional conditional compilation (#ifdef FPS60) controls whether the DMA destination address increment flag is set. The function csi2_enable_interrupt() is defined in the source but the test_case() function calls csi2_subsys_enable_interrupt() which is not defined in this source file. The macro MIZAR_MIPI_CSI2_RB_REG_BASE is used with an offset (+0xf4) for fracdiv output enable but has unresolved register mapping in Agent 4 output. The extern declaration 'extern int_pend' is missing the type specifier 'int' in the source."}]

# TestPlan columns
testplan_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers",
    "Validation / Acceptance Criteria", "Code Generation"
]

# MetaData columns
metadata_columns = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

# Create workbook
wb = Workbook()

# Rename default sheet to TestPlan
ws_tp = wb.active
ws_tp.title = "TestPlan"

# Create MetaData sheet
ws_md = wb.create_sheet("MetaData")

# Formatting styles
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_alignment = Alignment(wrap_text=True, vertical="top")

def populate_sheet(ws, columns, data, column_key_map=None):
    # Write headers
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment

    # Write data rows
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            key = col_name
            if column_key_map and col_name in column_key_map:
                key = column_key_map[col_name]
            value = row_data.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment

    # Freeze first row
    ws.freeze_panes = "A2"

    # Auto-size columns
    max_width = 60
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(data) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                lines = str(cell_value).split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
        adjusted_width = min(max_len + 2, max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

# Populate TestPlan sheet
populate_sheet(ws_tp, testplan_columns, json_data)

# Populate MetaData sheet
populate_sheet(ws_md, metadata_columns, json_data)

# Set MetaData sheet to veryHidden
ws_md.sheet_state = "veryHidden"

# Save
output_path = filename
wb.save(output_path)

# Validate
assert os.path.exists(output_path), "File does not exist"
assert os.path.getsize(output_path) > 0, "File is empty"
wb_check = load_workbook(output_path)
assert "TestPlan" in wb_check.sheetnames, "TestPlan sheet missing"
assert "MetaData" in wb_check.sheetnames, "MetaData sheet missing"

print(f"SUCCESS: {output_path}")
print(f"Size: {os.path.getsize(output_path)} bytes")
print(f"Sheets: {wb_check.sheetnames}")
print(f"TestPlan rows: {ws_tp.max_row - 1}")
print(f"MetaData rows: {ws_md.max_row - 1}")
