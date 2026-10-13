import json
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

IP_NAME = 'USB'
WORKBOOK_FILENAME = f'{IP_NAME}_TestPlan.xlsx'

_json_str = r"""
[
 {
 "Index": "1",
 "SS / Module": "USB",
 "Test Case Name": "usb_host_enumeration_hs",
 "Feature": "USB Host Enumeration - High Speed",
 "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n\"usb.h\"",
 "Meta Macros": "MIZAR_USB_BASE, MIZAR_USB_CONFIG",
 "Meta Arrays": "int data_in[512];\nint data_out[512];",
 "Speed": "High-Speed (HS)",
 "Mode": "Host",
 "Memory Start Offset": "NA",
 "Memory End Offset": "NA",
 "Meta Test Description": "USB host enumeration test for high-speed (HS) device.",
 "Test Description": "Verify USB host enumeration of a high-speed device.",
 "Meta Test Steps / Procedure": "1. extern int int_pend;\n2. int test_case() {...}",
 "Test Steps / Procedure": "1. Call NIC programming initialization.\n2. Enable all IRQs via GIC.\n3. Configure event ring.\n4. Start host controller.\n5. Perform enumeration.\n6. Call finish(0).",
 "Meta Impacted Registers": "CONFIG, DCBAAP_LO, DCBAAP_HI",
 "Impacted Registers": "CONFIG, DCBAAP_LO, DCBAAP_HI",
 "Meta Validation / Acceptance Criteria": "1. Port connect status high.\n2. finish(0) signals pass.",
 "Validation / Acceptance Criteria": "1. Port status change interrupt received.\n2. finish(0) indicates pass.",
 "Remarks": "High-speed enumeration test."
 },
 {
 "Index": "2",
 "SS / Module": "USB",
 "Test Case Name": "USB_FS_Device_Bulk_Transfer_test",
 "Feature": "USB Full-Speed Device Bulk Transfer",
 "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n\"usb.h\"",
 "Meta Macros": "MIZAR_USB_DEPCMD, MIZAR_USB_GCTL",
 "Meta Arrays": "int buf_data[16];",
 "Speed": "Full-Speed (FS)",
 "Mode": "Device",
 "Memory Start Offset": "NA",
 "Memory End Offset": "NA",
 "Meta Test Description": "USB full-speed device bulk transfer test.",
 "Test Description": "Verify USB full-speed device bulk data transfer.",
 "Meta Test Steps / Procedure": "1. extern int int_pend;\n2. int test_case() {...}",
 "Test Steps / Procedure": "1. Call NIC programming initialization.\n2. Enable all IRQs.\n3. Configure event buffer.\n4. Issue bulk transfer.\n5. Call finish(0).",
 "Meta Impacted Registers": "GCTL, GUCTL, DEVTEN",
 "Impacted Registers": "GCTL, GUCTL, DEVTEN",
 "Meta Validation / Acceptance Criteria": "1. DEPCMD commands complete.\n2. finish(0) signals pass.",
 "Validation / Acceptance Criteria": "1. All DEPCMD commands complete.\n2. finish(0) indicates pass.",
 "Remarks": "FS Device Bulk Transfer test."
 }
]
"""

json_data = json.loads(_json_str)

testplan_headers = []
metadata_headers = []
if len(json_data) > 0:
    all_headers = list(json_data[0].keys())
    testplan_headers = [h for h in all_headers if not h.startswith('Meta ')]
    metadata_headers = list(all_headers)

wb = Workbook()

header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
cell_alignment = Alignment(vertical='top', wrap_text=True)

def apply_header_format(ws, headers):
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

def auto_size_columns(ws, min_width=10, max_width=60):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_length = 0
        for cell in col:
            if cell.value:
                lines = str(cell.value).split('\n')
                max_line = max(len(line) for line in lines)
                max_length = max(max_length, max_line)
        adjusted = min(max(max_length + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted

def populate_worksheet(ws, headers, rows):
    apply_header_format(ws, headers)
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header, '')
            if isinstance(value, list):
                value = '\n'.join(str(item) for item in value)
            elif isinstance(value, dict):
                value = json.dumps(value, ensure_ascii=False)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_alignment
    ws.freeze_panes = 'A2'
    auto_size_columns(ws)

ws_testplan = wb.active
ws_testplan.title = 'TestPlan'
if len(testplan_headers) > 0:
    populate_worksheet(ws_testplan, testplan_headers, json_data)
else:
    ws_testplan.freeze_panes = 'A2'

ws_metadata = wb.create_sheet(title='MetaData')
if len(metadata_headers) > 0:
    populate_worksheet(ws_metadata, metadata_headers, json_data)
else:
    ws_metadata.freeze_panes = 'A2'
ws_metadata.sheet_state = 'veryHidden'

if os.path.exists(WORKBOOK_FILENAME):
    os.remove(WORKBOOK_FILENAME)

wb.save(WORKBOOK_FILENAME)

if os.path.exists(WORKBOOK_FILENAME):
    fsize = os.path.getsize(WORKBOOK_FILENAME)
    print(f'SUCCESS: {WORKBOOK_FILENAME} generated, size={fsize} bytes, xlsx={WORKBOOK_FILENAME.endswith(".xlsx")}')
else:
    print(f'ERROR: Failed to generate {WORKBOOK_FILENAME}')
