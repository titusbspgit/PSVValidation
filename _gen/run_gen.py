#!/usr/bin/env python3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os, base64

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "TestPlan"
ws.cell(row=1, column=1, value="Test")
wb.save("/tmp/test.xlsx")
print("OK")
