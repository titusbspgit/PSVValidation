import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# === Configuration ===
IP_NAME = "USB"
WORKBOOK_FILENAME = f"{IP_NAME}_TestPlan.xlsx"

# === JSON Data (preserved exactly as provided, loaded via json.loads) ===
_json_str = r"""
[
 {
 "Index": "1",
 "SS / Module": "USB",
 "Test Case Name": "usb_host_enumeration_hs",
 "Feature": "USB Host Enumeration - High Speed",
 "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n\"usb.h\"",
 "Meta Macros": "MIZAR_USB_BASE, MIZAR_USB_CONFIG, MIZAR_USB_DCBAAP_LO, MIZAR_USB_DCBAAP_HI, MIZAR_USB_ERSTSZ, MIZAR_USB_USBSTS, MIZAR_USB_IMAN, MIZAR_USB_ERDP_LO, MIZAR_USB_ERDP_HI, MIZAR_USB_ERSTBA_LO, MIZAR_USB_ERSTBA_HI, MIZAR_USB_IMOD, MIZAR_USB_USBCMD, MIZAR_USB_HCSPARAMS1, MIZAR_USB_SUPTPRT2_DW2, MIZAR_USB_SUPTPRT3_DW2, MIZAR_USB_PORTSC_20, MIZAR_USB_DBOFF, MIZAR_LSS_SYSREG_INTR_EN0, MIZAR_LSS_SYSREG_MSK_STS0, MIZAR_LSS_SYSREG_RAW_STCR0, USB_PORTSC_20_WCE, USB_PORTSC_20_WDE, USB_PORTSC_20_WOE, Device_Context_Base_Address_Array, Default_Event_Ring_Array, Event_Ring_Segment_Table, EP0_TR_Dequeue_Pointer, Default_Command_Ring, Default_Input_Context, MIZAR_USB_DB",
 "Meta Arrays": "int data_in[512];\nint data_out[512];",
 "Speed": "High-Speed (HS)",
 "Mode": "Host",
 "Memory Start Offset": "NA",
 "Memory End Offset": "NA",
 "Meta Test Description": "USB host enumeration test for high-speed (HS) device.",
 "Test Description": "Verify USB host enumeration of a high-speed device. The test initializes the USB host controller, configures the event ring and device context base address array, enables interrupts, starts the controller, detects device connection, performs port reset, assigns a device address via Enable Slot and Address Device commands, and enumerates the device by issuing GET_DESCRIPTOR control transfers on Endpoint 0.",
 "Meta Test Steps / Procedure": "1. extern int int_pend;\n2. int test_case() {...}",
 "Test Steps / Procedure": "1. Call NIC programming initialization.\n2. Enable all IRQs via GIC.\n3. Read and configure PHY control register.\n4. Read host controller structural parameters.\n5. Configure event ring and device context.\n6. Start host controller.\n7. Wait for port status change interrupt.\n8. Perform port reset.\n9. Call set_address().\n10. Call enumeration().\n11. Call finish(0).",
 "Meta Impacted Registers": "CONFIG, DCBAAP_LO, DCBAAP_HI, ERSTSZ, USBSTS, IMAN, ERDP_LO, ERDP_HI, ERSTBA_LO, ERSTBA_HI, IMOD, USBCMD, HCSPARAMS1, SUPTPRT2_DW2, SUPTPRT3_DW2, PORTSC_20, DB",
 "Impacted Registers": "CONFIG, DCBAAP_LO, DCBAAP_HI, ERSTSZ, USBSTS, IMAN, ERDP_LO, ERDP_HI, ERSTBA_LO, ERSTBA_HI, IMOD, USBCMD, HCSPARAMS1, SUPTPRT2_DW2, SUPTPRT3_DW2, PORTSC_20, DB",
 "Meta Validation / Acceptance Criteria": "1. Port connect status high after interrupt.\n2. Port reset completed.\n3. Event ring completion valid.\n4. finish(0) signals pass.",
 "Validation / Acceptance Criteria": "1. Port status change interrupt received.\n2. PORTSC_20 port connect status high.\n3. Port reset completion interrupt received.\n4. Event ring entries valid.\n5. GET_DESCRIPTOR completes.\n6. finish(0) indicates pass.",
 "Remarks": "High-speed enumeration test with set_address(). Default_IRQHandler uses && instead of &."
 },
 {
 "Index": "2",
 "SS / Module": "USB",
 "Test Case Name": "USB_FS_Device_Bulk_Transfer_test",
 "Feature": "USB Full-Speed Device Bulk Transfer",
 "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n\"usb.h\"",
 "Meta Macros": "MIZAR_USB_DEPCMD, MIZAR_USB_DEPCMDPAR0, MIZAR_USB_DEPCMDPAR1, MIZAR_USB_GCTL, MIZAR_USB_GUCTL, MIZAR_USB_DEVTEN, MIZAR_USB_GEVNTADRLO, MIZAR_USB_GEVNTADRHI, MIZAR_USB_GEVNTSIZ, MIZAR_USB_GEVNTCOUNT, MIZAR_USB_DCFG, MIZAR_USB_DSTS, MIZAR_USB_DCTL, MIZAR_USB_DALEPENA, MIZAR_USB_GUSB2PHYCFG, MIZAR_USB_GFLADJ, MIZAR_USB_BASE, MIZAR_LSS_SYSREG_MSK_STS0, MIZAR_LSS_SYSREG_RAW_STCR0, Buffer_PointerLO, Buffer_PointerLO_1, event_trb_addr, Default_Event_Ring_Array, DWORD",
 "Meta Arrays": "int buf_data[16];",
 "Speed": "Full-Speed (FS)",
 "Mode": "Device",
 "Memory Start Offset": "NA",
 "Memory End Offset": "NA",
 "Meta Test Description": "USB full-speed device bulk transfer test.",
 "Test Description": "Verify USB full-speed device bulk data transfer. The test initializes the USB device controller, configures global event buffer and control registers, performs device enumeration, handles GET_DESCRIPTOR requests, processes SET_ADDRESS, issues endpoint commands for bulk data transfer, and validates transfer completion through interrupt-driven event handling.",
 "Meta Test Steps / Procedure": "1. extern int int_pend;\n2. int event_counter;\n3. int test_case() {...}",
 "Test Steps / Procedure": "1. Call NIC programming initialization.\n2. Enable all IRQs via GIC.\n3. Clear buffer and event arrays.\n4. Configure global event buffer.\n5. Configure GCTL, GFLADJ, GUCTL.\n6. Configure PIPE control register.\n7. Begin enumeration.\n8. Handle GET_DESCRIPTOR.\n9. Process SET_ADDRESS.\n10. Issue bulk data transfer.\n11. Call finish(0).",
 "Meta Impacted Registers": "GCTL, GUCTL, DEVTEN, DSTS, DCTL, DALEPENA",
 "Impacted Registers": "GCTL, GUCTL, DEVTEN, DSTS, DCTL, DALEPENA",
 "Meta Validation / Acceptance Criteria": "1. DEPCMD commands complete.\n2. Event counter updated.\n3. int_pend loops resolved.\n4. Bulk transfer completes.\n5. finish(0) signals pass.",
 "Validation / Acceptance Criteria": "1. All DEPCMD commands complete.\n2. Event counter properly updated.\n3. All interrupt waits resolved.\n4. Bulk transfer with 64-byte length completes.\n5. finish(0) indicates pass.",
 "Remarks": "FS Device Bulk Transfer test. Uses wait_on(5) for polling. TRB control 0x813 for bulk. Default_IRQHandler uses && instead of &."
 }
]
"""

json_data = json.loads(_json_str)

# === Determine TestPlan and MetaData columns ===
testplan_headers = []
metadata_headers = []
if len(json_data) > 0:
 all_headers = list(json_data[0].keys())
 testplan_headers = [h for h in all_headers if not h.startswith("Meta ")]
 metadata_headers = list(all_headers)

# === Create Workbook ===
wb = Workbook()

# === Styles ===
header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
cell_alignment = Alignment(vertical='top', wrap_text=True)


def apply_header_format(ws, headers):
 for col_idx, header in enumerate(headers, 1):
     cell = ws.cell(row=1, column=col_idx, value=header)
     cell.font = header_font
     cell.fill = header_fill
     cell.alignment = header_alignment


def auto_size_columns(ws, min_width=10, max_width=60):
 for col in ws.columns:
     col_letter = get_column_letter(col[0].column)
     max_length = 0
     for cell in col:
         if cell.value:
             lines = str(cell.value).split('\n')
             max_line = max(len(line) for line in lines)
             max_length = max(max_length, max_line)
     adjusted = min(max(max_length + 2, min_width), max_width)
     ws.column_dimensions[col_letter].width = adjusted


def populate_worksheet(ws, headers, rows):
 apply_header_format(ws, headers)
 for row_idx, row_data in enumerate(rows, 2):
     for col_idx, header in enumerate(headers, 1):
         value = row_data.get(header, '')
         if isinstance(value, list):
             value = '\n'.join(str(item) for item in value)
         elif isinstance(value, dict):
             value = json.dumps(value, ensure_ascii=False)
         cell = ws.cell(row=row_idx, column=col_idx, value=value)
         cell.alignment = cell_alignment
 ws.freeze_panes = 'A2'
 auto_size_columns(ws)


# === TestPlan Worksheet ===
ws_testplan = wb.active
ws_testplan.title = 'TestPlan'
if len(testplan_headers) > 0:
 populate_worksheet(ws_testplan, testplan_headers, json_data)
else:
 ws_testplan.freeze_panes = 'A2'

# === MetaData Worksheet (veryHidden) ===
ws_metadata = wb.create_sheet(title='MetaData')
if len(metadata_headers) > 0:
 populate_worksheet(ws_metadata, metadata_headers, json_data)
else:
 ws_metadata.freeze_panes = 'A2'
ws_metadata.sheet_state = 'veryHidden'

# === Save Workbook (overwrite if exists) ===
if os.path.exists(WORKBOOK_FILENAME):
 os.remove(WORKBOOK_FILENAME)

wb.save(WORKBOOK_FILENAME)

# === Verify and Print ===
if os.path.exists(WORKBOOK_FILENAME):
 fsize = os.path.getsize(WORKBOOK_FILENAME)
 print(f'Workbook generated: {WORKBOOK_FILENAME}')
 print(f'File size: {fsize} bytes')
 print(f'Extension check: {WORKBOOK_FILENAME.endswith(".xlsx")}')
else:
 print(f'ERROR: Failed to generate {WORKBOOK_FILENAME}')
