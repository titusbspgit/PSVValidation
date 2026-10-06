#!/usr/bin/env python3
"""LPDDR4 TestPlan Excel Generator - Agent 7
Generates a real XLSX workbook with TestPlan and MetaData sheets.
"""
import os
import datetime
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# EMBEDDED JSON DATA - 4 LPDDR4 TEST CASES
# ============================================================
json_data = [
    {
        "Index": "1",
        "SS / Module": "lpddr4",
        "Test Case Name": "lpddr4_mem_access_test",
        "Feature": "Memory Access",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"lpddr4.h\"",
        "Meta Macros": "NA",
        "Meta Arrays": "unsigned long int exp_data_array[50]; /* declared but not statically initialized; unused in current source */",
        "Speed": "SG=3200 (default); SG=2667 (if SG2667 defined); SG=2133 (if SG2133 defined)",
        "Mode": "bus_width = 0 (Full Bus); DBI_EN = 0; DM_EN = 0",
        "Memory Start Offset": "port0_addr = 0 (default/MPS_DRAM); port1_addr = 0x11A0000000 (default/MPS_DRAM); port0_addr = 0 (APS_DRAM); port1_addr = 0x15A0000000 (APS_DRAM)",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase verifies LPDDR4 memory access at different data widths (8-bit, 16-bit, 32-bit, 64-bit) for both read and write operations. The test initializes the LPDDR4 controller and PHY, performs training, then calls check_mem_access() on port0_addr (conditionally, depending on APS_DRAM/MPS_DRAM/initiator defines) and on port1_addr + 0x100000. The check_mem_access() function generates a random 64-bit data value using rand(), writes it as a 64-bit value to the target address, then reads it back using 8-bit (8 reads), 16-bit (4 reads), 32-bit (2 reads), and 64-bit (1 read) accesses, comparing each read against the expected portion of the original data. Then it advances the address by 0x1000, generates a new random 64-bit data value, writes it using mixed-width write accesses (two 8-bit writes, one 16-bit write, one 32-bit write), reads back the full 64-bit value, and validates it matches the expected data. Errors are accumulated in err0 and the test concludes with finish(err0).",
        "Test Description": "Verify LPDDR4 memory access correctness across 8-bit, 16-bit, 32-bit, and 64-bit read and write widths. The test writes a known data pattern to memory and reads it back at each access width, validating data integrity. It also verifies mixed-width write operations by writing portions of data at different widths and reading back the full 64-bit value for comparison.",
        "Meta Test Steps / Procedure": "1. Entry: test_case() is called.\n2. Declare local variables: unsigned long int index, offset, exp_data_array[50], read_data.\n3. Call gpv_programming().\n4. Set err0 = 0.\n5. [If APS_DRAM defined]: Set port0_addr = 0, port1_addr = 0x15A0000000, ctl_base = 0x9EE03000, phy_base = 0x9F000000.\n6. [Else (default/MPS_DRAM)]: Set port0_addr = 0, port1_addr = 0x11A0000000, ctl_base = 0x11D003000, phy_base = 0x11D500000.\n7. Set bus_width = 0 (Full Bus).\n8. [If SG2667 defined]: Set SG = 2667.\n9. [Elif SG2133 defined]: Set SG = 2133.\n10. [Else]: Set SG = 3200.\n11. Set DBI_EN = 0.\n12. Set DM_EN = 0.\n13. Call lpddr4_training().\n14. [If NOT ((APS_DRAM && A53_INITIATOR) XOR (MPS_DRAM && AI_INITIATOR) XOR (MPS_DRAM && DSP_INITIATOR))]: Call check_mem_access(port0_addr).\n15. Call check_mem_access(port1_addr + 0x100000).\n16. Call finish(err0).\n\ncheck_mem_access(unsigned long int addr) internal flow:\n17. Declare local variables: unsigned int i; char *ptr8, data8; unsigned short int *ptr16, data16; unsigned int *ptr32, data32; unsigned long int *ptr64, data64, data.\n18. data = rand().\n19. data = data | ((unsigned long)rand() << 32). This produces a full 64-bit random value.\n20. Set ptr8 = &data.\n21. Set ptr16 = &data.\n22. Set ptr32 = &data.\n23. Set ptr64 = &data.\n24. printf(\"addr: 0x%lx, data: 0x%08x\\n\", addr, data).\n25. Call write_reg64(addr, data). Write the 64-bit random data to the target address.\n26. Loop i = 0 to 7 (8 iterations, 8-bit read verification):\n27. data8 = read_reg8(addr + i).\n28. If data8 != *(ptr8 + i): printf ERROR message with addr+i, data8, *(ptr8+i); err0++.\n29. End loop.\n30. Loop i = 0 to 3 (4 iterations, 16-bit read verification):\n31. data16 = read_reg16(addr + (i * 2)).\n32. If data16 != (ptr16 + i): printf ERROR message with addr+(i*2), data16, *(ptr16+i); err0++.\n33. End loop.\n34. Loop i = 0 to 1 (2 iterations, 32-bit read verification):\n35. data32 = read_reg(addr + (i * 4)).\n36. If data32 != (ptr32 + i): printf ERROR message with addr+(i*4), data32, *(ptr32+i); err0++.\n37. End loop.\n38. Call read_reg64(addr, &data64). Read 64-bit value.\n39. If data64 != *ptr64: printf ERROR message with addr, data64, *ptr64; err0++.\n40. addr = addr + 0x1000. Advance address by 0x1000.\n41. data = rand().\n42. data = data | ((unsigned long)rand() << 32). Generate new 64-bit random data.\n43. printf(\"addr: 0x%lx, data: 0x%016lx\\n\", addr, data).\n44. Call write_reg8(addr + 0, *(ptr8 + 0)). Write byte 0.\n45. Call write_reg8(addr + 1, *(ptr8 + 1)). Write byte 1.\n46. Call write_reg16(addr + 2, *(ptr16 + 1)). Write halfword at offset 2.\n47. Call write_reg(addr + 4, *(ptr32 + 1)). Write word at offset 4.\n48. Call read_reg64(addr, &data64). Read back full 64-bit value.\n49. If data64 != data: printf ERROR message with addr, data, data64.\n50. Return from check_mem_access().\n\nNote: Steps 17-50 execute once for each call to check_mem_access(). Step 14 calls check_mem_access(port0_addr) conditionally, and step 15 calls check_mem_access(port1_addr + 0x100000) unconditionally. Both calls execute the full internal flow (steps 17-50) independently with their respective address arguments.",
        "Test Steps / Procedure": "1. Call GPV programming initialization.\n2. Initialize error counter to zero.\n3. Configure port addresses, controller base, and PHY base addresses based on DRAM configuration (APS or MPS).\n4. Set bus width to Full Bus mode.\n5. Configure speed grade (3200, 2667, or 2133 based on configuration).\n6. Disable DBI and DM.\n7. Perform LPDDR4 training.\n8. Conditionally perform memory access check on port 0 address (based on DRAM and initiator configuration).\n9. Perform memory access check on port 1 address with offset 0x100000.\n10. For each memory access check: write a random 64-bit value to the target address, then read it back using 8-bit accesses (8 reads), 16-bit accesses (4 reads), 32-bit accesses (2 reads), and a single 64-bit read, validating each read against the expected data portion.\n11. Advance address by 0x1000, generate new random data, write it using mixed-width writes (two 8-bit, one 16-bit, one 32-bit), read back the full 64-bit value, and validate.\n12. Report test result based on accumulated error count.",
        "Meta Impacted Registers": "NA",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "1. 8-bit read verification: For i = 0 to 7, read_reg8(addr + i) must equal (ptr8 + i) where ptr8 points to the 64-bit data written via write_reg64. If mismatch, err0 is incremented and an error message is printed.\n2. 16-bit read verification: For i = 0 to 3, read_reg16(addr + (i*2)) must equal (ptr16 + i). If mismatch, err0 is incremented and an error message is printed.\n3. 32-bit read verification: For i = 0 to 1, read_reg(addr + (i*4)) must equal *(ptr32 + i). If mismatch, err0 is incremented and an error message is printed.\n4. 64-bit read verification: read_reg64(addr, &data64) must yield data64 == *ptr64. If mismatch, err0 is incremented and an error message is printed.\n5. Mixed-width write verification: After writing byte 0 via write_reg8(addr+0, *(ptr8+0)), byte 1 via write_reg8(addr+1, *(ptr8+1)), halfword via write_reg16(addr+2, (ptr16+1)), and word via write_reg(addr+4, (ptr32+1)), read_reg64(addr, &data64) must yield data64 == data (the full 64-bit random value). If mismatch, an error message is printed.\n6. Final pass/fail: finish(err0) is called. err0 == 0 indicates pass (no mismatches detected across all access widths and both address regions).",
        "Validation / Acceptance Criteria": "1. All 8-bit read-back values from the written memory location must match the corresponding bytes of the original 64-bit data.\n2. All 16-bit read-back values must match the corresponding halfwords of the original data.\n3. All 32-bit read-back values must match the corresponding words of the original data.\n4. The 64-bit read-back value must match the original written data.\n5. After performing mixed-width writes (8-bit, 16-bit, 32-bit) to compose a 64-bit value, the 64-bit read-back must match the expected composed data.\n6. The test passes with zero errors accumulated across all access-width verifications on all tested memory addresses.",
        "Remarks": "The testcase uses conditional compilation (#if defined) to select between APS_DRAM and MPS_DRAM configurations for port addresses and controller/PHY bases, and between speed grades SG2667, SG2133, and default 3200. The port0 memory access check is conditionally compiled based on DRAM type and initiator type combinations. The exp_data_array[50] is declared but not used in the current source. Agent 2, Agent 3, and Agent 4 provided no register-access tokens or mappings."
    },
    {
        "Index": "2",
        "SS / Module": "lpddr4",
        "Test Case Name": "lpddr4_mem_basic_wr_rd_test",
        "Feature": "Basic Memory Write/Read",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"lpddr4.h\"",
        "Meta Macros": "NA",
        "Meta Arrays": "unsigned long int exp_data_array[50]; / declared locally in test_case(), not statically initialized; populated at runtime via rand() in two separate phases: first 10 elements for port0/port1 conditional test, then 32 elements for port1 large-offset test /",
        "Speed": "SG=3200 (default); SG=2667 (if SG2667 defined); SG=2133 (if SG2133 defined)",
        "Mode": "bus_width = 0 (Full Bus); DBI_EN = 0; DM_EN = 0",
        "Memory Start Offset": "port0_addr = 0 (default/MPS_DRAM); port1_addr = 0x11A0000000 (default/MPS_DRAM); port0_addr = 0 (APS_DRAM); port1_addr = 0x15A0000000 (APS_DRAM)",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase verifies basic LPDDR4 memory write and read operations using 64-bit accesses. The test initializes the LPDDR4 controller and PHY, performs training, then executes two test phases. Phase 1 (conditional, excluded for certain APS_DRAM/MPS_DRAM and initiator combinations): generates 10 random 64-bit values, writes them to port0 at offset increments of 0x1000, reads them back from port0 and validates, then reads the same 10 locations from port1 and validates (testing port aliasing or mirroring). Phase 2 (unconditional): generates 32 random 64-bit values, writes them to port1 at offset increments of 0x20000000 (spanning a large address range), reads them back from port1 and validates. Errors are accumulated in err0 and the test concludes with finish(err0).",
        "Test Description": "Verify basic LPDDR4 memory write and read correctness using 64-bit accesses. The test writes random data patterns to memory via two ports at various address offsets and reads them back, validating data integrity. Phase 1 tests 10 locations on port 0 with 0x1000 spacing and cross-reads from port 1. Phase 2 tests 32 locations on port 1 with large address spacing (0x20000000) to exercise a wide address range.",
        "Meta Test Steps / Procedure": "1. Entry: test_case() is called.\n2. Declare local variables: unsigned long int index, offset, exp_data_array[50], read_data; unsigned int err0.\n3. Declare global variables: unsigned long int port0_addr, port1_addr.\n4. Call gpv_programming().\n5. Set err0 = 0.\n6. [If APS_DRAM defined]: Set port0_addr = 0, port1_addr = 0x15A0000000, ctl_base = 0x9EE03000, phy_base = 0x9F000000.\n7. [Else (default/MPS_DRAM)]: Set port0_addr = 0, port1_addr = 0x11A0000000, ctl_base = 0x11D003000, phy_base = 0x11D500000.\n8. Set bus_width = 0 (Full Bus).\n9. [If SG2667 defined]: Set SG = 2667.\n10. [Elif SG2133 defined]: Set SG = 2133.\n11. [Else]: Set SG = 3200.\n12. Set DBI_EN = 0.\n13. Set DM_EN = 0.\n14. Call lpddr4_training().\n15. Set offset = 0x1000.\n16. [If NOT ((defined(APS_DRAM) && defined(A53_INITIATOR)) XOR (defined(MPS_DRAM) && defined(AI_INITIATOR)) XOR (defined(MPS_DRAM) && defined(DSP_INITIATOR)))]: Begin Phase 1 (conditional block).\n17. Phase 1 - Port0 Write Loop: for index = 0; index < 10; index++:\n18. exp_data_array[index] = rand().\n19. exp_data_array[index] = exp_data_array[index] | ((unsigned long)rand() << 32). This produces a full 64-bit random value.\n20. Call write_reg64(port0_addr + ((unsigned long)index * offset), exp_data_array[index]). Write 64-bit value to port0_addr + index*0x1000.\n21. End Phase 1 Port0 Write Loop (10 iterations: index 0 through 9).\n22. Phase 1 - Port0 Read-Verify Loop: for index = 0; index < 10; index++:\n23. Call read_reg64(port0_addr + ((unsigned long)index * offset), &read_data). Read 64-bit value from port0_addr + index*0x1000.\n24. If read_data != exp_data_array[index]: printf(\"ERROR_0: port0 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data); err0++.\n25. End Phase 1 Port0 Read-Verify Loop (10 iterations: index 0 through 9).\n26. Phase 1 - Port1 Read-Verify Loop (cross-port read): for index = 0; index < 10; index++:\n27. Call read_reg64(port1_addr + ((unsigned long)index * offset), &read_data). Read 64-bit value from port1_addr + index*0x1000.\n28. If read_data != exp_data_array[index]: printf(\"ERROR_0: port1 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data); err0++.\n29. End Phase 1 Port1 Read-Verify Loop (10 iterations: index 0 through 9).\n30. End Phase 1 conditional block.\n31. Set offset = 0x20000000.\n32. Phase 2 - Port1 Write Loop: for index = 0; index < 32; index++:\n33. exp_data_array[index] = rand().\n34. exp_data_array[index] = exp_data_array[index] | ((unsigned long)rand() << 32). This produces a full 64-bit random value.\n35. Call write_reg64(port1_addr + ((unsigned long)index * offset), exp_data_array[index]). Write 64-bit value to port1_addr + index*0x20000000.\n36. End Phase 2 Port1 Write Loop (32 iterations: index 0 through 31).\n37. Phase 2 - Port1 Read-Verify Loop: for index = 0; index < 32; index++:\n38. Call read_reg64(port1_addr + ((unsigned long)index * offset), &read_data). Read 64-bit value from port1_addr + index*0x20000000.\n39. If read_data != exp_data_array[index]: printf(\"ERROR_1: port1 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data); err0++.\n40. End Phase 2 Port1 Read-Verify Loop (32 iterations: index 0 through 31).\n41. Call finish(err0).",
        "Test Steps / Procedure": "1. Call GPV programming initialization.\n2. Initialize error counter to zero.\n3. Configure port addresses, controller base, and PHY base addresses based on DRAM configuration (APS or MPS).\n4. Set bus width to Full Bus mode.\n5. Configure speed grade (3200, 2667, or 2133 based on configuration).\n6. Disable DBI and DM.\n7. Perform LPDDR4 training.\n8. Phase 1 (conditional): Generate 10 random 64-bit data values and write them to port 0 at 0x1000 address increments.\n9. Read back all 10 values from port 0 and validate each against the expected data.\n10. Read back all 10 values from port 1 at the same offsets and validate each against the expected data (cross-port verification).\n11. Phase 2 (unconditional): Generate 32 random 64-bit data values and write them to port 1 at 0x20000000 address increments, spanning a large address range.\n12. Read back all 32 values from port 1 and validate each against the expected data.\n13. Report test result based on accumulated error count.",
        "Meta Impacted Registers": "NA",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "1. Phase 1 - Port0 read-back verification: For index = 0 to 9, read_reg64(port0_addr + (index * 0x1000), &read_data) must yield read_data == exp_data_array[index]. If mismatch, printf(\"ERROR_0: port0 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data) is printed and err0 is incremented.\n2. Phase 1 - Port1 cross-port read-back verification: For index = 0 to 9, read_reg64(port1_addr + (index * 0x1000), &read_data) must yield read_data == exp_data_array[index]. If mismatch, printf(\"ERROR_0: port1 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data) is printed and err0 is incremented.\n3. Phase 2 - Port1 large-offset read-back verification: For index = 0 to 31, read_reg64(port1_addr + (index * 0x20000000), &read_data) must yield read_data == exp_data_array[index]. If mismatch, printf(\"ERROR_1: port1 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data) is printed and err0 is incremented.\n4. Final pass/fail: finish(err0) is called. err0 == 0 indicates pass (no mismatches detected across all phases and all indices).",
        "Validation / Acceptance Criteria": "1. All 10 read-back values from port 0 must match the corresponding written random data values.\n2. All 10 cross-port read-back values from port 1 (at the same offsets used for port 0 writes) must match the expected data.\n3. All 32 read-back values from port 1 at large address offsets must match the corresponding written random data values.\n4. The test passes with zero errors accumulated across all verification phases.",
        "Remarks": "The testcase uses conditional compilation (#if defined) to select between APS_DRAM and MPS_DRAM configurations for port addresses and controller/PHY bases, and between speed grades SG2667, SG2133, and default 3200. Phase 1 (port0 write and port0/port1 read verification with offset 0x1000) is conditionally compiled based on DRAM type and initiator type combinations. Phase 2 (port1 write/read with offset 0x20000000) is unconditional and exercises 32 locations spanning a large address range. The exp_data_array is declared with 50 elements but only the first 10 (Phase 1) and first 32 (Phase 2) are used. Agent 2, Agent 3, and Agent 4 provided no register-access tokens or mappings."
    },
    {
        "Index": "3",
        "SS / Module": "lpddr4",
        "Test Case Name": "lpddr4_mem_boundary_test",
        "Feature": "Memory Boundary",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"lpddr4.h\"",
        "Meta Macros": "NA",
        "Meta Arrays": "unsigned long int exp_data_array[150]; / declared locally in test_case(), not statically initialized; populated at runtime via rand() in two separate phases: Phase 1 populates elements 0 through N-1 where N is the number of iterations satisfying (index1*offset) < val_2g with offset = PORT0_OFF or 0x1000000 (default yields 128 iterations); Phase 2 populates elements 0 through M-1 where M is the number of iterations satisfying (index1*offset) < val_16g with offset = PORT1_OFF or 0x08000000 (default yields 128 iterations). Array size 150 accommodates the maximum iteration count. /",
        "Speed": "SG=3200 (default); SG=2667 (if SG2667 defined); SG=2133 (if SG2133 defined)",
        "Mode": "bus_width = 0 (Full Bus); DBI_EN = 0; DM_EN = 0",
        "Memory Start Offset": "port0_addr = 0 (default/MPS_DRAM); port1_addr = 0x11A0000000 (default/MPS_DRAM); port0_addr = 0 (APS_DRAM); port1_addr = 0x15A0000000 (APS_DRAM)",
        "Memory End Offset": "val_2g = 0x0000000080000000 (2GB boundary for port0 phase); val_16g = 0x0000000400000000 (16GB boundary for port1 phase)",
        "Meta Test Description": "This testcase verifies LPDDR4 memory boundary access by writing and reading 64-bit data across the full address range up to the 2GB boundary on port0 and up to the 16GB boundary on port1. The test initializes the LPDDR4 controller and PHY, performs training, then executes two test phases. Phase 1 (conditional, excluded for certain APS_DRAM/MPS_DRAM and initiator combinations): uses a configurable offset (PORT0_OFF or default 0x1000000) to iterate across the port0 address range up to val_2g (0x80000000). For each iteration, a random 64-bit value is generated, written to port0, then all values are read back from port0 and validated, and then all values are read back from port1 and validated (cross-port verification). Phase 2 (unconditional): uses a configurable offset (PORT1_OFF or default 0x08000000) to iterate across the port1 address range up to val_16g (0x400000000). For each iteration, a random 64-bit value is generated and written to port1, then all values are read back from port1 and validated. Errors are accumulated in err0 and the test concludes with finish(err0). With default offsets, each phase performs 128 write and 128 read iterations (Phase 1 also performs 128 cross-port reads).",
        "Test Description": "Verify LPDDR4 memory boundary access correctness by writing random 64-bit data patterns across the full address range up to the 2GB boundary on port 0 and up to the 16GB boundary on port 1, using configurable address offsets. The test validates data integrity by reading back all written values and comparing against expected data, including cross-port read verification on port 1 for data written to port 0.",
        "Meta Test Steps / Procedure": "1. Entry: test_case() is called.\n2. Declare local variables: unsigned long int index1, offset, exp_data_array[150], read_data; unsigned int err0; unsigned long val_2g; unsigned long val_16g.\n3. Set val_2g = 0x0000000080000000.\n4. Set val_16g = 0x0000000400000000.\n5. Call gpv_programming().\n6. Set err0 = 0.\n7. [If APS_DRAM defined]: Set port0_addr = 0, port1_addr = 0x15A0000000, ctl_base = 0x9EE03000, phy_base = 0x9F000000.\n8. [Else (default/MPS_DRAM)]: Set port0_addr = 0, port1_addr = 0x11A0000000, ctl_base = 0x11D003000, phy_base = 0x11D500000.\n9. Set bus_width = 0 (Full Bus).\n10. [If SG2667 defined]: Set SG = 2667.\n11. [Elif SG2133 defined]: Set SG = 2133.\n12. [Else]: Set SG = 3200.\n13. Set DBI_EN = 0.\n14. Set DM_EN = 0.\n15. Call lpddr4_training().\n16. [If PORT0_OFF defined]: Set offset = PORT0_OFF.\n17. [Else]: Set offset = 0x1000000.\n18. printf(\"port0 offset = 0x%x\\n\", offset).\n19. [If NOT ((defined(APS_DRAM) && defined(A53_INITIATOR)) XOR (defined(MPS_DRAM) && defined(AI_INITIATOR)) XOR (defined(MPS_DRAM) && defined(DSP_INITIATOR)))]: Begin Phase 1 (conditional block).\n20. Phase 1 - Port0 Write Loop: for index1 = 0; ((unsigned long)(index1 * offset)) < val_2g; index1++:\n21. exp_data_array[index1] = rand().\n22. exp_data_array[index1] = exp_data_array[index1] | ((unsigned long)rand() << 32). Produces a full 64-bit random value.\n23. Call write_reg64(port0_addr + ((unsigned long)(index1 * offset)), exp_data_array[index1]). Write 64-bit value to port0_addr + index1*offset.\n24. printf(\"port0_0 index1: %d write addr:0x%016lx\\n\", index1, port0_addr + ((unsigned long)(index1 * offset))).\n25. End Phase 1 Port0 Write Loop. (With default offset 0x1000000, loop runs 128 iterations: index1 = 0 through 127, since 127*0x1000000 = 0x7F000000 < 0x80000000 and 128*0x1000000 = 0x80000000 which is NOT < val_2g.)\n26. Phase 1 - Port0 Read-Verify Loop: for index1 = 0; ((unsigned long)(index1 * offset)) < val_2g; index1++:\n27. Call read_reg64(port0_addr + ((unsigned long)(index1 * offset)), &read_data). Read 64-bit value from port0_addr + index1*offset.\n28. printf(\"port0_0 index1: %d read addr:0x%016lx\\n\", index1, port0_addr + ((unsigned long)(index1 * offset))).\n29. If read_data != exp_data_array[index1]: printf(\"ERROR_0: port0 index1 = %d, exp_data = %lx, actual_data = %lx\\n\", index1, exp_data_array[index1], read_data); err0++.\n30. End Phase 1 Port0 Read-Verify Loop. (Same iteration count as write loop.)\n31. Phase 1 - Port1 Cross-Port Read-Verify Loop: for index1 = 0; ((unsigned long)(index1 * offset)) < val_2g; index1++:\n32. Call read_reg64(port1_addr + ((unsigned long)(index1 * offset)), &read_data). Read 64-bit value from port1_addr + index1*offset.\n33. printf(\"port1_0 index1: %d read addr:0x%016lx\\n\", index1, port1_addr + ((unsigned long)(index1 * offset))).\n34. If read_data != exp_data_array[index1]: printf(\"ERROR_0: port1 index1 = %d, exp_data = %lx, actual_data = %lx\\n\", index1, exp_data_array[index1], read_data); err0++.\n35. End Phase 1 Port1 Cross-Port Read-Verify Loop. (Same iteration count as write loop.)\n36. End Phase 1 conditional block.\n37. [If PORT1_OFF defined]: Set offset = PORT1_OFF.\n38. [Else]: Set offset = 0x08000000.\n39. printf(\"port1 offset = 0x%x\\n\", offset).\n40. Phase 2 - Port1 Write Loop: for index1 = 0; ((unsigned long)(index1 * offset)) < val_16g; index1++:\n41. exp_data_array[index1] = rand().\n42. exp_data_array[index1] = exp_data_array[index1] | ((unsigned long)rand() << 32). Produces a full 64-bit random value.\n43. Call write_reg64(port1_addr + ((unsigned long)(index1 * offset)), exp_data_array[index1]). Write 64-bit value to port1_addr + index1*offset.\n44. printf(\"port1_1 index1: %d write addr:0x%016lx\\n\", index1, port1_addr + ((unsigned long)(index1 * offset))).\n45. End Phase 2 Port1 Write Loop. (With default offset 0x08000000, loop runs 128 iterations: index1 = 0 through 127, since 127*0x08000000 = 0x3F8000000 < 0x400000000 and 128*0x08000000 = 0x400000000 which is NOT < val_16g.)\n46. Phase 2 - Port1 Read-Verify Loop: for index1 = 0; ((unsigned long)(index1 * offset)) < val_16g; index1++:\n47. Call read_reg64(port1_addr + ((unsigned long)(index1 * offset)), &read_data). Read 64-bit value from port1_addr + index1*offset.\n48. printf(\"port1_1 index1: %d read addr:0x%016lx\\n\", index1, port1_addr + ((unsigned long)(index1 * offset))).\n49. If read_data != exp_data_array[index1]: printf(\"ERROR_1: port1 index1 = %d, exp_data = %lx, actual_data = %lx\\n\", index1, exp_data_array[index1], read_data); err0++.\n50. End Phase 2 Port1 Read-Verify Loop. (Same iteration count as write loop.)\n51. Call finish(err0).",
        "Test Steps / Procedure": "1. Call GPV programming initialization.\n2. Initialize error counter to zero.\n3. Configure port addresses, controller base, and PHY base addresses based on DRAM configuration (APS or MPS).\n4. Set bus width to Full Bus mode.\n5. Configure speed grade (3200, 2667, or 2133 based on configuration).\n6. Disable DBI and DM.\n7. Perform LPDDR4 training.\n8. Set port 0 address offset (configurable via PORT0_OFF, default 0x1000000).\n9. Phase 1 (conditional): Generate random 64-bit data values and write them to port 0 at the configured offset increments, iterating across the address range up to the 2GB boundary.\n10. Read back all written values from port 0 and validate each against the expected data.\n11. Read back all values from port 1 at the same offsets and validate each against the expected data (cross-port verification).\n12. Set port 1 address offset (configurable via PORT1_OFF, default 0x08000000).\n13. Phase 2 (unconditional): Generate random 64-bit data values and write them to port 1 at the configured offset increments, iterating across the address range up to the 16GB boundary.\n14. Read back all written values from port 1 and validate each against the expected data.\n15. Report test result based on accumulated error count.",
        "Meta Impacted Registers": "NA",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "1. Phase 1 - Port0 read-back verification: For each index1 where (index1 * offset) < val_2g (0x80000000), read_reg64(port0_addr + (index1 * offset), &read_data) must yield read_data == exp_data_array[index1]. If mismatch, printf(\"ERROR_0: port0 index1 = %d, exp_data = %lx, actual_data = %lx\\n\", index1, exp_data_array[index1], read_data) is printed and err0 is incremented.\n2. Phase 1 - Port1 cross-port read-back verification: For each index1 where (index1 * offset) < val_2g (0x80000000), read_reg64(port1_addr + (index1 * offset), &read_data) must yield read_data == exp_data_array[index1]. If mismatch, printf(\"ERROR_0: port1 index1 = %d, exp_data = %lx, actual_data = %lx\\n\", index1, exp_data_array[index1], read_data) is printed and err0 is incremented.\n3. Phase 2 - Port1 read-back verification: For each index1 where (index1 * offset) < val_16g (0x400000000), read_reg64(port1_addr + (index1 * offset), &read_data) must yield read_data == exp_data_array[index1]. If mismatch, printf(\"ERROR_1: port1 index1 = %d, exp_data = %lx, actual_data = %lx\\n\", index1, exp_data_array[index1], read_data) is printed and err0 is incremented.\n4. Final pass/fail: finish(err0) is called. err0 == 0 indicates pass (no mismatches detected across all phases and all iterations).",
        "Validation / Acceptance Criteria": "1. All read-back values from port 0 across the 2GB address range must match the corresponding written random data values.\n2. All cross-port read-back values from port 1 (at the same offsets used for port 0 writes) must match the expected data across the 2GB range.\n3. All read-back values from port 1 across the 16GB address range must match the corresponding written random data values.\n4. The test passes with zero errors accumulated across all verification phases.",
        "Remarks": "The testcase uses conditional compilation (#if defined) to select between APS_DRAM and MPS_DRAM configurations for port addresses and controller/PHY bases, and between speed grades SG2667, SG2133, and default 3200. Phase 1 (port0 write and port0/port1 read verification up to 2GB boundary) is conditionally compiled based on DRAM type and initiator type combinations. Phase 2 (port1 write/read up to 16GB boundary) is unconditional. The port0 offset is configurable via PORT0_OFF (default 0x1000000) and the port1 offset is configurable via PORT1_OFF (default 0x08000000). With default offsets, each phase iterates 128 times. The exp_data_array is declared with 150 elements to accommodate up to 128 iterations per phase. The loop bounds are dynamic, controlled by val_2g (0x80000000) and val_16g (0x400000000) boundary values. Agent 2, Agent 3, and Agent 4 provided no register-access tokens or mappings."
    },
    {
        "Index": "4",
        "SS / Module": "lpddr4",
        "Test Case Name": "lpddr4_mem_quarter_data_bus_width_test",
        "Feature": "Quarter Data Bus Width",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"lpddr4.h\"",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "SG=3200 (default); SG=2667 (if SG2667 defined); SG=2133 (if SG2133 defined)",
        "Mode": "bus_width = 2 (Quarter Bus); DBI_EN = 0; DM_EN = 0",
        "Memory Start Offset": "port0_addr = 0 (default/MPS_DRAM); port1_addr = 0x11A0000000 (default/MPS_DRAM); port0_addr = 0 (APS_DRAM); port1_addr = 0x15A0000000 (APS_DRAM)",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase verifies LPDDR4 memory read access in quarter data bus width mode. The test initializes the LPDDR4 controller and PHY with bus_width = 2 (Quarter Bus), performs training, calls training_done(), then waits 1000 cycles via wait_on(1000). After the wait, it reads two 64-bit values from port0 at offsets 0x100 and 0x108 (conditionally, excluded for certain APS_DRAM/MPS_DRAM and initiator combinations) and validates them against expected values 0x3333333333333333 and 0x2222222222222222 respectively. It then unconditionally reads two 64-bit values from port1 at offsets 0x100 and 0x108 and validates them against the same expected values. The test does not perform any explicit writes; it reads pre-existing data expected to be present after training in quarter bus width mode. Errors are accumulated in err0 and the test concludes with finish(err0). The local variables index and offset are declared but unused in the current source.",
        "Test Description": "Verify LPDDR4 memory read access correctness in quarter data bus width mode. After initializing the controller and PHY in quarter bus width configuration, performing training, and waiting for stabilization, the test reads pre-existing 64-bit data from both port 0 and port 1 at specific offsets and validates the read values against known expected data patterns.",
        "Meta Test Steps / Procedure": "1. Entry: test_case() is called.\n2. Declare global variables: unsigned long int port0_addr, port1_addr.\n3. Declare local variables: unsigned long int index, offset, read_data0, read_data1, read_data2, read_data3; unsigned int err0.\n4. Call gpv_programming().\n5. Set err0 = 0.\n6. [If APS_DRAM defined]: Set port0_addr = 0, port1_addr = 0x15A0000000, ctl_base = 0x9EE03000, phy_base = 0x9F000000.\n7. [Else (default/MPS_DRAM)]: Set port0_addr = 0, port1_addr = 0x11A0000000, ctl_base = 0x11D003000, phy_base = 0x11D500000.\n8. Set bus_width = 2. (Quarter Bus: 0 -> Full Bus, 1 -> Half Bus, 2 -> Quarter bus.)\n9. [If SG2667 defined]: Set SG = 2667.\n10. [Elif SG2133 defined]: Set SG = 2133.\n11. [Else]: Set SG = 3200.\n12. Set DBI_EN = 0.\n13. Set DM_EN = 0.\n14. Call lpddr4_training().\n15. Call training_done().\n16. Call wait_on(1000). Wait for 1000 cycles.\n17. [If NOT ((defined(APS_DRAM) && defined(A53_INITIATOR)) XOR (defined(MPS_DRAM) && defined(AI_INITIATOR)) XOR (defined(MPS_DRAM) && defined(DSP_INITIATOR)))]: Begin conditional port0 read block.\n18. Call read_reg64(port0_addr + 0x100, &read_data0). Read 64-bit value from port0_addr + 0x100 into read_data0.\n19. Call read_reg64(port0_addr + 0x108, &read_data1). Read 64-bit value from port0_addr + 0x108 into read_data1.\n20. If (read_data0 != 0x3333333333333333) && (read_data1 != 0x2222222222222222): printf(\"ERROR: port0 read_data0 = %lx, read_data1 = %lx\\n\", read_data0, read_data1); err0++.\n21. End conditional port0 read block.\n22. Call read_reg64(port1_addr + 0x100, &read_data2). Read 64-bit value from port1_addr + 0x100 into read_data2.\n23. Call read_reg64(port1_addr + 0x108, &read_data3). Read 64-bit value from port1_addr + 0x108 into read_data3.\n24. If (read_data2 != 0x3333333333333333) && (read_data3 != 0x2222222222222222): printf(\"ERROR: port1 read_data2 = %lx, read_data3 = %lx\\n\", read_data2, read_data3); err0++.\n25. Call finish(err0).",
        "Test Steps / Procedure": "1. Call GPV programming initialization.\n2. Initialize error counter to zero.\n3. Configure port addresses, controller base, and PHY base addresses based on DRAM configuration (APS or MPS).\n4. Set bus width to Quarter Bus mode.\n5. Configure speed grade (3200, 2667, or 2133 based on configuration).\n6. Disable DBI and DM.\n7. Perform LPDDR4 training.\n8. Signal training completion.\n9. Wait for stabilization (1000 cycles).\n10. Conditionally read two 64-bit values from port 0 at offsets 0x100 and 0x108, and validate against expected data patterns.\n11. Unconditionally read two 64-bit values from port 1 at offsets 0x100 and 0x108, and validate against expected data patterns.\n12. Report test result based on accumulated error count.",
        "Meta Impacted Registers": "NA",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "1. Port0 conditional validation (step 20): read_data0 from read_reg64(port0_addr + 0x100, &read_data0) is checked against 0x3333333333333333, AND read_data1 from read_reg64(port0_addr + 0x108, &read_data1) is checked against 0x2222222222222222. The error condition is: if ((read_data0 != 0x3333333333333333) && (read_data1 != 0x2222222222222222)), then printf(\"ERROR: port0 read_data0 = %lx, read_data1 = %lx\\n\", read_data0, read_data1) is printed and err0 is incremented. Note: the condition uses logical AND (&&), meaning the error is reported only when BOTH values mismatch simultaneously.\n2. Port1 unconditional validation (step 24): read_data2 from read_reg64(port1_addr + 0x100, &read_data2) is checked against 0x3333333333333333, AND read_data3 from read_reg64(port1_addr + 0x108, &read_data3) is checked against 0x2222222222222222. The error condition is: if ((read_data2 != 0x3333333333333333) && (read_data3 != 0x2222222222222222)), then printf(\"ERROR: port1 read_data2 = %lx, read_data3 = %lx\\n\", read_data2, read_data3) is printed and err0 is incremented. Note: the condition uses logical AND (&&), meaning the error is reported only when BOTH values mismatch simultaneously.\n3. Final pass/fail: finish(err0) is called. err0 == 0 indicates pass (no mismatches detected).",
        "Validation / Acceptance Criteria": "1. The 64-bit value read from port 0 at offset 0x100 must equal 0x3333333333333333, or the value read from port 0 at offset 0x108 must equal 0x2222222222222222 (error is flagged only when both values mismatch simultaneously due to AND condition).\n2. The 64-bit value read from port 1 at offset 0x100 must equal 0x3333333333333333, or the value read from port 1 at offset 0x108 must equal 0x2222222222222222 (same AND condition logic).\n3. The test passes with zero errors accumulated across all read validations.",
        "Remarks": "This testcase operates in quarter data bus width mode (bus_width = 2), unlike the other LPDDR4 memory tests which use full bus width (bus_width = 0). The test calls training_done() and wait_on(1000) after lpddr4_training(), which are not present in the other LPDDR4 memory tests. The test performs only reads (no writes); it validates pre-existing data at specific memory offsets expected to be present after training in quarter bus width mode. The expected data patterns are 0x3333333333333333 at offset 0x100 and 0x2222222222222222 at offset 0x108. The validation uses logical AND (&&) rather than OR, meaning an error is only reported when both read values mismatch their expected values simultaneously. The local variables index and offset are declared but unused in the current source. Agent 2, Agent 3, and Agent 4 provided no register-access tokens or mappings."
    }
]

# ============================================================
# SHEET COLUMN DEFINITIONS
# ============================================================
TESTPLAN_COLUMNS = [
    "Index",
    "SS / Module",
    "Feature",
    "Test Case Name",
    "Test Description",
    "Speed",
    "Mode",
    "Memory Start Offset",
    "Memory End Offset",
    "Remarks",
    "Test Steps / Procedure",
    "Impacted Registers",
    "Validation / Acceptance Criteria",
    "Code Generation"
]

METADATA_COLUMNS = [
    "Index",
    "Test Case Name",
    "Meta Test Description",
    "Meta Test Steps / Procedure",
    "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria",
    "Meta Headers",
    "Meta Macros",
    "Meta Arrays"
]

# ============================================================
# FIELD MAPPINGS
# ============================================================
TESTPLAN_MAPPING = [
    "Index",
    "SS / Module",
    "Feature",
    "Test Case Name",
    "Test Description",
    "Speed",
    "Mode",
    "Memory Start Offset",
    "Memory End Offset",
    "Remarks",
    "Test Steps / Procedure",
    "Impacted Registers",
    "Validation / Acceptance Criteria"
]

METADATA_MAPPING = [
    "Index",
    "Test Case Name",
    "Meta Test Description",
    "Meta Test Steps / Procedure",
    "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria",
    "Meta Headers",
    "Meta Macros",
    "Meta Arrays"
]

def generate_workbook():
    """Generate the LPDDR4 TestPlan Excel workbook."""
    
    # Get IST timestamp
    import pytz
    try:
        ist = pytz.timezone('Asia/Kolkata')
        now_ist = datetime.datetime.now(ist)
    except Exception:
        # Fallback: UTC + 5:30
        now_utc = datetime.datetime.utcnow()
        now_ist = now_utc + datetime.timedelta(hours=5, minutes=30)
    
    timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
    filename = f"LPDDR4_TestPlan_{timestamp}.xlsx"
    output_dir = "Test_Output/LPDDR4/TestPlan"
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, filename)
    
    print(f"Generating workbook: {filepath}")
    print(f"IST Timestamp: {timestamp}")
    print(f"Number of test cases: {len(json_data)}")
    
    # Create workbook
    wb = Workbook()
    
    # ============================================================
    # TESTPLAN SHEET
    # ============================================================
    ws_tp = wb.active
    ws_tp.title = "TestPlan"
    
    # Header formatting
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    
    # Write TestPlan headers
    for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Write TestPlan data rows
    for row_idx, tc in enumerate(json_data, 2):
        for col_idx, field_key in enumerate(TESTPLAN_MAPPING, 1):
            value = tc.get(field_key, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_alignment
        # Code Generation column (14th) left blank
        cell = ws_tp.cell(row=row_idx, column=14, value="")
        cell.alignment = cell_alignment
    
    # Freeze first row
    ws_tp.freeze_panes = "A2"
    
    # Auto-size columns with max width cap
    for col_idx in range(1, len(TESTPLAN_COLUMNS) + 1):
        max_length = len(str(ws_tp.cell(row=1, column=col_idx).value or ""))
        for row_idx in range(2, len(json_data) + 2):
            cell_val = str(ws_tp.cell(row=row_idx, column=col_idx).value or "")
            # Use first line length for multi-line content
            first_line = cell_val.split('\n')[0] if '\n' in cell_val else cell_val
            max_length = max(max_length, min(len(first_line), 80))
        adjusted_width = min(max(max_length + 2, 12), 60)
        ws_tp.column_dimensions[get_column_letter(col_idx)].width = adjusted_width
    
    # ============================================================
    # METADATA SHEET
    # ============================================================
    ws_md = wb.create_sheet(title="MetaData")
    
    # Write MetaData headers
    for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Write MetaData data rows
    for row_idx, tc in enumerate(json_data, 2):
        for col_idx, field_key in enumerate(METADATA_MAPPING, 1):
            value = tc.get(field_key, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_alignment
    
    # Freeze first row
    ws_md.freeze_panes = "A2"
    
    # Auto-size columns with max width cap
    for col_idx in range(1, len(METADATA_COLUMNS) + 1):
        max_length = len(str(ws_md.cell(row=1, column=col_idx).value or ""))
        for row_idx in range(2, len(json_data) + 2):
            cell_val = str(ws_md.cell(row=row_idx, column=col_idx).value or "")
            first_line = cell_val.split('\n')[0] if '\n' in cell_val else cell_val
            max_length = max(max_length, min(len(first_line), 80))
        adjusted_width = min(max(max_length + 2, 12), 60)
        ws_md.column_dimensions[get_column_letter(col_idx)].width = adjusted_width
    
    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = 'veryHidden'
    
    # ============================================================
    # SAVE WORKBOOK
    # ============================================================
    wb.save(filepath)
    print(f"Workbook saved: {filepath}")
    
    # ============================================================
    # POST-SAVE VALIDATION
    # ============================================================
    # 1. File exists
    assert os.path.exists(filepath), f"File does not exist: {filepath}"
    
    # 2. File size > 0
    file_size = os.path.getsize(filepath)
    assert file_size > 0, f"File size is 0: {filepath}"
    print(f"File size: {file_size} bytes")
    
    # 3. Workbook opens successfully
    wb_verify = load_workbook(filepath)
    
    # 4. Contains TestPlan sheet
    assert "TestPlan" in wb_verify.sheetnames, "TestPlan sheet not found"
    
    # 5. Contains MetaData sheet
    assert "MetaData" in wb_verify.sheetnames, "MetaData sheet not found"
    
    # 6. Verify row counts
    ws_tp_v = wb_verify["TestPlan"]
    ws_md_v = wb_verify["MetaData"]
    tp_rows = ws_tp_v.max_row - 1  # Exclude header
    md_rows = ws_md_v.max_row - 1  # Exclude header
    assert tp_rows == len(json_data), f"TestPlan row count mismatch: {tp_rows} vs {len(json_data)}"
    assert md_rows == len(json_data), f"MetaData row count mismatch: {md_rows} vs {len(json_data)}"
    print(f"TestPlan rows: {tp_rows}, MetaData rows: {md_rows}")
    
    # 7. MetaData sheet is veryHidden
    assert ws_md_v.sheet_state == 'veryHidden', f"MetaData sheet state: {ws_md_v.sheet_state}"
    print("MetaData sheet state: veryHidden - VERIFIED")
    
    # ============================================================
    # META CONTENT VALIDATION
    # ============================================================
    meta_fields_to_validate = [
        "Meta Test Description",
        "Meta Test Steps / Procedure",
        "Meta Impacted Registers",
        "Meta Validation / Acceptance Criteria",
        "Meta Headers",
        "Meta Macros",
        "Meta Arrays"
    ]
    
    validation_passed = True
    for row_idx, tc in enumerate(json_data):
        excel_row = row_idx + 2  # 1-indexed, skip header
        for field in meta_fields_to_validate:
            col_idx = METADATA_MAPPING.index(field) + 1
            cell_value = ws_md_v.cell(row=excel_row, column=col_idx).value
            expected_value = tc.get(field, "")
            if cell_value != expected_value:
                print(f"VALIDATION FAILED: Row {row_idx+1}, Field '{field}'")
                print(f"  Expected length: {len(str(expected_value))}")
                print(f"  Got length: {len(str(cell_value))}")
                validation_passed = False
    
    if validation_passed:
        print("META CONTENT VALIDATION: PASSED")
    else:
        print("META CONTENT VALIDATION: FAILED")
        raise ValueError("Meta content validation failed")
    
    wb_verify.close()
    
    print(f"\n=== GENERATION COMPLETE ===")
    print(f"Filename: {filename}")
    print(f"Path: {filepath}")
    print(f"Size: {file_size} bytes")
    print(f"TestPlan rows: {tp_rows}")
    print(f"MetaData rows: {md_rows}")
    print(f"Validation: PASSED")
    
    return filepath, filename, tp_rows, md_rows

if __name__ == "__main__":
    filepath, filename, tp_rows, md_rows = generate_workbook()
    print(f"\nOUTPUT_FILE={filepath}")
    print(f"OUTPUT_FILENAME={filename}")
