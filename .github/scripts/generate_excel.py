import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone, timedelta
import os

ist = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(ist)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"Ethernet3_TestPlan_{timestamp}.xlsx"

output_dir = "Test_Output/Ethernet/Ethernet3/TestPlan"
os.makedirs(output_dir, exist_ok=True)
filepath = os.path.join(output_dir, filename)

row_data = {
    "Index": "1",
    "SS / Module": "Ethernet3",
    "Test Case Name": "ethernet3_reg_wr_rd_test",
    "Feature": "Register Write-Read Verification",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Test Description": "This test verifies the default (reset) values and write-read integrity of Ethernet3 MAC registers including MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_WD_JB_Timeout, MAC_Hash_Table_Reg0, and MAC_Hash_Table_Reg1. First, each register is read and its value is compared against the expected default value. Then, six distinct data patterns are written to each writable register and read back. The read-back value is validated against an expected value computed using the register's read mask, write mask, and default value. The test passes only if all default value checks and all write-read checks succeed for every register and every pattern.",
    "Test Steps / Procedure": "1. Read each target Ethernet3 MAC register and verify its value matches the expected default (reset) value.\n2. Skip registers that are not readable or are flagged to be skipped for default value checking.\n3. Record any default value mismatches.\n4. For each of six test data patterns, write the pattern to each writable target register.\n5. Skip registers that are not writable or are flagged to be skipped for write-read checking.\n6. Read back each register after writing.\n7. Compute the expected read-back value using the write mask, read mask, and default value for each register.\n8. Compare the actual read-back value against the computed expected value.\n9. Record any write-read mismatches.\n10. Report overall pass if no default value mismatches and no write-read mismatches occurred; otherwise report fail.",
    "Impacted Registers": "MAC_Configuration; MAC_Ext_Configuration; MAC_Packet_Filter; MAC_WD_JB_Timeout; MAC_Hash_Table_Reg0; MAC_Hash_Table_Reg1",
    "Validation / Acceptance Criteria": "1. All target registers must return their expected default values when read after reset.\n2. For each of six test data patterns (all-ones, alternating bits, inverse alternating bits, all-zeros, mixed pattern, upper-half pattern), each writable register must read back a value consistent with the written data as masked by the register's write mask and read mask combined with the default value for non-writable bits.\n3. The test passes only if zero default-value mismatches and zero write-read mismatches are recorded across all registers and all patterns.\n4. Any single mismatch causes the test to report failure.",
    "Remarks": "The soft_reset_chk() function is commented out in test_case() and is not executed. The CNT macro is defined as 434 but only 6 register addresses are populated in addr_array; remaining entries are implicitly zero-initialized. Six test data patterns are used for comprehensive bit-level coverage. Registers are skipped based on read mask, write mask, and skip arrays. The expected read-back value accounts for non-writable bits retaining their default values.",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_define.c\"; <test_common.h>; <ethernet3/ethernet3_def.h>; <ethernet3/ethernet3_offset.h>",
    "Meta Macros": "CNT",
    "Meta Arrays": "addr_array[434]; default_value_array[434]; read_mask_array[434]; write_mask_array[434]; skip_array[434]; skip_rst_array[434]; chk_val[6]",
    "Meta Test Description": "This testcase performs register default value verification and write-read verification for a set of Ethernet3 MAC registers. The test_case() function first calls chk_rst_val() which iterates over addr_array (containing mizar_ETHERNET3_MAC_CONFIGURATION, mizar_ETHERNET3_MAC_EXT_CONFIGURATION, mizar_ETHERNET3_MAC_PACKET_FILTER, mizar_ETHERNET3_MAC_WD_JB_TIMEOUT, mizar_ETHERNET3_MAC_HASH_TABLE_REG0, mizar_ETHERNET3_MAC_HASH_TABLE_REG1). For each register, it checks if read_mask_array[i] is 0x00000000 (skip if not readable) and if skip_rst_array[i] is 1 (skip if flagged). It then reads the register via read_reg(addr) and compares against default_value_array[i]. If mismatch, def_fail_cnt is incremented. Next, chk_rd_wr() is called which iterates over 6 test patterns in chk_val[6] = {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}. For each pattern, it writes data_wr to each register (skipping if skip_array[i]==1 or write_mask_array[i]==0x00000000), then reads back each register. The expected value is computed as exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])) where wr_n = (write_mask_array[i] ^ 0xffffffff). If data_rd != exp_val, wr_fail_cnt is incremented. Finally, if def_fail_cnt > 0 or wr_fail_cnt > 0, finish(1) is called (fail), otherwise finish(0) (pass).",
    "Meta Test Steps / Procedure": "1. test_case() calls chk_rst_val(). 2. chk_rst_val() iterates i from 0 to CNT-1. 3. For each i, addr = addr_array[i]. 4. If read_mask_array[i] == 0x00000000, skip (not readable). 5. If skip_rst_array[i] == 1, skip. 6. data_rd = read_reg(addr). 7. Compare data_rd with default_value_array[i]. 8. If mismatch, increment def_fail_cnt and print failure. 9. test_case() calls chk_rd_wr(). 10. chk_rd_wr() defines chk_val[6] = {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}. 11. Outer loop j from 0 to 5, data_wr = chk_val[j]. 12. Inner write loop: for each i from 0 to CNT-1, addr = addr_array[i]. 13. If skip_array[i] == 1, skip writing. 14. If write_mask_array[i] == 0x00000000, skip writing. 15. Else write_reg(addr, data_wr). 16. Inner read loop: for each i from 0 to CNT-1, addr = addr_array[i]. 17. If skip_array[i] == 1, skip reading. 18. If write_mask_array[i] == 0x00000000, skip reading. 19. If read_mask_array[i] == 0x00000000, skip reading. 20. Else data_rd = read_reg(addr). 21. Compute wr_n = write_mask_array[i] ^ 0xffffffff. 22. Compute exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])). 23. Compare data_rd with exp_val. 24. If mismatch, increment wr_fail_cnt and print failure. 25. After all loops, if def_fail_cnt > 0 or wr_fail_cnt > 0, call finish(1) (fail). 26. Else call finish(0) (pass).",
    "Meta Impacted Registers": "mizar_ETHERNET3_MAC_CONFIGURATION; mizar_ETHERNET3_MAC_EXT_CONFIGURATION; mizar_ETHERNET3_MAC_PACKET_FILTER; mizar_ETHERNET3_MAC_WD_JB_TIMEOUT; mizar_ETHERNET3_MAC_HASH_TABLE_REG0; mizar_ETHERNET3_MAC_HASH_TABLE_REG1",
    "Meta Validation / Acceptance Criteria": "In chk_rst_val(): For each register in addr_array where read_mask_array[i] != 0x00000000 and skip_rst_array[i] != 1, data_rd = read_reg(addr) must equal default_value_array[i]. Any mismatch increments def_fail_cnt. In chk_rd_wr(): For each of 6 patterns in chk_val[] = {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}, after write_reg(addr, data_wr), data_rd = read_reg(addr) must equal exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | ((write_mask_array[i] ^ 0xffffffff) & read_mask_array[i] & default_value_array[i])). Any mismatch increments wr_fail_cnt. Final pass condition: def_fail_cnt == 0 AND wr_fail_cnt == 0 results in finish(0). Otherwise finish(1)."
}

wb = openpyxl.Workbook()

# TestPlan Sheet
ws_tp = wb.active
ws_tp.title = "TestPlan"

tp_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

blue_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
white_font = Font(bold=True, color="FFFFFF")
wrap_align = Alignment(wrap_text=True, vertical="top")

for col_idx, col_name in enumerate(tp_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = white_font
    cell.fill = blue_fill
    cell.alignment = wrap_align

tp_row_values = [
    row_data.get("Index", ""),
    row_data.get("SS / Module", ""),
    row_data.get("Feature", ""),
    row_data.get("Test Case Name", ""),
    row_data.get("Test Description", ""),
    row_data.get("Speed", ""),
    row_data.get("Mode", ""),
    row_data.get("Memory Start Offset", ""),
    row_data.get("Memory End Offset", ""),
    row_data.get("Remarks", ""),
    row_data.get("Test Steps / Procedure", ""),
    row_data.get("Impacted Registers", ""),
    row_data.get("Validation / Acceptance Criteria", ""),
    ""
]

for col_idx, val in enumerate(tp_row_values, 1):
    cell = ws_tp.cell(row=2, column=col_idx, value=val)
    cell.alignment = wrap_align

ws_tp.freeze_panes = "A2"

for col_idx, col_name in enumerate(tp_columns, 1):
    max_len = len(col_name)
    for row in ws_tp.iter_rows(min_row=2, max_row=ws_tp.max_row, min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                for line in str(cell.value).split('\n'):
                    max_len = max(max_len, len(line))
    ws_tp.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 60)

# MetaData Sheet
ws_md = wb.create_sheet("MetaData")

md_columns = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

for col_idx, col_name in enumerate(md_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = white_font
    cell.fill = blue_fill
    cell.alignment = wrap_align

md_row_values = [
    row_data.get("Index", ""),
    row_data.get("Test Case Name", ""),
    row_data.get("Meta Test Description", ""),
    row_data.get("Meta Test Steps / Procedure", ""),
    row_data.get("Meta Impacted Registers", ""),
    row_data.get("Meta Validation / Acceptance Criteria", ""),
    row_data.get("Meta Headers", ""),
    row_data.get("Meta Macros", ""),
    row_data.get("Meta Arrays", "")
]

for col_idx, val in enumerate(md_row_values, 1):
    cell = ws_md.cell(row=2, column=col_idx, value=val)
    cell.alignment = wrap_align

ws_md.freeze_panes = "A2"

for col_idx, col_name in enumerate(md_columns, 1):
    max_len = len(col_name)
    for row in ws_md.iter_rows(min_row=2, max_row=ws_md.max_row, min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                for line in str(cell.value).split('\n'):
                    max_len = max(max_len, len(line))
    ws_md.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 60)

ws_md.sheet_state = "veryHidden"

wb.save(filepath)
print(f"Generated: {filepath}")
print(f"Filename: {filename}")
