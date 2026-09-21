#!/usr/bin/env python3
import json
import os
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime('%Y%m%d_%H%M%S')
filename = f'Ethernet1_TestPlan_{timestamp}.xlsx'

# JSON data
json_data = [
  {
    "Index": "1",
    "SS / Module": "Ethernet1",
    "Test Case Name": "ethernet1_reg_wr_rd_test",
    "Feature": "Register Write-Read Verification",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_define.c\"; <test_common.h>; <ethernet1/ethernet1_def.h>; <ethernet1/ethernet1_offset.h>",
    "Meta Macros": "SOFT_RST_REG_ADDRESS; SOFT_RST_REG_DATA; CNT",
    "Meta Arrays": "addr_array[434]; default_value_array[434]; read_mask_array[434]; write_mask_array[434]; skip_array[434]; skip_rst_array[434]; chk_val[6]",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase performs register default-value verification and write-read verification for Ethernet1 MAC registers. The test_case() function first calls chk_rst_val() to read each register address from addr_array[] and compare the read value against the corresponding default_value_array[] entry. It then calls chk_rd_wr() which iterates over six test patterns in chk_val[] (0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000). For each pattern, it writes the pattern to each register via write_reg(addr, data_wr), then reads back via read_reg(addr) and computes the expected value using: exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])), where wr_n = (write_mask_array[i] ^ 0xffffffff). Registers are skipped if their read_mask_array or write_mask_array entry is 0x00000000, or if skip_array or skip_rst_array entry is 1. The addr_array contains: mizar_ETHERNET1_MAC_CONFIGURATION, mizar_ETHERNET1_MAC_EXT_CONFIGURATION, mizar_ETHERNET1_MAC_PACKET_FILTER, mizar_ETHERNET1_MAC_WD_JB_TIMEOUT, mizar_ETHERNET1_MAC_HASH_TABLE_REG0, mizar_ETHERNET1_MAC_HASH_TABLE_REG1. Failure counters def_fail_cnt and wr_fail_cnt are incremented on mismatch. The test calls finish(1) if any failure count is greater than 0, otherwise finish(0). The soft_reset_chk() function is commented out and not executed.",
    "Test Description": "This test verifies the default reset values and write-read accessibility of Ethernet1 MAC registers. It first reads each register and compares the value against its expected default. Then it writes six distinct data patterns to each writable register and reads back the values, applying read and write masks to compute the expected result. The registers under test are MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_WD_JB_Timeout, MAC_Hash_Table_Reg0, and MAC_Hash_Table_Reg1. The test passes only if all default value checks and all write-read verifications succeed across all patterns.",
    "Meta Test Steps / Procedure": "1. test_case() is the entry point. 2. chk_rst_val() is called: iterates i from 0 to CNT-1 over addr_array[]. 3. For each register, checks if read_mask_array[i] == 0x00000000; if so, skips (not readable). 4. Checks if skip_rst_array[i] == 1; if so, skips. 5. Calls data_rd = read_reg(addr) for the register address. 6. Compares data_rd against default_value_array[i]; increments def_fail_cnt on mismatch. 7. chk_rd_wr() is called: iterates j from 0 to 5 over chk_val[] patterns {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}. 8. For each pattern, sets data_wr = chk_val[j]. 9. Write loop: iterates i from 0 to CNT-1; skips if skip_array[i] == 1 or write_mask_array[i] == 0x00000000; otherwise calls write_reg(addr, data_wr). 10. Read loop: iterates i from 0 to CNT-1; skips if skip_array[i] == 1, write_mask_array[i] == 0x00000000, or read_mask_array[i] == 0x00000000. 11. Calls data_rd = read_reg(addr). 12. Computes wr_n = (write_mask_array[i] ^ 0xffffffff). 13. Computes exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])). 14. Compares data_rd against exp_val; increments wr_fail_cnt on mismatch. 15. After all iterations, checks if def_fail_cnt > 0 or wr_fail_cnt > 0; calls finish(1) on failure, finish(0) on success.",
    "Test Steps / Procedure": "1. Read each Ethernet1 MAC register and verify the value matches its expected default reset value.\n2. Skip registers that are not readable based on their read mask configuration.\n3. For each of six distinct test data patterns, write the pattern to each writable register.\n4. Skip registers that are not writable based on their write mask configuration.\n5. Read back each written register and compute the expected value by applying the read mask, write mask, and default value.\n6. Compare the read-back value against the computed expected value for each register and each pattern.\n7. Track any mismatches in default value checks and write-read checks separately.\n8. Report test pass if all default value and write-read verifications succeed; report test fail if any mismatch is detected.",
    "Meta Impacted Registers": "mizar_ETHERNET1_MAC_CONFIGURATION; mizar_ETHERNET1_MAC_EXT_CONFIGURATION; mizar_ETHERNET1_MAC_PACKET_FILTER; mizar_ETHERNET1_MAC_WD_JB_TIMEOUT; mizar_ETHERNET1_MAC_HASH_TABLE_REG0; mizar_ETHERNET1_MAC_HASH_TABLE_REG1",
    "Impacted Registers": "MAC_Configuration; MAC_Ext_Configuration; MAC_Packet_Filter; MAC_WD_JB_Timeout; MAC_Hash_Table_Reg0; MAC_Hash_Table_Reg1",
    "Meta Validation / Acceptance Criteria": "In chk_rst_val(): data_rd = read_reg(addr) must equal default_value_array[i] for each register; def_fail_cnt is incremented on mismatch. In chk_rd_wr(): for each of 6 patterns in chk_val[] {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}, data_rd = read_reg(addr) must equal exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])) where wr_n = (write_mask_array[i] ^ 0xffffffff); wr_fail_cnt is incremented on mismatch. Final pass condition: def_fail_cnt == 0 AND wr_fail_cnt == 0 results in finish(0); otherwise finish(1).",
    "Validation / Acceptance Criteria": "1. All registers must return their expected default reset values when read after reset.\n2. For each of six test data patterns written to each writable register, the read-back value must match the expected value computed using the applicable read and write masks and default values.\n3. The test passes (finish with success) only if zero default-value mismatches and zero write-read mismatches are detected across all registers and all patterns.\n4. Any single mismatch in either the default value check or the write-read check causes the test to fail.",
    "Remarks": "The soft_reset_chk() function is commented out in the source and is not executed during this test. The addr_array is declared with size 434 but only 6 register entries are populated in the source. Six write patterns are used for comprehensive bit-level coverage. Registers are conditionally skipped based on read mask, write mask, skip array, and skip reset array configurations."
  },
  {
    "Index": "2",
    "SS / Module": "Ethernet1",
    "Test Case Name": "ethernet1_tx_basic_test",
    "Feature": "Basic TX DMA Transmission",
    "Meta Headers": "<stdio.h>; <stdlib.h>; <ethernet1/ethernet1_def.h>; <ethernet1/ethernet1_offset.h>; <test_common.h>; <ethernet0/ethernet0_programming_sequence.h>",
    "Meta Macros": "NA",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "Interrupt Mode",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase performs a basic Ethernet1 TX DMA transmission. The test_case() function sets int_pend = 1, then configures MAC address registers for 4 addresses (Address0-Address3) by writing high and low portions. MAC_PACKET_FILTER is configured with 0x80000400 to enable hash filtering and receive-all mode. MAC_CONFIGURATION is set to 0x2003 enabling TX, RX, and full duplex mode. MTL TX queue operation modes for Q0-Q3 are configured. DMA_SYSBUS_MODE, TX descriptor ring lengths, list addresses, tail pointers, channel control, and interrupt enables are configured for all 4 DMA channels. Only DMA CH0 TX is started. The test waits for 10 interrupt-driven transfer completions.",
    "Test Description": "This test verifies basic Ethernet1 TX DMA packet transmission using interrupt-driven completion. It configures four MAC addresses with DMA channel selection and address enable, sets up the MAC packet filter for hash filtering and receive-all mode, enables TX and RX in full duplex mode via MAC_Configuration, and configures MAC_Ext_Configuration. The MTL layer is configured with store-and-forward mode and 4096-byte TX queue sizes for all four queues, along with quantum/weight values and queue interrupt enables. The DMA system bus mode is configured for burst selection and outstanding request limits. TX descriptors are preloaded, and DMA channels 0 through 3 are configured with descriptor ring lengths, list addresses, tail pointers, and channel control values. DMA channel interrupts are enabled for all four channels, and an external interrupt register is enabled. TX transmission is started on DMA channel 0. The test waits for 10 transfer-complete interrupts, where each interrupt handler reads current TX descriptor and buffer pointers for all channels, reads and clears the DMA interrupt status, clears individual DMA channel status registers, and clears the external interrupt. The test passes after all 10 interrupt cycles complete successfully.",
    "Meta Test Steps / Procedure": "1. Set int_pend = 1. 2. Write MAC address registers for Address0-Address3. 3. Configure MAC_PACKET_FILTER, MAC_CONFIGURATION, MAC_EXT_CONFIGURATION. 4. Configure MTL TXQ operation modes, quantum/weights, interrupt control for Q0-Q3. 5. Configure DMA_SYSBUS_MODE. 6. Preload descriptors. 7. Configure DMA TX descriptor ring lengths, list addresses, tail pointers, channel control, interrupt enables for CH0-CH3. 8. Enable HSS autoreg interrupt. 9. Start TX on CH0. 10. Loop 10 times waiting for interrupts. 11. IRQ handler reads TX desc/buffer pointers, reads/clears DMA interrupt status and channel status, clears HSS interrupt. 12. After 10 cycles, wait and finish(0).",
    "Test Steps / Procedure": "1. Configure four MAC addresses (Address0 through Address3) with address enable, DMA channel selection, and upper/lower address bytes.\n2. Configure the MAC packet filter to enable hash filtering and receive-all mode.\n3. Configure the MAC configuration register to enable TX and RX in full duplex mode.\n4. Configure the MAC extended configuration register.\n5. Configure MTL TX queue operation modes for all four queues with store-and-forward mode and 4096-byte queue sizes.\n6. Configure MTL TX queue quantum/weight values for all four queues.\n7. Enable RX queue overflow and TX queue underflow interrupts for all four MTL queues.\n8. Configure the DMA system bus mode with burst selection and outstanding request limits.\n9. Preload TX descriptors into memory.\n10. Set the TX descriptor ring length for all four DMA channels.\n11. Program the TX descriptor list base addresses for all four DMA channels.\n12. Program the TX descriptor tail pointers for all four DMA channels.\n13. Configure DMA channel control registers for all four channels.\n14. Enable DMA channel interrupts for all four channels.\n15. Enable the external HSS autoreg Ethernet interrupt.\n16. Start TX transmission on DMA channel 0.\n17. Wait for a transfer-complete interrupt; upon interrupt, read current TX descriptor and buffer pointers for all four channels.\n18. Read the DMA interrupt status register to identify active interrupts.\n19. Clear individual DMA channel status registers for all four channels by writing to clear.\n20. Read the DMA interrupt status register again to verify clearing.\n21. Clear the external HSS autoreg interrupt.\n22. Repeat the interrupt wait and handling cycle for a total of 10 iterations.\n23. Wait for a final settling period and report test pass.",
    "Meta Impacted Registers": "mizar_ETHERNET1_MAC_ADDRESS3_HIGH; mizar_ETHERNET1_MAC_ADDRESS2_HIGH; mizar_ETHERNET1_MAC_ADDRESS1_HIGH; mizar_ETHERNET1_MAC_ADDRESS0_HIGH; mizar_ETHERNET1_MAC_ADDRESS3_LOW; mizar_ETHERNET1_MAC_ADDRESS2_LOW; mizar_ETHERNET1_MAC_ADDRESS1_LOW; mizar_ETHERNET1_MAC_ADDRESS0_LOW; mizar_ETHERNET1_MAC_PACKET_FILTER; mizar_ETHERNET1_MAC_CONFIGURATION; mizar_ETHERNET1_MAC_EXT_CONFIGURATION; mizar_ETHERNET1_MTL_TXQ3_OPERATION_MODE; mizar_ETHERNET1_MTL_TXQ2_OPERATION_MODE; mizar_ETHERNET1_MTL_TXQ1_OPERATION_MODE; mizar_ETHERNET1_MTL_TXQ0_OPERATION_MODE; mizar_ETHERNET1_MTL_TXQ3_QUANTUM_WEIGHT; mizar_ETHERNET1_MTL_TXQ2_QUANTUM_WEIGHT; mizar_ETHERNET1_MTL_TXQ1_QUANTUM_WEIGHT; mizar_ETHERNET1_MTL_TXQ0_QUANTUM_WEIGHT; mizar_ETHERNET1_MTL_Q3_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET1_MTL_Q2_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET1_MTL_Q1_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET1_MTL_Q0_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET1_DMA_SYSBUS_MODE; mizar_ETHERNET1_DMA_CH3_TXDESC_RING_LENGTH; mizar_ETHERNET1_DMA_CH2_TXDESC_RING_LENGTH; mizar_ETHERNET1_DMA_CH1_TXDESC_RING_LENGTH; mizar_ETHERNET1_DMA_CH0_TXDESC_RING_LENGTH; mizar_ETHERNET1_DMA_CH3_TXDESC_LIST_ADDRESS; mizar_ETHERNET1_DMA_CH2_TXDESC_LIST_ADDRESS; mizar_ETHERNET1_DMA_CH1_TXDESC_LIST_ADDRESS; mizar_ETHERNET1_DMA_CH0_TXDESC_LIST_ADDRESS; mizar_ETHERNET1_DMA_CH3_TXDESC_TAIL_POINTER; mizar_ETHERNET1_DMA_CH2_TXDESC_TAIL_POINTER; mizar_ETHERNET1_DMA_CH1_TXDESC_TAIL_POINTER; mizar_ETHERNET1_DMA_CH0_TXDESC_TAIL_POINTER; mizar_ETHERNET1_DMA_CH3_CONTROL; mizar_ETHERNET1_DMA_CH2_CONTROL; mizar_ETHERNET1_DMA_CH1_CONTROL; mizar_ETHERNET1_DMA_CH0_CONTROL; mizar_ETHERNET1_DMA_CH1_INTERRUPT_ENABLE; mizar_ETHERNET1_DMA_CH2_INTERRUPT_ENABLE; mizar_ETHERNET1_DMA_CH3_INTERRUPT_ENABLE; mizar_ETHERNET1_DMA_CH0_INTERRUPT_ENABLE; 0XE68C2058; mizar_ETHERNET1_DMA_CH0_TX_CONTROL; mizar_ETHERNET1_DMA_CH0_CURRENT_APP_TXDESC; mizar_ETHERNET1_DMA_CH0_CURRENT_APP_TXBUFFER; mizar_ETHERNET1_DMA_CH1_CURRENT_APP_TXDESC; mizar_ETHERNET1_DMA_CH1_CURRENT_APP_TXBUFFER; mizar_ETHERNET1_DMA_CH2_CURRENT_APP_TXDESC; mizar_ETHERNET1_DMA_CH2_CURRENT_APP_TXBUFFER; mizar_ETHERNET1_DMA_CH3_CURRENT_APP_TXDESC; mizar_ETHERNET1_DMA_CH3_CURRENT_APP_TXBUFFER; mizar_ETHERNET1_DMA_INTERRUPT_STATUS; mizar_ETHERNET1_DMA_CH0_STATUS; mizar_ETHERNET1_DMA_CH1_STATUS; mizar_ETHERNET1_DMA_CH2_STATUS; mizar_ETHERNET1_DMA_CH3_STATUS; 0XE68C2050",
    "Impacted Registers": "MAC_Address3_High; MAC_Address2_High; MAC_Address1_High; MAC_Address0_High; MAC_Address3_Low; MAC_Address2_Low; MAC_Address1_Low; MAC_Address0_Low; MAC_Packet_Filter; MAC_Configuration; MAC_Ext_Configuration; MTL_TxQ3_Operation_Mode; MTL_TxQ2_Operation_Mode; MTL_TxQ1_Operation_Mode; MTL_TxQ0_Operation_Mode; MTL_TxQ3_Quantum_Weight; MTL_TxQ2_Quantum_Weight; MTL_TxQ1_Quantum_Weight; MTL_TxQ0_Quantum_Weight; MTL_Q3_Interrupt_Control_Status; MTL_Q2_Interrupt_Control_Status; MTL_Q1_Interrupt_Control_Status; MTL_Q0_Interrupt_Control_Status; DMA_SysBus_Mode; DMA_CH3_TxDesc_Ring_Length; DMA_CH2_TxDesc_Ring_Length; DMA_CH1_TxDesc_Ring_Length; DMA_CH0_TxDesc_Ring_Length; DMA_CH3_TxDesc_List_Address; DMA_CH2_TxDesc_List_Address; DMA_CH1_TxDesc_List_Address; DMA_CH0_TxDesc_List_Address; DMA_CH3_TxDesc_Tail_Pointer; DMA_CH2_TxDesc_Tail_Pointer; DMA_CH1_TxDesc_Tail_Pointer; DMA_CH0_TxDesc_Tail_Pointer; DMA_CH3_Control; DMA_CH2_Control; DMA_CH1_Control; DMA_CH0_Control; DMA_CH1_Interrupt_Enable; DMA_CH2_Interrupt_Enable; DMA_CH3_Interrupt_Enable; DMA_CH0_Interrupt_Enable; DMA_CH0_Tx_Control; DMA_CH0_Current_App_TxDesc; DMA_CH0_Current_App_TxBuffer; DMA_CH1_Current_App_TxDesc; DMA_CH1_Current_App_TxBuffer; DMA_CH2_Current_App_TxDesc; DMA_CH2_Current_App_TxBuffer; DMA_CH3_Current_App_TxDesc; DMA_CH3_Current_App_TxBuffer; DMA_Interrupt_Status; DMA_CH0_Status; DMA_CH1_Status; DMA_CH2_Status; DMA_CH3_Status",
    "Meta Validation / Acceptance Criteria": "The test uses an interrupt-driven validation model. The int_pend flag is set to 1 before starting TX. The while(int_pend) loop with wait_on(10) polls until Default_IRQHandler() fires and sets int_pend = 0. This cycle repeats 10 times. DMA channel status registers are cleared by writing 0xffffffff. After all 10 interrupt cycles, finish(0) indicates test pass.",
    "Validation / Acceptance Criteria": "1. The test must successfully receive 10 transfer-complete interrupts from DMA channel 0 TX operations.\n2. Each interrupt handler must successfully read the current application TX descriptor and TX buffer pointers for all four DMA channels to confirm DMA progress.\n3. The DMA interrupt status register must reflect active channel interrupts upon each interrupt entry.\n4. Writing to clear the individual DMA channel status registers for all four channels must successfully clear the pending interrupt status.\n5. The DMA interrupt status register must show cleared status after individual channel status clearing.\n6. The external HSS autoreg interrupt must be successfully cleared in each interrupt cycle.\n7. The test completes with a pass indication after all 10 interrupt cycles and a final settling wait period.",
    "Remarks": "Only DMA channel 0 TX transmission is started; channels 1, 2, and 3 TX control writes are commented out in the source. Two hardcoded hexadecimal register addresses are used for HSS autoreg interrupt enable and clear operations, which could not be mapped to canonical register names. The test relies on preload_descriptor() calls to set up TX descriptors in memory before DMA transmission. The extern int_pend variable is used as a shared flag between the main test loop and the interrupt handler for synchronization. The test does not perform explicit data integrity checks on transmitted packets; validation is based on successful interrupt-driven completion of 10 TX cycles."
  }
]

# TestPlan sheet columns
testplan_columns = [
    'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
    'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
    'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
    'Code Generation'
]

# MetaData sheet columns
metadata_columns = [
    'Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
    'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
    'Meta Headers', 'Meta Macros', 'Meta Arrays'
]

# Create workbook
wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = 'TestPlan'

# Header formatting
header_font = Font(bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap_alignment = Alignment(wrap_text=True, vertical='top')

# Write TestPlan headers
for col_idx, col_name in enumerate(testplan_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# Write TestPlan data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(testplan_columns, 1):
        value = row_data.get(col_name, '')
        if value is None:
            value = ''
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# Freeze first row
ws_tp.freeze_panes = 'A2'

# Auto-size columns with max width cap
for col_idx, col_name in enumerate(testplan_columns, 1):
    max_len = len(col_name)
    for row in range(2, len(json_data) + 2):
        val = ws_tp.cell(row=row, column=col_idx).value
        if val:
            lines = str(val).split('\n')
            for line in lines:
                max_len = max(max_len, len(line))
    adjusted_width = min(max_len + 2, 60)
    ws_tp.column_dimensions[ws_tp.cell(row=1, column=col_idx).column_letter].width = adjusted_width

# --- MetaData Sheet ---
ws_md = wb.create_sheet('MetaData')

# Write MetaData headers
for col_idx, col_name in enumerate(metadata_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# Write MetaData data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(metadata_columns, 1):
        value = row_data.get(col_name, '')
        if value is None:
            value = ''
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# Freeze first row
ws_md.freeze_panes = 'A2'

# Auto-size columns with max width cap
for col_idx, col_name in enumerate(metadata_columns, 1):
    max_len = len(col_name)
    for row in range(2, len(json_data) + 2):
        val = ws_md.cell(row=row, column=col_idx).value
        if val:
            lines = str(val).split('\n')
            for line in lines:
                max_len = max(max_len, len(line))
    adjusted_width = min(max_len + 2, 60)
    ws_md.column_dimensions[ws_md.cell(row=1, column=col_idx).column_letter].width = adjusted_width

# Set MetaData sheet to veryHidden
ws_md.sheet_state = 'veryHidden'

# Save
output_dir = 'Test_Output/Ethernet/Ethernet1/TestPlan'
os.makedirs(output_dir, exist_ok=True)
output_path = os.path.join(output_dir, filename)
wb.save(output_path)

# Validate
wb2 = load_workbook(output_path)
assert 'TestPlan' in wb2.sheetnames
assert 'MetaData' in wb2.sheetnames
assert os.path.getsize(output_path) > 0
print(f'SUCCESS: {filename}')
print(f'PATH: {output_path}')
print(f'SIZE: {os.path.getsize(output_path)}')
print(f'SHEETS: {wb2.sheetnames}')
print(f'TESTPLAN_ROWS: {ws_tp.max_row - 1}')
print(f'METADATA_ROWS: {ws_md.max_row - 1}')
