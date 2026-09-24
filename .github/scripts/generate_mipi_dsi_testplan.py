#!/usr/bin/env python3
"""MIPI_DSI TestPlan Excel Generator - Agent 7 Fallback Automation"""
import json
import os
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp_str = now_ist.strftime('%Y%m%d_%H%M%S')

IP_NAME = 'MIPI_DSI'
OUTPUT_DIR = 'Test_Output/MIPI/TestPlan'
FILENAME = f'{IP_NAME}_TestPlan_{timestamp_str}.xlsx'
OUTPUT_PATH = os.path.join(OUTPUT_DIR, FILENAME)

json_data = [
    {
        "Index": "1",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_basic_test",
        "Feature": "NA",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs basic MIPI DSI operations involving DMAC, subsystem, and host register accesses. It writes to DMAC interrupt enable (MIZAR_MIPI_DSI_DMAC_INTEN), subsystem interrupt enable (MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE), host PHY interface configuration (MIZAR_MIPI_DSI_HOST_PHY_IF_CFG), host packet handler configuration (MIZAR_MIPI_DSI_HOST_PCKHDL_CFG), host clock manager configuration (MIZAR_MIPI_DSI_HOST_CLKMGR_CFG), and subsystem DPI control (MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL). It programs the DMAC debug instruction registers (MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1) and executes them via the DMAC debug command register (MIZAR_MIPI_DSI_DMAC_DBGCMD). It reads the subsystem interrupt mask (MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK) and DMAC masked interrupt status (MIZAR_MIPI_DSI_DMAC_INTMIS). It clears DMAC interrupts via MIZAR_MIPI_DSI_DMAC_INTCLR and subsystem raw interrupts via MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW.",
        "Test Description": "This testcase performs basic MIPI DSI operations across the DMAC, subsystem, and host domains. It configures the PHY interface via PHY_IF_CFG, the packet handler via PCKHDL_CFG, and the clock manager via CLKMGR_CFG. It enables interrupts at both the DMAC and subsystem levels via the interrupt_enable register. It disables DPI output via the dpi_control register. It programs and executes DMAC debug instructions. It reads the interrupt_mask register and DMAC masked interrupt status to check interrupt state. It clears interrupts via the DMAC interrupt clear register and the subsystem interrupt_raw register.",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_DSI_DMAC_INTEN to enable DMAC interrupts. 2. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE to enable subsystem-level interrupts. 3. Write to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG to configure the PHY interface. 4. Write to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG to configure the packet handler. 5. Write to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG to configure the clock manager. 6. Write to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to configure DPI control. 7. Write to MIZAR_MIPI_DSI_DMAC_DBGINST0 to load the first debug instruction word. 8. Write to MIZAR_MIPI_DSI_DMAC_DBGINST1 to load the second debug instruction word. 9. Write to MIZAR_MIPI_DSI_DMAC_DBGCMD to execute the loaded debug instruction. 10. Read MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK to check subsystem interrupt mask status. 11. Read MIZAR_MIPI_DSI_DMAC_INTMIS to check DMAC masked interrupt status. 12. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear DMAC interrupts. 13. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW to clear subsystem raw interrupts.",
        "Test Steps / Procedure": "1. Enable DMAC-level interrupts by writing to the DMAC interrupt enable register. 2. Enable subsystem-level interrupts by writing to the interrupt_enable register. 3. Configure the PHY interface by writing to PHY_IF_CFG. 4. Configure the packet handler by writing to PCKHDL_CFG. 5. Configure the clock manager by writing to CLKMGR_CFG. 6. Configure DPI control by writing to the dpi_control register. 7. Load the first DMAC debug instruction word into the DMAC debug instruction 0 register. 8. Load the second DMAC debug instruction word into the DMAC debug instruction 1 register. 9. Execute the DMAC debug instruction by writing to the DMAC debug command register. 10. Read the interrupt_mask register to verify subsystem interrupt mask status. 11. Read the DMAC masked interrupt status register to verify DMAC interrupt state. 12. Clear DMAC interrupts by writing to the DMAC interrupt clear register. 13. Clear subsystem raw interrupts by writing to the interrupt_raw register.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE; MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW",
        "Impacted Registers": "interrupt_enable; PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control; interrupt_mask; interrupt_raw",
        "Meta Validation / Acceptance Criteria": "The subsystem interrupt mask register (MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK) is read to verify interrupt configuration. The DMAC masked interrupt status register (MIZAR_MIPI_DSI_DMAC_INTMIS) is read to verify DMAC interrupt state. DMAC interrupts are cleared via MIZAR_MIPI_DSI_DMAC_INTCLR and subsystem raw interrupts are cleared via MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW.",
        "Validation / Acceptance Criteria": "The interrupt_mask register is read to confirm subsystem interrupt configuration is correct. The DMAC masked interrupt status register is read to confirm DMAC interrupt state. DMAC interrupts are successfully cleared via the DMAC interrupt clear register. Subsystem raw interrupts are successfully cleared via the interrupt_raw register.",
        "Remarks": "Six DMAC register macros (DMAC interrupt enable, DMAC debug instruction 0, DMAC debug instruction 1, DMAC debug command, DMAC masked interrupt status, DMAC interrupt clear) could not be mapped to canonical register names because no DMAC register specification document was provided. Source code files in the testcase folder did not contain MIPI DSI related content; testcase details are derived from Agent 2, Agent 3, and Agent 4 outputs only."
    },
    {
        "Index": "2",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_dbi_random_payload_test",
        "Feature": "NA",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a DBI random payload test on the MIPI DSI interface. It writes to the host PHY interface configuration register (MIZAR_MIPI_DSI_HOST_PHY_IF_CFG), host packet handler configuration register (MIZAR_MIPI_DSI_HOST_PCKHDL_CFG), and host clock manager configuration register (MIZAR_MIPI_DSI_HOST_CLKMGR_CFG) to set up the DSI host. It writes to the subsystem DPI control register (MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL) to configure DPI output. It writes to the DMAC interrupt enable register (MIZAR_MIPI_DSI_DMAC_INTEN) to enable DMAC interrupts. It programs DMAC debug instruction registers (MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1) and executes them via the DMAC debug command register (MIZAR_MIPI_DSI_DMAC_DBGCMD). It polls the DMAC masked interrupt status register (MIZAR_MIPI_DSI_DMAC_INTMIS) to wait for DMA transfer completion. It clears DMAC interrupts by writing to the DMAC interrupt clear register (MIZAR_MIPI_DSI_DMAC_INTCLR).",
        "Test Description": "This testcase performs a DBI random payload test on the MIPI DSI interface. It configures the DSI host by writing to PHY_IF_CFG, PCKHDL_CFG, and CLKMGR_CFG registers. It configures DPI output via the dpi_control register. It enables DMAC interrupts, programs DMAC debug instructions, and executes them via the DMAC debug command register. It polls the DMAC masked interrupt status register to wait for DMA transfer completion, then clears DMAC interrupts via the DMAC interrupt clear register.",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG to configure the PHY interface. 2. Write to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG to configure the packet handler. 3. Write to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG to configure the clock manager. 4. Write to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to configure DPI control. 5. Write to MIZAR_MIPI_DSI_DMAC_INTEN to enable DMAC interrupts. 6. Write to MIZAR_MIPI_DSI_DMAC_DBGINST0 to load the first debug instruction word. 7. Write to MIZAR_MIPI_DSI_DMAC_DBGINST1 to load the second debug instruction word. 8. Write to MIZAR_MIPI_DSI_DMAC_DBGCMD to execute the loaded debug instruction. 9. Poll MIZAR_MIPI_DSI_DMAC_INTMIS to wait for DMAC transfer completion. 10. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear DMAC interrupts.",
        "Test Steps / Procedure": "1. Configure the PHY interface by writing to PHY_IF_CFG. 2. Configure the packet handler by writing to PCKHDL_CFG. 3. Configure the clock manager by writing to CLKMGR_CFG. 4. Configure DPI output by writing to the dpi_control register. 5. Enable DMAC interrupts by writing to the DMAC interrupt enable register. 6. Load the first DMAC debug instruction word into the DMAC debug instruction 0 register. 7. Load the second DMAC debug instruction word into the DMAC debug instruction 1 register. 8. Execute the DMAC debug instruction by writing to the DMAC debug command register. 9. Poll the DMAC masked interrupt status register until DMA transfer completion is indicated. 10. Clear DMAC interrupts by writing to the DMAC interrupt clear register.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR",
        "Impacted Registers": "PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control",
        "Meta Validation / Acceptance Criteria": "The DMAC masked interrupt status register (MIZAR_MIPI_DSI_DMAC_INTMIS) is polled until the expected completion status is indicated, confirming that the DMA transfer completed successfully. DMAC interrupts are then cleared by writing to MIZAR_MIPI_DSI_DMAC_INTCLR.",
        "Validation / Acceptance Criteria": "The DMAC masked interrupt status register is polled until the expected completion status is indicated, confirming that the DMA transfer completed successfully. DMAC interrupts are then cleared via the DMAC interrupt clear register.",
        "Remarks": "Six DMAC register macros (DMAC interrupt enable, DMAC debug instruction 0, DMAC debug instruction 1, DMAC debug command, DMAC masked interrupt status, DMAC interrupt clear) could not be mapped to canonical register names because no DMAC register specification document was provided. The source code file in the testcase folder contained unrelated PCIe code; testcase details are derived from Agent 2, Agent 3, and Agent 4 outputs only."
    }
]

# TestPlan sheet columns
testplan_columns = [
    'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
    'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
    'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
    'Code Generation'
]

# MetaData sheet columns
metadata_columns = [
    'Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
    'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
    'Meta Headers', 'Meta Macros', 'Meta Arrays'
]

# Create workbook
wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = 'TestPlan'

header_font = Font(bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap_alignment = Alignment(wrap_text=True, vertical='top')

# Write headers
for col_idx, col_name in enumerate(testplan_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# Write data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(testplan_columns, 1):
        value = row_data.get(col_name, '')
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# Freeze first row
ws_tp.freeze_panes = 'A2'

# Auto-size columns
for col_idx, col_name in enumerate(testplan_columns, 1):
    max_len = len(col_name)
    for row in range(2, len(json_data) + 2):
        val = ws_tp.cell(row=row, column=col_idx).value
        if val:
            max_len = max(max_len, min(len(str(val)), 80))
    ws_tp.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

# --- MetaData Sheet ---
ws_md = wb.create_sheet('MetaData')

# Write headers
for col_idx, col_name in enumerate(metadata_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# Write data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(metadata_columns, 1):
        value = row_data.get(col_name, '')
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# Freeze first row
ws_md.freeze_panes = 'A2'

# Auto-size columns
for col_idx, col_name in enumerate(metadata_columns, 1):
    max_len = len(col_name)
    for row in range(2, len(json_data) + 2):
        val = ws_md.cell(row=row, column=col_idx).value
        if val:
            max_len = max(max_len, min(len(str(val)), 80))
    ws_md.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

# Set MetaData sheet to veryHidden
ws_md.sheet_state = 'veryHidden'

# Save workbook
os.makedirs(OUTPUT_DIR, exist_ok=True)
wb.save(OUTPUT_PATH)

# Validate
wb2 = load_workbook(OUTPUT_PATH)
assert 'TestPlan' in wb2.sheetnames, 'TestPlan sheet missing'
assert 'MetaData' in wb2.sheetnames, 'MetaData sheet missing'
assert os.path.getsize(OUTPUT_PATH) > 0, 'File is empty'

print(f'SUCCESS: {OUTPUT_PATH}')
print(f'FILENAME: {FILENAME}')
print(f'ROWS_TESTPLAN: {len(json_data)}')
print(f'ROWS_METADATA: {len(json_data)}')
print(f'FILE_SIZE: {os.path.getsize(OUTPUT_PATH)}')
