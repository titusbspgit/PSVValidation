#!/usr/bin/env python3
"""
USB TestPlan XLSX Generator - Agent 7
Run: python3 gen_xlsx.py
Output: /tmp/USB_TestPlan_20261006_131326.xlsx
Then base64 encode: base64 /tmp/USB_TestPlan_20261006_131326.xlsx > /tmp/usb_testplan.b64
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import json, os, base64

FILENAME = "USB_TestPlan_20261006_131326.xlsx"
OUTPATH = os.path.join("/tmp", FILENAME)

# Load JSON
with open("/tmp/usb_testcases.json", "r") as f:
    json_data = json.load(f)

# Create workbook
wb = openpyxl.Workbook()

# === TestPlan Sheet ===
ws_tp = wb.active
ws_tp.title = "TestPlan"

tp_headers = ["Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
              "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
              "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
              "Code Generation"]
tp_keys = tp_headers[:-1]  # Code Generation has no JSON key

# === MetaData Sheet ===
ws_md = wb.create_sheet("MetaData")

md_headers = ["Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
              "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
              "Meta Headers", "Meta Macros", "Meta Arrays"]

# Formatting
hdr_font = Font(bold=True, color="FFFFFF", size=11)
hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap = Alignment(wrap_text=True, vertical="top")

# Write TestPlan headers
for ci, h in enumerate(tp_headers, 1):
    c = ws_tp.cell(row=1, column=ci, value=h)
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = wrap

# Write MetaData headers
for ci, h in enumerate(md_headers, 1):
    c = ws_md.cell(row=1, column=ci, value=h)
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = wrap

# Populate rows
for ri, item in enumerate(json_data, 2):
    for ci, key in enumerate(tp_keys, 1):
        c = ws_tp.cell(row=ri, column=ci, value=item.get(key, ""))
        c.alignment = wrap
    ws_tp.cell(row=ri, column=len(tp_headers), value="").alignment = wrap

    for ci, key in enumerate(md_headers, 1):
        c = ws_md.cell(row=ri, column=ci, value=item.get(key, ""))
        c.alignment = wrap

# Freeze panes
ws_tp.freeze_panes = "A2"
ws_md.freeze_panes = "A2"

# Auto-size columns
for ws in [ws_tp, ws_md]:
    for col_cells in ws.columns:
        mx = 0
        letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                mx = max(mx, min(len(str(cell.value)), 60))
        ws.column_dimensions[letter].width = max(mx + 2, 15)

# MetaData veryHidden
ws_md.sheet_state = "veryHidden"

# Save
wb.save(OUTPATH)
sz = os.path.getsize(OUTPATH)
print(f"SAVED: {OUTPATH} ({sz} bytes)")

# Validate
wb2 = openpyxl.load_workbook(OUTPATH)
assert "TestPlan" in wb2.sheetnames, "TestPlan missing"
assert "MetaData" in wb2.sheetnames, "MetaData missing"
assert wb2["TestPlan"].max_row == len(json_data) + 1
assert wb2["MetaData"].max_row == len(json_data) + 1
# Validate meta content
for ri, item in enumerate(json_data):
    for ci, key in enumerate(md_headers):
        cell_val = wb2["MetaData"].cell(row=ri+2, column=ci+1).value or ""
        expected = item.get(key, "")
        assert str(cell_val) == str(expected), f"Mismatch row {ri+2} col {key}"
print("VALIDATION: PASSED")
print(f"ROWS_TP: {len(json_data)}")
print(f"ROWS_MD: {len(json_data)}")

# Base64 encode for GitHub upload
with open(OUTPATH, "rb") as f:
    b64 = base64.b64encode(f.read()).decode()
with open("/tmp/usb_testplan.b64", "w") as f:
    f.write(b64)
print(f"B64 written to /tmp/usb_testplan.b64 ({len(b64)} chars)")
