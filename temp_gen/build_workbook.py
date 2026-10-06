#!/usr/bin/env python3
"""
Agent 7 - USB TestPlan Excel Workbook Generator
Generates USB_TestPlan_20261006_130430.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os, sys, json

FILENAME = "USB_TestPlan_20261006_130430.xlsx"
OUTPUT_DIR = "."

# ── JSON Data (all 4 test cases) ──
json_data = json.loads(open("usb_testcases.json").read())

# ── Create Workbook ──
wb = openpyxl.Workbook()

# ── TestPlan Sheet ──
ws_tp = wb.active
ws_tp.title = "TestPlan"

tp_headers = [
    "Index", "SS / Module", "Feature", "Test Case Name",
    "Test Description", "Speed", "Mode", "Memory Start Offset",
    "Memory End Offset", "Remarks", "Test Steps / Procedure",
    "Impacted Registers", "Validation / Acceptance Criteria", "Code Generation"
]

tp_keys = [
    "Index", "SS / Module", "Feature", "Test Case Name",
    "Test Description", "Speed", "Mode", "Memory Start Offset",
    "Memory End Offset", "Remarks", "Test Steps / Procedure",
    "Impacted Registers", "Validation / Acceptance Criteria"
]

# ── MetaData Sheet ──
ws_md = wb.create_sheet("MetaData")

md_headers = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

md_keys = md_headers[:]

# ── Write Headers ──
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_align = Alignment(wrap_text=True, vertical="top")

for col_idx, h in enumerate(tp_headers, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

for col_idx, h in enumerate(md_headers, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

# ── Populate Data ──
for row_idx, item in enumerate(json_data, 2):
    for col_idx, key in enumerate(tp_keys, 1):
        val = item.get(key, "")
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = wrap_align
    # Code Generation column blank
    ws_tp.cell(row=row_idx, column=len(tp_headers)).alignment = wrap_align

    for col_idx, key in enumerate(md_keys, 1):
        val = item.get(key, "")
        cell = ws_md.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = wrap_align

# ── Freeze Panes ──
ws_tp.freeze_panes = "A2"
ws_md.freeze_panes = "A2"

# ── Auto-size Columns ──
for ws in [ws_tp, ws_md]:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, min(len(str(cell.value)), 80))
        ws.column_dimensions[col_letter].width = max(max_len + 2, 12)

# ── MetaData veryHidden ──
ws_md.sheet_state = "veryHidden"

# ── Save ──
filepath = os.path.join(OUTPUT_DIR, FILENAME)
wb.save(filepath)
print(f"Saved: {filepath}")
print(f"Size: {os.path.getsize(filepath)} bytes")

# ── Validate ──
wb2 = openpyxl.load_workbook(filepath)
assert "TestPlan" in wb2.sheetnames
assert "MetaData" in wb2.sheetnames
assert wb2["TestPlan"].max_row == len(json_data) + 1
assert wb2["MetaData"].max_row == len(json_data) + 1
print("Validation PASSED")
print(f"TestPlan rows: {len(json_data)}")
print(f"MetaData rows: {len(json_data)}")
