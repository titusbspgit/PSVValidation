#!/usr/bin/env python3
"""
USB TestPlan XLSX Generator - Agent 7
Generates workbook with all 4 test cases, validates, and outputs base64.
Usage: python3 executor.py > /tmp/usb_b64_output.txt
"""
import openpyxl, json, os, sys, base64
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FILENAME = "USB_TestPlan_20261006_131326.xlsx"
OUTPATH = os.path.join("/tmp", FILENAME)

json_data = json.loads(r'''JSONPLACEHOLDER''')

wb = openpyxl.Workbook()
ws_tp = wb.active; ws_tp.title = "TestPlan"
ws_md = wb.create_sheet("MetaData")
tp_h = ["Index","SS / Module","Feature","Test Case Name","Test Description","Speed","Mode","Memory Start Offset","Memory End Offset","Remarks","Test Steps / Procedure","Impacted Registers","Validation / Acceptance Criteria","Code Generation"]
md_h = ["Index","Test Case Name","Meta Test Description","Meta Test Steps / Procedure","Meta Impacted Registers","Meta Validation / Acceptance Criteria","Meta Headers","Meta Macros","Meta Arrays"]
hf=Font(bold=True,color="FFFFFF",size=11); hfl=PatternFill(start_color="4472C4",end_color="4472C4",fill_type="solid"); w=Alignment(wrap_text=True,vertical="top")
for i,h in enumerate(tp_h,1): c=ws_tp.cell(1,i,h); c.font=hf; c.fill=hfl; c.alignment=w
for i,h in enumerate(md_h,1): c=ws_md.cell(1,i,h); c.font=hf; c.fill=hfl; c.alignment=w
for ri,item in enumerate(json_data,2):
    for ci,k in enumerate(tp_h[:-1],1): c=ws_tp.cell(ri,ci,item.get(k,"")); c.alignment=w
    ws_tp.cell(ri,len(tp_h),"").alignment=w
    for ci,k in enumerate(md_h,1): c=ws_md.cell(ri,ci,item.get(k,"")); c.alignment=w
ws_tp.freeze_panes="A2"; ws_md.freeze_panes="A2"
for ws in [ws_tp,ws_md]:
    for cc in ws.columns:
        mx=0; lt=get_column_letter(cc[0].column)
        for cl in cc:
            if cl.value: mx=max(mx,min(len(str(cl.value)),60))
        ws.column_dimensions[lt].width=max(mx+2,15)
ws_md.sheet_state="veryHidden"
wb.save(OUTPATH)
sz=os.path.getsize(OUTPATH)
wb2=openpyxl.load_workbook(OUTPATH)
assert "TestPlan" in wb2.sheetnames and "MetaData" in wb2.sheetnames
assert wb2["TestPlan"].max_row==len(json_data)+1 and wb2["MetaData"].max_row==len(json_data)+1
for ri,item in enumerate(json_data):
    for ci,k in enumerate(md_h):
        cv=wb2["MetaData"].cell(ri+2,ci+1).value or ""
        assert str(cv)==str(item.get(k,""))
with open(OUTPATH,"rb") as f: b64=base64.b64encode(f.read()).decode()
print(f"FILE:{OUTPATH}")
print(f"SIZE:{sz}")
print(f"ROWS_TP:{len(json_data)}")
print(f"ROWS_MD:{len(json_data)}")
print(f"VALIDATION:PASSED")
print(f"B64LEN:{len(b64)}")
print(f"B64DATA:{b64}")
