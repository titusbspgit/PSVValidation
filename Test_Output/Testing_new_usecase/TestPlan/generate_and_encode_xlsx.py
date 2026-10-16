#!/usr/bin/env python3
"""
Self-contained MIPI_CSI TestPlan XLSX generator.
Generates the workbook and prints base64-encoded content for upload.
Run: python3 generate_and_encode_xlsx.py
"""
import base64
import io
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("ERROR: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

IP_NAME = "MIPI_CSI"
MODULE_NAME = "MIPI CSI-2"
FEATURE = "Test Pattern Generator"
TESTCASE_NAME = "mipi_csi2_test_pattern_generator"

TEST_DESCRIPTION = (
    "Configures and validates the MIPI CSI-2 internal test pattern generator. "
    "The test initializes the CSI-2 subsystem with virtual channel configuration, "
    "D-PHY initialization, DMA setup, enables the pattern generator, and verifies "
    "DMA transfer completion. Validates end-to-end data path from CSI-2 pattern "
    "generator through DMA to memory."
)

TEST_STEPS = (
    "Step 1: Declare global variables (data_rd, data_wr, def_fail_cnt, wr_fail_cnt, vcid_csi2_wrap_reg, gdma_path, csi_ctrl_data, gdma_int_rsts, gdma_ch0_rsts, gdma_ch1_rsts, tx_trnsfr_size, gdma_reg_base)\n"
    "Step 2: [INIT] Set int_pend = 1, configure vcid = 3, compute vcid_unselected_path = ((vcid + 1) & 0xf)\n"
    "Step 3: [INIT] Select GDMA path based on compile-time defines (GDMA3_PATH, GDMA2_PATH, GDMA1_PATH, or default GDMA0_PATH) and compute vcid_csi2_wrap_reg accordingly\n"
    "Step 4: [INIT] Compute gdma_reg_base = 0xE6A00000 + (gdma_path * 0x1000)\n"
    "Step 5: [INIT] Write MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with vcid_csi2_wrap_reg\n"
    "Step 6: [INIT] Write MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA with 0\n"
    "Step 7: [INIT] Call csi2_subsys_enable_interrupt() to enable CSI-2 subsystem interrupts\n"
    "Step 8: [INIT] Call snps_phy_init() to initialize SNPS D-PHY\n"
    "Step 9: [RUN] Read MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE register\n"
    "Step 10: [RUN] Poll PHY_STOPSTATE until value equals 0x1000f (with 100000 iteration timeout)\n"
    "Step 11: [RUN] Set hres=320, vres=16, valid_bits_per_pixel=24\n"
    "Step 12: [RUN] Compute csi2_data_trnsfr_size with 8-byte alignment\n"
    "Step 13: [RUN] Write DMA higher-order address registers: AR_CH0_DATA=0x100, AR_CH0_INSTRUCTION=0x0, AW_CH0_DATA=0x0, AW_CH0_INSTRUCTION=0x0\n"
    "Step 14: [RUN] Enable fracdiv output to CSI-2 subsystem: write (MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4) = 0x1\n"
    "Step 15: [RUN] Set dma_ch0_pc = 0xE6000000, configure dma_dest_addr_incr_flag based on FPS60 define\n"
    "Step 16: [RUN] Call dma_trnsfr_instn_preload_incr_addr() to program DMA transfer\n"
    "Step 17: [RUN] Call DMAGO_CSI() to issue DMAGO command for channel 0\n"
    "Step 18: [RUN] Call csi2_ctrlr_pg_enable(): PG_PATTERN_VRES=0x10, PG_PATTERN_HRES=0x70140, PG_CONFIG=0xe401, PG_ENABLE=1\n"
    "Step 19: [RUN] Busy-wait delay (adapted from DV wait_on(100))\n"
    "Step 20: [RUN] Disable pattern generator: write PG_ENABLE = 0\n"
    "Step 21: [RUN] Read DMA interrupt masked status (gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET)\n"
    "Step 22: [RUN] Poll DMA INTMIS until bit 0 is set (with 100000 iteration timeout)\n"
    "Step 23: [RUN] Busy-wait delay (adapted from DV wait_on(10000))\n"
    "Step 24: [RUN] Set out->status = 0 (test pass)\n"
    "Step 25: [TEARDOWN] No additional cleanup required, return 0"
)

IMPACTED_REGISTERS = (
    "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES\n"
    "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES\n"
    "MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG\n"
    "MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE\n"
    "MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID\n"
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED\n"
    "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE\n"
    "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL\n"
    "MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA\n"
    "MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA\n"
    "MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION\n"
    "MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA\n"
    "MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION\n"
    "MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4\n"
    "gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET"
)

VALIDATION_CRITERIA = (
    "1. PHY_STOPSTATE register must reach value 0x1000f within 100000 polling iterations\n"
    "2. DMA transfer size correctly computed with 8-byte alignment for hres=320, vres=16, bpp=24\n"
    "3. Pattern generator successfully enabled (PG_ENABLE=1) and disabled (PG_ENABLE=0)\n"
    "4. DMA INTMIS bit 0 set within 100000 iterations indicating transfer completion\n"
    "5. All interrupt mask registers programmed with correct mask values\n"
    "6. Test returns status=0 (PASS) on successful completion\n"
    "7. Test returns status=-1 (FAIL) on PHY_STOPSTATE or DMA INTMIS timeout\n"
    "8. NULL output pointer detected and rejected with return -1"
)

META_HEADERS = (
    "#include <test.h>\n"
    "#include <stdio.h>\n"
    "#include <stdlib.h>\n"
    '#include "test_common.h"\n'
    '#include "mipi_csi2.h"\n'
    '#include "log.h"\n'
    '#include "reg_access.h"\n'
    '#include "mipi_csi2_test_pattern_generator.h"'
)

META_MACROS = (
    "#define GDMA_CSI2_DATA_DEST_ADDR2 0xE6040080\n"
    "#define GDMA_CTRL_DATA_DEST_ADDR2 0xE6040000"
)

META_ARRAYS = (
    "int data_rd, data_wr;\n"
    "int def_fail_cnt = 0, wr_fail_cnt = 0;\n"
    "int vcid_csi2_wrap_reg;\n"
    "int gdma_path;\n"
    "long int csi_ctrl_data;\n"
    "int gdma_int_rsts, gdma_ch0_rsts, gdma_ch1_rsts;\n"
    "extern int int_pend;\n"
    "unsigned int tx_trnsfr_size;\n"
    "static unsigned long int gdma_reg_base;"
)

REMARKS = (
    "MANUAL_REVIEW: 'extern int_pend;' missing type - preserved as 'extern int int_pend;'. "
    "gdma_reg_base declared as file-scope static. "
    "csi2_subsys_enable_interrupt() from external header. "
    "Conditional: GDMA3/2/1/0_PATH. FPS60 controls incr_flag. "
    "DV wait_on() adapted to busy-wait."
)

CODE_GEN = "See generated files: mipi_csi2_test_pattern_generator.c, mipi_csi2_test_pattern_generator.h"

wb = Workbook()

# -- TestPlan --
ws_tp = wb.active
ws_tp.title = "TestPlan"
tp_headers = [
    "Index","SS / Module","Feature","Test Case Name",
    "Test Description","Speed","Mode","Memory Start Offset",
    "Memory End Offset","Remarks","Test Steps / Procedure",
    "Impacted Registers","Validation / Acceptance Criteria","Code Generation"
]
ws_tp.append(tp_headers)
ws_tp.append([1, MODULE_NAME, FEATURE, TESTCASE_NAME, TEST_DESCRIPTION,
              "", "", "", "", REMARKS, TEST_STEPS, IMPACTED_REGISTERS,
              VALIDATION_CRITERIA, CODE_GEN])

# -- MetaData --
ws_md = wb.create_sheet("MetaData")
md_headers = [
    "Index","Test Case Name","Meta Test Description",
    "Meta Test Steps / Procedure","Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria","Meta Headers",
    "Meta Macros","Meta Arrays"
]
ws_md.append(md_headers)
ws_md.append([1, TESTCASE_NAME, TEST_DESCRIPTION, TEST_STEPS,
              IMPACTED_REGISTERS, VALIDATION_CRITERIA, META_HEADERS,
              META_MACROS, META_ARRAYS])

# -- Formatting --
hf = Font(bold=True, color="FFFFFF", size=11)
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
ca = Alignment(vertical="top", wrap_text=True)
tb = Border(left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"))

for ws in [ws_tp, ws_md]:
    for cell in ws[1]:
        cell.font = hf; cell.fill = hfill; cell.alignment = ha; cell.border = tb
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = ca; cell.border = tb
    ws.freeze_panes = "A2"
    for col_cells in ws.columns:
        mx = 0; cl = col_cells[0].column_letter
        for c in col_cells:
            if c.value:
                for ln in str(c.value).split("\\n"):
                    mx = max(mx, len(ln))
        ws.column_dimensions[cl].width = max(min(mx + 4, 60), 12)

ws_md.sheet_state = "veryHidden"

# Save to buffer and print base64
buf = io.BytesIO()
wb.save(buf)
b64 = base64.b64encode(buf.getvalue()).decode("ascii")
print(b64)
