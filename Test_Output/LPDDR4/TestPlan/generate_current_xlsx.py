#!/usr/bin/env python3
"""
Agent 7 - LPDDR4 TestPlan XLSX Generator
Generates LPDDR4_TestPlan_YYYYMMDD_HHMMSS.xlsx using openpyxl
Run: python3 generate_current_xlsx.py
Requires: pip install openpyxl
Output: Prints base64-encoded XLSX content to stdout for GitHub upload
"""
import json
import os
import sys
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta

# IST timezone (GMT+05:30)
ist = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(ist)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
FILENAME = f"LPDDR4_TestPlan_{timestamp}.xlsx"

TESTPLAN_COLS = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

METADATA_COLS = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

JSON_DATA = [
    {
        "Index": "1",
        "SS / Module": "LPDDR4",
        "Test Case Name": "lpddr4_mem_dm_test",
        "Feature": "LPDDR4 Data Mask (DM)",
        "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include \"test_common.h\"\n#include \"lpddr4.h\"",
        "Meta Macros": "APS_DRAM (conditional compilation macro, controls port address and base address selection)\nSG2667 (conditional compilation macro, selects speed grade 2667)\nSG2133 (conditional compilation macro, selects speed grade 2133)\nA53_INITIATOR (conditional compilation macro, used in XOR-based port0 access guard)\nMPS_DRAM (conditional compilation macro, used in XOR-based port0 access guard)\nAI_INITIATOR (conditional compilation macro, used in XOR-based port0 access guard)\nDSP_INITIATOR (conditional compilation macro, used in XOR-based port0 access guard)",
        "Meta Arrays": "unsigned long int exp_data_array[50]; \u2014 Array of 50 elements of type unsigned long int. Used to store expected 64-bit random data values generated via rand() combined with ((unsigned long)rand() << 32). Populated in write loops for port0 (10 elements, offset 0x1000) and used for read-back comparison against read_data.",
        "Speed": "Configurable via preprocessor: SG2667 \u2192 SG = 2667, SG2133 \u2192 SG = 2133, default \u2192 SG = 3200",
        "Mode": "bus_width = 0 (Full Bus)",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "LPDDR4 Data Mask (DM) test. Function signature: int test_case(). Global variable declarations: unsigned long int port0_addr; unsigned long int port1_addr;. Local variable declarations: unsigned long int index; unsigned long int offset; unsigned long int exp_data_array[50]; unsigned long int read_data; unsigned int err0;. The test initializes GPV programming via gpv_programming(), sets err0 = 0, configures port addresses (port0_addr, port1_addr), controller base (ctl_base), and PHY base (phy_base) conditionally based on APS_DRAM define. Sets bus_width = 0 (Full Bus). Selects speed grade (SG) via conditional compilation: 2667 if SG2667 defined, 2133 if SG2133 defined, else 3200. Disables Data Bus Inversion (DBI_EN = 0). Enables Data Mask (DM_EN = 1). Calls lpddr4_training() to perform LPDDR4 memory training. After training, sets offset = 0x1000 and conditionally (guarded by XOR-based preprocessor condition on APS_DRAM/A53_INITIATOR, MPS_DRAM/AI_INITIATOR, MPS_DRAM/DSP_INITIATOR) writes 10 random 64-bit values to port0_addr using write_reg64(), then reads them back via read_reg64() comparing against exp_data_array[index]. Also performs port1 read-back verification loop of 10 iterations comparing read_data against exp_data_array[index]. Errors are tracked via err0 counter and reported via printf.",
        "Test Description": "This test validates the LPDDR4 Data Mask (DM) functionality. The test configures the memory controller and PHY base addresses, sets full bus width mode, selects the speed grade, disables Data Bus Inversion, enables Data Mask, performs LPDDR4 memory training, then writes random 64-bit data patterns to memory via port0 and reads them back via both port0 and port1 to verify data integrity with Data Mask enabled.",
        "Meta Test Steps / Procedure": "1. Call gpv_programming()\n2. err0 = 0\n3. Conditional on APS_DRAM:\n #if defined(APS_DRAM)\n port0_addr = 0\n port1_addr = 0x15A0000000\n ctl_base = 0x9EE03000\n phy_base = 0x9F000000\n #else\n port0_addr = 0\n port1_addr = 0x11A0000000\n ctl_base = 0x11D003000\n phy_base = 0x11D500000\n #endif\n4. bus_width = 0 //0 -> Full Bus, 1 -> Half Bus, 2 -> Quarter bus\n5. Conditional on SG2667/SG2133:\n #if defined(SG2667)\n SG = 2667\n #elif defined(SG2133)\n SG = 2133\n #else\n SG = 3200\n #endif\n6. DBI_EN = 0\n7. DM_EN = 1\n8. Call lpddr4_training()\n9. offset = 0x1000\n10. #if !((defined(APS_DRAM) && defined(A53_INITIATOR)) ^(defined(MPS_DRAM) && defined(AI_INITIATOR)) ^(defined(MPS_DRAM) && defined(DSP_INITIATOR)))\n Port0 Write Loop:\n for(index = 0; index < 10; index++)\n {\n exp_data_array[index] = rand();\n exp_data_array[index] = exp_data_array[index] | ((unsigned long)rand() << 32);\n write_reg64(port0_addr + ((unsigned long)index * offset), exp_data_array[index]);\n }\n11. Port0 Read-back Verification Loop:\n for(index = 0; index < 10; index++)\n {\n read_reg64(port0_addr + ((unsigned long)indexoffset), &read_data);\n if(read_data != exp_data_array[index])\n {\n printf(\"ERROR_0: port0 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data);\n err0++;\n }\n }\n #endif\n12. Port1 Read-back Verification Loop:\n for(index = 0; index < 10; index++)\n {\n read_reg64(port1_addr + ((unsigned long)index * offset), &read_data);\n if(read_data != exp_data_array[index])\n {\n printf(\"ERROR_0: port1 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data);\n err0++;\n }\n }",
        "Test Steps / Procedure": "1. Initialize GPV programming.\n2. Clear error counter (err0 = 0).\n3. Configure port addresses, controller base, and PHY base addresses (configuration varies based on APS_DRAM selection).\n4. Set bus width to Full Bus mode.\n5. Select speed grade (2667, 2133, or 3200 based on compile-time configuration).\n6. Disable Data Bus Inversion (DBI).\n7. Enable Data Mask (DM).\n8. Execute LPDDR4 memory training.\n9. Set memory access offset to 0x1000.\n10. Conditionally write 10 random 64-bit data values to port0 addresses using write_reg64 (guarded by initiator/DRAM type preprocessor condition).\n11. Read back data from port0 addresses using read_reg64 and compare against expected data array; increment error counter on mismatch.\n12. Read back data from port1 addresses using read_reg64 and compare against expected data array; increment error counter on mismatch.",
        "Meta Impacted Registers": "NA",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "err0 is initialized to 0 at the start of the test.\nPort0 read-back validation: for(index = 0; index < 10; index++) { read_reg64(port0_addr + ((unsigned long)indexoffset), &read_data); if(read_data != exp_data_array[index]) { printf(\"ERROR_0: port0 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data); err0++; } }\nPort1 read-back validation: for(index = 0; index < 10; index++) { read_reg64(port1_addr + ((unsigned long)index * offset), &read_data); if(read_data != exp_data_array[index]) { printf(\"ERROR_0: port1 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data); err0++; } }\nPass condition: err0 == 0 (no mismatches detected across all read-back comparisons).\nFail condition: err0 != 0 (one or more read_data != exp_data_array[index] mismatches detected).",
        "Validation / Acceptance Criteria": "The test passes if the error counter (err0) remains at 0, indicating all 64-bit data values read back from both port0 and port1 match the expected data stored in exp_data_array. The test fails if any read-back data does not match the expected data, with each mismatch incrementing the error counter and printing the index, expected value, and actual value.",
        "Remarks": "Source code retrieved from RAG is assembled from multiple fragments. The function int test_case() declares local variables: unsigned long int index; unsigned long int offset; unsigned long int exp_data_array[50]; unsigned long int read_data; unsigned int err0;. Global variables: unsigned long int port0_addr; unsigned long int port1_addr;. DM_EN = 1 is the distinguishing feature of this test (Data Mask enabled). DBI_EN = 0 (Data Bus Inversion disabled). The port0 write/read block is conditionally compiled using #if !((defined(APS_DRAM) && defined(A53_INITIATOR)) ^(defined(MPS_DRAM) && defined(AI_INITIATOR)) ^(defined(MPS_DRAM) && defined(DSP_INITIATOR))). All Agent 4 register mappings (APS_DRAM, SG2667, SG2133) are unresolved with register_name NA. The complete end-of-function return logic is not fully visible in the retrieved source context."
    }
]


def generate():
    wb = openpyxl.Workbook()

    # --- TestPlan Sheet ---
    ws_tp = wb.active
    ws_tp.title = "TestPlan"

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_align = Alignment(wrap_text=True, vertical="top")

    for col_idx, col_name in enumerate(TESTPLAN_COLS, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_align

    for row_idx, row_data in enumerate(JSON_DATA, 2):
        for col_idx, col_name in enumerate(TESTPLAN_COLS, 1):
            value = row_data.get(col_name, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align

    ws_tp.freeze_panes = "A2"

    for col_idx, col_name in enumerate(TESTPLAN_COLS, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(JSON_DATA) + 2):
            val = str(ws_tp.cell(row=row_idx, column=col_idx).value or "")
            for line in val.split("\n"):
                max_len = max(max_len, len(line))
        ws_tp.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 15), 60)

    # --- MetaData Sheet ---
    ws_md = wb.create_sheet("MetaData")

    for col_idx, col_name in enumerate(METADATA_COLS, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_align

    for row_idx, row_data in enumerate(JSON_DATA, 2):
        for col_idx, col_name in enumerate(METADATA_COLS, 1):
            value = row_data.get(col_name, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align

    ws_md.freeze_panes = "A2"

    for col_idx, col_name in enumerate(METADATA_COLS, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(JSON_DATA) + 2):
            val = str(ws_md.cell(row=row_idx, column=col_idx).value or "")
            for line in val.split("\n"):
                max_len = max(max_len, len(line))
        ws_md.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 15), 60)

    ws_md.sheet_state = "veryHidden"

    # Save
    wb.save(FILENAME)
    wb.close()

    # Validate
    file_size = os.path.getsize(FILENAME)
    wb2 = openpyxl.load_workbook(FILENAME)
    assert "TestPlan" in wb2.sheetnames, "TestPlan sheet missing"
    assert "MetaData" in wb2.sheetnames, "MetaData sheet missing"

    ws_v = wb2["MetaData"]
    meta_fields = {
        "Meta Test Description": 3,
        "Meta Test Steps / Procedure": 4,
        "Meta Impacted Registers": 5,
        "Meta Validation / Acceptance Criteria": 6,
        "Meta Headers": 7,
        "Meta Macros": 8,
        "Meta Arrays": 9
    }
    for field_name, col in meta_fields.items():
        actual = ws_v.cell(row=2, column=col).value
        expected = JSON_DATA[0][field_name]
        assert actual == expected, f"MISMATCH in {field_name}"

    tp_rows = wb2["TestPlan"].max_row - 1
    md_rows = wb2["MetaData"].max_row - 1
    wb2.close()

    # Output base64 for upload
    with open(FILENAME, "rb") as f:
        b64 = base64.b64encode(f.read()).decode("utf-8")

    print(f"FILENAME={FILENAME}")
    print(f"FILESIZE={file_size}")
    print(f"VALIDATION=PASSED")
    print(f"ROWS_TP={tp_rows}")
    print(f"ROWS_MD={md_rows}")
    # Print base64 content
    print(f"BASE64_START")
    print(b64)
    print(f"BASE64_END")


if __name__ == "__main__":
    generate()
