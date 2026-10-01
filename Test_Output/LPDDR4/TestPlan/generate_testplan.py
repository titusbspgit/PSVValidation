#!/usr/bin/env python3
"""
LPDDR4 Test Plan Generator
Converts the LPDDR4_TestPlan.csv to a properly formatted Excel file
with IST-timestamped naming convention.

Usage: python generate_testplan.py
Output: LPDDR4_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx
"""

import csv
import datetime
import os

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("WARNING: openpyxl not installed. Install with: pip install openpyxl")

# IST timezone offset
IST_OFFSET = datetime.timezone(datetime.timedelta(hours=5, minutes=30))

def generate_xlsx():
    if not HAS_OPENPYXL:
        print("Cannot generate .xlsx without openpyxl. Please install it.")
        return

    # Read CSV
    script_dir = os.path.dirname(os.path.abspath(__file__))
    csv_path = os.path.join(script_dir, "LPDDR4_TestPlan.csv")

    if not os.path.exists(csv_path):
        print(f"ERROR: {csv_path} not found")
        return

    rows = []
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(row)

    if len(rows) < 2:
        print("ERROR: CSV has no data rows")
        return

    headers = rows[0]
    data_rows = rows[1:]

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "LPDDR4 Test Plan"

    # Styles
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write data
    for row_idx, data_row in enumerate(data_rows, 2):
        for col_idx, value in enumerate(data_row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_alignment
            cell.border = thin_border

    # Auto-adjust column widths (capped)
    for col_idx in range(1, len(headers) + 1):
        max_len = len(str(headers[col_idx - 1]))
        for row_idx in range(2, len(data_rows) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)), 80))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 60)

    # Freeze top row
    ws.freeze_panes = "A2"

    # Generate filename with IST timestamp
    now_ist = datetime.datetime.now(IST_OFFSET)
    timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
    filename = f"LPDDR4_TestPlan_{timestamp}.xlsx"
    output_path = os.path.join(script_dir, filename)

    wb.save(output_path)
    print(f"Generated: {output_path}")
    return output_path

if __name__ == "__main__":
    generate_xlsx()
