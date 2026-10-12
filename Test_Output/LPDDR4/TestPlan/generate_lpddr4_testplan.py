#!/usr/bin/env python3
"""
LPDDR4 TestPlan Excel Generator
Run this script to generate the LPDDR4_TestPlan_YYYYMMDD_HHMMSS.xlsx workbook.
Requirements: pip install openpyxl
Usage: python3 generate_lpddr4_testplan.py
"""
import json
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# ============================================================
# JSON DATA (exact content from pipeline)
# ============================================================
json_data = [
  {
    "Index": "1",
    "SS / Module": "LPDDR4",
    "Test Case Name": "lpddr4_mem_half_data_bus_width_test",
    "Feature": "Half Data Bus Width",
    "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include \"test_common.h\"\n#include \"lpddr4.h\"",
    "Meta Macros": "APS_DRAM - Preprocessor conditional compilation flag for APS DRAM configuration (used in #if defined(APS_DRAM)); no #define value; externally defined.\nMPS_DRAM - Preprocessor conditional compilation flag for MPS DRAM configuration (used in #else path of APS_DRAM); no #define value; externally defined.\nSG2667 - Preprocessor conditional compilation flag for Speed Grade 2667 selection (used in #if defined(SG2667)); no #define value; externally defined.\nSG2133 - Preprocessor conditional compilation flag for Speed Grade 2133 selection (used in #elif defined(SG2133)); no #define value; externally defined.\nA53_INITIATOR - Preprocessor conditional compilation flag for A53 initiator selection (used in #if defined(A53_INITIATOR)); no #define value; externally defined.\nAI_INITIATOR - Preprocessor conditional compilation flag for AI initiator selection (used in #if defined(AI_INITIATOR)); no #define value; externally <NRP> conditional compilation flag for DSP initiator selection (used in #if defined(DSP_INITIATOR)); no #define value; externally defined.",
    "Meta Arrays": "NA",
    "Speed": "SG3200 (default), SG2667 (when SG2667 defined), SG2133 (when SG2133 defined)",
    "Mode": "Half Data Bus Width (bus_width = 1)",
    "Memory Start Offset": "0x100",
    "Memory End Offset": "0x108",
    "Meta Test Description": "This testcase verifies LPDDR4 memory data integrity when operating in half data bus width mode (bus_width = 1). The test configures the LPDDR4 controller and PHY base addresses based on the DRAM type (APS_DRAM or MPS_DRAM), sets bus_width to 1 (Half Bus), selects the speed grade (SG3200 default, SG2667, or SG2133 via preprocessor defines), disables Data Bus Inversion (<NRP> = 0) and <PERSON> (DM_EN = 0), performs LPDDR4 memory training via <PERSON>(), confirms training completion via training_done(), waits 1000 units for stabilization via wait_on(1000), then performs 64-bit read-back verification on <PERSON> (conditionally compiled based on XOR of initiator/DRAM combinations) and <PERSON> (unconditionally) at offsets 0x100 and 0x108. Read data is compared against expected patterns 0x3333333333333333 and 0x2222222222222222. Errors are accumulated in err0 and the test concludes with finish(err0).",
    "Test Description": "Verify LPDDR4 memory data integrity in half data bus width mode after training. The test configures the memory controller for half bus width operation, performs LPDDR4 training, waits for stabilization, then reads back data from both memory ports at specific offsets and verifies the data matches expected patterns. The test supports both APS_DRAM and MPS_DRAM configurations with multiple speed grades (3200, 2667, 2133).",
    "Meta Test Steps / Procedure": "1. File-scope declarations:\n   unsigned long int port0_addr;\n   unsigned long int port1_addr;\n\n2. Function: int test_case()\n\n3. Local variable declarations:\n   unsigned long int index;\n   unsigned long int offset;\n   unsigned long int read_data0;\n   unsigned long int read_data1;\n   unsigned long int read_data2;\n   unsigned long int read_data3;\n   unsigned int err0;\n\n4. Call gpv_programming();\n\n5. err0 = 0;\n\n6. Address configuration (preprocessor conditional):\n   #if defined(APS_DRAM)\n     port0_addr = 0;\n     port1_addr = 0x15A0000000;\n     ctl_base = 0x9EE03000;\n     phy_base = 0x9F000000;\n   #else\n     port0_addr = 0;\n     port1_addr = 0x11A0000000;\n     ctl_base = 0x11D003000;\n     phy_base = 0x11D500000;\n   #endif\n\n7. bus_width = 1; //0 -> Full Bus, 1 -> Half Bus, 2 -> Quarter bus\n\n8. Speed grade selection (preprocessor conditional):\n   #if defined(SG2667)\n     SG = 2667;\n   #elif defined(SG2133)\n     SG = 2133;\n   #else\n     SG = 3200;\n   #endif\n\n9. <NRP> = 0;\n\n10. DM_EN = 0;\n\n11. Call <PERSON>();\n\n12. Call training_done();\n\n13. Call wait_on(1000);\n\n14. Port0 read-back verification (preprocessor conditional):\n   #if !((defined(APS_DRAM) && defined(A53_INITIATOR)) ^(defined(MPS_DRAM) && defined(AI_INITIATOR)) ^(defined(MPS_DRAM) && defined(DSP_INITIATOR)))\n     read_reg64(port0_addr + 0x100, &read_data0);\n     read_reg64(port0_addr + 0x108, &read_data1);\n     if((read_data0 != 0x3333333333333333) && (read_data1 != 0x2222222222222222))\n     {\n       printf(\"ERROR: port0 read_data0 = %lx, read_data1 = %lx\\n\", read_data0, read_data1);\n       err0++;\n     }\n   #endif\n\n15. Port1 read-back verification (unconditional):\n   read_reg64(port1_addr + 0x100, &read_data2);\n   read_reg64(port1_addr + 0x108, &read_data3);\n   if((read_data2 != 0x3333333333333333) && (read_data3 != 0x2222222222222222))\n   {\n     printf(\"ERROR: port1 read_data2 = %lx, read_data3 = %lx\\n\", read_data2, read_data3);\n     err0++;\n   }\n\n16. Call finish(err0);",
    "Test Steps / Procedure": "1. Perform GPV programming to set up the bus interconnect.\n2. Initialize error counter to 0.\n3. Configure port addresses and controller/PHY base addresses based on DRAM type (APS_DRAM or MPS_DRAM).\n4. Set bus width to Half Bus mode (bus_width = 1).\n5. Select speed grade: SG3200 (default), SG2667, or SG2133.\n6. Disable Data Bus Inversion (<NRP> = 0).\n7. Disable Data Mask (DM_EN = 0).\n8. Execute LPDDR4 memory training sequence.\n9. Confirm training completion.\n10. Wait 1000 units for stabilization.\n11. Conditionally read 64-bit data from <LOCATION> at offsets 0x100 and 0x108 and verify against expected patterns (condition based on initiator/DRAM combination).\n12. Read 64-bit data from <LOCATION> at offsets 0x100 and 0x108 and verify against expected <LOCATION>. Report test result based on accumulated error count.",
    "Meta Impacted Registers": "NA",
    "Impacted Registers": "NA",
    "Meta Validation / Acceptance Criteria": "Port0 validation (conditionally compiled based on: #if !((defined(APS_DRAM) && defined(A53_INITIATOR)) ^(defined(MPS_DRAM) && defined(AI_INITIATOR)) ^(defined(MPS_DRAM) && defined(DSP_INITIATOR)))):\n  read_reg64(port0_addr + 0x100, &read_data0): expected read_data0 == 0x3333333333333333\n  read_reg64(port0_addr + 0x108, &read_data1): expected read_data1 == 0x2222222222222222\n  Fail condition: if((read_data0 != 0x3333333333333333) && (read_data1 != 0x2222222222222222)) then err0++ and printf error message\n\nPort1 validation (unconditional):\n  read_reg64(port1_addr + 0x100, &read_data2): expected read_data2 == <NRP>  read_reg64(port1_addr + 0x108, &read_data3): expected read_data3 == 0x2222222222222222\n  Fail condition: if((read_data2 != 0x3333333333333333) && (read_data3 != 0x2222222222222222)) then err0++ and printf error message\n\nPass condition: err0 == 0 (no mismatches detected)\nFail condition: err0 > 0 (one or more read-back mismatches)\nFinal call: finish(err0)\n\nNote: The comparison uses && (logical AND), meaning both values at offset 0x100 and 0x108 must mismatch simultaneously to trigger an error increment.",
    "Validation / Acceptance Criteria": "1. <LOCATION> (conditional): Data read from offset 0x100 must equal 0x3333333333333333 and data read from offset 0x108 must equal 0x2222222222222222.\n2. <LOCATION> (unconditional): Data read from offset 0x100 must equal 0x3333333333333333 and data read from offset 0x108 must equal 0x2222222222222222.\n3. The test passes when the error counter (err0) is 0, indicating all read-back data matched the expected patterns.\n4. The test fails when err0 is greater than 0, indicating one or more data mismatches were detected.\n5. Note: The error condition uses logical AND (&&), so both values at offsets 0x100 and 0x108 must mismatch simultaneously to trigger an error.",
    "Remarks": "All 7 macros (APS_DRAM, MPS_DRAM, SG2667, SG2133, A53_INITIATOR, AI_INITIATOR, DSP_INITIATOR) are compile-time configuration flags expected to be defined externally via compiler command-line options. Port0 read-back verification is conditionally compiled using an XOR-based preprocessor expression involving initiator and DRAM type combinations. Port1 read-back verification is always compiled. The variables ctl_base, phy_base, bus_width, SG, <NRP>, DM_EN are declared in included headers (lpddr4.h or test_common.h) and assigned within test_case(). The local variables index and offset are declared but their usage is not visible in the available source context."
  }
]

# ============================================================
# WORKBOOK GENERATION
# ============================================================
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"LPDDR4_TestPlan_{timestamp}.xlsx"
script_dir = os.path.dirname(os.path.abspath(__file__))
filepath = os.path.join(script_dir, filename)

# Create workbook
wb = openpyxl.Workbook()

# ---- TestPlan Sheet ----
ws_tp = wb.active
ws_tp.title = "TestPlan"

tp_headers = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

# Header formatting
blue_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
white_font = Font(color="FFFFFF", bold=True)
wrap_align = Alignment(wrap_text=True, vertical="top")

for col_idx, header in enumerate(tp_headers, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=header)
    cell.fill = blue_fill
    cell.font = white_font
    cell.alignment = wrap_align

# Populate data
for row in json_data:
    tp_row = [
        row.get("Index", ""),
        row.get("SS / Module", ""),
        row.get("Feature", ""),
        row.get("Test Case Name", ""),
        row.get("Test Description", ""),
        row.get("Speed", ""),
        row.get("Mode", ""),
        row.get("Memory Start Offset", ""),
        row.get("Memory End Offset", ""),
        row.get("Remarks", ""),
        row.get("Test Steps / Procedure", ""),
        row.get("Impacted Registers", ""),
        row.get("Validation / Acceptance Criteria", ""),
        ""  # Code Generation - blank
    ]
    ws_tp.append(tp_row)

# Apply wrap text to data rows
for row in ws_tp.iter_rows(min_row=2, max_row=ws_tp.max_row, min_col=1, max_col=len(tp_headers)):
    for cell in row:
        cell.alignment = wrap_align

# Freeze first row
ws_tp.freeze_panes = "A2"

# Auto-size columns with max width
for col_idx in range(1, len(tp_headers) + 1):
    max_len = 0
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    for row in ws_tp.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws_tp.max_row):
        for cell in row:
            if cell.value:
                lines = str(cell.value).split('\n')
                max_line = max(len(line) for line in lines)
                if max_line > max_len:
                    max_len = max_line
    adjusted_width = min(max_len + 2, 55)
    ws_tp.column_dimensions[col_letter].width = max(adjusted_width, 12)

# ---- MetaData Sheet ----
ws_md = wb.create_sheet("MetaData")

md_headers = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

for col_idx, header in enumerate(md_headers, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=header)
    cell.fill = blue_fill
    cell.font = white_font
    cell.alignment = wrap_align

for row in json_data:
    md_row = [
        row.get("Index", ""),
        row.get("Test Case Name", ""),
        row.get("Meta Test Description", ""),
        row.get("Meta Test Steps / Procedure", ""),
        row.get("Meta Impacted Registers", ""),
        row.get("Meta Validation / Acceptance Criteria", ""),
        row.get("Meta Headers", ""),
        row.get("Meta Macros", ""),
        row.get("Meta Arrays", "")
    ]
    ws_md.append(md_row)

for row in ws_md.iter_rows(min_row=2, max_row=ws_md.max_row, min_col=1, max_col=len(md_headers)):
    for cell in row:
        cell.alignment = wrap_align

ws_md.freeze_panes = "A2"

for col_idx in range(1, len(md_headers) + 1):
    max_len = 0
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    for row in ws_md.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws_md.max_row):
        for cell in row:
            if cell.value:
                lines = str(cell.value).split('\n')
                max_line = max(len(line) for line in lines)
                if max_line > max_len:
                    max_len = max_line
    adjusted_width = min(max_len + 2, 55)
    ws_md.column_dimensions[col_letter].width = max(adjusted_width, 12)

# Set MetaData sheet to veryHidden
ws_md.sheet_state = "veryHidden"

# ============================================================
# VALIDATION
# ============================================================
validation_passed = True
meta_fields = [
    "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

for field in meta_fields:
    col_idx = md_headers.index(field) + 1
    cell_value = ws_md.cell(row=2, column=col_idx).value
    expected = json_data[0].get(field, "")
    if cell_value != expected:
        print(f"VALIDATION FAILED: {field}")
        print(f"  Expected: {expected[:100]}...")
        print(f"  Got: {str(cell_value)[:100]}...")
        validation_passed = False

if validation_passed:
    print("META CONTENT VALIDATION: PASSED")
else:
    print("META CONTENT VALIDATION: FAILED")
    sys.exit(1)

# Save workbook
wb.save(filepath)
print(f"Workbook saved: {filepath}")

# Post-save validation
if not os.path.exists(filepath):
    print("ERROR: File does not exist after save")
    sys.exit(1)

file_size = os.path.getsize(filepath)
if file_size <= 0:
    print("ERROR: File size is 0")
    sys.exit(1)

print(f"File size: {file_size} bytes")

# Reopen validation
wb2 = openpyxl.load_workbook(filepath)
if "TestPlan" not in wb2.sheetnames:
    print("ERROR: TestPlan sheet not found on reopen")
    sys.exit(1)
if "MetaData" not in wb2.sheetnames:
    print("ERROR: MetaData sheet not found on reopen")
    sys.exit(1)

tp_rows = wb2["TestPlan"].max_row - 1
md_rows = wb2["MetaData"].max_row - 1
print(f"TestPlan rows: {tp_rows}")
print(f"MetaData rows: {md_rows}")
print(f"ALL VALIDATIONS PASSED")
print(f"FILENAME={filename}")

# Output base64 for upload
import base64
with open(filepath, "rb") as f:
    b64 = base64.b64encode(f.read()).decode("utf-8")
print(f"BASE64_LENGTH={len(b64)}")

# Write base64 to a text file for retrieval
b64_path = os.path.join(script_dir, "workbook_b64.txt")
with open(b64_path, "w") as f:
    f.write(b64)
print(f"Base64 written to: {b64_path}")

wb2.close()
wb.close()
