#!/usr/bin/env python3
"""
Agent 7 - Excel Generator Script
Generates LPDDR4 TestPlan XLSX workbook using openpyxl.
This script is meant to be executed to produce the actual .xlsx file.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone, timedelta
import os
import sys
import json
import base64

# IST timezone
ist = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(ist)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"LPDDR4_TestPlan_{timestamp}.xlsx"

# Input JSON data
json_data = {
    "Index": "1",
    "SS / Module": "LPDDR4",
    "Test Case Name": "lpddr4_mem_dm_test",
    "Feature": "LPDDR4 Data Mask (DM)",
    "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include \"test_common.h\"\n#include \"lpddr4.h\"",
    "Meta Macros": "APS_DRAM (conditional compilation macro, controls port address and base address selection)\nSG2667 (conditional compilation macro, selects speed grade 2667)\nSG2133 (conditional compilation macro, selects speed grade 2133)\nA53_INITIATOR (conditional compilation macro, used in XOR-based port0 access guard)\nMPS_DRAM (conditional compilation macro, used in XOR-based port0 access guard)\nAI_INITIATOR (conditional compilation macro, used in XOR-based port0 access guard)\nDSP_INITIATOR (conditional compilation macro, used in XOR-based port0 access guard)",
    "Meta Arrays": "unsigned long int exp_data_array[50]; \u2014 Array of 50 elements of type unsigned long int. Used to store expected 64-bit random data values generated via rand() combined with ((unsigned long)rand() << 32). Populated in write loops for port0 (10 elements, offset 0x1000) and used for read-back comparison against read_data.",
    "Speed": "Configurable via preprocessor: SG2667 \u2192 SG = 2667, SG2133 \u2192 SG = 2133, default \u2192 SG = 3200",
    "Mode": "bus_width = 0 (Full Bus)",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "LPDDR4 Data Mask (DM) test. Function signature: int test_case(). Global variable declarations: unsigned long int port0_addr; unsigned long int port1_addr;. Local variable declarations: unsigned long int index; unsigned long int offset; unsigned long int exp_data_array[50]; unsigned long int read_data; unsigned int err0;. The test initializes GPV programming via gpv_programming(), sets err0 = 0, configures port addresses (port0_addr, port1_addr), controller base (ctl_base), and PHY base (phy_base) conditionally based on APS_DRAM define. Sets bus_width = 0 (Full Bus). Selects speed grade (SG) via conditional compilation: 2667 if SG2667 defined, 2133 if SG2133 defined, else 3200. Disables Data Bus Inversion (DBI_EN = 0). Enables Data Mask (DM_EN = 1). Calls lpddr4_training() to perform LPDDR4 memory training. After training, sets offset = 0x1000 and conditionally (guarded by XOR-based preprocessor condition on APS_DRAM/A53_INITIATOR, MPS_DRAM/AI_INITIATOR, MPS_DRAM/DSP_INITIATOR) writes 10 random 64-bit values to port0_addr using write_reg64(), then reads them back via read_reg64() comparing against exp_data_array[index]. Also performs port1 read-back verification loop of 10 iterations comparing read_data against exp_data_array[index]. Errors are tracked via err0 counter and reported via printf.",
    "Test Description": "This test validates the LPDDR4 Data Mask (DM) functionality. The test configures the memory controller and PHY base addresses, sets full bus width mode, selects the speed grade, disables Data Bus Inversion, enables Data Mask, performs LPDDR4 memory training, then writes random 64-bit data patterns to memory via port0 and reads them back via both port0 and port1 to verify data integrity with Data Mask enabled.",
    "Meta Test Steps / Procedure": "1. Call gpv_programming()\n2. err0 = 0\n3. Conditional on APS_DRAM:\n #if defined(APS_DRAM)\n port0_addr = 0\n port1_addr = 0x15A0000000\n ctl_base = 0x9EE03000\n phy_base = 0x9F000000\n #else\n port0_addr = 0\n port1_addr = 0x11A0000000\n ctl_base = 0x11D003000\n phy_base = 0x11D500000\n #endif\n4. bus_width = 0 //0 -> Full Bus, 1 -> Half Bus, 2 -> Quarter bus\n5. Conditional on SG2667/SG2133:\n #if defined(SG2667)\n SG = 2667\n #elif defined(SG2133)\n SG = 2133\n #else\n SG = 3200\n #endif\n6. DBI_EN = 0\n7. DM_EN = 1\n8. Call lpddr4_training()\n9. offset = 0x1000\n10. #if !((defined(APS_DRAM) && defined(A53_INITIATOR)) ^(defined(MPS_DRAM) && defined(AI_INITIATOR)) ^(defined(MPS_DRAM) && defined(DSP_INITIATOR)))\n Port0 Write Loop:\n for(index = 0; index < 10; index++)\n {\n exp_data_array[index] = rand();\n exp_data_array[index] = exp_data_array[index] | ((unsigned long)rand() << 32);\n write_reg64(port0_addr + ((unsigned long)index * offset), exp_data_array[index]);\n }\n11. Port0 Read-back Verification Loop:\n for(index = 0; index < 10; index++)\n {\n read_reg64(port0_addr + ((unsigned long)indexoffset), &read_data);\n if(read_data != exp_data_array[index])\n {\n printf(\"ERROR_0: port0 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data);\n err0++;\n }\n }\n #endif\n12. Port1 Read-back Verification Loop:\n for(index = 0; index < 10; index++)\n {\n read_reg64(port1_addr + ((unsigned long)index * offset), &read_data);\n if(read_data != exp_data_array[index])\n {\n printf(\"ERROR_0: port1 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data);\n err0++;\n }\n }",
    "Test Steps / Procedure": "1. Initialize GPV programming.\n2. Clear error counter (err0 = 0).\n3. Configure port addresses, controller base, and PHY base addresses (configuration varies based on APS_DRAM selection).\n4. Set bus width to Full Bus mode.\n5. Select speed grade (2667, 2133, or 3200 based on compile-time configuration).\n6. Disable Data Bus Inversion (DBI).\n7. Enable Data Mask (DM).\n8. Execute LPDDR4 memory training.\n9. Set memory access offset to 0x1000.\n10. Conditionally write 10 random 64-bit data values to port0 addresses using write_reg64 (guarded by initiator/DRAM type preprocessor condition).\n11. Read back data from port0 addresses using read_reg64 and compare against expected data array; increment error counter on mismatch.\n12. Read back data from port1 addresses using read_reg64 and compare against expected data array; increment error counter on mismatch.",
    "Meta Impacted Registers": "NA",
    "Impacted Registers": "NA",
    "Meta Validation / Acceptance Criteria": "err0 is initialized to 0 at the start of the test.\nPort0 read-back validation: for(index = 0; index < 10; index++) { read_reg64(port0_addr + ((unsigned long)indexoffset), &read_data); if(read_data != exp_data_array[index]) { printf(\"ERROR_0: port0 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data); err0++; } }\nPort1 read-back validation: for(index = 0; index < 10; index++) { read_reg64(port1_addr + ((unsigned long)index * offset), &read_data); if(read_data != exp_data_array[index]) { printf(\"ERROR_0: port1 index = %d, exp_data = %lx, actual_data = %lx\\n\", index, exp_data_array[index], read_data); err0++; } }\nPass condition: err0 == 0 (no mismatches detected across all read-back comparisons).\nFail condition: err0 != 0 (one or more read_data != exp_data_array[index] mismatches detected).",
    "Validation / Acceptance Criteria": "The test passes if the error counter (err0) remains at 0, indicating all 64-bit data values read back from both port0 and port1 match the expected data stored in exp_data_array. The test fails if any read-back data does not match the expected data, with each mismatch incrementing the error counter and printing the index, expected value, and actual value.",
    "Remarks": "Source code retrieved from RAG is assembled from multiple fragments. The function int test_case() declares local variables: unsigned long int index; unsigned long int offset; unsigned long int exp_data_array[50]; unsigned long int read_data; unsigned int err0;. Global variables: unsigned long int port0_addr; unsigned long int port1_addr;. DM_EN = 1 is the distinguishing feature of this test (Data Mask enabled). DBI_EN = 0 (Data Bus Inversion disabled). The port0 write/read block is conditionally compiled using #if !((defined(APS_DRAM) && defined(A53_INITIATOR)) ^(defined(MPS_DRAM) && defined(AI_INITIATOR)) ^(defined(MPS_DRAM) && defined(DSP_INITIATOR))). All Agent 4 register mappings (APS_DRAM, SG2667, SG2133) are unresolved with register_name NA. The complete end-of-function return logic is not fully visible in the retrieved source context."
}

# Create workbook
wb = openpyxl.Workbook()

# ===== TestPlan Sheet =====
ws_tp = wb.active
ws_tp.title = "TestPlan"

tp_headers = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

# Write headers
for col_idx, header in enumerate(tp_headers, 1):
    ws_tp.cell(row=1, column=col_idx, value=header)

# Write data row
tp_row_data = [
    json_data["Index"],
    json_data["SS / Module"],
    json_data["Feature"],
    json_data["Test Case Name"],
    json_data["Test Description"],
    json_data["Speed"],
    json_data["Mode"],
    json_data["Memory Start Offset"],
    json_data["Memory End Offset"],
    json_data["Remarks"],
    json_data["Test Steps / Procedure"],
    json_data["Impacted Registers"],
    json_data["Validation / Acceptance Criteria"],
    ""  # Code Generation - not in JSON input
]

for col_idx, value in enumerate(tp_row_data, 1):
    ws_tp.cell(row=2, column=col_idx, value=value)

# ===== MetaData Sheet =====
ws_md = wb.create_sheet("MetaData")

md_headers = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

# Write headers
for col_idx, header in enumerate(md_headers, 1):
    ws_md.cell(row=1, column=col_idx, value=header)

# Write data row
md_row_data = [
    json_data["Index"],
    json_data["Test Case Name"],
    json_data["Meta Test Description"],
    json_data["Meta Test Steps / Procedure"],
    json_data["Meta Impacted Registers"],
    json_data["Meta Validation / Acceptance Criteria"],
    json_data["Meta Headers"],
    json_data["Meta Macros"],
    json_data["Meta Arrays"]
]

for col_idx, value in enumerate(md_row_data, 1):
    ws_md.cell(row=2, column=col_idx, value=value)

# ===== Formatting =====
blue_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
white_font = Font(color="FFFFFF", bold=True)
wrap_alignment = Alignment(wrap_text=True, vertical="top")

for ws in [ws_tp, ws_md]:
    # Header formatting
    for cell in ws[1]:
        cell.font = white_font
        cell.fill = blue_fill
        cell.alignment = wrap_alignment

    # Data cell formatting
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_alignment

    # Freeze first row
    ws.freeze_panes = "A2"

    # Auto-size columns with max width cap
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                lines = str(cell.value).split('\n')
                max_line = max(len(line) for line in lines)
                max_length = max(max_length, max_line)
        adjusted_width = min(max_length + 2, 60)
        ws.column_dimensions[col_letter].width = max(adjusted_width, 15)

# MetaData sheet veryHidden
ws_md.sheet_state = "veryHidden"

# ===== Save =====
filepath = filename
wb.save(filepath)
wb.close()

# ===== Post-Save Validation =====
validation_passed = True
errors = []

# Check file exists and size
if not os.path.exists(filepath):
    validation_passed = False
    errors.append("File does not exist")
else:
    file_size = os.path.getsize(filepath)
    if file_size == 0:
        validation_passed = False
        errors.append("File size is 0")

# Reopen and validate
try:
    wb2 = openpyxl.load_workbook(filepath)
    if "TestPlan" not in wb2.sheetnames:
        validation_passed = False
        errors.append("TestPlan sheet missing")
    if "MetaData" not in wb2.sheetnames:
        validation_passed = False
        errors.append("MetaData sheet missing")

    # Validate MetaData content
    ws_v = wb2["MetaData"]
    meta_checks = {
        "Meta Test Description": (3, json_data["Meta Test Description"]),
        "Meta Test Steps / Procedure": (4, json_data["Meta Test Steps / Procedure"]),
        "Meta Impacted Registers": (5, json_data["Meta Impacted Registers"]),
        "Meta Validation / Acceptance Criteria": (6, json_data["Meta Validation / Acceptance Criteria"]),
        "Meta Headers": (7, json_data["Meta Headers"]),
        "Meta Macros": (8, json_data["Meta Macros"]),
        "Meta Arrays": (9, json_data["Meta Arrays"])
    }

    for field_name, (col, expected) in meta_checks.items():
        actual = ws_v.cell(row=2, column=col).value
        if actual != expected:
            validation_passed = False
            errors.append(f"{field_name} mismatch")

    # Count rows
    tp_rows = ws_v = wb2["TestPlan"].max_row - 1  # exclude header
    md_rows = wb2["MetaData"].max_row - 1  # exclude header

    wb2.close()
except Exception as e:
    validation_passed = False
    errors.append(f"Reopen failed: {str(e)}")

# Output base64 for GitHub upload
if validation_passed:
    with open(filepath, "rb") as f:
        b64_content = base64.b64encode(f.read()).decode("utf-8")
    print(f"===FILENAME==={filename}===FILENAME===")
    print(f"===FILESIZE==={file_size}===FILESIZE===")
    print(f"===VALIDATION===PASSED===VALIDATION===")
    print(f"===ROWS_TP==={tp_rows}===ROWS_TP===")
    print(f"===ROWS_MD==={md_rows}===ROWS_MD===")
    print(f"===BASE64_START===")
    print(b64_content)
    print(f"===BASE64_END===")
else:
    print(f"===VALIDATION===FAILED===VALIDATION===")
    print(f"===ERRORS==={json.dumps(errors)}===ERRORS===")
    sys.exit(1)
