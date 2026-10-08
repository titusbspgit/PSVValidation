#!/usr/bin/env python3
"""Auto-generated MIPI_CSI TestPlan Excel Generator - Agent 7"""
import os, json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

json_data = [
  {
    "Index": "1",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_dphy_lanes_test",
    "Feature": "DPHY Lane Configuration and CSI-2 Data Transfer",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Test Description": "Verify MIPI CSI-2 DPHY lane configuration by iterating through all supported lane counts (4 to 1) and performing CSI-2 control and data packet transfers using DMA. The test enables all CSI-2 host interrupts, initializes the D-PHY, waits for PHY stop state, and for each lane configuration transfers control packets followed by conditional data packets, validating DMA completion via interrupt polling.",
    "Test Steps / Procedure": "1. Enable all CSI-2 host interrupts by reading the main interrupt status register to clear pending interrupts, then writing enable masks to PHY Fatal, Packet Fatal, PHY, Line, Boundary Frame Fatal, Sequence Frame Fatal, CRC Frame Fatal, Payload CRC Fatal, Data ID, and ECC Corrected interrupt mask registers.\n2. Configure the virtual channel register based on the selected GDMA path and VC_ID, and enable control data transfer.\n3. Write the virtual channel and control data registers a second time to confirm configuration.\n4. Initialize the D-PHY via the PHY initialization sequence.\n5. Poll the PHY Stop State register until all lanes and clock lane report stop state (expected value 0x1000f).\n6. Set DMA channel 0 program counter to 0xE6000000 and DMA channel 1 program counter to 0xE6000500.\n7. For each lane configuration (4 lanes down to 1 lane):\n   a. Write the lane count to the N_LANES register.\n   b. Calculate the control packet count as ((VRES * 3) + 2).\n   c. Trigger the CSI-2 sequence by writing the lane count to the trigger register.\n   d. For each packet in the control packet count:\n      i. Enable DMA interrupts for both channels.\n      ii. Program DMA channel 0 for an 8-byte control data transfer and issue DMAGO.\n      iii. Poll the DMA interrupt status register until channel 0 transfer completes (bit 0 set).\n      iv. Clear the DMA channel 0 interrupt.\n      v. Read the transferred CSI-2 control data.\n      vi. If the data type indicates a long packet (data type > 0xf), extract the word count, align to 8 bytes, program DMA channel 1 for the data transfer, issue DMAGO, poll for channel 1 completion (bit 1 set), then clear the channel 1 interrupt.\n8. Signal test completion.",
    "Impacted Registers": "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Validation / Acceptance Criteria": "1. The PHY Stop State register must read the expected value confirming all data lanes and the clock lane have entered stop state before proceeding.\n2. DMA channel 0 transfer must complete as indicated by the DMA interrupt status register bit 0 being set for each control packet transfer.\n3. The CSI-2 control data type field must be correctly evaluated to determine whether a long-packet data transfer is required.\n4. When a long packet is detected, DMA channel 1 transfer must complete as indicated by the DMA interrupt status register bit 1 being set.\n5. DMA interrupts must be properly cleared after each channel transfer completion.\n6. The test must iterate through all four lane configurations (4, 3, 2, 1 lanes) and process all control packets for each configuration.\n7. The test must complete successfully by calling the finish routine with a pass indication.",
    "Remarks": "The testcase uses conditional compilation for GDMA path selection (GDMA0_PATH, GDMA1_PATH, GDMA2_PATH, GDMA3_PATH) and resolution (GDMA0_FULL_MEM). External functions snps_phy_init(), dma_trnsfr_instn_preload(), DMAGO_CSI(), and finish() are not defined locally. The virtual channel register and control data register are written twice in sequence. Two unresolved hex addresses exist: 0xa0243ffc (CSI-2 sequence trigger) and 0xE6001000 (control data destination). DMA base+offset register accesses use gdma_reg_base (0xE6A00000) with MIPI_CSI2_DMA_INTEN_OFFSET, MIPI_CSI2_DMA_INTMIS_OFFSET, MIPI_CSI2_DMA_INTCLR_OFFSET, and hardcoded offset 0x28.",
    "Meta Headers": "#include <stdio.h>; #include <stdlib.h>; #include \"test_common.h\"; #include \"mipi_csi2.h\"",
    "Meta Macros": "#define GDMA_CSI2_DATA_DEST_ADDR2 0xE6040080\n#define GDMA_CTRL_DATA_DEST_ADDR2 0xE6040000\n#define LS_LE_EN 1\n#define TOTAL_FRAME 1\n#define VC_ID 3\n#define VRES 1080 (if GDMA0_FULL_MEM defined) or 3 (otherwise)\n#define HRES 1920 (if GDMA0_FULL_MEM defined) or 64 (otherwise)\n#define DATA_TYPE CSI2_RGB888",
    "Meta Arrays": "NA",
    "Meta Test Description": "This testcase validates MIPI CSI-2 DPHY lane configuration and data transfer across all lane counts (4 lanes down to 1 lane). The test begins by enabling all CSI-2 host interrupts via the csi2_enable_interrupt() helper function, which reads INT_ST_MAIN to clear pending interrupts and then writes mask values to 10 interrupt mask registers.",
    "Meta Test Steps / Procedure": "1. Enter test_case() function.\n2. Call csi2_enable_interrupt().\n3. Conditional GDMA path selection.\n4. Write virtual channel and control data registers.\n5. Call snps_phy_init().\n6. Poll PHY_STOPSTATE until 0x1000f.\n7. Outer loop: lane_num 3 to 0.\n8. Inner loop: packet transfers with DMA.\n9. Call finish(0).",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "1. PHY_STOPSTATE polling until rd_data == 0x1000f.\n2. DMA channel 0 completion polling bit 0.\n3. CSI-2 control data type check.\n4. DMA channel 1 completion polling bit 1.\n5. DMA interrupt clear operations.\n6. Test completion: finish(0)."
  },
  {
    "Index": "2",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_test_pattern_generator",
    "Feature": "CSI-2 Internal Test Pattern Generator and DMA Data Transfer",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Test Description": "Verify the MIPI CSI-2 internal test pattern generator by configuring the pattern generator with a specific resolution (320x16, 24 bits per pixel), enabling it to generate test data, transferring the generated data via DMA, and validating DMA transfer completion through interrupt polling. The test also configures virtual channels, initializes the D-PHY, and enables all CSI-2 host interrupts.",
    "Test Steps / Procedure": "1. Set the interrupt pending flag and configure the virtual channel ID (VC_ID=3) with the unselected path value for the selected GDMA path.\n2. Calculate the GDMA register base address based on the selected GDMA path.\n3. Write the virtual channel register with the computed VC ID mapping for the selected GDMA path.\n4. Disable control data transfer by writing 0 to the control data register.\n5. Enable all CSI-2 host and DMA interrupts by reading the main interrupt status register to clear pending interrupts, then writing enable masks to PHY Fatal, Packet Fatal, PHY, Line, Boundary Frame Fatal, Sequence Frame Fatal, CRC Frame Fatal, Payload CRC Fatal, Data ID, and ECC Corrected interrupt mask registers.\n6. Initialize the D-PHY via the PHY initialization sequence.\n7. Poll the PHY Stop State register until all data lanes and clock lane report stop state (expected value 0x1000f).\n8. Set the pattern generator resolution parameters: hres=320, vres=16, valid_bits_per_pixel=24.\n9. Calculate the DMA transfer size as an 8-byte aligned value of (valid_bits_per_pixel * hres / 8) * vres.\n10. Program the DMA higher-order AXI address registers for channel 0 read and write paths.\n11. Enable the fracdiv clock output to the CSI-2 subsystem by writing to the clock gating register.\n12. Program DMA channel 0 for data transfer with source address 0x00, destination address 0xE6001000, and the calculated transfer size.\n13. Issue DMAGO for DMA channel 0 to start the transfer.\n14. Enable the internal test pattern generator by configuring vertical resolution, horizontal resolution, pattern configuration (RGB888 data type), and enabling the generator.\n15. Wait for the pattern generator to produce data.\n16. Disable the pattern generator.\n17. Poll the DMA interrupt status register until channel 0 transfer completes (bit 0 set).\n18. Wait after DMA completion.\n19. Signal test completion with pass status.",
    "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; enableclkgating_csiphy; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Validation / Acceptance Criteria": "1. The PHY Stop State register must read the expected value confirming all data lanes and the clock lane have entered stop state before proceeding.\n2. The DMA transfer size must be correctly calculated as an 8-byte aligned value based on the configured resolution and bits per pixel.\n3. The pattern generator must be configured with the correct vertical resolution, horizontal resolution, and RGB888 data type before being enabled.\n4. The pattern generator must be disabled after the configured wait period.\n5. DMA channel 0 transfer must complete as indicated by the DMA interrupt status register bit 0 being set.\n6. The DMA higher-order AXI address registers must be correctly programmed before initiating the transfer.\n7. The clock gating register must be enabled to allow fracdiv output to the CSI-2 subsystem.\n8. The test must complete successfully by calling the finish routine with a pass indication.",
    "Remarks": "The testcase uses conditional compilation for GDMA path selection (GDMA0_PATH, GDMA1_PATH, GDMA2_PATH, GDMA3_PATH) and FPS60 mode. The test calls csi2_subsys_enable_interrupt() (external) which likely wraps the locally defined csi2_enable_interrupt(). External functions snps_phy_init(), dma_trnsfr_instn_preload_incr_addr(), DMAGO_CSI(), wait_on(), and finish() are not defined locally. The pattern generator uses RGB888 data type (0x24) with 320x16 resolution. The control_data register is explicitly set to 0 (disabled) unlike the dphy_lanes_test which sets it to 1. DMA base+offset register accesses use gdma_reg_base with MIPI_CSI2_DMA_INTMIS_OFFSET. The PPI_PG_ENABLE register is written twice: once with value 1 (enable) inside csi2_ctrlr_pg_enable() and once with value 0 (disable) in test_case() after wait_on(100).",
    "Meta Headers": "#include <stdio.h>; #include <stdlib.h>; #include \"test_common.h\"; #include \"mipi_csi2.h\"",
    "Meta Macros": "#define GDMA_CSI2_DATA_DEST_ADDR2 0xE6040080\n#define GDMA_CTRL_DATA_DEST_ADDR2 0xE6040000",
    "Meta Arrays": "NA",
    "Meta Test Description": "This testcase validates the MIPI CSI-2 internal test pattern generator functionality and DMA-based data transfer.",
    "Meta Test Steps / Procedure": "1. Enter test_case() function.\n2. Configure virtual channel.\n3. Call csi2_subsys_enable_interrupt().\n4. Call snps_phy_init().\n5. Poll PHY_STOPSTATE.\n6. Configure pattern generator.\n7. Program DMA.\n8. Enable/disable pattern generator.\n9. Poll DMA completion.\n10. Call finish(0).",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "1. PHY_STOPSTATE polling until 0x1000f.\n2. DMA transfer size calculation validation.\n3. Pattern generator configuration validation.\n4. DMA channel 0 completion polling.\n5. Pattern generator enable/disable sequence.\n6. DMA address register programming.\n7. Clock gating enable.\n8. Test completion: finish(0)."
  }
]

def generate():
    wb = Workbook()
    # TestPlan sheet
    ws_tp = wb.active
    ws_tp.title = 'TestPlan'
    tp_cols = ['Index','SS / Module','Feature','Test Case Name','Test Description','Speed','Mode','Memory Start Offset','Memory End Offset','Remarks','Test Steps / Procedure','Impacted Registers','Validation / Acceptance Criteria','Code Generation']
    # MetaData sheet
    ws_md = wb.create_sheet('MetaData')
    md_cols = ['Index','Test Case Name','Meta Test Description','Meta Test Steps / Procedure','Meta Impacted Registers','Meta Validation / Acceptance Criteria','Meta Headers','Meta Macros','Meta Arrays']
    ws_md.sheet_state = 'veryHidden'
    hdr_font = Font(bold=True, color='FFFFFF')
    hdr_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    wrap = Alignment(wrap_text=True, vertical='top')
    for ci, col in enumerate(tp_cols, 1):
        c = ws_tp.cell(row=1, column=ci, value=col)
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = wrap
    for ci, col in enumerate(md_cols, 1):
        c = ws_md.cell(row=1, column=ci, value=col)
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = wrap
    for ri, row in enumerate(json_data, 2):
        for ci, col in enumerate(tp_cols, 1):
            val = row.get(col, '')
            c = ws_tp.cell(row=ri, column=ci, value=val if val else '')
            c.alignment = wrap
        for ci, col in enumerate(md_cols, 1):
            val = row.get(col, '')
            c = ws_md.cell(row=ri, column=ci, value=val if val else '')
            c.alignment = wrap
    ws_tp.freeze_panes = 'A2'
    ws_md.freeze_panes = 'A2'
    for ws in [ws_tp, ws_md]:
        for col_cells in ws.columns:
            max_len = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 80))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)
    fname = 'MIPI_CSI_TestPlan_20261009_021000.xlsx'
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), fname)
    wb.save(out_path)
    print(f'Saved: {out_path}')
    print(f'Size: {os.path.getsize(out_path)} bytes')
    return out_path

if __name__ == '__main__':
    generate()
