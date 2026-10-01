#!/usr/bin/env python3
"""MIPI_DSI TestPlan Excel Generator - Agent 7 Fallback Automation
Generates a genuine Office Open XML (.xlsx) workbook using openpyxl.
Run: python generate_mipi_dsi_testplan_final.py
# Re-triggered: workflow_dispatch equivalent
"""
import json, os, sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment

# -- JSON DATA --
json_data = [
  {
    "Index": "1",
    "SS / Module": "MIPI_DSI",
    "Test Case Name": "mipi_dsi_basic_test",
    "Feature": "DBI Command Mode DMA Write",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_dsi.h\"",
    "Meta Macros": "NA",
    "Meta Arrays": "desc_t data_tdbdcb[2]; desc_t data_rebdcb[2]; desc_t cmd_tdbdcb[2]; desc_t cmd_rebdcb[2]",
    "Speed": "NA",
    "Mode": "DBI Command Mode",
    "Memory Start Offset": "0x0000; 0x10000; 0x3F000; 0x3F800",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase performs a basic MIPI DSI DBI write operation using DMA. It configures the DSI host PHY with 4 lanes (n_lanes=3 means 4 data lanes) and phy_stop_wait_time=0x40 via MIZAR_MIPI_DSI_HOST_PHY_IF_CFG. It sets PCKHDL_CFG to 0x3d and CLKMGR_CFG to 0x107. DBI mode is enabled by writing 0 to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL. PHY initialization is performed via phy_init(). DBI is configured with dbi_vcid=0x3, lut_size_conf=0x1, out_dbi_conf=0xb, in_dbi_conf=0x0, partitioning_en=0x1. The test generates 40 random pixels, converts them to byte count and wr_cmd_size via pixel_to_bytes_wr_cmd_size(). DPI clock is programmed based on dpi_clk_time_period=16.012400ns. dbi_config() is called to apply DBI configuration. A single DMA descriptor is set up for CH0 (data channel) sourcing from RAM_BASE+0x10000 and CH1 (command channel) sourcing from RAM_BASE+0x0000. DMA microcode programs are built at ch0_desc_addr (RAM_BASE+0x3F000) and ch1_desc_addr (RAM_BASE+0x3F800) using DMAMOV, program_data_num_bytes, DMAWMB, DMASEV, DMAEND. Random pixel data is loaded via load_rand_data() and write_memory_start command is loaded via load_wr_command() with DSI_WRITE_MEMORY_START. DMAC interrupts are enabled by writing 0x3 to MIZAR_MIPI_DSI_DMAC_INTEN and subsystem GDMA interrupt is enabled via MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE. DMA channels are started via the DMAC debug interface: MIZAR_MIPI_DSI_DMAC_DBGINST0 is written with 0x00A00000 for CH0 and 0x01A00000 for CH1, MIZAR_MIPI_DSI_DMAC_DBGINST1 is written with the respective descriptor addresses, and MIZAR_MIPI_DSI_DMAC_DBGCMD is written with 0x0 to execute. The test then waits in a while loop polling int_pend1 which is cleared by the ISR. The ISR (Default_IRQHandler) reads MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK to check for GDMA interrupt, reads MIZAR_MIPI_DSI_DMAC_INTMIS to identify which channel completed, clears the DMAC interrupt via MIZAR_MIPI_DSI_DMAC_INTCLR, and clears the subsystem interrupt via MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW. If unexpected interrupts occur, err0 is incremented. The test completes by calling finish(err0) where err0==0 indicates pass.",
    "Test Description": "This test validates a basic MIPI DSI DBI write operation using DMA transfer. It configures the DSI host PHY interface for 4 data lanes with a stop wait time, sets packet handling and clock manager configurations, enables DBI command mode, initializes the PHY, and configures DBI parameters including virtual channel ID, LUT size, output/input DBI format, and partitioning. The test prepares 40 random pixels of data and a write_memory_start command, builds single-descriptor DMA microcode programs for both data (CH0) and command (CH1) channels, and triggers DMA execution via the DMAC debug interface. Completion is verified through an interrupt-driven ISR that checks the subsystem interrupt mask for GDMA interrupt, reads the DMAC masked interrupt status to identify completed channels, clears both DMAC and subsystem-level interrupts, and tracks unexpected interrupt errors. The test passes when both DMA channels complete without unexpected interrupts.",
    "Meta Test Steps / Procedure": "1. Enable GIC IRQ for DSI_INTR_NO via GIC_EnableIRQ(DSI_INTR_NO). 2. Initialize int_pend=1, int_pend1=0x3, phy_stop_wait_time=0x40, n_lanes=3. 3. Write 0x3 to MIZAR_MIPI_DSI_DMAC_INTEN to enable interrupts for DMA CH0 and CH1. 4. Write MIPI_DSI_SUBSYS_INTERRUPT_ENABLE_GDMA_INTR to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE to enable GDMA interrupt at subsystem level. 5. Compute phy_if_cfg = n_lanes (3), then set phy_stop_wait_time field using set_data_mask with MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME mask and value 0x40. 6. Write computed phy_if_cfg to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG. 7. Write 0x3d to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG. 8. Write 0x107 to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG. 9. Write 0 to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to enable DBI mode. 10. Call phy_init() to initialize the DSI PHY. 11. Set DBI configuration variables. 12. Set num_of_pixel=40, call pixel_to_bytes_wr_cmd_size(40). 13. Set DBI command mode parameters. 14. Compute dpi_clk_freq, call program_dpi_clock(dpi_clk_freq). 15. Call dbi_config(). 16-25. Build DMA descriptors and microcode for CH0 and CH1. 26-31. Start DMA CH0 and CH1 via DMAC debug interface. 32-33. Poll int_pend1 until 0. 34-44. ISR handles GDMA interrupts. 45. Call finish(err0).",
    "Test Steps / Procedure": "1. Enable the DSI interrupt at the GIC level. 2. Enable DMAC interrupts for both DMA channels (CH0 and CH1) by writing to the INTEN register. 3. Enable the GDMA interrupt at the DSI subsystem level by writing to the interrupt_enable register. 4. Configure the PHY interface with 4 data lanes and stop wait time by writing to the PHY_IF_CFG register. 5. Configure packet handling by writing to the PCKHDL_CFG register. 6. Configure the clock manager by writing to the CLKMGR_CFG register. 7. Enable DBI command mode by writing to the dpi_control register. 8. Initialize the DSI PHY. 9. Configure DBI parameters including virtual channel ID, LUT size, output/input DBI format, partitioning, and command sizes. 10. Compute pixel-to-byte conversion for 40 pixels and set the write command size. 11. Program the DPI clock frequency. 12. Apply the DBI configuration to the DSI host. 13. Build DMA microcode descriptors for data channel (CH0) and command channel (CH1) with source/destination addresses and transfer lengths. 14. Load random pixel data into SRAM for the data channel. 15. Load the write_memory_start command into SRAM for the command channel. 16. Start DMA CH0 by writing the DMAGO instruction and descriptor address via the DBGINST0, DBGINST1, and DBGCMD registers. 17. Start DMA CH1 by writing the DMAGO instruction and descriptor address via the DBGINST0, DBGINST1, and DBGCMD registers. 18. Wait for both DMA channels to complete by polling the interrupt-driven completion flag. 19. Verify that the ISR correctly identifies GDMA interrupts via the interrupt_mask register, reads channel completion status from the INTMIS register, clears DMAC interrupts via the INTCLR register, and clears subsystem interrupts via the interrupt_raw register. 20. Confirm test passes with zero unexpected interrupt errors.",
    "Meta Impacted Registers": "MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE; MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW",
    "Impacted Registers": "INTEN; interrupt_enable; PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control; DBGINST0; DBGINST1; DBGCMD; interrupt_mask; INTMIS; INTCLR; interrupt_raw",
    "Meta Validation / Acceptance Criteria": "The ISR reads MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK and compares it against MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR. If matched, it reads MIZAR_MIPI_DSI_DMAC_INTMIS and checks if ch_mask_st is non-zero. If non-zero, the corresponding bit in int_pend1 is cleared, and the DMAC interrupt is cleared by writing ch_mask_st to MIZAR_MIPI_DSI_DMAC_INTCLR. The subsystem interrupt is cleared by writing dsi_subsys_mask_st to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW. The test passes when finish(err0) is called with err0==0.",
    "Validation / Acceptance Criteria": "The test passes when both DMA channels (CH0 and CH1) complete their transfers and generate interrupts. The ISR must correctly identify the GDMA interrupt source by reading the interrupt_mask register, confirm the specific channel completion by reading the INTMIS register, clear the DMAC interrupt via the INTCLR register, and clear the subsystem-level interrupt via the interrupt_raw register. Both channel completion flags must be received. No unexpected interrupt sources should be detected. The error counter must remain zero for the test to pass.",
    "Remarks": "The test uses an interrupt-driven model with a polling wait loop in the main function. The ISR handles both CH0 and CH1 completion interrupts independently. int_pend1 is initialized to 0x3 (bits for both channels) and each channel interrupt clears its respective bit. The test uses a single DMA descriptor per channel. DBI mode is selected by writing 0 to dpi_control. The DMAC debug interface is used to start DMA channels rather than a standard channel start mechanism."
  },
  {
    "Index": "2",
    "SS / Module": "MIPI_DSI",
    "Test Case Name": "mipi_dsi_dbi_random_payload_test",
    "Feature": "DBI Command Mode Random Payload DMA Write",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_dsi.h\"; <math.h>",
    "Meta Macros": "NA",
    "Meta Arrays": "desc_t data_tdbdcb[10]; desc_t data_rebdcb[10]; desc_t cmd_tdbdcb[10]; desc_t cmd_rebdcb[10]; unsigned int wr_cmd_arr[27]; unsigned int WR_CMD_SIZE[27]",
    "Speed": "NA",
    "Mode": "DBI Command Mode",
    "Memory Start Offset": "0x0000; 0x10000; 0x3F000; 0x3F800",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase performs a MIPI DSI DBI random payload stress test using DMA. It configures the DSI host PHY with 4 lanes (n_lanes=3 means 4 data lanes) and phy_stop_wait_time=0x40 via MIZAR_MIPI_DSI_HOST_PHY_IF_CFG. It sets PCKHDL_CFG to 0x3d and CLKMGR_CFG to 0x107. DBI mode is enabled by writing 0 to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL. The test runs 10 iterations with randomized pixel counts (multiples of 8, up to 8192). Each iteration reconfigures DBI wr_cmd_size, programs DMA, polls INTMIS for completion. Tests data path robustness.",
    "Test Description": "This test validates MIPI DSI DBI write operations with randomized payload sizes over multiple iterations using DMA transfer. It configures the DSI host PHY interface for 4 data lanes with a stop wait time, sets packet handling and clock manager configurations, enables DBI command mode, initializes the PHY, and configures DBI parameters including virtual channel ID, LUT size, output/input DBI format, and partitioning. The test runs 10 iterations, each with a randomly generated pixel count (multiples of 8, up to 8192). In each iteration, the DBI write command size is reconfigured, DMA microcode descriptors are rebuilt for both data (CH0) and command (CH1) channels, random pixel data and a write_memory_start command are loaded into SRAM, and DMA execution is triggered via the DMAC debug interface. Completion of both channels is verified by polling the DMAC masked interrupt status register until both channel completion bits are set, after which the interrupts are cleared. The test passes after all 10 iterations complete successfully.",
    "Meta Test Steps / Procedure": "1. Initialize phy_stop_wait_time=0x40, n_lanes=3. 2. Compute phy_if_cfg. 3-6. Write PHY_IF_CFG, PCKHDL_CFG, CLKMGR_CFG, DPI_CONTROL. 7. Call phy_init(). 8-15. Configure DBI and enable DMAC interrupts. 16-39. Loop 10 iterations: generate random pixels, reconfigure DBI, build DMA descriptors, start DMA, poll INTMIS, clear INTCLR. 40-41. Complete all iterations and call finish(0).",
    "Test Steps / Procedure": "1. Configure the PHY interface with 4 data lanes and stop wait time by writing to the PHY_IF_CFG register. 2. Configure packet handling by writing to the PCKHDL_CFG register. 3. Configure the clock manager by writing to the CLKMGR_CFG register. 4. Enable DBI command mode by writing to the dpi_control register. 5. Initialize the DSI PHY. 6. Configure DBI parameters including virtual channel ID, LUT size, output/input DBI format, partitioning, and allowed command size. 7. Compute pixel-to-byte conversion and set the initial write command size. 8. Program the DPI clock frequency. 9. Apply the DBI configuration to the DSI host. 10. Enable DMAC interrupts for both DMA channels (CH0 and CH1) by writing to the INTEN register. 11. Begin iterating 10 times with a randomly generated pixel count (multiples of 8, up to 8192) for each iteration. 12. Recompute the write command size and byte count for the new random pixel count. 13. Re-apply the DBI configuration with the updated write command size. 14. Build DMA microcode descriptors for data channel (CH0) and command channel (CH1) with source/destination addresses and 8-byte aligned transfer lengths. 15. Load random pixel data into SRAM for the data channel. 16. Load the write_memory_start command into SRAM for the command channel. 17. Start DMA CH0 by writing the DMAGO instruction and descriptor address via the DBGINST0, DBGINST1, and DBGCMD registers. 18. Start DMA CH1 by writing the DMAGO instruction and descriptor address via the DBGINST0, DBGINST1, and DBGCMD registers. 19. Poll the INTMIS register until both channel completion bits are set. 20. Clear the DMAC interrupts by writing to the INTCLR register. 21. Wait for settling delay before the next iteration. 22. Repeat steps 11-21 for all 10 iterations. 23. Confirm test passes after all iterations complete successfully.",
    "Meta Impacted Registers": "MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR",
    "Impacted Registers": "PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control; INTEN; DBGINST0; DBGINST1; DBGCMD; INTMIS; INTCLR",
    "Meta Validation / Acceptance Criteria": "In each of the 10 iterations, after starting both DMA channels via the DMAC debug interface, the test polls MIZAR_MIPI_DSI_DMAC_INTMIS in a while loop until rd_data equals 0x3. Once rd_data == 0x3, the test writes rd_data to MIZAR_MIPI_DSI_DMAC_INTCLR to clear both channel interrupts. The test unconditionally calls finish(0) after all 10 iterations complete.",
    "Validation / Acceptance Criteria": "For each of the 10 iterations, both DMA channels (CH0 and CH1) must complete their transfers, indicated by both completion bits being set in the INTMIS register. The interrupts must be successfully cleared via the INTCLR register after each iteration. All 10 iterations with varying random payload sizes must complete without the polling loop hanging. The test passes unconditionally after all iterations complete. A timeout or hang during polling indicates a DMA transfer failure.",
    "Remarks": "This test uses a polling-based model (not interrupt-driven ISR) to check DMA completion, unlike mipi_dsi_basic_test which uses an ISR. The test runs 10 iterations with randomized pixel counts generated as multiples of 8 (up to 8192). In each iteration, dbi_config() is re-called to update the wr_cmd_size for the new payload size. The DMA data length is 8-byte aligned using ceil((double)len/8)*8. The test uses a single DMA descriptor per channel per iteration. DBI mode is selected by writing 0 to dpi_control. partitioning_en is set to 0x0 in this test (unlike mipi_dsi_basic_test which uses 0x1). The DMAC debug interface is used to start DMA channels. finish(0) is called unconditionally, so the test always reports pass if all iterations complete."
  }
]

# -- SHEET DEFINITIONS --
TP_COLS = ["Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
           "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
           "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
           "Code Generation"]

MD_COLS = ["Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
           "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
           "Meta Headers", "Meta Macros", "Meta Arrays"]

def generate():
    IST = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(IST)
    ts = now_ist.strftime("%Y%m%d_%H%M%S")
    filename = f"MIPI_DSI_TestPlan_{ts}.xlsx"

    wb = Workbook()
    # -- TestPlan sheet --
    ws_tp = wb.active
    ws_tp.title = "TestPlan"
    ws_tp.append(TP_COLS)
    for row in json_data:
        ws_tp.append([row.get(c, "") for c in TP_COLS])

    # -- MetaData sheet --
    ws_md = wb.create_sheet("MetaData")
    ws_md.append(MD_COLS)
    for row in json_data:
        ws_md.append([row.get(c, "") for c in MD_COLS])

    # -- Formatting --
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")

    for ws in [ws_tp, ws_md]:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = wrap
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        # Auto-size columns
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    val = str(cell.value) if cell.value else ""
                    max_len = max(max_len, min(len(val), 80))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)

    # -- Hide MetaData --
    ws_md.sheet_state = "veryHidden"

    # -- Save --
    script_dir = os.path.dirname(os.path.abspath(__file__))
    filepath = os.path.join(script_dir, filename)
    wb.save(filepath)
    print(f"Saved: {filepath}")

    # -- Validate --
    assert os.path.exists(filepath), "File not found"
    assert os.path.getsize(filepath) > 0, "File is empty"
    vwb = load_workbook(filepath)
    assert "TestPlan" in vwb.sheetnames, "TestPlan sheet missing"
    assert "MetaData" in vwb.sheetnames, "MetaData sheet missing"
    assert vwb["MetaData"].sheet_state == "veryHidden", "MetaData not veryHidden"
    tp_rows = vwb["TestPlan"].max_row - 1
    md_rows = vwb["MetaData"].max_row - 1
    print(f"Validation PASSED: TestPlan={tp_rows} rows, MetaData={md_rows} rows, size={os.path.getsize(filepath)} bytes")
    print(f"FILENAME={filename}")
    return filename

if __name__ == "__main__":
    generate()
