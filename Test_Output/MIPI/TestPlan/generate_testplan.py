#!/usr/bin/env python3
"""MIPI_DSI TestPlan Excel Generator - Agent 7"""
import json
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

# ============================================================
# INPUT DATA
# ============================================================
json_data = [
    {
        "Index": "1",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_basic_test",
        "Feature": "NA",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs basic MIPI DSI validation by configuring DMAC, HOST, and SUBSYS register blocks. It writes to the DMAC interrupt enable register (MIZAR_MIPI_DSI_DMAC_INTEN), enables the GDMA interrupt in the subsystem interrupt enable register (MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE), configures the PHY interface (MIZAR_MIPI_DSI_HOST_PHY_IF_CFG), packet handler (MIZAR_MIPI_DSI_HOST_PCKHDL_CFG), and clock manager (MIZAR_MIPI_DSI_HOST_CLKMGR_CFG). It disables DPI control via the subsystem DPI control register (MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL). It programs the DMAC debug instruction registers (MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1) and executes the debug command (MIZAR_MIPI_DSI_DMAC_DBGCMD). It then reads the subsystem interrupt mask register (MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK) and DMAC masked interrupt status register (MIZAR_MIPI_DSI_DMAC_INTMIS) to check interrupt status. Finally, it clears the DMAC interrupt (MIZAR_MIPI_DSI_DMAC_INTCLR) and the subsystem raw interrupt (MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW).",
        "Test Description": "This testcase performs basic MIPI DSI validation by configuring the DMAC, HOST, and SUBSYS register blocks. It enables DMAC and subsystem interrupts via the interrupt_enable register, configures the PHY interface using PHY_IF_CFG, sets up the packet handler via PCKHDL_CFG, and programs the clock manager through CLKMGR_CFG. DPI control is disabled via the dpi_control register. The test programs DMAC debug instruction registers and executes a debug command to initiate a DMA transfer. It then reads the interrupt_mask register and a DMAC masked interrupt status register to verify interrupt assertion. Finally, it clears the DMAC interrupt and the subsystem raw interrupt via the interrupt_raw register.",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_DSI_DMAC_INTEN with value 0x3 to enable DMAC interrupts. 2. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE with MIPI_DSI_SUBSYS_INTERRUPT_ENABLE_GDMA_INTR to enable GDMA interrupt in subsystem. 3. Write to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG with computed phy_if_cfg value to configure PHY interface (using set_data_mask for PHY_STOP_WAIT_TIME field). 4. Write to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG with value 0x3d to configure packet handler. 5. Write to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG with value 0x107 to configure clock manager. 6. Write to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL with value 0 to disable DPI control. 7. Write to MIZAR_MIPI_DSI_DMAC_DBGINST0 with DMA instruction opcode to set up DMA channel 0 instruction. 8. Write to MIZAR_MIPI_DSI_DMAC_DBGINST1 with DMA descriptor address for channel 0. 9. Write to MIZAR_MIPI_DSI_DMAC_DBGCMD with value 0x0 to execute the debug command. 10. Read MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK to check subsystem interrupt mask status. 11. Read MIZAR_MIPI_DSI_DMAC_INTMIS to check DMAC masked interrupt status. 12. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear DMAC channel interrupt. 13. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW to clear subsystem raw interrupt.",
        "Test Steps / Procedure": "1. Enable DMAC interrupts by writing to the DMAC interrupt enable register. 2. Enable the GDMA interrupt in the subsystem by writing to the interrupt_enable register. 3. Configure the PHY interface by writing the PHY stop wait time field to the PHY_IF_CFG register. 4. Configure the packet handler by writing to the PCKHDL_CFG register. 5. Configure the clock manager by writing to the CLKMGR_CFG register. 6. Disable DPI control by writing zero to the dpi_control register. 7. Program the DMAC debug instruction registers with the DMA channel instruction opcode and descriptor address. 8. Execute the DMAC debug command to initiate the DMA transfer. 9. Read the interrupt_mask register to verify the subsystem interrupt mask status. 10. Read the DMAC masked interrupt status register to verify DMAC interrupt assertion. 11. Clear the DMAC interrupt by writing to the DMAC interrupt clear register. 12. Clear the subsystem raw interrupt by writing to the interrupt_raw register.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE; MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW",
        "Impacted Registers": "interrupt_enable; PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control; interrupt_mask; interrupt_raw",
        "Meta Validation / Acceptance Criteria": "After enabling DMAC and subsystem interrupts and executing the DMA debug command, the test reads MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK to verify the GDMA interrupt bit is asserted (compared against MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR). It also reads MIZAR_MIPI_DSI_DMAC_INTMIS to verify the DMAC masked interrupt status. The test clears interrupts by writing to MIZAR_MIPI_DSI_DMAC_INTCLR and MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW. The test passes if the expected interrupt bits are set in the mask and status registers after the DMA operation completes.",
        "Validation / Acceptance Criteria": "After enabling interrupts and executing the DMA debug command, the interrupt_mask register is read to verify the GDMA interrupt bit is asserted. The DMAC masked interrupt status register is read to confirm the DMAC interrupt fired. Interrupts are then cleared by writing to the DMAC interrupt clear register and the interrupt_raw register. The test passes if the expected interrupt bits are correctly set in the mask and status registers following the DMA operation.",
        "Remarks": "Six DMAC register macros (DMAC interrupt enable, DMAC debug instruction 0, DMAC debug instruction 1, DMAC debug command, DMAC masked interrupt status, DMAC interrupt clear) could not be mapped to canonical register names because the DMAC register specification document was not provided. Only HOST and SUBSYS registers were resolved via the supplied specification documents."
    },
    {
        "Index": "2",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_dbi_random_payload_test",
        "Feature": "NA",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs MIPI DSI DBI random payload validation. It configures the HOST block by writing to the PHY interface register (MIZAR_MIPI_DSI_HOST_PHY_IF_CFG), packet handler register (MIZAR_MIPI_DSI_HOST_PCKHDL_CFG), and clock manager register (MIZAR_MIPI_DSI_HOST_CLKMGR_CFG). It disables DPI control via the subsystem DPI control register (MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL). It enables DMAC interrupts by writing to the DMAC interrupt enable register (MIZAR_MIPI_DSI_DMAC_INTEN). It programs the DMAC debug instruction registers (MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1) and executes the debug command (MIZAR_MIPI_DSI_DMAC_DBGCMD) to initiate a DMA transfer with random payload data. It polls the DMAC masked interrupt status register (MIZAR_MIPI_DSI_DMAC_INTMIS) to wait for transfer completion, then clears the DMAC interrupt by writing to the interrupt clear register (MIZAR_MIPI_DSI_DMAC_INTCLR).",
        "Test Description": "This testcase validates MIPI DSI DBI interface operation with random payload data. It configures the PHY interface via PHY_IF_CFG, sets up the packet handler via PCKHDL_CFG, and programs the clock manager through CLKMGR_CFG. DPI control is disabled via the dpi_control register. DMAC interrupts are enabled, and the DMAC debug instruction registers are programmed to initiate a DMA transfer carrying random payload. The test polls the DMAC masked interrupt status register to confirm transfer completion and then clears the DMAC interrupt.",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG to configure the PHY interface. 2. Write to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG with value 0x3d to configure the packet handler. 3. Write to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG with value 0x107 to configure the clock manager. 4. Write to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL with value 0 to disable DPI control. 5. Write to MIZAR_MIPI_DSI_DMAC_INTEN with value 0x3 to enable DMAC interrupts. 6. Write to MIZAR_MIPI_DSI_DMAC_DBGINST0 with DMA instruction opcode to set up DMA channel 0 instruction. 7. Write to MIZAR_MIPI_DSI_DMAC_DBGINST1 with DMA descriptor address for channel 0. 8. Write to MIZAR_MIPI_DSI_DMAC_DBGCMD with value 0x0 to execute the debug command and start the DMA transfer. 9. Poll MIZAR_MIPI_DSI_DMAC_INTMIS in a loop to wait for the DMAC masked interrupt status to indicate transfer completion. 10. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear the DMAC interrupt after transfer completion.",
        "Test Steps / Procedure": "1. Configure the PHY interface by writing to the PHY_IF_CFG register. 2. Configure the packet handler by writing to the PCKHDL_CFG register. 3. Configure the clock manager by writing to the CLKMGR_CFG register. 4. Disable DPI control by writing zero to the dpi_control register. 5. Enable DMAC interrupts by writing to the DMAC interrupt enable register. 6. Program the DMAC debug instruction registers with the DMA channel instruction opcode and descriptor address for random payload transfer. 7. Execute the DMAC debug command to initiate the DMA transfer. 8. Poll the DMAC masked interrupt status register until the transfer completion interrupt is asserted. 9. Clear the DMAC interrupt by writing to the DMAC interrupt clear register.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR",
        "Impacted Registers": "PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control",
        "Meta Validation / Acceptance Criteria": "The test polls MIZAR_MIPI_DSI_DMAC_INTMIS in a loop to check for the DMAC masked interrupt status indicating DMA transfer completion. Once the expected interrupt bit is asserted, the test writes to MIZAR_MIPI_DSI_DMAC_INTCLR to clear the interrupt. The test passes if the DMAC interrupt fires after the random payload DMA transfer completes successfully.",
        "Validation / Acceptance Criteria": "The DMAC masked interrupt status register is polled until the transfer completion interrupt is asserted, confirming the random payload DMA transfer completed. The DMAC interrupt is then cleared via the interrupt clear register. The test passes if the expected interrupt bit is set after the DMA transfer.",
        "Remarks": "Six DMAC register macros (DMAC interrupt enable, DMAC debug instruction 0, DMAC debug instruction 1, DMAC debug command, DMAC masked interrupt status, DMAC interrupt clear) could not be mapped to canonical register names because the DMAC register specification document was not provided. Only HOST and SUBSYS registers were resolved via the supplied specification documents. The DMAC masked interrupt status register is polled in a loop, indicating a blocking wait for DMA completion."
    },
    {
        "Index": "3",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_subsys_reg_wr_rd_test",
        "Feature": "Subsystem Register Write-Read Verification",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs register write-read verification on MIPI DSI subsystem registers. It targets five SUBSYS registers accessed via read-modify-write operations: MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL (base 0xE6A41000, offset 0x0), MIZAR_MIPI_DSI_SUBSYS_LOW_PWR (base 0xE6A41000, offset 0x4), MIZAR_MIPI_DSI_SUBSYS_DBITE (base 0xE6A41000, offset 0x8), MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV (base 0xE6A41000, offset 0xC), and MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW (base 0xE6A41000, offset 0x10). For each register, the test reads the default reset value, verifies it against the expected default, then writes a test value and reads it back to confirm the write was successful. This validates both the reset state and read-write accessibility of each subsystem register.",
        "Test Description": "This testcase validates the register write-read accessibility of MIPI DSI subsystem registers. It verifies the reset default values and write-read functionality of the data_fifo_threshold_val, low_pwr, dbite, dbi_fdiv, and interrupt_raw registers. For each register, the test first reads and checks the default reset value, then writes a known test value and reads it back to confirm correct write-read behavior. This ensures all targeted subsystem registers are accessible and functioning as expected.",
        "Meta Test Steps / Procedure": "1. For each register in the set [MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL, MIZAR_MIPI_DSI_SUBSYS_LOW_PWR, MIZAR_MIPI_DSI_SUBSYS_DBITE, MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW], perform the following steps: 2. Read the register using read_reg() to obtain the current (reset default) value. 3. Compare the read value against the expected default value for that register. If mismatch, flag an error. 4. Write a test data value to the register using write_reg(). 5. Read the register back using read_reg() to obtain the written value. 6. Compare the read-back value (masked with the read mask) against the expected written value (masked with the write mask). If mismatch, flag an error. 7. Perform a soft reset to restore registers to default state between iterations if applicable. 8. After all registers are tested, report the total number of errors and pass/fail status.",
        "Test Steps / Procedure": "1. Iterate through each targeted MIPI DSI subsystem register (data_fifo_threshold_val, low_pwr, dbite, dbi_fdiv, interrupt_raw). 2. Read each register and verify its reset default value matches the expected default. 3. Write a known test value to each register. 4. Read back each register and verify the written value matches the expected value after applying the appropriate read and write masks. 5. Optionally perform a soft reset between register tests to restore default state. 6. Report the overall pass or fail result based on whether all register default checks and write-read checks passed without errors.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL; MIZAR_MIPI_DSI_SUBSYS_LOW_PWR; MIZAR_MIPI_DSI_SUBSYS_DBITE; MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW",
        "Impacted Registers": "data_fifo_threshold_val; low_pwr; dbite; dbi_fdiv; interrupt_raw",
        "Meta Validation / Acceptance Criteria": "For each register in the test set, two validations are performed: (1) The read-back value after reset must match the expected default value for that register. (2) After writing a test value, the read-back value (masked with the register's read mask) must match the written value (masked with the register's write mask). If any comparison fails for any register, the test reports an error for that register. The overall test passes only if all default value checks and all write-read checks pass for all five registers with zero errors.",
        "Validation / Acceptance Criteria": "For each subsystem register (data_fifo_threshold_val, low_pwr, dbite, dbi_fdiv, interrupt_raw): the reset default value read from the register must match the expected default, and after writing a test value, the read-back value must match the written value when appropriate read and write masks are applied. The test passes only if all default value verifications and all write-read verifications succeed with zero errors across all five registers.",
        "Remarks": "All five register macros were successfully resolved to the MIPI DSI SUBSYS register block (base 0xE6A41000) and matched to canonical register names from the subsystem specification document. The actual testcase source files in the repository contained unrelated PCIe code; therefore source-derived fields (Meta Headers, Meta Macros, Meta Arrays, Speed, Mode, Memory Start Offset, Memory End Offset) are marked as NA. The test description and steps are derived from the upstream agent outputs and the register write-read test pattern indicated by the testcase name and read_modify_write operations."
    }
]

# ============================================================
# SHEET COLUMN DEFINITIONS
# ============================================================
testplan_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

metadata_columns = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

# ============================================================
# GENERATE IST TIMESTAMP
# ============================================================
ist = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(ist)
timestamp_str = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"MIPI_DSI_TestPlan_{timestamp_str}.xlsx"
output_dir = os.path.dirname(os.path.abspath(__file__))
filepath = os.path.join(output_dir, filename)

print(f"Generating: {filename}")
print(f"IST Time: {now_ist.strftime('%Y-%m-%d %H:%M:%S %Z')}")

# ============================================================
# CREATE WORKBOOK
# ============================================================
wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = "TestPlan"

# --- MetaData Sheet ---
ws_md = wb.create_sheet(title="MetaData")

# ============================================================
# FORMATTING STYLES
# ============================================================
header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_alignment = Alignment(wrap_text=True, vertical="top")

# ============================================================
# POPULATE TESTPLAN SHEET
# ============================================================
# Headers
for col_idx, col_name in enumerate(testplan_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# Data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(testplan_columns, 1):
        value = row_data.get(col_name, "")
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# ============================================================
# POPULATE METADATA SHEET
# ============================================================
# Headers
for col_idx, col_name in enumerate(metadata_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# Data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(metadata_columns, 1):
        value = row_data.get(col_name, "")
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# ============================================================
# AUTO-SIZE COLUMNS
# ============================================================
def auto_size_columns(ws, columns, max_width=60):
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(col_name)
        for row in range(2, ws.max_row + 1):
            cell_val = ws.cell(row=row, column=col_idx).value
            if cell_val:
                lines = str(cell_val).split("\n")
                for line in lines:
                    max_len = max(max_len, len(line))
        adjusted_width = min(max_len + 2, max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

auto_size_columns(ws_tp, testplan_columns)
auto_size_columns(ws_md, metadata_columns)

# ============================================================
# FREEZE FIRST ROW
# ============================================================
ws_tp.freeze_panes = "A2"
ws_md.freeze_panes = "A2"

# ============================================================
# SET METADATA SHEET TO VERY HIDDEN
# ============================================================
ws_md.sheet_state = "veryHidden"

# ============================================================
# SAVE WORKBOOK
# ============================================================
wb.save(filepath)
print(f"Workbook saved: {filepath}")

# ============================================================
# POST-SAVE VALIDATION
# ============================================================
if not os.path.exists(filepath):
    print("VALIDATION FAILED: File does not exist")
    sys.exit(1)

file_size = os.path.getsize(filepath)
if file_size == 0:
    print("VALIDATION FAILED: File size is 0")
    sys.exit(1)

print(f"File size: {file_size} bytes")

# Verify workbook can be reopened
try:
    wb_verify = load_workbook(filepath)
    sheets = wb_verify.sheetnames
    print(f"Sheets found: {sheets}")
    if "TestPlan" not in sheets:
        print("VALIDATION FAILED: TestPlan sheet missing")
        sys.exit(1)
    if "MetaData" not in sheets:
        print("VALIDATION FAILED: MetaData sheet missing")
        sys.exit(1)
    tp_rows = wb_verify["TestPlan"].max_row - 1
    md_rows = wb_verify["MetaData"].max_row - 1
    print(f"TestPlan rows: {tp_rows}")
    print(f"MetaData rows: {md_rows}")
    print(f"MetaData sheet state: {wb_verify['MetaData'].sheet_state}")
    wb_verify.close()
except Exception as e:
    print(f"VALIDATION FAILED: {e}")
    sys.exit(1)

print("")
print("=" * 50)
print("VALIDATION: PASSED")
print(f"FILENAME: {filename}")
print(f"FILEPATH: {filepath}")
print(f"ROWS_TESTPLAN: {tp_rows}")
print(f"ROWS_METADATA: {md_rows}")
print("=" * 50)
