#!/usr/bin/env python3
"""MIPI_CSI TestPlan XLSX Generator - Agent 7 Final
Generates MIPI_CSI_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx with TestPlan and MetaData sheets.
Run: python generate_mipi_csi_testplan_xlsx_final.py
"""
import datetime, os, sys
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    os.system(f'{sys.executable} -m pip install openpyxl')
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment

json_data = [
  {
    "Index": "1",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_dphy_lanes_test",
    "Feature": "DPHY Lane Configuration and Data Transfer",
    "Meta Headers": '<stdio.h>; <stdlib.h>; "test_common.h"; "mipi_csi2.h"',
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000; LS_LE_EN = 1; TOTAL_FRAME = 1; VC_ID = 3; VRES = 3; HRES = 64; DATA_TYPE",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates MIPI CSI-2 DPHY lane configuration by iterating through lane counts from 4 down to 1. It begins by calling csi2_enable_interrupt() which reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear interrupts, then writes ten interrupt mask registers (MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL with 0x0000000f, MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL with 0x00000003, MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY with 0x000f000f, MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE with 0x000f000f, MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED with 0x0000ffff). The virtual channel register MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL is written with vcid_csi2_wrap_reg (derived from VC_ID shifted based on GDMA path selection). MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA is written with 1 to enable control data transfer. Both virtual channel and control data writes are performed twice. snps_phy_init() is called for D-PHY initialization. MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE is polled until it equals 0x1000f. A for loop iterates lane_num from 3 down to 0: MIZAR_MIPI_CSI2_HOST_N_LANES is written with lane_num, 0xa0243ffc is written with (lane_num+1) to trigger CSI-2 sequence. An inner loop iterates cntrl_pkt_cnt = ((VRES*3)+2) times. In each iteration: DMA interrupt enable register (gdma_reg_base+MIPI_CSI2_DMA_INTEN_OFFSET) is written with 0x3. dma_trnsfr_instn_preload() is called with ch0_preload_loc, gdma_reg_base, 0x8000 src_addr, 0xE6001000 dest_addr, 8 trnsfr_size, 0 irq_num. DMAGO_CSI() is called with gdma_reg_base, ch0_pc, 0x0 ch_num. DMA interrupt status (gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET) is polled until bit 0 is set. DMA interrupt clear (gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET) is written with 0x1. gdma_reg_base+0x28 is read. 0xE6001000 is read into csi_ctrl_data. If (csi_ctrl_data & 0x3f) > 0xf, word_count is extracted from bits [21:6], csi_data_size is computed as 8-byte aligned word_count. dma_trnsfr_instn_preload() is called with ch1_preload_loc, gdma_reg_base, 0x0000 src_addr, 0xE6002000 dest_addr, csi_data_size trnsfr_size, 1 irq_num. DMAGO_CSI() is called with gdma_reg_base, ch1_pc, 0x1 ch_num. DMA interrupt status is polled until bit 1 is set. DMA interrupt status is read again. DMA interrupt clear is written with 0x2. finish(0) is called at the end.",
    "Test Description": "This test validates MIPI CSI-2 DPHY lane configuration by iterating through all supported lane counts (4 down to 1). It first enables CSI-2 interrupts by clearing the main interrupt status register and configuring ten interrupt mask registers for PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected events. The virtual channel register is configured with the selected virtual channel ID and control data transfer is enabled. D-PHY initialization is performed and the PHY stop state register is polled until the PHY enters stop state. For each lane count, the N_LANES register is written with the lane number, and a trigger write initiates the CSI-2 sequence. An inner loop processes each control and data packet: DMA interrupts are enabled, a control data DMA transfer is programmed and executed on channel 0, the DMA interrupt status is polled for completion, and the interrupt is cleared. The control data is then read to determine if a CSI data payload transfer is needed. If the data type indicates a long packet, the word count is extracted, the data size is aligned to 8 bytes, and a data DMA transfer is programmed and executed on channel 1 with its own interrupt polling and clearing. The test completes after processing all packets for all lane configurations.",
    "Meta Test Steps / Procedure": "1. Enter test_case() function. 2. Call csi2_enable_interrupt(). 3. Inside csi2_enable_interrupt(): read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN) to clear interrupts, store in rd_data. 4. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, 0x0000000f). 5. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, 0x00000003). 6. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, 0x000f000f). 7. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, 0x000f000f). 8. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, 0x0000ffff). 9. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, 0x0000ffff). 10. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, 0x0000ffff). 11. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, 0x0000ffff). 12. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, 0x0000ffff). 13. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED, 0x0000ffff). 14. Return from csi2_enable_interrupt(). 15-51. [Full procedure as documented]",
    "Test Steps / Procedure": "1. Enable CSI-2 interrupts by reading the main interrupt status register to clear pending interrupts, then writing all ten interrupt mask registers. 2. Configure the virtual channel register with the selected virtual channel ID based on the GDMA path. 3. Enable control data transfer by writing to the control data register. 4. Repeat the virtual channel and control data register writes. 5. Perform D-PHY initialization. 6. Poll the PHY stop state register until the PHY enters stop state. 7. For each lane configuration (4 lanes down to 1 lane): a. Write the N_LANES register. b. Trigger the CSI-2 sequence. c. For each control and data packet: i. Enable DMA interrupts. ii. Program and start control data DMA transfer on channel 0. iii. Poll DMA interrupt status until channel 0 completes. iv. Clear DMA interrupt for channel 0. v. Read control data to determine packet type. vi. If long packet, program and start data DMA transfer on channel 1, poll and clear. 8. Complete the test.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000",
    "Impacted Registers": "INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED; virtual_channel; control_data; PHY_STOPSTATE; N_LANES",
    "Meta Validation / Acceptance Criteria": "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE is polled until rd_data equals 0x1000f. DMA interrupt status register is polled: for channel 0, bit 0 set; for channel 1, bit 1 set. The test calls finish(0) upon successful completion.",
    "Validation / Acceptance Criteria": "1. PHY stop state register must read expected value. 2. DMA interrupt status must indicate channel 0 completion. 3. DMA interrupt status must indicate channel 1 completion for long packets. 4. Control data read must correctly identify packet type. 5. Test must iterate through all four lane configurations. 6. Test completes successfully with pass indication.",
    "Remarks": "The test uses conditional compilation (DPHY_LANES_TEST, GDMA0_PATH/GDMA1_PATH/GDMA2_PATH/GDMA3_PATH, GDMA0_FULL_MEM). Default configuration uses VC_ID=3, VRES=3, HRES=64 with CSI2_RGB888 data type. Two hardcoded hex addresses (0xa0243ffc, 0xE6001000) could not be mapped to named registers. DMA register accesses use runtime base variable (gdma_reg_base=0xE6A00000) with offset macros."
  },
  {
    "Index": "2",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_test_pattern_generator",
    "Feature": "Test Pattern Generator and DMA Data Transfer",
    "Meta Headers": '<stdio.h>; <stdlib.h>; "test_common.h"; "mipi_csi2.h"',
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates the MIPI CSI-2 internal test pattern generator and DMA-based data transfer. It configures virtual channel, enables interrupts, initializes D-PHY, configures DMA address registers, programs DMA transfer, enables pattern generator (VRES=0x10, HRES=0x70140, CONFIG=0xe401), waits, disables pattern generator, polls DMA completion, and calls finish(0).",
    "Test Description": "This test validates the MIPI CSI-2 internal test pattern generator and DMA-based data transfer. It configures the virtual channel register with the selected virtual channel ID and disables control data transfer. CSI-2 and DMA interrupts are enabled by reading the main interrupt status register to clear pending interrupts and writing all ten interrupt mask registers. D-PHY initialization is performed and the PHY stop state register is polled until the PHY enters stop state. The DMA higher-order address registers for channel 0 read and write paths are configured. The clock gating enable register is written to enable fracdiv output to the CSI2 subsystem. A DMA transfer is programmed with the computed frame data size and started on channel 0. The pattern generator is then enabled by writing vertical resolution, horizontal resolution, configuration, and enable registers. After a wait, the pattern generator is disabled. The DMA interrupt status register is polled until the channel 0 transfer completes. The test completes after a final wait period.",
    "Meta Test Steps / Procedure": "1. Enter test_case(). 2-6. Configure GDMA path and vcid. 7. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg). 8. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 0). 9-21. Enable interrupts. 22-24. PHY init and stopstate poll. 25-26. Compute frame size. 27-30. Configure DMA address registers. 31. write_reg(MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4, 0x1). 32-36. Program and start DMA. 37-41. Enable pattern generator. 42-43. Wait and disable PG. 44-46. Poll DMA completion. 47-48. Wait and finish.",
    "Test Steps / Procedure": "1. Configure virtual channel and disable control data transfer. 2. Enable CSI-2 and DMA interrupts. 3. Perform D-PHY initialization. 4. Poll PHY stop state register. 5. Set frame parameters and compute data transfer size. 6. Configure DMA higher-order address registers. 7. Write clock gating enable register. 8. Program and start DMA data transfer on channel 0. 9. Enable pattern generator. 10. Wait, then disable pattern generator. 11. Poll DMA interrupt status until channel 0 completes. 12. Wait and complete the test.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; enableclkgating_csiphy; PHY_STOPSTATE; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE is polled until rd_data equals 0x1000f. DMA interrupt status register is polled: bit 0 must be set. Pattern generator is enabled then disabled. The test calls finish(0).",
    "Validation / Acceptance Criteria": "1. PHY stop state register must read expected value. 2. DMA interrupt status must indicate channel 0 completion. 3. Pattern generator must be successfully enabled and disabled. 4. Test completes successfully with pass indication.",
    "Remarks": "The test uses conditional compilation for GDMA path selection and FPS60 for destination address increment behavior. Pattern generator configured with VRES=0x10, HRES=0x70140, CONFIG=0xe401, valid_bits_per_pixel=24 (RGB888). The write to base+0xf4 is mapped to enableclkgating_csiphy register. DMA register accesses use runtime base variable with offset macros."
  }
]

def generate_xlsx():
    # IST timestamp
    import datetime
    utc_now = datetime.datetime.utcnow()
    ist_now = utc_now + datetime.timedelta(hours=5, minutes=30)
    ts = ist_now.strftime('%Y%m%d_%H%M%S')
    filename = f'MIPI_CSI_TestPlan_{ts}.xlsx'

    wb = Workbook()

    # --- TestPlan Sheet ---
    ws_tp = wb.active
    ws_tp.title = 'TestPlan'
    tp_cols = ['Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
               'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
               'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
               'Code Generation']

    # --- MetaData Sheet ---
    ws_md = wb.create_sheet('MetaData')
    md_cols = ['Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
               'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
               'Meta Headers', 'Meta Macros', 'Meta Arrays']

    # Formatting
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    wrap_align = Alignment(wrap_text=True, vertical='top')

    # Write TestPlan headers
    for col_idx, col_name in enumerate(tp_cols, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_align

    # Write TestPlan data
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(tp_cols, 1):
            val = row_data.get(col_name, '')
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=val if val else '')
            cell.alignment = wrap_align

    # Write MetaData headers
    for col_idx, col_name in enumerate(md_cols, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_align

    # Write MetaData data
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(md_cols, 1):
            val = row_data.get(col_name, '')
            cell = ws_md.cell(row=row_idx, column=col_idx, value=val if val else '')
            cell.alignment = wrap_align

    # Freeze first row
    ws_tp.freeze_panes = 'A2'
    ws_md.freeze_panes = 'A2'

    # Auto-size columns with max width cap
    for ws in [ws_tp, ws_md]:
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 80))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)

    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = 'veryHidden'

    # Save
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)
    wb.save(output_path)
    print(f'Generated: {output_path}')
    print(f'Filename: {filename}')

    # Verify
    wb2 = load_workbook(output_path)
    assert 'TestPlan' in wb2.sheetnames
    assert 'MetaData' in wb2.sheetnames
    assert wb2['MetaData'].sheet_state == 'veryHidden'
    print(f'Validation: PASSED')
    print(f'File size: {os.path.getsize(output_path)} bytes')
    return output_path, filename

if __name__ == '__main__':
    generate_xlsx()
