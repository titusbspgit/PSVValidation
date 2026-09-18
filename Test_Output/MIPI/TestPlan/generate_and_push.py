#!/usr/bin/env python3
"""
Standalone script to generate MIPI_DSI TestPlan Excel and push to GitHub.
Requirements: pip install openpyxl PyGithub
Usage: python generate_and_push.py
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import base64
import os

# ── Configuration ──
IP_NAME = "MIPI_DSI"
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"{IP_NAME}_TestPlan_{timestamp}.xlsx"

GITHUB_OWNER = "titusbspgit"
GITHUB_REPO = "PSVValidation"
GITHUB_BRANCH = "main"
OUTPUT_DIR = "Test_Output/MIPI/TestPlan"

# ── TestPlan columns ──
tp_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers",
    "Validation / Acceptance Criteria", "Code Generation"
]

tp_row = {
    "Index": "1",
    "SS / Module": "MIPI_DSI",
    "Feature": "NA",
    "Test Case Name": "mipi_dsi_basic_test",
    "Test Description": (
        "This testcase configures the MIPI DSI Host, Subsystem, and DMA Controller blocks "
        "for a basic DSI transfer operation. It configures the PHY interface via PHY_IF_CFG, "
        "sets up packet handling via PCKHDL_CFG, configures the clock manager via CLKMGR_CFG, "
        "configures DPI control via dpi_control, enables DMA and subsystem interrupts via "
        "interrupt_enable and DMAC interrupt enable registers, and initiates DMA channel "
        "operations via DMAC debug instruction registers. An interrupt service routine reads "
        "the interrupt_mask and DMAC masked interrupt status to identify interrupt sources, "
        "then clears interrupts by writing to the DMAC interrupt clear register and interrupt_raw."
    ),
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Remarks": (
        "Six DMAC registers (interrupt enable, masked interrupt status, interrupt clear, "
        "debug instruction 0, debug instruction 1, debug command) could not be mapped to "
        "canonical register names because the DMAC (PL330) register specification was not "
        "provided. Only DSI Host and DSI Subsystem register specifications were available. "
        "The source code in the repository folder does not contain the expected MIPI DSI "
        "testcase source; Agent 2, Agent 3, and Agent 4 outputs were used as the primary "
        "source for register-related fields."
    ),
    "Test Steps / Procedure": (
        "1. Enable DMA channel interrupts by writing to the DMAC interrupt enable register.\n"
        "2. Enable GDMA interrupt at the subsystem level by writing to the interrupt_enable register.\n"
        "3. Configure the DSI Host PHY interface parameters by writing to the PHY_IF_CFG register.\n"
        "4. Configure packet handling behavior by writing to the PCKHDL_CFG register.\n"
        "5. Configure the clock manager by writing to the CLKMGR_CFG register.\n"
        "6. Disable DPI control by writing to the dpi_control register.\n"
        "7. Issue DMA channel 0 transfer instructions via the DMAC debug instruction registers and execute the command.\n"
        "8. Issue DMA channel 1 transfer instructions via the DMAC debug instruction registers and execute the command.\n"
        "9. On interrupt, read the interrupt_mask register to identify the subsystem interrupt source.\n"
        "10. If GDMA interrupt is active, read the DMAC masked interrupt status register to identify the DMA channel source.\n"
        "11. Clear the DMA interrupt by writing to the DMAC interrupt clear register.\n"
        "12. Clear the subsystem interrupt by writing to the interrupt_raw register."
    ),
    "Impacted Registers": "interrupt_enable; PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control; interrupt_mask; interrupt_raw",
    "Validation / Acceptance Criteria": (
        "The interrupt service routine must correctly identify the GDMA interrupt source by "
        "reading the interrupt_mask register. The DMAC masked interrupt status must be read "
        "to determine the triggering DMA channel. DMA interrupts must be successfully cleared "
        "via the DMAC interrupt clear register. Subsystem interrupts must be successfully "
        "cleared via the interrupt_raw register. The test passes if the DMA transfer completes "
        "and all interrupts are properly serviced and cleared."
    ),
    "Code Generation": ""
}

# ── MetaData columns ──
md_columns = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

md_row = {
    "Index": "1",
    "Test Case Name": "mipi_dsi_basic_test",
    "Meta Test Description": (
        "This testcase configures the MIPI DSI Host, Subsystem, and DMAC blocks. "
        "It writes to MIZAR_MIPI_DSI_DMAC_INTEN to enable DMA interrupts, writes to "
        "MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE to enable subsystem-level GDMA interrupts, "
        "writes to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG to configure the PHY interface, writes to "
        "MIZAR_MIPI_DSI_HOST_PCKHDL_CFG to configure packet handling, writes to "
        "MIZAR_MIPI_DSI_HOST_CLKMGR_CFG to configure the clock manager, writes to "
        "MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to configure DPI control, writes to "
        "MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1, and "
        "MIZAR_MIPI_DSI_DMAC_DBGCMD to issue DMA channel instructions via the debug interface. "
        "In the interrupt service routine, it reads MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK and "
        "MIZAR_MIPI_DSI_DMAC_INTMIS to identify interrupt sources, then writes to "
        "MIZAR_MIPI_DSI_DMAC_INTCLR and MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW to clear the interrupts."
    ),
    "Meta Test Steps / Procedure": (
        "1. Write to MIZAR_MIPI_DSI_DMAC_INTEN with value 0x3 to enable DMA channel interrupts.\n"
        "2. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE with MIPI_DSI_SUBSYS_INTERRUPT_ENABLE_GDMA_INTR to enable GDMA interrupt at subsystem level.\n"
        "3. Configure PHY interface by computing phy_if_cfg using set_data_mask with MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME and writing to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG.\n"
        "4. Write 0x3d to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG to configure packet handling.\n"
        "5. Write 0x107 to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG to configure clock manager.\n"
        "6. Write 0 to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to configure DPI control.\n"
        "7. Write DMA channel 0 descriptor address to MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1, and execute via MIZAR_MIPI_DSI_DMAC_DBGCMD.\n"
        "8. Write DMA channel 1 descriptor address to MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1, and execute via MIZAR_MIPI_DSI_DMAC_DBGCMD.\n"
        "9. In ISR (Default_IRQHandler): read MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK to get subsystem interrupt mask status.\n"
        "10. If GDMA interrupt is indicated, read MIZAR_MIPI_DSI_DMAC_INTMIS to get DMA channel masked interrupt status.\n"
        "11. Write channel mask to MIZAR_MIPI_DSI_DMAC_INTCLR to clear DMA interrupt.\n"
        "12. Write subsystem mask to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW to clear subsystem interrupt."
    ),
    "Meta Impacted Registers": (
        "MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE; "
        "MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; "
        "MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; "
        "MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; "
        "MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK; "
        "MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR; "
        "MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW"
    ),
    "Meta Validation / Acceptance Criteria": (
        "The ISR reads MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK and checks if the value matches "
        "MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR to confirm the GDMA interrupt source. It then "
        "reads MIZAR_MIPI_DSI_DMAC_INTMIS to identify which DMA channel triggered the interrupt. "
        "Interrupts are cleared by writing to MIZAR_MIPI_DSI_DMAC_INTCLR and "
        "MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW. Successful completion implies that the DMA transfer "
        "completes and the interrupt is serviced and cleared without errors."
    ),
    "Meta Headers": "NA",
    "Meta Macros": "NA",
    "Meta Arrays": "NA"
}

# ── Formatting constants ──
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_alignment = Alignment(wrap_text=True, vertical="top")

def populate_sheet(ws, columns, data):
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    for col_idx, col_name in enumerate(columns, 1):
        val = data.get(col_name, "")
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.alignment = wrap_alignment
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(col_name)
        val = str(data.get(col_name, ""))
        for line in val.split("\n"):
            if len(line) > max_len:
                max_len = len(line)
        width = min(max_len + 4, 60)
        if width < 12:
            width = 12
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

# ── Build workbook ──
wb = openpyxl.Workbook()
ws_tp = wb.active
ws_tp.title = "TestPlan"
ws_md = wb.create_sheet("MetaData")

populate_sheet(ws_tp, tp_columns, tp_row)
populate_sheet(ws_md, md_columns, md_row)

ws_md.sheet_state = "veryHidden"

wb.save(filename)

# ── Verify ──
file_size = os.path.getsize(filename)
wb2 = openpyxl.load_workbook(filename)
sheets = wb2.sheetnames
assert "TestPlan" in sheets
assert "MetaData" in sheets
assert file_size > 0

print(f"File: {filename}")
print(f"Size: {file_size} bytes")
print(f"Sheets: {sheets}")
print(f"Validation: PASSED")

# ── Push to GitHub ──
try:
    from github import Github
    token = os.environ.get("GITHUB_TOKEN", "")
    if token:
        g = Github(token)
        repo = g.get_repo(f"{GITHUB_OWNER}/{GITHUB_REPO}")
        with open(filename, "rb") as f:
            content = f.read()
        path = f"{OUTPUT_DIR}/{filename}"
        result = repo.create_file(
            path=path,
            message="Added generated TestPlan Excel",
            content=content,
            branch=GITHUB_BRANCH
        )
        print(f"GitHub URL: {result['content'].html_url}")
        print(f"Commit SHA: {result['commit'].sha}")
    else:
        print("No GITHUB_TOKEN found. Upload manually.")
except ImportError:
    print("PyGithub not installed. Upload the file manually.")
    print(f"Target: {GITHUB_OWNER}/{GITHUB_REPO}/{OUTPUT_DIR}/{filename}")
