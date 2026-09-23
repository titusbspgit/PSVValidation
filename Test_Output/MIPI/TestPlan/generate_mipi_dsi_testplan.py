#!/usr/bin/env python3
"""MIPI_DSI TestPlan Excel Generator - Agent 7 Fallback Automation"""
import json
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

JSON_DATA = [
    {
        "Index": "1",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_basic_test",
        "Feature": "NA",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a basic MIPI DSI data transfer operation using the DMAC and DSI host/subsystem. It writes to the DMAC interrupt enable register (MIZAR_MIPI_DSI_DMAC_INTEN) to enable interrupts, configures the subsystem interrupt enable register (MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE), sets up the DSI host PHY interface configuration (MIZAR_MIPI_DSI_HOST_PHY_IF_CFG), packet handler configuration (MIZAR_MIPI_DSI_HOST_PCKHDL_CFG), and clock manager configuration (MIZAR_MIPI_DSI_HOST_CLKMGR_CFG). It disables DPI control via MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL. The DMAC is programmed using debug instruction registers (MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1) and executed via the debug command register (MIZAR_MIPI_DSI_DMAC_DBGCMD). The test reads the subsystem interrupt mask register (MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK) and DMAC masked interrupt status register (MIZAR_MIPI_DSI_DMAC_INTMIS) to check interrupt status. It clears DMAC interrupts via MIZAR_MIPI_DSI_DMAC_INTCLR and subsystem raw interrupts via MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW.",
        "Test Description": "This testcase validates a basic MIPI DSI data transfer using the DMA controller and DSI host/subsystem. It enables DMAC and subsystem interrupts, configures the DSI host PHY interface (PHY_IF_CFG), packet handler (PCKHDL_CFG), and clock manager (CLKMGR_CFG). It disables DPI control (dpi_control). The DMAC is programmed via debug instruction registers and executed via the debug command register. The test reads the subsystem interrupt mask (interrupt_mask) and DMAC masked interrupt status to verify interrupt assertion, then clears DMAC interrupts and subsystem raw interrupts (interrupt_raw).",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_DSI_DMAC_INTEN to enable DMAC interrupts. 2. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE to enable subsystem GDMA interrupt. 3. Write to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG to configure PHY interface with stop wait time. 4. Write to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG to configure packet handler. 5. Write to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG to configure clock manager. 6. Write 0 to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to disable DPI control. 7. Write to MIZAR_MIPI_DSI_DMAC_DBGINST0 with DMA debug instruction 0. 8. Write to MIZAR_MIPI_DSI_DMAC_DBGINST1 with DMA channel descriptor address. 9. Write to MIZAR_MIPI_DSI_DMAC_DBGCMD to execute the DMA debug instruction. 10. Read MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK to check subsystem interrupt mask status. 11. Read MIZAR_MIPI_DSI_DMAC_INTMIS to check DMAC masked interrupt status. 12. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear DMAC interrupt. 13. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW to clear subsystem raw interrupt.",
        "Test Steps / Procedure": "1. Enable DMAC interrupts by writing to the DMAC interrupt enable register. 2. Enable subsystem GDMA interrupt by writing to the interrupt_enable register. 3. Configure the DSI host PHY interface by writing to PHY_IF_CFG with the desired stop wait time. 4. Configure the DSI host packet handler by writing to PCKHDL_CFG. 5. Configure the DSI host clock manager by writing to CLKMGR_CFG. 6. Disable DPI control by writing zero to dpi_control. 7. Program the DMAC debug instruction registers with the DMA transfer instruction and channel descriptor address. 8. Execute the DMAC debug instruction by writing to the DMAC debug command register. 9. Read the interrupt_mask register to verify the subsystem interrupt mask status. 10. Read the DMAC masked interrupt status register to verify DMAC interrupt assertion. 11. Clear the DMAC interrupt by writing to the DMAC interrupt clear register. 12. Clear the subsystem raw interrupt by writing to the interrupt_raw register.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE; MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW",
        "Impacted Registers": "interrupt_enable; PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control; interrupt_mask; interrupt_raw",
        "Meta Validation / Acceptance Criteria": "After DMAC transfer is initiated, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK is read and checked for the GDMA interrupt bit (MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR). MIZAR_MIPI_DSI_DMAC_INTMIS is read to verify DMAC masked interrupt status is asserted. Interrupts are then cleared by writing to MIZAR_MIPI_DSI_DMAC_INTCLR and MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW. The test passes if the expected interrupt bits are set after the DMA transfer and are successfully cleared.",
        "Validation / Acceptance Criteria": "After the DMAC transfer is initiated, the interrupt_mask register is read and verified for the GDMA interrupt bit assertion. The DMAC masked interrupt status register is read to confirm the DMAC interrupt is asserted. Interrupts are then cleared by writing to the DMAC interrupt clear register and the interrupt_raw register. The test passes if the expected interrupt bits are set after the DMA transfer and are successfully cleared.",
        "Remarks": "Six DMAC-domain register macros could not be mapped to canonical register names due to the absence of a DMAC register specification document. The source files found in the testcase folder did not contain MIPI DSI code; the testcase details are derived from upstream Agent 2, Agent 3, and Agent 4 outputs."
    },
    {
        "Index": "2",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_dbi_random_payload_test",
        "Feature": "NA",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a MIPI DSI DBI random payload transfer using the DMAC and DSI host/subsystem. It configures the DSI host PHY interface (MIZAR_MIPI_DSI_HOST_PHY_IF_CFG), packet handler (MIZAR_MIPI_DSI_HOST_PCKHDL_CFG), and clock manager (MIZAR_MIPI_DSI_HOST_CLKMGR_CFG). It disables DPI control via MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL. DMAC interrupts are enabled by writing to MIZAR_MIPI_DSI_DMAC_INTEN. The DMAC is programmed using debug instruction registers (MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1) and executed via the debug command register (MIZAR_MIPI_DSI_DMAC_DBGCMD). The test polls the DMAC masked interrupt status register (MIZAR_MIPI_DSI_DMAC_INTMIS) to detect transfer completion, then clears the DMAC interrupt by writing to MIZAR_MIPI_DSI_DMAC_INTCLR.",
        "Test Description": "This testcase validates a MIPI DSI DBI random payload data transfer using the DMA controller and DSI host/subsystem. It configures the DSI host PHY interface (PHY_IF_CFG), packet handler (PCKHDL_CFG), and clock manager (CLKMGR_CFG). DPI control (dpi_control) is disabled. DMAC interrupts are enabled, and the DMAC is programmed via debug instruction registers and executed via the debug command register. The test polls the DMAC masked interrupt status register to detect transfer completion, then clears the DMAC interrupt.",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG to configure the PHY interface with stop wait time. 2. Write to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG to configure the packet handler. 3. Write to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG to configure the clock manager. 4. Write 0 to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to disable DPI control. 5. Write to MIZAR_MIPI_DSI_DMAC_INTEN to enable DMAC interrupts. 6. Write to MIZAR_MIPI_DSI_DMAC_DBGINST0 with DMA debug instruction 0. 7. Write to MIZAR_MIPI_DSI_DMAC_DBGINST1 with DMA channel descriptor address. 8. Write to MIZAR_MIPI_DSI_DMAC_DBGCMD to execute the DMA debug instruction. 9. Read/poll MIZAR_MIPI_DSI_DMAC_INTMIS to check DMAC masked interrupt status for transfer completion. 10. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear the DMAC interrupt.",
        "Test Steps / Procedure": "1. Configure the DSI host PHY interface by writing to PHY_IF_CFG with the desired stop wait time. 2. Configure the DSI host packet handler by writing to PCKHDL_CFG. 3. Configure the DSI host clock manager by writing to CLKMGR_CFG. 4. Disable DPI control by writing zero to dpi_control. 5. Enable DMAC interrupts by writing to the DMAC interrupt enable register. 6. Program the DMAC debug instruction registers with the DMA transfer instruction and channel descriptor address. 7. Execute the DMAC debug instruction by writing to the DMAC debug command register. 8. Poll the DMAC masked interrupt status register until the transfer completion interrupt is asserted. 9. Clear the DMAC interrupt by writing to the DMAC interrupt clear register.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR",
        "Impacted Registers": "PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control",
        "Meta Validation / Acceptance Criteria": "MIZAR_MIPI_DSI_DMAC_INTMIS is polled to detect the DMAC transfer completion interrupt. Once the expected interrupt bit is asserted, the interrupt is cleared by writing to MIZAR_MIPI_DSI_DMAC_INTCLR. The test passes if the DMAC masked interrupt status indicates successful transfer completion and the interrupt is successfully cleared.",
        "Validation / Acceptance Criteria": "The DMAC masked interrupt status register is polled until the transfer completion interrupt is asserted. The DMAC interrupt is then cleared by writing to the DMAC interrupt clear register. The test passes if the expected DMAC interrupt bit is set after the DMA transfer and is successfully cleared.",
        "Remarks": "Six DMAC-domain register macros could not be mapped to canonical register names due to the absence of a DMAC register specification document. The source files found in the testcase folder did not contain MIPI DSI code; the testcase details are derived from upstream Agent 2, Agent 3, and Agent 4 outputs."
    },
    {
        "Index": "3",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_subsys_reg_wr_rd_test",
        "Feature": "Register Write-Read Verification",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs register write-read verification on a set of MIPI DSI subsystem registers. It targets five subsystem registers via their address macros: MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL (base 0xE6A41000, offset 0x0), MIZAR_MIPI_DSI_SUBSYS_LOW_PWR (base 0xE6A41000, offset 0x4), MIZAR_MIPI_DSI_SUBSYS_DBITE (base 0xE6A41000, offset 0x8), MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV (base 0xE6A41000, offset 0xC), and MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW (base 0xE6A41000, offset 0x10). Each register undergoes a read-modify-write operation where a value is written and then read back to verify correctness.",
        "Test Description": "This testcase validates the write and read-back functionality of five MIPI DSI subsystem registers: data_fifo_threshold_val, low_pwr, dbite, dbi_fdiv, and interrupt_raw. For each register, a value is written and then read back to confirm the register correctly stores the written data. This verifies basic register accessibility and data integrity within the MIPI DSI subsystem address space.",
        "Meta Test Steps / Procedure": "1. For each register address in the address array (MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL, MIZAR_MIPI_DSI_SUBSYS_LOW_PWR, MIZAR_MIPI_DSI_SUBSYS_DBITE, MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW), perform a read-modify-write operation. 2. Write a test data value to the register using write_reg(addr_array[i], data_wr). 3. Read back the register value using read_reg(addr_array[i]). 4. Compare the read-back value against the expected value using the corresponding write mask and read mask from the mask arrays. 5. Repeat for all five registers in the array. 6. Report pass or fail based on comparison results.",
        "Test Steps / Procedure": "1. Iterate through the list of five target MIPI DSI subsystem registers. 2. Write a test data value to the data_fifo_threshold_val register and read it back to verify. 3. Write a test data value to the low_pwr register and read it back to verify. 4. Write a test data value to the dbite register and read it back to verify. 5. Write a test data value to the dbi_fdiv register and read it back to verify. 6. Write a test data value to the interrupt_raw register and read it back to verify. 7. For each register, compare the read-back value against the expected value using the applicable write and read masks. 8. Report the overall pass or fail result based on all register comparisons.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL; MIZAR_MIPI_DSI_SUBSYS_LOW_PWR; MIZAR_MIPI_DSI_SUBSYS_DBITE; MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW",
        "Impacted Registers": "data_fifo_threshold_val; low_pwr; dbite; dbi_fdiv; interrupt_raw",
        "Meta Validation / Acceptance Criteria": "For each register in the address array, after performing write_reg followed by read_reg, the read-back value is masked with the corresponding read mask and compared against the written value masked with the corresponding write mask. If all five registers (MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL, MIZAR_MIPI_DSI_SUBSYS_LOW_PWR, MIZAR_MIPI_DSI_SUBSYS_DBITE, MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW) return the expected values after write and read-back, the test passes. Any mismatch results in a test failure.",
        "Validation / Acceptance Criteria": "For each of the five subsystem registers (data_fifo_threshold_val, low_pwr, dbite, dbi_fdiv, interrupt_raw), the value read back after a write operation must match the expected value after applying the applicable write and read masks. The test passes if all five register write-read comparisons succeed. Any mismatch in any register results in test failure.",
        "Remarks": "All five register macros belong to the MIPI DSI subsystem domain (base address space) and were successfully mapped to canonical register names from the subsystem specification document. The source files found in the testcase folder did not contain MIPI DSI code; the testcase details are derived from upstream Agent 2, Agent 3, and Agent 4 outputs."
    }
]

TESTPLAN_COLUMNS = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

METADATA_COLUMNS = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

def generate_workbook():
    IST = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(IST)
    timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
    filename = f"MIPI_DSI_TestPlan_{timestamp}.xlsx"
    
    wb = Workbook()
    
    # --- TestPlan Sheet ---
    ws_tp = wb.active
    ws_tp.title = "TestPlan"
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    
    # Write TestPlan headers
    for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    
    # Write TestPlan data
    for row_idx, row_data in enumerate(JSON_DATA, 2):
        for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
            value = row_data.get(col_name, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
    
    # Freeze first row
    ws_tp.freeze_panes = "A2"
    
    # Auto-size columns
    for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(JSON_DATA) + 2):
            cell_val = ws_tp.cell(row=row_idx, column=col_idx).value
            if cell_val:
                max_len = max(max_len, min(len(str(cell_val)), 80))
        ws_tp.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)
    
    # --- MetaData Sheet ---
    ws_md = wb.create_sheet(title="MetaData")
    
    # Write MetaData headers
    for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    
    # Write MetaData data
    for row_idx, row_data in enumerate(JSON_DATA, 2):
        for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
            value = row_data.get(col_name, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
    
    # Freeze first row
    ws_md.freeze_panes = "A2"
    
    # Auto-size columns
    for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(JSON_DATA) + 2):
            cell_val = ws_md.cell(row=row_idx, column=col_idx).value
            if cell_val:
                max_len = max(max_len, min(len(str(cell_val)), 80))
        ws_md.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)
    
    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = "veryHidden"
    
    # Save workbook
    script_dir = os.path.dirname(os.path.abspath(__file__))
    filepath = os.path.join(script_dir, filename)
    wb.save(filepath)
    
    # Validate
    assert os.path.exists(filepath), f"File not found: {filepath}"
    assert os.path.getsize(filepath) > 0, f"File is empty: {filepath}"
    
    # Re-open to validate
    wb_check = load_workbook(filepath)
    assert "TestPlan" in wb_check.sheetnames, "TestPlan sheet missing"
    assert "MetaData" in wb_check.sheetnames, "MetaData sheet missing"
    tp_rows = wb_check["TestPlan"].max_row - 1
    md_rows = wb_check["MetaData"].max_row - 1
    assert tp_rows == 3, f"Expected 3 TestPlan rows, got {tp_rows}"
    assert md_rows == 3, f"Expected 3 MetaData rows, got {md_rows}"
    wb_check.close()
    
    print(f"SUCCESS: {filename}")
    print(f"FILEPATH: {filepath}")
    print(f"SIZE: {os.path.getsize(filepath)}")
    print(f"ROWS_TESTPLAN: {tp_rows}")
    print(f"ROWS_METADATA: {md_rows}")
    print(f"VALIDATION: PASSED")
    
    return filepath, filename

if __name__ == "__main__":
    generate_workbook()
