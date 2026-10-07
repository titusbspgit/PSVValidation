#!/usr/bin/env python3
"""Agent 7 - MIPI_CSI TestPlan XLSX Generator (Retry Run)
Generates a real openpyxl workbook with TestPlan and MetaData sheets.
"""
import os, sys, json, base64
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment

# ── IST timestamp ──
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp_str = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"MIPI_CSI_TestPlan_{timestamp_str}.xlsx"
output_dir = os.environ.get("OUTPUT_DIR", ".")
filepath = os.path.join(output_dir, filename)

# ── JSON Data ──
json_data = [
  {
    "Index": "1",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_dphy_lanes_test",
    "Feature": "DPHY Lane Configuration and CSI-2 Data Transfer",
    "Meta Headers": '<stdio.h>; <stdlib.h>; "test_common.h"; "mipi_csi2.h"',
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2; GDMA_CTRL_DATA_DEST_ADDR2; LS_LE_EN; TOTAL_FRAME; VC_ID; VRES; HRES; DATA_TYPE",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates MIPI CSI-2 DPHY lane configuration by iterating through all lane counts (4 down to 1). It first calls csi2_enable_interrupt() which reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts, then writes all CSI-2 host interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED) to enable various interrupt categories. The test then writes MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with the computed virtual channel ID and writes MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA with 1 to enable control data transfer. It calls snps_phy_init() for D-PHY initialization, then polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until the value equals 0x1000f indicating all lanes have entered stop state. For each lane count (3 down to 0), it writes MIZAR_MIPI_CSI2_HOST_N_LANES with the lane number, writes 0xa0243ffc with (lane_num+1) to trigger the CSI-2 sequence, then enters a packet processing loop. Inside the loop, it enables DMA interrupts via gdma_reg_base+MIPI_CSI2_DMA_INTEN_OFFSET, calls dma_trnsfr_instn_preload() and DMAGO_CSI() for control data transfer on channel 0, polls gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET for DMA completion (bit 0), clears the interrupt via gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET, reads 0xE6001000 for CSI control data, and if the data type field indicates image data (>0xf), extracts word_count, calculates csi_data_size with 8-byte alignment, programs DMA channel 1 for data transfer via dma_trnsfr_instn_preload() and DMAGO_CSI(), polls for DMA completion (bit 1), and clears the interrupt. The test calls finish(0) upon completion.",
    "Test Description": "This test validates MIPI CSI-2 DPHY lane configuration by iterating through all supported lane counts from 4 lanes down to 1 lane. For each lane configuration, the test enables all CSI-2 host interrupt masks, configures the virtual channel, enables control data transfer, initializes the D-PHY, and waits for all lanes to enter stop state. It then configures the number of active lanes via the N_LANES register and triggers a CSI-2 data sequence. For each received packet, the test uses DMA to transfer control data, reads the CSI control packet to determine the data type, and if image data is detected, calculates the transfer size with 8-byte alignment and performs a DMA data transfer. The test validates DMA completion for both control and data channels by polling interrupt status registers and clearing interrupts after each transfer.",
    "Meta Test Steps / Procedure": "1. Call csi2_enable_interrupt() which performs: read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN) to clear pending interrupts; write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, 0x0000000f); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, 0x00000003); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, 0x000f000f); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, 0x000f000f); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, 0x0000ffff); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, 0x0000ffff); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, 0x0000ffff); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, 0x0000ffff); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, 0x0000ffff); write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED, 0x0000ffff). 2. Compute vcid_csi2_wrap_reg based on GDMA path conditional compilation (GDMA3_PATH: VC_ID, GDMA2_PATH: VC_ID<<4, GDMA1_PATH: VC_ID<<8, GDMA0_PATH: VC_ID<<12). Set gdma_reg_base = 0xE6A00000. 3. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg). 4. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 1). 5. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg) (repeated). 6. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 1) (repeated). 7. Call snps_phy_init(). 8. rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE). 9. Poll: while(!(rd_data == 0x1000f)) { rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE); }. 10. Set ch0_pc = 0xE6000000, ch1_pc = 0xE6000500. 11. Loop for lane_num from 3 down to 0: write_reg(MIZAR_MIPI_CSI2_HOST_N_LANES, lane_num). 12. Compute cntrl_pkt_cnt = ((VRES*3) + 2). 13. write_reg(0xa0243ffc, (lane_num+1)). 14. Inner loop for i from 0 to cntrl_pkt_cnt-1: set ch0_preload_loc = ch0_pc, ch1_preload_loc = ch1_pc. 15. write_reg(gdma_reg_base+MIPI_CSI2_DMA_INTEN_OFFSET, 0x3). 16. dma_trnsfr_instn_preload(ch0_preload_loc, gdma_reg_base, 0x8000, 0xE6001000, 8, 0). 17. DMAGO_CSI(gdma_reg_base, ch0_pc, 0x0). 18. Poll: rd_data = 0; while((rd_data & 0x1) == 0x0) { rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET); }. 19. write_reg(gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET, 0x1). 20. rd_data = read_reg(gdma_reg_base + 0x28). 21. csi_ctrl_data = read_reg(0xE6001000). 22. If ((csi_ctrl_data & 0x3f) > 0xf): word_count = ((csi_ctrl_data >> 6) & 0xffff); csi_data_size = (word_count%8) ? (word_count/8 + 1)*8 : word_count. 23. dma_trnsfr_instn_preload(ch1_preload_loc, gdma_reg_base, 0x0000, 0xE6002000, csi_data_size, 1). 24. DMAGO_CSI(gdma_reg_base, ch1_pc, 0x1). 25. Poll: rd_data = 0; while((rd_data & 0x2) == 0x0) { rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET); }. 26. rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET). 27. write_reg(gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET, 0x2). 28. Call finish(0).",
    "Test Steps / Procedure": "1. Enable all CSI-2 host interrupts by reading the main interrupt status register to clear pending interrupts, then writing interrupt mask values to INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED registers. 2. Configure the virtual channel ID in the virtual_channel register based on the selected GDMA path. 3. Enable control data transfer by writing to the control_data register. 4. Initialize the D-PHY by calling the PHY initialization sequence. 5. Poll the PHY_STOPSTATE register until all lanes report stop state. 6. Iterate through lane counts from 4 down to 1 by writing the lane number to the N_LANES register. 7. Trigger the CSI-2 sequence for the current lane configuration. 8. For each expected packet, enable DMA interrupts, program DMA channel 0 for control data transfer, and start the DMA transfer. 9. Poll the DMA interrupt status register for channel 0 completion, then clear the DMA interrupt. 10. Read the CSI control data from the destination buffer to determine the data type and word count. 11. If the data type indicates image data, calculate the aligned transfer size, program DMA channel 1 for data transfer, and start the DMA transfer. 12. Poll the DMA interrupt status register for channel 1 completion, then clear the DMA interrupt. 13. Repeat steps 6 through 12 for all lane configurations. 14. Signal test completion.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Impacted Registers": "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "1. PHY stop state validation: read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) is polled until rd_data equals 0x1000f, confirming all DPHY lanes have entered stop state. 2. DMA channel 0 completion: read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) is polled until (rd_data & 0x1) != 0x0, confirming control data DMA transfer completion. 3. CSI control data type check: csi_ctrl_data = read_reg(0xE6001000), then (csi_ctrl_data & 0x3f) > 0xf determines if the packet contains image data requiring data channel DMA transfer. 4. DMA channel 1 completion: read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) is polled until (rd_data & 0x2) != 0x0, confirming data DMA transfer completion. 5. Word count extraction: word_count = ((csi_ctrl_data >> 6) & 0xffff) and csi_data_size is aligned to 8-byte boundary. 6. Test completion: finish(0) is called indicating successful test execution.",
    "Validation / Acceptance Criteria": "1. The PHY_STOPSTATE register must report all lanes in stop state before lane configuration begins. 2. DMA control data transfer on channel 0 must complete successfully as indicated by the DMA interrupt status register. 3. The CSI control data read from the destination buffer must contain a valid data type field to determine whether image data transfer is required. 4. When image data is present, DMA data transfer on channel 1 must complete successfully as indicated by the DMA interrupt status register. 5. The test must iterate through all four lane configurations (4, 3, 2, and 1 lane) and process all expected packets for each configuration. 6. The test must signal successful completion for all lane configurations.",
    "Remarks": "The test iterates lane counts from 4 down to 1 (lane_num 3 to 0) to validate DPHY operation across all supported lane configurations. DMA transfers use two channels: channel 0 for CSI-2 control data and channel 1 for image data. The virtual channel register write and control data enable write are performed twice in the source code. Two hardcoded hex addresses (used for triggering CSI-2 sequence and reading DMA fault status) could not be mapped to named registers in the specification. The GDMA path selection is compile-time conditional. DMA interrupt polling uses bit masking to distinguish channel 0 (bit 0) and channel 1 (bit 1) completion."
  },
  {
    "Index": "2",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_test_pattern_generator",
    "Feature": "Test Pattern Generator and CSI-2 Data Transfer",
    "Meta Headers": '<stdio.h>; <stdlib.h>; "test_common.h"; "mipi_csi2.h"',
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2; GDMA_CTRL_DATA_DEST_ADDR2",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates the MIPI CSI-2 internal test pattern generator functionality. The test_case() function sets vcid=3, computes vcid_unselected_path=((vcid+1)&0xf), then computes vcid_csi2_wrap_reg based on GDMA path conditional compilation (GDMA3_PATH, GDMA2_PATH, GDMA1_PATH, or GDMA0_PATH). It sets gdma_reg_base = 0xE6A00000 + ((gdma_path)*0x1000). It writes MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with vcid_csi2_wrap_reg and writes MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA with 0 to disable control data transfer. It calls csi2_subsys_enable_interrupt() and snps_phy_init(). It polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until rd_data equals 0x1000f. It sets hres=320, vres=16, valid_bits_per_pixel=24 and computes csi2_data_trnsfr_size with 8-byte alignment. It writes DMA address registers, enables fracdiv output, programs DMA transfer, enables the test pattern generator, waits, disables it, polls for DMA completion, and calls finish(0).",
    "Test Description": "This test validates the MIPI CSI-2 internal test pattern generator by configuring the pattern generator with a vertical resolution of 16 lines and horizontal resolution of 320 pixels at 24 bits per pixel. The test configures the virtual channel, disables control data transfer, enables CSI-2 and DMA interrupts, initializes the D-PHY, and waits for all lanes to enter stop state. It then programs the DMA address registers for channel 0 read and write paths, computes the aligned data transfer size, and programs the DMA transfer instructions. The test enables the pattern generator via the PPI_PG_ENABLE register, waits briefly, then disables it. It polls the DMA interrupt status for channel 0 completion to confirm the generated test pattern data was successfully transferred via DMA. The test signals completion after a final wait period.",
    "Meta Test Steps / Procedure": "1. Set int_pend = 1, vcid = 3, vcid_unselected_path = ((vcid + 1) & 0xf). 2. Compute vcid_csi2_wrap_reg based on GDMA path. 3. Set gdma_reg_base = 0xE6A00000 + ((gdma_path)*0x1000). 4. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg). 5. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 0). 6. Call csi2_subsys_enable_interrupt(). 7. Call snps_phy_init(). 8-9. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until 0x1000f. 10. Set hres=320, vres=16, valid_bits_per_pixel=24. 11. Compute csi2_data_trnsfr_size. 12-15. Write DMA address registers. 16. write_reg(MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4, 0x1). 17-19. Program and start DMA. 20. Enable pattern generator. 21. wait_on(100). 22. Disable pattern generator. 23-24. Poll DMA completion. 25. wait_on(10000). 26. finish(0).",
    "Test Steps / Procedure": "1. Configure the virtual channel ID in the virtual_channel register based on the selected GDMA path. 2. Disable control data transfer by writing 0 to the control_data register. 3. Enable CSI-2 and DMA interrupts by reading INT_ST_MAIN to clear pending interrupts, then writing interrupt mask values to all interrupt mask registers. 4. Initialize the D-PHY by calling the PHY initialization sequence. 5. Poll the PHY_STOPSTATE register until all lanes report stop state. 6. Set image parameters: horizontal resolution 320, vertical resolution 16, 24 bits per pixel. 7. Compute the 8-byte aligned data transfer size based on the image parameters. 8. Program the DMA channel 0 higher-order address registers. 9. Enable fracdiv output to the CSI-2 subsystem. 10. Program DMA transfer instructions and start DMA channel 0. 11. Enable the test pattern generator by writing to PPI_PG registers. 12. Wait briefly, then disable the pattern generator. 13. Poll the DMA interrupt status register for channel 0 completion. 14. Wait for a settling period and signal test completion.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "1. PHY stop state validation: read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) is polled until rd_data equals 0x1000f. 2. DMA channel 0 completion: polled until (rd_data & 0x1) != 0. 3. Pattern generator enable/disable sequence validated. 4. Test completion: finish(0) is called.",
    "Validation / Acceptance Criteria": "1. The PHY_STOPSTATE register must report all lanes in stop state before DMA and pattern generator configuration begins. 2. The test pattern generator must be successfully enabled and then disabled via the PPI_PG_ENABLE register. 3. DMA transfer of the generated test pattern data on channel 0 must complete successfully as indicated by the DMA interrupt status register. 4. The test must signal successful completion after the DMA transfer completes and a settling wait period elapses.",
    "Remarks": "The test uses the CSI-2 host internal test pattern generator instead of an external D-PHY source. Control data transfer is explicitly disabled by writing 0 to the control_data register, unlike the dphy_lanes_test which enables it. The pattern generator is configured for 320x16 resolution at 24 bits per pixel with a specific configuration value. The DMA higher-order address registers for channel 0 are programmed for both read and write paths. A write to an inline offset register at base+0xf4 enables fracdiv output to the CSI-2 subsystem."
  }
]

# ── TestPlan columns (14) ──
tp_cols = ["Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
           "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
           "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
           "Code Generation"]

# ── MetaData columns (9) ──
md_cols = ["Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
           "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
           "Meta Headers", "Meta Macros", "Meta Arrays"]

# ── Create workbook ──
wb = Workbook()
ws_tp = wb.active
ws_tp.title = "TestPlan"
ws_md = wb.create_sheet("MetaData")
ws_md.sheet_state = "veryHidden"

# ── Styles ──
hdr_font = Font(bold=True, color="FFFFFF", size=11)
hdr_fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")
wrap_align = Alignment(wrap_text=True, vertical="top")

def write_sheet(ws, columns, data):
    # Header
    for c, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=col_name)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = wrap_align
    # Data rows
    for r, row_data in enumerate(data, 2):
        for c, col_name in enumerate(columns, 1):
            val = row_data.get(col_name, "")
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = wrap_align
    # Freeze top row
    ws.freeze_panes = "A2"
    # Auto-size columns
    for c, col_name in enumerate(columns, 1):
        max_len = len(col_name)
        for r in range(2, len(data) + 2):
            val = str(ws.cell(row=r, column=c).value or "")
            lines = val.split("\n")
            for line in lines:
                if len(line) > max_len:
                    max_len = len(line)
        width = min(max_len + 2, 60)
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = width

write_sheet(ws_tp, tp_cols, json_data)
write_sheet(ws_md, md_cols, json_data)

# ── Save ──
os.makedirs(output_dir, exist_ok=True)
wb.save(filepath)
print(f"Workbook saved: {filepath}")

# ── Validate ──
assert os.path.exists(filepath), "File does not exist!"
assert os.path.getsize(filepath) > 0, "File is empty!"
wb2 = load_workbook(filepath)
assert "TestPlan" in wb2.sheetnames, "TestPlan sheet missing!"
assert "MetaData" in wb2.sheetnames, "MetaData sheet missing!"
assert wb2["MetaData"].sheet_state == "veryHidden", "MetaData not veryHidden!"
assert wb2["TestPlan"].max_row == 3, f"Expected 3 rows (1 header + 2 data), got {wb2['TestPlan'].max_row}"
assert wb2["MetaData"].max_row == 3, f"Expected 3 rows (1 header + 2 data), got {wb2['MetaData'].max_row}"
print(f"Validation PASSED. File size: {os.path.getsize(filepath)} bytes")
print(f"FILENAME={filename}")

# ── If running in GitHub Actions, output for next step ──
github_output = os.environ.get("GITHUB_OUTPUT")
if github_output:
    with open(github_output, "a") as f:
        f.write(f"filename={filename}\n")
        f.write(f"filepath={filepath}\n")
