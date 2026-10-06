#!/usr/bin/env python3
"""MIPI_CSI TestPlan Excel Generator - Self-contained script.
Generates MIPI_CSI_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx using openpyxl.
Usage: python3 excel_generator.py
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import os, sys, json, base64, subprocess

def generate_workbook():
    IST = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(IST)
    timestamp = now_ist.strftime('%Y%m%d_%H%M%S')
    filename = f'MIPI_CSI_TestPlan_{timestamp}.xlsx'
    script_dir = os.path.dirname(os.path.abspath(__file__))
    filepath = os.path.join(script_dir, filename)

    wb = openpyxl.Workbook()
    blue_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    white_bold = Font(color='FFFFFF', bold=True, size=11)
    wrap = Alignment(wrap_text=True, vertical='top')

    # ===== TestPlan Sheet =====
    ws = wb.active
    ws.title = 'TestPlan'
    tp_cols = ['Index','SS / Module','Feature','Test Case Name','Test Description','Speed','Mode','Memory Start Offset','Memory End Offset','Remarks','Test Steps / Procedure','Impacted Registers','Validation / Acceptance Criteria','Code Generation']
    for i, c in enumerate(tp_cols, 1):
        cell = ws.cell(1, i, c)
        cell.fill = blue_fill
        cell.font = white_bold
        cell.alignment = wrap

    # Row 1
    r1 = [
        '1', 'MIPI_CSI', 'DPHY Lane Configuration and CSI2 Data Transfer', 'mipi_csi2_dphy_lanes_test',
        'This test validates MIPI CSI-2 DPHY lane configuration by testing data reception across all supported lane counts (4 lanes down to 1 lane). The test enables CSI-2 interrupts by clearing pending interrupts via the main interrupt status register and configuring all interrupt mask registers for PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected events. It configures the virtual channel and enables control data transfer. After D-PHY initialization, it polls the PHY stop state register until all lanes and the clock lane confirm stop state entry. For each lane count (4 to 1), the N_LANES register is configured and a signal is sent to start the CSI-2 sequence. For each expected packet, DMA channel 0 transfers control data, the test waits for DMA completion via interrupt polling, reads the control packet, and if the data type indicates a long packet, programs DMA channel 1 to transfer the image data and waits for its completion. The test verifies successful data reception across all lane configurations.',
        'NA', 'NA', 'NA', 'NA',
        'The test iterates lane_num from 3 to 0, writing each value to N_LANES, effectively testing 4-lane, 3-lane, 2-lane, and 1-lane DPHY configurations. The virtual channel and control data registers are written twice (repeated writes). The D-PHY initialization is performed via snps_phy_init() whose implementation is external to this testcase. DMA transfer functions dma_trnsfr_instn_preload() and DMAGO_CSI() are also external. Two hex addresses (used for signaling and DMA status) could not be mapped to named registers in the specification documents. Default configuration uses VC_ID=3, VRES=3, HRES=64, DATA_TYPE=CSI2_RGB888 unless overridden by compile-time defines.',
        '1. Enable CSI-2 interrupts by reading the main interrupt status register to clear pending interrupts, then configure all interrupt mask registers (PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, ECC corrected) with appropriate enable masks. 2. Configure the virtual channel register with the selected virtual channel ID based on the GDMA path. 3. Enable control data transfer by writing to the control data register. 4. Perform D-PHY initialization sequence. 5. Poll the PHY stop state register until all data lanes and the clock lane confirm stop state entry (expected value indicates all lanes stopped). 6. For each lane configuration (4 lanes down to 1 lane): a. Write the lane count to the N_LANES register. b. Signal the start of the CSI-2 sequence for the current lane count. c. For each expected packet in the frame: i. Enable DMA interrupts for both channels. ii. Program and start DMA channel 0 for control data transfer. iii. Poll the DMA interrupt status register until channel 0 completion is indicated. iv. Clear the DMA channel 0 interrupt. v. Read the received control packet data. vi. If the control packet indicates a long packet (image data), calculate the transfer size, program and start DMA channel 1 for image data transfer, poll for channel 1 completion, and clear the channel 1 interrupt. 7. Verify test completes successfully for all lane configurations.',
        'virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED',
        '1. The PHY stop state register must read the expected value confirming all data lanes and the clock lane have entered stop state before proceeding with lane configuration. 2. For each DMA control data transfer, the DMA interrupt status must indicate channel 0 completion. 3. For each DMA image data transfer (when a long packet is detected), the DMA interrupt status must indicate channel 1 completion. 4. The test must successfully iterate through all lane configurations (4 lanes down to 1 lane) and complete all packet transfers for each configuration. 5. The test must call the finish routine with a pass indication (0) after all lane configurations are validated.',
        ''
    ]
    for i, v in enumerate(r1, 1):
        cell = ws.cell(2, i, v)
        cell.alignment = wrap

    # Row 2
    r2 = [
        '2', 'MIPI_CSI', 'Test Pattern Generator and CSI2 Data Reception', 'mipi_csi2_test_pattern_generator',
        'This test validates the MIPI CSI-2 internal test pattern generator and verifies data reception through the DMA path. The test configures the virtual channel and disables control data transfer. It enables all CSI-2 interrupts by clearing pending interrupts via the main interrupt status register and configuring all interrupt mask registers for PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected events. After D-PHY initialization, it polls the PHY stop state register until all data lanes and the clock lane confirm stop state entry. The test programs the DMA address registers for channel 0 read and write paths, enables the fractional divider output to the CSI-2 subsystem via the enableclkgating_csiphy register, and programs the DMA transfer with calculated transfer size based on resolution (320x16) and 24 bits per pixel. The pattern generator is configured with vertical resolution, horizontal resolution, and configuration parameters, then enabled. After a short wait, the pattern generator is disabled, and the test polls the DMA interrupt status until channel 0 transfer completion is confirmed. The test completes with a pass indication after a final wait.',
        'NA', 'NA', 'NA', 'NA',
        'The test uses the internal PPI test pattern generator to generate CSI-2 data internally, bypassing the need for an external D-PHY transmitter. The pattern generator is configured with vres=16, hres=320, 24 bits per pixel (RGB888 implied by config value 0xe401 with data type 0x24). The D-PHY initialization is performed via snps_phy_init() whose implementation is external to this testcase. DMA transfer functions dma_trnsfr_instn_preload_incr_addr() and DMAGO_CSI() are also external. The csi2_subsys_enable_interrupt() function is external but internally calls the locally defined csi2_enable_interrupt(). The destination address increment behavior is controlled by the FPS60 compile-time define. Control data transfer is explicitly disabled (written 0) unlike the DPHY lanes test. The enableclkgating_csiphy register is written to enable fractional divider output to the CSI-2 subsystem.',
        '1. Configure the virtual channel register with the selected virtual channel ID based on the GDMA path. 2. Disable control data transfer by writing to the control data register. 3. Enable CSI-2 interrupts by reading the main interrupt status register to clear pending interrupts, then configure all interrupt mask registers (PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, ECC corrected) with appropriate enable masks. 4. Perform D-PHY initialization sequence. 5. Poll the PHY stop state register until all data lanes and the clock lane confirm stop state entry. 6. Calculate the data transfer size based on resolution (320x16) and 24 bits per pixel with 8-byte alignment. 7. Program the DMA higher-order AXI address registers for channel 0 read and write paths. 8. Enable the fractional divider output to the CSI-2 subsystem via the enableclkgating_csiphy register. 9. Program and start DMA channel 0 for image data transfer with the calculated transfer size. 10. Enable the test pattern generator by configuring vertical resolution, horizontal resolution, pattern configuration, and enabling the PPI_PG_ENABLE register. 11. Wait for pattern generation to complete, then disable the pattern generator by clearing the PPI_PG_ENABLE register. 12. Poll the DMA interrupt status register until channel 0 transfer completion is indicated. 13. Wait for a final settling period and verify test completes successfully.',
        'PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; enableclkgating_csiphy; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED',
        '1. The PHY stop state register must read the expected value confirming all data lanes and the clock lane have entered stop state before proceeding. 2. The DMA interrupt status must indicate channel 0 transfer completion after the pattern generator has been enabled and disabled. 3. The test must complete the full sequence of pattern generator enable, wait, disable, and DMA completion polling without hanging. 4. The test must call the finish routine with a pass indication (0) after successful DMA transfer completion and final wait.',
        ''
    ]
    for i, v in enumerate(r2, 1):
        cell = ws.cell(3, i, v)
        cell.alignment = wrap

    col_widths = {'A':8,'B':15,'C':40,'D':40,'E':60,'F':8,'G':8,'H':22,'I':22,'J':60,'K':60,'L':55,'M':60,'N':15}
    for letter, width in col_widths.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = 'A2'

    # ===== MetaData Sheet =====
    ws2 = wb.create_sheet('MetaData')
    md_cols = ['Index','Test Case Name','Meta Test Description','Meta Test Steps / Procedure','Meta Impacted Registers','Meta Validation / Acceptance Criteria','Meta Headers','Meta Macros','Meta Arrays']
    for i, c in enumerate(md_cols, 1):
        cell = ws2.cell(1, i, c)
        cell.fill = blue_fill
        cell.font = white_bold
        cell.alignment = wrap

    # MetaData Row 1
    md_r1 = [
        '1', 'mipi_csi2_dphy_lanes_test',
        'This testcase validates MIPI CSI-2 DPHY lane configuration by iterating through all lane counts (4 lanes down to 1 lane). For each lane configuration, it performs a complete CSI-2 data reception sequence. The test begins by calling csi2_enable_interrupt() which reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts, then writes all interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED) to enable various CSI-2 interrupts. It then configures the virtual channel register MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with the VC_ID (default 3) shifted based on the selected GDMA path, and enables control data transfer by writing 1 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA. The D-PHY is initialized via snps_phy_init(), and the test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until it reads 0x1000f, confirming all lanes and the clock lane have entered stop state. A for-loop iterates lane_num from 3 down to 0, writing lane_num to MIZAR_MIPI_CSI2_HOST_N_LANES and signaling the lane count via write to 0xa0243ffc. For each lane configuration, an inner loop iterates over the expected number of control packets (cntrl_pkt_cnt = VRES*3 + 2). In each iteration, DMA channel 0 is programmed via dma_trnsfr_instn_preload() and started via DMAGO_CSI(gdma_reg_base, ch0_pc, 0x0) for control data transfer. The test polls gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET for bit 0 to confirm DMA channel 0 completion, then clears the interrupt via gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET. It reads the control data from 0xE6001000 and checks if the data type field (bits [5:0]) is greater than 0xf, indicating a long packet with image data. If so, it calculates the transfer size from word_count (bits [21:6]), programs DMA channel 1 via dma_trnsfr_instn_preload() and DMAGO_CSI(gdma_reg_base, ch1_pc, 0x1), polls for bit 1 in the DMA interrupt status, and clears the channel 1 interrupt. After all lane configurations are tested, finish(0) is called.',
        '1. Call csi2_enable_interrupt() function. 2. Inside csi2_enable_interrupt(): read MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts (rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN)). 3. Write 0x0000000f to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL to enable PHY fatal interrupts. 4. Write 0x00000003 to MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL to enable packet fatal interrupts. 5. Write 0x000f000f to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY to enable PHY interrupts. 6. Write 0x000f000f to MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE to enable line interrupts. 7. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL to enable boundary frame fatal interrupts. 8. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL to enable sequence frame fatal interrupts. 9. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL to enable CRC frame fatal interrupts. 10. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL to enable payload CRC fatal interrupts. 11. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID to enable data ID interrupts. 12. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED to enable ECC corrected interrupts. 13. Return from csi2_enable_interrupt(). 14. Determine vcid_csi2_wrap_reg based on GDMA path. 15. Set gdma_reg_base = 0xE6A00000. 16. Write vcid_csi2_wrap_reg to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL. 17. Write 1 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA. 18-19. Repeat writes. 20. Call snps_phy_init(). 21-22. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until 0x1000f. 23. Set DMA PC addresses. 24-47. Lane iteration and DMA transfer loop. 48. Call finish(0).',
        'MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED',
        'The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until rd_data equals 0x1000f, confirming all 4 data lanes (bits [3:0] = 0xf) and the clock lane (bit 16 = 0x10000) have entered stop state. For each packet transfer, the DMA interrupt status register is polled: bit 0 must be set for DMA channel 0 completion, and bit 1 must be set for DMA channel 1 completion. The test calls finish(0) upon successful completion of all lane configurations, indicating a pass.',
        '<stdio.h>; <stdlib.h>; "test_common.h"; "mipi_csi2.h"',
        'GDMA_CSI2_DATA_DEST_ADDR2; GDMA_CTRL_DATA_DEST_ADDR2; LS_LE_EN; TOTAL_FRAME; VC_ID; VRES; HRES; DATA_TYPE',
        'NA'
    ]
    for i, v in enumerate(md_r1, 1):
        cell = ws2.cell(2, i, v)
        cell.alignment = wrap

    # MetaData Row 2
    md_r2 = [
        '2', 'mipi_csi2_test_pattern_generator',
        'This testcase validates the MIPI CSI-2 internal test pattern generator (PPI PG) and verifies data reception through the DMA path. The test configures the virtual channel register MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with a VC ID of 3 (shifted based on the selected GDMA path) and disables control data transfer by writing 0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA. It calls csi2_subsys_enable_interrupt() to clear pending interrupts and enable all interrupt mask registers. D-PHY initialization is performed via snps_phy_init(), then MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE is polled until it reads 0x1000f. The test calculates the data transfer size based on hres=320, vres=16, valid_bits_per_pixel=24, with 8-byte alignment. DMA address registers are programmed. A write to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 with value 0x1 enables sending fracdiv output to the CSI-2 subsystem. DMA channel 0 is programmed and started. The pattern generator is enabled by calling csi2_ctrlr_pg_enable() which writes PPI_PG_PATTERN_VRES, PPI_PG_PATTERN_HRES, PPI_PG_CONFIG, and PPI_PG_ENABLE. After wait_on(100), the pattern generator is disabled. The DMA interrupt status is polled until bit 0 is set. After wait_on(10000), finish(0) is called.',
        '1. Set int_pend = 1. 2. Set vcid = 3 and calculate vcid_unselected_path. 3. Determine vcid_csi2_wrap_reg based on GDMA path. 4. Calculate gdma_reg_base. 5. Write vcid_csi2_wrap_reg to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL. 6. Write 0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA. 7-19. Enable interrupts. 20-22. D-PHY init and poll PHY_STOPSTATE. 23-24. Calculate transfer size. 25-28. Program DMA address registers. 29. Write 0x1 to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4. 30-33. Program and start DMA. 34-39. Enable pattern generator. 40-41. Wait and disable PG. 42-43. Poll DMA completion. 44-45. Final wait and finish(0).',
        'MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED',
        'The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until rd_data equals 0x1000f. The DMA interrupt status register is polled until bit 0 is set, confirming DMA channel 0 transfer completion. The test calls finish(0) upon successful completion, indicating a pass.',
        '<stdio.h>; <stdlib.h>; "test_common.h"; "mipi_csi2.h"',
        'GDMA_CSI2_DATA_DEST_ADDR2; GDMA_CTRL_DATA_DEST_ADDR2',
        'NA'
    ]
    for i, v in enumerate(md_r2, 1):
        cell = ws2.cell(3, i, v)
        cell.alignment = wrap

    for letter in ['A','B','C','D','E','F','G','H','I']:
        ws2.column_dimensions[letter].width = 45
    ws2.column_dimensions['A'].width = 8
    ws2.column_dimensions['B'].width = 40
    ws2.freeze_panes = 'A2'
    ws2.sheet_state = 'veryHidden'

    wb.save(filepath)
    file_size = os.path.getsize(filepath)
    print(f'Generated: {filename}')
    print(f'Path: {filepath}')
    print(f'Size: {file_size} bytes')

    # Verify
    wb2 = openpyxl.load_workbook(filepath)
    assert 'TestPlan' in wb2.sheetnames, 'TestPlan sheet missing'
    assert 'MetaData' in wb2.sheetnames, 'MetaData sheet missing'
    assert wb2['TestPlan'].max_row == 3, f'Expected 3 rows, got {wb2["TestPlan"].max_row}'
    assert wb2['MetaData'].max_row == 3, f'Expected 3 rows, got {wb2["MetaData"].max_row}'
    print('Verification: PASSED')
    print(f'TestPlan rows: {wb2["TestPlan"].max_row - 1}')
    print(f'MetaData rows: {wb2["MetaData"].max_row - 1}')
    wb2.close()

    return filepath, filename

if __name__ == '__main__':
    filepath, filename = generate_workbook()
    print(f'\nSUCCESS: {filename}')
