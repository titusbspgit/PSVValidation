#!/usr/bin/env python3
"""
Agent 7 - Excel Generator: MIPI_CSI TestPlan XLSX Generator
Generates MIPI_CSI_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx with TestPlan and MetaData sheets.
Run: python generate_mipi_csi_testplan_20261009.py
Requires: pip install openpyxl
After running, commit the generated .xlsx to this same directory.
"""
import json, os, sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
FILENAME = f"MIPI_CSI_TestPlan_{now_ist.strftime('%Y%m%d_%H%M%S')}.xlsx"

JSON_DATA = [
    {
        "Index": "1",
        "SS / Module": "MIPI_CSI",
        "Test Case Name": "mipi_csi2_dphy_lanes_test",
        "Feature": "DPHY Lane Configuration",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_csi2.h\"",
        "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000; LS_LE_EN = 1; TOTAL_FRAME = 1; VC_ID = 3; VRES = 3; HRES = 64; DATA_TYPE",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates MIPI CSI-2 DPHY lane configuration by iterating through lane counts from 4 down to 1. It begins by calling csi2_enable_interrupt() which reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts, then writes all ten interrupt mask registers (MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL with 0x0000000f, MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL with 0x00000003, MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY with 0x000f000f, MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE with 0x000f000f, MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID with 0x0000ffff, MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED with 0x0000ffff) to enable all CSI-2 host interrupts. The virtual channel register MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL is written with the computed vcid_csi2_wrap_reg value (based on VC_ID shifted by GDMA path), and MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA is written with 1 to enable control data transfer. These two writes are performed twice. snps_phy_init() is called for D-PHY initialization. MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE is then polled in a while loop until it equals 0x1000f, indicating all lanes have entered stop state. A for loop iterates lane_num from 3 down to 0: MIZAR_MIPI_CSI2_HOST_N_LANES is written with lane_num, 0xa0243ffc is written with (lane_num+1) to trigger the CSI-2 sequence. An inner loop iterates cntrl_pkt_cnt = (VRES*3)+2 times. In each iteration, gdma_reg_base+MIPI_CSI2_DMA_INTEN_OFFSET is written with 0x3 to enable DMA IRQs. dma_trnsfr_instn_preload() is called for channel 0 control data transfer (src_addr=0x8000, dest_addr=0xE6001000, trnsfr_size=8, irq_num=0). DMAGO_CSI() is called for channel 0. gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET is polled until bit 0 is set. gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET is written with 0x1 to clear the IRQ. gdma_reg_base+0x28 is read. 0xE6001000 is read to get csi_ctrl_data. If (csi_ctrl_data & 0x3f) > 0xf, word_count is extracted from bits [21:6], csi_data_size is computed as 8-byte aligned word_count, dma_trnsfr_instn_preload() is called for channel 1 data transfer (src_addr=0x0000, dest_addr=0xE6002000, trnsfr_size=csi_data_size, irq_num=1), DMAGO_CSI() is called for channel 1, gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET is polled until bit 1 is set, then read again, and gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET is written with 0x2 to clear the IRQ. The test ends with finish(0).",
        "Test Description": "This test validates MIPI CSI-2 DPHY lane configuration by iterating through all supported lane counts (4 lanes down to 1 lane). It first enables all CSI-2 host interrupts by reading the main interrupt status register to clear pending interrupts, then writing all ten interrupt mask registers with appropriate enable masks. The virtual channel register is configured with the selected virtual channel ID, and control data transfer is enabled. D-PHY initialization is performed, and the PHY stop state register is polled until all lanes enter stop state. For each lane configuration (4 to 1), the number of active lanes is set in the N_LANES register, and a trigger write initiates the CSI-2 sequence. For each control packet, DMA interrupt enable is configured, a DMA control data transfer is initiated on channel 0, and the DMA interrupt status is polled until the transfer completes. The DMA interrupt is then cleared. The control data destination is read to determine the CSI data type and word count. If the data type indicates a long packet, the word count is extracted, aligned to 8 bytes, and a DMA data transfer is initiated on channel 1. The DMA interrupt status is polled for channel 1 completion, and the interrupt is cleared. This process repeats for all packets across all lane configurations. The test passes if all DMA transfers complete successfully for every lane count.",
        "Meta Test Steps / Procedure": "1. Entry: test_case() is called. 2. Call csi2_enable_interrupt(): read INT_ST_MAIN, write all 10 interrupt mask registers. 3. GDMA path selection and vcid computation. 4. Write virtual_channel and control_data registers (twice). 5. Call snps_phy_init(). 6. Poll PHY_STOPSTATE until 0x1000f. 7. Loop lane_num 3..0: write N_LANES, write trigger, inner loop for DMA transfers on ch0 and ch1. 8. finish(0).",
        "Test Steps / Procedure": "1. Enable CSI-2 host interrupts by reading the main interrupt status register to clear pending interrupts, then writing all ten interrupt mask registers. 2. Configure the virtual channel register. 3. Enable control data transfer. Repeat writes. 4. Perform D-PHY initialization. 5. Poll PHY stop state register. 6. Loop lane count 4 to 1: set N_LANES, trigger CSI-2 sequence. 7. For each packet: enable DMA IRQs, preload and start DMA ch0, poll completion, clear IRQ, read control data. 8. If long packet: preload and start DMA ch1, poll completion, clear IRQ. 9. Verify test completes successfully.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000",
        "Impacted Registers": "INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED; virtual_channel; control_data; PHY_STOPSTATE; N_LANES",
        "Meta Validation / Acceptance Criteria": "1. PHY_STOPSTATE polled until 0x1000f. 2. DMA ch0 INTMIS bit 0 polled. 3. csi_ctrl_data checked for long packet. 4. DMA ch1 INTMIS bit 1 polled. 5. finish(0) called on success.",
        "Validation / Acceptance Criteria": "1. PHY stop state register must indicate all lanes stopped. 2. DMA ch0 completion bit must be set. 3. Control data checked for long packet type. 4. DMA ch1 completion bit must be set. 5. All transfers must complete for all 4 lane configs. 6. Test passes with finish(0).",
        "Remarks": "Virtual channel and control data writes performed twice. snps_phy_init() not available in testcase folder. DMA offset macros used as compound expressions. Two hardcoded hex addresses unmapped."
    },
    {
        "Index": "2",
        "SS / Module": "MIPI_CSI",
        "Test Case Name": "mipi_csi2_test_pattern_generator",
        "Feature": "Test Pattern Generator",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_csi2.h\"",
        "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates the MIPI CSI-2 internal test pattern generator functionality. The test configures the CSI-2 subsystem virtual channel register, disables control data transfer, enables all CSI-2 host interrupts, performs D-PHY initialization, polls PHY stop state, configures DMA address registers, enables clock gating, sets up DMA transfer, enables then disables the pattern generator, and polls DMA completion.",
        "Test Description": "This test validates the MIPI CSI-2 internal test pattern generator by configuring the subsystem virtual channel register, disabling control data transfer, enabling all CSI-2 host interrupts, performing D-PHY initialization, and waiting for PHY stop state. DMA higher-order address registers are configured for channel 0 data read and write paths. The clock gating register is written to enable fracdiv output to the CSI-2 subsystem. A DMA data transfer is set up with computed transfer size based on resolution (320x16) and pixel depth (24 bits per pixel) with 8-byte alignment. DMA channel 0 is started. The pattern generator is then enabled by configuring vertical resolution, horizontal resolution, pattern generator configuration, and pattern generator enable registers. After a short wait, the pattern generator is disabled. The DMA interrupt masked status is polled until the channel 0 transfer completes. The test passes after a final wait period.",
        "Meta Test Steps / Procedure": "1. Entry: test_case() called. 2. vcid=3, compute vcid mapping. 3. Write virtual_channel, write control_data=0. 4. Enable interrupts (read INT_ST_MAIN, write 10 mask regs). 5. snps_phy_init(). 6. Poll PHY_STOPSTATE until 0x1000f. 7. Compute transfer size (320x16x24bpp, 8-byte aligned). 8. Write DMA address registers. 9. Write enableclkgating_csiphy. 10. Preload DMA ch0 and start. 11. Enable PG (VRES, HRES, CONFIG, ENABLE). 12. Wait, disable PG. 13. Poll DMA INTMIS bit 0. 14. Wait, finish(0).",
        "Test Steps / Procedure": "1. Set virtual channel ID and compute mapping. Write virtual channel register. 2. Disable control data transfer. 3. Enable all CSI-2 host interrupts. 4. Perform D-PHY initialization. 5. Poll PHY stop state register. 6. Compute DMA transfer size. 7. Configure DMA address registers. 8. Write clock gating register. 9. Preload and start DMA ch0. 10. Enable pattern generator. 11. Wait, disable pattern generator. 12. Poll DMA completion. 13. Wait, verify test completes.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
        "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; enableclkgating_csiphy; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
        "Meta Validation / Acceptance Criteria": "1. PHY_STOPSTATE polled until 0x1000f. 2. DMA INTMIS bit 0 polled after PG disable. 3. finish(0) called. 4. No explicit data comparison.",
        "Validation / Acceptance Criteria": "1. PHY stop state must indicate all lanes stopped. 2. DMA ch0 completion bit must be set after PG disable. 3. Test passes with finish(0). 4. No explicit data comparison; DMA completion is acceptance criterion.",
        "Remarks": "Control data disabled (written 0). PG enabled then disabled after wait_on(100). DMA uses incremental address support. FPS60 controls dest addr increment. MIZAR_MIPI_CSI2_RB_REG_BASE+0xf4 maps to enableclkgating_csiphy."
    }
]

TESTPLAN_COLS = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

METADATA_COLS = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

def create_workbook():
    wb = Workbook()
    # --- TestPlan Sheet ---
    ws_tp = wb.active
    ws_tp.title = "TestPlan"
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_align = Alignment(wrap_text=True, vertical="top")

    for col_idx, col_name in enumerate(TESTPLAN_COLS, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_align

    for row_idx, row_data in enumerate(JSON_DATA, 2):
        for col_idx, col_name in enumerate(TESTPLAN_COLS, 1):
            val = row_data.get(col_name, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = wrap_align

    ws_tp.freeze_panes = "A2"

    # Auto-size columns
    for col_idx, col_name in enumerate(TESTPLAN_COLS, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(JSON_DATA) + 2):
            val = str(ws_tp.cell(row=row_idx, column=col_idx).value or "")
            lines = val.split("\n")
            for line in lines:
                max_len = max(max_len, len(line))
        ws_tp.column_dimensions[ws_tp.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 80)

    # --- MetaData Sheet ---
    ws_md = wb.create_sheet("MetaData")
    for col_idx, col_name in enumerate(METADATA_COLS, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_align

    for row_idx, row_data in enumerate(JSON_DATA, 2):
        for col_idx, col_name in enumerate(METADATA_COLS, 1):
            val = row_data.get(col_name, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = wrap_align

    ws_md.freeze_panes = "A2"

    for col_idx, col_name in enumerate(METADATA_COLS, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(JSON_DATA) + 2):
            val = str(ws_md.cell(row=row_idx, column=col_idx).value or "")
            lines = val.split("\n")
            for line in lines:
                max_len = max(max_len, len(line))
        ws_md.column_dimensions[ws_md.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 80)

    ws_md.sheet_state = "veryHidden"

    # Save
    script_dir = os.path.dirname(os.path.abspath(__file__))
    filepath = os.path.join(script_dir, FILENAME)
    wb.save(filepath)
    print(f"Workbook saved: {filepath}")

    # Validate
    wb2 = load_workbook(filepath)
    assert "TestPlan" in wb2.sheetnames, "TestPlan sheet missing"
    assert "MetaData" in wb2.sheetnames, "MetaData sheet missing"
    assert wb2["MetaData"].sheet_state == "veryHidden", "MetaData not veryHidden"
    tp_rows = wb2["TestPlan"].max_row - 1
    md_rows = wb2["MetaData"].max_row - 1
    fsize = os.path.getsize(filepath)
    print(f"Validation PASSED: {tp_rows} TestPlan rows, {md_rows} MetaData rows, {fsize} bytes")
    print(f"Filename: {FILENAME}")
    wb2.close()
    return filepath

if __name__ == "__main__":
    create_workbook()
