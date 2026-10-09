#!/usr/bin/env python3
"""MIPI CSI TestPlan XLSX Generator - Agent 7 Direct Execution"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os, sys, json, base64, datetime

def generate_workbook():
    # IST timestamp
    from datetime import timezone, timedelta
    ist = timezone(timedelta(hours=5, minutes=30))
    now = datetime.datetime.now(ist)
    timestamp = now.strftime('%Y%m%d_%H%M%S')
    filename = f'MIPI_CSI_TestPlan_{timestamp}.xlsx'

    # Input data
    row_data = {
        "Index": "1",
        "SS / Module": "mipi_csi2_subsys",
        "Feature": "DPHY Lanes",
        "Test Case Name": "mipi_csi2_dphy_lanes_test",
        "Test Description": "Verify MIPI CSI-2 D-PHY lane configuration by iterating through all supported lane counts (4 down to 1), enabling CSI-2 interrupts, initializing the D-PHY, waiting for PHY stop state, configuring virtual channels, and performing DMA-based control data and CSI image data transfers for each packet in a frame at each lane configuration.",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "0xE6000000; 0xE6000500; 0xE6001000; 0xE6002000; 0xE6A00000",
        "Memory End Offset": "NA",
        "Remarks": "The functions snps_phy_init(), dma_trnsfr_instn_preload(), DMAGO_CSI(), and finish() are external functions not defined within the testcase folder; their implementations are assumed to be provided by the test framework or included via test_common.h and mipi_csi2.h. The macros MIPI_CSI2_DMA_INTEN_OFFSET, MIPI_CSI2_DMA_INTMIS_OFFSET, and MIPI_CSI2_DMA_INTCLR_OFFSET are expected to be defined in mipi_csi2.h. The GDMA path selection (GDMA0_PATH, GDMA1_PATH, GDMA2_PATH, GDMA3_PATH) and GDMA0_FULL_MEM are compile-time configuration macros that determine the virtual channel register value and resolution parameters. Two hardcoded addresses (0xa0243ffc and 0xE6001000) remain unresolved in the Agent 4 register mapping. The virtual_channel and control_data registers are written twice in sequence, which appears intentional in the source. The DPHY_LANES_TEST conditional compilation guard must be defined for the test-specific macros to take effect.",
        "Test Steps / Procedure": "1. Enable all CSI-2 host controller interrupts by reading the main interrupt status register to clear pending interrupts, then writing mask values to enable PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected interrupt categories.\n2. Configure the virtual channel register based on the selected GDMA path and enable control data transfer.\n3. Repeat the virtual channel and control data enable writes a second time.\n4. Initialize the SNPS D-PHY.\n5. Poll the PHY_STOPSTATE register until all 4 data lanes and the clock lane report stop state (expected value 0x1000f).\n6. Set DMA program counter addresses for channel 0 and channel 1.\n7. Iterate through lane configurations from 4 lanes (lane_num=3) down to 1 lane (lane_num=0).\n8. For each lane configuration, write the lane count to the N_LANES register and trigger the CSI-2 sequence.\n9. For each packet in the frame (total packets = VRES*3 + 2), perform the following:\n a. Enable DMA interrupts for both channels.\n b. Preload DMA channel 0 transfer instructions for control data transfer (8 bytes from source to destination).\n c. Start DMA channel 0 and poll the DMA interrupt status register until channel 0 transfer completes.\n d. Clear the DMA channel 0 interrupt.\n e. Read the control data and check the data type field.\n f. If the data type indicates CSI image data, compute the word count and aligned transfer size, preload DMA channel 1 transfer instructions, start DMA channel 1, poll for channel 1 completion, and clear the channel 1 interrupt.\n10. After all lane configurations are tested, signal test completion.",
        "Impacted Registers": "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
        "Validation / Acceptance Criteria": "1. All CSI-2 interrupt mask registers are correctly programmed with the specified enable values.\n2. The PHY_STOPSTATE register confirms all 4 data lanes and the clock lane are in stop state before proceeding.\n3. The N_LANES register is correctly configured for each lane count iteration (4 lanes down to 1 lane).\n4. DMA channel 0 control data transfers complete successfully for each packet, confirmed by DMA interrupt status polling.\n5. DMA channel 1 CSI image data transfers complete successfully when the control data indicates a valid CSI data type.\n6. DMA interrupts are properly cleared after each transfer completion.\n7. The test completes with a pass status via finish(0) after all lane configurations are exercised.",
        "Code Generation": "",
        "Meta Test Description": "This testcase validates MIPI CSI-2 D-PHY lane configuration by iterating through all lane counts (4 lanes down to 1 lane). For each lane configuration, the test programs the CSI-2 host controller N_LANES register, triggers a CSI-2 sequence via a write to hardcoded address 0xa0243ffc, and then performs DMA-based control data and CSI image data transfers for each packet in a frame. The test first enables all CSI-2 interrupts via csi2_enable_interrupt(), configures the virtual channel register and control data enable, initializes the SNPS D-PHY via snps_phy_init(), and polls PHY_STOPSTATE until all 4 data lanes and clock lane are in stop state (value 0x1000f). For each lane count iteration (lane_num=3 down to 0), the test computes cntrl_pkt_cnt = (VRES3)+2 and loops through that many packets. Each packet iteration programs DMA interrupt enable, preloads DMA channel 0 transfer instructions for control data (src=0x8000, dest=0xE6001000, size=8 bytes), issues DMAGO_CSI for channel 0, polls DMA interrupt status register (MIPI_CSI2_DMA_INTMIS_OFFSET) until bit 0 is set, clears DMA interrupt, reads control data from 0xE6001000, and if the data type field (bits[5:0]) > 0xf, computes word_count from bits[21:6], calculates csi_data_size aligned to 8 bytes, preloads DMA channel 1 transfer instructions for CSI data (src=0x0000, dest=0xE6002000, size=csi_data_size), issues DMAGO_CSI for channel 1, polls DMA interrupt status until bit 1 is set, reads DMA interrupt status again, and clears DMA interrupt bit 1. The test completes with finish(0).",
        "Meta Test Steps / Procedure": "1. Enter test_case() function.\n2. Declare local variables: rx_desc, tx_desc (long long int), gdma_tx_trnsfr_size, gdma_trnsfr_size (long long int), cntrl_pkt_cnt (int).\n3. printf(\"start line\\n\").\n4. Call csi2_enable_interrupt() to enable CSI-2 interrupts.\n5. [Inside csi2_enable_interrupt()]: Declare local int rd_data.\n6. [Inside csi2_enable_interrupt()]: rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN) \u2014 read INT_ST_MAIN register to clear interrupts.\n7. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, 0x0000000f) \u2014 enable phy_fatal interrupts for all 4 lanes.\n8. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, 0x00000003) \u2014 enable pkt_fatal interrupts (err_ecc_double and shorter_payload).\n9. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, 0x000f000f) \u2014 enable phy interrupts (errsoths[3:0] and erresc[19:16]).\n10. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, 0x000f000f) \u2014 enable line interrupts (l_bndry_match[3:0] and l_seq[19:16]).\n11. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, 0x0000ffff) \u2014 enable boundary frame fatal interrupts for all 16 VCs.\n12. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, 0x0000ffff) \u2014 enable seq frame fatal interrupts for all 16 VCs.\n13. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, 0x0000ffff) \u2014 enable CRC frame fatal interrupts for all 16 VCs.\n14. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, 0x0000ffff) \u2014 enable payload CRC fatal interrupts for all 16 VCs.\n15. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, 0x0000ffff) \u2014 enable data ID interrupts for all 16 VCs.\n16. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED, 0x0000ffff) \u2014 enable ECC corrected interrupts for all 16 VCs.\n17. Return from csi2_enable_interrupt().\n18. Conditional compilation: set vcid_csi2_wrap_reg based on GDMA path. For GDMA3_PATH: vcid_csi2_wrap_reg = VC_ID. For GDMA2_PATH: vcid_csi2_wrap_reg = (VC_ID << 4). For GDMA1_PATH: vcid_csi2_wrap_reg = (VC_ID << 8). For GDMA0_PATH: vcid_csi2_wrap_reg = (VC_ID << 12), printf(\"VC_ID=%d\\n\", VC_ID), gdma_path = 0.\n19. Set gdma_reg_base = 0xE6A00000.\n20. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg) \u2014 first write to CSI-2 virtual channel register.\n21. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 1) \u2014 first enable control data transfer.\n22. printf(\"vcid_csi2_wrap_reg=%0x\\n\", vcid_csi2_wrap_reg).\n23. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg) \u2014 second write to CSI-2 virtual channel register.\n24. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 1) \u2014 second enable control data transfer.\n25. Call snps_phy_init() \u2014 D-PHY initialization sequence (external function, implementation not in testcase folder).\n26. rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) \u2014 first read of PHY stop state.\n27. Enter while loop: condition !(rd_data == 0x1000f).\n28. Inside while loop: rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) \u2014 poll PHY stop state until all 4 data lanes (bits[3:0]) and clock lane (bit[16]) are in stop state (value 0x1000f).\n29. Exit while loop when rd_data == 0x1000f.\n30. Set ch0_pc = 0xE6000000.\n31. Set ch1_pc = 0xE6000500.\n32. Enter outer for loop: lane_num = 3 down to 0 (lane_num >= 0, lane_num--).\n33. write_reg(MIZAR_MIPI_CSI2_HOST_N_LANES, lane_num) \u2014 configure number of lanes in CSI-2 controller (lane_num=3 means 4 lanes, lane_num=0 means 1 lane).\n34. Compute cntrl_pkt_cnt = ((VRES * 3) + 2).\n35. printf(\"DEBUG: cntrl_pkt_cnt=%d\\n\", cntrl_pkt_cnt).\n36. write_reg(0xa0243ffc, (lane_num + 1)) \u2014 trigger CSI-2 sequence with lane count value.\n37. Enter inner for loop: i = 0 to cntrl_pkt_cnt - 1.\n38. Set ch0_preload_loc = ch0_pc.\n39. Set ch1_preload_loc = ch1_pc.\n40. write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTEN_OFFSET, 0x3) \u2014 enable dma_irq[1] and dma_irq[0].\n41. Call dma_trnsfr_instn_preload(ch0_preload_loc, gdma_reg_base, 0x8000, 0xE6001000, 8, 0) \u2014 preload DMA channel 0 transfer instructions for control data (src_addr=0x8000, dest_addr=0xE6001000, trnsfr_size=8, irq_num=0).\n42. Call DMAGO_CSI(gdma_reg_base, ch0_pc, 0x0) \u2014 start DMA channel 0 (ch_num=0).\n43. Set rd_data = 0.\n44. Enter while loop: condition (rd_data & 0x1) == 0x0.\n45. Inside while loop: rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) \u2014 poll DMA interrupt status register for channel 0 completion (bit 0).\n46. Inside while loop: printf(\"polling irq; irq_status_reg rd_data =%0x\\n\", rd_data).\n47. Exit while loop when (rd_data & 0x1) != 0x0.\n48. write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET, 0x1) \u2014 clear DMA IRQ bit 0.\n49. rd_data = read_reg(gdma_reg_base + 0x28) \u2014 read DMA interrupt status register at offset 0x28.\n50. printf(\"DEBUG : irq polling completed ch0; rd_data = %0x \\n\", rd_data).\n51. csi_ctrl_data = read_reg(0xE6001000) \u2014 read control data from address 0xE6001000.\n52. printf(\"DEBUG: csi_ctrl_data=%0x\\n\", csi_ctrl_data).\n53. Check condition: if ((csi_ctrl_data & 0x3f) > 0xf) \u2014 check if data type field indicates CSI image data.\n54. [If true]: word_count = ((csi_ctrl_data >> 6) & 0xffff) \u2014 extract word count from control data bits[21:6].\n55. [If true]: csi_data_size = (word_count % 8) ? (word_count / 8 + 1) * 8 : word_count \u2014 align transfer size to 8-byte boundary.\n56. [If true]: Call dma_trnsfr_instn_preload(ch1_preload_loc, gdma_reg_base, 0x0000, 0xE6002000, csi_data_size, 1) \u2014 preload DMA channel 1 transfer instructions for CSI data (src_addr=0x0000, dest_addr=0xE6002000, trnsfr_size=csi_data_size, irq_num=1).\n57. [If true]: Call DMAGO_CSI(gdma_reg_base, ch1_pc, 0x1) \u2014 start DMA channel 1 (ch_num=1).\n58. [If true]: Set rd_data = 0.\n59. [If true]: Enter while loop: condition (rd_data & 0x2) == 0x0.\n60. [If true]: Inside while loop: rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) \u2014 poll DMA interrupt status register for channel 1 completion (bit 1).\n61. [If true]: Inside while loop: printf(\"polling irq; irq_status_reg rd_data =%0x\\n\", rd_data).\n62. [If true]: Exit while loop when (rd_data & 0x2) != 0x0.\n63. [If true]: rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) \u2014 additional read of DMA interrupt status after loop exit.\n64. [If true]: write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET, 0x2) \u2014 clear DMA IRQ bit 1.\n65. [If true]: printf(\"DEBUG : irq polling completed ch1; rd_data=%0x\\n\", rd_data).\n66. End of if block.\n67. End of inner for loop (next i iteration).\n68. End of outer for loop (next lane_num iteration, decrementing from 3 to 0).\n69. Call finish(0) \u2014 test completion with pass status.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
        "Meta Validation / Acceptance Criteria": "1. After calling csi2_enable_interrupt(), verify INT_ST_MAIN register (offset 0xC) was read to clear any pending interrupts. All status bits (status_int_st_phy_fatal[0], status_int_st_pkt_fatal[1], status_int_st_bndry_frame_fatal[2], status_int_st_seq_frame_fatal[3], status_int_st_crc_frame_fatal[4], status_int_st_pld_crc_fatal[5], status_int_st_data_id[6], status_int_st_ecc_corrected[7], status_int_st_phy[16], status_int_st_line[17], status_int_st_ipi_fatal[18]) are RC (read-clear).\n2. Verify INT_MSK_PHY_FATAL register (offset 0xE4) was written with 0x0000000f, enabling mask_phy_errsotsynchs_0[0], mask_phy_errsotsynchs_1[1], mask_phy_errsotsynchs_2[2], mask_phy_errsotsynchs_3[3].\n3. Verify INT_MSK_PKT_FATAL register (offset 0xF4) was written with 0x00000003, enabling mask_err_ecc_double[0] and mask_shorter_payload[1].\n4. Verify INT_MSK_PHY register (offset 0x114) was written with 0x000f000f, enabling mask_phy_errsoths_0[0], mask_phy_errsoths_1[1], mask_phy_errsoths_2[2], mask_phy_errsoths_3[3], mask_phy_erresc_0[16], mask_phy_erresc_1[17], mask_phy_erresc_2[18], mask_phy_erresc_3[19].\n5. Verify INT_MSK_LINE register (offset 0x134) was written with 0x000f000f, enabling mask_err_l_bndry_match_di0[0], mask_err_l_bndry_match_di1[1], mask_err_l_bndry_match_di2[2], mask_err_l_bndry_match_di3[3], mask_err_l_seq_di0[16], mask_err_l_seq_di1[17], mask_err_l_seq_di2[18], mask_err_l_seq_di3[19].\n6. Verify INT_MSK_BNDRY_FRAME_FATAL register (offset 0x284) was written with 0x0000ffff, enabling err_f_bndry_match_vc0[0] through err_f_bndry_match_vc15[15].\n7. Verify INT_MSK_SEQ_FRAME_FATAL register (offset 0x294) was written with 0x0000ffff, enabling err_f_seq_vc0[0] through err_f_seq_vc15[15].\n8. Verify INT_MSK_CRC_FRAME_FATAL register (offset 0x2A4) was written with 0x0000ffff, enabling err_frame_data_vc0[0] through err_frame_data_vc15[15].\n9. Verify INT_MSK_PLD_CRC_FATAL register (offset 0x2B4) was written with 0x0000ffff, enabling err_crc_vc0[0] through err_crc_vc15[15].\n10. Verify INT_MSK_DATA_ID register (offset 0x2C4) was written with 0x0000ffff, enabling err_id_vc0[0] through err_id_vc15[15].\n11. Verify INT_MSK_ECC_CORRECTED register (offset 0x2D4) was written with 0x0000ffff, enabling err_ecc_corrected_vc0[0] through err_ecc_corrected_vc15[15].\n12. Verify PHY_STOPSTATE register (offset 0x4C) polling exits when rd_data == 0x1000f, confirming phy_stopstatedata_0[0]=1, phy_stopstatedata_1[1]=1, phy_stopstatedata_2[2]=1, phy_stopstatedata_3[3]=1, and phy_stopstateclk[16]=1.\n13. For each lane_num iteration (3, 2, 1, 0): verify N_LANES register (offset 0x4) is written with lane_num value (n_lanes[2:0] field).\n14. For each lane_num iteration: verify write to 0xa0243ffc with value (lane_num + 1) triggers CSI-2 sequence.\n15. For each packet iteration: verify DMA channel 0 interrupt polling loop exits when (rd_data & 0x1) != 0x0, confirming control data transfer completion.\n16. For each packet iteration: verify DMA interrupt is cleared by writing 0x1 to gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET.\n17. For each packet iteration where (csi_ctrl_data & 0x3f) > 0xf: verify DMA channel 1 interrupt polling loop exits when (rd_data & 0x2) != 0x0, confirming CSI data transfer completion.\n18. For each packet iteration where (csi_ctrl_data & 0x3f) > 0xf: verify DMA interrupt is cleared by writing 0x2 to gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET.\n19. Verify word_count extraction: word_count = ((csi_ctrl_data >> 6) & 0xffff).\n20. Verify csi_data_size alignment: csi_data_size = (word_count % 8) ? (word_count / 8 + 1) * 8 : word_count.\n21. Verify test completes successfully via finish(0).",
        "Meta Headers": "#include <stdio.h>; #include <stdlib.h>; #include \"test_common.h\"; #include \"mipi_csi2.h\"",
        "Meta Macros": "#define GDMA_CSI2_DATA_DEST_ADDR2 0xE6040080\n#define GDMA_CTRL_DATA_DEST_ADDR2 0xE6040000\n#define LS_LE_EN 1\n#define TOTAL_FRAME 1\n#define VC_ID 3\n#define VRES 1080 (if GDMA0_FULL_MEM defined) or 3 (default)\n#define HRES 1920 (if GDMA0_FULL_MEM defined) or 64 (default)\n#define DATA_TYPE CSI2_RGB888",
        "Meta Arrays": "NA"
    }

    # Create workbook
    wb = openpyxl.Workbook()

    # ===== TestPlan Sheet =====
    ws_tp = wb.active
    ws_tp.title = 'TestPlan'

    tp_columns = [
        'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
        'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
        'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
        'Code Generation'
    ]

    # Header formatting
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell_alignment = Alignment(vertical='top', wrap_text=True)

    # Write TestPlan headers
    for col_idx, col_name in enumerate(tp_columns, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Write TestPlan data row
    tp_data = [
        row_data.get('Index', ''),
        row_data.get('SS / Module', ''),
        row_data.get('Feature', ''),
        row_data.get('Test Case Name', ''),
        row_data.get('Test Description', ''),
        row_data.get('Speed', ''),
        row_data.get('Mode', ''),
        row_data.get('Memory Start Offset', ''),
        row_data.get('Memory End Offset', ''),
        row_data.get('Remarks', ''),
        row_data.get('Test Steps / Procedure', ''),
        row_data.get('Impacted Registers', ''),
        row_data.get('Validation / Acceptance Criteria', ''),
        row_data.get('Code Generation', '')
    ]

    for col_idx, value in enumerate(tp_data, 1):
        cell = ws_tp.cell(row=2, column=col_idx, value=value)
        cell.alignment = cell_alignment

    # Freeze first row
    ws_tp.freeze_panes = 'A2'

    # Auto-size columns with max width cap
    for col_idx in range(1, len(tp_columns) + 1):
        max_length = len(str(tp_columns[col_idx - 1]))
        for row in ws_tp.iter_rows(min_row=2, max_row=ws_tp.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        if len(line) > max_length:
                            max_length = len(line)
        adjusted_width = min(max_length + 2, 60)
        ws_tp.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    # ===== MetaData Sheet =====
    ws_md = wb.create_sheet('MetaData')

    md_columns = [
        'Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
        'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
        'Meta Headers', 'Meta Macros', 'Meta Arrays'
    ]

    # Write MetaData headers
    for col_idx, col_name in enumerate(md_columns, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Write MetaData data row
    md_data = [
        row_data.get('Index', ''),
        row_data.get('Test Case Name', ''),
        row_data.get('Meta Test Description', ''),
        row_data.get('Meta Test Steps / Procedure', ''),
        row_data.get('Meta Impacted Registers', ''),
        row_data.get('Meta Validation / Acceptance Criteria', ''),
        row_data.get('Meta Headers', ''),
        row_data.get('Meta Macros', ''),
        row_data.get('Meta Arrays', '')
    ]

    for col_idx, value in enumerate(md_data, 1):
        cell = ws_md.cell(row=2, column=col_idx, value=value)
        cell.alignment = cell_alignment

    # Freeze first row
    ws_md.freeze_panes = 'A2'

    # Auto-size columns with max width cap
    for col_idx in range(1, len(md_columns) + 1):
        max_length = len(str(md_columns[col_idx - 1]))
        for row in ws_md.iter_rows(min_row=2, max_row=ws_md.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        if len(line) > max_length:
                            max_length = len(line)
        adjusted_width = min(max_length + 2, 60)
        ws_md.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = 'veryHidden'

    # Save workbook
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)
    wb.save(output_path)
    print(f'GENERATED:{filename}')
    print(f'PATH:{output_path}')
    print(f'SIZE:{os.path.getsize(output_path)}')

    # Validate
    wb2 = openpyxl.load_workbook(output_path)
    sheets = wb2.sheetnames
    print(f'SHEETS:{sheets}')
    assert 'TestPlan' in sheets
    assert 'MetaData' in sheets
    assert wb2['MetaData'].sheet_state == 'veryHidden'
    print('VALIDATION:PASSED')

    # Output base64 for GitHub upload
    import base64
    with open(output_path, 'rb') as f:
        b64content = base64.b64encode(f.read()).decode('utf-8')
    
    # Write base64 to a sidecar file for the workflow to pick up
    b64_path = output_path + '.b64'
    with open(b64_path, 'w') as f:
        f.write(b64content)
    print(f'B64_PATH:{b64_path}')
    print(f'FILENAME:{filename}')
    
    return filename, output_path, b64content

if __name__ == '__main__':
    filename, path, b64 = generate_workbook()
    print(f'SUCCESS: {filename}')
