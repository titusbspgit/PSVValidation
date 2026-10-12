#!/usr/bin/env python3
"""Agent 7 - MIPI DSI TestPlan XLSX Generator
Generates a genuine Office Open XML workbook (.xlsx) using openpyxl.
"""
import os
import json
from datetime import datetime, timezone, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── IST timestamp ──────────────────────────────────────────────
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")

# ── Input JSON data ────────────────────────────────────────────
json_data = [
  {
    "Index": "1",
    "SS / Module": "mipi_dsi",
    "Test Case Name": "mipi_dsi_basic_test",
    "Feature": "NA",
    "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include \"test_common.h\"\n#include \"mipi_dsi.h\"",
    "Meta Macros": "MIZAR_MIPI_DSI_DMAC_INTEN, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE, MIPI_DSI_SUBSYS_INTERRUPT_ENABLE_GDMA_INTR, MIZAR_MIPI_DSI_HOST_PHY_IF_CFG, MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME, MIZAR_MIPI_DSI_HOST_PCKHDL_CFG, MIZAR_MIPI_DSI_HOST_CLKMGR_CFG, MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL, MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1, MIZAR_MIPI_DSI_DMAC_DBGCMD, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK, MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR, MIZAR_MIPI_DSI_DMAC_INTMIS, MIZAR_MIPI_DSI_DMAC_INTCLR, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW, DSI_INTR_NO, RAM_BASE, DMA_SAR, DMA_DAR, DSI_WRITE_MEMORY_START",
    "Meta Arrays": "desc_t data_tdbdcb[2];\ndesc_t data_rebdcb[2];\ndesc_t cmd_tdbdcb[2];\ndesc_t cmd_rebdcb[2];",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase configures the MIPI DSI subsystem for a basic DBI (Display Bus Interface) write operation. It enables GIC and DMAC interrupts, configures PHY interface (n_lanes=3, phy_stop_wait_time=0x40), sets PCKHDL_CFG to 0x3d and CLKMGR_CFG to 0x107, disables DPI control, initializes PHY via phy_init(), configures DBI parameters (vcid=0x3, lut_size_conf=0x1, out_dbi_conf=0xb, in_dbi_conf=0x0, partitioning_en=0x1, allowed_cmd_size=0x7, num_of_pixel=40), programs DPI clock, sets up DMA descriptors for data channel (CH0) and command channel (CH1) with random data payload and write memory start command, launches both DMA channels via debug instruction registers (DBGINST0, DBGINST1, DBGCMD), waits for interrupt-driven completion by polling int_pend1, and finishes with err0 as the error count. The ISR (Default_IRQHandler) reads interrupt_mask to identify GDMA interrupt source, reads INTMIS for channel mask status, clears pending channels from int_pend1, clears DMAC interrupt via INTCLR, clears subsystem-level interrupt via interrupt_raw, and increments err0 on unexpected interrupts.",
    "Test Description": "Basic MIPI DSI DBI write test that configures PHY interface, DBI parameters, programs DMA descriptors for data and command channels, initiates DMA transfers via debug instruction registers, and validates completion through interrupt-driven handshake with error tracking.",
    "Meta Test Steps / Procedure": "1. GIC_EnableIRQ(DSI_INTR_NO);\n2. int_pend = 1;\n3. int_pend1 = 0x3;\n4. phy_stop_wait_time = 0x40;\n5. n_lanes = 3;\n6. write_reg(MIZAR_MIPI_DSI_DMAC_INTEN, 0x3);\n7. write_reg(MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE, MIPI_DSI_SUBSYS_INTERRUPT_ENABLE_GDMA_INTR);\n8. unsigned int phy_if_cfg = n_lanes;\n9. phy_if_cfg = set_data_mask(phy_if_cfg, MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME, phy_stop_wait_time);\n10. write_reg(MIZAR_MIPI_DSI_HOST_PHY_IF_CFG, phy_if_cfg);\n11. write_reg(MIZAR_MIPI_DSI_HOST_PCKHDL_CFG, 0x3d);\n12. write_reg(MIZAR_MIPI_DSI_HOST_CLKMGR_CFG, 0x107);\n13. write_reg(MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL, 0);\n14. phy_init();\n15. dbi_vcid = 0x3;\n16. load_cmd_or_data_to_sram = 1;\n17. lut_size_conf = 0x1;\n18. out_dbi_conf = 0xb;\n19. in_dbi_conf = 0x0;\n20. partitioning_en = 0x1;\n21. allowed_cmd_size = 0x7;\n22. wr_cmd_size = 192 + 1;\n23. unsigned int num_of_pixel = 40;\n24. pixel_num_bytes_wr_cmd_size_t pixel_attr = pixel_to_bytes_wr_cmd_size(num_of_pixel);\n25. wr_cmd_size = pixel_attr.wr_cmd_size;\n26. unsigned int num_bytes = pixel_attr.num_bytes;\n27. max_rd_pkt_size = 0x0;\n28. dcs_lw_tx = 0x0;\n29. dcs_sr_0p_tx = 0x0;\n30. dcs_sw_1p_tx = 0x0;\n31. dcs_sw_0p_tx = 0x0;\n32. gen_lw_tx = 0x0;\n33. gen_sr_2p_tx = 0x0;\n34. gen_sr_1p_tx = 0x0;\n35. gen_sr_0p_tx = 0x0;\n36. gen_sw_2p_tx = 0x0;\n37. gen_sw_1p_tx = 0x0;\n38. gen_sw_0p_tx = 0x0;\n39. ack_rqst_en = 0x0;\n40. tear_fx_en = 0x1;\n41. generic_vc_id = 0x2;\n42. dpi_clk_time_period = 16.012400;\n43. dpi_clk_freq = 1000000000/dpi_clk_time_period;\n44. program_dpi_clock(dpi_clk_freq);\n45. dbi_config();\n46. unsigned int num_descriptors = 1;\n47. ch0_desc_addr = RAM_BASE + 0x3F000;\n48. ch1_desc_addr = RAM_BASE + 0x3F800;\n49. ch0_desc_addr_act = ch0_desc_addr;\n50. ch1_desc_addr_act = ch1_desc_addr;\n51. for(unsigned int itter = 0; itter < num_descriptors; itter++) {\n52. data_tdbdcb[itter].addr = RAM_BASE + 0x10000;\n53. data_tdbdcb[itter].len = num_bytes;\n54. data_tdbdcb[itter].eop = (itter == (num_descriptors - 1));\n55. data_rebdcb[itter].addr = 0x10000000000;\n56. data_rebdcb[itter].len = num_bytes;\n57. cmd_tdbdcb[itter].addr = RAM_BASE + 0x0000;\n58. cmd_tdbdcb[itter].len = 4;\n59. cmd_tdbdcb[itter].eop = (itter == (num_descriptors - 1));\n60. cmd_rebdcb[itter].addr = 0x10000008000;\n61. cmd_rebdcb[itter].len = 4;\n62. DMAMOV(&ch0_desc_addr, DMA_SAR, data_tdbdcb[itter].addr);\n63. DMAMOV(&ch0_desc_addr, DMA_DAR, data_rebdcb[itter].addr);\n64. program_data_num_bytes(&ch0_desc_addr, data_tdbdcb[itter].len);\n65. DMAWMB(&ch0_desc_addr);\n66. DMASEV(&ch0_desc_addr, 0);\n67. DMAEND(&ch0_desc_addr);\n68. DMAMOV(&ch1_desc_addr, DMA_SAR, cmd_tdbdcb[itter].addr);\n69. DMAMOV(&ch1_desc_addr, DMA_DAR, cmd_rebdcb[itter].addr);\n70. program_data_num_bytes(&ch1_desc_addr, cmd_tdbdcb[itter].len);\n71. DMAWMB(&ch1_desc_addr);\n72. DMASEV(&ch1_desc_addr, 1);\n73. DMAEND(&ch1_desc_addr);\n74. load_rand_data(data_tdbdcb[itter]);\n75. load_wr_command(cmd_tdbdcb[itter].addr, num_bytes, DSI_WRITE_MEMORY_START);\n76. }\n77. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST0, 0x00A00000);\n78. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST1, ch0_desc_addr_act);\n79. write_reg(MIZAR_MIPI_DSI_DMAC_DBGCMD, 0x0);\n80. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST0, 0x01A00000);\n81. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST1, ch1_desc_addr_act);\n82. write_reg(MIZAR_MIPI_DSI_DMAC_DBGCMD, 0x0);\n83. while(int_pend1) { wait_on(10); }\n84. wait_on(10000);\n85. finish(err0);\n\nISR (void Default_IRQHandler()):\n86. unsigned int ch_mask_st;\n87. unsigned int dsi_subsys_mask_st;\n88. int_pend = 0;\n89. dsi_subsys_mask_st = read_reg(MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK);\n90. if(dsi_subsys_mask_st == MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR) {\n91. ch_mask_st = read_reg(MIZAR_MIPI_DSI_DMAC_INTMIS);\n92. if(ch_mask_st) {\n93. int_pend1 = int_pend1 & (~ch_mask_st);\n94. write_reg(MIZAR_MIPI_DSI_DMAC_INTCLR, ch_mask_st);\n95. } else {\n96. err0++;\n97. }\n98. write_reg(MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW, dsi_subsys_mask_st);\n99. } else {\n100. err0++;\n101. }\n102. GIC_ClearIRQ(DSI_INTR_NO);\n103. int_pend = 1;",
    "Test Steps / Procedure": "1. Enable GIC interrupt for DSI.\n2. Initialize interrupt pending flags (int_pend=1, int_pend1=0x3).\n3. Set PHY stop wait time to 0x40 and number of lanes to 3.\n4. Enable DMAC interrupts by writing 0x3 to INTEN register.\n5. Enable GDMA interrupt at subsystem level via interrupt_enable register.\n6. Configure PHY_IF_CFG register with n_lanes and phy_stop_wait_time.\n7. Configure PCKHDL_CFG register with 0x3d.\n8. Configure CLKMGR_CFG register with 0x107.\n9. Disable DPI by writing 0 to dpi_control register.\n10. Initialize PHY.\n11. Configure DBI parameters (vcid=0x3, LUT size, output/input DBI config, partitioning, command size, pixel count=40).\n12. Set DCS/generic packet transmission parameters to 0, enable tear effect.\n13. Program DPI clock based on 16.012400ns period.\n14. Apply DBI configuration.\n15. Set up DMA descriptor addresses for data channel (CH0) and command channel (CH1).\n16. For each descriptor: populate source/destination addresses, lengths, EOP flags; program DMA move, write memory barrier, send event, and end instructions; load random data payload; load write memory start command.\n17. Launch CH0 DMA by writing 0x00A00000 to DBGINST0, CH0 descriptor address to DBGINST1, and 0x0 to DBGCMD.\n18. Launch CH1 DMA by writing 0x01A00000 to DBGINST0, CH1 descriptor address to DBGINST1, and 0x0 to DBGCMD.\n19. Poll int_pend1 with wait_on(10) until both channels complete (cleared by ISR).\n20. Wait additional 10000 cycles.\n21. Finish test with err0 as error count.\n22. ISR: Read interrupt_mask to identify GDMA source, read INTMIS for channel status, clear pending channels from int_pend1, clear DMAC interrupt via INTCLR, clear subsystem interrupt via interrupt_raw, increment err0 on unexpected interrupts, clear GIC IRQ.",
    "Meta Impacted Registers": "INTEN, interrupt_enable, PHY_IF_CFG, PCKHDL_CFG, CLKMGR_CFG, dpi_control, DBGINST0, DBGINST1, DBGCMD, interrupt_mask, INTMIS, INTCLR, interrupt_raw",
    "Impacted Registers": "INTEN, interrupt_enable, PHY_IF_CFG, PCKHDL_CFG, CLKMGR_CFG, dpi_control, DBGINST0, DBGINST1, DBGCMD, interrupt_mask, INTMIS, INTCLR, interrupt_raw",
    "Meta Validation / Acceptance Criteria": "1. ISR reads dsi_subsys_mask_st from MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK and compares against MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR; if not equal, err0 is incremented and printf(\"ERROR1: Unexpected interrupt dsi_subsys_mask_st = 0x%x\\n\", dsi_subsys_mask_st) is called.\n2. ISR reads ch_mask_st from MIZAR_MIPI_DSI_DMAC_INTMIS; if ch_mask_st is 0 (no channel interrupt pending), err0 is incremented and printf(\"ERROR2: Unexpected interrupt ch_mask_st = 0x%x \\n\", ch_mask_st) is called.\n3. If ch_mask_st is non-zero, int_pend1 = int_pend1 & (~ch_mask_st) clears the completed channel bits.\n4. DMAC interrupt is cleared by write_reg(MIZAR_MIPI_DSI_DMAC_INTCLR, ch_mask_st).\n5. Subsystem interrupt is cleared by write_reg(MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW, dsi_subsys_mask_st).\n6. Main loop exits when int_pend1 becomes 0 (both CH0 and CH1 interrupts received and cleared).\n7. Test passes when finish(err0) is called with err0 == 0 (no unexpected interrupts occurred).",
    "Validation / Acceptance Criteria": "1. The ISR must identify the interrupt source as GDMA via the interrupt_mask register; unexpected sources increment the error count.\n2. The ISR must read a non-zero channel mask from INTMIS; a zero value indicates an unexpected interrupt and increments the error count.\n3. Both DMA channels (CH0 and CH1) must complete and trigger interrupts that clear the corresponding bits in int_pend1.\n4. DMAC interrupts must be properly cleared via the INTCLR register.\n5. Subsystem-level interrupts must be properly cleared via the interrupt_raw register.\n6. The test passes when err0 is 0 at finish, indicating no unexpected interrupts occurred during the DBI write transfer.",
    "Remarks": "The testcase uses two DMA channels: CH0 for data transfer and CH1 for command transfer. int_pend1 is initialized to 0x3 representing both channels pending. The ISR clears individual channel bits as their interrupts arrive. The test uses desc_t typed descriptor arrays (data_tdbdcb, data_rebdcb, cmd_tdbdcb, cmd_rebdcb) each of size 2, though only 1 descriptor is used (num_descriptors=1). DPI clock period is set to 16.012400ns. The data destination address is 0x10000000000 and command destination address is 0x10000008000. extern int int_pend is declared as extern int; unsigned int int_pend1 and unsigned int err0 are file-scope unsigned int variables.",
    "Code Generation": ""
  }
]

# ── Configuration ──────────────────────────────────────────────
IP_NAME = "MIPI_DSI"
filename = f"{IP_NAME}_TestPlan_{timestamp}.xlsx"
output_dir = "Test_Output/MIPI/TestPlan"
os.makedirs(output_dir, exist_ok=True)
filepath = os.path.join(output_dir, filename)

# ── Column definitions ─────────────────────────────────────────
tp_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name",
    "Test Description", "Speed", "Mode", "Memory Start Offset",
    "Memory End Offset", "Remarks", "Test Steps / Procedure",
    "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

md_columns = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

# ── Create workbook ────────────────────────────────────────────
wb = openpyxl.Workbook()

# TestPlan sheet
ws_tp = wb.active
ws_tp.title = "TestPlan"

# MetaData sheet
ws_md = wb.create_sheet("MetaData")

# ── Formatting ─────────────────────────────────────────────────
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_alignment = Alignment(wrap_text=True, vertical="top")

# ── Write TestPlan headers ─────────────────────────────────────
for col_idx, col_name in enumerate(tp_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# ── Write MetaData headers ─────────────────────────────────────
for col_idx, col_name in enumerate(md_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# ── Populate TestPlan data ─────────────────────────────────────
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(tp_columns, 1):
        value = row_data.get(col_name, "")
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# ── Populate MetaData data ─────────────────────────────────────
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(md_columns, 1):
        value = row_data.get(col_name, "")
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# ── Freeze first row ───────────────────────────────────────────
ws_tp.freeze_panes = "A2"
ws_md.freeze_panes = "A2"

# ── Auto-size columns ──────────────────────────────────────────
MAX_WIDTH = 60
for ws in [ws_tp, ws_md]:
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                min_row=1, max_row=ws.max_row):
            for cell in row:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        max_len = max(max_len, len(line))
        adjusted = min(max_len + 2, MAX_WIDTH)
        if adjusted < 10:
            adjusted = 10
        ws.column_dimensions[col_letter].width = adjusted

# ── MetaData sheet: veryHidden ─────────────────────────────────
ws_md.sheet_state = "veryHidden"

# ── Save workbook ──────────────────────────────────────────────
wb.save(filepath)
wb.close()

# ── Post-save validation ──────────────────────────────────────
file_exists = os.path.exists(filepath)
file_size = os.path.getsize(filepath) if file_exists else 0
validation_passed = False

if file_exists and file_size > 0:
    vwb = openpyxl.load_workbook(filepath)
    sheet_names = vwb.sheetnames
    has_tp = "TestPlan" in sheet_names
    has_md = "MetaData" in sheet_names
    tp_rows = vwb["TestPlan"].max_row - 1
    md_rows = vwb["MetaData"].max_row - 1
    vwb.close()
    if has_tp and has_md:
        validation_passed = True

print(f"FILENAME={filename}")
print(f"FILEPATH={filepath}")
print(f"FILE_EXISTS={file_exists}")
print(f"FILE_SIZE={file_size}")
print(f"VALIDATION={'PASSED' if validation_passed else 'FAILED'}")
print(f"TP_ROWS={tp_rows}")
print(f"MD_ROWS={md_rows}")
print(f"TIMESTAMP={timestamp}")
print(f"\nGeneration complete: {filepath}")
