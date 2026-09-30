#!/usr/bin/env python3
"""Agent 7 - MIPI_CSI TestPlan XLSX Generator (Regeneration)
Generates a genuine openpyxl XLSX workbook with TestPlan and MetaData sheets.
"""
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

# ============================================================
# JSON DATA - 2 testcase rows
# ============================================================
json_data = [
    {
        "Index": "1",
        "SS / Module": "MIPI_CSI",
        "Test Case Name": "mipi_csi2_dphy_lanes_test",
        "Feature": "D-PHY Lane Configuration",
        "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>; \"pcie.h\"",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates MIPI CSI2 D-PHY lane configuration by iterating through different lane counts. The test writes to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to configure virtual channel settings, writes to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to set control data, writes to MIZAR_MIPI_CSI2_HOST_N_LANES to configure the number of active D-PHY lanes, and polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE to verify PHY stop state. It reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to check the main interrupt status. It masks all interrupt sources by writing to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, and MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED. A write to 0xa0243ffc and a read from 0xE6001000 are also performed as part of the DMA or system-level configuration.",
        "Test Description": "This testcase validates MIPI CSI2 D-PHY lane configuration by iterating through different lane counts. The test configures the virtual channel register and control data register in the CSI2 subsystem wrapper. It writes the number of active D-PHY lanes to the N_LANES register and polls the PHY_STOPSTATE register to verify that the PHY has entered the stop state for the configured lanes. The main interrupt status register INT_ST_MAIN is read to check for any pending interrupts. All interrupt mask registers are configured to mask interrupt sources including PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected interrupts. Additional system-level register accesses are performed for DMA or configuration purposes.",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to configure virtual channel settings for the CSI2 subsystem wrapper. 2. Write to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to set the control data configuration. 3. Write to MIZAR_MIPI_CSI2_HOST_N_LANES to configure the number of active D-PHY data lanes. 4. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE to wait until the PHY stop state is achieved for the configured lane count. 5. Write to 0xa0243ffc for system-level or DMA configuration. 6. Read from 0xE6001000 for system-level status or configuration verification. 7. Read MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to check the main interrupt status register. 8. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL to mask PHY fatal interrupts. 9. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL to mask packet fatal interrupts. 10. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY to mask PHY interrupts. 11. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE to mask line interrupts. 12. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL to mask boundary frame fatal interrupts. 13. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL to mask sequence frame fatal interrupts. 14. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL to mask CRC frame fatal interrupts. 15. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL to mask payload CRC fatal interrupts. 16. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID to mask data ID interrupts. 17. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED to mask ECC corrected interrupts. 18. Repeat steps 1-17 for each lane configuration iteration (e.g., 4 lanes down to 1 lane). 19. Validate that no errors occurred during the lane configuration iterations.",
        "Test Steps / Procedure": "1. Configure the virtual channel register in the CSI2 subsystem wrapper with the desired virtual channel settings. 2. Configure the control data register in the CSI2 subsystem wrapper. 3. Write the number of active D-PHY data lanes to the N_LANES register. 4. Poll the PHY_STOPSTATE register until the PHY stop state is confirmed for the configured lane count. 5. Perform system-level DMA or configuration register access. 6. Read a system-level status register for configuration verification. 7. Read the INT_ST_MAIN register to check for any pending main interrupt status. 8. Mask all CSI2 host interrupt sources by writing to the interrupt mask registers: INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED. 9. Repeat the lane configuration and validation sequence for each supported lane count. 10. Verify that the test completes without errors for all lane configurations.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
        "Impacted Registers": "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
        "Meta Validation / Acceptance Criteria": "The test validates that MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE returns the expected stop state value for each configured lane count after writing to MIZAR_MIPI_CSI2_HOST_N_LANES. MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN is read to confirm no unexpected interrupts are asserted. All interrupt mask registers are written to ensure interrupt sources are properly masked during the test. The polling loop on MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE must complete successfully without timeout for each lane configuration. The test passes if all lane configurations complete without errors.",
        "Validation / Acceptance Criteria": "The PHY_STOPSTATE register must reflect the correct stop state for each configured lane count after writing the lane configuration to N_LANES. The INT_ST_MAIN register must not indicate unexpected interrupt conditions. All interrupt mask registers must be successfully written to mask the respective interrupt sources. The polling of PHY_STOPSTATE must complete without timeout for each lane configuration. The test passes if all D-PHY lane configurations from maximum to minimum lanes complete successfully without errors.",
        "Remarks": "Two register addresses (used for system-level or DMA configuration) could not be mapped to canonical register names from the provided specification documents. The source file retrieved from the repository did not contain MIPI CSI2-specific code; the register-level analysis was derived from upstream agent outputs. The SOFT_RST_REG_ADDRESS macro was excluded per instructions."
    },
    {
        "Index": "2",
        "SS / Module": "MIPI_CSI",
        "Test Case Name": "mipi_csi2_test_pattern_generator",
        "Feature": "Internal Test Pattern Generator (PPI PG)",
        "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>; \"pcie.h\"",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates the MIPI CSI2 internal test pattern generator (PPI PG) functionality. The test configures the pattern generator vertical resolution by writing 0x10 to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES, configures horizontal resolution by writing 0x70140 to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES, configures the pattern generator mode and data type by writing 0xe401 to MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG, and enables the pattern generator by writing 1 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE. The CSI2 subsystem wrapper is configured by writing virtual channel settings to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL and writing 0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA. The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE to wait for the PHY stop state. DMA address registers are configured by writing to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA (value 0x100), MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION (value 0x0), MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA (value 0x0), and MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION (value 0x0). A write to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 with value 0x1 is performed to enable clock gating for the CSI PHY. The pattern generator is then disabled by writing 0 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE. The DMA interrupt status is polled via gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET. The main interrupt status is read from MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN. All interrupt mask registers are written to mask interrupt sources: MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, and MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED.",
        "Test Description": "This testcase validates the MIPI CSI2 internal test pattern generator (PPI PG) functionality. The test configures the pattern generator with a vertical resolution of 16 lines and a horizontal resolution with RGB888 data type, then enables the pattern generator via the PPI_PG_ENABLE register. The CSI2 subsystem wrapper is configured by writing virtual channel settings to the virtual_channel register and clearing the control_data register. The PHY_STOPSTATE register is polled to confirm the PHY has entered stop state. DMA address registers for channel 0 read and write paths are configured via dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, and dma_m0_addr_aw_ch0_Instruction. A subsystem wrapper register is written to enable clock gating for the CSI PHY. The pattern generator is subsequently disabled. The DMA interrupt masked status is polled to confirm DMA transfer completion. The INT_ST_MAIN register is read to check the main interrupt status. All CSI2 host interrupt mask registers are configured to mask interrupt sources including PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected interrupts.",
        "Meta Test Steps / Procedure": "1. Write 0x10 to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES to set the pattern generator vertical resolution to 16 lines. 2. Write 0x70140 to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES to set the pattern generator horizontal resolution (HRES=320, data type RGB888). 3. Write 0xe401 to MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG to configure the pattern generator mode and data type settings. 4. Write 1 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE to enable the internal test pattern generator. 5. Write virtual channel ID configuration to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to set the virtual channel mapping in the CSI2 subsystem wrapper. 6. Write 0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to clear the control data register. 7. Read MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE to check the current PHY stop state. 8. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while loop until the expected stop state value is achieved for the configured lanes. 9. Write 0x100 to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA to configure the DMA channel 0 read address for data. 10. Write 0x0 to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION to configure the DMA channel 0 read address for instructions. 11. Write 0x0 to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA to configure the DMA channel 0 write address for data. 12. Write 0x0 to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION to configure the DMA channel 0 write address for instructions. 13. Write 0x1 to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 to enable clock gating for the CSI PHY. 14. Write 0 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE to disable the internal test pattern generator after data generation. 15. Poll gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET to wait for the DMA interrupt masked status indicating transfer completion. 16. Read MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to check the main interrupt status register for any pending interrupts. 17. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL to mask PHY fatal interrupts. 18. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL to mask packet fatal interrupts. 19. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY to mask PHY interrupts. 20. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE to mask line interrupts. 21. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL to mask boundary frame fatal interrupts. 22. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL to mask sequence frame fatal interrupts. 23. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL to mask CRC frame fatal interrupts. 24. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL to mask payload CRC fatal interrupts. 25. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID to mask data ID interrupts. 26. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED to mask ECC corrected interrupts. 27. Validate that the DMA transfer completed successfully and no unexpected errors occurred.",
        "Test Steps / Procedure": "1. Configure the pattern generator vertical resolution to 16 lines by writing to the PPI_PG_PATTERN_VRES register. 2. Configure the pattern generator horizontal resolution and data type (RGB888, HRES=320) by writing to the PPI_PG_PATTERN_HRES register. 3. Configure the pattern generator mode and data type settings by writing to the PPI_PG_CONFIG register. 4. Enable the internal test pattern generator by writing to the PPI_PG_ENABLE register. 5. Configure the virtual channel mapping in the CSI2 subsystem wrapper by writing to the virtual_channel register. 6. Clear the control_data register in the CSI2 subsystem wrapper. 7. Poll the PHY_STOPSTATE register until the PHY stop state is confirmed for the configured lanes. 8. Configure DMA channel 0 read and write address registers by writing to dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, and dma_m0_addr_aw_ch0_Instruction. 9. Enable clock gating for the CSI PHY by writing to the subsystem wrapper clock gating register. 10. Disable the internal test pattern generator by clearing the PPI_PG_ENABLE register. 11. Poll the DMA interrupt masked status to confirm DMA transfer completion. 12. Read the INT_ST_MAIN register to verify the main interrupt status. 13. Mask all CSI2 host interrupt sources by writing to the interrupt mask registers: INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED. 14. Verify that the test completes without errors.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
        "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
        "Meta Validation / Acceptance Criteria": "The test validates that MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE returns the expected stop state value after the pattern generator is configured and enabled. The polling loop on MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE must complete successfully without timeout. The DMA interrupt masked status (gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) must indicate successful DMA transfer completion. MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN is read to confirm no unexpected interrupts are asserted after the pattern generation and DMA transfer. All interrupt mask registers are written to ensure interrupt sources are properly masked. The pattern generator is disabled after data generation by writing 0 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE. The test passes if the pattern generator configuration, DMA data path, and interrupt masking all complete without errors.",
        "Validation / Acceptance Criteria": "The PHY_STOPSTATE register must reflect the correct stop state after the pattern generator is configured and enabled. The polling of PHY_STOPSTATE must complete without timeout. The DMA transfer must complete successfully as indicated by the DMA interrupt masked status. The INT_ST_MAIN register must not indicate unexpected interrupt conditions after pattern generation and DMA transfer. All interrupt mask registers must be successfully written to mask the respective interrupt sources. The pattern generator must be successfully disabled after data generation via the PPI_PG_ENABLE register. The test passes if the internal test pattern generator configuration, DMA data path verification, and interrupt masking all complete without errors.",
        "Remarks": "The source file retrieved from the repository contained PCIe test code rather than MIPI CSI2 test pattern generator code; the register-level analysis was derived from upstream agent outputs which analyzed the correct MIPI CSI2 source. One Agent 2 token (MIZAR_MIPI_CSI2_RB_REG_BASE) could not be mapped to a canonical register name by Agent 4 as it is a base address macro used in an expression with offset to target the clock gating register. The SOFT_RST_REG_ADDRESS macro was excluded per instructions. DMA interrupt polling uses a variable base address (gdma_reg_base) which is not captured as a standalone Agent 2 token."
    }
]

# ============================================================
# COLUMN DEFINITIONS
# ============================================================
TESTPLAN_COLUMNS = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

METADATA_COLUMNS = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

def create_workbook():
    """Create and return a formatted openpyxl workbook."""
    wb = Workbook()

    # --- TestPlan Sheet ---
    ws_tp = wb.active
    ws_tp.title = "TestPlan"

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_align = Alignment(wrap_text=True, vertical="top")

    # Write TestPlan headers
    for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_align

    # Write TestPlan data rows
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
            value = row_data.get(col_name, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align

    # Freeze first row
    ws_tp.freeze_panes = "A2"

    # Auto-size columns
    for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(json_data) + 2):
            cell_val = str(ws_tp.cell(row=row_idx, column=col_idx).value or "")
            max_len = max(max_len, min(len(cell_val), 80))
        col_letter = get_column_letter(col_idx)
        ws_tp.column_dimensions[col_letter].width = min(max_len + 4, 60)

    # --- MetaData Sheet ---
    ws_md = wb.create_sheet(title="MetaData")

    # Write MetaData headers
    for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_align

    # Write MetaData data rows
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
            value = row_data.get(col_name, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align

    # Freeze first row
    ws_md.freeze_panes = "A2"

    # Auto-size columns
    for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(json_data) + 2):
            cell_val = str(ws_md.cell(row=row_idx, column=col_idx).value or "")
            max_len = max(max_len, min(len(cell_val), 80))
        col_letter = get_column_letter(col_idx)
        ws_md.column_dimensions[col_letter].width = min(max_len + 4, 60)

    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = "veryHidden"

    return wb


def main():
    # IST timestamp
    IST = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(IST)
    timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
    filename = f"MIPI_CSI_TestPlan_{timestamp}.xlsx"

    output_dir = os.environ.get("OUTPUT_DIR", ".")
    filepath = os.path.join(output_dir, filename)

    print(f"Generating workbook: {filename}")
    print(f"Output path: {filepath}")

    wb = create_workbook()
    wb.save(filepath)
    print(f"Workbook saved: {filepath}")

    # Validate
    file_size = os.path.getsize(filepath)
    print(f"File size: {file_size} bytes")
    assert file_size > 0, "File size is 0!"

    wb2 = load_workbook(filepath)
    sheet_names = wb2.sheetnames
    print(f"Sheet names: {sheet_names}")
    assert "TestPlan" in sheet_names, "TestPlan sheet missing!"
    assert "MetaData" in sheet_names, "MetaData sheet missing!"

    tp_rows = wb2["TestPlan"].max_row - 1
    md_rows = wb2["MetaData"].max_row - 1
    print(f"TestPlan data rows: {tp_rows}")
    print(f"MetaData data rows: {md_rows}")
    assert tp_rows == 2, f"Expected 2 TestPlan rows, got {tp_rows}"
    assert md_rows == 2, f"Expected 2 MetaData rows, got {md_rows}"

    md_state = wb2["MetaData"].sheet_state
    print(f"MetaData sheet_state: {md_state}")
    assert md_state == "veryHidden", f"Expected veryHidden, got {md_state}"

    print(f"VALIDATION=PASSED")
    print(f"FILENAME={filename}")
    print(f"FILEPATH={filepath}")
    print(f"FILESIZE={file_size}")

    return filename, filepath


if __name__ == "__main__":
    main()
