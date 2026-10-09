#!/usr/bin/env python3
"""Generate MIPI_CSI TestPlan Excel workbook using openpyxl."""
import os, sys, datetime, base64, json
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

def generate():
    ist = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
    now = datetime.datetime.now(ist)
    ts = now.strftime('%Y%m%d_%H%M%S')
    filename = f'MIPI_CSI_TestPlan_{ts}.xlsx'

    # Data
    row = {
        'Index': '1',
        'SS_Module': 'mipi_csi2_subsys',
        'Feature': 'DPHY Lanes',
        'Test_Case_Name': 'mipi_csi2_dphy_lanes_test',
        'Test_Description': 'Verify MIPI CSI-2 D-PHY lane configuration by iterating through all supported lane counts (4 down to 1), enabling CSI-2 interrupts, initializing the D-PHY, waiting for PHY stop state, configuring virtual channels, and performing DMA-based control data and CSI image data transfers for each packet in a frame at each lane configuration.',
        'Speed': 'NA',
        'Mode': 'NA',
        'Memory_Start_Offset': '0xE6000000; 0xE6000500; 0xE6001000; 0xE6002000; 0xE6A00000',
        'Memory_End_Offset': 'NA',
        'Remarks': 'The functions snps_phy_init(), dma_trnsfr_instn_preload(), DMAGO_CSI(), and finish() are external functions not defined within the testcase folder; their implementations are assumed to be provided by the test framework or included via test_common.h and mipi_csi2.h. The macros MIPI_CSI2_DMA_INTEN_OFFSET, MIPI_CSI2_DMA_INTMIS_OFFSET, and MIPI_CSI2_DMA_INTCLR_OFFSET are expected to be defined in mipi_csi2.h. The GDMA path selection (GDMA0_PATH, GDMA1_PATH, GDMA2_PATH, GDMA3_PATH) and GDMA0_FULL_MEM are compile-time configuration macros that determine the virtual channel register value and resolution parameters. Two hardcoded addresses (0xa0243ffc and 0xE6001000) remain unresolved in the Agent 4 register mapping. The virtual_channel and control_data registers are written twice in sequence, which appears intentional in the source. The DPHY_LANES_TEST conditional compilation guard must be defined for the test-specific macros to take effect.',
        'Test_Steps': """1. Enable all CSI-2 host controller interrupts by reading the main interrupt status register to clear pending interrupts, then writing mask values to enable PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected interrupt categories.
2. Configure the virtual channel register based on the selected GDMA path and enable control data transfer.
3. Repeat the virtual channel and control data enable writes a second time.
4. Initialize the SNPS D-PHY.
5. Poll the PHY_STOPSTATE register until all 4 data lanes and the clock lane report stop state (expected value 0x1000f).
6. Set DMA program counter addresses for channel 0 and channel 1.
7. Iterate through lane configurations from 4 lanes (lane_num=3) down to 1 lane (lane_num=0).
8. For each lane configuration, write the lane count to the N_LANES register and trigger the CSI-2 sequence.
9. For each packet in the frame (total packets = VRES*3 + 2), perform the following:
 a. Enable DMA interrupts for both channels.
 b. Preload DMA channel 0 transfer instructions for control data transfer (8 bytes from source to destination).
 c. Start DMA channel 0 and poll the DMA interrupt status register until channel 0 transfer completes.
 d. Clear the DMA channel 0 interrupt.
 e. Read the control data and check the data type field.
 f. If the data type indicates CSI image data, compute the word count and aligned transfer size, preload DMA channel 1 transfer instructions, start DMA channel 1, poll for channel 1 completion, and clear the channel 1 interrupt.
10. After all lane configurations are tested, signal test completion.""",
        'Impacted_Registers': 'virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED',
        'Validation': """1. All CSI-2 interrupt mask registers are correctly programmed with the specified enable values.
2. The PHY_STOPSTATE register confirms all 4 data lanes and the clock lane are in stop state before proceeding.
3. The N_LANES register is correctly configured for each lane count iteration (4 lanes down to 1 lane).
4. DMA channel 0 control data transfers complete successfully for each packet, confirmed by DMA interrupt status polling.
5. DMA channel 1 CSI image data transfers complete successfully when the control data indicates a valid CSI data type.
6. DMA interrupts are properly cleared after each transfer completion.
7. The test completes with a pass status via finish(0) after all lane configurations are exercised.""",
        'Meta_Test_Description': """This testcase validates MIPI CSI-2 D-PHY lane configuration by iterating through all lane counts (4 lanes down to 1 lane). For each lane configuration, the test programs the CSI-2 host controller N_LANES register, triggers a CSI-2 sequence via a write to hardcoded address 0xa0243ffc, and then performs DMA-based control data and CSI image data transfers for each packet in a frame. The test first enables all CSI-2 interrupts via csi2_enable_interrupt(), configures the virtual channel register and control data enable, initializes the SNPS D-PHY via snps_phy_init(), and polls PHY_STOPSTATE until all 4 data lanes and clock lane are in stop state (value 0x1000f). For each lane count iteration (lane_num=3 down to 0), the test computes cntrl_pkt_cnt = (VRES*3)+2 and loops through that many packets. Each packet iteration programs DMA interrupt enable, preloads DMA channel 0 transfer instructions for control data (src=0x8000, dest=0xE6001000, size=8 bytes), issues DMAGO_CSI for channel 0, polls DMA interrupt status register (MIPI_CSI2_DMA_INTMIS_OFFSET) until bit 0 is set, clears DMA interrupt, reads control data from 0xE6001000, and if the data type field (bits[5:0]) > 0xf, computes word_count from bits[21:6], calculates csi_data_size aligned to 8 bytes, preloads DMA channel 1 transfer instructions for CSI data (src=0x0000, dest=0xE6002000, size=csi_data_size), issues DMAGO_CSI for channel 1, polls DMA interrupt status until bit 1 is set, reads DMA interrupt status again, and clears DMA interrupt bit 1. The test completes with finish(0).""",
        'Meta_Test_Steps': """1. Enter test_case() function.
2. Declare local variables: rx_desc, tx_desc (long long int), gdma_tx_trnsfr_size, gdma_trnsfr_size (long long int), cntrl_pkt_cnt (int).
3. printf("start line\\n").
4. Call csi2_enable_interrupt() to enable CSI-2 interrupts.
5. [Inside csi2_enable_interrupt()]: Declare local int rd_data.
6. [Inside csi2_enable_interrupt()]: rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN) \u2014 read INT_ST_MAIN register to clear interrupts.
7. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, 0x0000000f) \u2014 enable phy_fatal interrupts for all 4 lanes.
8. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, 0x00000003) \u2014 enable pkt_fatal interrupts (err_ecc_double and shorter_payload).
9. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, 0x000f000f) \u2014 enable phy interrupts (errsoths[3:0] and erresc[19:16]).
10. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, 0x000f000f) \u2014 enable line interrupts (l_bndry_match[3:0] and l_seq[19:16]).
11. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, 0x0000ffff) \u2014 enable boundary frame fatal interrupts for all 16 VCs.
12. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, 0x0000ffff) \u2014 enable seq frame fatal interrupts for all 16 VCs.
13. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, 0x0000ffff) \u2014 enable CRC frame fatal interrupts for all 16 VCs.
14. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, 0x0000ffff) \u2014 enable payload CRC fatal interrupts for all 16 VCs.
15. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, 0x0000ffff) \u2014 enable data ID interrupts for all 16 VCs.
16. [Inside csi2_enable_interrupt()]: write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED, 0x0000ffff) \u2014 enable ECC corrected interrupts for all 16 VCs.
17. Return from csi2_enable_interrupt().
18. Conditional compilation: set vcid_csi2_wrap_reg based on GDMA path. For GDMA3_PATH: vcid_csi2_wrap_reg = VC_ID. For GDMA2_PATH: vcid_csi2_wrap_reg = (VC_ID << 4). For GDMA1_PATH: vcid_csi2_wrap_reg = (VC_ID << 8). For GDMA0_PATH: vcid_csi2_wrap_reg = (VC_ID << 12), printf("VC_ID=%d\\n", VC_ID), gdma_path = 0.
19. Set gdma_reg_base = 0xE6A00000.
20. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg) \u2014 first write to CSI-2 virtual channel register.
21. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 1) \u2014 first enable control data transfer.
22. printf("vcid_csi2_wrap_reg=%0x\\n", vcid_csi2_wrap_reg).
23. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg) \u2014 second write to CSI-2 virtual channel register.
24. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 1) \u2014 second enable control data transfer.
25. Call snps_phy_init() \u2014 D-PHY initialization sequence (external function, implementation not in testcase folder).
26. rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) \u2014 first read of PHY stop state.
27. Enter while loop: condition !(rd_data == 0x1000f).
28. Inside while loop: rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) \u2014 poll PHY stop state until all 4 data lanes (bits[3:0]) and clock lane (bit[16]) are in stop state (value 0x1000f).
29. Exit while loop when rd_data == 0x1000f.
30. Set ch0_pc = 0xE6000000.
31. Set ch1_pc = 0xE6000500.
32. Enter outer for loop: lane_num = 3 down to 0 (lane_num >= 0, lane_num--).
33. write_reg(MIZAR_MIPI_CSI2_HOST_N_LANES, lane_num) \u2014 configure number of lanes in CSI-2 controller (lane_num=3 means 4 lanes, lane_num=0 means 1 lane).
34. Compute cntrl_pkt_cnt = ((VRES * 3) + 2).
35. printf("DEBUG: cntrl_pkt_cnt=%d\\n", cntrl_pkt_cnt).
36. write_reg(0xa0243ffc, (lane_num + 1)) \u2014 trigger CSI-2 sequence with lane count value.
37. Enter inner for loop: i = 0 to cntrl_pkt_cnt - 1.
38. Set ch0_preload_loc = ch0_pc.
39. Set ch1_preload_loc = ch1_pc.
40. write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTEN_OFFSET, 0x3) \u2014 enable dma_irq[1] and dma_irq[0].
41. Call dma_trnsfr_instn_preload(ch0_preload_loc, gdma_reg_base, 0x8000, 0xE6001000, 8, 0) \u2014 preload DMA channel 0 transfer instructions for control data.
42. Call DMAGO_CSI(gdma_reg_base, ch0_pc, 0x0) \u2014 start DMA channel 0 (ch_num=0).
43. Set rd_data = 0.
44. Enter while loop: condition (rd_data & 0x1) == 0x0.
45. Inside while loop: rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) \u2014 poll DMA interrupt status register for channel 0 completion (bit 0).
46. Inside while loop: printf("polling irq; irq_status_reg rd_data =%0x\\n", rd_data).
47. Exit while loop when (rd_data & 0x1) != 0x0.
48. write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET, 0x1) \u2014 clear DMA IRQ bit 0.
49. rd_data = read_reg(gdma_reg_base + 0x28) \u2014 read DMA interrupt status register at offset 0x28.
50. printf("DEBUG : irq polling completed ch0; rd_data = %0x \\n", rd_data).
51. csi_ctrl_data = read_reg(0xE6001000) \u2014 read control data from address 0xE6001000.
52. printf("DEBUG: csi_ctrl_data=%0x\\n", csi_ctrl_data).
53. Check condition: if ((csi_ctrl_data & 0x3f) > 0xf) \u2014 check if data type field indicates CSI image data.
54. [If true]: word_count = ((csi_ctrl_data >> 6) & 0xffff) \u2014 extract word count from control data bits[21:6].
55. [If true]: csi_data_size = (word_count % 8) ? (word_count / 8 + 1) * 8 : word_count \u2014 align transfer size to 8-byte boundary.
56. [If true]: Call dma_trnsfr_instn_preload(ch1_preload_loc, gdma_reg_base, 0x0000, 0xE6002000, csi_data_size, 1) \u2014 preload DMA channel 1 transfer instructions for CSI data.
57. [If true]: Call DMAGO_CSI(gdma_reg_base, ch1_pc, 0x1) \u2014 start DMA channel 1 (ch_num=1).
58. [If true]: Set rd_data = 0.
59. [If true]: Enter while loop: condition (rd_data & 0x2) == 0x0.
60. [If true]: Inside while loop: rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) \u2014 poll DMA interrupt status register for channel 1 completion (bit 1).
61. [If true]: Inside while loop: printf("polling irq; irq_status_reg rd_data =%0x\\n", rd_data).
62. [If true]: Exit while loop when (rd_data & 0x2) != 0x0.
63. [If true]: rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) \u2014 additional read of DMA interrupt status after loop exit.
64. [If true]: write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET, 0x2) \u2014 clear DMA IRQ bit 1.
65. [If true]: printf("DEBUG : irq polling completed ch1; rd_data=%0x\\n", rd_data).
66. End of if block.
67. End of inner for loop (next i iteration).
68. End of outer for loop (next lane_num iteration, decrementing from 3 to 0).
69. Call finish(0) \u2014 test completion with pass status.""",
        'Meta_Impacted_Registers': 'MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED',
        'Meta_Validation': """1. After calling csi2_enable_interrupt(), verify INT_ST_MAIN register (offset 0xC) was read to clear any pending interrupts. All status bits (status_int_st_phy_fatal[0], status_int_st_pkt_fatal[1], status_int_st_bndry_frame_fatal[2], status_int_st_seq_frame_fatal[3], status_int_st_crc_frame_fatal[4], status_int_st_pld_crc_fatal[5], status_int_st_data_id[6], status_int_st_ecc_corrected[7], status_int_st_phy[16], status_int_st_line[17], status_int_st_ipi_fatal[18]) are RC (read-clear).
2. Verify INT_MSK_PHY_FATAL register (offset 0xE4) was written with 0x0000000f, enabling mask_phy_errsotsynchs_0[0], mask_phy_errsotsynchs_1[1], mask_phy_errsotsynchs_2[2], mask_phy_errsotsynchs_3[3].
3. Verify INT_MSK_PKT_FATAL register (offset 0xF4) was written with 0x00000003, enabling mask_err_ecc_double[0] and mask_shorter_payload[1].
4. Verify INT_MSK_PHY register (offset 0x114) was written with 0x000f000f, enabling mask_phy_errsoths_0[0], mask_phy_errsoths_1[1], mask_phy_errsoths_2[2], mask_phy_errsoths_3[3], mask_phy_erresc_0[16], mask_phy_erresc_1[17], mask_phy_erresc_2[18], mask_phy_erresc_3[19].
5. Verify INT_MSK_LINE register (offset 0x134) was written with 0x000f000f, enabling mask_err_l_bndry_match_di0[0], mask_err_l_bndry_match_di1[1], mask_err_l_bndry_match_di2[2], mask_err_l_bndry_match_di3[3], mask_err_l_seq_di0[16], mask_err_l_seq_di1[17], mask_err_l_seq_di2[18], mask_err_l_seq_di3[19].
6. Verify INT_MSK_BNDRY_FRAME_FATAL register (offset 0x284) was written with 0x0000ffff, enabling err_f_bndry_match_vc0[0] through err_f_bndry_match_vc15[15].
7. Verify INT_MSK_SEQ_FRAME_FATAL register (offset 0x294) was written with 0x0000ffff, enabling err_f_seq_vc0[0] through err_f_seq_vc15[15].
8. Verify INT_MSK_CRC_FRAME_FATAL register (offset 0x2A4) was written with 0x0000ffff, enabling err_frame_data_vc0[0] through err_frame_data_vc15[15].
9. Verify INT_MSK_PLD_CRC_FATAL register (offset 0x2B4) was written with 0x0000ffff, enabling err_crc_vc0[0] through err_crc_vc15[15].
10. Verify INT_MSK_DATA_ID register (offset 0x2C4) was written with 0x0000ffff, enabling err_id_vc0[0] through err_id_vc15[15].
11. Verify INT_MSK_ECC_CORRECTED register (offset 0x2D4) was written with 0x0000ffff, enabling err_ecc_corrected_vc0[0] through err_ecc_corrected_vc15[15].
12. Verify PHY_STOPSTATE register (offset 0x4C) polling exits when rd_data == 0x1000f, confirming phy_stopstatedata_0[0]=1, phy_stopstatedata_1[1]=1, phy_stopstatedata_2[2]=1, phy_stopstatedata_3[3]=1, and phy_stopstateclk[16]=1.
13. For each lane_num iteration (3, 2, 1, 0): verify N_LANES register (offset 0x4) is written with lane_num value (n_lanes[2:0] field).
14. For each lane_num iteration: verify write to 0xa0243ffc with value (lane_num + 1) triggers CSI-2 sequence.
15. For each packet iteration: verify DMA channel 0 interrupt polling loop exits when (rd_data & 0x1) != 0x0, confirming control data transfer completion.
16. For each packet iteration: verify DMA interrupt is cleared by writing 0x1 to gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET.
17. For each packet iteration where (csi_ctrl_data & 0x3f) > 0xf: verify DMA channel 1 interrupt polling loop exits when (rd_data & 0x2) != 0x0, confirming CSI data transfer completion.
18. For each packet iteration where (csi_ctrl_data & 0x3f) > 0xf: verify DMA interrupt is cleared by writing 0x2 to gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET.
19. Verify word_count extraction: word_count = ((csi_ctrl_data >> 6) & 0xffff).
20. Verify csi_data_size alignment: csi_data_size = (word_count % 8) ? (word_count / 8 + 1) * 8 : word_count.
21. Verify test completes successfully via finish(0).""",
        'Meta_Headers': '#include <stdio.h>; #include <stdlib.h>; #include "test_common.h"; #include "mipi_csi2.h"',
        'Meta_Macros': '#define GDMA_CSI2_DATA_DEST_ADDR2 0xE6040080\n#define GDMA_CTRL_DATA_DEST_ADDR2 0xE6040000\n#define LS_LE_EN 1\n#define TOTAL_FRAME 1\n#define VC_ID 3\n#define VRES 1080 (if GDMA0_FULL_MEM defined) or 3 (default)\n#define HRES 1920 (if GDMA0_FULL_MEM defined) or 64 (default)\n#define DATA_TYPE CSI2_RGB888',
        'Meta_Arrays': 'NA'
    }

    wb = Workbook()

    # --- TestPlan Sheet ---
    ws_tp = wb.active
    ws_tp.title = 'TestPlan'
    tp_headers = ['Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
                  'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
                  'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria', 'Code Generation']
    tp_data = [row['Index'], row['SS_Module'], row['Feature'], row['Test_Case_Name'],
               row['Test_Description'], row['Speed'], row['Mode'], row['Memory_Start_Offset'],
               row['Memory_End_Offset'], row['Remarks'], row['Test_Steps'],
               row['Impacted_Registers'], row['Validation'], 'Required']

    # --- MetaData Sheet ---
    ws_md = wb.create_sheet('MetaData')
    md_headers = ['Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
                  'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
                  'Meta Headers', 'Meta Macros', 'Meta Arrays']
    md_data = [row['Index'], row['Test_Case_Name'], row['Meta_Test_Description'],
               row['Meta_Test_Steps'], row['Meta_Impacted_Registers'], row['Meta_Validation'],
               row['Meta_Headers'], row['Meta_Macros'], row['Meta_Arrays']]

    # Formatting
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    wrap_align = Alignment(wrap_text=True, vertical='top')

    for sheet, headers, data_row in [(ws_tp, tp_headers, tp_data), (ws_md, md_headers, md_data)]:
        for col_idx, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap_align
        for col_idx, value in enumerate(data_row, 1):
            cell = sheet.cell(row=2, column=col_idx, value=value)
            cell.alignment = wrap_align
        sheet.freeze_panes = 'A2'
        # Auto-size columns
        for col_idx in range(1, len(headers) + 1):
            max_len = len(str(headers[col_idx - 1]))
            if col_idx <= len(data_row) and data_row[col_idx - 1]:
                lines = str(data_row[col_idx - 1]).split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
            col_width = min(max_len + 4, 80)
            if col_width < 12:
                col_width = 12
            col_letter = sheet.cell(row=1, column=col_idx).column_letter
            sheet.column_dimensions[col_letter].width = col_width

    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = 'veryHidden'

    # Save
    script_dir = os.path.dirname(os.path.abspath(__file__))
    filepath = os.path.join(script_dir, filename)
    wb.save(filepath)
    print(f'SUCCESS: {filepath}')
    print(f'FILENAME: {filename}')

    # Verify
    from openpyxl import load_workbook
    vwb = load_workbook(filepath)
    assert 'TestPlan' in vwb.sheetnames
    assert 'MetaData' in vwb.sheetnames
    fsize = os.path.getsize(filepath)
    assert fsize > 0
    print(f'VALIDATED: size={fsize}, sheets={vwb.sheetnames}')
    return filepath, filename

if __name__ == '__main__':
    generate()
