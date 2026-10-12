#!/usr/bin/env python3
"""MIPI_DSI TestPlan Excel Generator - Agent 7 Direct Execution
Generates MIPI_DSI_TestPlan_<timestamp>.xlsx with TestPlan and MetaData sheets.
Run: python generate_mipi_dsi_testplan_20261013_190000.py
"""
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")

# Output config
IP_NAME = "MIPI_DSI"
output_dir = os.path.dirname(os.path.abspath(__file__))
filename = f"{IP_NAME}_TestPlan_{timestamp}.xlsx"
filepath = os.path.join(output_dir, filename)

# ============ DATA ============
json_data = [
    {
        "Index": "1",
        "SS / Module": "mipi_dsi",
        "Test Case Name": "mipi_dsi_dbi_random_payload_test",
        "Feature": "NA",
        "Meta Headers": '#include <stdio.h>\n#include <stdlib.h>\n#include "test_common.h"\n#include "mipi_dsi.h"\n#include <math.h>',
        "Meta Macros": "MIZAR_MIPI_DSI_HOST_PHY_IF_CFG\nMIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME\nMIZAR_MIPI_DSI_HOST_PCKHDL_CFG\nMIZAR_MIPI_DSI_HOST_CLKMGR_CFG\nMIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL\nMIZAR_MIPI_DSI_DMAC_INTEN\nMIZAR_MIPI_DSI_DMAC_DBGINST0\nMIZAR_MIPI_DSI_DMAC_DBGINST1\nMIZAR_MIPI_DSI_DMAC_DBGCMD\nMIZAR_MIPI_DSI_DMAC_INTMIS\nMIZAR_MIPI_DSI_DMAC_INTCLR\nRAM_BASE\nDMA_SAR\nDMA_DAR\nDSI_WRITE_MEMORY_START\nPROC_MODE",
        "Meta Arrays": "desc_t data_tdbdcb[10];\ndesc_t data_rebdcb[10];\ndesc_t cmd_tdbdcb[10];\ndesc_t cmd_rebdcb[10];\nunsigned int wr_cmd_arr[27];\nunsigned int WR_CMD_SIZE[27];",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase configures the MIPI DSI host PHY interface, packet handler, and clock manager registers. It disables DPI control (DBI enable mode), initializes the PHY via phy_init(), configures DBI parameters (VCID, LUT size, output/input DBI config, partitioning, command size, pixel attributes, tear effect enable, generic VC ID, and various DCS/generic TX settings). It programs the DPI clock based on a calculated frequency from dpi_clk_time_period = 16.012400. It then enables DMAC interrupts and iterates 10 times, each iteration generating a random number of pixels (rand() % 1024) * 8, recalculating wr_cmd_size and num_bytes via pixel_to_bytes_wr_cmd_size(), reconfiguring DBI, setting up DMA descriptors for data channel (CH0) and command channel (CH1) with source addresses in SRAM and destination addresses at 0x10000000000 and 0x10000008000 respectively, programming DMA move/write-memory-barrier/signal-event/end instructions, loading random data and write commands into SRAM, executing DMA transfers via DBGINST0/DBGINST1/DBGCMD register writes for both channels, polling INTMIS until both channel interrupts (0x3) are asserted, clearing interrupts via INTCLR, and waiting via wait_on(100000). After all iterations, finish(0) is called.",
        "Test Description": "Configure MIPI DSI host PHY, packet handler, and clock manager. Initialize PHY. Configure DBI mode parameters including VCID, LUT size, command size, pixel attributes, and tear effect. Program DPI clock. Enable DMAC interrupts. Iterate 10 times with random pixel counts, setting up DMA descriptors for data and command channels, loading random payload data and write commands into SRAM, executing DMA transfers, polling for completion interrupts on both channels, and clearing interrupts. Finish test.",
        "Meta Test Steps / Procedure": "1. int test_case() {\n2. pixel_num_bytes_wr_cmd_size_t pixel_attr;\n3. unsigned int phy_if_cfg;\n4. unsigned long int gen_pld_data;\n5. unsigned int num_bytes;\n6. unsigned int num_descriptors;\n7. unsigned int num_of_pixel;\n8. unsigned int itter;\n9. unsigned int wr_cmd_arr[27];\n10. unsigned int WR_CMD_SIZE[27];\n11. unsigned int i;\n12. unsigned int total_num_bytes;\n13. unsigned int num_wr_cmds;\n14. unsigned int rd_data;\n15. phy_stop_wait_time = 0x40;\n16. n_lanes = 3;\n17. phy_if_cfg = n_lanes;\n18. phy_if_cfg = set_data_mask(phy_if_cfg, MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME, phy_stop_wait_time);\n19. write_reg(MIZAR_MIPI_DSI_HOST_PHY_IF_CFG, phy_if_cfg);\n20. write_reg(MIZAR_MIPI_DSI_HOST_PCKHDL_CFG, 0x3d);\n21. write_reg(MIZAR_MIPI_DSI_HOST_CLKMGR_CFG, 0x107);\n22. write_reg(MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL, 0);\n23. phy_init();\n24. dbi_vcid = 0x3;\n25. load_cmd_or_data_to_sram = 1;\n26. lut_size_conf = 0x1;\n27. out_dbi_conf = 0xb;\n28. in_dbi_conf = 0x0;\n29. partitioning_en = 0x0;\n30. allowed_cmd_size = 0x25;\n31. wr_cmd_size = 192 + 1;\n32. num_of_pixel = 40;\n33. pixel_attr = pixel_to_bytes_wr_cmd_size(num_of_pixel);\n34. wr_cmd_size = pixel_attr.wr_cmd_size;\n35. num_bytes = pixel_attr.num_bytes;\n36. max_rd_pkt_size = 0x0;\n37. dcs_lw_tx = 0x0;\n38. dcs_sr_0p_tx = 0x0;\n39. dcs_sw_1p_tx = 0x0;\n40. dcs_sw_0p_tx = 0x0;\n41. gen_lw_tx = 0x0;\n42. gen_sr_2p_tx = 0x0;\n43. gen_sr_1p_tx = 0x0;\n44. gen_sr_0p_tx = 0x0;\n45. gen_sw_2p_tx = 0x0;\n46. gen_sw_1p_tx = 0x0;\n47. gen_sw_0p_tx = 0x0;\n48. ack_rqst_en = 0x0;\n49. tear_fx_en = 0x1;\n50. generic_vc_id = 0x2;\n51. dpi_clk_time_period = 16.012400;\n52. dpi_clk_freq = 1000000000/dpi_clk_time_period;\n53. program_dpi_clock(dpi_clk_freq);\n54. dbi_config();\n55. num_descriptors = 1;\n56. write_reg(MIZAR_MIPI_DSI_DMAC_INTEN, 0x3);\n57. for(i = 0; i < 10; i++) {\n58. num_of_pixel = ((rand() % 1024)) * 8;\n59. pixel_attr = pixel_to_bytes_wr_cmd_size(num_of_pixel);\n60. wr_cmd_size = pixel_attr.wr_cmd_size;\n61. num_bytes = pixel_attr.num_bytes;\n62. printf(\"i = %d, wr_cmd_size = %d\\n\", i, wr_cmd_size);\n63. dbi_config();\n64. ch0_desc_addr = RAM_BASE + 0x3F000;\n65. ch1_desc_addr = RAM_BASE + 0x3F800;\n66. ch0_desc_addr_act = ch0_desc_addr;\n67. ch1_desc_addr_act = ch1_desc_addr;\n68. for(itter = 0; itter < num_descriptors; itter++) {\n69. data_tdbdcb[itter].addr = RAM_BASE + 0x10000 + (itter * 0x1000);\n70. data_tdbdcb[itter].len = num_bytes;\n71. data_tdbdcb[itter].eop = (itter == (num_descriptors - 1));\n72. data_rebdcb[itter].addr = 0x10000000000;\n73. data_rebdcb[itter].len = num_bytes;\n74. cmd_tdbdcb[itter].addr = RAM_BASE + 0x0000;\n75. cmd_tdbdcb[itter].len = 4;\n76. cmd_tdbdcb[itter].eop = (itter == (num_descriptors - 1));\n77. cmd_rebdcb[itter].addr = 0x10000008000;\n78. cmd_rebdcb[itter].len = 4;\n79. DMAMOV(&ch0_desc_addr, DMA_SAR, data_tdbdcb[itter].addr);\n80. DMAMOV(&ch0_desc_addr, DMA_DAR, data_rebdcb[itter].addr);\n81. program_data_num_bytes(&ch0_desc_addr, (ceil((double)data_tdbdcb[itter].len/8))*8);\n82. DMAWMB(&ch0_desc_addr);\n83. DMASEV(&ch0_desc_addr, 0);\n84. DMAEND(&ch0_desc_addr);\n85. DMAMOV(&ch1_desc_addr, DMA_SAR, cmd_tdbdcb[itter].addr);\n86. DMAMOV(&ch1_desc_addr, DMA_DAR, cmd_rebdcb[itter].addr);\n87. program_data_num_bytes(&ch1_desc_addr, cmd_tdbdcb[itter].len);\n88. DMAWMB(&ch1_desc_addr);\n89. DMASEV(&ch1_desc_addr, 1);\n90. DMAEND(&ch1_desc_addr);\n91. load_rand_data(data_tdbdcb[itter]);\n92. load_wr_command(cmd_tdbdcb[itter].addr, num_bytes, DSI_WRITE_MEMORY_START);\n93. }\n94. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST0, 0x00A00000);\n95. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST1, ch0_desc_addr_act);\n96. write_reg(MIZAR_MIPI_DSI_DMAC_DBGCMD, 0x0);\n97. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST0, 0x01A00000);\n98. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST1, ch1_desc_addr_act);\n99. write_reg(MIZAR_MIPI_DSI_DMAC_DBGCMD, 0x0);\n100. rd_data = read_reg(MIZAR_MIPI_DSI_DMAC_INTMIS);\n101. while(rd_data != 0x3) { rd_data = read_reg(MIZAR_MIPI_DSI_DMAC_INTMIS); }\n102. write_reg(MIZAR_MIPI_DSI_DMAC_INTCLR, rd_data);\n103. wait_on(100000);\n104. }\n105. finish(0);",
        "Test Steps / Procedure": "1. Configure PHY_IF_CFG with n_lanes=3 and phy_stop_wait_time=0x40.\n2. Write 0x3d to PCKHDL_CFG.\n3. Write 0x107 to CLKMGR_CFG.\n4. Write 0 to dpi_control to select DBI mode.\n5. Initialize PHY via phy_init().\n6. Set DBI configuration parameters: VCID=0x3, LUT size=0x1, output DBI config=0xb, input DBI config=0x0, partitioning disabled, allowed_cmd_size=0x25, initial wr_cmd_size=193, num_of_pixel=40.\n7. Compute pixel attributes via pixel_to_bytes_wr_cmd_size() and update wr_cmd_size and num_bytes.\n8. Disable all DCS and generic TX command types. Enable tear effect (tear_fx_en=1). Set generic_vc_id=0x2.\n9. Program DPI clock based on period 16.012400ns.\n10. Call dbi_config() for initial DBI configuration.\n11. Set num_descriptors=1.\n12. Write 0x3 to INTEN to enable interrupts for both DMA channels.\n13. Loop 10 iterations:\n a. Generate random num_of_pixel = (rand() % 1024) * 8.\n b. Recompute pixel attributes, wr_cmd_size, and num_bytes.\n c. Reconfigure DBI via dbi_config().\n d. Set CH0 and CH1 descriptor addresses.\n e. Populate data and command descriptors with source/destination addresses and lengths.\n f. Program DMA instructions (DMAMOV, program_data_num_bytes, DMAWMB, DMASEV, DMAEND) for both channels.\n g. Load random data into SRAM for data descriptor.\n h. Load write command into SRAM for command descriptor.\n i. Execute CH0 DMA: write 0x00A00000 to DBGINST0, CH0 descriptor address to DBGINST1, 0x0 to DBGCMD.\n j. Execute CH1 DMA: write 0x01A00000 to DBGINST0, CH1 descriptor address to DBGINST1, 0x0 to DBGCMD.\n k. Poll INTMIS until value equals 0x3 (both channels complete).\n l. Clear interrupts by writing read value to INTCLR.\n m. Wait via wait_on(100000).\n14. Call finish(0).",
        "Meta Impacted Registers": "PHY_IF_CFG\nPCKHDL_CFG\nCLKMGR_CFG\ndpi_control\nINTEN\nDBGINST0\nDBGINST1\nDBGCMD\nINTMIS\nINTCLR",
        "Impacted Registers": "PHY_IF_CFG, PCKHDL_CFG, CLKMGR_CFG, dpi_control, INTEN, DBGINST0, DBGINST1, DBGCMD, INTMIS, INTCLR",
        "Meta Validation / Acceptance Criteria": "1. Poll read_reg(MIZAR_MIPI_DSI_DMAC_INTMIS) in a while loop until rd_data == 0x3, indicating both DMA channel 0 and channel 1 interrupts are asserted.\n2. After polling succeeds, write rd_data (0x3) to MIZAR_MIPI_DSI_DMAC_INTCLR to clear both interrupt flags.\n3. This polling and clearing is performed in each of the 10 loop iterations.\n4. finish(0) is called after all iterations complete, indicating test pass with code 0.",
        "Validation / Acceptance Criteria": "1. For each of the 10 iterations, INTMIS must read 0x3 indicating both DMA channels completed successfully.\n2. Interrupts are cleared by writing to INTCLR after each iteration.\n3. Test completes successfully via finish(0) after all 10 iterations.",
        "Remarks": "The testcase uses random pixel counts per iteration ((rand() % 1024) * 8), so payload sizes vary across runs. Data channel (CH0) destination address is 0x10000000000 and command channel (CH1) destination address is 0x10000008000. The data length is ceil-aligned to 8-byte boundaries via (ceil((double)data_tdbdcb[itter].len/8))*8. Variables gen_pld_data, wr_cmd_arr, WR_CMD_SIZE, total_num_bytes, and num_wr_cmds are declared but not used in the visible source.",
    }
]

# ============ SHEET DEFINITIONS ============
testplan_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

metadata_columns = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

# ============ CREATE WORKBOOK ============
wb = Workbook()

# TestPlan sheet
ws_tp = wb.active
ws_tp.title = "TestPlan"

# MetaData sheet
ws_md = wb.create_sheet(title="MetaData")

# ============ FORMATTING ============
header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_alignment = Alignment(wrap_text=True, vertical="top")

def write_sheet(ws, columns, data, key_map=None):
    # Write headers
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    
    # Write data rows
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            k = col_name if key_map is None else key_map.get(col_name, col_name)
            value = row_data.get(k, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
    
    # Auto-size columns
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(str(col_name))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        max_len = max(max_len, len(line))
        # Cap width
        adjusted_width = min(max_len + 4, 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width
    
    # Freeze first row
    ws.freeze_panes = "A2"

# Write TestPlan sheet
write_sheet(ws_tp, testplan_columns, json_data)

# Write MetaData sheet
write_sheet(ws_md, metadata_columns, json_data)

# Set MetaData sheet to veryHidden
ws_md.sheet_state = "veryHidden"

# ============ SAVE ============
wb.save(filepath)
print(f"Workbook saved: {filepath}")
print(f"Filename: {filename}")

# ============ VALIDATE ============
assert os.path.exists(filepath), "File does not exist!"
assert os.path.getsize(filepath) > 0, "File is empty!"

# Reopen and validate
wb_check = load_workbook(filepath)
assert "TestPlan" in wb_check.sheetnames, "TestPlan sheet missing!"
assert "MetaData" in wb_check.sheetnames, "MetaData sheet missing!"
assert wb_check["MetaData"].sheet_state == "veryHidden", "MetaData not veryHidden!"
tp_rows = wb_check["TestPlan"].max_row - 1  # minus header
md_rows = wb_check["MetaData"].max_row - 1
print(f"Validation PASSED")
print(f"TestPlan rows: {tp_rows}")
print(f"MetaData rows: {md_rows}")
print(f"File size: {os.path.getsize(filepath)} bytes")
wb_check.close()
