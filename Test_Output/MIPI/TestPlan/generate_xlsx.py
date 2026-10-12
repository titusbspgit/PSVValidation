#!/usr/bin/env python3
"""MIPI_DSI TestPlan Excel Generator - Agent 7"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone, timedelta
import os, json, sys

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime('%Y%m%d_%H%M%S')
filename = f'MIPI_DSI_TestPlan_{timestamp}.xlsx'

# Input JSON data
json_data = [
{
"Index":"1",
"SS / Module":"mipi_dsi",
"Test Case Name":"mipi_dsi_basic_test",
"Feature":"NA",
"Meta Headers":"#include <stdio.h>\n#include <stdlib.h>\n#include \"test_common.h\"\n#include \"mipi_dsi.h\"",
"Meta Macros":"MIZAR_MIPI_DSI_DMAC_INTEN, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE, MIPI_DSI_SUBSYS_INTERRUPT_ENABLE_GDMA_INTR, MIZAR_MIPI_DSI_HOST_PHY_IF_CFG, MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME, MIZAR_MIPI_DSI_HOST_PCKHDL_CFG, MIZAR_MIPI_DSI_HOST_CLKMGR_CFG, MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL, MIZAR_MIPI_DSI_DMAC_DBGINST0, MIZAR_MIPI_DSI_DMAC_DBGINST1, MIZAR_MIPI_DSI_DMAC_DBGCMD, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK, MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR, MIZAR_MIPI_DSI_DMAC_INTMIS, MIZAR_MIPI_DSI_DMAC_INTCLR, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW, DSI_INTR_NO, RAM_BASE, DMA_SAR, DMA_DAR, DSI_WRITE_MEMORY_START",
"Meta Arrays":"desc_t data_tdbdcb[2];\ndesc_t data_rebdcb[2];\ndesc_t cmd_tdbdcb[2];\ndesc_t cmd_rebdcb[2];",
"Speed":"NA",
"Mode":"NA",
"Memory Start Offset":"NA",
"Memory End Offset":"NA",
"Meta Test Description":"This testcase configures the MIPI DSI subsystem for a basic DBI (Display Bus Interface) write operation. It enables GIC and DMAC interrupts, configures PHY interface (n_lanes=3, phy_stop_wait_time=0x40), sets PCKHDL_CFG to 0x3d and CLKMGR_CFG to 0x107, disables DPI control, initializes PHY via phy_init(), configures DBI parameters (vcid=0x3, lut_size_conf=0x1, out_dbi_conf=0xb, in_dbi_conf=0x0, partitioning_en=0x1, allowed_cmd_size=0x7, num_of_pixel=40), programs DPI clock, sets up DMA descriptors for data channel (CH0) and command channel (CH1) with random data payload and write memory start command, launches both DMA channels via debug instruction registers (DBGINST0, DBGINST1, DBGCMD), waits for interrupt-driven completion by polling int_pend1, and finishes with err0 as the error count. The ISR (Default_IRQHandler) reads interrupt_mask to identify GDMA interrupt source, reads INTMIS for channel mask status, clears pending channels from int_pend1, clears DMAC interrupt via INTCLR, clears subsystem-level interrupt via interrupt_raw, and increments err0 on unexpected interrupts.",
"Test Description":"Basic MIPI DSI DBI write test that configures PHY interface, DBI parameters, programs DMA descriptors for data and command channels, initiates DMA transfers via debug instruction registers, and validates completion through interrupt-driven handshake with error tracking.",
"Meta Test Steps / Procedure":"1. GIC_EnableIRQ(DSI_INTR_NO);\n2. int_pend = 1;\n3. int_pend1 = 0x3;\n4. phy_stop_wait_time = 0x40;\n5. n_lanes = 3;\n6. write_reg(MIZAR_MIPI_DSI_DMAC_INTEN, 0x3);\n7. write_reg(MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE, MIPI_DSI_SUBSYS_INTERRUPT_ENABLE_GDMA_INTR);\n8. unsigned int phy_if_cfg = n_lanes;\n9. phy_if_cfg = set_data_mask(phy_if_cfg, MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME, phy_stop_wait_time);\n10. write_reg(MIZAR_MIPI_DSI_HOST_PHY_IF_CFG, phy_if_cfg);\n11. write_reg(MIZAR_MIPI_DSI_HOST_PCKHDL_CFG, 0x3d);\n12. write_reg(MIZAR_MIPI_DSI_HOST_CLKMGR_CFG, 0x107);\n13. write_reg(MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL, 0);\n14. phy_init();\n15. dbi_vcid = 0x3;\n16. load_cmd_or_data_to_sram = 1;\n17. lut_size_conf = 0x1;\n18. out_dbi_conf = 0xb;\n19. in_dbi_conf = 0x0;\n20. partitioning_en = 0x1;\n21. allowed_cmd_size = 0x7;\n22. wr_cmd_size = 192 + 1;\n23. unsigned int num_of_pixel = 40;\n24. pixel_num_bytes_wr_cmd_size_t pixel_attr = pixel_to_bytes_wr_cmd_size(num_of_pixel);\n25. wr_cmd_size = pixel_attr.wr_cmd_size;\n26. unsigned int num_bytes = pixel_attr.num_bytes;\n27. max_rd_pkt_size = 0x0;\n28. dcs_lw_tx = 0x0;\n29. dcs_sr_0p_tx = 0x0;\n30. dcs_sw_1p_tx = 0x0;\n31. dcs_sw_0p_tx = 0x0;\n32. gen_lw_tx = 0x0;\n33. gen_sr_2p_tx = 0x0;\n34. gen_sr_1p_tx = 0x0;\n35. gen_sr_0p_tx = 0x0;\n36. gen_sw_2p_tx = 0x0;\n37. gen_sw_1p_tx = 0x0;\n38. gen_sw_0p_tx = 0x0;\n39. ack_rqst_en = 0x0;\n40. tear_fx_en = 0x1;\n41. generic_vc_id = 0x2;\n42. dpi_clk_time_period = 16.012400;\n43. dpi_clk_freq = 1000000000/dpi_clk_time_period;\n44. program_dpi_clock(dpi_clk_freq);\n45. dbi_config();\n46. unsigned int num_descriptors = 1;\n47. ch0_desc_addr = RAM_BASE + 0x3F000;\n48. ch1_desc_addr = RAM_BASE + 0x3F800;\n49. ch0_desc_addr_act = ch0_desc_addr;\n50. ch1_desc_addr_act = ch1_desc_addr;\n51. for(unsigned int itter = 0; itter < num_descriptors; itter++) {\n52. data_tdbdcb[itter].addr = RAM_BASE + 0x10000;\n53. data_tdbdcb[itter].len = num_bytes;\n54. data_tdbdcb[itter].eop = (itter == (num_descriptors - 1));\n55. data_rebdcb[itter].addr = 0x10000000000;\n56. data_rebdcb[itter].len = num_bytes;\n57. cmd_tdbdcb[itter].addr = RAM_BASE + 0x0000;\n58. cmd_tdbdcb[itter].len = 4;\n59. cmd_tdbdcb[itter].eop = (itter == (num_descriptors - 1));\n60. cmd_rebdcb[itter].addr = 0x10000008000;\n61. cmd_rebdcb[itter].len = 4;\n62. DMAMOV(&ch0_desc_addr, DMA_SAR, data_tdbdcb[itter].addr);\n63. DMAMOV(&ch0_desc_addr, DMA_DAR, data_rebdcb[itter].addr);\n64. program_data_num_bytes(&ch0_desc_addr, data_tdbdcb[itter].len);\n65. DMAWMB(&ch0_desc_addr);\n66. DMASEV(&ch0_desc_addr, 0);\n67. DMAEND(&ch0_desc_addr);\n68. DMAMOV(&ch1_desc_addr, DMA_SAR, cmd_tdbdcb[itter].addr);\n69. DMAMOV(&ch1_desc_addr, DMA_DAR, cmd_rebdcb[itter].addr);\n70. program_data_num_bytes(&ch1_desc_addr, cmd_tdbdcb[itter].len);\n71. DMAWMB(&ch1_desc_addr);\n72. DMASEV(&ch1_desc_addr, 1);\n73. DMAEND(&ch1_desc_addr);\n74. load_rand_data(data_tdbdcb[itter]);\n75. load_wr_command(cmd_tdbdcb[itter].addr, num_bytes, DSI_WRITE_MEMORY_START);\n76. }\n77. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST0, 0x00A00000);\n78. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST1, ch0_desc_addr_act);\n79. write_reg(MIZAR_MIPI_DSI_DMAC_DBGCMD, 0x0);\n80. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST0, 0x01A00000);\n81. write_reg(MIZAR_MIPI_DSI_DMAC_DBGINST1, ch1_desc_addr_act);\n82. write_reg(MIZAR_MIPI_DSI_DMAC_DBGCMD, 0x0);\n83. while(int_pend1) { wait_on(10); }\n84. wait_on(10000);\n85. finish(err0);\n\nISR (void Default_IRQHandler()):\n86. unsigned int ch_mask_st;\n87. unsigned int dsi_subsys_mask_st;\n88. int_pend = 0;\n89. dsi_subsys_mask_st = read_reg(MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK);\n90. if(dsi_subsys_mask_st == MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR) {\n91. ch_mask_st = read_reg(MIZAR_MIPI_DSI_DMAC_INTMIS);\n92. if(ch_mask_st) {\n93. int_pend1 = int_pend1 & (~ch_mask_st);\n94. write_reg(MIZAR_MIPI_DSI_DMAC_INTCLR, ch_mask_st);\n95. } else {\n96. err0++;\n97. }\n98. write_reg(MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW, dsi_subsys_mask_st);\n99. } else {\n100. err0++;\n101. }\n102. GIC_ClearIRQ(DSI_INTR_NO);\n103. int_pend = 1;",
"Test Steps / Procedure":"1. Enable GIC interrupt for DSI.\n2. Initialize interrupt pending flags (int_pend=1, int_pend1=0x3).\n3. Set PHY stop wait time to 0x40 and number of lanes to 3.\n4. Enable DMAC interrupts by writing 0x3 to INTEN register.\n5. Enable GDMA interrupt at subsystem level via interrupt_enable register.\n6. Configure PHY_IF_CFG register with n_lanes and phy_stop_wait_time.\n7. Configure PCKHDL_CFG register with 0x3d.\n8. Configure CLKMGR_CFG register with 0x107.\n9. Disable DPI by writing 0 to dpi_control register.\n10. Initialize PHY.\n11. Configure DBI parameters (vcid=0x3, LUT size, output/input DBI config, partitioning, command size, pixel count=40).\n12. Set DCS/generic packet transmission parameters to 0, enable tear effect.\n13. Program DPI clock based on 16.012400ns period.\n14. Apply DBI configuration.\n15. Set up DMA descriptor addresses for data channel (CH0) and command channel (CH1).\n16. For each descriptor: populate source/destination addresses, lengths, EOP flags; program DMA move, write memory barrier, send event, and end instructions; load random data payload; load write memory start command.\n17. Launch CH0 DMA by writing 0x00A00000 to DBGINST0, CH0 descriptor address to DBGINST1, and 0x0 to DBGCMD.\n18. Launch CH1 DMA by writing 0x01A00000 to DBGINST0, CH1 descriptor address to DBGINST1, and 0x0 to DBGCMD.\n19. Poll int_pend1 with wait_on(10) until both channels complete (cleared by ISR).\n20. Wait additional 10000 cycles.\n21. Finish test with err0 as error count.\n22. ISR: Read interrupt_mask to identify GDMA source, read INTMIS for channel status, clear pending channels from int_pend1, clear DMAC interrupt via INTCLR, clear subsystem interrupt via interrupt_raw, increment err0 on unexpected interrupts, clear GIC IRQ.",
"Meta Impacted Registers":"INTEN, interrupt_enable, PHY_IF_CFG, PCKHDL_CFG, CLKMGR_CFG, dpi_control, DBGINST0, DBGINST1, DBGCMD, interrupt_mask, INTMIS, INTCLR, interrupt_raw",
"Impacted Registers":"INTEN, interrupt_enable, PHY_IF_CFG, PCKHDL_CFG, CLKMGR_CFG, dpi_control, DBGINST0, DBGINST1, DBGCMD, interrupt_mask, INTMIS, INTCLR, interrupt_raw",
"Meta Validation / Acceptance Criteria":"1. ISR reads dsi_subsys_mask_st from MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK and compares against MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR; if not equal, err0 is incremented and printf(\"ERROR1: Unexpected interrupt dsi_subsys_mask_st = 0x%x\\n\", dsi_subsys_mask_st) is called.\n2. ISR reads ch_mask_st from MIZAR_MIPI_DSI_DMAC_INTMIS; if ch_mask_st is 0 (no channel interrupt pending), err0 is incremented and printf(\"ERROR2: Unexpected interrupt ch_mask_st = 0x%x \\n\", ch_mask_st) is called.\n3. If ch_mask_st is non-zero, int_pend1 = int_pend1 & (~ch_mask_st) clears the completed channel bits.\n4. DMAC interrupt is cleared by write_reg(MIZAR_MIPI_DSI_DMAC_INTCLR, ch_mask_st).\n5. Subsystem interrupt is cleared by write_reg(MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW, dsi_subsys_mask_st).\n6. Main loop exits when int_pend1 becomes 0 (both CH0 and CH1 interrupts received and cleared).\n7. Test passes when finish(err0) is called with err0 == 0 (no unexpected interrupts occurred).",
"Validation / Acceptance Criteria":"1. The ISR must identify the interrupt source as GDMA via the interrupt_mask register; unexpected sources increment the error count.\n2. The ISR must read a non-zero channel mask from INTMIS; a zero value indicates an unexpected interrupt and increments the error count.\n3. Both DMA channels (CH0 and CH1) must complete and trigger interrupts that clear the corresponding bits in int_pend1.\n4. DMAC interrupts must be properly cleared via the INTCLR register.\n5. Subsystem-level interrupts must be properly cleared via the interrupt_raw register.\n6. The test passes when err0 is 0 at finish, indicating no unexpected interrupts occurred during the DBI write transfer.",
"Remarks":"The testcase uses two DMA channels: CH0 for data transfer and CH1 for command transfer. int_pend1 is initialized to 0x3 representing both channels pending. The ISR clears individual channel bits as their interrupts arrive. The test uses desc_t typed descriptor arrays (data_tdbdcb, data_rebdcb, cmd_tdbdcb, cmd_rebdcb) each of size 2, though only 1 descriptor is used (num_descriptors=1). DPI clock period is set to 16.012400ns. The data destination address is 0x10000000000 and command destination address is 0x10000008000. extern int int_pend is declared as extern int; unsigned int int_pend1 and unsigned int err0 are file-scope unsigned int variables."
}
]

# TestPlan columns
tp_cols = [
    'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
    'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
    'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
    'Code Generation'
]

# MetaData columns
md_cols = [
    'Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
    'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
    'Meta Headers', 'Meta Macros', 'Meta Arrays'
]

# Create workbook
wb = openpyxl.Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = 'TestPlan'

header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap_align = Alignment(wrap_text=True, vertical='top')

# Write headers
for col_idx, col_name in enumerate(tp_cols, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

# Write data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(tp_cols, 1):
        val = row_data.get(col_name, '')
        if val is None:
            val = ''
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = wrap_align

# Freeze first row
ws_tp.freeze_panes = 'A2'

# Auto-size columns with max width cap
MAX_WIDTH = 60
for col_idx, col_name in enumerate(tp_cols, 1):
    max_len = len(str(col_name))
    for row in ws_tp.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                lines = str(cell.value).split('\n')
                longest_line = max(len(l) for l in lines)
                if longest_line > max_len:
                    max_len = longest_line
    adjusted = min(max_len + 2, MAX_WIDTH)
    ws_tp.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = adjusted

# --- MetaData Sheet ---
ws_md = wb.create_sheet('MetaData')

# Write headers
for col_idx, col_name in enumerate(md_cols, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

# Write data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(md_cols, 1):
        val = row_data.get(col_name, '')
        if val is None:
            val = ''
        cell = ws_md.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = wrap_align

# Freeze first row
ws_md.freeze_panes = 'A2'

# Auto-size columns with max width cap
for col_idx, col_name in enumerate(md_cols, 1):
    max_len = len(str(col_name))
    for row in ws_md.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                lines = str(cell.value).split('\n')
                longest_line = max(len(l) for l in lines)
                if longest_line > max_len:
                    max_len = longest_line
    adjusted = min(max_len + 2, MAX_WIDTH)
    ws_md.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = adjusted

# Set MetaData sheet to veryHidden
ws_md.sheet_state = 'veryHidden'

# Save workbook
output_path = filename
wb.save(output_path)
wb.close()

# Post-save validation
assert os.path.exists(output_path), 'File does not exist'
assert os.path.getsize(output_path) > 0, 'File is empty'

# Reopen and validate
wb_check = openpyxl.load_workbook(output_path)
assert 'TestPlan' in wb_check.sheetnames, 'TestPlan sheet missing'
assert 'MetaData' in wb_check.sheetnames, 'MetaData sheet missing'
assert wb_check['MetaData'].sheet_state == 'veryHidden', 'MetaData not veryHidden'
tp_rows = wb_check['TestPlan'].max_row - 1
md_rows = wb_check['MetaData'].max_row - 1
wb_check.close()

print(f'SUCCESS: {output_path}')
print(f'Size: {os.path.getsize(output_path)} bytes')
print(f'TestPlan rows: {tp_rows}')
print(f'MetaData rows: {md_rows}')
print(f'Validation: PASSED')
