#!/usr/bin/env python3
"""
Agent 7 - Direct Excel Generator for MIPI_CSI TestPlan
Generates MIPI_CSI_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx using openpyxl
Timestamp: IST (GMT+05:30)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import os

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"MIPI_CSI_TestPlan_{timestamp}.xlsx"
output_dir = os.path.dirname(os.path.abspath(__file__))
filepath = os.path.join(output_dir, filename)

# ── TestPlan columns ──
tp_cols = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers",
    "Validation / Acceptance Criteria", "Code Generation"
]

# ── MetaData columns ──
md_cols = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

# ══════════════════════════════════════════════════════════════
# ROW 1 DATA — mipi_csi2_dphy_lanes_test
# ══════════════════════════════════════════════════════════════
r1_test_desc = (
    "This test validates the MIPI CSI-2 receiver operation across all D-PHY lane configurations from 4 lanes down to 1 lane. "
    "It first enables all CSI-2 host interrupt masks for PHY fatal, packet fatal, PHY, line, boundary frame fatal, "
    "sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected categories. "
    "The virtual channel register is configured based on the selected GDMA path, and control data transfer is enabled. "
    "The D-PHY is initialized and the test waits for all PHY lanes to reach stop state by polling the PHY_STOPSTATE register. "
    "For each lane configuration (4, 3, 2, 1 lane), the N_LANES register is programmed and a CSI-2 sequence is triggered. "
    "For each control packet, a DMA transfer is initiated to receive the control data. After DMA completion (confirmed via "
    "interrupt polling), the received control data is examined. If the data type indicates a long packet (greater than short "
    "packet threshold), the word count is extracted, the data size is aligned to 8 bytes, and a second DMA transfer is "
    "initiated to receive the image data payload. DMA completion for the data channel is also confirmed via interrupt polling. "
    "The test passes after successfully completing all lane configurations."
)

r1_test_steps = (
    "1. Enable all CSI-2 host interrupts by reading the main interrupt status register to clear pending interrupts, "
    "then writing appropriate mask values to all interrupt mask registers (PHY fatal, packet fatal, PHY, line, "
    "boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, ECC corrected).\n"
    "2. Configure the virtual channel register with the appropriate virtual channel ID based on the selected GDMA path.\n"
    "3. Enable control data transfer by writing to the control_data register.\n"
    "4. Initialize the D-PHY by calling the PHY initialization sequence.\n"
    "5. Poll the PHY_STOPSTATE register until all lanes report stop state.\n"
    "6. For each lane configuration (4 lanes down to 1 lane):\n"
    "   a. Write the lane count to the N_LANES register.\n"
    "   b. Trigger the CSI-2 sequence for the current lane configuration.\n"
    "7. For each control packet in the current lane configuration:\n"
    "   a. Enable DMA interrupts for both channels.\n"
    "   b. Program and start DMA channel 0 to transfer the control packet (8 bytes).\n"
    "   c. Poll the DMA interrupt status register until channel 0 transfer completes.\n"
    "   d. Clear the DMA channel 0 interrupt.\n"
    "   e. Read the received control data to determine the packet type.\n"
    "8. If the received packet is a long packet (data type indicates image data):\n"
    "   a. Extract the word count from the control data.\n"
    "   b. Compute the 8-byte-aligned data transfer size.\n"
    "   c. Program and start DMA channel 1 to transfer the image data payload.\n"
    "   d. Poll the DMA interrupt status register until channel 1 transfer completes.\n"
    "   e. Clear the DMA channel 1 interrupt.\n"
    "9. Repeat steps 7-8 for all packets in the current lane configuration.\n"
    "10. Repeat steps 6-9 for all lane configurations (4, 3, 2, 1 lane).\n"
    "11. Verify the test completes successfully for all lane configurations."
)

r1_impacted_regs = (
    "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; "
    "INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; "
    "INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; "
    "INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED"
)

r1_validation = (
    "1. The PHY_STOPSTATE register must read the expected stop-state value confirming all lanes are in stop state before proceeding.\n"
    "2. For each DMA control-packet transfer, the DMA interrupt status must indicate channel 0 completion.\n"
    "3. For long packets, the DMA interrupt status must indicate channel 1 completion after the data transfer.\n"
    "4. The control data received must be correctly parsed to distinguish short packets from long packets based on the data type field.\n"
    "5. All four lane configurations (4, 3, 2, 1 lane) must complete their respective packet transfers without timeout or error.\n"
    "6. The test must reach the final pass condition after all lane iterations complete successfully."
)

r1_remarks = (
    "The test iterates through all D-PHY lane configurations from 4 lanes down to 1 lane. "
    "The virtual channel ID and GDMA path are selected via compile-time conditional defines "
    "(GDMA0_PATH, GDMA1_PATH, GDMA2_PATH, GDMA3_PATH). The virtual channel register and control data register "
    "are written twice (repeated writes in source). DMA-based transfers are used for both control packets and data packets. "
    "The polling on PHY_STOPSTATE expects the value 0x1000f which includes the clock lane stop state. "
    "Two hardcoded addresses are used that could not be mapped to named registers from the provided specifications. "
    "The snps_phy_init() and dma_trnsfr_instn_preload() and DMAGO_CSI() function implementations are not available "
    "in the supplied testcase folder."
)

r1_meta_test_desc = (
    "This testcase validates the MIPI CSI-2 receiver across all D-PHY lane configurations (4 lanes down to 1 lane). "
    "It begins by calling csi2_enable_interrupt() which reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts, "
    "then writes interrupt mask registers (INT_MSK_PHY_FATAL with 0x0000000f, INT_MSK_PKT_FATAL with 0x00000003, "
    "INT_MSK_PHY with 0x000f000f, INT_MSK_LINE with 0x000f000f, INT_MSK_BNDRY_FRAME_FATAL with 0x0000ffff, "
    "INT_MSK_SEQ_FRAME_FATAL with 0x0000ffff, INT_MSK_CRC_FRAME_FATAL with 0x0000ffff, INT_MSK_PLD_CRC_FATAL with 0x0000ffff, "
    "INT_MSK_DATA_ID with 0x0000ffff, INT_MSK_ECC_CORRECTED with 0x0000ffff). The virtual channel ID is computed based on "
    "the GDMA path (GDMA0/1/2/3) and written to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL. Control data transfer is enabled "
    "by writing 1 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA. The D-PHY is initialized via snps_phy_init(). The test then polls "
    "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until it reads 0x1000f, confirming all lanes have entered stop state. A for-loop "
    "iterates lane_num from 3 down to 0: for each lane configuration, MIZAR_MIPI_CSI2_HOST_N_LANES is written with lane_num, "
    "and a trigger write to 0xa0243ffc with (lane_num+1) starts the CSI-2 sequence. An inner loop iterates over control packet "
    "count ((VRES*3)+2). In each iteration, DMA interrupt enable is written (gdma_reg_base+MIPI_CSI2_DMA_INTEN_OFFSET with 0x3), "
    "DMA control-data transfer instructions are preloaded via dma_trnsfr_instn_preload() with src_addr=0x8000, dest_addr=0xE6001000, "
    "trnsfr_size=8, irq_num=0. DMAGO_CSI is called to start DMA channel 0. The test polls gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET "
    "until bit 0 is set, then clears the interrupt by writing 0x1 to gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET. It reads "
    "gdma_reg_base+0x28 for debug status. The control data is read from 0xE6001000. If the data type field (bits[5:0]) is greater "
    "than 0xf, word_count is extracted from bits[21:6], csi_data_size is computed (8-byte aligned), and a second DMA transfer is "
    "programmed via dma_trnsfr_instn_preload() with src_addr=0x0000, dest_addr=0xE6002000, trnsfr_size=csi_data_size, irq_num=1. "
    "DMAGO_CSI starts DMA channel 1. The test polls gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET until bit 1 is set, reads the "
    "interrupt status again, then clears the interrupt by writing 0x2 to gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET. After all "
    "lane iterations complete, finish(0) is called to end the test with pass status."
)

r1_meta_test_steps = (
    "1. Enter test_case() function.\n"
    "2. Call csi2_enable_interrupt().\n"
    "3. Inside csi2_enable_interrupt(): Read MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts.\n"
    "4. Write 0x0000000f to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL to enable PHY fatal interrupts.\n"
    "5. Write 0x00000003 to MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL to enable packet fatal interrupts.\n"
    "6. Write 0x000f000f to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY to enable PHY interrupts.\n"
    "7. Write 0x000f000f to MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE to enable line interrupts.\n"
    "8. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL to enable boundary frame fatal interrupts.\n"
    "9. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL to enable sequence frame fatal interrupts.\n"
    "10. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL to enable CRC frame fatal interrupts.\n"
    "11. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL to enable payload CRC fatal interrupts.\n"
    "12. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID to enable data ID interrupts.\n"
    "13. Write 0x0000ffff to MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED to enable ECC corrected interrupts.\n"
    "14. Return from csi2_enable_interrupt().\n"
    "15. Compute vcid_csi2_wrap_reg based on GDMA path conditional compilation.\n"
    "16. Set gdma_path variable accordingly (0, 1, 2, or 3).\n"
    "17. Set gdma_reg_base = 0xE6A00000.\n"
    "18. Write vcid_csi2_wrap_reg to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL.\n"
    "19. Write 1 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to enable control data transfer.\n"
    "20. Write vcid_csi2_wrap_reg to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL again (repeated write).\n"
    "21. Write 1 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA again (repeated write).\n"
    "22. Call snps_phy_init() to initialize the D-PHY.\n"
    "23. Read MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE into rd_data.\n"
    "24. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while loop until rd_data equals 0x1000f.\n"
    "25. Set ch0_pc = 0xE6000000 and ch1_pc = 0xE6000500.\n"
    "26. Begin outer for-loop: lane_num iterates from 3 down to 0.\n"
    "27. Write lane_num to MIZAR_MIPI_CSI2_HOST_N_LANES.\n"
    "28. Compute cntrl_pkt_cnt = (VRES * 3) + 2.\n"
    "29. Write (lane_num + 1) to 0xa0243ffc to trigger the CSI-2 sequence.\n"
    "30. Begin inner for-loop: i iterates from 0 to cntrl_pkt_cnt - 1.\n"
    "31. Set ch0_preload_loc = ch0_pc and ch1_preload_loc = ch1_pc.\n"
    "32. Write 0x3 to gdma_reg_base + MIPI_CSI2_DMA_INTEN_OFFSET.\n"
    "33. Call dma_trnsfr_instn_preload() for DMA channel 0.\n"
    "34. Call DMAGO_CSI() to start DMA channel 0.\n"
    "35. Set rd_data = 0.\n"
    "36. Poll gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET until bit 0 is set.\n"
    "37. Write 0x1 to gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET.\n"
    "38. Read gdma_reg_base + 0x28 for debug status.\n"
    "39. Read 0xE6001000 into csi_ctrl_data.\n"
    "40. Check if (csi_ctrl_data & 0x3f) > 0xf for long packet.\n"
    "41. If long packet: Extract word_count from bits[21:6].\n"
    "42. Compute csi_data_size with 8-byte alignment.\n"
    "43. Call dma_trnsfr_instn_preload() for DMA channel 1.\n"
    "44. Call DMAGO_CSI() to start DMA channel 1.\n"
    "45. Set rd_data = 0.\n"
    "46. Poll gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET until bit 1 is set.\n"
    "47. Read gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET again.\n"
    "48. Write 0x2 to gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET.\n"
    "49. End of long-packet conditional block.\n"
    "50. End of inner for-loop.\n"
    "51. End of outer for-loop.\n"
    "52. Call finish(0) to end the test with pass status."
)

r1_meta_impacted_regs = (
    "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; "
    "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; "
    "MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED"
)

r1_meta_validation = (
    "The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until rd_data equals 0x1000f, confirming all D-PHY lanes have entered "
    "stop state. For each packet transfer, the DMA interrupt status register (gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) is "
    "polled: bit 0 must be set for DMA channel 0 completion, and bit 1 must be set for DMA channel 1 completion. The control "
    "data read from 0xE6001000 is checked: if (csi_ctrl_data & 0x3f) > 0xf, the packet is treated as a long packet and a data "
    "DMA transfer is initiated. The test calls finish(0) upon successful completion of all lane iterations (lane_num from 3 to 0), "
    "indicating pass. If any polling loop does not complete (hangs), the test fails by timeout."
)

# ══════════════════════════════════════════════════════════════
# ROW 2 DATA — mipi_csi2_test_pattern_generator
# ══════════════════════════════════════════════════════════════
r2_test_desc = (
    "This test validates the MIPI CSI-2 internal PPI Pattern Generator functionality with DMA-based data reception. "
    "The test configures the virtual channel register based on the selected GDMA path and disables control data transfer. "
    "CSI-2 and DMA interrupts are enabled. The D-PHY is initialized and the test waits for all PHY lanes to reach stop state "
    "by polling the PHY_STOPSTATE register. The DMA address mapping registers for channel 0 are programmed for read and write "
    "address paths (both data and instruction). A fractional divider output enable register is written. The DMA transfer is "
    "programmed with a computed frame size based on 320 horizontal pixels, 16 vertical lines, and 24 bits per pixel (RGB888), "
    "with 8-byte alignment per line. The DMA channel 0 is started. The Pattern Generator is then enabled with vertical resolution "
    "of 16 lines, horizontal resolution of 320 pixels, and RGB888 data type configuration. After a short wait, the Pattern "
    "Generator is disabled. The test polls the DMA interrupt status until channel 0 transfer completes, waits for a settling "
    "period, and then ends with a pass status."
)

r2_test_steps = (
    "1. Configure the virtual channel register with the appropriate virtual channel ID based on the selected GDMA path.\n"
    "2. Disable control data transfer by writing to the control_data register.\n"
    "3. Enable all CSI-2 host interrupts by reading the main interrupt status register to clear pending interrupts, "
    "then writing appropriate mask values to all interrupt mask registers (PHY fatal, packet fatal, PHY, line, "
    "boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, ECC corrected).\n"
    "4. Initialize the D-PHY by calling the PHY initialization sequence.\n"
    "5. Poll the PHY_STOPSTATE register until all lanes report stop state.\n"
    "6. Set frame parameters: horizontal resolution to 320 pixels, vertical resolution to 16 lines, and 24 bits per pixel (RGB888).\n"
    "7. Compute the total DMA transfer size based on the frame parameters with 8-byte alignment per line.\n"
    "8. Program the DMA address mapping registers for channel 0: set the higher-order AXI read address for data, "
    "read address for instructions, write address for data, and write address for instructions.\n"
    "9. Enable the fractional divider output to the CSI-2 subsystem via the subsystem register.\n"
    "10. Program and start DMA channel 0 with the computed transfer size, source address, and destination address.\n"
    "11. Enable the Pattern Generator by configuring vertical resolution, horizontal resolution, RGB888 data type, and enabling the PG.\n"
    "12. Wait for a short period, then disable the Pattern Generator.\n"
    "13. Poll the DMA interrupt status register until channel 0 transfer completes.\n"
    "14. Wait for a settling period.\n"
    "15. Verify the test completes successfully."
)

r2_impacted_regs = (
    "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; "
    "virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; "
    "dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; "
    "INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; "
    "INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; "
    "INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED"
)

r2_validation = (
    "1. The PHY_STOPSTATE register must read the expected stop-state value confirming all lanes including the clock lane are in stop state before proceeding.\n"
    "2. The Pattern Generator must be successfully enabled with the configured vertical resolution, horizontal resolution, and RGB888 data type.\n"
    "3. The DMA interrupt status must indicate channel 0 transfer completion after the Pattern Generator generates and transmits the frame data.\n"
    "4. The test must reach the final pass condition after the DMA transfer completes and the settling wait period elapses.\n"
    "5. No timeout must occur during PHY stop-state polling or DMA interrupt polling."
)

r2_remarks = (
    "The test uses the internal PPI Pattern Generator to generate a test frame (VRES=16, HRES=320, RGB888 24bpp) without "
    "requiring an external CSI-2 transmitter. The virtual channel ID and GDMA path are selected via compile-time conditional "
    "defines (GDMA0_PATH, GDMA1_PATH, GDMA2_PATH, GDMA3_PATH). Control data transfer is explicitly disabled (written 0) unlike "
    "the D-PHY lanes test which enables it. The DMA destination address increment flag is controlled by the FPS60 compile-time define. "
    "The Pattern Generator is enabled and then disabled after a short wait, allowing it to generate a single frame. "
    "A register at subsystem base + 0xf4 is written to enable fractional divider output. "
    "The csi2_subsys_enable_interrupt() function is called in the test but only csi2_enable_interrupt() is defined in the source file. "
    "The snps_phy_init(), dma_trnsfr_instn_preload_incr_addr(), DMAGO_CSI(), and wait_on() function implementations are not available "
    "in the supplied testcase folder."
)

r2_meta_test_desc = (
    "This testcase validates the MIPI CSI-2 internal PPI Pattern Generator with DMA-based data reception. "
    "The test begins by setting int_pend = 1 and vcid = 3. vcid_unselected_path is computed as ((vcid + 1) & 0xf). "
    "Based on conditional compilation (GDMA0_PATH/GDMA1_PATH/GDMA2_PATH/GDMA3_PATH), vcid_csi2_wrap_reg is computed by "
    "shifting vcid and vcid_unselected_path into the appropriate nibble positions, and gdma_path is set (0, 1, 2, or 3). "
    "gdma_reg_base is computed as 0xE6A00000 + (gdma_path * 0x1000). The virtual channel register "
    "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL is written with vcid_csi2_wrap_reg. Control data transfer is disabled by writing "
    "0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA. csi2_subsys_enable_interrupt() is called. snps_phy_init() is called. "
    "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE is polled until rd_data equals 0x1000f. hres=320, vres=16, valid_bits_per_pixel=24. "
    "csi2_data_trnsfr_size is computed as 15360 bytes. DMA address registers are programmed. "
    "MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 is written with 0x1. DMA channel 0 is started. "
    "csi2_ctrlr_pg_enable() writes PPI_PG_PATTERN_VRES with 0x10, PPI_PG_PATTERN_HRES with 0x70140, "
    "PPI_PG_CONFIG with 0xe401, PPI_PG_ENABLE with 1. wait_on(100). PPI_PG_ENABLE written with 0. "
    "DMA interrupt polled. wait_on(10000). finish(0)."
)

r2_meta_test_steps = (
    "1. Enter test_case() function.\n"
    "2. Set int_pend = 1.\n"
    "3. Set vcid = 3.\n"
    "4. Compute vcid_unselected_path = ((vcid + 1) & 0xf) = 0x4.\n"
    "5. Based on GDMA path conditional compilation, compute vcid_csi2_wrap_reg and set gdma_path.\n"
    "6. Compute gdma_reg_base = 0xE6A00000 + (gdma_path * 0x1000).\n"
    "7. Write vcid_csi2_wrap_reg to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL.\n"
    "8. Write 0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to disable control data transfer.\n"
    "9. Call csi2_subsys_enable_interrupt().\n"
    "10. Call snps_phy_init().\n"
    "11. Read MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE into rd_data.\n"
    "12. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until rd_data equals 0x1000f.\n"
    "13. Set hres = 320, vres = 16, valid_bits_per_pixel = 24.\n"
    "14. Compute csi2_data_trnsfr_size = 15360 bytes.\n"
    "15. Write 0x100 to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA.\n"
    "16. Write 0x0 to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION.\n"
    "17. Write 0x0 to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA.\n"
    "18. Write 0x0 to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION.\n"
    "19. Write 0x1 to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4.\n"
    "20. Set dma_ch0_pc = 0xE6000000.\n"
    "21. Set dma_dest_addr_incr_flag based on FPS60 conditional.\n"
    "22. Call dma_trnsfr_instn_preload_incr_addr().\n"
    "23. Call DMAGO_CSI() to start DMA channel 0.\n"
    "24. Call csi2_ctrlr_pg_enable().\n"
    "25. Write 0x10 to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES.\n"
    "26. Write 0x70140 to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES.\n"
    "27. Write 0xe401 to MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG.\n"
    "28. Write 1 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE.\n"
    "29. Return from csi2_ctrlr_pg_enable().\n"
    "30. Call wait_on(100).\n"
    "31. Write 0 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE.\n"
    "32. Read gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET.\n"
    "33. Poll until bit 0 is set.\n"
    "34. Call wait_on(10000).\n"
    "35. Call finish(0).\n"
    "36-46. csi2_enable_interrupt() writes all interrupt mask registers."
)

r2_meta_impacted_regs = (
    "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; "
    "MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; "
    "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; "
    "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; "
    "MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; "
    "MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; "
    "MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; "
    "MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED"
)

r2_meta_validation = (
    "The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until rd_data equals 0x1000f. "
    "The DMA interrupt status register is polled until bit 0 is set. "
    "The pattern generator is enabled then disabled after wait_on(100). "
    "The test calls finish(0) upon successful completion. "
    "If any polling loop does not complete, the test fails by timeout."
)

# ══════════════════════════════════════════════════════════════
# BUILD WORKBOOK
# ══════════════════════════════════════════════════════════════
wb = openpyxl.Workbook()

# ── TestPlan Sheet ──
ws_tp = wb.active
ws_tp.title = "TestPlan"

# Header row
for col_idx, col_name in enumerate(tp_cols, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")

# Row 1 data
r1_tp = [
    1, "MIPI_CSI", "D-PHY Lane Configuration", "mipi_csi2_dphy_lanes_test",
    r1_test_desc, "NA", "DMA Mode", "NA", "NA", r1_remarks,
    r1_test_steps, r1_impacted_regs, r1_validation, ""
]
for col_idx, val in enumerate(r1_tp, 1):
    cell = ws_tp.cell(row=2, column=col_idx, value=val)
    cell.alignment = Alignment(wrap_text=True, vertical="top")

# Row 2 data
r2_tp = [
    2, "MIPI_CSI", "PPI Pattern Generator", "mipi_csi2_test_pattern_generator",
    r2_test_desc, "NA", "DMA Mode", "NA", "NA", r2_remarks,
    r2_test_steps, r2_impacted_regs, r2_validation, ""
]
for col_idx, val in enumerate(r2_tp, 1):
    cell = ws_tp.cell(row=3, column=col_idx, value=val)
    cell.alignment = Alignment(wrap_text=True, vertical="top")

# Freeze first row
ws_tp.freeze_panes = "A2"

# Auto-size columns
for col_idx in range(1, len(tp_cols) + 1):
    max_len = len(str(tp_cols[col_idx - 1]))
    for row in ws_tp.iter_rows(min_row=2, max_row=3, min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                lines = str(cell.value).split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
    adjusted_width = min(max_len + 2, 60)
    ws_tp.column_dimensions[get_column_letter(col_idx)].width = max(adjusted_width, 12)

# ── MetaData Sheet ──
ws_md = wb.create_sheet("MetaData")

# Header row
for col_idx, col_name in enumerate(md_cols, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")

# Row 1 metadata
r1_md = [
    1, "mipi_csi2_dphy_lanes_test", r1_meta_test_desc, r1_meta_test_steps,
    r1_meta_impacted_regs, r1_meta_validation,
    '<stdio.h>; <stdlib.h>; "test_common.h"; "mipi_csi2.h"',
    "GDMA_CSI2_DATA_DEST_ADDR2; GDMA_CTRL_DATA_DEST_ADDR2; LS_LE_EN; TOTAL_FRAME; VC_ID; VRES; HRES; DATA_TYPE",
    "NA"
]
for col_idx, val in enumerate(r1_md, 1):
    cell = ws_md.cell(row=2, column=col_idx, value=val)
    cell.alignment = Alignment(wrap_text=True, vertical="top")

# Row 2 metadata
r2_md = [
    2, "mipi_csi2_test_pattern_generator", r2_meta_test_desc, r2_meta_test_steps,
    r2_meta_impacted_regs, r2_meta_validation,
    '<stdio.h>; <stdlib.h>; "test_common.h"; "mipi_csi2.h"',
    "GDMA_CSI2_DATA_DEST_ADDR2; GDMA_CTRL_DATA_DEST_ADDR2",
    "NA"
]
for col_idx, val in enumerate(r2_md, 1):
    cell = ws_md.cell(row=3, column=col_idx, value=val)
    cell.alignment = Alignment(wrap_text=True, vertical="top")

# Freeze first row
ws_md.freeze_panes = "A2"

# Auto-size columns
for col_idx in range(1, len(md_cols) + 1):
    max_len = len(str(md_cols[col_idx - 1]))
    for row in ws_md.iter_rows(min_row=2, max_row=3, min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                lines = str(cell.value).split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
    adjusted_width = min(max_len + 2, 60)
    ws_md.column_dimensions[get_column_letter(col_idx)].width = max(adjusted_width, 12)

# Set MetaData sheet to veryHidden
ws_md.sheet_state = "veryHidden"

# ── Save ──
wb.save(filepath)
print(f"SUCCESS: {filename}")
print(f"PATH: {filepath}")
print(f"SIZE: {os.path.getsize(filepath)} bytes")

# ── Validate ──
wb2 = openpyxl.load_workbook(filepath)
assert "TestPlan" in wb2.sheetnames, "TestPlan sheet missing"
assert "MetaData" in wb2.sheetnames, "MetaData sheet missing"
assert wb2["TestPlan"].max_row == 3, f"Expected 3 rows in TestPlan, got {wb2['TestPlan'].max_row}"
assert wb2["MetaData"].max_row == 3, f"Expected 3 rows in MetaData, got {wb2['MetaData'].max_row}"
assert wb2["MetaData"].sheet_state == "veryHidden", "MetaData not veryHidden"
print("VALIDATION: PASSED")
wb2.close()
