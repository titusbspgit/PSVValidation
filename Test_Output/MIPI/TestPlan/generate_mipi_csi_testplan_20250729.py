#!/usr/bin/env python3
"""MIPI_CSI TestPlan Excel Generator - Agent 7
Generates MIPI_CSI_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx with TestPlan and MetaData sheets.
Timestamp uses IST (GMT+05:30).
"""
import json, os, sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment

# === IST Timestamp ===
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
FILENAME = f"MIPI_CSI_TestPlan_{timestamp}.xlsx"
OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(OUTPUT_DIR, FILENAME)

# === JSON Data ===
json_data = [
  {
    "Index": "1",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_dphy_lanes_test",
    "Feature": "DPHY Lane Configuration and DMA Data Transfer",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_csi2.h\"",
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000; LS_LE_EN = 1; TOTAL_FRAME = 1; VC_ID = 3; VRES = 1080; VRES = 3; HRES = 1920; HRES = 64; DATA_TYPE",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates MIPI CSI-2 DPHY lane configuration by iterating through lane counts from 4 down to 1. For each lane configuration, it enables CSI-2 interrupts by reading MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts, then writes all interrupt mask registers (INT_MSK_PHY_FATAL with 0x0000000f, INT_MSK_PKT_FATAL with 0x00000003, INT_MSK_PHY with 0x000f000f, INT_MSK_LINE with 0x000f000f, INT_MSK_BNDRY_FRAME_FATAL with 0x0000ffff, INT_MSK_SEQ_FRAME_FATAL with 0x0000ffff, INT_MSK_CRC_FRAME_FATAL with 0x0000ffff, INT_MSK_PLD_CRC_FATAL with 0x0000ffff, INT_MSK_DATA_ID with 0x0000ffff, INT_MSK_ECC_CORRECTED with 0x0000ffff). It configures the virtual channel register MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with vcid_csi2_wrap_reg (derived from VC_ID shifted based on GDMA path selection), enables control data transfer by writing 1 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, calls snps_phy_init() for D-PHY initialization, then polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until it equals 0x1000f. In the lane loop, it writes MIZAR_MIPI_CSI2_HOST_N_LANES with the current lane_num (3 down to 0), writes 0xa0243ffc with (lane_num+1) to trigger the CSI-2 sequence, then for each control packet iterates: preloads DMA channel 0 instructions via dma_trnsfr_instn_preload(), triggers DMA via DMAGO_CSI(), polls gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET for bit 0, clears DMA IRQ via gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET with 0x1, reads control data from 0xE6001000, and if data_type > 0xf, calculates word_count and csi_data_size, preloads DMA channel 1 for data transfer, triggers DMAGO_CSI for channel 1, polls MIPI_CSI2_DMA_INTMIS_OFFSET for bit 1, and clears DMA IRQ with 0x2. The test completes by calling finish(0).",
    "Test Description": "This test validates MIPI CSI-2 DPHY lane configuration by iterating through lane counts from 4 lanes down to 1 lane. For each lane configuration, it enables all CSI-2 interrupt masks, configures the virtual channel register and enables control data transfer, initializes the D-PHY, and polls the PHY stop state register until the PHY enters stop state. It then configures the number of active lanes via the N_LANES register and triggers a CSI-2 sequence. For each control packet, it performs a DMA transfer of control data (channel 0), polls for DMA completion, reads the received control data, and if the data type indicates image data, performs a second DMA transfer of pixel data (channel 1) with calculated transfer size based on word count. The test verifies successful DMA completion for each transfer across all lane configurations.",
    "Meta Test Steps / Procedure": "1. Entry: test_case() is called.\n2. Call csi2_enable_interrupt():\n   2a. rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN) \u2014 read to clear pending interrupts.\n   2b. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, 0x0000000f) \u2014 enable PHY fatal interrupts.\n   2c. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, 0x00000003) \u2014 enable packet fatal interrupts.\n   2d. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, 0x000f000f) \u2014 enable PHY interrupts.\n   2e. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, 0x000f000f) \u2014 enable line interrupts.\n   2f. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, 0x0000ffff) \u2014 enable boundary frame fatal interrupts.\n   2g. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, 0x0000ffff) \u2014 enable sequence frame fatal interrupts.\n   2h. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, 0x0000ffff) \u2014 enable CRC frame fatal interrupts.\n   2i. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, 0x0000ffff) \u2014 enable payload CRC fatal interrupts.\n   2j. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, 0x0000ffff) \u2014 enable data ID interrupts.\n   2k. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED, 0x0000ffff) \u2014 enable ECC corrected interrupts.\n3. Determine vcid_csi2_wrap_reg based on GDMA path selection (GDMA3_PATH: VC_ID, GDMA2_PATH: VC_ID << 4, GDMA1_PATH: VC_ID << 8, GDMA0_PATH: VC_ID << 12). Set gdma_path accordingly.\n4. Set gdma_reg_base = 0xE6A00000.\n5. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg) \u2014 configure virtual channel.\n6. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 1) \u2014 enable control data transfer.\n7. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg) \u2014 re-write virtual channel register.\n8. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 1) \u2014 re-enable control data transfer.\n9. Call snps_phy_init() \u2014 D-PHY initialization sequence.\n10. rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) \u2014 initial read of PHY stop state.\n11. Poll: while rd_data != 0x1000f, repeatedly read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) until PHY enters stop state (value equals 0x1000f).\n12. Set ch0_pc = 0xE6000000, ch1_pc = 0xE6000500.\n13. Begin outer loop: for lane_num = 3 down to 0 (4 iterations for 4-lane, 3-lane, 2-lane, 1-lane):\n   13a. write_reg(MIZAR_MIPI_CSI2_HOST_N_LANES, lane_num) \u2014 configure number of active lanes.\n   13b. Calculate cntrl_pkt_cnt = (VRES * 3) + 2.\n   13c. write_reg(0xa0243ffc, (lane_num + 1)) \u2014 trigger CSI-2 sequence for current lane count.\n   13d. Begin inner loop: for i = 0 to cntrl_pkt_cnt - 1 (iterate for each packet):\n      13d-i. Set ch0_preload_loc = ch0_pc, ch1_preload_loc = ch1_pc.\n      13d-ii. write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTEN_OFFSET, 0x3) \u2014 enable DMA IRQ[1] and IRQ[0].\n      13d-iii. Call dma_trnsfr_instn_preload(ch0_preload_loc, gdma_reg_base, 0x8000, 0xE6001000, 8, 0).\n      13d-iv. Call DMAGO_CSI(gdma_reg_base, ch0_pc, 0x0) \u2014 trigger DMA GO for channel 0.\n      13d-v. Poll: while (rd_data & 0x1) == 0x0, read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET).\n      13d-vi. write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET, 0x1) \u2014 clear DMA IRQ for channel 0.\n      13d-vii. rd_data = read_reg(gdma_reg_base + 0x28).\n      13d-viii. csi_ctrl_data = read_reg(0xE6001000) \u2014 read received control data.\n      13d-ix. Condition: if (csi_ctrl_data & 0x3f) > 0xf (data type indicates image data):\n         13d-ix-a. Calculate word_count = (csi_ctrl_data >> 6) & 0xffff.\n         13d-ix-b. Calculate csi_data_size = (word_count % 8) ? (word_count / 8 + 1) * 8 : word_count.\n         13d-ix-c. Call dma_trnsfr_instn_preload(ch1_preload_loc, gdma_reg_base, 0x0000, 0xE6002000, csi_data_size, 1).\n         13d-ix-d. Call DMAGO_CSI(gdma_reg_base, ch1_pc, 0x1) \u2014 trigger DMA GO for channel 1.\n         13d-ix-e. Poll: while (rd_data & 0x2) == 0x0, read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET).\n         13d-ix-f. rd_data = read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET).\n         13d-ix-g. write_reg(gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET, 0x2) \u2014 clear DMA IRQ for channel 1.\n14. Call finish(0) \u2014 test completion with pass status.",
    "Test Steps / Procedure": "1. Enable all CSI-2 interrupt masks by reading the main interrupt status register to clear pending interrupts, then writing enable values to all interrupt mask registers (PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, ECC corrected).\n2. Configure the virtual channel register with the appropriate virtual channel ID based on the selected DMA path.\n3. Enable control data transfer by writing to the control data register.\n4. Initialize the D-PHY by calling the PHY initialization sequence.\n5. Poll the PHY stop state register until the PHY enters the expected stop state.\n6. For each lane configuration (4 lanes down to 1 lane):\n   a. Write the number of active lanes to the N_LANES register.\n   b. Trigger the CSI-2 sequence for the current lane count.\n   c. For each expected control packet:\n      i. Enable DMA interrupts for both channels.\n      ii. Preload DMA channel 0 instructions and trigger DMA transfer for control data.\n      iii. Poll the DMA interrupt status register until channel 0 transfer completes.\n      iv. Clear the DMA interrupt for channel 0.\n      v. Read the received control data.\n      vi. If the data type indicates image data, calculate the data transfer size, preload DMA channel 1 instructions, trigger DMA transfer for pixel data, poll until channel 1 completes, and clear the DMA interrupt for channel 1.\n7. Complete the test with pass status.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Impacted Registers": "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "1. MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE is polled until rd_data equals 0x1000f, confirming PHY has entered stop state on all lanes and clock lane.\n2. DMA channel 0 completion is validated by polling gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET until (rd_data & 0x1) != 0x0, confirming control data DMA transfer completed.\n3. DMA channel 1 completion is validated by polling gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET until (rd_data & 0x2) != 0x0, confirming pixel data DMA transfer completed.\n4. Control data is read from 0xE6001000 and checked: if (csi_ctrl_data & 0x3f) > 0xf, the data type indicates image data requiring a channel 1 DMA transfer.\n5. word_count is extracted as (csi_ctrl_data >> 6) & 0xffff and csi_data_size is aligned to 8-byte boundary.\n6. DMA IRQ is cleared after each transfer: 0x1 for channel 0, 0x2 for channel 1.\n7. The test iterates through all 4 lane configurations (lane_num 3 to 0) and processes (VRES*3)+2 packets per configuration.\n8. Test passes by calling finish(0) after all lane configurations and all packet transfers complete successfully.",
    "Validation / Acceptance Criteria": "1. The PHY stop state register must indicate that all data lanes and the clock lane have entered stop state before proceeding.\n2. DMA channel 0 transfer must complete successfully for each control packet, confirmed by the DMA interrupt status register indicating channel 0 completion.\n3. When the received control data indicates image data (data type greater than a threshold), DMA channel 1 transfer must also complete successfully, confirmed by the DMA interrupt status register indicating channel 1 completion.\n4. DMA interrupts must be properly cleared after each transfer completion.\n5. All four lane configurations (4-lane, 3-lane, 2-lane, 1-lane) must complete their respective packet transfers without error.\n6. The test must complete with a pass status after all lane configurations are processed.",
    "Remarks": "The test uses conditional compilation (GDMA0_PATH, GDMA1_PATH, GDMA2_PATH, GDMA3_PATH) to select the DMA path and corresponding virtual channel ID shift. VRES and HRES values depend on GDMA0_FULL_MEM compilation flag (1080x1920 or 3x64). Two hex addresses (0xa0243ffc and 0xE6001000) could not be mapped to named registers. The DMA base address gdma_reg_base is set to 0xE6A00000. DMA offset macros (MIPI_CSI2_DMA_INTEN_OFFSET, MIPI_CSI2_DMA_INTMIS_OFFSET, MIPI_CSI2_DMA_INTCLR_OFFSET) are used for DMA register access but were not included in Agent 2 tokens. The snps_phy_init(), dma_trnsfr_instn_preload(), and DMAGO_CSI() function implementations are external to this testcase source file."
  },
  {
    "Index": "2",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_test_pattern_generator",
    "Feature": "Internal Test Pattern Generator (PPI PG) and DMA Data Transfer",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_csi2.h\"",
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates the MIPI CSI-2 internal test pattern generator (PPI PG) functionality with DMA data transfer. The test configures the PPI pattern generator to produce a 320x16 RGB888 pattern (24 bits per pixel). It sets up the CSI-2 subsystem by configuring the virtual channel register MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with vcid_csi2_wrap_reg (derived from vcid=3 shifted based on GDMA path selection), disables control data transfer by writing 0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, enables CSI-2 and DMA interrupts via csi2_subsys_enable_interrupt(), initializes the D-PHY via snps_phy_init(), and polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until it equals 0x1000f. It then calculates the data transfer size based on resolution and bits per pixel (aligned to 8-byte boundary), programs DMA address registers (MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA with 0x100, MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION with 0x0, MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA with 0x0, MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION with 0x0), enables fracdiv output by writing 0x1 to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4, preloads DMA channel 0 instructions via dma_trnsfr_instn_preload_incr_addr(), triggers DMA via DMAGO_CSI(), enables the pattern generator by calling csi2_ctrlr_pg_enable() which writes MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES with 0x10, MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES with 0x70140, MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG with 0xe401, and MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE with 1. After a wait_on(100), the pattern generator is disabled by writing 0 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE. The test then polls gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET until bit 0 is set (DMA channel 0 completion), waits wait_on(10000), and calls finish(0). The csi2_enable_interrupt() function reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts, then writes all interrupt mask registers with specific enable values.",
    "Test Description": "This test validates the MIPI CSI-2 internal test pattern generator (PPI PG) with DMA data transfer. It configures the CSI-2 subsystem by setting the virtual channel, disabling control data transfer, enabling all CSI-2 interrupt masks, and initializing the D-PHY. After polling the PHY stop state register until the PHY enters stop state, it calculates the DMA transfer size for a 320x16 pixel RGB888 pattern (24 bits per pixel, 8-byte aligned). It programs the DMA address registers for channel 0 read and write paths, enables the fractional divider output, and preloads DMA channel 0 transfer instructions. The DMA transfer is triggered, then the pattern generator is enabled with the configured vertical resolution, horizontal resolution, and pattern configuration. After a short wait, the pattern generator is disabled. The test polls the DMA interrupt status register until channel 0 transfer completes, confirming successful reception of the generated test pattern data via DMA.",
    "Meta Test Steps / Procedure": "1. Entry: test_case() is called.\n2. Set int_pend = 1.\n3. Set vcid = 3, vcid_unselected_path = ((vcid + 1) & 0xf) = 4.\n4. Determine vcid_csi2_wrap_reg based on GDMA path selection.\n5. Set gdma_reg_base = 0xE6A00000 + ((gdma_path) * 0x1000).\n6. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg).\n7. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 0).\n8. Call csi2_subsys_enable_interrupt().\n9. Call csi2_enable_interrupt():\n   9a. rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN).\n   9b-9k. write_reg all interrupt mask registers.\n10. Call snps_phy_init().\n11. rd_data = read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE).\n12. Poll: while rd_data != 0x1000f.\n13. Set hres = 320, vres = 16, valid_bits_per_pixel = 24.\n14. Calculate csi2_data_trnsfr_size.\n15-18. write_reg DMA address registers.\n19. write_reg(MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4, 0x1).\n20-23. DMA preload and DMAGO.\n24. Call csi2_ctrlr_pg_enable():\n   24a-24d. Configure PG registers and enable.\n25. wait_on(100).\n26. write_reg(MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE, 0).\n27-28. Poll DMA completion.\n29. wait_on(10000).\n30. finish(0).",
    "Test Steps / Procedure": "1. Configure the virtual channel register with the appropriate virtual channel ID based on the selected DMA path.\n2. Disable control data transfer by writing to the control data register.\n3. Enable all CSI-2 interrupt masks by reading the main interrupt status register to clear pending interrupts, then writing enable values to all interrupt mask registers.\n4. Initialize the D-PHY by calling the PHY initialization sequence.\n5. Poll the PHY stop state register until the PHY enters the expected stop state.\n6. Calculate the DMA transfer size for a 320x16 pixel pattern at 24 bits per pixel, aligned to 8-byte boundary.\n7. Program the DMA address registers for channel 0 read and write paths.\n8. Enable the fractional divider output to the CSI-2 subsystem.\n9. Preload DMA channel 0 transfer instructions with source, destination, and transfer size, then trigger DMA GO for channel 0.\n10. Enable the internal test pattern generator by configuring vertical resolution (16 lines), horizontal resolution (320 pixels), pattern configuration (RGB888 format), and setting the enable bit.\n11. Wait briefly, then disable the pattern generator.\n12. Poll the DMA interrupt status register until channel 0 transfer completes (bit 0 set).\n13. Wait for post-transfer settling, then complete the test with pass status.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "1. MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE is polled until rd_data equals 0x1000f.\n2. DMA channel 0 completion is validated by polling gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET until (rd_data & 0x1) != 0.\n3. The pattern generator is enabled with MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE = 1 and subsequently disabled with MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE = 0 after wait_on(100).\n4. The PG configuration uses VRES=0x10 (16 lines), HRES=0x70140, CONFIG=0xe401 (RGB888 data type 0x24).\n5. DMA transfer size is calculated as (valid_bits_per_pixel * hres / 8) per line, aligned to 8-byte boundary, multiplied by vres.\n6. Test passes by calling finish(0) after DMA transfer completion and wait_on(10000) settling period.",
    "Validation / Acceptance Criteria": "1. The PHY stop state register must indicate that all data lanes and the clock lane have entered stop state before proceeding with pattern generation.\n2. DMA channel 0 transfer must complete successfully, confirmed by the DMA interrupt status register indicating channel 0 completion (bit 0 set).\n3. The pattern generator must be enabled and then disabled in a controlled sequence, producing a 320x16 pixel RGB888 test pattern.\n4. The DMA transfer size must match the expected data volume based on the configured resolution and bits per pixel.\n5. The test must complete with a pass status after DMA transfer completion and post-transfer settling.",
    "Remarks": "The test uses conditional compilation (GDMA0_PATH, GDMA1_PATH, GDMA2_PATH, GDMA3_PATH) to select the DMA path and corresponding virtual channel ID shift. The FPS60 flag controls whether the DMA destination address increment is enabled. The pattern generator produces RGB888 data (data type 0x24) at 320x16 resolution with 24 bits per pixel. The MIZAR_MIPI_CSI2_RB_REG_BASE token used in the expression (MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4) could not be mapped to a named register by Agent 4. The csi2_subsys_enable_interrupt() function is called from test_case() but its implementation is external; csi2_enable_interrupt() is defined locally. The snps_phy_init(), dma_trnsfr_instn_preload_incr_addr(), DMAGO_CSI(), wait_on(), and finish() function implementations are external to this testcase source file."
  }
]

# === TestPlan Sheet Columns ===
testplan_cols = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

# === MetaData Sheet Columns ===
metadata_cols = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

# === Create Workbook ===
wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = "TestPlan"

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_align = Alignment(wrap_text=True, vertical="top")

# Write headers
for col_idx, col_name in enumerate(testplan_cols, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

# Write data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(testplan_cols, 1):
        value = row_data.get(col_name, "")
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_align

# --- MetaData Sheet ---
ws_md = wb.create_sheet(title="MetaData")

for col_idx, col_name in enumerate(metadata_cols, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(metadata_cols, 1):
        value = row_data.get(col_name, "")
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_align

# === Freeze first row ===
ws_tp.freeze_panes = "A2"
ws_md.freeze_panes = "A2"

# === Auto-size columns ===
def auto_size(ws, max_width=80):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                lines = str(cell.value).split('\n')
                for line in lines:
                    max_len = max(max_len, len(line))
        adjusted = min(max_len + 2, max_width)
        ws.column_dimensions[col_letter].width = max(adjusted, 12)

auto_size(ws_tp)
auto_size(ws_md)

# === Set MetaData sheet to veryHidden ===
ws_md.sheet_state = "veryHidden"

# === Save ===
wb.save(OUTPUT_PATH)
print(f"Workbook saved: {OUTPUT_PATH}")
print(f"Filename: {FILENAME}")

# === Validate ===
assert os.path.exists(OUTPUT_PATH), "File does not exist!"
assert os.path.getsize(OUTPUT_PATH) > 0, "File is empty!"
wb2 = load_workbook(OUTPUT_PATH)
assert "TestPlan" in wb2.sheetnames, "TestPlan sheet missing!"
assert "MetaData" in wb2.sheetnames, "MetaData sheet missing!"
print(f"Validation PASSED. Size: {os.path.getsize(OUTPUT_PATH)} bytes")
print(f"TestPlan rows: {ws_tp.max_row - 1}")
print(f"MetaData rows: {ws_md.max_row - 1}")
