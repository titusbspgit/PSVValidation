#!/usr/bin/env python3
"""MIPI_CSI TestPlan Excel Generator - Agent 7 Fallback Automation"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone, timedelta
import os
import sys

def main():
    ist = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(ist)
    timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
    filename = f"MIPI_CSI_TestPlan_{timestamp}.xlsx"
    output_dir = os.environ.get("OUTPUT_DIR", ".")
    filepath = os.path.join(output_dir, filename)

    json_data = [
      {
        "Index": "1",
        "SS / Module": "MIPI_CSI",
        "Test Case Name": "mipi_csi2_dphy_lanes_test",
        "Feature": "DPHY Lane Configuration and Interrupt Enablement",
        "Meta Headers": "NA",
        "Meta Macros": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED; MIPI_CSI2_DMA_INTEN_OFFSET; MIPI_CSI2_DMA_INTMIS_OFFSET; MIPI_CSI2_DMA_INTCLR_OFFSET",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates MIPI CSI2 DPHY lane configuration and CSI2 host interrupt enablement. The test_case() function first writes to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to configure the virtual channel, then writes to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to set control data. A second set of writes to the same registers is performed. The test then polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while loop to wait for the PHY stop state condition. Once the PHY stop state is confirmed, MIZAR_MIPI_CSI2_HOST_N_LANES is written to configure the number of active DPHY lanes. A write to hardcoded address 0xa0243ffc is performed. The test calls csi2_enable_interrupt() which reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to check the main interrupt status, then writes to all CSI2 host interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED) to enable all interrupt masks. DMA interrupt enable, masked interrupt status read, and interrupt clear operations are performed using gdma_reg_base with offsets MIPI_CSI2_DMA_INTEN_OFFSET, MIPI_CSI2_DMA_INTMIS_OFFSET, and MIPI_CSI2_DMA_INTCLR_OFFSET. A read from hardcoded address 0xE6001000 is also performed. The DMA masked interrupt status is polled and cleared in a loop.",
        "Test Description": "This test validates the MIPI CSI2 DPHY lane configuration and CSI2 host interrupt enablement flow. The test configures the virtual channel and control data registers in the CSI2 register block. It then polls the PHY_STOPSTATE register to wait for the PHY lanes to reach the stop state. Once confirmed, the N_LANES register is written to configure the number of active DPHY lanes. The test enables all CSI2 host interrupt masks by writing to INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED registers after reading the INT_ST_MAIN register. DMA interrupt enable, status polling, and interrupt clear operations are also performed. The test verifies that the DPHY lane configuration and interrupt setup complete successfully.",
        "Meta Test Steps / Procedure": "1. Enter test_case() function. 2. Write to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to configure the virtual channel setting. 3. Write to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to set the control data value. 4. Write to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL again with a second virtual channel configuration. 5. Write to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA again with a second control data value. 6. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while loop, waiting for the PHY stop state condition to be met. 7. Write to MIZAR_MIPI_CSI2_HOST_N_LANES to configure the number of active DPHY lanes. 8. Write to hardcoded address 0xa0243ffc. 9. Write to gdma_reg_base + MIPI_CSI2_DMA_INTEN_OFFSET to enable DMA interrupts. 10. Call csi2_enable_interrupt() function. 11. Inside csi2_enable_interrupt(): Read MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to check the main interrupt status. 12. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL to enable PHY fatal interrupt mask. 13. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL to enable packet fatal interrupt mask. 14. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY to enable PHY interrupt mask. 15. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE to enable line interrupt mask. 16. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL to enable boundary frame fatal interrupt mask. 17. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL to enable sequence frame fatal interrupt mask. 18. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL to enable CRC frame fatal interrupt mask. 19. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL to enable payload CRC fatal interrupt mask. 20. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID to enable data ID interrupt mask. 21. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED to enable ECC corrected interrupt mask. 22. Return from csi2_enable_interrupt(). 23. Read from gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET to check DMA masked interrupt status. 24. Read from gdma_reg_base + 0x28 to read a DMA-related register. 25. Read from hardcoded address 0xE6001000. 26. Poll gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET in a loop to wait for DMA interrupt completion. 27. Write to gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET to clear DMA interrupts. 28. Read gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET again to verify DMA interrupt cleared. 29. Write to gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET again if needed. 30. Test completes and returns result.",
        "Test Steps / Procedure": "1. Configure the virtual channel register with the first virtual channel setting. 2. Write the first control data value to the control data register. 3. Configure the virtual channel register with a second virtual channel setting. 4. Write the second control data value to the control data register. 5. Poll the PHY_STOPSTATE register until the PHY lanes reach the stop state. 6. Configure the N_LANES register to set the number of active DPHY lanes. 7. Perform a write to an external configuration address. 8. Enable DMA interrupts via the DMA interrupt enable register. 9. Read the INT_ST_MAIN register to check the main interrupt status. 10. Enable all CSI2 host interrupt masks by writing to INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED registers. 11. Read the DMA masked interrupt status register to check for DMA completion. 12. Read an external status register. 13. Poll the DMA masked interrupt status register until DMA transfer completes. 14. Clear DMA interrupts by writing to the DMA interrupt clear register. 15. Verify DMA interrupts are cleared by re-reading the DMA masked interrupt status register. 16. Confirm test completion.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
        "Impacted Registers": "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
        "Meta Validation / Acceptance Criteria": "The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while loop to validate that the PHY lanes have reached the stop state before proceeding with lane configuration. After enabling all CSI2 host interrupt masks, the DMA masked interrupt status register (gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) is polled in a loop to wait for DMA transfer completion. The DMA interrupt is then cleared via gdma_reg_base + MIPI_CSI2_DMA_INTCLR_OFFSET, and the masked interrupt status is re-read to verify the interrupt has been cleared. The test passes if PHY stop state is achieved, lane configuration completes, all interrupt masks are enabled, and DMA transfer completes with interrupts properly cleared.",
        "Validation / Acceptance Criteria": "The test passes when: 1) The PHY_STOPSTATE register indicates that all configured DPHY lanes have reached the stop state. 2) The N_LANES register is successfully configured with the desired lane count. 3) All CSI2 host interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED) are enabled. 4) DMA transfer completes as indicated by the DMA masked interrupt status register. 5) DMA interrupts are properly cleared after completion.",
        "Remarks": "The test uses polling on PHY_STOPSTATE to wait for DPHY lane readiness before configuring the number of lanes. Two hardcoded hex addresses (0xa0243ffc for write and 0xE6001000 for read) could not be mapped to canonical register names from the provided specification documents. DMA operations use a base address variable (gdma_reg_base) with offset macros for interrupt enable, masked interrupt status, and interrupt clear registers. The test enables a comprehensive set of CSI2 host interrupt masks covering PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected interrupt categories."
      },
      {
        "Index": "2",
        "SS / Module": "MIPI_CSI",
        "Test Case Name": "mipi_csi2_test_pattern_generator",
        "Feature": "Test Pattern Generation",
        "Meta Headers": "NA",
        "Meta Macros": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED; MIPI_CSI2_DMA_INTMIS_OFFSET",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates the MIPI CSI2 internal test pattern generator functionality. The csi2_ctrlr_pg_enable() function configures the pattern generator by writing the vertical resolution to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES, the horizontal resolution to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES, the pattern generator configuration to MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG, and enables the pattern generator by writing to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE. The test_case() function writes to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to configure the virtual channel and writes to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to set control data. It polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while loop to wait for the PHY stop state condition. DMA channel 0 read and write address registers are configured by writing to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA, MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION, MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA, and MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION. A write to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 is performed. The pattern generator is disabled by writing 0 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE. The csi2_enable_interrupt() function reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN and writes to all CSI2 host interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED). DMA masked interrupt status is polled via gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET.",
        "Test Description": "This test validates the MIPI CSI2 internal test pattern generator. The test configures the pattern generator vertical resolution, horizontal resolution, and configuration registers, then enables the pattern generator via the PPI_PG_ENABLE register. The virtual_channel and control_data registers are configured. The PHY_STOPSTATE register is polled to wait for PHY lane readiness. DMA channel 0 address registers (dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, dma_m0_addr_aw_ch0_Instruction) are configured for data transfer. All CSI2 host interrupt masks are enabled by writing to INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED after reading INT_ST_MAIN. The pattern generator is then disabled and DMA interrupt status is polled for completion.",
        "Meta Test Steps / Procedure": "1. Enter csi2_ctrlr_pg_enable() function. 2. Write to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES to set the vertical resolution for the test pattern. 3. Write to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES to set the horizontal resolution for the test pattern. 4. Write to MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG to configure the pattern generator settings. 5. Write to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE to enable the pattern generator. 6. Return from csi2_ctrlr_pg_enable(). 7. Enter test_case() function. 8. Write to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to configure the virtual channel. 9. Write to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to set the control data value. 10. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while loop, waiting for the PHY stop state condition to be met. 11. Write to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA to configure DMA channel 0 read address for data. 12. Write to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION to configure DMA channel 0 read address for instructions. 13. Write to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA to configure DMA channel 0 write address for data. 14. Write to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION to configure DMA channel 0 write address for instructions. 15. Write to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 to configure an additional register block register. 16. Call csi2_enable_interrupt() function. 17. Inside csi2_enable_interrupt(): Read MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to check the main interrupt status. 18. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL to enable PHY fatal interrupt mask. 19. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL to enable packet fatal interrupt mask. 20. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY to enable PHY interrupt mask. 21. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE to enable line interrupt mask. 22. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL to enable boundary frame fatal interrupt mask. 23. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL to enable sequence frame fatal interrupt mask. 24. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL to enable CRC frame fatal interrupt mask. 25. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL to enable payload CRC fatal interrupt mask. 26. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID to enable data ID interrupt mask. 27. Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED to enable ECC corrected interrupt mask. 28. Return from csi2_enable_interrupt(). 29. Write 0 to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE to disable the pattern generator. 30. Read gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET to check DMA masked interrupt status. 31. Poll gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET in a loop to wait for DMA completion. 32. Test completes and returns result.",
        "Test Steps / Procedure": "1. Configure the test pattern generator vertical resolution by writing to the PPI_PG_PATTERN_VRES register. 2. Configure the test pattern generator horizontal resolution by writing to the PPI_PG_PATTERN_HRES register. 3. Configure the test pattern generator settings by writing to the PPI_PG_CONFIG register. 4. Enable the test pattern generator by writing to the PPI_PG_ENABLE register. 5. Configure the virtual_channel register with the desired virtual channel setting. 6. Write the control data value to the control_data register. 7. Poll the PHY_STOPSTATE register until the PHY lanes reach the stop state. 8. Configure DMA channel 0 read address registers by writing to dma_m0_addr_ar_ch0_data and dma_m0_addr_ar_ch0_Instruction registers. 9. Configure DMA channel 0 write address registers by writing to dma_m0_addr_aw_ch0_data and dma_m0_addr_aw_ch0_Instruction registers. 10. Configure an additional register block register for DMA operation. 11. Read the INT_ST_MAIN register to check the main interrupt status. 12. Enable all CSI2 host interrupt masks by writing to INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED registers. 13. Disable the test pattern generator by writing to the PPI_PG_ENABLE register. 14. Poll the DMA masked interrupt status register until DMA transfer completes. 15. Confirm test completion.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
        "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
        "Meta Validation / Acceptance Criteria": "The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while loop to validate that the PHY lanes have reached the stop state before proceeding with DMA and pattern generator configuration. After enabling all CSI2 host interrupt masks, the DMA masked interrupt status register (gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) is polled in a loop to wait for DMA transfer completion. The pattern generator is enabled via MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE and later disabled by writing 0 to the same register. The test passes if the PHY stop state is achieved, the pattern generator is successfully configured and enabled, DMA addresses are set, all interrupt masks are enabled, and DMA transfer completes successfully.",
        "Validation / Acceptance Criteria": "The test passes when: 1) The PHY_STOPSTATE register indicates that all configured DPHY lanes have reached the stop state. 2) The test pattern generator is successfully configured with the desired vertical resolution (PPI_PG_PATTERN_VRES), horizontal resolution (PPI_PG_PATTERN_HRES), and configuration (PPI_PG_CONFIG), and enabled via PPI_PG_ENABLE. 3) DMA channel 0 address registers (dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, dma_m0_addr_aw_ch0_Instruction) are properly configured. 4) All CSI2 host interrupt mask registers are enabled. 5) DMA transfer completes as indicated by the DMA masked interrupt status register. 6) The pattern generator is successfully disabled after the test.",
        "Remarks": "The test uses polling on PHY_STOPSTATE to wait for DPHY lane readiness before configuring DMA and pattern generator. The pattern generator is enabled and then disabled within the test flow, validating the enable/disable cycle. DMA operations use a base address variable (gdma_reg_base) with offset macros for masked interrupt status. A write to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 targets an additional register in the register block. The test enables a comprehensive set of CSI2 host interrupt masks covering PHY fatal, packet fatal, PHY, line, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected interrupt categories. Source files in the repository folder did not contain the expected MIPI CSI2 test pattern generator code; testcase details are derived from upstream agent analysis."
      }
    ]

    testplan_columns = [
        "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
        "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
        "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
        "Code Generation"
    ]

    metadata_columns = [
        "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
        "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
        "Meta Headers", "Meta Macros", "Meta Arrays"
    ]

    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    # TestPlan sheet
    ws_tp = wb.active
    ws_tp.title = "TestPlan"
    for col_idx, col_name in enumerate(testplan_columns, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(testplan_columns, 1):
            value = row_data.get(col_name, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
    ws_tp.freeze_panes = "A2"
    for col_idx, col_name in enumerate(testplan_columns, 1):
        max_len = len(col_name)
        for row in ws_tp.iter_rows(min_row=2, max_row=ws_tp.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    for line in str(cell.value).split('\n'):
                        if len(line) > max_len:
                            max_len = len(line)
        ws_tp.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 2, 60)

    # MetaData sheet
    ws_md = wb.create_sheet("MetaData")
    for col_idx, col_name in enumerate(metadata_columns, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(metadata_columns, 1):
            value = row_data.get(col_name, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
    ws_md.freeze_panes = "A2"
    for col_idx, col_name in enumerate(metadata_columns, 1):
        max_len = len(col_name)
        for row in ws_md.iter_rows(min_row=2, max_row=ws_md.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    for line in str(cell.value).split('\n'):
                        if len(line) > max_len:
                            max_len = len(line)
        ws_md.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 2, 60)
    ws_md.sheet_state = "veryHidden"

    wb.save(filepath)
    print(f"SUCCESS: {filepath}")
    print(f"FILENAME: {filename}")

    # Validation
    wb_check = openpyxl.load_workbook(filepath)
    assert "TestPlan" in wb_check.sheetnames
    assert "MetaData" in wb_check.sheetnames
    ws_tp_c = wb_check["TestPlan"]
    ws_md_c = wb_check["MetaData"]
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(testplan_columns, 1):
            expected = row_data.get(col_name, "")
            actual = ws_tp_c.cell(row=row_idx, column=col_idx).value or ""
            assert str(expected) == str(actual), f"TestPlan mismatch at row {row_idx}, col {col_name}"
        for col_idx, col_name in enumerate(metadata_columns, 1):
            expected = row_data.get(col_name, "")
            actual = ws_md_c.cell(row=row_idx, column=col_idx).value or ""
            assert str(expected) == str(actual), f"MetaData mismatch at row {row_idx}, col {col_name}"
    wb_check.close()
    print("VALIDATION: PASSED")
    return filename

if __name__ == "__main__":
    main()
