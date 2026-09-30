#!/usr/bin/env python3
"""Agent 7 - MIPI_CSI TestPlan XLSX Generator
Generates MIPI_CSI_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx using openpyxl.
Timestamp uses IST (GMT+05:30).
"""
import json
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

# ---- INPUT DATA ----
json_data = [
    {
        "Index": "1",
        "SS / Module": "MIPI_CSI",
        "Test Case Name": "mipi_csi2_dphy_lanes_test",
        "Feature": "D-PHY Lane Configuration",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates MIPI CSI2 D-PHY lane configuration. It writes to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to configure virtual channel mapping, writes to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to configure control data, polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE to wait for PHY stop state, writes to MIZAR_MIPI_CSI2_HOST_N_LANES to configure the number of active D-PHY lanes, reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to check main interrupt status, and writes to all CSI2 host interrupt mask registers (MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED) to enable interrupt masks. Additionally, a write to address 0xa0243ffc and a read from address 0xE6001000 are performed.",
        "Test Description": "This test validates MIPI CSI2 D-PHY lane configuration. It configures the virtual channel mapping via the virtual_channel register, sets up the control_data register, polls the PHY_STOPSTATE register to confirm the PHY has reached stop state, and writes to the N_LANES register to configure the number of active D-PHY data lanes. The test reads the INT_ST_MAIN register to check the main interrupt status and enables all CSI2 host interrupt masks including INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED. Two additional register accesses are performed to unresolved hardware addresses.",
        "Meta Test Steps / Procedure": "Step 1: Write to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL (base 0xE6A04000, offset 0x0) to configure the virtual channel mapping for CSI2 reception. Step 2: Write to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA (base 0xE6A04000, offset 0x20) to configure the control data settings. Step 3: Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE (base 0xE6A05000, offset 0x4C) to wait until the D-PHY reaches stop state, indicating the PHY lanes are idle and ready for reconfiguration. Step 4: Write to MIZAR_MIPI_CSI2_HOST_N_LANES (base 0xE6A05000, offset 0x4) to configure the number of active D-PHY data lanes. Step 5: Write to address 0xa0243ffc to perform an auxiliary configuration write. Step 6: Read from address 0xE6001000 to retrieve status or configuration data. Step 7: Read MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN (base 0xE6A05000, offset 0xC) to check the main interrupt status register for any pending interrupts. Step 8: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL (base 0xE6A05000, offset 0xE4) to enable PHY fatal interrupt mask. Step 9: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL (base 0xE6A05000, offset 0xF4) to enable packet fatal interrupt mask. Step 10: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY (base 0xE6A05000, offset 0x114) to enable PHY interrupt mask. Step 11: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE (base 0xE6A05000, offset 0x134) to enable line interrupt mask. Step 12: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL (base 0xE6A05000, offset 0x284) to enable boundary frame fatal interrupt mask. Step 13: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL (base 0xE6A05000, offset 0x294) to enable sequence frame fatal interrupt mask. Step 14: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL (base 0xE6A05000, offset 0x2A4) to enable CRC frame fatal interrupt mask. Step 15: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL (base 0xE6A05000, offset 0x2B4) to enable payload CRC fatal interrupt mask. Step 16: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID (base 0xE6A05000, offset 0x2C4) to enable data ID interrupt mask. Step 17: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED (base 0xE6A05000, offset 0x2D4) to enable ECC corrected interrupt mask.",
        "Test Steps / Procedure": "1. Configure the virtual channel mapping by writing to the virtual_channel register.\n2. Configure the control data settings by writing to the control_data register.\n3. Poll the PHY_STOPSTATE register until the D-PHY reaches stop state, confirming lanes are idle.\n4. Write to the N_LANES register to set the desired number of active D-PHY data lanes.\n5. Perform an auxiliary configuration write to an external hardware address.\n6. Read from an external hardware address to retrieve status information.\n7. Read the INT_ST_MAIN register to check for any pending main interrupt status.\n8. Enable the PHY fatal interrupt mask by writing to INT_MSK_PHY_FATAL.\n9. Enable the packet fatal interrupt mask by writing to INT_MSK_PKT_FATAL.\n10. Enable the PHY interrupt mask by writing to INT_MSK_PHY.\n11. Enable the line interrupt mask by writing to INT_MSK_LINE.\n12. Enable the boundary frame fatal interrupt mask by writing to INT_MSK_BNDRY_FRAME_FATAL.\n13. Enable the sequence frame fatal interrupt mask by writing to INT_MSK_SEQ_FRAME_FATAL.\n14. Enable the CRC frame fatal interrupt mask by writing to INT_MSK_CRC_FRAME_FATAL.\n15. Enable the payload CRC fatal interrupt mask by writing to INT_MSK_PLD_CRC_FATAL.\n16. Enable the data ID interrupt mask by writing to INT_MSK_DATA_ID.\n17. Enable the ECC corrected interrupt mask by writing to INT_MSK_ECC_CORRECTED.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
        "Impacted Registers": "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
        "Meta Validation / Acceptance Criteria": "The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE to validate that the D-PHY has entered stop state before proceeding with lane reconfiguration. MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN is read to verify the main interrupt status. All interrupt mask registers are written to enable the corresponding interrupt sources. Successful completion requires the PHY stop state poll to resolve and all register writes to complete without error.",
        "Validation / Acceptance Criteria": "The PHY_STOPSTATE register must indicate that the D-PHY lanes have reached stop state before lane configuration proceeds. The INT_ST_MAIN register read must return a valid interrupt status. All interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED) must be successfully written to enable the corresponding interrupt sources. The N_LANES register must accept the configured lane count. The test passes if all register operations complete successfully and the PHY stop state condition is met.",
        "Remarks": "The source code in the repository for this testcase folder does not contain the actual MIPI CSI2 D-PHY lanes test implementation; the program.c file contains unrelated PCIe test code. The testcase JSON was generated using Agent 2, Agent 3, and Agent 4 outputs only. Two register accesses (0xa0243ffc write and 0xE6001000 read) could not be resolved to canonical register names by Agent 4. The PHY_STOPSTATE register is accessed via polling, indicating a wait condition for D-PHY lane idle state."
    },
    {
        "Index": "2",
        "SS / Module": "MIPI_CSI",
        "Test Case Name": "mipi_csi2_test_pattern_generator",
        "Feature": "Test Pattern Generator",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates the MIPI CSI2 internal test pattern generator functionality. It writes to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES (base 0xE6A05000, offset 0x60) to configure the vertical resolution of the test pattern, writes to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES (base 0xE6A05000, offset 0x64) to configure the horizontal resolution, writes to MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG (base 0xE6A05000, offset 0x68) to configure the pattern generator settings including data type and virtual channel, and writes to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE (base 0xE6A05000, offset 0x6C) to enable or disable the pattern generator. The test writes to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL (base 0xE6A04000, offset 0x0) to configure virtual channel mapping and writes to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA (base 0xE6A04000, offset 0x20) to configure control data. It polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE (base 0xE6A05000, offset 0x4C) to wait for the D-PHY to reach stop state. DMA address registers are configured: MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA (offset 0x6C), MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION (offset 0x70), MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA (offset 0x8C), and MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION (offset 0x90). The test reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN (base 0xE6A05000, offset 0xC) to check main interrupt status and writes to all CSI2 host interrupt mask registers: MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL (offset 0xE4), MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL (offset 0xF4), MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY (offset 0x114), MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE (offset 0x134), MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL (offset 0x284), MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL (offset 0x294), MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL (offset 0x2A4), MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL (offset 0x2B4), MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID (offset 0x2C4), and MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED (offset 0x2D4) to enable interrupt masks.",
        "Test Description": "This test validates the MIPI CSI2 internal test pattern generator. It configures the pattern generator vertical resolution via PPI_PG_PATTERN_VRES, horizontal resolution via PPI_PG_PATTERN_HRES, pattern configuration via PPI_PG_CONFIG, and enables the generator via PPI_PG_ENABLE. The virtual channel mapping is set through the virtual_channel register and control data is configured via the control_data register. The test polls PHY_STOPSTATE to confirm the D-PHY has reached stop state. DMA address registers are configured including dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, and dma_m0_addr_aw_ch0_Instruction for data transfer. The main interrupt status is checked via INT_ST_MAIN, and all CSI2 host interrupt masks are enabled including INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, and INT_MSK_ECC_CORRECTED.",
        "Meta Test Steps / Procedure": "Step 1: Write to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES (base 0xE6A05000, offset 0x60) to configure the vertical resolution of the test pattern. Step 2: Write to MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES (base 0xE6A05000, offset 0x64) to configure the horizontal resolution of the test pattern. Step 3: Write to MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG (base 0xE6A05000, offset 0x68) to configure the pattern generator settings including data type and virtual channel selection. Step 4: Write to MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE (base 0xE6A05000, offset 0x6C) to enable the test pattern generator. Step 5: Write to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL (base 0xE6A04000, offset 0x0) to configure the virtual channel mapping for CSI2 reception. Step 6: Write to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA (base 0xE6A04000, offset 0x20) to configure the control data settings. Step 7: Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE (base 0xE6A05000, offset 0x4C) to wait until the D-PHY reaches stop state, indicating the PHY lanes are idle and ready. Step 8: Write to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA (base 0xE6A04000, offset 0x6C) to configure the DMA read address for channel 0 data. Step 9: Write to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION (base 0xE6A04000, offset 0x70) to configure the DMA read address for channel 0 instruction. Step 10: Write to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA (base 0xE6A04000, offset 0x8C) to configure the DMA write address for channel 0 data. Step 11: Write to MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION (base 0xE6A04000, offset 0x90) to configure the DMA write address for channel 0 instruction. Step 12: Read MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN (base 0xE6A05000, offset 0xC) to check the main interrupt status register for any pending interrupts. Step 13: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL (base 0xE6A05000, offset 0xE4) to enable PHY fatal interrupt mask. Step 14: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL (base 0xE6A05000, offset 0xF4) to enable packet fatal interrupt mask. Step 15: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY (base 0xE6A05000, offset 0x114) to enable PHY interrupt mask. Step 16: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE (base 0xE6A05000, offset 0x134) to enable line interrupt mask. Step 17: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL (base 0xE6A05000, offset 0x284) to enable boundary frame fatal interrupt mask. Step 18: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL (base 0xE6A05000, offset 0x294) to enable sequence frame fatal interrupt mask. Step 19: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL (base 0xE6A05000, offset 0x2A4) to enable CRC frame fatal interrupt mask. Step 20: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL (base 0xE6A05000, offset 0x2B4) to enable payload CRC fatal interrupt mask. Step 21: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID (base 0xE6A05000, offset 0x2C4) to enable data ID interrupt mask. Step 22: Write to MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED (base 0xE6A05000, offset 0x2D4) to enable ECC corrected interrupt mask.",
        "Test Steps / Procedure": "1. Configure the test pattern vertical resolution by writing to PPI_PG_PATTERN_VRES.\n2. Configure the test pattern horizontal resolution by writing to PPI_PG_PATTERN_HRES.\n3. Configure the pattern generator settings including data type and virtual channel by writing to PPI_PG_CONFIG.\n4. Enable the test pattern generator by writing to PPI_PG_ENABLE.\n5. Configure the virtual channel mapping by writing to the virtual_channel register.\n6. Configure the control data settings by writing to the control_data register.\n7. Poll the PHY_STOPSTATE register until the D-PHY reaches stop state, confirming lanes are idle.\n8. Configure the DMA read address for channel 0 data by writing to dma_m0_addr_ar_ch0_data.\n9. Configure the DMA read address for channel 0 instruction by writing to dma_m0_addr_ar_ch0_Instruction.\n10. Configure the DMA write address for channel 0 data by writing to dma_m0_addr_aw_ch0_data.\n11. Configure the DMA write address for channel 0 instruction by writing to dma_m0_addr_aw_ch0_Instruction.\n12. Read the INT_ST_MAIN register to check for any pending main interrupt status.\n13. Enable the PHY fatal interrupt mask by writing to INT_MSK_PHY_FATAL.\n14. Enable the packet fatal interrupt mask by writing to INT_MSK_PKT_FATAL.\n15. Enable the PHY interrupt mask by writing to INT_MSK_PHY.\n16. Enable the line interrupt mask by writing to INT_MSK_LINE.\n17. Enable the boundary frame fatal interrupt mask by writing to INT_MSK_BNDRY_FRAME_FATAL.\n18. Enable the sequence frame fatal interrupt mask by writing to INT_MSK_SEQ_FRAME_FATAL.\n19. Enable the CRC frame fatal interrupt mask by writing to INT_MSK_CRC_FRAME_FATAL.\n20. Enable the payload CRC fatal interrupt mask by writing to INT_MSK_PLD_CRC_FATAL.\n21. Enable the data ID interrupt mask by writing to INT_MSK_DATA_ID.\n22. Enable the ECC corrected interrupt mask by writing to INT_MSK_ECC_CORRECTED.",
        "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
        "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
        "Meta Validation / Acceptance Criteria": "The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE to validate that the D-PHY has entered stop state before proceeding with test pattern generation and DMA configuration. MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN is read to verify the main interrupt status. The pattern generator registers MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES, MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES, MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG, and MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE must be written successfully to configure and enable the test pattern generator. All interrupt mask registers must be written to enable the corresponding interrupt sources. DMA address registers for channel 0 read and write paths must be configured successfully. Successful completion requires the PHY stop state poll to resolve, the pattern generator to be enabled, DMA addresses to be configured, and all register writes to complete without error.",
        "Validation / Acceptance Criteria": "The PHY_STOPSTATE register must indicate that the D-PHY lanes have reached stop state before test pattern generation proceeds. The pattern generator must be successfully configured with the desired vertical resolution via PPI_PG_PATTERN_VRES, horizontal resolution via PPI_PG_PATTERN_HRES, and pattern settings via PPI_PG_CONFIG, and enabled via PPI_PG_ENABLE. The DMA address registers (dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, dma_m0_addr_aw_ch0_Instruction) must accept the configured addresses. The INT_ST_MAIN register read must return a valid interrupt status. All interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED) must be successfully written. The test passes if all register operations complete successfully, the PHY stop state condition is met, and the pattern generator is enabled.",
        "Remarks": "The source code in the repository for this testcase folder contains unrelated PCIe test code rather than the actual MIPI CSI2 test pattern generator implementation. The testcase JSON was generated using Agent 2, Agent 3, and Agent 4 outputs only. All 22 register-access tokens were successfully resolved by Agent 3 and mapped to canonical register names by Agent 4 with no unresolved tokens. The PHY_STOPSTATE register is accessed via polling, indicating a wait condition for D-PHY lane idle state. The SOFT_RST_REG_ADDRESS macro was excluded per instructions."
    }
]

# ---- CONFIGURATION ----
IP_NAME = "MIPI_CSI"
OUTPUT_DIR = "Test_Output/MIPI/TestPlan"

TESTPLAN_COLUMNS = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

METADATA_COLUMNS = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

def generate_xlsx():
    IST = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(IST)
    timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
    filename = f"{IP_NAME}_TestPlan_{timestamp}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename) if os.path.isdir(OUTPUT_DIR) else filename

    wb = Workbook()

    # ---- TestPlan Sheet ----
    ws_tp = wb.active
    ws_tp.title = "TestPlan"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_alignment = Alignment(wrap_text=True, vertical="top")

    for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment

    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
            value = row_data.get(col_name, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment

    ws_tp.freeze_panes = "A2"

    # ---- MetaData Sheet ----
    ws_md = wb.create_sheet(title="MetaData")

    for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment

    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
            value = row_data.get(col_name, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment

    ws_md.freeze_panes = "A2"
    ws_md.sheet_state = "veryHidden"

    # ---- Auto-size columns ----
    MAX_WIDTH = 60
    for ws in [ws_tp, ws_md]:
        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            col_letter = get_column_letter(col_idx)
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=False):
                for cell in row:
                    if cell.value:
                        lines = str(cell.value).split("\n")
                        for line in lines:
                            max_len = max(max_len, len(line))
            adjusted = min(max_len + 2, MAX_WIDTH)
            ws.column_dimensions[col_letter].width = max(adjusted, 12)

    # ---- Save ----
    wb.save(filepath)
    print(f"Workbook saved: {filepath}")
    print(f"Filename: {filename}")

    # ---- Validate ----
    assert os.path.exists(filepath), f"File not found: {filepath}"
    assert os.path.getsize(filepath) > 0, "File is empty"
    vwb = load_workbook(filepath)
    assert "TestPlan" in vwb.sheetnames, "TestPlan sheet missing"
    assert "MetaData" in vwb.sheetnames, "MetaData sheet missing"
    tp_rows = vwb["TestPlan"].max_row - 1
    md_rows = vwb["MetaData"].max_row - 1
    print(f"Validation PASSED: TestPlan rows={tp_rows}, MetaData rows={md_rows}")
    print(f"File size: {os.path.getsize(filepath)} bytes")
    vwb.close()

    return filepath, filename

if __name__ == "__main__":
    fp, fn = generate_xlsx()
    print(f"SUCCESS: {fn}")
