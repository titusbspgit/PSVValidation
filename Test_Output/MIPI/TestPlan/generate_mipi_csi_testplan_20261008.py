#!/usr/bin/env python3
"""
MIPI_CSI TestPlan XLSX Generator - Agent 7
==========================================
Run this script to generate the genuine XLSX workbook:
  python3 generate_mipi_csi_testplan_20261008.py

Output: MIPI_CSI_TestPlan_YYYYMMDD_HHMMSS.xlsx (IST timestamp)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import os, sys

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
ts = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"MIPI_CSI_TestPlan_{ts}.xlsx"
out_dir = os.path.dirname(os.path.abspath(__file__))
filepath = os.path.join(out_dir, filename)

# ═══════════════════════════════════════════════════════════════
# TEST PLAN DATA (2 test cases)
# ═══════════════════════════════════════════════════════════════
json_data = [
  {
    "Index": "1",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_dphy_lanes_test",
    "Feature": "D-PHY Lane Configuration",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_csi2.h\"",
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000; LS_LE_EN = 1; TOTAL_FRAME = 1; VC_ID = 3; VRES; HRES; DATA_TYPE",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates MIPI CSI-2 D-PHY lane configuration by iterating through lane counts from 4 lanes down to 1 lane. It first calls csi2_enable_interrupt() which reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts, then writes all CSI-2 host interrupt mask registers (INT_MSK_PHY_FATAL with 0x0000000f, INT_MSK_PKT_FATAL with 0x00000003, INT_MSK_PHY with 0x000f000f, INT_MSK_LINE with 0x000f000f, INT_MSK_BNDRY_FRAME_FATAL with 0x0000ffff, INT_MSK_SEQ_FRAME_FATAL with 0x0000ffff, INT_MSK_CRC_FRAME_FATAL with 0x0000ffff, INT_MSK_PLD_CRC_FATAL with 0x0000ffff, INT_MSK_DATA_ID with 0x0000ffff, INT_MSK_ECC_CORRECTED with 0x0000ffff). The test then configures the virtual channel register MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with the computed vcid_csi2_wrap_reg value based on the selected GDMA path (VC_ID shifted by 0, 4, 8, or 12 bits depending on GDMA3/2/1/0_PATH), and writes 1 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to enable control data transfer. These two writes are performed twice. snps_phy_init() is called for D-PHY initialization. The test then polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until it reads 0x1000f, confirming PHY stop-state entry. A for loop iterates lane_num from 3 down to 0: each iteration writes lane_num to MIZAR_MIPI_CSI2_HOST_N_LANES, writes (lane_num+1) to address 0xa0243ffc to trigger the CSI-2 sequence, then enters an inner loop for cntrl_pkt_cnt = (VRES*3)+2 iterations. In each inner iteration: DMA interrupt enable is written (0x3 to gdma_reg_base+MIPI_CSI2_DMA_INTEN_OFFSET), dma_trnsfr_instn_preload is called for channel 0 with src_addr=0x8000, dest_addr=0xE6001000, trnsfr_size=8, irq_num=0, then DMAGO_CSI is called for channel 0. The test polls gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET until bit 0 is set, then clears the interrupt by writing 0x1 to gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET. It reads gdma_reg_base+0x28 for debug status, then reads 0xE6001000 for csi_ctrl_data. If (csi_ctrl_data & 0x3f) > 0xf, it extracts word_count from bits [21:6], calculates csi_data_size aligned to 8 bytes, calls dma_trnsfr_instn_preload for channel 1 with src_addr=0x0000, dest_addr=0xE6002000, trnsfr_size=csi_data_size, irq_num=1, then DMAGO_CSI for channel 1. It polls MIPI_CSI2_DMA_INTMIS_OFFSET until bit 1 is set, reads INTMIS again, then clears by writing 0x2 to INTCLR. After all iterations, finish(0) is called.",
    "Test Description": "This test validates MIPI CSI-2 D-PHY lane configuration by iterating through lane counts from 4 lanes down to 1 lane. It enables all CSI-2 host interrupts by reading INT_ST_MAIN to clear pending status and writing enable masks to all interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED). It configures the virtual_channel register with the appropriate virtual channel ID and enables control data transfer via the control_data register. After D-PHY initialization, it polls PHY_STOPSTATE to confirm the PHY has entered stop-state. For each lane configuration (4 to 1), it writes the lane count to N_LANES, triggers the CSI-2 sequence, and performs DMA-based control and data packet transfers. Each packet transfer uses DMA channel 0 for control data, polling for DMA completion via interrupt status. If the received control data indicates a long packet (data type greater than short-packet threshold), a data transfer is initiated on DMA channel 1, again polling for completion. The test verifies successful DMA transfers across all lane configurations.",
    "Meta Test Steps / Procedure": "1. Entry: test_case() is called. 2. Call csi2_enable_interrupt(). 3. Inside csi2_enable_interrupt(): read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN) to clear pending interrupts, store in rd_data. 4. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL, 0x0000000f) to enable PHY fatal interrupts. 5. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL, 0x00000003) to enable packet fatal interrupts. 6. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY, 0x000f000f) to enable PHY interrupts. 7. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE, 0x000f000f) to enable line interrupts. 8. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL, 0x0000ffff) to enable boundary frame fatal interrupts. 9. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL, 0x0000ffff) to enable sequence frame fatal interrupts. 10. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL, 0x0000ffff) to enable CRC frame fatal interrupts. 11. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL, 0x0000ffff) to enable payload CRC fatal interrupts. 12. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID, 0x0000ffff) to enable data ID interrupts. 13. write_reg(MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED, 0x0000ffff) to enable ECC corrected interrupts. 14. Return from csi2_enable_interrupt(). 15. Conditional compilation selects GDMA path. 16. gdma_reg_base = 0xE6A00000. 17-20. Write virtual_channel and control_data registers twice. 21. Call snps_phy_init(). 22-23. Poll PHY_STOPSTATE until 0x1000f. 24-51. Lane loop (3 to 0) with DMA transfers, long packet handling. 51. Call finish(0).",
    "Test Steps / Procedure": "1. Enable CSI-2 host interrupts by reading INT_ST_MAIN to clear pending status, then writing enable masks to all interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED). 2. Configure the virtual_channel register with the virtual channel ID for the selected DMA path. 3. Enable control data transfer by writing to the control_data register. 4. Perform D-PHY initialization. 5. Poll PHY_STOPSTATE register until the PHY enters stop-state on all configured lanes. 6. For each lane configuration from 4 lanes down to 1 lane: write the lane count to N_LANES register and trigger the CSI-2 sequence. 7. For each control packet in the current lane configuration: enable DMA interrupts, preload DMA channel 0 transfer instructions for control data, and start the DMA transfer. 8. Poll DMA interrupt status for channel 0 completion, then clear the DMA interrupt. 9. Read the received CSI control data from the destination address. 10. If the control data indicates a long packet, extract the word count, compute the aligned data size, preload DMA channel 1 transfer instructions for CSI data, and start the DMA transfer. 11. Poll DMA interrupt status for channel 1 completion, then clear the DMA interrupt. 12. Repeat steps 7-11 for all packets in the current lane configuration. 13. Repeat steps 6-12 for all lane configurations. 14. Verify test completion with a pass indication.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000",
    "Impacted Registers": "INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED; virtual_channel; control_data; PHY_STOPSTATE; N_LANES",
    "Meta Validation / Acceptance Criteria": "1. PHY_STOPSTATE polling: read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) must return 0x1000f to confirm all lanes and clock lane are in stop-state before proceeding. The test loops until rd_data == 0x1000f. 2. DMA channel 0 completion: read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) is polled until (rd_data & 0x1) != 0x0, indicating DMA channel 0 transfer is complete. 3. DMA channel 1 completion (conditional): read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) is polled until (rd_data & 0x2) != 0x0, indicating DMA channel 1 transfer is complete. 4. Long packet detection: (csi_ctrl_data & 0x3f) > 0xf determines whether the received control data represents a long packet requiring data transfer on channel 1. 5. Test pass: finish(0) is called after all lane configurations and all packet transfers complete successfully, indicating overall test pass.",
    "Validation / Acceptance Criteria": "1. PHY_STOPSTATE register must read the expected stop-state value confirming all lanes and clock lane are in stop-state before lane configuration begins. 2. DMA channel 0 interrupt status must indicate completion for each control packet transfer. 3. If the received control data indicates a long packet, DMA channel 1 interrupt status must indicate completion for each data packet transfer. 4. The test iterates through all four lane configurations (4, 3, 2, 1 lanes) and all packets per configuration without failure. 5. The test completes successfully with a pass indication after all lane configurations are validated.",
    "Remarks": "The virtual_channel and control_data registers are written twice in sequence, which appears intentional in the source. The snps_phy_init() and dma_trnsfr_instn_preload() and DMAGO_CSI() function implementations are not available in this testcase folder. DMA register accesses use a variable base (gdma_reg_base) with offset macros (MIPI_CSI2_DMA_INTEN_OFFSET, MIPI_CSI2_DMA_INTMIS_OFFSET, MIPI_CSI2_DMA_INTCLR_OFFSET) and a hardcoded offset (0x28). Two direct hex addresses (0xa0243ffc and 0xE6001000) could not be mapped to named registers in the provided specification documents. VRES and HRES are conditionally compiled based on GDMA0_FULL_MEM."
  },
  {
    "Index": "2",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_test_pattern_generator",
    "Feature": "Test Pattern Generator",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_csi2.h\"",
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates the MIPI CSI-2 internal test pattern generator functionality. It configures the CSI-2 subsystem to receive data generated by the host controller's built-in pattern generator and transfers the received data via DMA. The test begins by setting int_pend = 1, vcid = 3, and computing vcid_unselected_path = ((vcid + 1) & 0xf). Based on conditional compilation (GDMA3_PATH, GDMA2_PATH, GDMA1_PATH, or GDMA0_PATH), vcid_csi2_wrap_reg is computed by shifting vcid and vcid_unselected_path into the appropriate nibble positions, and gdma_path is set accordingly. gdma_reg_base is set to 0xE6A00000 + (gdma_path * 0x1000). The test writes vcid_csi2_wrap_reg to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL and writes 0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to disable control data transfer. csi2_subsys_enable_interrupt() is called (which internally calls csi2_enable_interrupt()). Inside csi2_enable_interrupt(): read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN) clears pending interrupts, then writes all interrupt mask registers. snps_phy_init() is called for D-PHY initialization. The test polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until it reads 0x1000f. hres is set to 320, vres to 16, valid_bits_per_pixel to 24. csi2_data_trnsfr_size is computed. DMA address registers are programmed. Then 0x1 is written to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 to enable sending fracdiv output to the CSI-2 subsystem. DMA channel 0 is started. csi2_ctrlr_pg_enable() writes PPI_PG_PATTERN_VRES=0x10, PPI_PG_PATTERN_HRES=0x70140, PPI_PG_CONFIG=0xe401, PPI_PG_ENABLE=1. wait_on(100) is called. Then PPI_PG_ENABLE is written 0. DMA completion is polled. finish(0) is called.",
    "Test Description": "This test validates the MIPI CSI-2 internal test pattern generator by configuring the pattern generator to produce a 320x16 pixel RGB888 frame and transferring the received data via DMA. It configures the virtual_channel register with the virtual channel ID for the selected DMA path and disables control data transfer via the control_data register. All CSI-2 host interrupts are enabled by reading INT_ST_MAIN to clear pending status and writing enable masks to all interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED). After D-PHY initialization, it polls PHY_STOPSTATE to confirm stop-state entry. DMA address mapping registers (dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, dma_m0_addr_aw_ch0_Instruction) are programmed. The enableclkgating_csiphy register is written to enable fracdiv output. DMA channel 0 is configured and started for the data transfer. The pattern generator is enabled by writing vertical resolution to PPI_PG_PATTERN_VRES, horizontal resolution to PPI_PG_PATTERN_HRES, configuration to PPI_PG_CONFIG, and enabling via PPI_PG_ENABLE. After a wait, the pattern generator is disabled by writing to PPI_PG_ENABLE. The test polls DMA interrupt status for channel 0 completion and finishes with a pass indication.",
    "Meta Test Steps / Procedure": "1. Entry: test_case() is called. 2. Set int_pend = 1. 3. Set vcid = 3. 4. Compute vcid_unselected_path = ((vcid + 1) & 0xf). 5. Conditional compilation selects GDMA path. 6. gdma_reg_base = 0xE6A00000 + (gdma_path * 0x1000). 7. write_reg(MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL, vcid_csi2_wrap_reg). 8. write_reg(MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA, 0). 9. Call csi2_subsys_enable_interrupt(). 10-20. Enable all interrupt masks. 21. Return from csi2_enable_interrupt(). 22. Call snps_phy_init(). 23-24. Poll PHY_STOPSTATE until 0x1000f. 25-26. Compute transfer size. 27-30. Program DMA address registers. 31. write_reg(MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4, 0x1). 32-36. Start DMA transfer. 37-42. Enable pattern generator. 43-44. Wait and disable PG. 45-46. Poll DMA completion. 47-48. Wait and finish(0).",
    "Test Steps / Procedure": "1. Configure the virtual_channel register with the virtual channel ID for the selected DMA path and disable control data transfer via the control_data register. 2. Enable CSI-2 host interrupts by reading INT_ST_MAIN to clear pending status, then writing enable masks to all interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED). 3. Perform D-PHY initialization. 4. Poll PHY_STOPSTATE register until the PHY enters stop-state on all configured lanes. 5. Compute the total frame data transfer size for a 320x16 RGB888 frame aligned to 8 bytes. 6. Program DMA address mapping registers (dma_m0_addr_ar_ch0_data, dma_m0_addr_ar_ch0_Instruction, dma_m0_addr_aw_ch0_data, dma_m0_addr_aw_ch0_Instruction) for higher-order AXI address bits. 7. Write to the enableclkgating_csiphy register to enable fracdiv output to the CSI-2 subsystem. 8. Preload DMA channel 0 transfer instructions and start the DMA transfer. 9. Enable the pattern generator by writing vertical resolution to PPI_PG_PATTERN_VRES, horizontal resolution to PPI_PG_PATTERN_HRES, configuration to PPI_PG_CONFIG, and enabling via PPI_PG_ENABLE. 10. Wait, then disable the pattern generator by writing to PPI_PG_ENABLE. 11. Poll DMA interrupt status for channel 0 completion. 12. Wait and verify test completion with a pass indication.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; enableclkgating_csiphy; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "1. PHY_STOPSTATE polling: read_reg(MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE) must return 0x1000f to confirm all lanes and clock lane are in stop-state before proceeding. The test loops until rd_data == 0x1000f. 2. DMA channel 0 completion: read_reg(gdma_reg_base + MIPI_CSI2_DMA_INTMIS_OFFSET) is polled until (rd_data & 0x1) != 0, indicating DMA channel 0 transfer is complete. 3. Pattern generator enable/disable sequence: MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE is written with 1 to enable and then 0 to disable the pattern generator, with a wait_on(100) delay between enable and disable. 4. Test pass: finish(0) is called after DMA completion and a wait_on(10000) delay, indicating overall test pass.",
    "Validation / Acceptance Criteria": "1. PHY_STOPSTATE register must read the expected stop-state value confirming all lanes and clock lane are in stop-state before pattern generator activation. 2. DMA channel 0 interrupt status must indicate completion, confirming the pattern-generated frame data was successfully transferred. 3. The pattern generator must be enabled and then disabled in sequence with proper timing. 4. The test completes successfully with a pass indication after DMA transfer completion.",
    "Remarks": "The test uses the internal test pattern generator (PPI_PG) to generate a 320x16 pixel RGB888 frame (VRES=16, HRES=320, valid_bits_per_pixel=24). The pattern generator configuration value 0xe401 encodes data type RGB888 and pattern mode settings. The MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 composite access is mapped to the enableclkgating_csiphy register. The snps_phy_init(), dma_trnsfr_instn_preload_incr_addr(), DMAGO_CSI(), csi2_subsys_enable_interrupt(), wait_on(), and finish() function implementations are not available in this testcase folder. DMA register accesses use a variable base (gdma_reg_base) with offset macro (MIPI_CSI2_DMA_INTMIS_OFFSET). The control_data register is written with 0 to disable control data transfer, unlike the dphy_lanes_test which enables it. VRES and HRES values are hardcoded in test_case() as 16 and 320 respectively. The dma_dest_addr_incr_flag is conditionally compiled based on FPS60."
  }
]

# ═══════════════════════════════════════════════════════════════
# COLUMN DEFINITIONS
# ═══════════════════════════════════════════════════════════════
tp_cols = ["Index","SS / Module","Feature","Test Case Name","Test Description",
           "Speed","Mode","Memory Start Offset","Memory End Offset","Remarks",
           "Test Steps / Procedure","Impacted Registers",
           "Validation / Acceptance Criteria","Code Generation"]

md_cols = ["Index","Test Case Name","Meta Test Description",
           "Meta Test Steps / Procedure","Meta Impacted Registers",
           "Meta Validation / Acceptance Criteria","Meta Headers",
           "Meta Macros","Meta Arrays"]

# ═══════════════════════════════════════════════════════════════
# CREATE WORKBOOK
# ═══════════════════════════════════════════════════════════════
wb = openpyxl.Workbook()
ws_tp = wb.active
ws_tp.title = "TestPlan"
ws_md = wb.create_sheet("MetaData")

hf = Font(bold=True, color="FFFFFF", size=11)
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wa = Alignment(wrap_text=True, vertical="top")

def write_sheet(ws, columns, data):
    for ci, cn in enumerate(columns, 1):
        c = ws.cell(row=1, column=ci, value=cn)
        c.font = hf; c.fill = hfill; c.alignment = wa
    for ri, rd in enumerate(data, 2):
        for ci, cn in enumerate(columns, 1):
            c = ws.cell(row=ri, column=ci, value=rd.get(cn, ""))
            c.alignment = wa
    ws.freeze_panes = "A2"
    for ci, cn in enumerate(columns, 1):
        ml = len(cn)
        for ri in range(2, len(data)+2):
            v = ws.cell(row=ri, column=ci).value
            if v:
                for ln in str(v).split('\n'):
                    ml = max(ml, len(ln))
        ws.column_dimensions[get_column_letter(ci)].width = min(ml+4, 80)

write_sheet(ws_tp, tp_cols, json_data)
write_sheet(ws_md, md_cols, json_data)
ws_md.sheet_state = "veryHidden"

wb.save(filepath)
wb.close()

# ═══════════════════════════════════════════════════════════════
# VERIFY
# ═══════════════════════════════════════════════════════════════
sz = os.path.getsize(filepath)
wb2 = openpyxl.load_workbook(filepath)
print(f"SUCCESS: {filename} ({sz} bytes)")
print(f"Sheets: {wb2.sheetnames}")
print(f"TestPlan rows: {ws_tp.max_row - 1}")
print(f"MetaData rows: {ws_md.max_row - 1}")
wb2.close()
