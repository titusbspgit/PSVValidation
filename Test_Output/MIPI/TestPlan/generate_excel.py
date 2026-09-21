#!/usr/bin/env python3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone, timedelta
import os
import json
import sys

json_data = [
    {
        "Index": "1",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_basic_test",
        "Feature": "DBI Write Memory via DMA",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a basic MIPI DSI DBI write memory operation using DMA. It enables DMAC interrupts by writing to MIZAR_MIPI_DSI_DMAC_INTEN, enables subsystem-level GDMA interrupt via MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE, configures the DSI host PHY interface via MIZAR_MIPI_DSI_HOST_PHY_IF_CFG (including PHY stop wait time), configures packet handling via MIZAR_MIPI_DSI_HOST_PCKHDL_CFG, sets up the clock manager via MIZAR_MIPI_DSI_HOST_CLKMGR_CFG, and controls DPI via MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL. The DMA channel is programmed using debug instruction registers MIZAR_MIPI_DSI_DMAC_DBGINST0 and MIZAR_MIPI_DSI_DMAC_DBGINST1, then executed via MIZAR_MIPI_DSI_DMAC_DBGCMD. The test reads MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK to verify the GDMA interrupt mask status, reads MIZAR_MIPI_DSI_DMAC_INTMIS to check DMA interrupt completion, clears the DMA interrupt via MIZAR_MIPI_DSI_DMAC_INTCLR, and clears the subsystem interrupt via MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW.",
        "Test Description": "This test validates the basic MIPI DSI DBI write memory transfer using DMA. It configures the DSI host PHY interface (PHY_IF_CFG), packet handling (PCKHDL_CFG), and clock manager (CLKMGR_CFG). The subsystem interrupt enable register (interrupt_enable) is configured to enable the GDMA interrupt. The DPI control register (dpi_control) is configured for the transfer. DMA channel instructions are programmed and executed via DMAC debug instruction and command registers. The test verifies completion by reading the subsystem interrupt mask register (interrupt_mask) to confirm the GDMA interrupt status, reads the DMAC masked interrupt status register to confirm DMA completion, then clears both the DMAC interrupt and the subsystem raw interrupt.",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_DSI_DMAC_INTEN to enable DMA controller interrupts. 2. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE with MIPI_DSI_SUBSYS_INTERRUPT_ENABLE_GDMA_INTR to enable GDMA interrupt at subsystem level. 3. Write to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG using set_data_mask with MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME to configure PHY stop wait time. 4. Write to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG to configure packet handling. 5. Write to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG to configure the clock manager. 6. Write to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to configure DPI control settings. 7. Write to MIZAR_MIPI_DSI_DMAC_DBGINST0 with DMA channel instruction byte 0 (including DSI_WRITE_MEMORY_START command, DMA source address DMA_SAR, and destination address DMA_DAR). 8. Write to MIZAR_MIPI_DSI_DMAC_DBGINST1 with DMA channel instruction byte 1. 9. Write to MIZAR_MIPI_DSI_DMAC_DBGCMD to execute the DMA debug instruction. 10. Read MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK and compare with MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR to verify GDMA interrupt mask status. 11. Read MIZAR_MIPI_DSI_DMAC_INTMIS to check DMA masked interrupt status for completion. 12. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear the DMA interrupt. 13. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW to clear the subsystem raw interrupt.",
        "Test Steps / Procedure": "1. Enable DMA controller interrupts by writing to the DMAC interrupt enable register. 2. Enable the GDMA interrupt at the subsystem level by writing to the interrupt_enable register. 3. Configure the DSI host PHY interface by writing the PHY stop wait time to the PHY_IF_CFG register. 4. Configure packet handling by writing to the PCKHDL_CFG register. 5. Configure the clock manager by writing to the CLKMGR_CFG register. 6. Configure DPI control settings by writing to the dpi_control register. 7. Program the DMA channel by writing source and destination addresses and the DBI write memory start command to the DMAC debug instruction registers. 8. Execute the DMA transfer by writing to the DMAC debug command register. 9. Read the interrupt_mask register and verify the GDMA interrupt mask status. 10. Read the DMAC masked interrupt status register to confirm DMA transfer completion. 11. Clear the DMA interrupt by writing to the DMAC interrupt clear register. 12. Clear the subsystem raw interrupt by writing to the interrupt_raw register.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_ENABLE; MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW",
        "Impacted Registers": "interrupt_enable; PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control; interrupt_mask; interrupt_raw",
        "Meta Validation / Acceptance Criteria": "1. After enabling GDMA interrupt, read MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_MASK and compare the GDMA interrupt bit (MIPI_DSI_SUBSYS_INTERRUPT_MASK_GDMA_INTR) to confirm the interrupt is properly masked/enabled. 2. After executing the DMA transfer via MIZAR_MIPI_DSI_DMAC_DBGCMD, read MIZAR_MIPI_DSI_DMAC_INTMIS to verify that the DMA transfer completed successfully by checking the masked interrupt status. 3. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear the DMA interrupt and verify it is acknowledged. 4. Write to MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW to clear the subsystem-level raw interrupt.",
        "Validation / Acceptance Criteria": "1. The interrupt_mask register must reflect the correct GDMA interrupt mask status after enabling the GDMA interrupt in the interrupt_enable register. 2. The DMAC masked interrupt status register must indicate successful DMA transfer completion after executing the DMA debug command. 3. The DMAC interrupt must be successfully cleared by writing to the DMAC interrupt clear register. 4. The subsystem raw interrupt must be successfully cleared by writing to the interrupt_raw register. 5. The test passes if the DBI write memory DMA transfer completes and all interrupt status checks confirm expected behavior.",
        "Remarks": "Six DMAC block registers (DMAC interrupt enable, DMAC debug instruction 0, DMAC debug instruction 1, DMAC debug command, DMAC masked interrupt status, DMAC interrupt clear) could not be mapped to canonical register names because no DMAC specification document was provided. The test involves interrupt-driven DMA completion verification with both DMAC-level and subsystem-level interrupt handling."
    },
    {
        "Index": "2",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_dbi_random_payload_test",
        "Feature": "DBI Random Payload Transfer via DMA",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a MIPI DSI DBI random payload transfer using DMA. It configures the DSI host PHY interface via MIZAR_MIPI_DSI_HOST_PHY_IF_CFG (including PHY stop wait time), configures packet handling via MIZAR_MIPI_DSI_HOST_PCKHDL_CFG, sets up the clock manager via MIZAR_MIPI_DSI_HOST_CLKMGR_CFG, and controls DPI via MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL. DMAC interrupts are enabled by writing to MIZAR_MIPI_DSI_DMAC_INTEN. The DMA channel is programmed using debug instruction registers MIZAR_MIPI_DSI_DMAC_DBGINST0 and MIZAR_MIPI_DSI_DMAC_DBGINST1, then executed via MIZAR_MIPI_DSI_DMAC_DBGCMD. The test polls MIZAR_MIPI_DSI_DMAC_INTMIS to check DMA interrupt completion, then clears the DMA interrupt via MIZAR_MIPI_DSI_DMAC_INTCLR. The payload data used for the DBI write memory command is random.",
        "Test Description": "This test validates a MIPI DSI DBI transfer with random payload data using DMA. It configures the DSI host PHY interface (PHY_IF_CFG), packet handling (PCKHDL_CFG), and clock manager (CLKMGR_CFG). The DPI control register (dpi_control) is configured for the transfer. DMA controller interrupts are enabled, and the DMA channel is programmed with debug instructions and executed via the DMAC debug command register. The test polls the DMAC masked interrupt status register to confirm DMA transfer completion, then clears the DMA interrupt. The test verifies that a DBI write memory operation with random payload data completes successfully through DMA.",
        "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_DSI_HOST_PHY_IF_CFG using set_data_mask with MIPI_DSI_HOST_PHY_IF_CFG_PHY_STOP_WAIT_TIME to configure PHY stop wait time. 2. Write to MIZAR_MIPI_DSI_HOST_PCKHDL_CFG to configure packet handling. 3. Write to MIZAR_MIPI_DSI_HOST_CLKMGR_CFG to configure the clock manager. 4. Write to MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL to configure DPI control settings. 5. Write to MIZAR_MIPI_DSI_DMAC_INTEN to enable DMA controller interrupts. 6. Write to MIZAR_MIPI_DSI_DMAC_DBGINST0 with DMA channel instruction byte 0 (including DSI_WRITE_MEMORY_START command, DMA source address DMA_SAR, and destination address DMA_DAR for random payload). 7. Write to MIZAR_MIPI_DSI_DMAC_DBGINST1 with DMA channel instruction byte 1. 8. Write to MIZAR_MIPI_DSI_DMAC_DBGCMD to execute the DMA debug instruction. 9. Poll MIZAR_MIPI_DSI_DMAC_INTMIS in a loop to wait for DMA masked interrupt status indicating transfer completion. 10. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear the DMA interrupt.",
        "Test Steps / Procedure": "1. Configure the DSI host PHY interface by writing the PHY stop wait time to the PHY_IF_CFG register. 2. Configure packet handling by writing to the PCKHDL_CFG register. 3. Configure the clock manager by writing to the CLKMGR_CFG register. 4. Configure DPI control settings by writing to the dpi_control register. 5. Enable DMA controller interrupts by writing to the DMAC interrupt enable register. 6. Program the DMA channel by writing source and destination addresses and the DBI write memory start command with random payload to the DMAC debug instruction registers. 7. Execute the DMA transfer by writing to the DMAC debug command register. 8. Poll the DMAC masked interrupt status register to wait for DMA transfer completion. 9. Clear the DMA interrupt by writing to the DMAC interrupt clear register.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_HOST_PHY_IF_CFG; MIZAR_MIPI_DSI_HOST_PCKHDL_CFG; MIZAR_MIPI_DSI_HOST_CLKMGR_CFG; MIZAR_MIPI_DSI_SUBSYS_DPI_CONTROL; MIZAR_MIPI_DSI_DMAC_INTEN; MIZAR_MIPI_DSI_DMAC_DBGINST0; MIZAR_MIPI_DSI_DMAC_DBGINST1; MIZAR_MIPI_DSI_DMAC_DBGCMD; MIZAR_MIPI_DSI_DMAC_INTMIS; MIZAR_MIPI_DSI_DMAC_INTCLR",
        "Impacted Registers": "PHY_IF_CFG; PCKHDL_CFG; CLKMGR_CFG; dpi_control",
        "Meta Validation / Acceptance Criteria": "1. After executing the DMA transfer via MIZAR_MIPI_DSI_DMAC_DBGCMD, poll MIZAR_MIPI_DSI_DMAC_INTMIS to verify that the DMA transfer completed successfully by checking the masked interrupt status. 2. Write to MIZAR_MIPI_DSI_DMAC_INTCLR to clear the DMA interrupt and verify it is acknowledged. 3. The random payload data must be successfully transferred via the DBI write memory command through DMA without errors.",
        "Validation / Acceptance Criteria": "1. The DMAC masked interrupt status register must indicate successful DMA transfer completion after executing the DMA debug command. 2. The DMA interrupt must be successfully cleared by writing to the DMAC interrupt clear register. 3. The test passes if the DBI write memory operation with random payload data completes successfully through DMA and the interrupt status confirms expected behavior.",
        "Remarks": "Six DMAC block registers (DMAC interrupt enable, DMAC debug instruction 0, DMAC debug instruction 1, DMAC debug command, DMAC masked interrupt status, DMAC interrupt clear) could not be mapped to canonical register names because no DMAC specification document was provided. The test uses random payload data for the DBI write memory transfer, distinguishing it from the basic test. The DMAC masked interrupt status register is polled in a loop to wait for transfer completion."
    },
    {
        "Index": "3",
        "SS / Module": "MIPI_DSI",
        "Test Case Name": "mipi_dsi_subsys_reg_wr_rd_test",
        "Feature": "Subsystem Register Write-Read Verification",
        "Meta Headers": "NA",
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs register write-read verification on MIPI DSI subsystem registers. It uses an array-driven approach where register addresses are stored in addr_array[] and iterated over. For each register, the test first reads the register via read_reg() to check the reset default value (chk_rst_val), then performs a write via write_reg() followed by a read-back via read_reg() to verify the written value (chk_rd_wr). The registers under test are: MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL, MIZAR_MIPI_DSI_SUBSYS_LOW_PWR, MIZAR_MIPI_DSI_SUBSYS_DBITE, MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV, and MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW. Each register undergoes read_modify_write operations with corresponding default value, read mask, and write mask validation.",
        "Test Description": "This test verifies the write-read accessibility of MIPI DSI subsystem registers. It iterates over a set of subsystem registers including data_fifo_threshold_val, low_pwr, dbite, dbi_fdiv, and interrupt_raw. For each register, the test first reads and verifies the reset default value, then writes a test value and reads it back to confirm the register is writable and retains the expected data. The test validates both reset-value correctness and write-read integrity for all targeted subsystem registers.",
        "Meta Test Steps / Procedure": "1. Initialize addr_array[] with register address macros: MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL, MIZAR_MIPI_DSI_SUBSYS_LOW_PWR, MIZAR_MIPI_DSI_SUBSYS_DBITE, MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV, MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW. 2. For each register in addr_array, call chk_rst_val(): read the register using read_reg(addr_array[i]) and compare the read value against the expected default value using the corresponding read mask. 3. For each register in addr_array, call chk_rd_wr(): write a test data value using write_reg(addr_array[i], data_wr) with the write mask applied, then read back using read_reg(addr_array[i]) and compare the read value against the expected written value using the read mask. 4. Track errors in err1 (reset value check errors) and err2 (write-read check errors). 5. Report pass or fail based on accumulated error counts.",
        "Test Steps / Procedure": "1. Initialize the register address array with the five target MIPI DSI subsystem register addresses. 2. For each register in the array, read the current value and verify it matches the expected reset default value. 3. For each register in the array, write a test value to the register applying the appropriate write mask. 4. Read back each register after writing and verify the read value matches the expected written value using the appropriate read mask. 5. Accumulate error counts for reset-value mismatches and write-read mismatches. 6. Report the test result as pass if no errors are detected, or fail if any mismatch is found.",
        "Meta Impacted Registers": "MIZAR_MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL; MIZAR_MIPI_DSI_SUBSYS_LOW_PWR; MIZAR_MIPI_DSI_SUBSYS_DBITE; MIZAR_MIPI_DSI_SUBSYS_DBI_FDIV; MIZAR_MIPI_DSI_SUBSYS_INTERRUPT_RAW",
        "Impacted Registers": "data_fifo_threshold_val; low_pwr; dbite; dbi_fdiv; interrupt_raw",
        "Meta Validation / Acceptance Criteria": "1. For each register in addr_array, read_reg(addr_array[i]) must return a value that, when masked with the corresponding read mask, matches the expected default value (e.g., MIPI_DSI_SUBSYS_DATA_FIFO_THRESHOLD_VAL_DEFAULT_VAL, MIPI_DSI_SUBSYS_LOW_PWR_DEFAULT_VAL, MIPI_DSI_SUBSYS_DBITE_DEFAULT_VAL, MIPI_DSI_SUBSYS_DBI_FDIV_DEFAULT_VAL, MIPI_DSI_SUBSYS_INTERRUPT_RAW_DEFAULT_VAL). 2. For each register, after write_reg(addr_array[i], data_wr), the subsequent read_reg(addr_array[i]) value masked with the read mask must match the written value masked with the write mask. 3. err1 must be 0 (all reset value checks pass). 4. err2 must be 0 (all write-read checks pass).",
        "Validation / Acceptance Criteria": "1. Each subsystem register (data_fifo_threshold_val, low_pwr, dbite, dbi_fdiv, interrupt_raw) must return its expected reset default value when read after power-on or reset. 2. After writing a test value to each register, the read-back value must match the written value within the applicable read and write mask constraints. 3. The reset-value error count must be zero, indicating all registers hold correct default values. 4. The write-read error count must be zero, indicating all registers correctly accept and retain written data. 5. The test passes only if both reset-value verification and write-read verification complete without any mismatches.",
        "Remarks": "All five registers belong to the MIPI DSI subsystem block (base 0xE6A41000) and were successfully mapped to canonical register names from the subsystem specification document. The test uses an array-driven iteration pattern with separate reset-value and write-read verification phases. Read masks and write masks are applied during comparisons to account for read-only or reserved bit fields. The actual MIPI DSI source code was not present in the repository folder; the testcase details are derived from the authoritative upstream agent outputs."
    }
]

def generate_excel():
    IST = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(IST)
    timestamp = now_ist.strftime('%Y%m%d_%H%M%S')
    filename = f'MIPI_DSI_TestPlan_{timestamp}.xlsx'
    
    wb = openpyxl.Workbook()
    
    # TestPlan sheet
    ws_tp = wb.active
    ws_tp.title = 'TestPlan'
    
    tp_columns = [
        'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
        'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
        'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
        'Code Generation'
    ]
    
    # MetaData sheet
    ws_md = wb.create_sheet('MetaData')
    
    md_columns = [
        'Index', 'Test Case Name', 'Meta Test Description',
        'Meta Test Steps / Procedure', 'Meta Impacted Registers',
        'Meta Validation / Acceptance Criteria', 'Meta Headers',
        'Meta Macros', 'Meta Arrays'
    ]
    
    # Header formatting
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    wrap_alignment = Alignment(wrap_text=True, vertical='top')
    
    # Write TestPlan headers
    for col_idx, col_name in enumerate(tp_columns, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    
    # Write MetaData headers
    for col_idx, col_name in enumerate(md_columns, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    
    # Populate TestPlan rows
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(tp_columns, 1):
            value = row_data.get(col_name, '')
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
    
    # Populate MetaData rows
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(md_columns, 1):
            value = row_data.get(col_name, '')
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
    
    # Auto-size columns with max width cap
    for ws in [ws_tp, ws_md]:
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 2, 60)
            ws.column_dimensions[col_letter].width = max(adjusted_width, 12)
    
    # Freeze first row
    ws_tp.freeze_panes = 'A2'
    ws_md.freeze_panes = 'A2'
    
    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = 'veryHidden'
    
    # Save
    output_dir = os.path.dirname(os.path.abspath(__file__))
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    
    # Validate
    assert os.path.exists(filepath), 'File does not exist'
    assert os.path.getsize(filepath) > 0, 'File is empty'
    wb2 = openpyxl.load_workbook(filepath)
    assert 'TestPlan' in wb2.sheetnames, 'TestPlan sheet missing'
    assert 'MetaData' in wb2.sheetnames, 'MetaData sheet missing'
    wb2.close()
    
    print(f'SUCCESS: {filename}')
    print(f'PATH: {filepath}')
    print(f'SIZE: {os.path.getsize(filepath)}')
    return filename, filepath

if __name__ == '__main__':
    generate_excel()
