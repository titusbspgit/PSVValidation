#!/usr/bin/env python3
"""
Standalone XLSX generator for MIPI_CSI TestPlan - Agent 7
Run: python3 _temp_gen_xlsx.py
Output: MIPI_CSI_TestPlan_YYYYMMDD_HHMMSS.xlsx in same directory
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import os

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
ts = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"MIPI_CSI_TestPlan_{ts}.xlsx"
filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)

json_data = [
  {
    "Index": "1",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_dphy_lanes_test",
    "Feature": "D-PHY Lane Configuration",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_csi2.h\"",
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000; LS_LE_EN = 1; TOTAL_FRAME = 1; VC_ID = 3; VRES; HRES; DATA_TYPE",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates MIPI CSI-2 D-PHY lane configuration by iterating through lane counts from 4 lanes down to 1 lane. It first calls csi2_enable_interrupt() which reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to clear pending interrupts, then writes all CSI-2 host interrupt mask registers (INT_MSK_PHY_FATAL with 0x0000000f, INT_MSK_PKT_FATAL with 0x00000003, INT_MSK_PHY with 0x000f000f, INT_MSK_LINE with 0x000f000f, INT_MSK_BNDRY_FRAME_FATAL with 0x0000ffff, INT_MSK_SEQ_FRAME_FATAL with 0x0000ffff, INT_MSK_CRC_FRAME_FATAL with 0x0000ffff, INT_MSK_PLD_CRC_FATAL with 0x0000ffff, INT_MSK_DATA_ID with 0x0000ffff, INT_MSK_ECC_CORRECTED with 0x0000ffff). The test then configures the virtual channel register MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL with the computed vcid_csi2_wrap_reg value based on the selected GDMA path, and writes 1 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to enable control data transfer. These two writes are performed twice. snps_phy_init() is called for D-PHY initialization. The test then polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE until it reads 0x1000f, confirming PHY stop-state entry. A for loop iterates lane_num from 3 down to 0: each iteration writes lane_num to MIZAR_MIPI_CSI2_HOST_N_LANES, writes (lane_num+1) to address 0xa0243ffc to trigger the CSI-2 sequence, then enters an inner loop for cntrl_pkt_cnt = (VRES*3)+2 iterations performing DMA transfers. After all iterations, finish(0) is called.",
    "Test Description": "This test validates MIPI CSI-2 D-PHY lane configuration by iterating through lane counts from 4 lanes down to 1 lane. It enables all CSI-2 host interrupts by reading INT_ST_MAIN to clear pending status and writing enable masks to all interrupt mask registers (INT_MSK_PHY_FATAL, INT_MSK_PKT_FATAL, INT_MSK_PHY, INT_MSK_LINE, INT_MSK_BNDRY_FRAME_FATAL, INT_MSK_SEQ_FRAME_FATAL, INT_MSK_CRC_FRAME_FATAL, INT_MSK_PLD_CRC_FATAL, INT_MSK_DATA_ID, INT_MSK_ECC_CORRECTED). It configures the virtual_channel register with the appropriate virtual channel ID and enables control data transfer via the control_data register. After D-PHY initialization, it polls PHY_STOPSTATE to confirm the PHY has entered stop-state. For each lane configuration (4 to 1), it writes the lane count to N_LANES, triggers the CSI-2 sequence, and performs DMA-based control and data packet transfers. Each packet transfer uses DMA channel 0 for control data, polling for DMA completion via interrupt status. If the received control data indicates a long packet, a data transfer is initiated on DMA channel 1, again polling for completion. The test verifies successful DMA transfers across all lane configurations.",
    "Meta Test Steps / Procedure": "1. Entry: test_case() is called. 2. Call csi2_enable_interrupt(). 3. read_reg(MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN) to clear pending interrupts. 4-13. Write all interrupt mask registers. 14. Return from csi2_enable_interrupt(). 15-16. Configure GDMA path. 17-20. Write virtual_channel and control_data registers twice. 21. Call snps_phy_init(). 22-23. Poll PHY_STOPSTATE until 0x1000f. 24-51. Lane loop with DMA transfers. 51. Call finish(0).",
    "Test Steps / Procedure": "1. Enable CSI-2 host interrupts by reading INT_ST_MAIN to clear pending status, then writing enable masks to all interrupt mask registers. 2. Configure the virtual_channel register. 3. Enable control data transfer. 4. Perform D-PHY initialization. 5. Poll PHY_STOPSTATE register. 6. For each lane configuration: write N_LANES and trigger CSI-2 sequence. 7-11. DMA transfer operations with polling. 12-14. Repeat for all packets and lane configurations. Verify test completion.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000",
    "Impacted Registers": "INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED; virtual_channel; control_data; PHY_STOPSTATE; N_LANES",
    "Meta Validation / Acceptance Criteria": "1. PHY_STOPSTATE polling: must return 0x1000f. 2. DMA channel 0 completion: polled until bit 0 set. 3. DMA channel 1 completion (conditional): polled until bit 1 set. 4. Long packet detection: (csi_ctrl_data & 0x3f) > 0xf. 5. Test pass: finish(0) called.",
    "Validation / Acceptance Criteria": "1. PHY_STOPSTATE register must read the expected stop-state value. 2. DMA channel 0 interrupt status must indicate completion. 3. If long packet, DMA channel 1 must complete. 4. All four lane configurations complete without failure. 5. Test completes with pass indication.",
    "Remarks": "The virtual_channel and control_data registers are written twice in sequence. DMA register accesses use a variable base with offset macros. Two direct hex addresses (0xa0243ffc and 0xE6001000) could not be mapped to named registers."
  },
  {
    "Index": "2",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_test_pattern_generator",
    "Feature": "Test Pattern Generator",
    "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"mipi_csi2.h\"",
    "Meta Macros": "GDMA_CSI2_DATA_DEST_ADDR2 = 0xE6040080; GDMA_CTRL_DATA_DEST_ADDR2 = 0xE6040000",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates the MIPI CSI-2 internal test pattern generator functionality. It configures the CSI-2 subsystem to receive data generated by the host controller's built-in pattern generator and transfers the received data via DMA. The test writes vcid_csi2_wrap_reg to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL and writes 0 to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA. csi2_enable_interrupt() is called. snps_phy_init() is called. PHY_STOPSTATE is polled until 0x1000f. DMA address registers are programmed. 0x1 is written to MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4. DMA channel 0 is started. Pattern generator is enabled then disabled. DMA completion is polled. finish(0) is called.",
    "Test Description": "This test validates the MIPI CSI-2 internal test pattern generator by configuring the pattern generator to produce a 320x16 pixel RGB888 frame and transferring the received data via DMA. It configures the virtual_channel register and disables control data transfer. All CSI-2 host interrupts are enabled. After D-PHY initialization, it polls PHY_STOPSTATE. DMA address mapping registers are programmed. The enableclkgating_csiphy register is written. DMA channel 0 is configured and started. The pattern generator is enabled then disabled. The test polls DMA interrupt status for completion and finishes with a pass indication.",
    "Meta Test Steps / Procedure": "1. Entry: test_case() is called. 2-6. Set variables and compute GDMA path. 7. write_reg(VIRTUAL_CHANNEL). 8. write_reg(CONTROL_DATA, 0). 9-21. Enable interrupts. 22. Call snps_phy_init(). 23-24. Poll PHY_STOPSTATE. 25-30. Program DMA address registers. 31. write_reg(RB_REG_BASE + 0xf4, 0x1). 32-36. Start DMA. 37-42. Enable pattern generator. 43-44. Wait and disable PG. 45-46. Poll DMA completion. 47-48. Wait and finish(0).",
    "Test Steps / Procedure": "1. Configure virtual_channel and disable control data transfer. 2. Enable CSI-2 host interrupts. 3. Perform D-PHY initialization. 4. Poll PHY_STOPSTATE register. 5. Compute frame data transfer size. 6. Program DMA address mapping registers. 7. Write enableclkgating_csiphy register. 8. Start DMA channel 0 transfer. 9. Enable pattern generator. 10. Wait, then disable pattern generator. 11. Poll DMA interrupt status for completion. 12. Verify test completion.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; enableclkgating_csiphy; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "1. PHY_STOPSTATE polling: must return 0x1000f. 2. DMA channel 0 completion: polled until bit 0 set. 3. Pattern generator enable/disable sequence with wait. 4. Test pass: finish(0) called.",
    "Validation / Acceptance Criteria": "1. PHY_STOPSTATE register must read expected stop-state value. 2. DMA channel 0 interrupt status must indicate completion. 3. Pattern generator must be enabled and disabled in sequence. 4. Test completes with pass indication.",
    "Remarks": "The test uses the internal test pattern generator (PPI_PG) to generate a 320x16 pixel RGB888 frame. The MIZAR_MIPI_CSI2_RB_REG_BASE + 0xf4 composite access is mapped to the enableclkgating_csiphy register. DMA register accesses use a variable base with offset macro."
  }
]

tp_cols = ["Index","SS / Module","Feature","Test Case Name","Test Description",
           "Speed","Mode","Memory Start Offset","Memory End Offset","Remarks",
           "Test Steps / Procedure","Impacted Registers",
           "Validation / Acceptance Criteria","Code Generation"]

md_cols = ["Index","Test Case Name","Meta Test Description",
           "Meta Test Steps / Procedure","Meta Impacted Registers",
           "Meta Validation / Acceptance Criteria","Meta Headers",
           "Meta Macros","Meta Arrays"]

wb = openpyxl.Workbook()
ws_tp = wb.active
ws_tp.title = "TestPlan"
ws_md = wb.create_sheet("MetaData")

hf = Font(bold=True, color="FFFFFF", size=11)
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wa = Alignment(wrap_text=True, vertical="top")

def write_sheet(ws, columns, data):
    for ci, cn in enumerate(columns, 1):
        c = ws.cell(row=1, column=ci, value=cn)
        c.font = hf; c.fill = hfill; c.alignment = wa
    for ri, rd in enumerate(data, 2):
        for ci, cn in enumerate(columns, 1):
            c = ws.cell(row=ri, column=ci, value=rd.get(cn, ""))
            c.alignment = wa
    ws.freeze_panes = "A2"
    for ci, cn in enumerate(columns, 1):
        ml = len(cn)
        for ri in range(2, len(data)+2):
            v = ws.cell(row=ri, column=ci).value
            if v:
                for ln in str(v).split('\n'):
                    ml = max(ml, len(ln))
        ws.column_dimensions[get_column_letter(ci)].width = min(ml+4, 80)

write_sheet(ws_tp, tp_cols, json_data)
write_sheet(ws_md, md_cols, json_data)
ws_md.sheet_state = "veryHidden"

wb.save(filepath)
wb.close()

# Verify
sz = os.path.getsize(filepath)
wb2 = openpyxl.load_workbook(filepath)
print(f"OK: {filename} ({sz} bytes) sheets={wb2.sheetnames}")
wb2.close()
