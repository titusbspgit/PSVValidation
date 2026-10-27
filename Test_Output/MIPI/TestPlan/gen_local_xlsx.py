#!/usr/bin/env python3
import openpyxl, os, base64
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
hf = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
hfill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wa = Alignment(wrap_text=True, vertical='top')

ws = wb.active; ws.title = "TestPlan"
for ci, cn in enumerate(["Index","SS / Module","Feature","Test Case Name","Speed","Mode","Memory Start Offset","Memory End Offset","Test Description","Test Steps / Procedure","Impacted Registers","Validation / Acceptance Criteria","Remarks","Code Generation (Required / Not)"], 1):
    c = ws.cell(row=1, column=ci, value=cn); c.font = hf; c.fill = hfill; c.alignment = wa

ws.cell(row=2, column=1, value=1).alignment = wa
ws.cell(row=2, column=2, value="MIPI_CSI").alignment = wa
ws.cell(row=2, column=3, value="Register Read/Write Validation (RB Register Block)").alignment = wa
ws.cell(row=2, column=4, value="mipi_csi2_rb_reg_wr_rd_test").alignment = wa
ws.cell(row=2, column=5, value="NA").alignment = wa
ws.cell(row=2, column=6, value="NA").alignment = wa
ws.cell(row=2, column=7, value="NA").alignment = wa
ws.cell(row=2, column=8, value="NA").alignment = wa
ws.cell(row=2, column=9, value="This test performs register-level read/write validation on 60 MIPI CSI2 Receive Block (RB) registers.").alignment = wa
ws.cell(row=2, column=10, value="Step 1: Default value check. Step 2: Write/Read validation. Step 3: Evaluate result.").alignment = wa
ws.cell(row=2, column=11, value="virtual_channel, line_info, fifo_threshold_val, lane_clk, mem, interrupt_raw, interrupt_mask, interrupt_enable, control_data, flush, fifo_interrupt_raw, fifo_interrupt_enable, fifo_interrupt_mask, soft_rst, null_blank, dmac_boot, boot_addr_VC0-VC3, boot_irq_ns_VC0-VC3, dphy_rx_cfgclk, dphy_rx_hs, dphy_IO_continuity, dma_m0-m3_addr_ar/aw_ch0/ch1_data/Instruction, dphy_rx").alignment = wa
ws.cell(row=2, column=12, value="Pass: def_fail_cnt==0 AND wr_fail_cnt==0. Fail: either >0.").alignment = wa
ws.cell(row=2, column=13, value="Soft reset check commented out. 6 test patterns. 60 registers covered.").alignment = wa

dv = DataValidation(type="list", formula1='"Required,Not Required"', allow_blank=True)
dv.add("N2:N2"); ws.add_data_validation(dv)
for i, w in enumerate([8,15,35,35,8,8,20,20,60,60,60,60,60,25], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

wm = wb.create_sheet("MetaData")
for ci, cn in enumerate(["Test Case Name","Meta Headers","Meta Macros","Meta Arrays","Meta Types","Meta Test Description","Meta Test Steps / Procedure","Meta Impacted Registers","Meta Validation / Acceptance Criteria"], 1):
    c = wm.cell(row=1, column=ci, value=cn); c.font = hf; c.fill = hfill; c.alignment = wa
wm.cell(row=2, column=1, value="mipi_csi2_rb_reg_wr_rd_test").alignment = wa
wm.cell(row=2, column=2, value="program.c: stdio.h, stdlib.h, test_define.c, test_common.h; test_define.c: mipi_csi2_rb_reg.h").alignment = wa
wm.cell(row=2, column=3, value="249 macros total").alignment = wa
wm.cell(row=2, column=4, value="addr_array[60], read_mask_array[60], write_mask_array[60], default_value_array[60], skip_array[60], chk_val[6]").alignment = wa
wm.cell(row=2, column=5, value="int, unsigned long int, const int").alignment = wa
wm.cell(row=2, column=6, value="Register-level read/write validation on 60 MIPI CSI2 RB registers.").alignment = wa
wm.cell(row=2, column=7, value="1. chk_rst_val() 2. chk_rd_wr() 3. test_case()").alignment = wa
wm.cell(row=2, column=8, value="virtual_channel, line_info, fifo_threshold_val, lane_clk, mem, interrupt_raw, interrupt_mask, interrupt_enable, control_data, flush, fifo_interrupt_raw, fifo_interrupt_enable, fifo_interrupt_mask, soft_rst, null_blank, dmac_boot, boot_addr_VC0-VC3, boot_irq_ns_VC0-VC3, dphy_rx_cfgclk, dphy_rx_hs, dphy_IO_continuity, dma registers, dphy_rx").alignment = wa
wm.cell(row=2, column=9, value="Pass: def_fail_cnt==0 AND wr_fail_cnt==0 -> finish(0). Fail: finish(1).").alignment = wa
for i, w in enumerate([35,50,60,60,30,60,60,60,60], 1):
    wm.column_dimensions[get_column_letter(i)].width = w
wm.freeze_panes = "A2"; wm.sheet_state = 'veryHidden'

wb.save("/tmp/out.xlsx")
print(f"SIZE:{os.path.getsize('/tmp/out.xlsx')}")
