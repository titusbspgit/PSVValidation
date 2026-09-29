#!/usr/bin/env python3
"""Standalone XLSX generator for MIPI_CSI TestPlan."""
import os, sys, json, base64
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime('%Y%m%d_%H%M%S')
filename = f'MIPI_CSI_TestPlan_{timestamp}.xlsx'

json_data = [
  {
    "Index": "1",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_dphy_lanes_test",
    "Feature": "DPHY Lane Configuration and Interrupt Setup",
    "Meta Headers": "NA",
    "Meta Macros": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED; MIPI_CSI2_DMA_INTEN_OFFSET; MIPI_CSI2_DMA_INTMIS_OFFSET; MIPI_CSI2_DMA_INTCLR_OFFSET",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates MIPI CSI2 DPHY lane configuration and interrupt handling. In the test_case() function, it first writes to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL to configure the virtual channel, then writes to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA to set control data. It polls MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE in a while-loop to wait for the PHY data lanes to reach stop state. It writes to MIZAR_MIPI_CSI2_HOST_N_LANES to configure the number of active DPHY lanes. A write to hardcoded address 0xa0243ffc is performed for auxiliary configuration. DMA-related registers are accessed via base+offset expressions (gdma_reg_base+MIPI_CSI2_DMA_INTEN_OFFSET for interrupt enable, gdma_reg_base+MIPI_CSI2_DMA_INTMIS_OFFSET for masked interrupt status read, gdma_reg_base+MIPI_CSI2_DMA_INTCLR_OFFSET for interrupt clear, and gdma_reg_base+0x28 for DMA status). A read from hardcoded address 0xE6001000 is performed. The csi2_enable_interrupt() function reads MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN to check main interrupt status, then writes to all CSI2 host interrupt mask registers to enable all interrupt sources.",
    "Test Description": "This testcase validates MIPI CSI2 DPHY lane configuration and comprehensive interrupt setup. It configures the virtual channel and control data registers in the CSI2 register block, then polls the PHY_STOPSTATE register to confirm that the DPHY data lanes have entered stop state. The number of active DPHY lanes is configured via the N_LANES register. DMA interrupt enable, status, and clear operations are performed. The test enables all CSI2 host interrupt masks by writing to the interrupt mask registers for PHY fatal, packet fatal, PHY errors, line errors, boundary frame fatal, sequence frame fatal, CRC frame fatal, payload CRC fatal, data ID, and ECC corrected interrupt categories. The main interrupt status register is read to verify the interrupt state.",
    "Meta Test Steps / Procedure": "1. Write to MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL (base 0xE6A04000, offset 0x0) to configure virtual channel. 2. Write to MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA (base 0xE6A04000, offset 0x20) to set control data parameters. 3. Poll MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE (base 0xE6A05000, offset 0x4C) in a while-loop waiting for PHY stop state indication on data lanes. 4. Write to MIZAR_MIPI_CSI2_HOST_N_LANES (base 0xE6A05000, offset 0x4) to configure the number of active DPHY lanes. 5-21. Configure all interrupt mask registers.",
    "Test Steps / Procedure": "1. Configure the virtual channel register in the CSI2 register block to set the desired virtual channel. 2. Configure the control data register in the CSI2 register block with the required control parameters. 3. Poll the PHY_STOPSTATE register until the DPHY data lanes indicate they have entered stop state. 4. Configure the N_LANES register to set the number of active DPHY lanes for the test. 5. Perform an auxiliary configuration write to an external peripheral register. 6. Enable DMA interrupts by writing to the DMA interrupt enable register. 7. Read the DMA masked interrupt status register to verify DMA interrupt state. 8. Clear any pending DMA interrupts by writing to the DMA interrupt clear register. 9. Read the DMA status register to confirm DMA channel state. 10. Read an external system register for status verification. 11. Read the INT_ST_MAIN register to check the current main interrupt status of the CSI2 host. 12. Enable all CSI2 host interrupt masks. 13. Verify that all interrupt mask registers are configured correctly and the DPHY lane configuration is complete.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_HOST_N_LANES; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PKT_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PHY; MIZAR_MIPI_CSI2_HOST_INT_MSK_LINE; MIZAR_MIPI_CSI2_HOST_INT_MSK_BNDRY_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_SEQ_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_CRC_FRAME_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_PLD_CRC_FATAL; MIZAR_MIPI_CSI2_HOST_INT_MSK_DATA_ID; MIZAR_MIPI_CSI2_HOST_INT_MSK_ECC_CORRECTED",
    "Impacted Registers": "virtual_channel; control_data; PHY_STOPSTATE; N_LANES; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "The test validates that all register writes complete successfully and all 10 interrupt mask registers are written successfully to enable the corresponding interrupt sources.",
    "Validation / Acceptance Criteria": "1. The virtual_channel register is written successfully. 2. The control_data register is written successfully. 3. The PHY_STOPSTATE register polling completes. 4. The N_LANES register is configured correctly. 5. DMA operations complete without errors. 6. The INT_ST_MAIN register read returns expected status. 7. All 10 CSI2 host interrupt mask registers are configured correctly. 8. The test completes without timeout.",
    "Remarks": "Two hardcoded register addresses could not be mapped to canonical register names. The PHY_STOPSTATE register is read-only and is polled in a while-loop. DMA registers are accessed via base+offset expressions."
  },
  {
    "Index": "2",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_rb_reg_wr_rd_test",
    "Feature": "Register Block Read/Write Verification",
    "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>",
    "Meta Macros": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_LINE_INFO; MIZAR_MIPI_CSI2_RB_REG_FIFO_THRESHOLD_VAL; MIZAR_MIPI_CSI2_RB_REG_LANE_CLK; MIZAR_MIPI_CSI2_RB_REG_MEM; MIZAR_MIPI_CSI2_RB_REG_INTERRUPT_RAW; MIZAR_MIPI_CSI2_RB_REG_INTERRUPT_MASK; MIZAR_MIPI_CSI2_RB_REG_INTERRUPT_ENABLE; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_RB_REG_FLUSH",
    "Meta Arrays": "addr_array[]; default_val_array[]; read_mask_array[]; write_mask_array[]; skip_array[]",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase performs register reset-value verification and write/read-back validation for all 10 registers in the MIPI CSI2 RB REG block (base 0xE6A04000).",
    "Test Description": "This testcase validates the reset default values and write/read-back behavior of all registers in the MIPI CSI2 register block. It first reads each register and compares the masked value against the expected reset default to verify correct power-on state. Then, for writable registers, it writes a test pattern, reads back the value, and compares the masked result against the written data to confirm write/read-back integrity. Two registers (interrupt_raw and flush) are read-only and are skipped during the write/read-back phase. All 10 registers in the block are covered.",
    "Meta Test Steps / Procedure": "1. Define addr_array with 10 register address macros. 2-5. Define arrays for default values, read masks, write masks, and skip flags. 6. Call chk_rst_val() to verify reset values. 7. Call chk_rd_wr() to verify write/read-back. 8. Check error counters.",
    "Test Steps / Procedure": "1. Initialize the test with arrays containing register addresses, expected reset default values, read masks, write masks, and skip flags for all 10 registers. 2. Perform reset value verification for all registers. 3. Perform write/read-back verification for writable registers. 4. Skip write/read-back testing for read-only registers. 5. Evaluate the error counters for pass/fail.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_LINE_INFO; MIZAR_MIPI_CSI2_RB_REG_FIFO_THRESHOLD_VAL; MIZAR_MIPI_CSI2_RB_REG_LANE_CLK; MIZAR_MIPI_CSI2_RB_REG_MEM; MIZAR_MIPI_CSI2_RB_REG_INTERRUPT_RAW; MIZAR_MIPI_CSI2_RB_REG_INTERRUPT_MASK; MIZAR_MIPI_CSI2_RB_REG_INTERRUPT_ENABLE; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_RB_REG_FLUSH",
    "Impacted Registers": "virtual_channel; line_info; fifo_threshold_val; lane_clk; mem; interrupt_raw; interrupt_mask; interrupt_enable; control_data; flush",
    "Meta Validation / Acceptance Criteria": "The test passes if both err1 == 0 and err2 == 0 after all iterations. Registers at skip_array indices 5 and 9 are excluded from write/read-back validation.",
    "Validation / Acceptance Criteria": "1. All 10 registers must return expected reset default values. 2. For 8 writable registers, write/read-back must match. 3. interrupt_raw and flush must be correctly skipped. 4. Zero errors in both phases.",
    "Remarks": "The test uses parallel arrays to iterate over all 10 registers. Two registers are marked as read-only via skip_array. The soft_reset_chk() function is excluded per instructions."
  },
  {
    "Index": "3",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_csi2_test_pattern_generator",
    "Feature": "Test Pattern Generation and DMA Configuration",
    "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>",
    "Meta Macros": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; and 10 interrupt mask macros; MIPI_CSI2_DMA_INTMIS_OFFSET",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates the MIPI CSI2 internal test pattern generator functionality along with DMA channel configuration and comprehensive interrupt setup.",
    "Test Description": "This testcase validates the MIPI CSI2 internal test pattern generator along with DMA channel configuration and interrupt setup. It configures the pattern generator vertical and horizontal resolution, sets the pattern configuration, and enables the generator. The virtual_channel and control_data registers are configured. The PHY_STOPSTATE register is polled. DMA channel 0 read and write address registers are configured. The enableclkgating_csiphy register is written. After test pattern transmission, the pattern generator is disabled. DMA interrupt status is checked, and all CSI2 host interrupt mask registers are enabled.",
    "Meta Test Steps / Procedure": "1-4. Configure pattern generator registers. 5-7. Configure virtual channel, control data, poll PHY. 8-12. Configure DMA and clock gating. 13-25. Disable PG, check DMA status, enable all interrupt masks.",
    "Test Steps / Procedure": "1. Configure the test pattern generator by setting vertical and horizontal resolution. 2. Set the pattern configuration. 3. Enable the test pattern generator. 4. Configure virtual_channel and control_data. 5. Poll PHY_STOPSTATE. 6-8. Configure DMA channel 0 registers. 9. Enable CSI PHY clock gating. 10. Disable the test pattern generator. 11-14. Check DMA status and enable all interrupt masks.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_VRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_PATTERN_HRES; MIZAR_MIPI_CSI2_HOST_PPI_PG_CONFIG; MIZAR_MIPI_CSI2_HOST_PPI_PG_ENABLE; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AR_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_DATA; MIZAR_MIPI_CSI2_RB_REG_DMA_M0_ADDR_AW_CH0_INSTRUCTION; MIZAR_MIPI_CSI2_RB_REG_BASE; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; and 10 interrupt mask macros",
    "Impacted Registers": "PPI_PG_PATTERN_VRES; PPI_PG_PATTERN_HRES; PPI_PG_CONFIG; PPI_PG_ENABLE; virtual_channel; control_data; PHY_STOPSTATE; dma_m0_addr_ar_ch0_data; dma_m0_addr_ar_ch0_Instruction; dma_m0_addr_aw_ch0_data; dma_m0_addr_aw_ch0_Instruction; enableclkgating_csiphy; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "All register writes complete successfully. Pattern generator enable/disable cycle works. All 10 interrupt mask registers are written successfully.",
    "Validation / Acceptance Criteria": "1. PPI_PG registers are configured correctly. 2. PPI_PG_ENABLE toggles correctly. 3. virtual_channel and control_data are configured. 4. PHY_STOPSTATE polling completes. 5. DMA channel 0 registers are configured. 6. enableclkgating_csiphy is written. 7. DMA status read completes. 8. All interrupt mask registers are configured. 9. No timeout during PHY polling.",
    "Remarks": "The test pattern generator is enabled and then disabled within the same test flow. The PHY_STOPSTATE register is polled in a while-loop. The test exercises three functional areas: test pattern generation, DMA channel configuration, and interrupt mask setup."
  },
  {
    "Index": "4",
    "SS / Module": "MIPI_CSI",
    "Test Case Name": "mipi_dphy_idi_test",
    "Feature": "DPHY IDI Data Transfer with Virtual Channel Iteration",
    "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>",
    "Meta Macros": "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_RB_REG_NULL_BLANK; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; 10 interrupt mask macros; MIPI_CSI2_DMA_INTEN_OFFSET; MIPI_CSI2_DMA_INTMIS_OFFSET; MIPI_CSI2_DMA_INTCLR_OFFSET; GDMA_CSI2_DATA_DEST_ADDR2; GDMA_CTRL_DATA_DEST_ADDR2",
    "Meta Arrays": "NA",
    "Speed": "NA",
    "Mode": "NA",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "This testcase validates MIPI CSI2 DPHY IDI data transfer with virtual channel iteration, null/blank packet configuration, DMA interrupt handling, and comprehensive CSI2 host interrupt setup.",
    "Test Description": "This testcase validates MIPI CSI2 DPHY IDI data transfer with virtual channel iteration and comprehensive interrupt configuration. It polls the PHY_STOPSTATE register to confirm DPHY data lanes have entered stop state. The control_data register is configured to enable data control, and the null_blank register is configured to enable null/blank packet insertion. The virtual_channel register is written with the virtual channel ID for each iteration. DMA interrupt operations are performed. All CSI2 host interrupt masks are enabled.",
    "Meta Test Steps / Procedure": "1. Poll PHY_STOPSTATE. 2-4. Configure control_data, null_blank, virtual_channel. 5. Write to external register. 6-10. DMA operations. 11-21. Enable all interrupt masks.",
    "Test Steps / Procedure": "1. Poll the PHY_STOPSTATE register. 2. Configure control_data. 3. Configure null_blank. 4. Write virtual channel ID. 5. Write to external peripheral register. 6-10. DMA interrupt operations. 11. Read INT_ST_MAIN. 12. Enable all CSI2 host interrupt masks. 13. Verify all configurations complete without errors.",
    "Meta Impacted Registers": "MIZAR_MIPI_CSI2_HOST_PHY_STOPSTATE; MIZAR_MIPI_CSI2_RB_REG_CONTROL_DATA; MIZAR_MIPI_CSI2_RB_REG_NULL_BLANK; MIZAR_MIPI_CSI2_RB_REG_VIRTUAL_CHANNEL; 0xa0243ffc; 0xE6001000; MIZAR_MIPI_CSI2_HOST_INT_ST_MAIN; and 10 interrupt mask macros",
    "Impacted Registers": "PHY_STOPSTATE; control_data; null_blank; virtual_channel; INT_ST_MAIN; INT_MSK_PHY_FATAL; INT_MSK_PKT_FATAL; INT_MSK_PHY; INT_MSK_LINE; INT_MSK_BNDRY_FRAME_FATAL; INT_MSK_SEQ_FRAME_FATAL; INT_MSK_CRC_FRAME_FATAL; INT_MSK_PLD_CRC_FATAL; INT_MSK_DATA_ID; INT_MSK_ECC_CORRECTED",
    "Meta Validation / Acceptance Criteria": "PHY_STOPSTATE polling exits successfully. All register writes complete. All 10 interrupt mask registers are written successfully.",
    "Validation / Acceptance Criteria": "1. PHY_STOPSTATE polling completes. 2. control_data is written successfully. 3. null_blank is written successfully. 4. virtual_channel is configured correctly. 5. External register writes complete. 6. DMA operations complete. 7. INT_ST_MAIN returns expected status. 8. All 10 interrupt mask registers are configured. 9. No timeout during polling.",
    "Remarks": "Two hardcoded register addresses could not be mapped. The null_blank register is unique to this testcase. The test iterates over virtual channel IDs. Two macros (GDMA_CSI2_DATA_DEST_ADDR2 and GDMA_CTRL_DATA_DEST_ADDR2) are defined but not used in register access calls."
  }
]

testplan_columns = [
    'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
    'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
    'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
    'Code Generation'
]

metadata_columns = [
    'Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
    'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
    'Meta Headers', 'Meta Macros', 'Meta Arrays'
]

wb = Workbook()
ws_tp = wb.active
ws_tp.title = 'TestPlan'

header_font = Font(bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap_alignment = Alignment(wrap_text=True, vertical='top')

for col_idx, col_name in enumerate(testplan_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(testplan_columns, 1):
        value = row_data.get(col_name, '')
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

ws_tp.freeze_panes = 'A2'

for col_idx, col_name in enumerate(testplan_columns, 1):
    max_len = len(col_name)
    for row_idx in range(2, len(json_data) + 2):
        val = ws_tp.cell(row=row_idx, column=col_idx).value
        if val:
            max_len = max(max_len, min(len(str(val)), 80))
    ws_tp.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

ws_md = wb.create_sheet('MetaData')

for col_idx, col_name in enumerate(metadata_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(metadata_columns, 1):
        value = row_data.get(col_name, '')
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

ws_md.freeze_panes = 'A2'

for col_idx, col_name in enumerate(metadata_columns, 1):
    max_len = len(col_name)
    for row_idx in range(2, len(json_data) + 2):
        val = ws_md.cell(row=row_idx, column=col_idx).value
        if val:
            max_len = max(max_len, min(len(str(val)), 80))
    ws_md.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

ws_md.sheet_state = 'veryHidden'

wb.save(filename)

# Validate
wb2 = load_workbook(filename)
assert 'TestPlan' in wb2.sheetnames
assert 'MetaData' in wb2.sheetnames
file_size = os.path.getsize(filename)
assert file_size > 0

print(f'FILENAME={filename}')
print(f'FILE_SIZE={file_size}')
print(f'VALIDATION=PASSED')
print(f'ROWS_TESTPLAN={len(json_data)}')
print(f'ROWS_METADATA={len(json_data)}')

# Output base64 for GitHub upload
import base64
with open(filename, 'rb') as f:
    b64 = base64.b64encode(f.read()).decode()
print(f'BASE64_START')
print(b64)
print(f'BASE64_END')
