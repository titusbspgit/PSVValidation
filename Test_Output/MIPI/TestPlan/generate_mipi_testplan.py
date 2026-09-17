#!/usr/bin/env python3
"""
MIPI TestPlan Excel Generator
Generates MIPI_TestPlan_YYYYMMDD_HHMMSS.xlsx with openpyxl
Usage: pip install openpyxl && python3 generate_mipi_testplan.py
"""
import os, sys, json
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: pip install openpyxl"); sys.exit(1)

IST = timezone(timedelta(hours=5, minutes=30))
ts = datetime.now(IST).strftime("%Y%m%d_%H%M%S")
fn = f"MIPI_TestPlan_{ts}.xlsx"
script_dir = os.path.dirname(os.path.abspath(__file__))
fp = os.path.join(script_dir, fn)

json_data = json.loads(open(os.path.join(script_dir, 'testplan_data.json')).read())

wb = Workbook()
ws1 = wb.active
ws1.title = "TestPlan"
tp_cols = ["Index","SS / Module","Feature","Test Case Name","Test Description","Speed","Mode","Memory Start Offset","Memory End Offset","Remarks","Test Steps / Procedure","Impacted Registers","Validation / Acceptance Criteria","Code Generation"]
ws1.append(tp_cols)

ws2 = wb.create_sheet("MetaData")
md_cols = ["Index","Test Case Name","Meta Test Description","Meta Test Steps / Procedure","Meta Impacted Registers","Meta Validation / Acceptance Criteria","Meta Headers","Meta Macros","Meta Arrays"]
ws2.append(md_cols)

for row in json_data:
    ws1.append([row.get(c, "") for c in tp_cols])
    ws2.append([row.get(c, "") for c in md_cols])

hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap = Alignment(wrap_text=True, vertical="top")

for ws in [ws1, ws2]:
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = wrap
    ws.freeze_panes = "A2"
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    cell.alignment = wrap
                    max_len = max(max_len, min(len(str(cell.value)), 80))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)

ws2.sheet_state = 'veryHidden'
wb.save(fp)
print(f"SUCCESS: {fn} ({os.path.getsize(fp)} bytes)")
print(f"PATH: {fp}")
