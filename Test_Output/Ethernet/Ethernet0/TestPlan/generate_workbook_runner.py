#!/usr/bin/env python3
"""
Runner script - generates the Excel workbook and outputs base64 for upload.
"""
import json, sys, os, base64, io
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
except ImportError:
    print("pip install openpyxl required")
    sys.exit(1)

# Run the generator
exec(open("Ethernet_TestPlan_Generator_v2.py").read())
