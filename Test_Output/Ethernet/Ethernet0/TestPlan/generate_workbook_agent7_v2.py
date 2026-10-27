#!/usr/bin/env python3
"""
Agent 7: Excel Workbook Generator v2
Reads JSON data from external files and generates the workbook.
Requires: pip install openpyxl
Usage: python generate_workbook_agent7_v2.py
"""
import json, os, sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: pip install openpyxl"); sys.exit(1)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

def load_json(filename):
    path = os.path.join(SCRIPT_DIR, filename)
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)

def generate():
    TESTPLAN_JSON = load_json('agent7_testplan_input.json')
    META_JSON = load_json('agent7_metadata_input.json')
    IP_NAME = "Ethernet"

    # Cross-dataset validation
    tp_names = {tc['Test Case Name'] for tc in TESTPLAN_JSON}
    meta_names = {m['Test Case Name'] for m in META_JSON}
    for name in tp_names:
        if name not in meta_names:
            print(f"VALIDATION FAILED: '{name}' in TestPlan but not in MetaData")
            sys.exit(1)

    IST = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(IST)
    ts = now_ist.strftime("%Y%m%d_%H%M%S")
    fn = f"{IP_NAME}_TestPlan_{ts}.xlsx"
    output_path = os.path.join(SCRIPT_DIR, fn)

    wb = Workbook()

    # ---- TestPlan Sheet ----
    ws = wb.active
    ws.title = "TestPlan"
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wa = Alignment(wrap_text=True, vertical="top")

    tp_headers = [
        "Index", "SS / Module", "Feature", "Test Case Name", "Speed", "Mode",
        "Memory Start Offset", "Memory End Offset", "Test Description",
        "Test Steps / Procedure", "Impacted Registers",
        "Validation / Acceptance Criteria", "Remarks",
        "Code Generation (Required / Not)"
    ]
    tp_keys = [
        "Index", "SS / Module", "Feature", "Test Case Name", "Speed", "Mode",
        "Memory Start Offset", "Memory End Offset", "Test Description",
        "Test Steps / Procedure", "Impacted Registers",
        "Validation / Acceptance Criteria", "Remarks"
    ]

    for ci, h in enumerate(tp_headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hf; c.fill = hfill; c.alignment = wa

    for ri, tc in enumerate(TESTPLAN_JSON, 2):
        for ci, k in enumerate(tp_keys, 1):
            ws.cell(row=ri, column=ci, value=tc.get(k, "")).alignment = wa
        # Column 14 (Code Generation) intentionally left blank

    # Data validation dropdown
    dv = DataValidation(type="list", formula1='"Required,Not Required"', allow_blank=True)
    dv.error = "Please select Required or Not Required"
    dv.errorTitle = "Invalid Entry"
    dv.prompt = "Select Code Generation option"
    dv.promptTitle = "Code Generation"
    lr = 1 + len(TESTPLAN_JSON)
    dv.add(f"N2:N{lr}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "A2"

    # Auto-size columns
    for ci in range(1, len(tp_headers) + 1):
        max_len = len(tp_headers[ci - 1])
        for row in ws.iter_rows(min_row=2, max_row=lr, min_col=ci, max_col=ci):
            for cell in row:
                if cell.value:
                    for line in str(cell.value).split('\n'):
                        max_len = max(max_len, len(line))
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 2, 80)

    # ---- MetaData Sheet ----
    wm = wb.create_sheet("MetaData")
    md_headers = [
        "Test Case Name", "Meta Headers", "Meta Macros", "Meta Arrays",
        "Meta Types", "Meta Test Description", "Meta Test Steps / Procedure",
        "Meta Impacted Registers", "Meta Validation / Acceptance Criteria"
    ]

    for ci, h in enumerate(md_headers, 1):
        c = wm.cell(row=1, column=ci, value=h)
        c.font = hf; c.fill = hfill; c.alignment = wa

    for ri, m in enumerate(META_JSON, 2):
        for ci, k in enumerate(md_headers, 1):
            wm.cell(row=ri, column=ci, value=m.get(k, "")).alignment = wa

    wm.freeze_panes = "A2"
    wm.sheet_state = "veryHidden"

    md_lr = 1 + len(META_JSON)
    for ci in range(1, len(md_headers) + 1):
        max_len = len(md_headers[ci - 1])
        for row in wm.iter_rows(min_row=2, max_row=md_lr, min_col=ci, max_col=ci):
            for cell in row:
                if cell.value:
                    for line in str(cell.value).split('\n'):
                        max_len = max(max_len, len(line))
        wm.column_dimensions[get_column_letter(ci)].width = min(max_len + 2, 80)

    wb.save(output_path)
    print(f"Workbook saved: {output_path}")

    # Post-save validation
    wb2 = load_workbook(output_path)
    assert "TestPlan" in wb2.sheetnames, "TestPlan sheet missing"
    assert "MetaData" in wb2.sheetnames, "MetaData sheet missing"
    tp_rows = wb2["TestPlan"].max_row - 1
    md_rows = wb2["MetaData"].max_row - 1
    assert tp_rows == len(TESTPLAN_JSON), f"TestPlan rows mismatch: {tp_rows} vs {len(TESTPLAN_JSON)}"
    assert md_rows == len(META_JSON), f"MetaData rows mismatch: {md_rows} vs {len(META_JSON)}"
    fsize = os.path.getsize(output_path)
    assert fsize > 0, "File size is 0"
    print(f"Validation PASSED: TestPlan={tp_rows}, MetaData={md_rows}, Size={fsize} bytes")
    print(f"Filename: {fn}")
    return output_path

if __name__ == "__main__":
    generate()
