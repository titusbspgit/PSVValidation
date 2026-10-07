#!/usr/bin/env python3
"""Ethernet TestPlan Excel Generator - Agent 7 Direct Generation"""
import json, os, sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# === JSON DATA (embedded from Agent pipeline) ===
json_data = [
  {"Index":"1","SS / Module":"Ethernet","Test Case Name":"ethernet0_reg_wr_rd_test","Feature":"Register Write-Read Verification","Meta Headers":"<stdio.h>; <stdlib.h>; \"test_define.c\"; <test_common.h>; <ethernet0/ethernet0_def.h>; <ethernet0/ethernet0_offset.h>","Meta Macros":"#define SOFT_RST_REG_ADDRESS 0x00000000\n#define SOFT_RST_REG_DATA 0x00000000\n#define CNT 434","Meta Arrays":"const unsigned long int addr_array[434] = {\n  mizar_ETHERNET0_MAC_CONFIGURATION,\n  mizar_ETHERNET0_MAC_EXT_CONFIGURATION,\n  mizar_ETHERNET0_MAC_PACKET_FILTER,\n  mizar_ETHERNET0_MAC_WD_JB_TIMEOUT,\n  mizar_ETHERNET0_MAC_HASH_TABLE_REG0\n};\n\nconst int default_value_array[434] = {\n  ETHERNET0_MAC_CONFIGURATION_DEFAULT_VAL,\n  ETHERNET0_MAC_EXT_CONFIGURATION_DEFAULT_VAL,\n  ETHERNET0_MAC_PACKET_FILTER_DEFAULT_VAL,\n  ETHERNET0_MAC_WD_JB_TIMEOUT_DEFAULT_VAL,\n  ETHERNET0_MAC_HASH_TABLE_REG0_DEFAULT_VAL\n};\n\nconst int read_mask_array[434] = {\n  ETHERNET0_MAC_CONFIGURATION_READ_MASK,\n  ETHERNET0_MAC_EXT_CONFIGURATION_READ_MASK,\n  ETHERNET0_MAC_PACKET_FILTER_READ_MASK,\n  ETHERNET0_MAC_WD_JB_TIMEOUT_READ_MASK,\n  ETHERNET0_MAC_HASH_TABLE_REG0_READ_MASK\n};\n\nconst int write_mask_array[434] = {\n  ETHERNET0_MAC_CONFIGURATION_WRITE_MASK,\n  ETHERNET0_MAC_EXT_CONFIGURATION_WRITE_MASK,\n  ETHERNET0_MAC_PACKET_FILTER_WRITE_MASK,\n  ETHERNET0_MAC_WD_JB_TIMEOUT_WRITE_MASK,\n  ETHERNET0_MAC_HASH_TABLE_REG0_WRITE_MASK\n};\n\nconst int skip_array[434] = {0,0,0,0,0};\nconst int skip_rst_array[434] = {0,0,0,0,0};\nint chk_val[6] = {0xffffffff,0xaaaaaaaa,0x55555555,0x00000000,0xA5A5A5A5,0xffff0000};","Speed":"NA","Mode":"NA","Memory Start Offset":"NA","Memory End Offset":"NA","Meta Test Description":"This testcase performs a comprehensive register write-read verification for the Ethernet0 MAC IP block.","Test Description":"This test verifies the default reset values and write-read accessibility of Ethernet0 MAC registers.","Meta Test Steps / Procedure":"Steps 1-43 covering chk_rst_val and chk_rd_wr phases","Test Steps / Procedure":"1. Initialize global fail counters to zero.\n2. Execute default value check phase.\n3. Execute write-read check phase with six test patterns.\n4. Evaluate overall result.","Meta Impacted Registers":"mizar_ETHERNET0_MAC_CONFIGURATION; mizar_ETHERNET0_MAC_EXT_CONFIGURATION; mizar_ETHERNET0_MAC_PACKET_FILTER; mizar_ETHERNET0_MAC_WD_JB_TIMEOUT; mizar_ETHERNET0_MAC_HASH_TABLE_REG0","Impacted Registers":"MAC_Configuration; MAC_Ext_Configuration; MAC_Packet_Filter; MAC_WD_JB_Timeout; MAC_Hash_Table_Reg0","Meta Validation / Acceptance Criteria":"Phase 1 and Phase 2 validation criteria","Validation / Acceptance Criteria":"All registers must pass default value and write-read checks.","Remarks":"CNT=434 but only 5 entries initialized. soft_reset_chk() commented out."},
  {"Index":"2","SS / Module":"Ethernet","Test Case Name":"ethernet0_rx_basic_test","Feature":"Basic RX DMA Reception","Meta Headers":"<stdio.h>; <stdlib.h>; <math.h>; <ethernet0/ethernet0_def.h>; <test_common.h>; <ethernet0/ethernet0_offset.h>; <ethernet0/ethernet0_funcs.h>; <ethernet0/ethernet0_programming_sequence.h>","Meta Macros":"NA","Meta Arrays":"NA","Speed":"1Gbps (default); 100Mbps (ETH_100M); 10Mbps (ETH_10M)","Mode":"Full Duplex (default); Half Duplex (HALF_DUPLEX)","Memory Start Offset":"NA","Memory End Offset":"NA","Meta Test Description":"This testcase verifies basic Ethernet0 RX DMA reception across 4 DMA channels.","Test Description":"This test verifies basic Ethernet0 receive DMA operation with full MAC/MTL/DMA configuration.","Meta Test Steps / Procedure":"Steps 1-130 covering full RX configuration and ISR","Test Steps / Procedure":"1. Configure GIC.\n2. Configure MAC/MTL/DMA.\n3. Preload RX descriptors.\n4. Start RX and wait for interrupts.","Meta Impacted Registers":"91 registers across MAC, MTL, and DMA blocks","Impacted Registers":"87 Ethernet IP registers + 4 external hex addresses","Meta Validation / Acceptance Criteria":"Interrupt-driven completion with ISR validation","Validation / Acceptance Criteria":"All transfer-complete interrupts must be received.","Remarks":"Conditional compilation for speed/duplex/interface. 4 external hex addresses unresolved."},
  {"Index":"3","SS / Module":"Ethernet","Test Case Name":"ethernet0_tx_basic_test","Feature":"Basic TX DMA Transmission","Meta Headers":"<stdio.h>; <stdlib.h>; <math.h>; <test_common.h>; <ethernet0/ethernet0_def.h>; <ethernet0/ethernet0_offset.h>; <ethernet0/ethernet0_funcs.h>; <ethernet0/ethernet0_programming_sequence.h>","Meta Macros":"NA","Meta Arrays":"NA","Speed":"1Gbps (default); 100Mbps (ETH_100M); 10Mbps (ETH_10M)","Mode":"Full Duplex (default); Half Duplex (HALF_DUPLEX)","Memory Start Offset":"NA","Memory End Offset":"NA","Meta Test Description":"This testcase verifies basic Ethernet0 TX DMA transmission on CH0.","Test Description":"This test verifies basic Ethernet0 transmit DMA operation with full configuration.","Meta Test Steps / Procedure":"Steps 1-81 covering TX configuration and ISR","Test Steps / Procedure":"1. Configure GIC.\n2. Configure MAC/MTL/DMA.\n3. Preload TX descriptors.\n4. Start TX on CH0 and wait for interrupts.","Meta Impacted Registers":"51 registers across MAC, MTL, and DMA blocks","Impacted Registers":"47 Ethernet IP registers + 4 external hex addresses","Meta Validation / Acceptance Criteria":"Interrupt-driven TX completion","Validation / Acceptance Criteria":"All TX transfer-complete interrupts must be received.","Remarks":"Only CH0 TX started. CH1-3 TX_CONTROL commented out. non_secure_prot_nic() called."},
  {"Index":"4","SS / Module":"Ethernet","Test Case Name":"ethernet0_tx_rx_multi_chnl_test","Feature":"Multi-Channel TX/RX DMA Transfer","Meta Headers":"<stdio.h>; <stdlib.h>; <math.h>; <test_common.h>; <ethernet0/ethernet0_def.h>; <ethernet0/ethernet0_offset.h>; <ethernet0/ethernet0_funcs.h>; <ethernet0/ethernet0_programming_sequence.h>","Meta Macros":"NA","Meta Arrays":"NA","Speed":"1Gbps (default); 100Mbps (ETH_100M); 10Mbps (ETH_10M)","Mode":"Full Duplex (default); Half Duplex (HALF_DUPLEX)","Memory Start Offset":"NA","Memory End Offset":"NA","Meta Test Description":"This testcase verifies multi-channel TX and RX DMA operation across all four DMA channels.","Test Description":"This test verifies multi-channel Ethernet0 TX and RX DMA with sequential channel activation.","Meta Test Steps / Procedure":"Steps 1-184 covering full multi-channel TX/RX configuration, 4 phases, and ISR","Test Steps / Procedure":"1. Configure GIC.\n2. Configure MAC/MTL/DMA for all 4 channels.\n3. Preload TX and RX descriptors.\n4. Execute 4 sequential phases (CH3,CH2,CH1,CH0).","Meta Impacted Registers":"96 registers across MAC, MTL, DMA, and MMC blocks","Impacted Registers":"92 Ethernet IP registers + 4 external hex addresses","Meta Validation / Acceptance Criteria":"4-phase packet count validation (10,20,30,40)","Validation / Acceptance Criteria":"All 40 TX and 40 RX packets must complete across all 4 channels.","Remarks":"Sequential channel activation CH3->CH2->CH1->CH0. MTL_RXQ_DMA_MAP0 dynamically updated. ISR manages RX desc tail pointer wrap-around."}
]

# === IST Timestamp ===
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"Ethernet_TestPlan_{timestamp}.xlsx"

# === Styling ===
header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
cell_align = Alignment(vertical='top', wrap_text=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

# === Create Workbook ===
wb = Workbook()

# === TestPlan Sheet ===
ws_tp = wb.active
ws_tp.title = "TestPlan"

tp_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria"
]

tp_keys = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria"
]

# Write headers
for col_idx, col_name in enumerate(tp_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Write data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, key in enumerate(tp_keys, 1):
        val = row_data.get(key, "")
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = cell_align
        cell.border = thin_border

# Auto-size columns
for col_idx, col_name in enumerate(tp_columns, 1):
    max_len = len(col_name)
    for row_idx in range(2, len(json_data) + 2):
        val = ws_tp.cell(row=row_idx, column=col_idx).value
        if val:
            lines = str(val).split('\n')
            for line in lines:
                max_len = max(max_len, len(line))
    width = min(max_len + 4, 80)
    ws_tp.column_dimensions[ws_tp.cell(row=1, column=col_idx).column_letter].width = width

ws_tp.freeze_panes = 'A2'

# === MetaData Sheet ===
ws_md = wb.create_sheet(title="MetaData")

md_columns = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

md_keys = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

# Write headers
for col_idx, col_name in enumerate(md_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Write data rows
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, key in enumerate(md_keys, 1):
        val = row_data.get(key, "")
        cell = ws_md.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = cell_align
        cell.border = thin_border

# Auto-size columns
for col_idx, col_name in enumerate(md_columns, 1):
    max_len = len(col_name)
    for row_idx in range(2, len(json_data) + 2):
        val = ws_md.cell(row=row_idx, column=col_idx).value
        if val:
            lines = str(val).split('\n')
            for line in lines:
                max_len = max(max_len, len(line))
    width = min(max_len + 4, 80)
    ws_md.column_dimensions[ws_md.cell(row=1, column=col_idx).column_letter].width = width

ws_md.freeze_panes = 'A2'
ws_md.sheet_state = 'veryHidden'

# === Save ===
script_dir = os.path.dirname(os.path.abspath(__file__))
output_path = os.path.join(script_dir, filename)
wb.save(output_path)

# === Validate ===
file_size = os.path.getsize(output_path)
print(f"SUCCESS: Generated {filename}")
print(f"Path: {output_path}")
print(f"Size: {file_size} bytes")
print(f"Rows TestPlan: {len(json_data)}")
print(f"Rows MetaData: {len(json_data)}")
print(f"FILENAME={filename}")

# Verify reopenable
from openpyxl import load_workbook
wb2 = load_workbook(output_path)
assert "TestPlan" in wb2.sheetnames
assert "MetaData" in wb2.sheetnames
print("VALIDATION=PASSED")
