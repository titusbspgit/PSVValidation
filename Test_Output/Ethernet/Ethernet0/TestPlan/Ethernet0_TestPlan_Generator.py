#!/usr/bin/env python3
"""
Ethernet0 TestPlan Excel Generator
Generates Ethernet0_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx
Run: pip install openpyxl && python Ethernet0_TestPlan_Generator.py
"""
import json
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"Ethernet0_TestPlan_{timestamp}.xlsx"

# Output directory
output_dir = os.path.dirname(os.path.abspath(__file__))
output_path = os.path.join(output_dir, filename)

# JSON Data
json_data = [
    {
        "Index": "1",
        "SS / Module": "ETHERNET0",
        "Test Case Name": "ethernet0_reg_wr_rd_test",
        "Feature": "Register Write/Read Verification",
        "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include \"test_define.c\"\n#include <test_common.h>\n#include <ethernet0/ethernet0_def.h>\n#include <ethernet0/ethernet0_offset.h>",
        "Meta Macros": "#define CNT 434\n#define SOFT_RST_REG_ADDRESS 0x00000000\n#define SOFT_RST_REG_DATA 0x00000000\nETHERNET0_MAC_CONFIGURATION_DEFAULT_VAL\nETHERNET0_MAC_EXT_CONFIGURATION_DEFAULT_VAL\nETHERNET0_MAC_PACKET_FILTER_DEFAULT_VAL\nETHERNET0_MAC_WD_JB_TIMEOUT_DEFAULT_VAL\nETHERNET0_MAC_HASH_TABLE_REG0_DEFAULT_VAL\nETHERNET0_MAC_HASH_TABLE_REG1_DEFAULT_VAL\nETHERNET0_MAC_VLAN_TAG_CTRL_DEFAULT_VAL\nETHERNET0_MAC_VLAN_TAG_DATA_DEFAULT_VAL\nETHERNET0_MAC_VLAN_HASH_TABLE_DEFAULT_VAL\nETHERNET0_MAC_Q0_TX_FLOW_CTRL_DEFAULT_VAL\nETHERNET0_MAC_CONFIGURATION_WRITE_MASK\nETHERNET0_MAC_EXT_CONFIGURATION_WRITE_MASK\nETHERNET0_MAC_PACKET_FILTER_WRITE_MASK\nETHERNET0_MAC_WD_JB_TIMEOUT_WRITE_MASK\nETHERNET0_MAC_HASH_TABLE_REG0_WRITE_MASK\nETHERNET0_MAC_HASH_TABLE_REG1_WRITE_MASK\nETHERNET0_MAC_VLAN_TAG_CTRL_WRITE_MASK\nETHERNET0_MAC_VLAN_TAG_DATA_WRITE_MASK\nETHERNET0_MAC_VLAN_HASH_TABLE_WRITE_MASK\nETHERNET0_MAC_Q0_TX_FLOW_CTRL_WRITE_MASK\nETHERNET0_MAC_CONFIGURATION_READ_MASK\nETHERNET0_MAC_EXT_CONFIGURATION_READ_MASK\nETHERNET0_MAC_PACKET_FILTER_READ_MASK\nETHERNET0_MAC_WD_JB_TIMEOUT_READ_MASK\nETHERNET0_MAC_HASH_TABLE_REG0_READ_MASK\nETHERNET0_MAC_HASH_TABLE_REG1_READ_MASK\nETHERNET0_MAC_VLAN_TAG_CTRL_READ_MASK\nETHERNET0_MAC_VLAN_TAG_DATA_READ_MASK\nETHERNET0_MAC_VLAN_HASH_TABLE_READ_MASK\nETHERNET0_MAC_Q0_TX_FLOW_CTRL_READ_MASK\nETHERNET0_MAC_Q1_TX_FLOW_CTRL_READ_MASK\nETHERNET0_MAC_Q2_TX_FLOW_CTRL_READ_MASK\nETHERNET0_MAC_Q3_TX_FLOW_CTRL_READ_MASK\nETHERNET0_MAC_RX_FLOW_CTRL_READ_MASK\nDEBUG_DISPLAY\nETH_10M\nETH_100M\nHALF_DUPLEX",
        "Meta Arrays": "const int addr_array[434] = {mizar_ETHERNET0_MAC_CONFIGURATION, mizar_ETHERNET0_MAC_EXT_CONFIGURATION, mizar_ETHERNET0_MAC_PACKET_FILTER, mizar_ETHERNET0_MAC_WD_JB_TIMEOUT, mizar_ETHERNET0_MAC_HASH_TABLE_REG0, mizar_ETHERNET0_MAC_HASH_TABLE_REG1, mizar_ETHERNET0_MAC_VLAN_TAG_CTRL, mizar_ETHERNET0_MAC_VLAN_TAG_DATA, mizar_ETHERNET0_MAC_VLAN_HASH_TABLE, mizar_ETHERNET0_MAC_Q0_TX_FLOW_CTRL, mizar_ETHERNET0_MAC_Q1_TX_FLOW_CTRL, mizar_ETHERNET0_MAC_Q2_TX_FLOW_CTRL, mizar_ETHERNET0_MAC_Q3_TX_FLOW_CTRL, mizar_ETHERNET0_MAC_RX_FLOW_CTRL, ...};\nconst int default_value_array[434] = {ETHERNET0_MAC_CONFIGURATION_DEFAULT_VAL, ETHERNET0_MAC_EXT_CONFIGURATION_DEFAULT_VAL, ETHERNET0_MAC_PACKET_FILTER_DEFAULT_VAL, ETHERNET0_MAC_WD_JB_TIMEOUT_DEFAULT_VAL, ETHERNET0_MAC_HASH_TABLE_REG0_DEFAULT_VAL, ETHERNET0_MAC_HASH_TABLE_REG1_DEFAULT_VAL, ETHERNET0_MAC_VLAN_TAG_CTRL_DEFAULT_VAL, ETHERNET0_MAC_VLAN_TAG_DATA_DEFAULT_VAL, ETHERNET0_MAC_VLAN_HASH_TABLE_DEFAULT_VAL, ETHERNET0_MAC_Q0_TX_FLOW_CTRL_DEFAULT_VAL, ...};\nconst int write_mask_array[434] = {ETHERNET0_MAC_CONFIGURATION_WRITE_MASK, ETHERNET0_MAC_EXT_CONFIGURATION_WRITE_MASK, ETHERNET0_MAC_PACKET_FILTER_WRITE_MASK, ETHERNET0_MAC_WD_JB_TIMEOUT_WRITE_MASK, ETHERNET0_MAC_HASH_TABLE_REG0_WRITE_MASK, ETHERNET0_MAC_HASH_TABLE_REG1_WRITE_MASK, ETHERNET0_MAC_VLAN_TAG_CTRL_WRITE_MASK, ETHERNET0_MAC_VLAN_TAG_DATA_WRITE_MASK, ETHERNET0_MAC_VLAN_HASH_TABLE_WRITE_MASK, ETHERNET0_MAC_Q0_TX_FLOW_CTRL_WRITE_MASK, ...};\nconst int read_mask_array[434] = {ETHERNET0_MAC_CONFIGURATION_READ_MASK, ETHERNET0_MAC_EXT_CONFIGURATION_READ_MASK, ETHERNET0_MAC_PACKET_FILTER_READ_MASK, ETHERNET0_MAC_WD_JB_TIMEOUT_READ_MASK, ETHERNET0_MAC_HASH_TABLE_REG0_READ_MASK, ETHERNET0_MAC_HASH_TABLE_REG1_READ_MASK, ETHERNET0_MAC_VLAN_TAG_CTRL_READ_MASK, ETHERNET0_MAC_VLAN_TAG_DATA_READ_MASK, ETHERNET0_MAC_VLAN_HASH_TABLE_READ_MASK, ETHERNET0_MAC_Q0_TX_FLOW_CTRL_READ_MASK, ETHERNET0_MAC_Q1_TX_FLOW_CTRL_READ_MASK, ETHERNET0_MAC_Q2_TX_FLOW_CTRL_READ_MASK, ...};\nconst int skip_array[434] = {0, 0, 0, 0, 0, 0, 1, 1, 0, 1, ...};\nint chk_val[6] = {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000};",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs register write/read verification for 434 Ethernet0 registers. It includes three phases: (1) chk_rst_val() - reads all 434 registers and compares against expected default/reset values stored in default_value_array[434], incrementing def_fail_cnt on mismatch; (2) chk_rd_wr() - iterates over 6 test patterns (0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000) writing each pattern masked with write_mask_array[i] to each register, reading back and comparing against expected value (data_wr & write_mask_array[i] & read_mask_array[i]), incrementing wr_fail_cnt on mismatch; (3) soft_reset_chk() - triggers a soft reset and re-verifies register values. Registers flagged in skip_array (value 1) are skipped. Registers with write_mask_array value of 0x00000000 are considered non-writable and skipped for writing. The test uses addr_array[434] containing mizar_ETHERNET0_* register address macros. Global variables: int data_rd, data_wr; int def_fail_cnt = 0, wr_fail_cnt = 0. The main entry point is int test_case() which orchestrates the three phases and returns pass/fail based on accumulated failure counts.",
        "Test Description": "Verify default reset values and write/read accessibility of 434 Ethernet0 registers. The test reads each register after reset and compares against expected default values. Then it writes six different test patterns to each writable register, reads back, and verifies correctness using write and read masks. Certain registers (e.g., MAC_VLAN_Tag_Ctrl, MAC_VLAN_Tag_Data) are skipped per the skip array. A soft reset check is also performed to verify registers return to default values after reset. The test passes if no default value mismatches and no write/read mismatches are detected.",
        "Meta Test Steps / Procedure": "1. #include <stdio.h>\n2. #include <stdlib.h>\n3. #include \"test_define.c\"\n4. #include <test_common.h>\n5. test_define.c includes: #include <ethernet0/ethernet0_def.h>, #include <ethernet0/ethernet0_offset.h>, #define CNT 434\n6. Global declarations: int data_rd, data_wr; int def_fail_cnt = 0, wr_fail_cnt = 0;\n7. #define SOFT_RST_REG_ADDRESS 0x00000000\n8. #define SOFT_RST_REG_DATA 0x00000000\n9. const int addr_array[434] declared with mizar_ETHERNET0_* register address macros\n10. const int default_value_array[434] declared with ETHERNET0_DEFAULT_VAL macros\n11. const int write_mask_array[434] declared with ETHERNET0WRITE_MASK macros\n12. const int read_mask_array[434] declared with ETHERNET0*_READ_MASK macros\n13. const int skip_array[434] declared with 0/1 values (indices 6,7,9 are 1 meaning skip)\n14. int test_case() entry point called\n15. Phase 1 - chk_rst_val(): for(i=0; i<CNT; i++) { addr=addr_array[i]; if(skip_array[i]==1) continue; data_rd = read_reg(addr); if(data_rd == default_value_array[i]) { #ifdef DEBUG_DISPLAY printf(\"RST : PASS Reading Default value from Address :0x%x Expected : 0x%x\\tRead_data : 0x%x\\n\", addr, default_value_array[i], data_rd); #endif } else { def_fail_cnt++; printf(\"RST : Failed Default value mismatch Addr :0x%x Expected : 0x%x\\tRead_data : 0x%x\\n\", addr, default_value_array[i], data_rd); } }\n16. Phase 2 - chk_rd_wr(): int chk_val[6] = {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}; for(j=0; j<6; j++) { data_wr = chk_val[j]; for(i=0; i<CNT; i++) { addr=addr_array[i]; if(skip_array[i]==1) continue; if(write_mask_array[i]==0x00000000) continue; write_reg(addr, data_wr & write_mask_array[i]); data_rd = read_reg(addr); exp_val = data_wr & write_mask_array[i] & read_mask_array[i]; if(data_rd != exp_val) { wr_fail_cnt++; } } }\n17. Phase 3 - soft_reset_chk(): writes to SOFT_RST_REG_ADDRESS with SOFT_RST_REG_DATA and re-verifies register reset values\n18. test_case() returns pass/fail based on def_fail_cnt and wr_fail_cnt",
        "Test Steps / Procedure": "1. Initialize failure counters (def_fail_cnt, wr_fail_cnt) to zero.\n2. Phase 1 - Default Value Check (chk_rst_val): Iterate over all 434 registers using addr_array. Skip registers flagged in skip_array. Read each register and compare against expected default value from default_value_array. Log pass/fail for each register. Increment def_fail_cnt on mismatch.\n3. Phase 2 - Write/Read Check (chk_rd_wr): For each of 6 test patterns (0xFFFFFFFF, 0xAAAAAAAA, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xFFFF0000), iterate over all 434 registers. Skip registers flagged in skip_array. Skip non-writable registers (write mask = 0x00000000). Write test pattern masked with write mask to each register. Read back register value. Compare read value against expected value (test pattern AND write mask AND read mask). Increment wr_fail_cnt on mismatch.\n4. Phase 3 - Soft Reset Check (soft_reset_chk): Trigger soft reset and verify registers return to default values.\n5. Return overall pass/fail based on accumulated failure counts.",
        "Meta Impacted Registers": "MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_WD_JB_Timeout, MAC_Hash_Table_Reg0, MAC_Hash_Table_Reg1, MAC_VLAN_Tag_Ctrl, MAC_VLAN_Tag_Data, MAC_VLAN_Hash_Table, MAC_Q0_Tx_Flow_Ctrl, MAC_Q1_Tx_Flow_Ctrl, MAC_Q2_Tx_Flow_Ctrl, MAC_Q3_Tx_Flow_Ctrl, MAC_Rx_Flow_Ctrl, MAC_Address0_High, MAC_Address0_Low, MAC_Address1_Low, MAC_Address2_Low, MAC_Address3_Low, DMA_CH0_TxDesc_List_Address, DMA_CH1_TxDesc_List_Address, DMA_CH2_TxDesc_List_Address, DMA_CH3_TxDesc_List_Address, DMA_CH0_RxDesc_List_Address, DMA_CH1_RxDesc_List_Address, DMA_CH2_RxDesc_List_Address, DMA_CH3_RxDesc_List_Address, MAC_RxQ_Ctrl0, MAC_RxQ_Ctrl1, MAC_RxQ_Ctrl2, MAC_RxQ_Ctrl4",
        "Impacted Registers": "MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_WD_JB_Timeout, MAC_Hash_Table_Reg0, MAC_Hash_Table_Reg1, MAC_VLAN_Tag_Ctrl, MAC_VLAN_Tag_Data, MAC_VLAN_Hash_Table, MAC_Q0_Tx_Flow_Ctrl, MAC_Q1_Tx_Flow_Ctrl, MAC_Q2_Tx_Flow_Ctrl, MAC_Q3_Tx_Flow_Ctrl, MAC_Rx_Flow_Ctrl, MAC_Address0_High, MAC_Address0_Low, MAC_Address1_Low, MAC_Address2_Low, MAC_Address3_Low, DMA_CH0_TxDesc_List_Address, DMA_CH1_TxDesc_List_Address, DMA_CH2_TxDesc_List_Address, DMA_CH3_TxDesc_List_Address, DMA_CH0_RxDesc_List_Address, DMA_CH1_RxDesc_List_Address, DMA_CH2_RxDesc_List_Address, DMA_CH3_RxDesc_List_Address, MAC_RxQ_Ctrl0, MAC_RxQ_Ctrl1, MAC_RxQ_Ctrl2, MAC_RxQ_Ctrl4",
        "Meta Validation / Acceptance Criteria": "Phase 1 (chk_rst_val): For each register i (0 to 433), if skip_array[i] != 1: data_rd = read_reg(addr_array[i]); if(data_rd == default_value_array[i]) then PASS, else def_fail_cnt++ and FAIL.\nPhase 2 (chk_rd_wr): For each test pattern chk_val[j] (j=0..5) and each register i (0 to 433), if skip_array[i] != 1 and write_mask_array[i] != 0x00000000: write_reg(addr, data_wr & write_mask_array[i]); data_rd = read_reg(addr); exp_val = data_wr & write_mask_array[i] & read_mask_array[i]; if(data_rd != exp_val) then wr_fail_cnt++.\nPhase 3 (soft_reset_chk): After soft reset, verify registers return to default values.\nOverall pass condition: def_fail_cnt == 0 and wr_fail_cnt == 0.",
        "Validation / Acceptance Criteria": "1. All 434 Ethernet0 registers (excluding those in the skip array) must read back their expected default values after reset. No default value mismatches allowed (def_fail_cnt must be 0).\n2. For each of the 6 test patterns (0xFFFFFFFF, 0xAAAAAAAA, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xFFFF0000), all writable registers must correctly store and read back the written value when masked with write and read masks. No write/read mismatches allowed (wr_fail_cnt must be 0).\n3. After soft reset, all registers must return to their default values.\n4. Registers flagged in skip_array (indices 6, 7, 9 corresponding to MAC_VLAN_Tag_Ctrl, MAC_VLAN_Tag_Data, and others) are excluded from verification.\n5. Non-writable registers (write_mask = 0x00000000) are excluded from write/read verification.",
        "Remarks": "The test covers 434 Ethernet0 registers defined in addr_array. Three registers are skipped (skip_array indices 6, 7, 9 set to 1). The test uses conditional compilation macros ETH_10M, ETH_100M, and HALF_DUPLEX for speed/duplex configuration of MAC_Configuration register. DEBUG_DISPLAY macro enables verbose pass/skip logging. The test is structured in three phases: default value check, write/read pattern check with 6 patterns, and soft reset verification."
    },
    {
        "Index": "2",
        "SS / Module": "ETHERNET0",
        "Test Case Name": "ethernet0_tx_basic_test",
        "Feature": "Basic TX Transmission",
        "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include <math.h>\n#include <ethernet0/ethernet0_def.h>\n#include <test_common.h>\n#include <ethernet0/ethernet0_offset.h>\n#include <ethernet0/ethernet0_funcs.h>\n#include <ethernet0/ethernet0_programming_sequence.h>",
        "Meta Macros": "ETH_10M\nETH_100M\nHALF_DUPLEX\nSEL_ENET0\nSEL_ENET1\nSEL_ENET2\nSEL_ENET3\nDEBUG_DISPLAY",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a basic Ethernet0 TX (transmit) operation on DMA Channel 0. It initializes the Ethernet interface via enet_intf_sel() based on conditional compilation macros SEL_ENET0/SEL_ENET1/SEL_ENET2/SEL_ENET3 (default channel 4 if none defined). It configures the MAC_CONFIGURATION register with speed and duplex settings based on ETH_10M, ETH_100M, and HALF_DUPLEX macros. Only DMA_CH0_TX_CONTROL is actively written with 0x00100007 (START TX); CH1, CH2, CH3 TX_CONTROL writes are commented out.",
        "Test Description": "Perform a basic Ethernet0 transmit operation using DMA Channel 0. The test configures the MAC layer (speed, duplex, MAC address), MTL layer (TX queue operation mode with 4096-byte queue and store-and-forward), and DMA layer (descriptor list addresses, tail pointers, channel control). Only DMA Channel 0 TX is actively started; Channels 1-3 TX control writes are commented out. The test transmits packets and polls for transfer completion interrupts, verifying successful transmission.",
        "Meta Test Steps / Procedure": "1. #include <stdio.h>\n2. #include <stdlib.h>\n3. #include <math.h>\n4. #include headers\n5. Global declarations\n6. int test_case() entry point\n7. Ethernet interface selection\n8. MAC Configuration with speed/duplex\n9. write_enet_reg MAC_EXT_CONFIGURATION, MAC_ADDRESS0_HIGH, MAC_ADDRESS0_LOW\n10. DMA channel control init\n11. MTL_TXQ0_OPERATION_MODE = 0x000f000a\n12. DMA TX descriptor list addresses configured\n13. DMA TX descriptor tail pointers written\n14. DMA_CH0_TX_CONTROL = 0x00100007 (START TX)\n15. VIP sequencer triggered\n16. Poll for transfer completion interrupts",
        "Test Steps / Procedure": "1. Select Ethernet interface using enet_intf_sel() based on SEL_ENETx compile-time macro (default channel 4).\n2. If ETH_10M mode, set trns_count to 6.\n3. Configure MAC_Configuration register with appropriate speed and duplex settings (10M/100M/1G, Half/Full Duplex).\n4. Write MAC_Ext_Configuration register with 0x0.\n5. Program MAC_Address0_High with 0x80000607 and MAC_Address0_Low with 0x08090a0b.\n6. Initialize DMA channel control registers (DMA_CH0_Control through DMA_CH3_Control) to 0x0.\n7. Configure MTL_TxQ0_Operation_Mode with 4096-byte TX queue size, store-and-forward enabled, TX Q0 enabled.\n8. Set DMA TX descriptor list base addresses for all 4 channels.\n9. Write DMA TX descriptor tail pointers for all 4 channels.\n10. Start TX on DMA Channel 0 only by writing DMA_CH0_Tx_Control with 0x00100007.\n11. Trigger ETH VIP sequencer start.\n12. Poll for transfer completion interrupts in a loop for trns_count iterations.\n13. Determine pass/fail based on failure counters.",
        "Meta Impacted Registers": "MAC_Configuration, MAC_Ext_Configuration, MAC_Address0_High, DMA_CH0_Control, DMA_CH1_Control, DMA_CH2_Control, DMA_CH3_Control, DMA_CH0_Tx_Control, DMA_CH0_TxDesc_List_Address, DMA_CH1_TxDesc_List_Address, DMA_CH2_TxDesc_List_Address, DMA_CH3_TxDesc_List_Address, DMA_CH0_TxDesc_Tail_Pointer, DMA_CH1_TxDesc_Tail_Pointer, DMA_CH2_TxDesc_Tail_Pointer, DMA_CH3_TxDesc_Tail_Pointer, MTL_TxQ0_Operation_Mode",
        "Impacted Registers": "MAC_Configuration, MAC_Ext_Configuration, MAC_Address0_High, DMA_CH0_Control, DMA_CH1_Control, DMA_CH2_Control, DMA_CH3_Control, DMA_CH0_Tx_Control, DMA_CH0_TxDesc_List_Address, DMA_CH1_TxDesc_List_Address, DMA_CH2_TxDesc_List_Address, DMA_CH3_TxDesc_List_Address, DMA_CH0_TxDesc_Tail_Pointer, DMA_CH1_TxDesc_Tail_Pointer, DMA_CH2_TxDesc_Tail_Pointer, DMA_CH3_TxDesc_Tail_Pointer, MTL_TxQ0_Operation_Mode",
        "Meta Validation / Acceptance Criteria": "Transfer completion is validated by polling extern int int_pend in a for loop iterating trns_count times (default 10, or 6 when ETH_10M defined). Pass condition: def_fail_cnt == 0 and wr_fail_cnt == 0.",
        "Validation / Acceptance Criteria": "1. All trns_count (10 default, 6 for 10M mode) transmit transfer interrupts must be received and acknowledged via the int_pend flag.\n2. No default value check failures (def_fail_cnt must be 0).\n3. No write/read verification failures (wr_fail_cnt must be 0).\n4. DMA Channel 0 must successfully complete TX operations as indicated by interrupt completion.\n5. Test passes (finish(0)) only if all failure counters remain at zero.",
        "Remarks": "Only DMA Channel 0 TX is actively started; DMA Channels 1-3 TX_CONTROL writes are commented out with value 0x00100007. The test uses conditional compilation for speed selection (ETH_10M: 10Mbps, ETH_100M: 100Mbps, default: 1Gbps) and duplex mode (HALF_DUPLEX or Full Duplex)."
    },
    {
        "Index": "3",
        "SS / Module": "ETHERNET0",
        "Test Case Name": "ethernet0_rx_basic_test",
        "Feature": "Basic RX Reception",
        "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include <math.h>\n#include <ethernet0/ethernet0_def.h>\n#include <test_common.h>\n#include <ethernet0/ethernet0_offset.h>\n#include <ethernet0/ethernet0_funcs.h>\n#include <ethernet0/ethernet0_programming_sequence.h>",
        "Meta Macros": "ETH_10M\nETH_100M\nHALF_DUPLEX\nDEBUG_DISPLAY\nPRELOAD_AGAIN\nSEL_ENET0\nSEL_ENET1\nSEL_ENET2\nSEL_ENET3",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a basic Ethernet0 RX (receive) operation. It initializes the Ethernet interface via enet_intf_sel() and configures MAC, MTL, and DMA layers for receive operations with VLAN tag control, packet filter, RX queue control, and interrupt handling.",
        "Test Description": "Perform a basic Ethernet0 receive operation using DMA Channel 0. The test configures the MAC layer (speed, duplex, MAC addresses for 4 address slots with DMA channel selection, VLAN tag control, packet filter with hash filter and receive-all enabled, RX queue control with DCB/Generic mode and queue priorities), MTL layer (TX queue operation mode with 4096-byte queue and store-and-forward, RX queue to DMA channel mapping, queue interrupt control for overflow/underflow), and DMA layer (descriptor list addresses for TX and RX, channel control, RX control with buffer size, interrupt enables). The test triggers the VIP sequencer and uses an interrupt handler (Default_IRQHandler) to monitor RX descriptor and buffer status across all 4 DMA channels, with optional descriptor preloading when PRELOAD_AGAIN is defined.",
        "Meta Test Steps / Procedure": "1. Include headers\n2. Global declarations\n3. int test_case() entry point\n4. ETH_10M trns_count = 6\n5. Ethernet interface selection\n6. MAC Configuration with speed/duplex\n7. MAC_EXT_CONFIGURATION = 0x0\n8. MAC_VLAN_TAG_CTRL = 0x01600000\n9. MAC_PACKET_FILTER = 0x80000400\n10. MAC_RXQ_CTRL0-4 configured\n11. MAC_ADDRESS HIGH/LOW for 4 slots\n12. MTL_TXQ0_OPERATION_MODE = 0x000f000a\n13. MTL_RXQ_DMA_MAP0 = 0x03020100\n14. MTL Q interrupt control\n15. DMA channel control and descriptors\n16. DMA interrupt enables\n17. DMA RX control\n18. Power and VIP sequencer\n19. Polling loop\n20. Default_IRQHandler with PRELOAD_AGAIN logic",
        "Test Steps / Procedure": "1. Set trns_count to 6 if ETH_10M mode is defined.\n2. Select Ethernet interface using enet_intf_sel() based on SEL_ENETx compile-time macro (default channel 4).\n3. Configure MAC_Configuration register with appropriate speed and duplex settings (10M/100M/1G, Half/Full Duplex).\n4. Write MAC_Ext_Configuration register with 0x0.\n5. Configure MAC_VLAN_Tag_Ctrl with 0x01600000 to enable VLAN TAG in status and VLAN TAG RX stripping.\n6. Configure MAC_Packet_Filter with 0x80000400 to enable hash filter and receive all packets.\n7. Program MAC_RxQ_Ctrl0 with 0x00aa, MAC_RxQ_Ctrl1 with 0x000, MAC_RxQ_Ctrl2 with 0x08040201, MAC_RxQ_Ctrl4 with 0x0.\n8. Program MAC address HIGH registers for 4 address slots with DMA channel selection and address enable.\n9. Program MAC address LOW registers with lower 32 bits of MAC addresses.\n10. Configure MTL_TxQ0_Operation_Mode with 0x000f000a.\n11. Configure MTL_RxQ_DMA_Map0 with 0x03020100.\n12. Enable RX queue overflow and TX queue underflow interrupts.\n13. Initialize DMA_CH0_Control to 0x00000000.\n14. Set DMA TX descriptor list base addresses.\n15. Enable DMA channel interrupts.\n16. Configure DMA_CH0_RX_CONTROL and DMA_CH0_RX_CONTROL2.\n17. Set DMA_CH0_RXDESC_LIST_ADDRESS to 0xE6040000.\n18. Write power configuration register.\n19. Trigger ETH VIP sequencer start.\n20. Poll for transfer completion interrupts via int_pend.\n21. In Default_IRQHandler: read current app RX descriptor and RX buffer for all 4 DMA channels.",
        "Meta Impacted Registers": "MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_VLAN_Tag_Ctrl, MAC_Address0_Low, MAC_Address1_Low, MAC_Address2_Low, MAC_Address3_Low, MAC_RxQ_Ctrl0, MAC_RxQ_Ctrl1, MAC_RxQ_Ctrl2, MAC_RxQ_Ctrl4, MTL_TxQ0_Operation_Mode, MTL_RxQ_DMA_Map0, MTL_Q1_Interrupt_Control_Status, MTL_Q2_Interrupt_Control_Status, MTL_Q3_Interrupt_Control_Status, DMA_CH0_TxDesc_List_Address, DMA_CH1_TxDesc_List_Address, DMA_CH2_TxDesc_List_Address, DMA_CH3_TxDesc_List_Address, DMA_CH1_Interrupt_Enable, DMA_CH2_Interrupt_Enable, DMA_CH3_Interrupt_Enable, DMA_CH0_RxDesc_Tail_Pointer, DMA_CH3_Current_App_RxDesc, DMA_CH3_Current_App_RxBuffer, DMA_Interrupt_Status",
        "Impacted Registers": "MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_VLAN_Tag_Ctrl, MAC_Address0_Low, MAC_Address1_Low, MAC_Address2_Low, MAC_Address3_Low, MAC_RxQ_Ctrl0, MAC_RxQ_Ctrl1, MAC_RxQ_Ctrl2, MAC_RxQ_Ctrl4, MTL_TxQ0_Operation_Mode, MTL_RxQ_DMA_Map0, MTL_Q1_Interrupt_Control_Status, MTL_Q2_Interrupt_Control_Status, MTL_Q3_Interrupt_Control_Status, DMA_CH0_TxDesc_List_Address, DMA_CH1_TxDesc_List_Address, DMA_CH2_TxDesc_List_Address, DMA_CH3_TxDesc_List_Address, DMA_CH1_Interrupt_Enable, DMA_CH2_Interrupt_Enable, DMA_CH3_Interrupt_Enable, DMA_CH0_RxDesc_Tail_Pointer, DMA_CH3_Current_App_RxDesc, DMA_CH3_Current_App_RxBuffer, DMA_Interrupt_Status",
        "Meta Validation / Acceptance Criteria": "In Default_IRQHandler: rdata_tailpointer = read_enet_reg(DMA_CH0_RXDESC_TAIL_POINTER); if(rdata_tailpointer == rdata_desc) and reaches 0xE60400f0, preload_descriptor_rx is called if PRELOAD_AGAIN defined. Pass condition: def_fail_cnt == 0 and wr_fail_cnt == 0.",
        "Validation / Acceptance Criteria": "1. All transfer completion interrupts must be received and acknowledged via the int_pend flag for trns_count iterations (10 default, 6 for 10M mode).\n2. No default value check failures (def_fail_cnt must be 0).\n3. No write/read verification failures (wr_fail_cnt must be 0).\n4. DMA_CH0_RxDesc_Tail_Pointer is monitored in the interrupt handler; when it reaches 0xE60400f0, RX descriptors are optionally preloaded if PRELOAD_AGAIN is defined.\n5. DMA_Interrupt_Status is read in the interrupt handler to verify interrupt source.\n6. Current application RX descriptor and RX buffer addresses are read for all 4 DMA channels to confirm RX progress.",
        "Remarks": "The test configures 4 MAC address slots with different DMA channel selections. MAC_VLAN_Tag_Ctrl is configured with 0x01600000 for VLAN tag stripping. MAC_Packet_Filter uses 0x80000400 (hash filter + receive all). RX queue priorities are set via MAC_RxQ_Ctrl2=0x08040201. MTL RX queue to DMA channel mapping is 1:1 via MTL_RxQ_DMA_Map0=0x03020100. The interrupt handler includes PRELOAD_AGAIN conditional logic to reload RX descriptors when the tail pointer reaches the end of the ring."
    },
    {
        "Index": "4",
        "SS / Module": "ETHERNET0",
        "Test Case Name": "ethernet0_tx_rx_multi_chnl_test",
        "Feature": "Multi-Channel TX/RX Transmission and Reception",
        "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include <math.h>\n#include <ethernet0/ethernet0_def.h>\n#include <test_common.h>\n#include <ethernet0/ethernet0_offset.h>\n#include <ethernet0/ethernet0_funcs.h>\n#include <ethernet0/ethernet0_programming_sequence.h>",
        "Meta Macros": "ETH_10M\nETH_100M\nHALF_DUPLEX\nDEBUG_DISPLAY\nPRELOAD_AGAIN\nSEL_ENET0\nSEL_ENET1\nSEL_ENET2\nSEL_ENET3",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs multi-channel TX and RX operations on Ethernet0 across all 4 DMA channels (CH0-CH3) sequentially. It configures MAC, MTL, and DMA layers for all channels and activates them one by one.",
        "Test Description": "Perform multi-channel Ethernet0 TX and RX operations across all 4 DMA channels (CH0-CH3) sequentially. The test first configures the MAC layer (speed, duplex, MAC addresses with per-channel DMA selection, VLAN tag control, packet filter with hash filter and receive-all), MTL layer (TX queue operation modes for Q0-Q2, RX queue operation modes for Q0-Q3, quantum weights for Q1-Q3, RX queue to DMA channel mapping, queue interrupt control), and DMA layer (channel control, TX/RX control for all 4 channels, descriptor list addresses, tail pointers, interrupt enables, DMA mode). After initial bulk configuration, the test activates each channel sequentially: Channel 0 runs first until 20 TX and 20 RX packets complete, then Channel 1 is activated with updated TX control, RXQ DMA mapping, RX descriptor tail pointer, and RX control until 20 TX/20 RX packets complete, then Channel 2 is activated similarly until 10 TX/10 RX packets complete, then Channel 3 is activated. Each channel activation triggers the VIP sequencer with a unique magic value.",
        "Meta Test Steps / Procedure": "1. Include headers\n2. Global declarations\n3. int test_case() entry point\n4. ETH_10M trns_count = 6\n5. Ethernet interface selection\n6. MAC Configuration\n7. Override enable\n8. MAC_EXT_CONFIGURATION, RXQ_CTRL, PACKET_FILTER, VLAN_TAG_CTRL\n9. MAC addresses HIGH/LOW\n10. MTL operation modes and quantum weights\n11. MTL RXQ DMA MAP and interrupt control\n12. DMA channel control, TX/RX control\n13. DMA descriptor addresses and tail pointers\n14. VIP sequencer\n15. Channel 0 loop (20 TX/20 RX)\n16. Channel 1 activation and loop\n17. Channel 2 activation and loop\n18. Channel 3 activation\n19. Default_IRQHandler\n20. Pass/fail determination",
        "Test Steps / Procedure": "1. Select Ethernet interface using enet_intf_sel() based on SEL_ENETx compile-time macro (default channel 4).\n2. If ETH_10M mode, set trns_count to 6.\n3. Configure MAC_Configuration register with appropriate speed and duplex settings. Default 1G Full Duplex uses 0x00002003.\n4. Enable override by writing 0x1A08 to address (0xE68C205C + (enet_sel * 4)), then wait_on(300).\n5. Write MAC_Ext_Configuration with 0x0.\n6. Program MAC_RxQ_Ctrl0=0x000000aa, MAC_RxQ_Ctrl1=0x00000000, MAC_RxQ_Ctrl2=0x08040201.\n7. Configure MAC_Packet_Filter with 0x80000408 (hash filter + receive all).\n8. Configure MAC_VLAN_Tag_Ctrl with 0x01600000.\n9. Program MAC address HIGH registers for 4 address slots with DMA channel selection.\n10. Program MAC address LOW registers.\n11. Configure MTL_Operation_Mode=0x00000000.\n12. Configure MTL TX queue operation modes.\n13. Configure MTL RX queue operation modes.\n14. Program MTL TX queue quantum weights.\n15. Set initial MTL_RxQ_DMA_Map0=0x00000083.\n16. Enable MTL queue interrupts.\n17. Initialize DMA channel control registers to 0x0.\n18. Bulk configure DMA TX control for all 4 channels with 0x00100006.\n19. Bulk configure DMA RX control for all 4 channels with 0x00101000.\n20. Configure DMA_CH0_Rx_Control2=0x9, DMA_CH0_Interrupt_Enable=0xf0df, DMA_Mode=0x00000000.\n21. Set DMA TX descriptor list addresses.\n22. Set DMA RX descriptor list addresses.\n23. Set DMA TX descriptor tail pointers.\n24. Set DMA RX descriptor tail pointers.\n25. Trigger VIP sequencer.\n26. Channel 0: Loop up to 20 iterations; break when rxpkt==20 && txpkt==20.\n27. Channel 1 activation and loop.\n28. Channel 2 activation and loop.\n29. Channel 3 activation.\n30. Determine pass/fail based on def_fail_cnt and wr_fail_cnt.",
        "Meta Impacted Registers": "MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_VLAN_Tag_Ctrl, MAC_Address0_High, MAC_Address1_High, MAC_Address2_High, MAC_Address3_High, MAC_Address0_Low, MAC_Address1_Low, MAC_Address2_Low, MAC_Address3_Low, MAC_RxQ_Ctrl0, MAC_RxQ_Ctrl1, MAC_RxQ_Ctrl2, DMA_Mode, DMA_CH0_Control, DMA_CH1_Control, DMA_CH2_Control, DMA_CH3_Control, DMA_CH0_Tx_Control, DMA_CH1_Tx_Control, DMA_CH2_Tx_Control, DMA_CH3_Tx_Control, DMA_CH0_Rx_Control, DMA_CH1_Rx_Control, DMA_CH2_Rx_Control, DMA_CH3_Rx_Control, DMA_CH0_Rx_Control2, DMA_CH0_TxDesc_List_Address, DMA_CH1_TxDesc_List_Address, DMA_CH2_TxDesc_List_Address, DMA_CH3_TxDesc_List_Address, DMA_CH0_RxDesc_List_Address, DMA_CH1_RxDesc_List_Address, DMA_CH2_RxDesc_List_Address, DMA_CH3_RxDesc_List_Address, DMA_CH0_TxDesc_Tail_Pointer, DMA_CH1_TxDesc_Tail_Pointer, DMA_CH2_TxDesc_Tail_Pointer, DMA_CH3_TxDesc_Tail_Pointer, DMA_CH0_RxDesc_Tail_Pointer, DMA_CH0_Interrupt_Enable, DMA_CH1_Interrupt_Enable, DMA_CH2_Interrupt_Enable, DMA_CH3_Interrupt_Enable, DMA_CH0_Current_App_RxDesc, DMA_CH3_Current_App_TxDesc, DMA_CH3_Current_App_TxBuffer, MTL_Operation_Mode, MTL_TxQ0_Operation_Mode, MTL_TxQ1_Operation_Mode, MTL_TxQ2_Operation_Mode, MTL_RxQ0_Operation_Mode, MTL_RxQ1_Operation_Mode, MTL_RxQ2_Operation_Mode, MTL_RxQ3_Operation_Mode, MTL_TxQ1_Quantum_Weight, MTL_TxQ2_Quantum_Weight, MTL_TxQ3_Quantum_Weight, MTL_RxQ_DMA_Map0, MTL_Q1_Interrupt_Control_Status, MTL_Q2_Interrupt_Control_Status, MTL_Q3_Interrupt_Control_Status",
        "Impacted Registers": "MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_VLAN_Tag_Ctrl, MAC_Address0_High, MAC_Address1_High, MAC_Address2_High, MAC_Address3_High, MAC_Address0_Low, MAC_Address1_Low, MAC_Address2_Low, MAC_Address3_Low, MAC_RxQ_Ctrl0, MAC_RxQ_Ctrl1, MAC_RxQ_Ctrl2, DMA_Mode, DMA_CH0_Control, DMA_CH1_Control, DMA_CH2_Control, DMA_CH3_Control, DMA_CH0_Tx_Control, DMA_CH1_Tx_Control, DMA_CH2_Tx_Control, DMA_CH3_Tx_Control, DMA_CH0_Rx_Control, DMA_CH1_Rx_Control, DMA_CH2_Rx_Control, DMA_CH3_Rx_Control, DMA_CH0_Rx_Control2, DMA_CH0_TxDesc_List_Address, DMA_CH1_TxDesc_List_Address, DMA_CH2_TxDesc_List_Address, DMA_CH3_TxDesc_List_Address, DMA_CH0_RxDesc_List_Address, DMA_CH1_RxDesc_List_Address, DMA_CH2_RxDesc_List_Address, DMA_CH3_RxDesc_List_Address, DMA_CH0_TxDesc_Tail_Pointer, DMA_CH1_TxDesc_Tail_Pointer, DMA_CH2_TxDesc_Tail_Pointer, DMA_CH3_TxDesc_Tail_Pointer, DMA_CH0_RxDesc_Tail_Pointer, DMA_CH0_Interrupt_Enable, DMA_CH1_Interrupt_Enable, DMA_CH2_Interrupt_Enable, DMA_CH3_Interrupt_Enable, DMA_CH0_Current_App_RxDesc, DMA_CH3_Current_App_TxDesc, DMA_CH3_Current_App_TxBuffer, MTL_Operation_Mode, MTL_TxQ0_Operation_Mode, MTL_TxQ1_Operation_Mode, MTL_TxQ2_Operation_Mode, MTL_RxQ0_Operation_Mode, MTL_RxQ1_Operation_Mode, MTL_RxQ2_Operation_Mode, MTL_RxQ3_Operation_Mode, MTL_TxQ1_Quantum_Weight, MTL_TxQ2_Quantum_Weight, MTL_TxQ3_Quantum_Weight, MTL_RxQ_DMA_Map0, MTL_Q1_Interrupt_Control_Status, MTL_Q2_Interrupt_Control_Status, MTL_Q3_Interrupt_Control_Status",
        "Meta Validation / Acceptance Criteria": "Channel 0: rxpkt==20 && txpkt==20. Channel 1: rxpkt==20 && txpkt==20. Channel 2: rxpkt==10 && txpkt==10. Default_IRQHandler sets int_pend=0. Pass condition: def_fail_cnt == 0 and wr_fail_cnt == 0.",
        "Validation / Acceptance Criteria": "1. Channel 0 must complete 20 TX and 20 RX packets within 20 interrupt-driven iterations.\n2. Channel 1 must complete 20 TX and 20 RX packets within 20 interrupt-driven iterations.\n3. Channel 2 must complete 10 TX and 10 RX packets within 20 interrupt-driven iterations.\n4. Channel 3 RX is activated and must process packets as configured.\n5. The interrupt handler must correctly clear int_pend and update tail pointers.\n6. No default value check failures (def_fail_cnt must be 0).\n7. No write/read verification failures (wr_fail_cnt must be 0).\n8. Test passes (finish(0)) only if all failure counters remain at zero.",
        "Remarks": "The test uses sequential per-channel activation: Channel 0 runs first, then Channel 1, then Channel 2, then Channel 3. Each channel activation updates DMA TX control to 0x00100007 (start TX), updates MTL_RxQ_DMA_Map0 with per-channel mapping, updates RX descriptor tail pointer, and sets RX control to 0x00101001. Each channel uses a unique VIP sequencer magic value. Channel 0 and Channel 1 require 20 TX/20 RX packets; Channel 2 requires 10 TX/10 RX packets."
    }
]

# TestPlan columns
testplan_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

# MetaData columns
metadata_columns = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

def create_workbook():
    wb = Workbook()
    
    # Create TestPlan sheet
    ws_tp = wb.active
    ws_tp.title = "TestPlan"
    
    # Create MetaData sheet
    ws_md = wb.create_sheet("MetaData")
    
    # Formatting
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    wrap_alignment = Alignment(wrap_text=True, vertical='top')
    
    # Write TestPlan headers
    for col_idx, col_name in enumerate(testplan_columns, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    
    # Write TestPlan data
    for row_idx, record in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(testplan_columns, 1):
            value = record.get(col_name, "")
            if value is None:
                value = ""
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=str(value))
            cell.alignment = wrap_alignment
    
    # Write MetaData headers
    for col_idx, col_name in enumerate(metadata_columns, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    
    # Write MetaData data
    for row_idx, record in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(metadata_columns, 1):
            value = record.get(col_name, "")
            if value is None:
                value = ""
            cell = ws_md.cell(row=row_idx, column=col_idx, value=str(value))
            cell.alignment = wrap_alignment
    
    # Auto-size columns with max width cap
    max_width = 80
    for ws in [ws_tp, ws_md]:
        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            col_letter = get_column_letter(col_idx)
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws.max_row):
                for cell in row:
                    if cell.value:
                        lines = str(cell.value).split('\n')
                        for line in lines:
                            if len(line) > max_len:
                                max_len = len(line)
            adjusted_width = min(max_len + 2, max_width)
            if adjusted_width < 12:
                adjusted_width = 12
            ws.column_dimensions[col_letter].width = adjusted_width
    
    # Freeze first row
    ws_tp.freeze_panes = 'A2'
    ws_md.freeze_panes = 'A2'
    
    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = 'veryHidden'
    
    # Save workbook
    wb.save(output_path)
    print(f"Workbook saved: {output_path}")
    
    # Validate
    if os.path.exists(output_path) and os.path.getsize(output_path) > 0:
        wb_check = load_workbook(output_path)
        sheets = wb_check.sheetnames
        if 'TestPlan' in sheets and 'MetaData' in sheets:
            print(f"VALIDATION: PASSED")
            print(f"File size: {os.path.getsize(output_path)} bytes")
            print(f"Sheets: {sheets}")
            print(f"TestPlan rows: {ws_tp.max_row - 1}")
            print(f"MetaData rows: {ws_md.max_row - 1}")
            print(f"Filename: {filename}")
            return True
    print("VALIDATION: FAILED")
    return False

if __name__ == '__main__':
    create_workbook()
