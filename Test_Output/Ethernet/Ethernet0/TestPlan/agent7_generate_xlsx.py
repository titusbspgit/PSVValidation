#!/usr/bin/env python3
"""Agent 7: Generates Ethernet_TestPlan_<IST_timestamp>.xlsx from JSON data files."""
import json, os, sys, datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    os.system(f'{sys.executable} -m pip install openpyxl')
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

script_dir = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(script_dir, 'agent7_testplan_data.json'), 'r') as f:
    TESTPLAN_JSON = json.load(f)
with open(os.path.join(script_dir, 'agent7_metadata.json'), 'r') as f:
    META_JSON = json.load(f)

IP_NAME = 'Ethernet'
now = datetime.datetime.utcnow() + datetime.timedelta(hours=5, minutes=30)
timestamp = now.strftime('%Y%m%d_%H%M%S')
filename = f'{IP_NAME}_TestPlan_{timestamp}.xlsx'

# Validation
tp_names = {tc['Test Case Name'] for tc in TESTPLAN_JSON}
mt_names = {mc['Test Case Name'] for mc in META_JSON}
for n in tp_names:
    assert n in mt_names, f'FAIL: {n} not in META'

wb = Workbook()
hf = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
hfill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
ha = Alignment(horizontal='center', vertical='center', wrap_text=True)
ca = Alignment(vertical='top', wrap_text=True)
bdr = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

# TestPlan sheet
ws = wb.active
ws.title = 'TestPlan'
tp_cols = ['Index','SS / Module','Feature','Test Case Name','Speed','Mode',
           'Memory Start Offset','Memory End Offset','Test Description',
           'Test Steps / Procedure','Impacted Registers',
           'Validation / Acceptance Criteria','Remarks',
           'Code Generation (Required / Not)']
for ci, cn in enumerate(tp_cols, 1):
    c = ws.cell(row=1, column=ci, value=cn)
    c.font = hf; c.fill = hfill; c.alignment = ha; c.border = bdr
data_keys = tp_cols[:-1]  # all except Code Generation
for ri, tc in enumerate(TESTPLAN_JSON, 2):
    for ci, key in enumerate(data_keys, 1):
        c = ws.cell(row=ri, column=ci, value=tc.get(key, ''))
        c.alignment = ca; c.border = bdr
    c = ws.cell(row=ri, column=14, value='')
    c.alignment = ca; c.border = bdr
dv = DataValidation(type='list', formula1='"Required,Not Required"', allow_blank=True)
dv.error = 'Please select Required or Not Required'
dv.errorTitle = 'Invalid Entry'
lr = 1 + len(TESTPLAN_JSON)
dv.add(f'N2:N{lr}')
ws.add_data_validation(dv)
ws.freeze_panes = 'A2'
for col in range(1, 15):
    ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(80, max(12, len(str(ws.cell(row=1, column=col).value or '')) + 4))

# MetaData sheet
wm = wb.create_sheet(title='MetaData')
meta_cols = ['Test Case Name','Meta Headers','Meta Macros','Meta Arrays','Meta Types',
             'Meta Function Inventory','Meta Test Description',
             'Meta Test Steps / Procedure','Meta Impacted Registers',
             'Meta Validation / Acceptance Criteria']
for ci, cn in enumerate(meta_cols, 1):
    c = wm.cell(row=1, column=ci, value=cn)
    c.font = hf; c.fill = hfill; c.alignment = ha; c.border = bdr
for ri, mc in enumerate(META_JSON, 2):
    for ci, key in enumerate(meta_cols, 1):
        c = wm.cell(row=ri, column=ci, value=mc.get(key, ''))
        c.alignment = ca; c.border = bdr
wm.freeze_panes = 'A2'
wm.sheet_state = 'veryHidden'

fpath = os.path.join(script_dir, filename)
wb.save(fpath)
print(f'Generated: {filename}')
print(f'Size: {os.path.getsize(fpath)} bytes')
print(f'Path: {fpath}')
