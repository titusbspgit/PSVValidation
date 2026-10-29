#!/usr/bin/env python3
"""
Build and upload Ethernet1 TestPlan Excel Workbook.

Run this script from any machine with Python 3 + openpyxl + requests:
    pip install openpyxl requests
    python build_xlsx.py

It will:
1. Read testplan_data.json and metadata_data.json from the same directory
2. Generate a formatted .xlsx workbook
3. Upload it to GitHub via the API
"""
import json, os, sys, datetime, base64

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

GITHUB_OWNER = "titusbspgit"
GITHUB_REPO = "PSVValidation"
GITHUB_BRANCH = "main"
GITHUB_DIR = "Test_Output/Ethernet/Ethernet1/TestPlan"

def load_json(filename):
    with open(os.path.join(SCRIPT_DIR, filename), 'r', encoding='utf-8') as f:
        return json.load(f)

def auto_width(ws, max_width=80):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                lines = str(cell.value).split('\n')
                longest = max(len(l) for l in lines)
                if longest > max_len:
                    max_len = longest
        ws.column_dimensions[col_letter].width = min(max_len + 4, max_width)

def generate_workbook():
    tp = load_json('testplan_data.json')
    md = load_json('metadata_data.json')

    # Cross-validate
    tp_names = {r['Test Case Name'] for r in tp}
    md_names = {r['Test Case Name'] for r in md}
    if not tp_names.issubset(md_names):
        print(f"VALIDATION FAILED: TestPlan names not in MetaData: {tp_names - md_names}")
        return None, None, 0, 0

    # IST timestamp
    ist = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
    ts = datetime.datetime.now(ist).strftime('%Y%m%d_%H%M%S')
    filename = f'Ethernet1_TestPlan_{ts}.xlsx'

    wb = Workbook()

    # ---- TestPlan Sheet ----
    ws_tp = wb.active
    ws_tp.title = 'TestPlan'

    tp_cols = [
        'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Speed', 'Mode',
        'Memory Start Offset', 'Memory End Offset', 'Test Description',
        'Test Steps / Procedure', 'Impacted Registers',
        'Validation / Acceptance Criteria', 'Remarks',
        'Code Generation (Required / Not)'
    ]

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    wrap = Alignment(wrap_text=True, vertical='top')

    for c, name in enumerate(tp_cols, 1):
        cell = ws_tp.cell(row=1, column=c, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    for r, rec in enumerate(tp, 2):
        ws_tp.cell(row=r, column=1, value=rec.get('Index', '')).alignment = wrap
        ws_tp.cell(row=r, column=2, value=rec.get('SS / Module', '')).alignment = wrap
        ws_tp.cell(row=r, column=3, value=rec.get('Feature', '')).alignment = wrap
        ws_tp.cell(row=r, column=4, value=rec.get('Test Case Name', '')).alignment = wrap
        ws_tp.cell(row=r, column=5, value=rec.get('Speed', '')).alignment = wrap
        ws_tp.cell(row=r, column=6, value=rec.get('Mode', '')).alignment = wrap
        ws_tp.cell(row=r, column=7, value=rec.get('Memory Start Offset', '')).alignment = wrap
        ws_tp.cell(row=r, column=8, value=rec.get('Memory End Offset', '')).alignment = wrap
        ws_tp.cell(row=r, column=9, value=rec.get('Test Description', '')).alignment = wrap
        ws_tp.cell(row=r, column=10, value=rec.get('Test Steps / Procedure', '')).alignment = wrap
        ws_tp.cell(row=r, column=11, value=rec.get('Impacted Registers', '')).alignment = wrap
        ws_tp.cell(row=r, column=12, value=rec.get('Validation / Acceptance Criteria', '')).alignment = wrap
        ws_tp.cell(row=r, column=13, value=rec.get('Remarks', '')).alignment = wrap
        # Column 14 intentionally left blank for user dropdown selection

    # Data Validation dropdown for Code Generation column
    dv = DataValidation(
        type='list',
        formula1='"Required,Not Required"',
        allow_blank=True,
        showDropDown=False
    )
    dv.error = 'Please select Required or Not Required'
    dv.errorTitle = 'Invalid Entry'
    dv.prompt = 'Select Code Generation status'
    dv.promptTitle = 'Code Generation'
    last_row = len(tp) + 1
    dv.add(f'N2:N{last_row}')
    ws_tp.add_data_validation(dv)

    ws_tp.freeze_panes = 'A2'
    auto_width(ws_tp)

    # ---- MetaData Sheet (veryHidden) ----
    ws_md = wb.create_sheet('MetaData')

    md_cols = [
        'Test Case Name', 'Meta Headers', 'Meta Macros', 'Meta Arrays',
        'Meta Types', 'Meta Test Description', 'Meta Test Steps / Procedure',
        'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria'
    ]

    for c, name in enumerate(md_cols, 1):
        cell = ws_md.cell(row=1, column=c, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    for r, rec in enumerate(md, 2):
        ws_md.cell(row=r, column=1, value=rec.get('Test Case Name', '')).alignment = wrap
        ws_md.cell(row=r, column=2, value=rec.get('Meta Headers', '')).alignment = wrap
        ws_md.cell(row=r, column=3, value=rec.get('Meta Macros', '')).alignment = wrap
        ws_md.cell(row=r, column=4, value=rec.get('Meta Arrays', '')).alignment = wrap
        ws_md.cell(row=r, column=5, value=rec.get('Meta Types', '')).alignment = wrap
        ws_md.cell(row=r, column=6, value=rec.get('Meta Test Description', '')).alignment = wrap
        ws_md.cell(row=r, column=7, value=rec.get('Meta Test Steps / Procedure', '')).alignment = wrap
        ws_md.cell(row=r, column=8, value=rec.get('Meta Impacted Registers', '')).alignment = wrap
        ws_md.cell(row=r, column=9, value=rec.get('Meta Validation / Acceptance Criteria', '')).alignment = wrap

    ws_md.sheet_state = 'veryHidden'
    ws_md.freeze_panes = 'A2'
    auto_width(ws_md)

    # Save locally
    out_path = os.path.join(SCRIPT_DIR, filename)
    wb.save(out_path)

    # Validate
    from openpyxl import load_workbook
    wb2 = load_workbook(out_path)
    assert 'TestPlan' in wb2.sheetnames
    assert 'MetaData' in wb2.sheetnames
    assert wb2['TestPlan'].max_row - 1 == len(tp)
    assert wb2['MetaData'].max_row - 1 == len(md)
    wb2.close()

    print(f"Workbook generated: {filename}")
    print(f"TestPlan rows: {len(tp)}")
    print(f"MetaData rows: {len(md)}")
    print("Post-save validation: PASSED")

    return out_path, filename, len(tp), len(md)

def upload_to_github(filepath, filename):
    try:
        import requests
    except ImportError:
        print("WARNING: requests not installed. Skipping GitHub upload.")
        return None, None

    token = os.environ.get('GITHUB_TOKEN')
    if not token:
        print("WARNING: GITHUB_TOKEN not set. Skipping GitHub upload.")
        return None, None

    with open(filepath, 'rb') as f:
        content_b64 = base64.b64encode(f.read()).decode()

    api_url = f"https://api.github.com/repos/{GITHUB_OWNER}/{GITHUB_REPO}/contents/{GITHUB_DIR}/{filename}"
    headers = {
        "Authorization": f"token {token}",
        "Accept": "application/vnd.github.v3+json"
    }
    payload = {
        "message": "Generated TestPlan Excel Workbook",
        "content": content_b64,
        "branch": GITHUB_BRANCH
    }

    resp = requests.put(api_url, json=payload, headers=headers)
    if resp.status_code in (200, 201):
        data = resp.json()
        html_url = data['content']['html_url']
        sha = data['commit']['sha']
        print(f"Uploaded to GitHub: {html_url}")
        print(f"Commit SHA: {sha}")
        return html_url, sha
    else:
        print(f"GitHub upload failed: {resp.status_code} {resp.text}")
        return None, None

if __name__ == '__main__':
    result = generate_workbook()
    if result[0]:
        out_path, filename, tp_rows, md_rows = result
        gh_url, commit_sha = upload_to_github(out_path, filename)
        print(json.dumps({
            "status": "SUCCESS" if gh_url else "IN_PROGRESS",
            "output_file_path": out_path,
            "github_url": gh_url or "",
            "filename": filename,
            "rows_testplan": tp_rows,
            "rows_metadata": md_rows,
            "commit_sha": commit_sha or "",
            "validation": "PASSED"
        }, indent=2))
    else:
        print(json.dumps({
            "status": "FAILURE",
            "output_file_path": "",
            "github_url": "",
            "filename": "",
            "rows_testplan": 0,
            "rows_metadata": 0,
            "commit_sha": "",
            "validation": "FAILED"
        }, indent=2))
