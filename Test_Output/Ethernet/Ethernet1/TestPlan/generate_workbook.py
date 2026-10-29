#!/usr/bin/env python3
"""Generate Ethernet1 TestPlan Excel Workbook.
Run: pip install openpyxl && python generate_workbook.py
"""
import json, os
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
IP_NAME = "Ethernet1"
FILENAME = f"{IP_NAME}_TestPlan_{now_ist.strftime('%Y%m%d_%H%M%S')}.xlsx"

script_dir = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(script_dir, 'testplan_data.json'), 'r', encoding='utf-8') as f:
    testplan_data = json.load(f)
with open(os.path.join(script_dir, 'metadata_data.json'), 'r', encoding='utf-8') as f:
    meta_data = json.load(f)

TP_COLS = ['Index','SS / Module','Feature','Test Case Name','Speed','Mode',
           'Memory Start Offset','Memory End Offset','Test Description',
           'Test Steps / Procedure','Impacted Registers',
           'Validation / Acceptance Criteria','Remarks',
           'Code Generation (Required / Not)']
MD_COLS = ['Test Case Name','Meta Headers','Meta Macros','Meta Arrays',
           'Meta Types','Meta Test Description','Meta Test Steps / Procedure',
           'Meta Impacted Registers','Meta Validation / Acceptance Criteria']

header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = 'TestPlan'
for ci, col_name in enumerate(TP_COLS, 1):
    cell = ws_tp.cell(row=1, column=ci, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap
for ri, row_data in enumerate(testplan_data, 2):
    for ci, col_name in enumerate(TP_COLS, 1):
        if col_name == 'Code Generation (Required / Not)':
            ws_tp.cell(row=ri, column=ci, value='').alignment = wrap
        else:
            val = row_data.get(col_name, '')
            ws_tp.cell(row=ri, column=ci, value=val).alignment = wrap

dv = DataValidation(type='list', formula1='"Required,Not Required"', allow_blank=True)
dv.error = 'Please select Required or Not Required'
dv.errorTitle = 'Invalid Entry'
dv.prompt = 'Select Code Generation status'
dv.promptTitle = 'Code Generation'
cg_col = len(TP_COLS)
for ri in range(2, len(testplan_data) + 2):
    dv.add(ws_tp.cell(row=ri, column=cg_col))
ws_tp.add_data_validation(dv)
ws_tp.freeze_panes = 'A2'
for ci in range(1, len(TP_COLS) + 1):
    max_len = len(str(ws_tp.cell(row=1, column=ci).value))
    for ri in range(2, len(testplan_data) + 2):
        val = str(ws_tp.cell(row=ri, column=ci).value or '')
        lines = val.split('\n')
        max_line = max((len(l) for l in lines), default=0)
        if max_line > max_len:
            max_len = max_line
    width = min(max_len + 4, 80)
    ws_tp.column_dimensions[get_column_letter(ci)].width = width

# --- MetaData Sheet ---
ws_md = wb.create_sheet('MetaData')
ws_md.sheet_state = 'veryHidden'
for ci, col_name in enumerate(MD_COLS, 1):
    cell = ws_md.cell(row=1, column=ci, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap
for ri, row_data in enumerate(meta_data, 2):
    for ci, col_name in enumerate(MD_COLS, 1):
        val = row_data.get(col_name, '')
        ws_md.cell(row=ri, column=ci, value=val).alignment = wrap
for ci in range(1, len(MD_COLS) + 1):
    max_len = len(str(ws_md.cell(row=1, column=ci).value))
    for ri in range(2, len(meta_data) + 2):
        val = str(ws_md.cell(row=ri, column=ci).value or '')
        lines = val.split('\n')
        max_line = max((len(l) for l in lines), default=0)
        if max_line > max_len:
            max_len = max_line
    width = min(max_len + 4, 80)
    ws_md.column_dimensions[get_column_letter(ci)].width = width
ws_md.freeze_panes = 'A2'

outpath = os.path.join(script_dir, FILENAME)
wb.save(outpath)
print(f'Workbook saved: {outpath}')
print(f'TestPlan rows: {len(testplan_data)}')
print(f'MetaData rows: {len(meta_data)}')

# Validate
from openpyxl import load_workbook
wb2 = load_workbook(outpath)
assert 'TestPlan' in wb2.sheetnames
assert 'MetaData' in wb2.sheetnames
assert wb2['MetaData'].sheet_state == 'veryHidden'
assert wb2['TestPlan'].max_row - 1 == len(testplan_data)
assert wb2['MetaData'].max_row - 1 == len(meta_data)
print('Validation PASSED')
