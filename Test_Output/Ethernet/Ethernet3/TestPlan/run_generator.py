#!/usr/bin/env python3
"""
Standalone generator: creates the Ethernet3 TestPlan XLSX and prints its
base64 representation so it can be captured and committed separately.
"""
import openpyxl, base64, io, json, sys
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
ts = now_ist.strftime('%Y%m%d_%H%M%S')
fname = f'Ethernet3_TestPlan_{ts}.xlsx'

wb = openpyxl.Workbook()
ws_tp = wb.active
ws_tp.title = 'TestPlan'
ws_md = wb.create_sheet('MetaData')

tp_cols = ['Index','SS / Module','Feature','Test Case Name','Test Description',
           'Speed','Mode','Memory Start Offset','Memory End Offset','Remarks',
           'Test Steps / Procedure','Impacted Registers',
           'Validation / Acceptance Criteria','Code Generation']

md_cols = ['Index','Test Case Name','Meta Test Description',
           'Meta Test Steps / Procedure','Meta Impacted Registers',
           'Meta Validation / Acceptance Criteria','Meta Headers',
           'Meta Macros','Meta Arrays']

d = {
 'Index':'1',
 'SS / Module':'Ethernet3',
 'Test Case Name':'ethernet3_reg_wr_rd_test',
 'Feature':'Register Read/Write Verification',
 'Speed':'NA','Mode':'NA',
 'Memory Start Offset':'0x0','Memory End Offset':'0x14',
 'Test Description':'This test verifies the default (reset) values and write-read accessibility of Ethernet3 MAC registers. First, each register is read and its value is compared against the expected default value, respecting read masks and skip conditions. Then, six distinct data patterns are written to each writable register and read back. The read-back value is validated against an expected value computed using the write mask, read mask, and default value to account for read-only and write-only bit fields. The tested registers are MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_WD_JB_Timeout, MAC_Hash_Table_Reg0, and MAC_Hash_Table_Reg1. The test passes only if all default value checks and all write-read checks succeed with zero mismatches.',
 'Test Steps / Procedure':'1. Read each Ethernet3 MAC register and verify that the value matches the expected default (reset) value, skipping registers that are not readable or are marked to be skipped.\n2. For each of six test data patterns (all-ones, alternating 0xAA, alternating 0x55, all-zeros, 0xA5A5A5A5, 0xFFFF0000), write the pattern to each writable register, skipping registers that are not writable or are marked to be skipped.\n3. After writing each pattern, read back each register and compute the expected value using the write mask, read mask, and default value to account for read-only and write-only bit fields.\n4. Compare the read-back value against the computed expected value for each register and each pattern.\n5. Track all default-value mismatches and write-read mismatches using separate failure counters.\n6. If any mismatch is detected across all checks, report the test as failed; otherwise report the test as passed.',
 'Impacted Registers':'MAC_Configuration; MAC_Ext_Configuration; MAC_Packet_Filter; MAC_WD_JB_Timeout; MAC_Hash_Table_Reg0; MAC_Hash_Table_Reg1',
 'Validation / Acceptance Criteria':'1. All Ethernet3 MAC registers must return their expected default values when read after reset.\n2. For each of the six write test patterns, the read-back value from each register must match the expected value computed by applying the write mask and read mask to the written data and the default value.\n3. The test passes (finish with success) only when both the default-value failure count and the write-read failure count are zero.\n4. Any single mismatch in either the default value check or the write-read check causes the test to fail.',
 'Remarks':'The soft_reset_chk() function is commented out in test_case() and is not executed during this test. The addr_array is declared with size 434 but only 6 register entries are populated. Six distinct bit patterns are used to exercise various bit combinations across all writable registers. Skip arrays (skip_array and skip_rst_array) are all set to 0, meaning no registers are skipped in this configuration. The expected read-back value computation accounts for read-only bits retaining their default values and write-only bits not being readable.',
 'Code Generation':'',
 'Meta Headers':'<stdio.h>; <stdlib.h>; "test_define.c"; <test_common.h>; <ethernet3/ethernet3_def.h>; <ethernet3/ethernet3_offset.h>',
 'Meta Macros':'SOFT_RST_REG_ADDRESS; SOFT_RST_REG_DATA; CNT; DEBUG_DISPLAY',
 'Meta Arrays':'addr_array[434]; default_value_array[434]; read_mask_array[434]; write_mask_array[434]; skip_array[434]; skip_rst_array[434]; chk_val[6]',
 'Meta Test Description':'This testcase performs register default value verification and write-read verification for a set of Ethernet3 MAC registers defined in addr_array. The test has three phases: (1) chk_rst_val() reads each register via read_reg(addr) using addresses from addr_array and compares the read value against default_value_array entries, skipping registers with read_mask_array == 0x00000000 or skip_rst_array == 1. (2) chk_rd_wr() iterates over 6 test patterns in chk_val[6] = {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}. For each pattern, it writes data_wr to each register via write_reg(addr, data_wr), skipping registers with skip_array == 1 or write_mask_array == 0x00000000. It then reads back each register via read_reg(addr), computes the expected value as exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])) where wr_n = (write_mask_array[i] ^ 0xffffffff), and compares data_rd against exp_val. (3) soft_reset_chk() is commented out in test_case(). Failure counters def_fail_cnt and wr_fail_cnt track mismatches. The test calls finish(1) if any failures occur, otherwise finish(0). The registers tested are: mizar_ETHERNET3_MAC_CONFIGURATION, mizar_ETHERNET3_MAC_EXT_CONFIGURATION, mizar_ETHERNET3_MAC_PACKET_FILTER, mizar_ETHERNET3_MAC_WD_JB_TIMEOUT, mizar_ETHERNET3_MAC_HASH_TABLE_REG0, mizar_ETHERNET3_MAC_HASH_TABLE_REG1.',
 'Meta Test Steps / Procedure':'1. test_case() calls chk_rst_val(). 2. chk_rst_val(): For i=0 to CNT-1, addr = addr_array[i]. If read_mask_array[i] == 0x00000000, skip (not readable). If skip_rst_array[i] == 1, skip. Otherwise call data_rd = read_reg(addr). Compare data_rd == default_value_array[i]. If mismatch, increment def_fail_cnt and print failure. 3. test_case() calls chk_rd_wr(). 4. chk_rd_wr(): Define chk_val[6] = {0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}. For j=0 to 5, set data_wr = chk_val[j]. 5. Write phase: For i=0 to CNT-1, addr = addr_array[i]. If skip_array[i] == 1, skip. If write_mask_array[i] == 0x00000000, skip (not writable). Otherwise call write_reg(addr, data_wr). 6. Read-back phase: For i=0 to CNT-1, addr = addr_array[i]. If skip_array[i] == 1, skip. If write_mask_array[i] == 0x00000000, skip. If read_mask_array[i] == 0x00000000, skip. Otherwise call data_rd = read_reg(addr). Compute wr_n = (write_mask_array[i] ^ 0xffffffff). Compute exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])). Compare data_rd == exp_val. If mismatch, increment wr_fail_cnt and print failure. 7. After both phases, test_case() checks if def_fail_cnt > 0 or wr_fail_cnt > 0. If true, call finish(1) (fail). Otherwise call finish(0) (pass). 8. soft_reset_chk() is commented out and not executed.',
 'Meta Impacted Registers':'mizar_ETHERNET3_MAC_CONFIGURATION; mizar_ETHERNET3_MAC_EXT_CONFIGURATION; mizar_ETHERNET3_MAC_PACKET_FILTER; mizar_ETHERNET3_MAC_WD_JB_TIMEOUT; mizar_ETHERNET3_MAC_HASH_TABLE_REG0; mizar_ETHERNET3_MAC_HASH_TABLE_REG1',
 'Meta Validation / Acceptance Criteria':'In chk_rst_val(): For each register in addr_array (where read_mask_array[i] != 0x00000000 and skip_rst_array[i] != 1), data_rd = read_reg(addr) must equal default_value_array[i]. Any mismatch increments def_fail_cnt. In chk_rd_wr(): For each of 6 patterns in chk_val[] and each register in addr_array (where skip_array[i] != 1, write_mask_array[i] != 0x00000000, and read_mask_array[i] != 0x00000000), data_rd = read_reg(addr) must equal exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | ((write_mask_array[i] ^ 0xffffffff) & read_mask_array[i] & default_value_array[i])). Any mismatch increments wr_fail_cnt. Final pass condition: def_fail_cnt == 0 AND wr_fail_cnt == 0 results in finish(0). Otherwise finish(1).'
}

hf = Font(bold=True, color='FFFFFF', size=11)
hfill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wa = Alignment(wrap_text=True, vertical='top')

def pop(ws, cols, row):
    for ci, cn in enumerate(cols, 1):
        c = ws.cell(row=1, column=ci, value=cn)
        c.font = hf; c.fill = hfill; c.alignment = wa
    for ci, cn in enumerate(cols, 1):
        v = row.get(cn, '')
        c = ws.cell(row=2, column=ci, value=v)
        c.alignment = wa
    for ci in range(1, len(cols)+1):
        ml = len(str(cols[ci-1]))
        for r in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=ci, max_col=ci):
            for c in r:
                if c.value: ml = max(ml, min(len(str(c.value)), 80))
        ws.column_dimensions[get_column_letter(ci)].width = min(ml+4, 60)
    ws.freeze_panes = 'A2'

pop(ws_tp, tp_cols, d)
pop(ws_md, md_cols, d)
ws_md.sheet_state = 'veryHidden'

buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
raw = buf.read()
print(json.dumps({"filename": fname, "size": len(raw), "b64": base64.b64encode(raw).decode('ascii')}))
