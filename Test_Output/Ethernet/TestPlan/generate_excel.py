#!/usr/bin/env python3
"""Generate Ethernet TestPlan Excel workbook using openpyxl."""
import json
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime('%Y%m%d_%H%M%S')
filename = f'Ethernet_TestPlan_{timestamp}.xlsx'

# JSON data
json_data = [
    {
        "Index": "1",
        "SS / Module": "Ethernet",
        "Test Case Name": "ethernet0_reg_wr_rd_test",
        "Feature": "Register Write-Read Verification",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_define.c\"; <test_common.h>; <ethernet0/ethernet0_def.h>; <ethernet0/ethernet0_offset.h>",
        "Meta Macros": "CNT",
        "Meta Arrays": "addr_array[434]; default_value_array[434]; read_mask_array[434]; write_mask_array[434]; skip_array[434]; skip_rst_array[434]; chk_val[6]",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs register default value verification and write-read verification on a set of Ethernet MAC registers defined in addr_array. The test_case() function first calls chk_rst_val() which iterates over addr_array (CNT=434 entries, 5 populated), reads each register using read_reg(addr), and compares the read value against the corresponding default_value_array entry. Registers with read_mask_array[i]==0x00000000 or skip_rst_array[i]==1 are skipped. Then chk_rd_wr() is called, which iterates over 6 test data patterns in chk_val[6]={0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}. For each pattern, it writes the pattern to each register using write_reg(addr, data_wr), skipping registers where skip_array[i]==1 or write_mask_array[i]==0x00000000. It then reads back each register using read_reg(addr), computes the expected value as exp_val=((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])) where wr_n=(write_mask_array[i] ^ 0xffffffff), and compares data_rd against exp_val. Mismatches increment def_fail_cnt or wr_fail_cnt. The test passes (finish(0)) if both fail counts are zero, otherwise fails (finish(1)). The registers tested are: mizar_ETHERNET0_MAC_CONFIGURATION, mizar_ETHERNET0_MAC_EXT_CONFIGURATION, mizar_ETHERNET0_MAC_PACKET_FILTER, mizar_ETHERNET0_MAC_WD_JB_TIMEOUT, mizar_ETHERNET0_MAC_HASH_TABLE_REG0.",
        "Test Description": "This test verifies the default (reset) values and write-read accessibility of Ethernet MAC registers. It first reads each register and confirms the value matches the expected default. Then it writes six distinct data patterns to each register, reads back the values, and verifies the read-back matches the expected value computed using the register's read mask, write mask, and default value. The registers under test are MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_WD_JB_Timeout, and MAC_Hash_Table_Reg0. The test passes only if all default value checks and all write-read checks succeed with zero mismatches.",
        "Meta Test Steps / Procedure": "1. Enter test_case(). 2. Call chk_rst_val(): iterate i from 0 to CNT-1 over addr_array. For each entry, set addr=addr_array[i]. If read_mask_array[i]==0x00000000, skip (not readable). If skip_rst_array[i]==1, skip. Otherwise call data_rd=read_reg(addr). Compare data_rd against default_value_array[i]. If mismatch, increment def_fail_cnt and print failure. 3. Call chk_rd_wr(): define chk_val[6]={0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000}. For each j from 0 to 5, set data_wr=chk_val[j]. 4. Write phase: iterate i from 0 to CNT-1. Set addr=addr_array[i]. If skip_array[i]==1, skip. If write_mask_array[i]==0x00000000, skip (not writable). Otherwise call write_reg(addr, data_wr). 5. Read-back phase: iterate i from 0 to CNT-1. Set addr=addr_array[i]. If skip_array[i]==1, skip. If write_mask_array[i]==0x00000000, skip. If read_mask_array[i]==0x00000000, skip. Otherwise call data_rd=read_reg(addr). Compute wr_n=(write_mask_array[i] ^ 0xffffffff). Compute exp_val=((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])). Compare data_rd against exp_val. If mismatch, increment wr_fail_cnt and print failure. 6. After all patterns processed, check if def_fail_cnt>0 or wr_fail_cnt>0. If yes, call finish(1) (fail). Otherwise call finish(0) (pass).",
        "Test Steps / Procedure": "1. Read each Ethernet MAC register and verify the value matches its expected default (reset) value, skipping registers that are not readable.\n2. For each of six distinct test data patterns, write the pattern to each writable register.\n3. After writing each pattern, read back each register and compute the expected value using the register's read mask, write mask, and default value.\n4. Compare the read-back value against the computed expected value for each register.\n5. Track the number of default-value mismatches and write-read mismatches separately.\n6. If all comparisons pass with zero mismatches, report test pass; otherwise report test failure.",
        "Meta Impacted Registers": "mizar_ETHERNET0_MAC_CONFIGURATION; mizar_ETHERNET0_MAC_EXT_CONFIGURATION; mizar_ETHERNET0_MAC_PACKET_FILTER; mizar_ETHERNET0_MAC_WD_JB_TIMEOUT; mizar_ETHERNET0_MAC_HASH_TABLE_REG0",
        "Impacted Registers": "MAC_Configuration; MAC_Ext_Configuration; MAC_Packet_Filter; MAC_WD_JB_Timeout; MAC_Hash_Table_Reg0",
        "Meta Validation / Acceptance Criteria": "In chk_rst_val(): for each register, data_rd=read_reg(addr) must equal default_value_array[i]. If data_rd != default_value_array[i], def_fail_cnt is incremented. In chk_rd_wr(): for each of 6 test patterns (0xffffffff, 0xaaaaaaaa, 0x55555555, 0x00000000, 0xA5A5A5A5, 0xffff0000), after write_reg(addr, data_wr) and data_rd=read_reg(addr), data_rd must equal exp_val=((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])) where wr_n=(write_mask_array[i] ^ 0xffffffff). If data_rd != exp_val, wr_fail_cnt is incremented. Final pass condition: def_fail_cnt==0 AND wr_fail_cnt==0 results in finish(0). Any non-zero fail count results in finish(1).",
        "Validation / Acceptance Criteria": "Each register must return its expected default value when read after reset. After writing each of six test data patterns, the read-back value from each register must match the expected value computed from the write data, read mask, write mask, and default value. The test passes only if all default value checks and all write-read checks produce zero mismatches across all registers (MAC_Configuration, MAC_Ext_Configuration, MAC_Packet_Filter, MAC_WD_JB_Timeout, MAC_Hash_Table_Reg0).",
        "Remarks": "The test uses six hardcoded data patterns for write-read verification. Registers are skipped based on read mask, write mask, and skip arrays. The soft_reset_chk() function is commented out and not executed. The addr_array is declared with 434 entries but only 5 registers are populated."
    },
    {
        "Index": "2",
        "SS / Module": "Ethernet",
        "Test Case Name": "ethernet0_rx_basic_test",
        "Feature": "Basic RX Data Path",
        "Meta Headers": "<stdio.h>; <stdlib.h>; <math.h>; <ethernet0/ethernet0_def.h>; <test_common.h>; <ethernet0/ethernet0_offset.h>; <ethernet0/ethernet0_funcs.h>; <ethernet0/ethernet0_programming_sequence.h>",
        "Meta Macros": "ETH_10M; ETH_100M; SEL_ENET0; SEL_ENET1; SEL_ENET2; SEL_ENET3; HALF_DUPLEX; DEBUG_DISPLAY",
        "Meta Arrays": "NA",
        "Speed": "ETH_10M; ETH_100M",
        "Mode": "HALF_DUPLEX; Full Duplex",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a basic Ethernet RX path validation. The test_case() function sets int_pend=1, initializes GIC via GIC_Set() and GIC_EnableAllIRQ(). Speed is selected via conditional compilation: enet_10m_speed() for ETH_10M, enet_100m_speed() for ETH_100M. Interface selection is done via enet_intf_sel() based on SEL_ENET0/1/2/3 macros. MAC_CONFIGURATION is written with speed and duplex mode settings. MAC RX queue control registers, VLAN tag control, packet filter, MAC addresses, MAC_EXT_CONFIGURATION, MTL TX and RX queue operation modes, quantum weights, MTL operation mode, RXQ_DMA_MAP0, MTL queue interrupt control status, RX queue controls, DMA TX and RX controls, descriptor list addresses, ring lengths, system bus mode, interrupt enables, DMA mode are all configured. RX descriptors are preloaded. RX is started by writing RX descriptor tail pointer and enabling CH0 RX. VIP sequencer is triggered. Test loops waiting for transfer-complete interrupts. IRQ handler reads DMA current app RX descriptor and buffer for CH0-CH3, checks DMA interrupt status, clears DMA channel status registers.",
        "Test Description": "This test validates the basic Ethernet receive data path. It configures the MAC for the selected speed and duplex mode, programs four MAC addresses, sets up RX queue controls and priorities, configures MTL TX and RX queue operation modes with store-and-forward and appropriate queue sizes, maps RX queues to DMA channels, enables MTL queue overflow and underflow interrupts, programs DMA channel TX and RX controls with burst lengths and buffer sizes, sets up TX and RX descriptor list addresses and ring lengths for all four DMA channels, configures the DMA system bus mode, preloads RX descriptors, enables DMA channel interrupts and MAC interrupts, starts the RX path by writing the RX descriptor tail pointer and enabling RX DMA, and triggers an external VIP sequencer. The test then waits for a configurable number of transfer-complete interrupts. In the interrupt handler, DMA current application RX descriptor and buffer pointers are read for all channels, DMA interrupt status is checked, and DMA channel status registers are cleared. The test passes after all expected RX transfers complete successfully.",
        "Meta Test Steps / Procedure": "1. Set int_pend=1, initialize trns_count=10. 2. Call GIC_Set() and GIC_EnableAllIRQ(). 3-31. Configure all MAC, MTL, DMA registers. 32. Loop trns_count times: wait for int_pend, set int_pend=1. 33. wait_on(2000). 34. finish(0). IRQ Handler: 35-41. Read/clear DMA status.",
        "Test Steps / Procedure": "1. Initialize the GIC and enable all interrupts.\n2. Select the Ethernet speed and interface.\n3. Configure MAC, MTL TX/RX queues, DMA channels for all four channels.\n4. Preload RX descriptors.\n5. Enable interrupts and start RX path.\n6. Trigger VIP sequencer.\n7. Wait for expected transfer-complete interrupts.\n8. In IRQ handler, read DMA pointers and clear status.\n9. Verify test passes after all transfers.",
        "Meta Impacted Registers": "mizar_ETHERNET0_MAC_CONFIGURATION; mizar_ETHERNET0_MAC_RXQ_CTRL0; mizar_ETHERNET0_MAC_RXQ_CTRL1; mizar_ETHERNET0_MAC_RXQ_CTRL2; mizar_ETHERNET0_MAC_RXQ_CTRL4; mizar_ETHERNET0_MAC_VLAN_TAG_CTRL; mizar_ETHERNET0_MAC_PACKET_FILTER; mizar_ETHERNET0_MAC_ADDRESS3_HIGH; mizar_ETHERNET0_MAC_ADDRESS2_HIGH; mizar_ETHERNET0_MAC_ADDRESS1_HIGH; mizar_ETHERNET0_MAC_ADDRESS0_HIGH; mizar_ETHERNET0_MAC_ADDRESS3_LOW; mizar_ETHERNET0_MAC_ADDRESS2_LOW; mizar_ETHERNET0_MAC_ADDRESS1_LOW; mizar_ETHERNET0_MAC_ADDRESS0_LOW; mizar_ETHERNET0_MAC_EXT_CONFIGURATION; mizar_ETHERNET0_MTL_TXQ3_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ2_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ1_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ0_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ3_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_TXQ2_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_TXQ1_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_OPERATION_MODE; mizar_ETHERNET0_MTL_RXQ3_OPERATION_MODE; mizar_ETHERNET0_MTL_RXQ2_OPERATION_MODE; mizar_ETHERNET0_MTL_RXQ1_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ0_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_RXQ0_OPERATION_MODE; mizar_ETHERNET0_MTL_RXQ_DMA_MAP0; mizar_ETHERNET0_MTL_Q3_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_Q2_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_Q1_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_RXQ3_CONTROL; mizar_ETHERNET0_MTL_RXQ2_CONTROL; mizar_ETHERNET0_MTL_RXQ1_CONTROL; mizar_ETHERNET0_DMA_CH3_TX_CONTROL; mizar_ETHERNET0_DMA_CH2_TX_CONTROL; mizar_ETHERNET0_DMA_CH1_TX_CONTROL; mizar_ETHERNET0_DMA_CH3_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH2_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH1_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH3_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH2_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH1_TXDESC_RING_LENGTH; mizar_ETHERNET0_MTL_Q0_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_DMA_CH3_CONTROL; mizar_ETHERNET0_DMA_CH2_CONTROL; mizar_ETHERNET0_DMA_CH1_CONTROL; mizar_ETHERNET0_MTL_RXQ0_CONTROL; mizar_ETHERNET0_DMA_CH3_RX_CONTROL; mizar_ETHERNET0_DMA_CH2_RX_CONTROL; mizar_ETHERNET0_DMA_CH1_RX_CONTROL; mizar_ETHERNET0_DMA_SYSBUS_MODE; mizar_ETHERNET0_DMA_CH3_RXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH2_RXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH1_RXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH3_RX_CONTROL2; mizar_ETHERNET0_DMA_CH2_RX_CONTROL2; mizar_ETHERNET0_DMA_CH1_RX_CONTROL2; mizar_ETHERNET0_DMA_CH0_TX_CONTROL; mizar_ETHERNET0_DMA_CH0_TXDESC_LIST_ADDRESS; 0XE68C2058; mizar_ETHERNET0_DMA_CH1_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH2_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH3_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH0_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH0_CONTROL; mizar_ETHERNET0_DMA_CH0_RX_CONTROL; mizar_ETHERNET0_DMA_CH0_RXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH0_RX_CONTROL2; mizar_ETHERNET0_DMA_CH0_INTERRUPT_ENABLE; mizar_ETHERNET0_MAC_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_MODE; mizar_ETHERNET0_DMA_CH0_RXDESC_TAIL_POINTER; 0xA0243ffc; 0xA0243ff8; mizar_ETHERNET0_DMA_CH0_CURRENT_APP_RXDESC; mizar_ETHERNET0_DMA_CH0_CURRENT_APP_RXBUFFER; mizar_ETHERNET0_DMA_CH1_CURRENT_APP_RXDESC; mizar_ETHERNET0_DMA_CH1_CURRENT_APP_RXBUFFER; mizar_ETHERNET0_DMA_CH2_CURRENT_APP_RXDESC; mizar_ETHERNET0_DMA_CH2_CURRENT_APP_RXBUFFER; mizar_ETHERNET0_DMA_CH3_CURRENT_APP_RXDESC; mizar_ETHERNET0_DMA_CH3_CURRENT_APP_RXBUFFER; mizar_ETHERNET0_DMA_INTERRUPT_STATUS; mizar_ETHERNET0_DMA_CH0_STATUS; mizar_ETHERNET0_DMA_CH1_STATUS; mizar_ETHERNET0_DMA_CH2_STATUS; mizar_ETHERNET0_DMA_CH3_STATUS; 0XE68C2050",
        "Impacted Registers": "MAC_Configuration; MAC_RxQ_Ctrl0; MAC_RxQ_Ctrl1; MAC_RxQ_Ctrl2; MAC_RxQ_Ctrl4; MAC_VLAN_Tag_Ctrl; MAC_Packet_Filter; MAC_Address3_High; MAC_Address2_High; MAC_Address1_High; MAC_Address0_High; MAC_Address3_Low; MAC_Address2_Low; MAC_Address1_Low; MAC_Address0_Low; MAC_Ext_Configuration; MTL_TxQ3_Operation_Mode; MTL_TxQ2_Operation_Mode; MTL_TxQ1_Operation_Mode; MTL_TxQ0_Operation_Mode; MTL_TxQ3_Quantum_Weight; MTL_TxQ2_Quantum_Weight; MTL_TxQ1_Quantum_Weight; MTL_Operation_Mode; MTL_RxQ3_Operation_Mode; MTL_RxQ2_Operation_Mode; MTL_RxQ1_Operation_Mode; MTL_TxQ0_Quantum_Weight; MTL_RxQ0_Operation_Mode; MTL_RxQ_DMA_Map0; MTL_Q3_Interrupt_Control_Status; MTL_Q2_Interrupt_Control_Status; MTL_Q1_Interrupt_Control_Status; MTL_RxQ3_Control; MTL_RxQ2_Control; MTL_RxQ1_Control; DMA_CH3_Tx_Control; DMA_CH2_Tx_Control; DMA_CH1_Tx_Control; DMA_CH3_TxDesc_List_Address; DMA_CH2_TxDesc_List_Address; DMA_CH1_TxDesc_List_Address; DMA_CH3_TxDesc_Ring_Length; DMA_CH2_TxDesc_Ring_Length; DMA_CH1_TxDesc_Ring_Length; MTL_Q0_Interrupt_Control_Status; DMA_CH3_Control; DMA_CH2_Control; DMA_CH1_Control; MTL_RxQ0_Control; DMA_CH3_Rx_Control; DMA_CH2_Rx_Control; DMA_CH1_Rx_Control; DMA_SysBus_Mode; DMA_CH3_RxDesc_List_Address; DMA_CH2_RxDesc_List_Address; DMA_CH1_RxDesc_List_Address; DMA_CH3_Rx_Control2; DMA_CH2_Rx_Control2; DMA_CH1_Rx_Control2; DMA_CH0_Tx_Control; DMA_CH0_TxDesc_List_Address; DMA_CH1_Interrupt_Enable; DMA_CH2_Interrupt_Enable; DMA_CH3_Interrupt_Enable; DMA_CH0_TxDesc_Ring_Length; DMA_CH0_Control; DMA_CH0_Rx_Control; DMA_CH0_RxDesc_List_Address; DMA_CH0_Rx_Control2; DMA_CH0_Interrupt_Enable; MAC_Interrupt_Enable; DMA_Mode; DMA_CH0_RxDesc_Tail_Pointer; DMA_CH0_Current_App_RxDesc; DMA_CH0_Current_App_RxBuffer; DMA_CH1_Current_App_RxDesc; DMA_CH1_Current_App_RxBuffer; DMA_CH2_Current_App_RxDesc; DMA_CH2_Current_App_RxBuffer; DMA_CH3_Current_App_RxDesc; DMA_CH3_Current_App_RxBuffer; DMA_Interrupt_Status; DMA_CH0_Status; DMA_CH1_Status; DMA_CH2_Status; DMA_CH3_Status",
        "Meta Validation / Acceptance Criteria": "The test loops trns_count times (10 by default, 6 for ETH_10M). In each iteration, it waits for int_pend to become 0, which is set in Default_IRQHandler(). The IRQ handler reads DMA current app RX descriptor and buffer registers for CH0-CH3, reads DMA_INTERRUPT_STATUS, clears DMA channel status registers for CH0-CH3 by writing 0xffffffff, reads DMA_INTERRUPT_STATUS again. After all trns_count interrupts are serviced, wait_on(2000) is called and finish(0) is called indicating test pass.",
        "Validation / Acceptance Criteria": "The test passes when all expected transfer-complete interrupts (10 for default speed, 6 for 10M speed) are received and serviced. In each interrupt, the DMA current application RX descriptor and buffer pointers for all four channels must be readable. The DMA interrupt status must indicate a valid interrupt source. After clearing all DMA channel status registers, the DMA interrupt status must reflect the cleared state. The test completes successfully by calling finish with a pass indication after all RX transfers are processed.",
        "Remarks": "The test supports conditional compilation for three speed modes (10M, 100M, 1G) and two duplex modes (half duplex, full duplex). The transfer count is reduced to 6 for 10M speed. Four hardcoded hex addresses are used for external system register writes and VIP sequencer triggering. All four DMA channels and all four MTL queues are configured."
    },
    {
        "Index": "3",
        "SS / Module": "Ethernet",
        "Test Case Name": "ethernet0_tx_basic_test",
        "Feature": "Basic TX Data Path",
        "Meta Headers": "<stdio.h>; <stdlib.h>; <math.h>; <test_common.h>; <ethernet0/ethernet0_def.h>; <ethernet0/ethernet0_offset.h>; <ethernet0/ethernet0_funcs.h>; <ethernet0/ethernet0_programming_sequence.h>",
        "Meta Macros": "ETH_10M; ETH_100M; SEL_ENET0; SEL_ENET1; SEL_ENET2; SEL_ENET3; HALF_DUPLEX; DEBUG_DISPLAY; CLRIRQ_42",
        "Meta Arrays": "NA",
        "Speed": "ETH_10M; ETH_100M",
        "Mode": "HALF_DUPLEX; Full Duplex",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a basic Ethernet TX path validation. It configures MAC addresses, packet filter, MAC configuration with speed/duplex settings, triggers VIP sequencer, configures MTL TX queue operation modes, quantum weights, MTL queue interrupt control, DMA system bus mode, preloads TX descriptors, programs TX descriptor ring lengths, list addresses, tail pointers for all 4 DMA channels, enables DMA interrupts, and starts TX on CH0. Loops waiting for transfer-complete interrupts. IRQ handler reads DMA interrupt status, clears CH0 status.",
        "Test Description": "This test validates the basic Ethernet transmit data path using DMA channel 0. It configures the MAC for the selected speed and duplex mode, programs four MAC addresses with address enable and DMA channel selection, sets the packet filter for hash filtering and receive-all mode, configures MTL TX queue operation modes for all four queues with store-and-forward and 4096-byte queue sizes, programs TX queue quantum weights for scheduling, enables MTL queue overflow and underflow interrupts, configures the DMA system bus mode, preloads TX descriptors into memory, programs TX descriptor ring lengths, list addresses, and tail pointers for all four DMA channels, enables DMA channel interrupts, and starts TX on DMA channel 0. The test triggers an external VIP sequencer and then waits for a configurable number of transfer-complete interrupts.",
        "Meta Test Steps / Procedure": "1-26. Configure all MAC, MTL, DMA registers. 27. Start CH0 TX. 28. Loop trns_count times. 29. finish(0). IRQ Handler: Read/clear DMA status.",
        "Test Steps / Procedure": "1. Initialize the GIC and enable all interrupts.\n2. Select speed and interface.\n3. Program MAC addresses.\n4. Configure packet filter and MAC Configuration.\n5. Trigger VIP sequencer.\n6. Configure MTL TX queues and DMA channels.\n7. Preload TX descriptors.\n8. Enable interrupts and start TX on CH0.\n9. Wait for transfer-complete interrupts.\n10. In IRQ handler, read DMA status and clear CH0 status.\n11. Verify test passes.",
        "Meta Impacted Registers": "mizar_ETHERNET0_MAC_ADDRESS3_HIGH; mizar_ETHERNET0_MAC_ADDRESS2_HIGH; mizar_ETHERNET0_MAC_ADDRESS1_HIGH; mizar_ETHERNET0_MAC_ADDRESS0_HIGH; mizar_ETHERNET0_MAC_ADDRESS3_LOW; mizar_ETHERNET0_MAC_ADDRESS2_LOW; mizar_ETHERNET0_MAC_ADDRESS1_LOW; mizar_ETHERNET0_MAC_ADDRESS0_LOW; mizar_ETHERNET0_MAC_PACKET_FILTER; mizar_ETHERNET0_MAC_CONFIGURATION; 0xA0243ffc; 0xA0243ff8; mizar_ETHERNET0_MAC_EXT_CONFIGURATION; mizar_ETHERNET0_MTL_TXQ3_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ2_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ1_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ0_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ3_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_TXQ2_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_TXQ1_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_TXQ0_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_Q3_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_Q2_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_Q1_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_Q0_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_DMA_SYSBUS_MODE; mizar_ETHERNET0_DMA_CH3_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH2_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH1_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH0_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH3_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH2_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH1_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH0_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH3_TXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH2_TXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH1_TXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH0_TXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH3_CONTROL; mizar_ETHERNET0_DMA_CH2_CONTROL; mizar_ETHERNET0_DMA_CH1_CONTROL; mizar_ETHERNET0_DMA_CH0_CONTROL; mizar_ETHERNET0_DMA_CH1_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH2_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH3_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH0_INTERRUPT_ENABLE; 0XE68C2058; mizar_ETHERNET0_DMA_CH0_TX_CONTROL; mizar_ETHERNET0_DMA_INTERRUPT_STATUS; mizar_ETHERNET0_DMA_CH0_STATUS; 0XE68C2050",
        "Impacted Registers": "MAC_Address3_High; MAC_Address2_High; MAC_Address1_High; MAC_Address0_High; MAC_Address3_Low; MAC_Address2_Low; MAC_Address1_Low; MAC_Address0_Low; MAC_Packet_Filter; MAC_Configuration; MAC_Ext_Configuration; MTL_TxQ3_Operation_Mode; MTL_TxQ2_Operation_Mode; MTL_TxQ1_Operation_Mode; MTL_TxQ0_Operation_Mode; MTL_TxQ3_Quantum_Weight; MTL_TxQ2_Quantum_Weight; MTL_TxQ1_Quantum_Weight; MTL_TxQ0_Quantum_Weight; MTL_Q3_Interrupt_Control_Status; MTL_Q2_Interrupt_Control_Status; MTL_Q1_Interrupt_Control_Status; MTL_Q0_Interrupt_Control_Status; DMA_SysBus_Mode; DMA_CH3_TxDesc_Ring_Length; DMA_CH2_TxDesc_Ring_Length; DMA_CH1_TxDesc_Ring_Length; DMA_CH0_TxDesc_Ring_Length; DMA_CH3_TxDesc_List_Address; DMA_CH2_TxDesc_List_Address; DMA_CH1_TxDesc_List_Address; DMA_CH0_TxDesc_List_Address; DMA_CH3_TxDesc_Tail_Pointer; DMA_CH2_TxDesc_Tail_Pointer; DMA_CH1_TxDesc_Tail_Pointer; DMA_CH0_TxDesc_Tail_Pointer; DMA_CH3_Control; DMA_CH2_Control; DMA_CH1_Control; DMA_CH0_Control; DMA_CH1_Interrupt_Enable; DMA_CH2_Interrupt_Enable; DMA_CH3_Interrupt_Enable; DMA_CH0_Interrupt_Enable; DMA_CH0_Tx_Control; DMA_Interrupt_Status; DMA_CH0_Status",
        "Meta Validation / Acceptance Criteria": "The test loops trns_count times. IRQ handler reads DMA_INTERRUPT_STATUS, writes 0xffffffff to DMA_CH0_STATUS, reads DMA_INTERRUPT_STATUS again. After all interrupts, finish(0) is called.",
        "Validation / Acceptance Criteria": "The test passes when all expected transfer-complete interrupts are received and serviced. DMA Interrupt Status must indicate valid source. After clearing CH0 status, DMA Interrupt Status must reflect cleared state. Test completes with pass after all TX transfers are processed.",
        "Remarks": "Supports three speed modes and two duplex modes. Only CH0 TX is started; CH1-CH3 TX start is commented out. Four hardcoded hex addresses used for external system register writes."
    },
    {
        "Index": "4",
        "SS / Module": "Ethernet",
        "Test Case Name": "ethernet0_tx_rx_multi_chnl_test",
        "Feature": "Multi-Channel TX/RX Data Path",
        "Meta Headers": "<stdio.h>; <stdlib.h>; <math.h>; <test_common.h>; <ethernet0/ethernet0_def.h>; <ethernet0/ethernet0_offset.h>; <ethernet0/ethernet0_funcs.h>; <ethernet0/ethernet0_programming_sequence.h>",
        "Meta Macros": "ETH_10M; ETH_100M; SEL_ENET0; SEL_ENET1; SEL_ENET2; SEL_ENET3; HALF_DUPLEX; PRELOAD_AGAIN",
        "Meta Arrays": "NA",
        "Speed": "ETH_10M; ETH_100M",
        "Mode": "HALF_DUPLEX; Full Duplex",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs a multi-channel Ethernet TX and RX data path validation using all four DMA channels (CH0-CH3) sequentially. The test operates in four phases: Phase 1 starts CH3, Phase 2 starts CH2, Phase 3 starts CH1, Phase 4 starts CH0, with dynamic RXQ_DMA_MAP0 remapping and polling of packet count registers until cumulative counts reach 10, 20, 30, 40 respectively.",
        "Test Description": "This test validates the multi-channel Ethernet transmit and receive data path by sequentially activating all four DMA channels (CH3, CH2, CH1, CH0). It configures MAC, MTL, and DMA registers for all four channels and operates in four sequential phases, each starting TX/RX on one channel, dynamically remapping RX queue to DMA channel, and polling packet count registers until expected cumulative counts are reached. The test passes after all four phases complete with 40 total TX and RX packets.",
        "Meta Test Steps / Procedure": "1-38. Configure all MAC, MTL, DMA registers for all 4 channels. Phase 1: Start CH3 TX/RX, poll until 10 packets. Phase 2: Start CH2 TX/RX, poll until 20. Phase 3: Start CH1 TX/RX, poll until 30. Phase 4: Start CH0 TX/RX, poll until 40. IRQ: Advance RX tail pointers, clear DMA status.",
        "Test Steps / Procedure": "1. Initialize GIC and enable interrupts.\n2. Configure MAC speed, duplex, addresses, packet filter, MMC IPC RX interrupt mask.\n3. Configure MTL TX/RX queues, quantum weights, interrupt control, RXQ-DMA mapping, RX queue controls.\n4. Configure DMA TX/RX controls, descriptor addresses, ring lengths, system bus mode.\n5. Preload TX and RX descriptors for all channels.\n6. Enable DMA interrupts and configure DMA mode.\n7. Phase 1: Start CH3 TX/RX, poll until 10 packets.\n8. Phase 2: Start CH2 TX/RX, poll until 20 packets.\n9. Phase 3: Start CH1 TX/RX, poll until 30 packets.\n10. Phase 4: Start CH0 TX/RX, poll until 40 packets.\n11. In IRQ handler, advance RX tail pointers and clear DMA status.\n12. Verify test passes after all phases.",
        "Meta Impacted Registers": "0xA0243ffc; mizar_ETHERNET0_MAC_CONFIGURATION; mizar_ETHERNET0_MAC_EXT_CONFIGURATION; mizar_ETHERNET0_MAC_RXQ_CTRL0; mizar_ETHERNET0_MAC_RXQ_CTRL1; mizar_ETHERNET0_MAC_RXQ_CTRL2; mizar_ETHERNET0_MAC_VLAN_TAG_CTRL; mizar_ETHERNET0_MAC_PACKET_FILTER; mizar_ETHERNET0_MMC_IPC_RX_INTERRUPT_MASK; mizar_ETHERNET0_MAC_ADDRESS3_HIGH; mizar_ETHERNET0_MAC_ADDRESS2_HIGH; mizar_ETHERNET0_MAC_ADDRESS1_HIGH; mizar_ETHERNET0_MAC_ADDRESS0_HIGH; mizar_ETHERNET0_MAC_ADDRESS3_LOW; mizar_ETHERNET0_MAC_ADDRESS2_LOW; mizar_ETHERNET0_MAC_ADDRESS1_LOW; mizar_ETHERNET0_MAC_ADDRESS0_LOW; mizar_ETHERNET0_MTL_TXQ3_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ2_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ1_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ0_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ3_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_TXQ2_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_TXQ1_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_OPERATION_MODE; mizar_ETHERNET0_MTL_TXQ0_QUANTUM_WEIGHT; mizar_ETHERNET0_MTL_RXQ3_OPERATION_MODE; mizar_ETHERNET0_MTL_RXQ2_OPERATION_MODE; mizar_ETHERNET0_MTL_RXQ1_OPERATION_MODE; mizar_ETHERNET0_MTL_RXQ0_OPERATION_MODE; mizar_ETHERNET0_MTL_Q3_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_Q2_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_Q1_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_Q0_INTERRUPT_CONTROL_STATUS; mizar_ETHERNET0_MTL_RXQ_DMA_MAP0; mizar_ETHERNET0_MTL_RXQ3_CONTROL; mizar_ETHERNET0_MTL_RXQ2_CONTROL; mizar_ETHERNET0_MTL_RXQ1_CONTROL; mizar_ETHERNET0_MTL_RXQ0_CONTROL; mizar_ETHERNET0_DMA_CH3_TX_CONTROL; mizar_ETHERNET0_DMA_CH2_TX_CONTROL; mizar_ETHERNET0_DMA_CH1_TX_CONTROL; mizar_ETHERNET0_DMA_CH0_TX_CONTROL; mizar_ETHERNET0_DMA_CH3_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH2_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH1_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH0_TXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH3_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH2_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH1_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH3_CONTROL; mizar_ETHERNET0_DMA_CH2_CONTROL; mizar_ETHERNET0_DMA_CH1_CONTROL; mizar_ETHERNET0_DMA_SYSBUS_MODE; mizar_ETHERNET0_DMA_CH3_RX_CONTROL; mizar_ETHERNET0_DMA_CH2_RX_CONTROL; mizar_ETHERNET0_DMA_CH1_RX_CONTROL; mizar_ETHERNET0_DMA_CH0_RX_CONTROL; mizar_ETHERNET0_DMA_CH3_RXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH2_RXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH1_RXDESC_LIST_ADDRESS; mizar_ETHERNET0_DMA_CH0_RXDESC_LIST_ADDRESS; 0XE68C2058; mizar_ETHERNET0_DMA_CH3_RX_CONTROL2; mizar_ETHERNET0_DMA_CH2_RX_CONTROL2; mizar_ETHERNET0_DMA_CH1_RX_CONTROL2; mizar_ETHERNET0_DMA_CH0_TXDESC_RING_LENGTH; mizar_ETHERNET0_DMA_CH3_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH2_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH1_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_CH2_TXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH3_TXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH0_CONTROL; mizar_ETHERNET0_DMA_CH0_RX_CONTROL2; mizar_ETHERNET0_DMA_CH0_INTERRUPT_ENABLE; mizar_ETHERNET0_DMA_MODE; mizar_ETHERNET0_DMA_CH3_RXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH2_RXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH1_RXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH0_RXDESC_TAIL_POINTER; 0xA0243ff8; mizar_ETHERNET0_DMA_CH1_TXDESC_TAIL_POINTER; mizar_ETHERNET0_DMA_CH0_TXDESC_TAIL_POINTER; mizar_ETHERNET0_RX_PACKETS_COUNT_GOOD_BAD; mizar_ETHERNET0_TX_PACKET_COUNT_GOOD_BAD; mizar_ETHERNET0_DMA_CH0_CURRENT_APP_RXDESC; mizar_ETHERNET0_DMA_CH1_CURRENT_APP_RXDESC; mizar_ETHERNET0_DMA_CH2_CURRENT_APP_RXDESC; mizar_ETHERNET0_DMA_CH3_CURRENT_APP_RXDESC; mizar_ETHERNET0_DMA_INTERRUPT_STATUS; mizar_ETHERNET0_DMA_CH0_STATUS; mizar_ETHERNET0_DMA_CH1_STATUS; mizar_ETHERNET0_DMA_CH2_STATUS; mizar_ETHERNET0_DMA_CH3_STATUS; 0XE68C2050",
        "Impacted Registers": "MAC_Configuration; MAC_Ext_Configuration; MAC_RxQ_Ctrl0; MAC_RxQ_Ctrl1; MAC_RxQ_Ctrl2; MAC_VLAN_Tag_Ctrl; MAC_Packet_Filter; MMC_IPC_Rx_Interrupt_Mask; MAC_Address3_High; MAC_Address2_High; MAC_Address1_High; MAC_Address0_High; MAC_Address3_Low; MAC_Address2_Low; MAC_Address1_Low; MAC_Address0_Low; MTL_TxQ3_Operation_Mode; MTL_TxQ2_Operation_Mode; MTL_TxQ1_Operation_Mode; MTL_TxQ0_Operation_Mode; MTL_TxQ3_Quantum_Weight; MTL_TxQ2_Quantum_Weight; MTL_TxQ1_Quantum_Weight; MTL_Operation_Mode; MTL_TxQ0_Quantum_Weight; MTL_RxQ3_Operation_Mode; MTL_RxQ2_Operation_Mode; MTL_RxQ1_Operation_Mode; MTL_RxQ0_Operation_Mode; MTL_Q3_Interrupt_Control_Status; MTL_Q2_Interrupt_Control_Status; MTL_Q1_Interrupt_Control_Status; MTL_Q0_Interrupt_Control_Status; MTL_RxQ_DMA_Map0; MTL_RxQ3_Control; MTL_RxQ2_Control; MTL_RxQ1_Control; MTL_RxQ0_Control; DMA_CH3_Tx_Control; DMA_CH2_Tx_Control; DMA_CH1_Tx_Control; DMA_CH0_Tx_Control; DMA_CH3_TxDesc_List_Address; DMA_CH2_TxDesc_List_Address; DMA_CH1_TxDesc_List_Address; DMA_CH0_TxDesc_List_Address; DMA_CH3_TxDesc_Ring_Length; DMA_CH2_TxDesc_Ring_Length; DMA_CH1_TxDesc_Ring_Length; DMA_CH3_Control; DMA_CH2_Control; DMA_CH1_Control; DMA_SysBus_Mode; DMA_CH3_Rx_Control; DMA_CH2_Rx_Control; DMA_CH1_Rx_Control; DMA_CH0_Rx_Control; DMA_CH3_RxDesc_List_Address; DMA_CH2_RxDesc_List_Address; DMA_CH1_RxDesc_List_Address; DMA_CH0_RxDesc_List_Address; DMA_CH3_Rx_Control2; DMA_CH2_Rx_Control2; DMA_CH1_Rx_Control2; DMA_CH0_TxDesc_Ring_Length; DMA_CH3_Interrupt_Enable; DMA_CH2_Interrupt_Enable; DMA_CH1_Interrupt_Enable; DMA_CH2_TxDesc_Tail_Pointer; DMA_CH3_TxDesc_Tail_Pointer; DMA_CH0_Control; DMA_CH0_Rx_Control2; DMA_CH0_Interrupt_Enable; DMA_Mode; DMA_CH3_RxDesc_Tail_Pointer; DMA_CH2_RxDesc_Tail_Pointer; DMA_CH1_RxDesc_Tail_Pointer; DMA_CH0_RxDesc_Tail_Pointer; DMA_CH1_TxDesc_Tail_Pointer; DMA_CH0_TxDesc_Tail_Pointer; Rx_Packets_Count_Good_Bad; Tx_Packet_Count_Good_Bad; DMA_CH0_Current_App_RxDesc; DMA_CH1_Current_App_RxDesc; DMA_CH2_Current_App_RxDesc; DMA_CH3_Current_App_RxDesc; DMA_Interrupt_Status; DMA_CH0_Status; DMA_CH1_Status; DMA_CH2_Status; DMA_CH3_Status",
        "Meta Validation / Acceptance Criteria": "Phase 1 polls until rxpkt==10 && txpkt==10. Phase 2 polls until rxpkt==20 && txpkt==20. Phase 3 polls until rxpkt==30 && txpkt==30. Phase 4 polls until rxpkt==40 && txpkt==40. IRQ handler advances RX tail pointers and clears all DMA channel status registers. After all phases, finish(0) indicates pass.",
        "Validation / Acceptance Criteria": "The test passes when all four phases complete: Phase 1 with 10 TX/RX packets on CH3, Phase 2 with cumulative 20 on CH2, Phase 3 with cumulative 30 on CH1, Phase 4 with cumulative 40 on CH0. DMA RX tail pointers must be correctly advanced in the IRQ handler. All DMA channel status registers must be cleared after each interrupt. The test completes with pass after 40 total packets.",
        "Remarks": "Supports three speed modes and two duplex modes. Channels activated in reverse order (CH3 first). Dynamic RXQ_DMA_MAP0 remapping per phase. Circular buffer RX tail pointer management with boundary wrapping. PRELOAD_AGAIN macro enables optional RX descriptor re-preloading. Four hardcoded hex addresses for external system register writes."
    }
]

# TestPlan sheet columns
tp_columns = [
    'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
    'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
    'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
    'Code Generation'
]

# MetaData sheet columns
md_columns = [
    'Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
    'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
    'Meta Headers', 'Meta Macros', 'Meta Arrays'
]

# Create workbook
wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = 'TestPlan'

# Header formatting
header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap_alignment = Alignment(wrap_text=True, vertical='top')
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Write TestPlan headers
for col_idx, col_name in enumerate(tp_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment
    cell.border = thin_border

# Write TestPlan data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(tp_columns, 1):
        value = row_data.get(col_name, '')
        if value is None:
            value = ''
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment
        cell.border = thin_border

# Freeze first row
ws_tp.freeze_panes = 'A2'

# Auto-size columns with max width cap
MAX_WIDTH = 60
for col_idx in range(1, len(tp_columns) + 1):
    max_len = len(str(ws_tp.cell(row=1, column=col_idx).value or ''))
    for row_idx in range(2, len(json_data) + 2):
        cell_val = str(ws_tp.cell(row=row_idx, column=col_idx).value or '')
        lines = cell_val.split('\n')
        for line in lines:
            max_len = max(max_len, len(line))
    adjusted_width = min(max_len + 4, MAX_WIDTH)
    ws_tp.column_dimensions[get_column_letter(col_idx)].width = max(adjusted_width, 12)

# --- MetaData Sheet ---
ws_md = wb.create_sheet('MetaData')

# Write MetaData headers
for col_idx, col_name in enumerate(md_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment
    cell.border = thin_border

# Write MetaData data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(md_columns, 1):
        value = row_data.get(col_name, '')
        if value is None:
            value = ''
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment
        cell.border = thin_border

# Freeze first row
ws_md.freeze_panes = 'A2'

# Auto-size MetaData columns
for col_idx in range(1, len(md_columns) + 1):
    max_len = len(str(ws_md.cell(row=1, column=col_idx).value or ''))
    for row_idx in range(2, len(json_data) + 2):
        cell_val = str(ws_md.cell(row=row_idx, column=col_idx).value or '')
        lines = cell_val.split('\n')
        for line in lines:
            max_len = max(max_len, len(line))
    adjusted_width = min(max_len + 4, MAX_WIDTH)
    ws_md.column_dimensions[get_column_letter(col_idx)].width = max(adjusted_width, 12)

# Set MetaData sheet to veryHidden
ws_md.sheet_state = 'veryHidden'

# Save workbook
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)
wb.save(output_path)

# Validate
assert os.path.exists(output_path), f'File not found: {output_path}'
assert os.path.getsize(output_path) > 0, f'File is empty: {output_path}'

# Re-open to validate
wb_check = load_workbook(output_path)
assert 'TestPlan' in wb_check.sheetnames, 'TestPlan sheet missing'
assert 'MetaData' in wb_check.sheetnames, 'MetaData sheet missing'
assert wb_check['TestPlan'].max_row == 5, f'Expected 5 rows (1 header + 4 data), got {wb_check["TestPlan"].max_row}'
assert wb_check['MetaData'].max_row == 5, f'Expected 5 rows (1 header + 4 data), got {wb_check["MetaData"].max_row}'

print(f'SUCCESS: Generated {filename}')
print(f'Path: {output_path}')
print(f'Size: {os.path.getsize(output_path)} bytes')
print(f'TestPlan rows: {wb_check["TestPlan"].max_row - 1}')
print(f'MetaData rows: {wb_check["MetaData"].max_row - 1}')
print(f'Validation: PASSED')
