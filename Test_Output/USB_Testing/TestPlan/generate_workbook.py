#!/usr/bin/env python3
"""
Auto-generated workbook generator for USB TestPlan.
Execute this script to produce the final XLSX workbook.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone, timedelta
import os, json

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
filename = f"USB_TestPlan_{now_ist.strftime('%Y%m%d')}_{now_ist.strftime('%H%M%S')}.xlsx"
output_dir = os.path.dirname(os.path.abspath(__file__))
filepath = os.path.join(output_dir, filename)

json_data = [
  {
    "Index": "1",
    "SS / Module": "USB",
    "Test Case Name": "usb_host_enumeration_ls",
    "Feature": "NA",
    "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include \"usb.h\"",
    "Meta Macros": "MIZAR_USB_BASE, <PERSON>, <LOCATION>, <LOCATION>, MIZAR_USB_ERDP_LO, MIZAR_USB_ERDP_HI, MIZAR_USB_ERSTBA_LO, MIZAR_USB_ERSTBA_HI, MIZAR_USB_ERSTSZ, MIZAR_USB_USBCMD, MIZAR_USB_CONFIG, MIZAR_USB_DCBAAP_LO, MIZAR_USB_DCBAAP_HI, MIZAR_USB_HCSPARAMS1, <LOCATION>, <PERSON>, USB_PORTSC_20_WCE, USB_PORTSC_20_WDE, USB_PORTSC_20_WOE, MIZAR_USB_DBOFF, MIZAR_USB_CRCR_LO, MIZAR_USB_CRCR_HI, MIZAR_USB_DB, MIZAR_USB_GEVNTCOUNT, MIZAR_LSS_SYSREG_MSK_STS0, MIZAR_LSS_SYSREG_RAW_STCR0, MIZAR_LSS_SYSREG_INTR_EN0, DWORD, Default_Event_Ring_Array, Default_Input_Context, Event_Ring_Segment_Table, EP0_TR_Dequeue_Pointer, Default_Command_Ring, <LOCATION>, Device_Context_Array",
    "Meta Arrays": "int data_in[512];\nint data_out[512];",
    "Speed": "Low-Speed (LS)",
    "Mode": "Host",
    "Memory Start Offset": "NA",
    "Memory End Offset": "NA",
    "Meta Test Description": "USB host enumeration test for a Low-Speed (LS) device. The test initializes the xHCI host controller by programming NIC, enabling all IRQs via GIC, configuring PHY control register, reading <PERSON> for port count, reading supported protocol capability registers (<PERSON>, SUPTPRT3_DW2), and reading PORTSC_20. It then sets up the Device Context Base Address Array (DCBAAP), Command Ring (CRCR), Event Ring Segment Table (ERSTBA), Event Ring Dequeue Pointer (ERDP), Interrupt Moderation (IMOD), Interrupt Management (IMAN), Event Ring Segment Table Size (ERSTSZ), and CONFIG register. Wake enable bits (WCE, WDE, WOE) are set on PORTSC_20. The doorbell offset (DBOFF) is read. The host controller is started via USBCMD. An interrupt-driven polling mechanism (int_pend with wait_on(100)) is used to wait for port connect events and port reset completion. An Enable Slot Command TRB is issued on the Command Ring and the doorbell is rung. The Input Context (Slot Context and Endpoint 0 Context) is configured for a low-speed device with max packet size of 8 bytes. An Address Device command is issued. The enumeration() function performs GET_DESCRIPTOR (Device Descriptor, 18 bytes) and GET_DESCRIPTOR (Configuration Descriptor, 24 bytes) control transfers via EP0 Transfer Ring TRBs (Setup, Data, Status stages), ringing the doorbell to trigger processing. Event completion is read from the Event Ring. The test passes via finish(0) if count is not greater than 0, or fails via finish(1) if count > 0.",
    "Test Description": "This test verifies USB host-side enumeration of a Low-Speed device. The xHCI host controller is initialized, interrupts are enabled, and the PHY is configured. The controller\u2019s data structures (Device Context Base Address Array, Command Ring, Event Ring) are set up. Port wake enables are configured. The host controller is started and waits for a device connect event. Upon detection, a port reset is performed. An Enable Slot command is issued, followed by Input Context configuration for a low-speed device and an Address Device command. The enumeration sequence then performs GET_DESCRIPTOR requests for Device and Configuration descriptors via control transfers on Endpoint 0. Event completion is verified from the Event Ring, and the test result is determined based on error count.",
    "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC/controller initialization.\n2. Call <PERSON>() to enable all IRQs (GIC_EnableIRQ(84) is commented out).\n3. Read PHY control register: rd_data = read_reg(MIZAR_USB_BASE + 0xc200).\n4. Write PHY control register: write_reg(MIZAR_USB_BASE + 0xc200, 0x102407).\n5. Read port count: port_count = read_reg(MIZAR_USB_HCSPARAMS1).\n6. Read supported protocol USB 2.0: rd_data = read_reg(MIZAR_USB_SUPTPRT2_DW2).\n7. Read supported protocol USB 3.0: rd_data = read_reg(MIZAR_USB_SUPTPRT3_DW2).\n8. Read port status: rd_data = read_reg(MIZAR_USB_PORTSC_20).\n9. Read CONFIG: rd_data = read_reg(MIZAR_USB_CONFIG).\n10. Write CONFIG with CIE: write_reg(MIZAR_USB_CONFIG, 0x110).\n11. Write DCBAAP_LO: write_reg(MIZAR_USB_DCBAAP_LO, Device_Context_Base_Address_Array).\n12. Write DCBAAP_HI: write_reg(MIZAR_USB_DCBAAP_HI, <LOCATION>. Write ERSTSZ: write_reg(MIZAR_USB_ERSTSZ, 0x1).\n14. Write ERDP_LO: <LOCATION>, Default_Event_Ring_Array).\n15. Write ERDP_HI: write_reg(MIZAR_USB_ERDP_HI, 0x0).\n16. Write Event_Ring_Segment_Table entries: write_reg(Event_Ring_Segment_Table, Default_Event_Ring_Array), write_reg(Event_Ring_Segment_Table + DWORD, 0x0), write_reg(Event_Ring_Segment_Table + 2DWORD, 0x30).\n17. Write ERSTBA_LO: write_reg(MIZAR_USB_ERSTBA_LO, Event_Ring_Segment_Table).\n18. Write ERSTBA_HI: write_reg(MIZAR_USB_ERSTBA_HI, <PERSON>. Write IMOD: write_reg(MIZAR_USB_IMOD, 0x0).\n20. Write <PERSON>: write_reg(MIZAR_USB_IMAN, 0x2).\n21. Write USBCMD: write_reg(MIZAR_USB_USBCMD, <PERSON>. Write PORTSC_20 wake enables: write_reg(MIZAR_USB_PORTSC_20, set_data(read_reg(MIZAR_USB_PORTSC_20), USB_PORTSC_20_WCE, 1)), write_reg(MIZAR_USB_PORTSC_20, set_data(read_reg(MIZAR_USB_PORTSC_20), USB_PORTSC_20_WDE, 1)), write_reg(MIZAR_USB_PORTSC_20, set_data(read_reg(MIZAR_USB_PORTSC_20), USB_PORTSC_20_WOE, 1)).\n23. Write PORTSC_20 consolidated: write_reg(MIZAR_USB_PORTSC_20, 0xe0002a0).\n24. Read doorbell offset: db_offset = <PERSON>. Write USBCMD to start host controller: write_reg(MIZAR_USB_USBCMD, 0x5).\n26. Set int_pend = 1, poll while(int_pend) { wait_on(100); } \u2014 wait for port connect event.\n27. Read USB status: usb_status = read_reg(MIZAR_USB_USBSTS).\n28. Clear USBSTS: write_reg(MIZAR_USB_USBSTS, <LOCATION>. Write <PERSON>: write_reg(MIZAR_USB_IMAN, 0x2).\n30. Read PORTSC_20: port_status = read_reg(MIZAR_USB_PORTSC_20).\n31. Update ERDP_LO: <LOCATION>, Default_Event_Ring_Array + 0x18).\n32. Write PORTSC_20 for port reset: write_reg(MIZAR_USB_PORTSC_20, 0xe0006f1).\n33. Clear USBSTS: write_reg(MIZAR_USB_USBSTS, <LOCATION>. Write <PERSON>: write_reg(MIZAR_USB_IMAN, 0x2).\n35. Write PORTSC_20 post-reset: write_reg(MIZAR_USB_PORTSC_20, 0xe220200).\n36. Set int_pend = 1, poll while(int_pend) { wait_on(100); } \u2014 wait for port reset completion.\n37. Write USBSTS: write_reg(MIZAR_USB_USBSTS, 0x8).\n38. Write <PERSON>: write_reg(MIZAR_USB_IMAN, <LOCATION>. Write PORTSC_20: write_reg(MIZAR_USB_PORTSC_20, 0xe200e01).\n40. Write Device_Context_Base_Address_Array entries: write_reg(Device_Context_Base_Address_Array + 2DWORD, Device_Context_Array + 0x100), write_reg(Device_Context_Base_Address_Array + 3DWORD, 0x0), write_reg(Device_Context_Base_Address_Array + 4DWORD, Device_Context_Array + 0x0d00), write_reg(Device_Context_Base_Address_Array + 5*DWORD, 0x0).\n41. Write CONFIG MaxSlotsEn: write_reg(MIZAR_USB_CONFIG, 0x10).\n42. Write Enable Slot Command TRB to Default_Command_Ring: write_reg(Default_Command_Ring + 0x0, 0x0), write_reg(Default_Command_Ring + 0x4, 0x0), write_reg(Default_Command_Ring + 0x8, 0x0), write_reg(Default_Command_Ring + 0xc, 0x00002401).\n43. Write <NRP>: write_reg(MIZAR_USB_CRCR_LO, Default_Command_Ring + 0x1).\n44. Write CRCR_HI: write_reg(MIZAR_USB_CRCR_HI, 0x0).\n45. Ring doorbell: write_reg(MIZAR_USB_DB, <LOCATION>. Set int_pend = 1, poll while(int_pend) { wait_on(100); } \u2014 wait for Enable Slot completion.\n47. Read USBSTS: usb_status = read_reg(MIZAR_USB_USBSTS).\n48. Clear USBSTS: write_reg(MIZAR_USB_USBSTS, <LOCATION>. Write <PERSON>: write_reg(MIZAR_USB_IMAN, 0x2).\n50. Update <PERSON>: write_reg(MIZAR_USB_ERDP_HI, 0x0), <LOCATION>, Default_Event_Ring_Array + 0x28).\n51. Write Input Control Context: write_reg(Default_Input_Context, 0x0), write_reg(Default_Input_Context + DWORD, 0x3).\n52. Write Slot Context: write_reg(Default_Input_Context + 0x40, 0x08200000), write_reg(Default_Input_Context + 0x44, 0x00010000).\n53. Write Endpoint 0 Context: write_reg(Default_Input_Context + 0x80, 0x00), write_reg(Default_Input_Context + 0x84, 0x00080020), write_reg(Default_Input_Context + 0x88, EP0_TR_Dequeue_Pointer | 0x1), write_reg(Default_Input_Context + 0x90, 0x08).\n54. Issue Address Device command on Command Ring and ring doorbell.\n55. Set int_pend = 1, poll while(int_pend) { wait_on(100); } \u2014 wait for Address Device completion.\n56. Call enumeration() function.\n57. enumeration(): Write GET_DESCRIPTOR (Device) Setup TRB: write_reg(EP0_TR_Dequeue_Pointer, 0x01000680), write_reg(EP0_TR_Dequeue_Pointer + 0x4, 0x00120000), write_reg(EP0_TR_Dequeue_Pointer + 0x8, 0x08), write_reg(EP0_TR_Dequeue_Pointer + 0xc, 0x00030861).\n58. Write Data Stage TRB: write_reg(EP0_TR_Dequeue_Pointer + 0x10, EP0_TR_Dequeue_Pointer + 0x200), write_reg(EP0_TR_Dequeue_Pointer + 0x18, 0x12), write_reg(EP0_TR_Dequeue_Pointer + 0x1c, 0x00010c27).\n59. Write GET_DESCRIPTOR (Configuration) Setup TRB: write_reg(EP0_TR_Dequeue_Pointer + 0xB0, 0x02000680), write_reg(EP0_TR_Dequeue_Pointer + 0xB4, 0x00180000), write_reg(EP0_TR_Dequeue_Pointer + 0xB8, 0x08), write_reg(EP0_TR_Dequeue_Pointer + 0xBC, 0x00030861).\n60. Write Data Stage TRB for Configuration Descriptor: write_reg(EP0_TR_Dequeue_Pointer + 0xC0, EP0_TR_Dequeue_Pointer + 0x560), write_reg(EP0_TR_Dequeue_Pointer + 0xC8, 0x18), write_reg(EP0_TR_Dequeue_Pointer + 0xCC, 0x00010C25).\n61. Write Status Stage TRB: write_reg(EP0_TR_Dequeue_Pointer + 0xdc, 0x00001023).\n62. Ring doorbell for EP0: write_reg(MIZAR_USB_BASE + 0x484, 0x1).\n63. Clear USBSTS: write_reg(MIZAR_USB_USBSTS, 0x8).\n64. Write <PERSON>: write_reg(MIZAR_USB_IMAN, 0x2).\n65. Update <PERSON>: write_reg(MIZAR_USB_ERDP_HI, 0x0), <LOCATION>, Default_Event_Ring_Array + 0x68).\n66. Ring doorbell again: write_reg(MIZAR_USB_BASE + 0x484, 0x2).\n67. Set int_pend = 1, poll while(int_pend) { wait_on(100); } \u2014 wait for enumeration completion.\n68. Read event completion: event_completion = read_reg(Default_Event_Ring_Array + 0x60).\n69. Print: printf(\"event_completion is %x Default_Event_Ring_Array %x\\n\", event_completion, Default_Event_Ring_Array).\n70. Validation: if (count > 0x0) { finish(1); } else { finish(0); }.\n\nDefault_IRQHandler():\n1. Set int_pend = 0.\n2. Read MIZAR_LSS_SYSREG_MSK_STS0.\n3. Read MIZAR_LSS_SYSREG_RAW_STCR0.\n4. If (rd_data && 0x80000000): <LOCATION>, <LOCATION>), write_reg(MIZAR_LSS_SYSREG_RAW_STCR0, 0x80000000).\n5. Call GIC_ClearIRQ(84).",
    "Test Steps / Procedure": "1. Initialize NIC controller programming.\n2. Enable all IRQs via GIC.\n3. Configure PHY control register.\n4. Read HCSPARAMS1 to determine port count.\n5. Read Supported Protocol Capability registers (SUPTPRT2_DW2, SUPTPRT3_DW2).\n6. Read PORTSC_20 for initial port status.\n7. Configure CONFIG register with CIE bit.\n8. Set up DCBAAP_LO and DCBAAP_HI with Device Context Base Address Array pointer.\n9. Configure Event Ring: set ERSTSZ to 1, initialize ERDP_LO/HI, populate Event Ring Segment Table, set ERSTBA_LO/HI.\n10. Configure IMOD to 0 and IMAN to 0x2.\n11. Write USBCMD to enable interrupts.\n12. Enable PORTSC_20 wake enables (WCE, WDE, WOE).\n13. Read DBOFF for doorbell offset.\n14. Start host controller by writing USBCMD with Run/Stop bit.\n15. Wait for port connect event via interrupt-driven polling.\n16. Acknowledge USB status and interrupt management registers.\n17. Read PORTSC_20 to verify port connect.\n18. Update ERDP_LO to advance Event Ring Dequeue Pointer.\n19. Initiate port reset via PORTSC_20.\n20. Clear status and interrupt registers, configure post-reset port settings.\n21. Wait for port reset completion via interrupt-driven polling.\n22. Set up Device Context Base Address Array entries for device slots.\n23. Configure CONFIG register with MaxSlotsEn = 16.\n24. Issue Enable Slot Command TRB on Command Ring.\n25. Set up CRCR_LO/HI and ring host controller doorbell.\n26. Wait for Enable Slot completion.\n27. Acknowledge status, update ERDP.\n28. Configure Input Context: Input Control Context (Add Slot + EP0), Slot Context (low-speed, port 1), Endpoint 0 Context (Control EP, max packet size 8, Transfer Ring pointer with DCS=1, average TRB length 8).\n29. Issue Address Device command and ring doorbell.\n30. Wait for Address Device completion.\n31. Call enumeration() to perform GET_DESCRIPTOR (Device Descriptor, 18 bytes) and GET_DESCRIPTOR (Configuration Descriptor, 24 bytes) control transfers via EP0 TRBs.\n32. Ring doorbell for EP0 to trigger transfer processing.\n33. Clear status, update ERDP, ring doorbell again.\n34. Wait for enumeration completion.\n35. Read event completion from Event Ring.\n36. Determine test result: finish(0) for pass, finish(1) for fail based on error count.",
    "Meta Impacted Registers": "USBSTS, <PERSON>, IMOD, ERDP_LO, ERDP_HI, ERSTBA_LO, ERSTBA_HI, ERSTSZ, USBCMD, CONFIG, DCBAAP_LO, DCBAAP_HI, <PERSON>, <PERSON>, SUPTPRT3_DW2, PORTSC_20, DBOFF",
    "Impacted Registers": "USBSTS, <PERSON>, IMOD, ERDP_LO, ERDP_HI, ERSTBA_LO, ERSTBA_HI, ERSTSZ, USBCMD, CONFIG, DCBAAP_LO, DCBAAP_HI, <PERSON>, <PERSON>, SUPTPRT3_DW2, PORTSC_20, DBOFF",
    "Meta Validation / Acceptance Criteria": "1. Event completion is read from the Event Ring: event_completion = read_reg(Default_Event_Ring_Array + 0x60).\n2. Debug print: printf(\"event_completion is %x Default_Event_Ring_Array %x\\n\", event_completion, Default_Event_Ring_Array).\n3. Pass condition: if count is not greater than 0x0, call finish(0) \u2014 test PASS.\n4. Fail condition: if count > 0x0, call finish(1) \u2014 test FAIL.\n5. Port status is read after connect event: port_status = read_reg(MIZAR_USB_PORTSC_20) \u2014 port connect status should be high.\n6. USB status is read after host controller start: usb_status = read_reg(MIZAR_USB_USBSTS).\n7. Interrupt-driven synchronization: int_pend flag is set to 1 before each wait, and cleared to 0 by Default_IRQHandler upon interrupt receipt.\n8. Default_IRQHandler checks bit 31 (0x80000000) of MIZAR_LSS_SYSREG_RAW_STCR0 and acknowledges <PERSON> with <LOCATION> and clears raw status if set.",
    "Validation / Acceptance Criteria": "1. Event completion status is successfully read from the Event Ring after enumeration.\n2. The test passes (finish(0)) when no errors are detected (count equals 0).\n3. The test fails (finish(1)) if any error count is greater than 0.\n4. PORTSC_20 is read to verify port connect status is asserted after device connection.\n5. USBSTS is read to verify USB status after host controller is started.\n6. Interrupt-driven synchronization confirms that each command/event completes successfully via the Default_IRQHandler clearing the pending flag.\n7. IMAN is acknowledged in the interrupt handler to confirm interrupt source.",
    "Remarks": "The test uses an interrupt-driven polling mechanism with int_pend flag and wait_on(100) for synchronization between main flow and Default_IRQHandler. GIC_EnableIRQ(84) is commented out in favor of <PERSON>(). The doorbell for EP0 transfers is accessed via MIZAR_USB_BASE + 0x484 rather than through the DBOFF-derived address. The Slot Context encodes low-speed device configuration (0x08200000) with root hub port number 1 (0x00010000). EP0 max packet size is 8 bytes, typical for low-speed USB devices."
  }
]

# Create workbook
wb = openpyxl.Workbook()

# ---- TestPlan Sheet ----
ws_tp = wb.active
ws_tp.title = "TestPlan"

tp_headers = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_alignment = Alignment(wrap_text=True, vertical="top")

for col_idx, header in enumerate(tp_headers, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

for row_data in json_data:
    row_values = [
        row_data.get("Index", ""),
        row_data.get("SS / Module", ""),
        row_data.get("Feature", ""),
        row_data.get("Test Case Name", ""),
        row_data.get("Test Description", ""),
        row_data.get("Speed", ""),
        row_data.get("Mode", ""),
        row_data.get("Memory Start Offset", ""),
        row_data.get("Memory End Offset", ""),
        row_data.get("Remarks", ""),
        row_data.get("Test Steps / Procedure", ""),
        row_data.get("Impacted Registers", ""),
        row_data.get("Validation / Acceptance Criteria", ""),
        ""
    ]
    ws_tp.append(row_values)
    for col_idx in range(1, len(row_values) + 1):
        ws_tp.cell(row=ws_tp.max_row, column=col_idx).alignment = wrap_alignment

ws_tp.freeze_panes = "A2"

for col in ws_tp.columns:
    max_length = 0
    col_letter = col[0].column_letter
    for cell in col:
        if cell.value:
            max_length = max(max_length, min(len(str(cell.value)), 60))
    ws_tp.column_dimensions[col_letter].width = max(12, min(max_length + 2, 60))

# ---- MetaData Sheet ----
ws_md = wb.create_sheet("MetaData")

md_headers = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

for col_idx, header in enumerate(md_headers, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

for row_data in json_data:
    row_values = [
        row_data.get("Index", ""),
        row_data.get("Test Case Name", ""),
        row_data.get("Meta Test Description", ""),
        row_data.get("Meta Test Steps / Procedure", ""),
        row_data.get("Meta Impacted Registers", ""),
        row_data.get("Meta Validation / Acceptance Criteria", ""),
        row_data.get("Meta Headers", ""),
        row_data.get("Meta Macros", ""),
        row_data.get("Meta Arrays", "")
    ]
    ws_md.append(row_values)
    for col_idx in range(1, len(row_values) + 1):
        ws_md.cell(row=ws_md.max_row, column=col_idx).alignment = wrap_alignment

ws_md.freeze_panes = "A2"
ws_md.sheet_state = "veryHidden"

for col in ws_md.columns:
    max_length = 0
    col_letter = col[0].column_letter
    for cell in col:
        if cell.value:
            max_length = max(max_length, min(len(str(cell.value)), 60))
    ws_md.column_dimensions[col_letter].width = max(12, min(max_length + 2, 60))

# Save
wb.save(filepath)
print(f"SAVED: {filepath}")

# Validate
wb2 = openpyxl.load_workbook(filepath)
assert "TestPlan" in wb2.sheetnames
assert "MetaData" in wb2.sheetnames
tp = wb2["TestPlan"]
md = wb2["MetaData"]
assert tp.max_row == 2
assert md.max_row == 2

# Meta validation
for row_data in json_data:
    assert md.cell(row=2, column=3).value == row_data["Meta Test Description"]
    assert md.cell(row=2, column=4).value == row_data["Meta Test Steps / Procedure"]
    assert md.cell(row=2, column=5).value == row_data["Meta Impacted Registers"]
    assert md.cell(row=2, column=6).value == row_data["Meta Validation / Acceptance Criteria"]
    assert md.cell(row=2, column=7).value == row_data["Meta Headers"]
    assert md.cell(row=2, column=8).value == row_data["Meta Macros"]
    assert md.cell(row=2, column=9).value == row_data["Meta Arrays"]

print(f"VALIDATION: PASSED")
print(f"FILENAME: {filename}")
print(f"ROWS_TESTPLAN: 1")
print(f"ROWS_METADATA: 1")
wb2.close()
