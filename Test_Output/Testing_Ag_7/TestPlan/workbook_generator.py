#!/usr/bin/env python3
"""
Agent 7 - USB TestPlan Excel Workbook Generator
Generates: USB_TestPlan_20261013_002556.xlsx
Run this script to generate the workbook locally, then upload the .xlsx file.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os

# ============================================================
# FULL JSON DATA (exact values from pipeline)
# ============================================================

row = {
    "Index": "2",
    "SS / Module": "USB",
    "Test Case Name": "USB_FS_Device_Bulk_Transfer_test",
    "Feature": "USB DWC3 Device Controller Full-Speed Bulk Transfer",
    "Speed": "Full-Speed (FS)",
    "Mode": "Device",
    "Memory Start Offset": "RAM+0xd00",
    "Memory End Offset": "RAM+0x3000",
    "Test Description": "Configure the USB DWC3 device controller in device mode for full-speed operation. Initialize the event buffer (GEVNTADRLO/HI/SIZ/COUNT), configure global control (GCTL) for device mode, set device configuration (DCFG), enable device events (DEVTEN), configure user control (GUCTL), issue DEPSTARTCFG command via set_configuration, configure transfer resources for 8 endpoints using DEPXFERCFG loop, enable physical endpoints 0 and 1 (DALEPENA), set device control (DCTL), enable system-level interrupt (INTR_EN0). Perform USB enumeration: handle setup/status stages, process SET_CONFIGURATION and GET_DESCRIPTOR (Configuration) requests with 60-byte configuration descriptor containing 6 endpoint descriptors (bulk and interrupt). Configure USB2 PHY (GUSB2PHYCFG). Set device address via DCFG. Execute bulk data transfer on EP4 with 64-byte max packet size using Buffer_PointerLO_1 and StartXfer command (0x506). Verify transfer completion via interrupt-driven event processing.",
    "Remarks": "The Default_IRQHandler uses logical AND (&&) instead of bitwise AND (&) for checking bit 31 of MIZAR_LSS_SYSREG_RAW_STCR0, which is a potential bug \u2014 it evaluates true for any non-zero value. Four system/device register macros are unresolved: MIZAR_LSS_SYSREG_INTR_EN0, MIZAR_LSS_SYSREG_MSK_STS0, MIZAR_LSS_SYSREG_RAW_STCR0 (lss_sysreg headers not in corpus), MIZAR_USB_DCFG (macro definition not found in header corpus though USB_DCFG_OFFSET=0xC700 exists). Seven register macros have ambiguous offsets due to dual definitions in headers: MIZAR_USB_GEVNTADRLO, MIZAR_USB_GEVNTADRHI, MIZAR_USB_GEVNTSIZ, MIZAR_USB_GEVNTCOUNT, MIZAR_USB_DEPCMDPAR1, MIZAR_USB_DEPCMDPAR0, MIZAR_USB_DEPCMD. RAM-based memory macros (Default_Event_Ring_Array, Buffer_PointerLO, Buffer_PointerLO_1, event_trb_addr) have ambiguous base addresses due to conditional RAM definition in headers. MIZAR_USB_GUSB2PHYCFG is used in source but was not in the original 22-macro list from Agent 2. The variable event_comletion appears to be a typo for event_completion. Debug markers 0xdeadbee4 and 0xdeadbee5 are written to address 0xA0243ffc to track enumeration progress. The DCFG register is written twice: first with 0x480801 during init, then with 0x480809 during SET ADDRESS.",
    "Test Steps / Procedure": """1. Call nic_programming() to initialize NIC subsystem.
2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.
3. Clear 20 DWORD entries at Buffer_PointerLO and event_trb_addr to zero.
4. Configure event buffer: write GEVNTADRLO=Default_Event_Ring_Array, GEVNTADRHI=0x0, GEVNTSIZ=0x30 (48 bytes), GEVNTCOUNT=0x0.
5. Read GCTL, then write GCTL=0x30c12214 to set device mode port direction.
6. Read DCFG, then write DCFG=0x480801 for device configuration.
7. Write DEVTEN=0x1f to enable device events (USB reset, connect done, disconnect, etc.).
8. Read GUCTL, then write GUCTL=0xa400010 for user control settings.
9. Call set_configuration(0, 0, 0, 0x409) to issue DEPSTARTCFG (Start New Configuration) command.
10. Loop i=0 to 7: For each endpoint, write DEPCMDPAR0+(i0x10)=0x1, write DEPCMD+(i0x10)=0x402 (DEPXFERCFG), poll DEPCMD until command completes.
11. Write DALEPENA=0x3 to enable physical endpoints 0 and 1.
12. Write DCTL=0x80f00000 to configure and start device controller.
13. Write MIZAR_LSS_SYSREG_INTR_EN0=0x80000000 to enable system-level interrupt.
14. Wait for interrupt (int_pend polling loop).
15. Issue StartXfer on EP1: write DEPCMDPAR1+0x10=event_trb_addr, DEPCMDPAR0+0x10=0x0, DEPCMD+0x10=0x506, poll until completion.
16. Wait for three successive interrupts.
17. Call status_stage() to complete prior control transfer.
18. Call setup_stage() for USB_SET_CONFIGURATION_OR_RESET_TT, write debug marker 0xdeadbee4.
19. Wait for interrupt.
20. Call setup_stage(), then write GUSB2PHYCFG=0x40002547 to configure USB2 PHY.
21. Check event_counter<=0x4 twice, waiting for interrupt each time if true.
22. Call setup_stage() for GET_DESCRIPTOR USB CONFIGURATION, write debug marker 0xdeadbee5.
23. Wait for interrupt.
24. Data stage: configure TRB at event_trb_addr with Buffer_PointerLO, length=0x9, control=0x853. Write configuration descriptor data to Buffer_PointerLO: 0x003c0209, 0xe0000101, 0x00000032, 0x00000000.
25. Write DEPCMDPAR1+0x10=event_trb_addr, DEPCMDPAR0+0x10=0x0.
26. SET ADDRESS: write DCFG=0x480809.
27. Configure TRB at event_trb_addr with Buffer_PointerLO, length=0x0, control=0x853. Write DEPCMDPAR1+0x10=event_trb_addr, DEPCMDPAR0+0x10=0x0.
28. Bulk transfer: configure TRB at event_trb_addr with Buffer_PointerLO_1, length=0x40 (64 bytes), control=0x813.
29. Write DEPCMDPAR1+0x40=event_trb_addr, DEPCMDPAR0+0x40=0x0, DEPCMD+0x40=0x506 (StartXfer on EP4), poll until completion.
30. Wait for interrupt to confirm bulk transfer completion.""",
    "Impacted Registers": "GCTL (offset 0xC110, read then write 0x30c12214 for device mode), DCTL (offset 0xC704, write 0x80f00000 to start device controller), DALEPENA (offset 0xC720, write 0x3 to enable physical EP 0 and 1), DEVTEN (offset 0xC708, write 0x1f to enable device events), GUCTL (offset 0xC12C, read then write 0xa400010), DSTS (offset 0xC70C, read in Default_IRQHandler for status)",
    "Validation / Acceptance Criteria": """1. All 8 endpoint DEPXFERCFG commands (0x402) must complete successfully (DEPCMD register value changes from 0x402 for each endpoint).
2. All StartXfer commands (0x506) on EP0, EP1, and EP4 must complete successfully (DEPCMD register value changes from 0x506).
3. DEPSTARTCFG command (0x409) issued via set_configuration must complete.
4. Interrupt handler must correctly read GEVNTCOUNT, store to event_counter, and write back to acknowledge events.
5. Interrupt handler must identify USB interrupt source from MIZAR_LSS_SYSREG_RAW_STCR0 bit 31 and clear via GIC_ClearIRQ(84).
6. Configuration descriptor (60 bytes) must be correctly written to Buffer_PointerLO with valid USB descriptor format.
7. SET ADDRESS must succeed by writing DCFG=0x480809.
8. Bulk transfer on EP4 with 64-byte payload using Buffer_PointerLO_1 must complete via interrupt-driven event processing.
9. USB2 PHY must be configured with GUSB2PHYCFG=0x40002547.
10. All interrupt-driven wait loops (int_pend) must be serviced by Default_IRQHandler.""",
    "Meta Headers": '#include <stdio.h>\n#include <stdlib.h>\n#include "usb.h"',
    "Meta Macros": "MIZAR_USB_GEVNTADRLO, MIZAR_USB_GEVNTADRHI, MIZAR_USB_GEVNTSIZ, MIZAR_USB_GEVNTCOUNT, MIZAR_USB_GCTL, MIZAR_USB_DCFG, MIZAR_USB_DSTS, MIZAR_USB_DCTL, MIZAR_USB_DALEPENA, MIZAR_USB_DEVTEN, MIZAR_USB_GUCTL, MIZAR_USB_DEPCMDPAR1, MIZAR_USB_DEPCMDPAR0, MIZAR_USB_DEPCMD, MIZAR_LSS_SYSREG_INTR_EN0, MIZAR_LSS_SYSREG_MSK_STS0, MIZAR_LSS_SYSREG_RAW_STCR0, Default_Event_Ring_Array, Buffer_PointerLO, Buffer_PointerLO_1, event_trb_addr, DWORD",
    "Meta Arrays": "int buf_data[16];",
    "Meta Test Description": """Configure USB DWC3 device controller in device mode for full-speed bulk transfer. Clear buffer and TRB memory by zeroing 20 DWORD entries at Buffer_PointerLO and event_trb_addr. Configure event buffer: write MIZAR_USB_GEVNTADRLO=Default_Event_Ring_Array, MIZAR_USB_GEVNTADRHI=0x0, MIZAR_USB_GEVNTSIZ=0x30, MIZAR_USB_GEVNTCOUNT=0x0. Read MIZAR_USB_GCTL, then write MIZAR_USB_GCTL=0x30c12214 (device mode port direction). Read MIZAR_USB_DCFG, then write MIZAR_USB_DCFG=0x480801. Write MIZAR_USB_DEVTEN=0x1f to enable device events. Read MIZAR_USB_GUCTL, then write MIZAR_USB_GUCTL=0xa400010. Call set_configuration(0, 0, 0, 0x409) for START NEW CONFIGURATION (DEPSTARTCFG). Loop i=0 to 7: write MIZAR_USB_DEPCMDPAR0+(i0x10)=0x1, write MIZAR_USB_DEPCMD+(i0x10)=0x402 (DEPXFERCFG), poll MIZAR_USB_DEPCMD+(i0x10) until completion. Write MIZAR_USB_DALEPENA=0x3 (enable physical EP 0 and 1). Write MIZAR_USB_DCTL=0x80f00000. Write MIZAR_LSS_SYSREG_INTR_EN0=0x80000000. Wait for interrupt. Issue StartXfer on EP1: write MIZAR_USB_DEPCMDPAR1+0x10=event_trb_addr, MIZAR_USB_DEPCMDPAR0+0x10=0x0, MIZAR_USB_DEPCMD+0x10=0x506, poll completion. Wait for interrupts. Call status_stage(). Call setup_stage() for USB_SET_CONFIGURATION_OR_RESET_TT, write marker 0xdeadbee4. Wait for interrupt. Write MIZAR_USB_GUSB2PHYCFG=0x40002547. Check event_counter<=0x4 twice with interrupt waits. Call setup_stage() for GET_DESCRIPTOR USB CONFIGURATION, write marker 0xdeadbee5. Wait for interrupt. Data stage: write event_trb_addr=Buffer_PointerLO, event_trb_addr+0x8=0x9, event_trb_addr+0xc=0x853. Write configuration descriptor to Buffer_PointerLO: 0x003c0209, 0xe0000101, 0x00000032, 0x00000000. Write MIZAR_USB_DEPCMDPAR1+0x10=event_trb_addr, MIZAR_USB_DEPCMDPAR0+0x10=0x0. SET ADDRESS: write MIZAR_USB_DCFG=0x480809. Write event_trb_addr=Buffer_PointerLO, event_trb_addr+0x8=0x0, event_trb_addr+0xc=0x853. Write MIZAR_USB_DEPCMDPAR1+0x10=event_trb_addr, MIZAR_USB_DEPCMDPAR0+0x10=0x0. Bulk transfer: write event_trb_addr=Buffer_PointerLO_1, event_trb_addr+0x8=0x40, event_trb_addr+0xc=0x813. Write MIZAR_USB_DEPCMDPAR1+0x40=event_trb_addr, MIZAR_USB_DEPCMDPAR0+0x40=0x0, MIZAR_USB_DEPCMD+0x40=0x506 (StartXfer on EP4). Poll MIZAR_USB_DEPCMD+0x40 until completion. Wait for interrupt.""",
    "Meta Impacted Registers": "GCTL, DCTL, DALEPENA, DEVTEN, GUCTL, DSTS",
    "Meta Validation / Acceptance Criteria": """1. MIZAR_USB_DEPCMD+(i0x10) polled after writing 0x402 \u2014 DEPXFERCFG command must complete (register value changes from 0x402) for each of 8 endpoints.
2. MIZAR_USB_DEPCMD+0x10 polled after writing 0x506 \u2014 StartXfer command on EP1 must complete (register value changes from 0x506).
3. MIZAR_USB_DEPCMD polled after writing 0x506 in setup_stage() \u2014 StartXfer command must complete.
4. MIZAR_USB_DEPCMD polled after writing 0x506 in status_stage() \u2014 StartXfer command must complete.
5. MIZAR_USB_DEPCMD+trb_address polled after writing cmd in set_configuration() \u2014 command must complete (register value changes from cmd).
6. MIZAR_USB_DEPCMD+0x40 polled after writing 0x506 \u2014 StartXfer command on EP4 for bulk transfer must complete.
7. Default_IRQHandler reads event_count = read_reg(MIZAR_USB_GEVNTCOUNT), stores to event_counter, writes back event_count to MIZAR_USB_GEVNTCOUNT to acknowledge events.
8. Default_IRQHandler checks rd_data && 0x80000000 from MIZAR_LSS_SYSREG_RAW_STCR0 to confirm USB interrupt source (bit 31).
9. Default_IRQHandler writes MIZAR_LSS_SYSREG_RAW_STCR0=0x80000000 to acknowledge interrupt.
10. GIC_ClearIRQ(84) called in Default_IRQHandler to clear IRQ 84.
11. event_counter checked against 0x4 threshold \u2014 conditional interrupt waits based on event count.
12. Configuration descriptor data written to Buffer_PointerLO: wTotalLength=0x003c (60 bytes), bNumInterfaces=1, bConfigurationValue=1.
13. Bulk transfer TRB configured with Buffer_PointerLO_1, length=0x40 (64 bytes), control=0x813.""",
    "Meta Test Steps / Procedure": """1. extern int int_pend;
2. int event_counter;
3. int test_case() {
4. int rd_data, wr_data, port_count, db_offset;
5. int buf_data[16], i, bulk, intr, event_comletion, j;
6. int hand_shake;
7. nic_programming();
8. GIC_EnableAllIRQ();
9. for(j=0; j<20; j++) {
10. write_reg(Buffer_PointerLO + jDWORD, 0x0);
11. write_reg(event_trb_addr + jDWORD, 0x0);
12. }
13. write_reg(MIZAR_USB_GEVNTADRLO, Default_Event_Ring_Array);
14. write_reg(MIZAR_USB_GEVNTADRHI, 0x0);
15. write_reg(MIZAR_USB_GEVNTSIZ, 0x30);
16. write_reg(MIZAR_USB_GEVNTCOUNT, 0x0);
17. rd_data = read_reg(MIZAR_USB_GCTL);
18. wr_data = 0x30c12214;
19. write_reg(MIZAR_USB_GCTL, wr_data);
20. rd_data = read_reg(MIZAR_USB_DCFG);
21. write_reg(MIZAR_USB_DCFG, 0x480801);
22. write_reg(MIZAR_USB_DEVTEN, 0x1f);
23. rd_data = read_reg(MIZAR_USB_GUCTL);
24. write_reg(MIZAR_USB_GUCTL, 0xa400010);
25. set_configuration(0, 0, 0, 0x409);
26. for(i=0; i<8; i++) {
27. write_reg(MIZAR_USB_DEPCMDPAR0 + (i0x10), 0x1);
28. write_reg(MIZAR_USB_DEPCMD + (i0x10), 0x402);
29. rd_data = read_reg(MIZAR_USB_DEPCMD + (i0x10));
30. while(rd_data == 0x402) { wait_on(10); rd_data = read_reg(MIZAR_USB_DEPCMD + (i0x10)); }
31. }
32. write_reg(MIZAR_USB_DALEPENA, 0x3);
33. write_reg(MIZAR_USB_DCTL, 0x80f00000);
34. write_reg(MIZAR_LSS_SYSREG_INTR_EN0, 0x80000000);
35. int_pend = 1;
36. while(int_pend) { wait_on(5); }
37. write_reg(MIZAR_USB_DEPCMDPAR1 + 0x10, event_trb_addr);
38. write_reg(MIZAR_USB_DEPCMDPAR0 + 0x10, 0x0);
39. write_reg(MIZAR_USB_DEPCMD + 0x10, 0x506);
40. rd_data = read_reg(MIZAR_USB_DEPCMD + 0x10);
41. while(rd_data == 0x506) { wait_on(10); rd_data = read_reg(MIZAR_USB_DEPCMD + 0x10); }
42. int_pend = 1;
43. while(int_pend) { wait_on(5); }
44. int_pend = 1;
45. while(int_pend) { wait_on(5); }
46. int_pend = 1;
47. while(int_pend) { wait_on(5); }
48. status_stage();
49. setup_stage();
50. write_reg(0xA0243ffc, 0xdeadbee4);
51. int_pend = 1;
52. while(int_pend) { wait_on(5); }
53. setup_stage();
54. write_reg(MIZAR_USB_GUSB2PHYCFG, 0x40002547);
55. if(event_counter <= 0x4) { int_pend = 1; while(int_pend) { wait_on(5); } }
56. if(event_counter <= 0x4) { int_pend = 1; while(int_pend) { wait_on(5); } }
57. setup_stage();
58. write_reg(0xA0243ffc, 0xdeadbee5);
59. int_pend = 1;
60. while(int_pend) { wait_on(5); }
61. write_reg(event_trb_addr, Buffer_PointerLO);
62. write_reg(event_trb_addr + 0x8, 0x9);
63. write_reg(event_trb_addr + 0xc, 0x853);
64. write_reg(Buffer_PointerLO, 0x003c0209);
65. write_reg(Buffer_PointerLO + 0x4, 0xe0000101);
66. write_reg(Buffer_PointerLO + 0x8, 0x00000032);
67. write_reg(Buffer_PointerLO + 0xc, 0x00000000);
68. write_reg(MIZAR_USB_DEPCMDPAR1 + 0x10, event_trb_addr);
69. write_reg(MIZAR_USB_DEPCMDPAR0 + 0x10, 0x0);
70. write_reg(MIZAR_USB_DCFG, 0x480809);
71. write_reg(event_trb_addr, Buffer_PointerLO);
72. write_reg(event_trb_addr + 0x8, 0x0);
73. write_reg(event_trb_addr + 0xc, 0x853);
74. write_reg(MIZAR_USB_DEPCMDPAR1 + 0x10, event_trb_addr);
75. write_reg(MIZAR_USB_DEPCMDPAR0 + 0x10, 0x0);
76. write_reg(event_trb_addr, Buffer_PointerLO_1);
77. write_reg(event_trb_addr + 0x8, 0x40);
78. write_reg(event_trb_addr + 0xc, 0x813);
79. write_reg(MIZAR_USB_DEPCMDPAR1 + 0x40, event_trb_addr);
80. write_reg(MIZAR_USB_DEPCMDPAR0 + 0x40, 0x0);
81. write_reg(MIZAR_USB_DEPCMD + 0x40, 0x506);
82. rd_data = read_reg(MIZAR_USB_DEPCMD + 0x40);
83. while(rd_data == 0x506) { wait_on(10); rd_data = read_reg(MIZAR_USB_DEPCMD + 0x40); }
84. int_pend = 1;
85. while(int_pend) { wait_on(5); }
86. }
87. void Default_IRQHandler() {
88. int rd_data, sysreg_rd_data, event_count;
89. int_pend = 0;
90. rd_data = read_reg(MIZAR_LSS_SYSREG_MSK_STS0);
91. rd_data = read_reg(MIZAR_LSS_SYSREG_RAW_STCR0);
92. event_count = read_reg(MIZAR_USB_GEVNTCOUNT);
93. event_counter = event_count;
94. write_reg(MIZAR_USB_GEVNTCOUNT, event_count);
95. if(rd_data && 0x80000000) {
96. write_reg(MIZAR_LSS_SYSREG_RAW_STCR0, 0x80000000);
97. }
98. GIC_ClearIRQ(84);
99. }
100. void setup_stage() {
101. int rd_data;
102. write_reg(event_trb_addr, Buffer_PointerLO);
103. write_reg(event_trb_addr + 0x8, 0x8);
104. write_reg(event_trb_addr + 0xc, 0x823);
105. write_reg(MIZAR_USB_DEPCMDPAR1, event_trb_addr);
106. write_reg(MIZAR_USB_DEPCMDPAR0, 0x0);
107. write_reg(MIZAR_USB_DEPCMD, 0x506);
108. rd_data = read_reg(MIZAR_USB_DEPCMD);
109. while(rd_data == 0x506) { wait_on(10); rd_data = read_reg(MIZAR_USB_DEPCMD); }
110. int_pend = 1;
111. while(int_pend) { wait_on(5); }
112. }
113. void status_stage() {
114. int rd_data;
115. write_reg(event_trb_addr, Buffer_PointerLO);
116. write_reg(event_trb_addr + 0x8, 0x0);
117. write_reg(event_trb_addr + 0xc, 0x843);
118. write_reg(MIZAR_USB_DEPCMDPAR1, event_trb_addr);
119. write_reg(MIZAR_USB_DEPCMDPAR0, 0x0);
120. write_reg(MIZAR_USB_DEPCMD, 0x506);
121. rd_data = read_reg(MIZAR_USB_DEPCMD);
122. while(rd_data == 0x506) { wait_on(10); rd_data = read_reg(MIZAR_USB_DEPCMD); }
123. int_pend = 1;
124. while(int_pend) { wait_on(5); }
125. wait_on(5);
126. }
127. void set_configuration(int trb_address, int parameter0, int parameter1, int cmd) {
128. int read_data;
129. write_reg(MIZAR_USB_DEPCMDPAR1 + trb_address, parameter1);
130. write_reg(MIZAR_USB_DEPCMDPAR0 + trb_address, parameter0);
131. write_reg(MIZAR_USB_DEPCMD + trb_address, cmd);
132. read_data = read_reg(MIZAR_USB_DEPCMD + trb_address);
133. while(read_data == cmd) { wait_on(30); read_data = read_reg(MIZAR_USB_DEPCMD + trb_address); }
134. }"""
}

FILENAME = "USB_TestPlan_20261013_002556.xlsx"
OUTPUT_PATH = FILENAME

# ============================================================
# CREATE WORKBOOK
# ============================================================
wb = openpyxl.Workbook()

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_align = Alignment(wrap_text=True, vertical="top")

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = "TestPlan"

tp_headers = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

tp_field_map = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    None  # Code Generation - not in JSON
]

for col_idx, h in enumerate(tp_headers, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

for col_idx, field in enumerate(tp_field_map, 1):
    val = row.get(field, "") if field else ""
    cell = ws_tp.cell(row=2, column=col_idx, value=val)
    cell.alignment = wrap_align

ws_tp.freeze_panes = "A2"

# Auto-size columns
for col_idx in range(1, len(tp_headers) + 1):
    max_len = len(str(ws_tp.cell(row=1, column=col_idx).value or ""))
    data_val = str(ws_tp.cell(row=2, column=col_idx).value or "")
    lines = data_val.split('\n')
    longest_line = max((len(l) for l in lines), default=0)
    col_width = min(max(max_len, longest_line, 12) + 2, 60)
    ws_tp.column_dimensions[get_column_letter(col_idx)].width = col_width

# --- MetaData Sheet ---
ws_md = wb.create_sheet("MetaData")

md_headers = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

md_field_map = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

for col_idx, h in enumerate(md_headers, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

for col_idx, field in enumerate(md_field_map, 1):
    val = row.get(field, "")
    cell = ws_md.cell(row=2, column=col_idx, value=val)
    cell.alignment = wrap_align

ws_md.freeze_panes = "A2"

for col_idx in range(1, len(md_headers) + 1):
    max_len = len(str(ws_md.cell(row=1, column=col_idx).value or ""))
    data_val = str(ws_md.cell(row=2, column=col_idx).value or "")
    lines = data_val.split('\n')
    longest_line = max((len(l) for l in lines), default=0)
    col_width = min(max(max_len, longest_line, 12) + 2, 60)
    ws_md.column_dimensions[get_column_letter(col_idx)].width = col_width

# Set MetaData sheet to veryHidden
ws_md.sheet_state = "veryHidden"

# Save
wb.save(OUTPUT_PATH)
print(f"File saved: {OUTPUT_PATH}")
print(f"File size: {os.path.getsize(OUTPUT_PATH)} bytes")

# Validate
wb2 = openpyxl.load_workbook(OUTPUT_PATH)
assert "TestPlan" in wb2.sheetnames
assert "MetaData" in wb2.sheetnames
assert wb2["TestPlan"].max_row == 2
assert wb2["MetaData"].max_row == 2

# Validate meta content
for col_idx, field in enumerate(md_field_map, 1):
    cell_val = wb2["MetaData"].cell(row=2, column=col_idx).value or ""
    expected = row.get(field, "")
    assert cell_val == expected, f"Mismatch in {field}"

wb2.close()
print("All validations PASSED")
print(f"TestPlan rows: 1")
print(f"MetaData rows: 1")
