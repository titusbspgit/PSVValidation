#!/usr/bin/env python3
"""
PCIE TestPlan XLSX Generator
Generates: PCIE_TestPlan_20250717_183000.xlsx
Run: python3 generate_PCIE_TestPlan_20250717.py
Requires: pip install openpyxl
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ============================================================
# JSON DATA
# ============================================================
json_data = [
    {
        "Index": "1",
        "SS / Module": "PCIE",
        "Test Case Name": "pcie_device_enumerate_test",
        "Feature": "Device Enumeration",
        "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>; \"pcie.h\"",
        "Meta Macros": "DM0_RC; DM1_RC; DM0_EP; DM1_EP; DEBUG_DISPLAY",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "DM0_RC; DM1_RC; DM0_EP; DM1_EP",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs PCIe device enumeration. It begins by writing 0x0 to address 0xE6004100. It then invokes link training for DM0 or DM1 in RC or EP mode (conditionally compiled via DM0_RC, DM1_RC, DM0_EP, DM1_EP) using link_training_dm0_x4(4) or link_training_dm1_x4(4). Cache coherency programming is performed by reading mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF, modifying bit fields [11:14], [3:6], [27:30], [19:22] using set_data() with value 0xf, and writing back for both PCIE0 and PCIE1 instances. A wait_on(20) delay is inserted between cache programming phases. The SII0 register at offset 0xC0 is polled via read_sii0_reg(0xC0) until bits matching mask 0xD1 equal 0xD1, confirming link-up status. The same polling is done for SII1 via read_sii1_reg(0xC0). Under DM0_RC, the Vendor ID is read from pcie_slv0 offset 0x0, command register at offset 0x4 is written with 0x7, and mem_base_program_dm0_x4() and mem_base_program_dm1_x4() are called. Six registers at addresses 0xE690000C, 0xE6900010, 0xE6900014, 0xE6900018, 0xE6900030, 0xE6900034 are written with 0x1. Cache coherency is then disabled by writing 0x0 to bit fields [27:30] and [19:22] for both PCIE0 and PCIE1 coherency control registers. BAR probing is performed on pcie_slv1 and pcie_slv0 at offsets 0x10, 0x14, 0x18, 0x1c, 0x20, 0x24 by writing 0xFFFFFFFF, reading back, then programming with specific base addresses (0x0, 0x4, 0x20000000, 0x40000000, 0x60000000, 0x80000000) and reading back again. Finally, address 0xE6004100 is polled until its value equals 0x12345678, with wait_on(5) delays between iterations, and finish(0) is called on success.",
        "Test Description": "This test performs PCIe device enumeration for dual-mode controllers (DM0/DM1) in Root Complex or Endpoint configurations. It initiates link training in x4 mode, programs cache coherency control registers for both PCIE0 and PCIE1 instances by setting specific bit fields, and polls the SII link status registers until link-up is confirmed. Under Root Complex mode, it reads the Vendor ID from the downstream device, enables memory and bus master commands, and programs memory base addresses for both controllers. It then enables a set of configuration registers. Cache coherency is subsequently disabled for both PCIE instances. BAR (Base Address Register) enumeration is performed on both slave ports by probing all six BARs with all-ones patterns, reading back to determine BAR sizes, and then programming each BAR with specific base addresses. The test concludes by polling a synchronization register until a completion handshake value is received.",
        "Meta Test Steps / Procedure": "1. write_reg(0xE6004100, 0x0) -- initialize synchronization register to zero.\n2. Conditionally call link_training_dm0_x4(4) or link_training_dm1_x4(4) based on DM0_RC, DM1_RC, DM0_EP, DM1_EP compile flags.\n3. CACHE PROGRAMMING: read mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF, set_data bits [11:14]=0xf, [3:6]=0xf, write back.\n4. Read mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF again, set_data bits [27:30]=0xf, [19:22]=0xf, write back.\n5. Repeat steps 3-4 for mizar_PCIE1_DBI_DSP_COHERENCY_CONTROL_3_OFF.\n6. wait_on(20).\n7. Read mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF, set_data bits [11:14]=0xf, [3:6]=0xf, [27:30]=0xf, [19:22]=0xf, write back.\n8. Repeat step 7 for mizar_PCIE1_DBI_DSP_COHERENCY_CONTROL_3_OFF.\n9. read_sii0_reg(0xC0) -- read SII0 link status.\n10. Poll read_sii0_reg(0xC0) in while loop until (data_rd & 0xD1) == 0xD1.\n11. read_sii1_reg(0xC0) -- poll SII1 link status until (data_rd & 0xD1) == 0xD1.\n12. Under DM0_RC: read_pcie_slv0_reg(0x0) to get Vendor ID; write_pcie_slv0_reg(0x4, 0x7) to enable command; call mem_base_program_dm0_x4() and mem_base_program_dm1_x4(); wait_on(10).\n13. Write 0x1 to 0xE690000C, 0xE6900010, 0xE6900014, 0xE6900018, 0xE6900030, 0xE6900034.\n14. DISABLE_CACHE PROGRAMMING: read mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF, set_data bits [11:14]=0xf, [3:6]=0xf, write back; then set bits [27:30]=0xf, [19:22]=0x0, write back.\n15. Repeat step 14 for mizar_PCIE1_DBI_DSP_COHERENCY_CONTROL_3_OFF.\n16. wait_on(10).\n17. Read mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF, set bits [11:14]=0xf, [3:6]=0xf, [27:30]=0x0, [19:22]=0x0, write back.\n18. Repeat step 17 for mizar_PCIE1_DBI_DSP_COHERENCY_CONTROL_3_OFF.\n19. wait_on(30).\n20. BAR probing on pcie_slv1: write 0xFFFFFFFF to offsets 0x10-0x24, read back all, write specific base addresses, read back all.\n21. BAR probing on pcie_slv0: same sequence as step 20.\n22. wait_on(10), then poll read_reg(0xE6004100) until value == 0x12345678 with wait_on(5) between iterations.\n23. Call finish(0).",
        "Test Steps / Procedure": "1. Initialize the synchronization register to zero.\n2. Perform PCIe link training in x4 mode for the selected controller (DM0 or DM1) in the configured role (Root Complex or Endpoint).\n3. Enable cache coherency for PCIE0 by reading the coherency control register, setting bit fields [11:14], [3:6], [27:30], and [19:22] to enable values, and writing back.\n4. Enable cache coherency for PCIE1 using the same bit-field programming as PCIE0.\n5. Wait for coherency settings to take effect.\n6. Re-apply full cache coherency enable for both PCIE0 and PCIE1 coherency control registers.\n7. Poll the SII0 link status register until the link-up condition (mask match) is confirmed.\n8. Poll the SII1 link status register until the link-up condition is confirmed.\n9. In Root Complex mode, read the Vendor ID from the downstream device on slave port 0.\n10. Enable memory space access and bus master in the downstream device command register.\n11. Program memory base addresses for both DM0 and DM1 controllers.\n12. Enable a set of six PCIe configuration registers by writing enable values.\n13. Disable cache coherency for both PCIE0 and PCIE1 by clearing the relevant bit fields in the coherency control registers.\n14. Wait for cache disable to take effect, then re-confirm cache disable for both instances.\n15. Probe all six BARs on slave port 1 by writing all-ones, reading back BAR sizes, then programming specific base addresses and verifying.\n16. Probe all six BARs on slave port 0 using the same enumeration sequence.\n17. Poll the synchronization register until the expected completion handshake value is received.\n18. Complete the test successfully.",
        "Meta Impacted Registers": "0xE6004100; mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF; mizar_PCIE1_DBI_DSP_COHERENCY_CONTROL_3_OFF; 0xE690000C; 0xE6900010; 0xE6900014; 0xE6900018; 0xE6900030; 0xE6900034",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "1. SII0 link status polling: read_sii0_reg(0xC0) is polled until (data_rd & 0xD1) == 0xD1, confirming PCIE0 link-up.\n2. SII1 link status polling: read_sii1_reg(0xC0) is polled until (data_rd & 0xD1) == 0xD1, confirming PCIE1 link-up.\n3. Vendor ID read: read_pcie_slv0_reg(0x0) returns the device Vendor ID (printed for verification).\n4. BAR probing: write 0xFFFFFFFF to BAR offsets 0x10-0x24 on pcie_slv1 and pcie_slv0, read back to determine BAR sizes, then write specific base addresses and read back to confirm programming.\n5. Synchronization handshake: read_reg(0xE6004100) is polled until value equals 0x12345678.\n6. Test passes by calling finish(0) after all polling conditions are met.",
        "Validation / Acceptance Criteria": "1. The SII0 link status register must indicate link-up by matching the expected bit pattern before proceeding.\n2. The SII1 link status register must indicate link-up by matching the expected bit pattern before proceeding.\n3. The Vendor ID read from the downstream PCIe device on slave port 0 must return a valid value.\n4. All six BARs on both slave ports must respond correctly to the all-ones probe pattern and accept the programmed base addresses on read-back.\n5. The synchronization register must eventually return the expected completion handshake value.\n6. The test completes successfully by calling the finish routine with a pass status.",
        "Remarks": "The test uses conditional compilation (DM0_RC, DM1_RC, DM0_EP, DM1_EP) to select the controller and role. Cache coherency is enabled before enumeration and disabled afterward. SII link status polling uses a busy-wait loop. BAR enumeration follows the standard PCI probe sequence. All nine Agent 4 register mappings are unresolved because the register specification provided (gp0_autoreg.xlsx) covers GPIO, not PCIE."
    },
    {
        "Index": "2",
        "SS / Module": "PCIE",
        "Test Case Name": "pcie_dma_write_test",
        "Feature": "DMA Write and Read",
        "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>; \"pcie.h\"",
        "Meta Macros": "DM0_RC; DM1_RC; DM0_EP; DM1_EP; DEBUG_DISPLAY",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "DM0_RC; DM1_RC; DM0_EP; DM1_EP",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs PCIe DMA write and read operations across 4 channels with interrupt-based completion handling.",
        "Test Description": "This test validates PCIe DMA write and read data transfer operations across all four DMA channels using interrupt-driven completion. It performs link training, polls the SII link status register until link-up is confirmed, reads the downstream device Vendor ID, programs BARs and memory base addresses, and then preloads source memory with known data patterns. The GIC is configured to enable DMA interrupts. For each of the four DMA channels, a write transfer is programmed from local source memory to the PCIe write address, triggered via the DMA write doorbell register, and the test waits for the interrupt handler to signal completion. After all four write channels complete, four read-back transfers are similarly programmed. The test completes successfully after all DMA transfers finish.",
        "Meta Test Steps / Procedure": "1. write_reg(0xE6004100, 0x0)\n2. Link training\n3. SII polling\n4. Vendor ID read, BAR/mem base programming\n5. Sync handshake polling\n6. Source memory preload with 0xC0DEBEED and 0xF00DDEAF\n7. GIC setup\n8. DMA write/read across 4 channels with interrupt completion",
        "Test Steps / Procedure": "1. Initialize the synchronization register to zero.\n2. Perform PCIe link training in x4 mode for the selected controller.\n3. Poll the SII link status register until link-up is confirmed.\n4. Read the Vendor ID, enable commands, program BARs and memory base addresses.\n5. Wait for synchronization handshake.\n6. Configure DMA transfer parameters.\n7. Preload source memory with known data patterns.\n8. Configure GIC for interrupt handling.\n9. Unmask DMA interrupts.\n10. Execute DMA write transfers on all 4 channels with interrupt completion.\n11. Execute DMA read transfers on all 4 channels with interrupt completion.\n12. Interrupt handler clears status and signals completion.\n13. Complete the test successfully.",
        "Meta Impacted Registers": "0xE6004100; mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF; mizar_PCIE1_DBI_DSP_COHERENCY_CONTROL_3_OFF; mizar_PCIE0_DBI_DSP_DMA_WRITE_INT_MASK_OFF; mizar_PCIE0_DBI_DSP_DMA_READ_INT_MASK_OFF; mizar_PCIE0_DBI_DSP_DMA_WRITE_DOORBELL_OFF; mizar_PCIE0_DBI_DSP_DMA_READ_DOORBELL_OFF; mizar_PCIE0_DBI_DSP_DMA_WRITE_INT_STATUS_OFF; mizar_PCIE0_DBI_DSP_DMA_READ_INT_STATUS_OFF; mizar_PCIE0_DBI_DSP_DMA_WRITE_INT_CLEAR_OFF; mizar_PCIE0_DBI_DSP_DMA_READ_INT_CLEAR_OFF; mizar_PCIE0_DBI_DSP_MSI_CAP_OFF_08H_REG; mizar_PCIE0_DBI_DSP_MSI_CAP_OFF_10H_REG; mizar_PCIE0_DBI_DSP_FILTER_MASK_2_OFF; mizar_PCIE0_DBI_DSP_AXI_MSTR_MSG_ADDR_HIGH_OFF; mizar_PCIE0_DBI_DSP_UTILITY_OFF; 0xE68860B8; 0xE68862B8; 0xE68864B8; 0xE68A60B8; 0xE68A62B8; 0xE68A64B8",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "1. SII link-up confirmed\n2. Sync handshake received\n3. Source memory preloaded\n4. All 4 DMA write channels complete via interrupt\n5. All 4 DMA read channels complete via interrupt\n6. Interrupt handler correctly clears status\n7. GIC interrupts cleared\n8. finish(0) called",
        "Validation / Acceptance Criteria": "1. SII link status must indicate link-up.\n2. Synchronization register must return expected handshake value.\n3. Source memory must be correctly preloaded.\n4. Each DMA write channel must complete successfully.\n5. Each DMA read channel must complete successfully.\n6. Interrupt handler must correctly clear interrupts.\n7. GIC interrupt must be properly cleared.\n8. Test completes successfully.",
        "Remarks": "DMA transfers use interrupt-driven completion. The DM1_RC path uses PCIE1 macros with GIC IRQ 0x23. All register mappings are unresolved (GPIO spec, not PCIE)."
    },
    {
        "Index": "3",
        "SS / Module": "PCIE",
        "Test Case Name": "pcie_mem_wr_rd_test",
        "Feature": "Memory Write and Read",
        "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>; \"pcie.h\"",
        "Meta Macros": "DM0_RC; DM1_RC; DM0_EP; DM1_EP; DM0; DM1; DEBUG_DISPLAY",
        "Meta Arrays": "NA",
        "Speed": "NA",
        "Mode": "DM0_RC; DM1_RC; DM0_EP; DM1_EP",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs PCIe memory write and read operations via slave ports for both RC and EP configurations.",
        "Test Description": "This test validates PCIe memory write and read operations through the PCIe slave ports for both Root Complex and Endpoint configurations. It performs link training, programs cache coherency, polls SII link status, programs BARs and memory bases, then performs memory write/read operations at multiple offsets with distinct data patterns. The test concludes by polling a synchronization register.",
        "Meta Test Steps / Procedure": "1. write_reg(0xE6004100, 0x0)\n2. Link training\n3. Cache coherency programming\n4. SII polling\n5. BAR/mem base programming\n6. Cache disable\n7. Memory write/read operations with various patterns\n8. Sync handshake polling\n9. finish(0)",
        "Test Steps / Procedure": "1. Initialize synchronization register.\n2. Perform link training.\n3. Enable cache coherency for PCIE0 and PCIE1.\n4. Poll SII link status until link-up.\n5. Program BARs and memory base addresses.\n6. Signal readiness.\n7. Disable cache coherency.\n8. Perform memory write/read operations.\n9. Poll synchronization register.\n10. Complete test successfully.",
        "Meta Impacted Registers": "0xE6004100; mizar_PCIE0_DBI_DSP_COHERENCY_CONTROL_3_OFF; mizar_PCIE1_DBI_DSP_COHERENCY_CONTROL_3_OFF",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "1. SII link-up confirmed\n2. Vendor ID valid\n3. Memory write/read patterns match\n4. Sync handshake received\n5. finish(0) called",
        "Validation / Acceptance Criteria": "1. SII link status must indicate link-up.\n2. Vendor ID must be valid (RC mode).\n3. All memory write/read operations must pass.\n4. Synchronization register must return expected value.\n5. Test completes successfully.",
        "Remarks": "Cache coherency is enabled before enumeration and disabled before memory operations. The pcie_slv*_mem_wr_rd() functions perform write followed by read-back comparison internally. All register mappings are unresolved (GPIO spec, not PCIE)."
    },
    {
        "Index": "4",
        "SS / Module": "PCIE",
        "Test Case Name": "pcie_reg_wr_rd_test",
        "Feature": "Register Write and Read",
        "Meta Headers": "<stdlib.h>; <stdio.h>; <test_common.h>; <pcie.h>",
        "Meta Macros": "NA",
        "Meta Arrays": "rc0_ctl_addr[5]; rc1_ctl_addr[5]; ctl_default[5]; sii0_addr[3]; sii1_addr[3]; sii_default[3]; sii0_write_mask[3]; sii1_write_mask[3]; phy0_addr[3]; phy1_addr[3]; phy0_default[3]; phy1_default[3]; phy0_write_mask[3]; phy1_write_mask[3]; chk_val[6]; chk_val_phy[3]",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase performs register reset value verification and write/read-back testing for DBI, SII, and PHY registers across PCIE0 and PCIE1.",
        "Test Description": "This test verifies the reset default values and write/read-back functionality of PCIE registers across three register categories: DBI controller registers, SII registers, and PHY registers, for both PCIE0 and PCIE1 instances. The test passes only if all checks succeed with zero errors.",
        "Meta Test Steps / Procedure": "1. chk_rst_val() for DBI, SII, PHY registers\n2. PHY reset control write\n3. chk_rd_wr() with patterns 0xFFFFFFFF, 0xAAAAAAAA, 0x55555555\n4. PHY patterns 0x7BAF, 0x1, 0x003B\n5. finish(err2 || err1)",
        "Test Steps / Procedure": "1. Read all DBI registers and verify reset defaults.\n2. Read all SII registers and verify reset defaults.\n3. Take PHY out of reset.\n4. Read all PHY registers and verify reset defaults.\n5. Write test patterns to all registers.\n6. Read back and verify all registers.\n7. Repeat for all test patterns.\n8. Verify zero errors.",
        "Meta Impacted Registers": "mizar_PCIE0_DBI_DSP_MSI_CAP_OFF_08H_REG; mizar_PCIE0_DBI_DSP_MSI_CAP_OFF_10H_REG; mizar_PCIE0_DBI_DSP_FILTER_MASK_2_OFF; mizar_PCIE0_DBI_DSP_AXI_MSTR_MSG_ADDR_HIGH_OFF; mizar_PCIE0_DBI_DSP_UTILITY_OFF; mizar_PCIE1_DBI_DSP_MSI_CAP_OFF_08H_REG; mizar_PCIE1_DBI_DSP_MSI_CAP_OFF_10H_REG; mizar_PCIE1_DBI_DSP_FILTER_MASK_2_OFF; mizar_PCIE1_DBI_DSP_AXI_MSTR_MSG_ADDR_HIGH_OFF; mizar_PCIE1_DBI_DSP_UTILITY_OFF; mizar_PCIE0_SII_PCIE0_TRANSMIT_HEADER2; mizar_PCIE0_SII_PCIE0_TRANSMIT_HEADER3; mizar_PCIE0_SII_PHY_CONTROL_23; mizar_PCIE1_SII_PCIE1_TRANSMIT_HEADER2; mizar_PCIE1_SII_PCIE1_TRANSMIT_HEADER3; mizar_PCIE1_SII_PHY_CONTROL_23; mizar_PCIE0_SII_PHY_RST_CONTROL; mizar_PCIE1_SII_PHY_RST_CONTROL; 0xE68860B8; 0xE68862B8; 0xE68864B8; 0xE68A60B8; 0xE68A62B8; 0xE68A64B8",
        "Impacted Registers": "NA",
        "Meta Validation / Acceptance Criteria": "1. All reset values match expected defaults (0x0)\n2. All write/read-back patterns match\n3. SII write masks applied correctly\n4. PHY alignment-based extraction correct\n5. Zero errors across both PCIE instances",
        "Validation / Acceptance Criteria": "1. All DBI registers must read back reset default of zero.\n2. All SII registers must read back reset default of zero.\n3. All PHY registers must read back reset default of zero.\n4. Write/read-back must match for all patterns.\n5. Test passes with zero errors.",
        "Remarks": "The test covers DBI (5 per instance), SII (3 per instance), and PHY (3 per instance) registers. PHY registers use 13-bit write mask and alignment-based extraction. All register mappings are unresolved (GPIO spec, not PCIE)."
    }
]

# ============================================================
# COLUMN DEFINITIONS
# ============================================================
TESTPLAN_COLUMNS = [
    "Index", "SS / Module", "Feature", "Test Case Name",
    "Test Description", "Speed", "Mode", "Memory Start Offset",
    "Memory End Offset", "Remarks", "Test Steps / Procedure",
    "Impacted Registers", "Validation / Acceptance Criteria", "Code Generation"
]

METADATA_COLUMNS = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

# ============================================================
# WORKBOOK GENERATION
# ============================================================
def generate_workbook(output_path):
    wb = Workbook()

    # -- TestPlan Sheet --
    ws_tp = wb.active
    ws_tp.title = "TestPlan"

    # -- MetaData Sheet --
    ws_md = wb.create_sheet("MetaData")

    # Header styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Write TestPlan headers
    for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
        cell.border = thin_border

    # Write MetaData headers
    for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
        cell.border = thin_border

    # Populate data rows
    for row_idx, item in enumerate(json_data, 2):
        # TestPlan row
        for col_idx, col_name in enumerate(TESTPLAN_COLUMNS, 1):
            value = item.get(col_name, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
            cell.border = thin_border

        # MetaData row
        for col_idx, col_name in enumerate(METADATA_COLUMNS, 1):
            value = item.get(col_name, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment
            cell.border = thin_border

    # Auto-size columns
    for ws in [ws_tp, ws_md]:
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        max_length = max(max_length, len(line))
            adjusted_width = min(max(max_length + 2, 12), 60)
            ws.column_dimensions[col_letter].width = adjusted_width

    # Freeze first row
    ws_tp.freeze_panes = "A2"
    ws_md.freeze_panes = "A2"

    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = "veryHidden"

    # Save
    wb.save(output_path)
    print(f"Workbook saved: {output_path}")
    print(f"File size: {os.path.getsize(output_path)} bytes")
    return output_path

if __name__ == "__main__":
    output = generate_workbook("PCIE_TestPlan_20250717_183000.xlsx")
    print(f"Generated: {output}")
