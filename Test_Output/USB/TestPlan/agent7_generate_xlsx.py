#!/usr/bin/env python3
"""Agent 7 - USB TestPlan Excel Generator
Generates a real .xlsx workbook using openpyxl with TestPlan and MetaData sheets.
IST timestamp used for filename.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import os
import json
import base64
import sys

# IST timezone offset
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp_str = now_ist.strftime('%Y%m%d_%H%M%S')
filename = f'USB_TestPlan_{timestamp_str}.xlsx'

# Test case data
test_cases = [
    {
        "Index": "1",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Bulk_Transfer_test",
        "Feature": "USB Full-Speed Device Bulk Transfer",
        "Meta Headers": '<stdio.h>; <stdlib.h>; \"usb.h\"',
        "Meta Macros": "NA",
        "Meta Arrays": "int buf_data[16]; /* declared local array, not statically initialized, used as local variable */",
        "Speed": "Full-Speed",
        "Mode": "Device",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Full-Speed Device mode Bulk Transfer operation. The test performs the following high-level sequence: NIC programming and GIC IRQ enable, clearing of Buffer_PointerLO and event_trb_addr memory regions (20 DWORDs each), USB controller soft reset via MIZAR_USB_DCTL and polling for reset completion, USB2 PHY configuration via MIZAR_USB_GUSB2PHYCFG, event ring setup (GEVNTADRLO, GEVNTADRHI, GEVNTSIZ, GEVNTCOUNT), global controller configuration via MIZAR_USB_GCTL, device configuration via MIZAR_USB_DCFG, device event enable via MIZAR_USB_DEVTEN, GUCTL configuration, endpoint configuration using set_configuration() for 9 endpoints (START_NEW_CONFIGURATION + 8 endpoint configs), endpoint TX resource configuration loop (8 iterations writing DEPCMDPAR0 and DEPCMD per endpoint with polling), enabling physical endpoints 0 and 1 via DALEPENA, starting device controller run via DCTL, enabling sysreg interrupt, waiting for link state connect/reset events via interrupt waits, enumeration sequence including: setup_stage for control transfers, SET_ADDRESS phase with TRB and DEPCMD operations, handshake polling at 0xa0243ff4, full enumeration() call which performs GET_DEVICE_DESCRIPTOR (18 bytes device descriptor data), GET_CONFIGURATION_DESCRIPTOR (9 bytes short read then 60 bytes full read with complete configuration descriptor data), SET_CONFIGURATION with status stage, handshake polling at 0xa0243ff8, then bulk OUT transfer TRB setup on physical endpoint 4 (offset 0x40) with Buffer_PointerLO_1 and size 0x40, followed by bulk IN transfer TRB setup on physical endpoint 5 (offset 0x50) with Buffer_PointerLO_1 and size 0x40, with interrupt waits after each transfer, and final marker write 0xdeadbee7 to 0xA0243ffc before finish(0). The IRQ handler (Default_IRQHandler) clears int_pend, reads sysreg masked status and raw status/clear registers, reads and writes back GEVNTCOUNT to acknowledge events, clears sysreg interrupt bit 31, and clears GIC IRQ 84.",
        "Test Description": "Verify USB Full-Speed Device mode Bulk Transfer by performing controller soft reset, PHY and event ring configuration, endpoint configuration for control and bulk endpoints, device enumeration (GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION), SET_ADDRESS, and executing bulk OUT and bulk IN transfers with interrupt-driven completion and handshake synchronization.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.\n3. Loop j=0 to 19: write_reg(Buffer_PointerLO + j*DWORD, 0x0); write_reg(event_trb_addr + j*DWORD, 0x0);\n4. wr_data = 0x40f00000; write_reg(MIZAR_USB_DCTL, 0x40f00000) - soft reset.\n5. rd_data = read_reg(MIZAR_USB_DCTL); while(rd_data != 0xf00000) { wait_on(5); rd_data = read_reg(MIZAR_USB_DCTL); } - poll for soft reset completion.\n6-74. [Full procedure as specified in test case]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Clear buffer pointer and event TRB address memory regions (20 DWORDs each).\n3. Perform USB controller soft reset via DCTL and poll for reset completion.\n4. Configure USB2 PHY via GUSB2PHYCFG.\n5. Set up event ring by configuring event address low, event address high, event size, and event count registers.\n6. Configure global controller register GCTL for port direction.\n7. Configure device configuration register DCFG.\n8. Enable device events via DEVTEN register.\n9. Configure GUCTL register.\n10. Issue START_NEW_CONFIGURATION command and configure 8 endpoints using endpoint command registers with parameter registers.\n11. Configure TX resources for 8 endpoints by issuing transfer resource commands and polling for completion.\n12. Enable physical endpoints 0 and 1 via DALEPENA.\n13. Start device controller run via DCTL.\n14. Enable system register interrupt.\n15. Wait for link state connect/reset events via interrupt-driven waits.\n16. Re-configure DCFG and wait for additional events if needed.\n17. Write enumeration marker, read DCFG and DSTS, reconfigure DCFG and DCTL, enable all endpoints via DALEPENA.\n18. Execute setup stage for initial control transfer and wait for completion.\n19. Reconfigure USB2 PHY and wait for events.\n20. Set device address in DCFG, prepare status TRB, issue endpoint command on EP1 IN, and poll for completion.\n21. Wait for interrupt events and poll handshake register for synchronization.\n22. Execute full enumeration sequence: GET_DEVICE_DESCRIPTOR with 18-byte device descriptor data response, GET_CONFIGURATION_DESCRIPTOR short read (9 bytes) and full read (60 bytes) with complete configuration descriptor data, SET_CONFIGURATION with status data stage.\n23. Poll second handshake register for synchronization.\n24. Prepare and execute Bulk OUT transfer on physical endpoint 4 with 64-byte TRB using endpoint command, wait for interrupt completion.\n25. Prepare and execute Bulk IN transfer on physical endpoint 5 with 64-byte TRB using endpoint command, wait for interrupt completion.\n26. Write final completion marker and wait for final interrupt.\n27. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; 0xa0243ff4; 0xa0243ff8; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Meta Validation / Acceptance Criteria": "1. Soft reset completion: read_reg(MIZAR_USB_DCTL) must return 0xf00000 (bit 30 cleared) to exit polling loop.\n2. Endpoint command completion for set_configuration: read_reg(MIZAR_USB_DEPCMD+trb_address) must not equal the issued cmd value to exit polling loop.\n3-20. [Full validation criteria as specified]",
        "Validation / Acceptance Criteria": "1. USB controller soft reset completes successfully as indicated by DCTL register clearing the soft reset bit.\n2. All endpoint configuration commands complete successfully by polling the endpoint command register until the command active bit clears.\n3. All endpoint TX resource configuration commands complete for all 8 endpoints.\n4. Link state connect and reset events are received via interrupt-driven completion.\n5. Handshake synchronization registers return non-zero values indicating host-side readiness.\n6. Device enumeration completes successfully including GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short and full reads), and SET_CONFIGURATION control transfers.\n7. Each control transfer stage (setup, data, status) completes with endpoint command completion and interrupt-driven event notification.\n8. Bulk OUT transfer on physical endpoint 4 completes with 64-byte TRB and interrupt notification.\n9. Bulk IN transfer on physical endpoint 5 completes with 64-byte TRB and interrupt notification.\n10. IRQ handler correctly reads and acknowledges event count, clears system register interrupt, and clears GIC IRQ.\n11. Test completes with finish(0) indicating successful pass.",
        "Remarks": "The testcase uses interrupt-driven synchronization with int_pend flag cleared by Default_IRQHandler. Multiple conditional interrupt waits depend on event_counter <= 0x4. Handshake polling at addresses 0xa0243ff4 and 0xa0243ff8 provides host-side synchronization. The IRQ handler condition uses logical AND (&&) instead of bitwise AND (&) for bit 31 check on MIZAR_LSS_SYSREG_RAW_STCR0, which may be a source code bug. Buffer_PointerLO, Buffer_PointerLO_1, event_trb_addr, Default_Event_Ring_Array, DWORD, and several MIZAR_USB register macros are defined externally in usb.h which is not available in the testcase folder. The status_stage() function contains two separate interrupt waits with a wait_on(5) delay between them."
    },
    {
        "Index": "2",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Interrupt_Transfer_test",
        "Feature": "USB Full-Speed Device Interrupt Transfer",
        "Meta Headers": '<stdio.h>; <stdlib.h>; \"usb.h\"',
        "Meta Macros": "NA",
        "Meta Arrays": "int buf_data[16]; /* declared local array, not statically initialized */",
        "Speed": "Full-Speed",
        "Mode": "Device",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Full-Speed Device mode Interrupt Transfer operation. The test performs: NIC programming and GIC IRQ enable, clearing of Buffer_PointerLO and event_trb_addr memory regions (20 DWORDs each), USB controller soft reset via MIZAR_USB_DCTL and polling for reset completion (polling with wait_on(100)), USB2 PHY configuration via MIZAR_USB_GUSB2PHYCFG, event ring setup, global controller configuration, device configuration, endpoint configuration, enumeration sequence, then Interrupt OUT transfer TRB setup on physical endpoint 2 (offset 0x20) with Buffer_PointerLO_1 and size 0x40 and TRB control 0x815, followed by Interrupt IN transfer TRB setup on physical endpoint 3 (offset 0x30) with Buffer_PointerLO_1 and size 0x40 and TRB control 0x813, with DSTS frame number polling.",
        "Test Description": "Verify USB Full-Speed Device mode Interrupt Transfer by performing controller soft reset, PHY and event ring configuration, endpoint configuration for control and interrupt endpoints, device enumeration (GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION), SET_ADDRESS, and executing interrupt OUT and interrupt IN transfers with interrupt-driven completion, frame number validation, and handshake synchronization.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.\n3-74. [Full procedure as specified in test case]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Clear buffer pointer and event TRB address memory regions (20 DWORDs each).\n3. Perform USB controller soft reset via DCTL and poll for reset completion.\n4. Configure USB2 PHY via GUSB2PHYCFG.\n5. Set up event ring by configuring event address low, event address high, event size, and event count registers.\n6. Configure global controller register GCTL for port direction.\n7. Configure device configuration register DCFG.\n8. Enable device events via DEVTEN register.\n9. Configure GUCTL register.\n10. Issue START_NEW_CONFIGURATION command and configure 8 endpoints using endpoint command registers with parameter registers.\n11. Configure TX resources for 8 endpoints by issuing transfer resource commands and polling for completion.\n12. Enable physical endpoints 0 and 1 via DALEPENA.\n13. Start device controller run via DCTL.\n14. Enable system register interrupt.\n15. Wait for link state connect/reset events via interrupt-driven waits.\n16. Re-configure DCFG and wait for additional events if needed.\n17. Write enumeration marker, read DCFG and DSTS, reconfigure DCFG and DCTL, enable all endpoints via DALEPENA.\n18. Execute setup stage for initial control transfer and wait for completion.\n19. Reconfigure USB2 PHY and wait for events.\n20. Set device address in DCFG, prepare status TRB, issue endpoint command on EP1 IN, and poll for completion.\n21. Wait for interrupt events and poll handshake register for synchronization.\n22. Execute full enumeration sequence: GET_DEVICE_DESCRIPTOR with 18-byte device descriptor data response, GET_CONFIGURATION_DESCRIPTOR short read (9 bytes) and full read (60 bytes) with complete configuration descriptor data, SET_CONFIGURATION with status data stage.\n23. Poll second handshake register for synchronization.\n24. Prepare and execute Interrupt OUT transfer on physical endpoint 2 with 64-byte interrupt TRB using endpoint command, wait for interrupt completion.\n25. Prepare and execute Interrupt IN transfer on physical endpoint 3 with 64-byte TRB using endpoint command, wait for interrupt completion.\n26. Poll DSTS register for frame number to become non-zero.\n27. Write final completion marker and wait for final interrupt.\n28. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; 0xa0243ff4; 0xa0243ff8; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Meta Validation / Acceptance Criteria": "1. Soft reset completion: read_reg(MIZAR_USB_DCTL) must return 0xf00000.\n2-21. [Full validation criteria as specified]",
        "Validation / Acceptance Criteria": "1. USB controller soft reset completes successfully as indicated by DCTL register clearing the soft reset bit.\n2. All endpoint configuration commands complete successfully by polling the endpoint command register until the command active bit clears.\n3. All endpoint TX resource configuration commands complete for all 8 endpoints.\n4. Link state connect and reset events are received via interrupt-driven completion.\n5. Handshake synchronization registers return non-zero values indicating host-side readiness.\n6. Device enumeration completes successfully including GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short and full reads), and SET_CONFIGURATION control transfers.\n7. Each control transfer stage (setup, data, status) completes with endpoint command completion and interrupt-driven event notification.\n8. Interrupt OUT transfer on physical endpoint 2 completes with 64-byte TRB and interrupt notification.\n9. Interrupt IN transfer on physical endpoint 3 completes with 64-byte TRB and interrupt notification.\n10. DSTS register frame number field (bits [11:3]) becomes non-zero indicating valid frame number.\n11. IRQ handler correctly reads and acknowledges event count, clears system register interrupt, and clears GIC IRQ.\n12. Test completes with finish(0) indicating successful pass.",
        "Remarks": "The testcase uses interrupt-driven synchronization with int_pend flag cleared by Default_IRQHandler. Multiple conditional interrupt waits depend on event_counter <= 0x4. Handshake polling at addresses 0xa0243ff4 and 0xa0243ff8 provides host-side synchronization. The IRQ handler condition uses logical AND (&&) instead of bitwise AND (&) for bit 31 check on MIZAR_LSS_SYSREG_RAW_STCR0, which may be a source code bug. Interrupt OUT transfer uses TRB control value 0x815 (interrupt type) on physical endpoint 2 (offset 0x20), while Interrupt IN transfer uses TRB control value 0x813 on physical endpoint 3 (offset 0x30). After the Interrupt IN transfer, the test polls DSTS for frame number validity using mask 0x00000FF8."
    },
    {
        "Index": "3",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Isochronous_Transfer_test",
        "Feature": "USB Full-Speed Device Isochronous Transfer",
        "Meta Headers": '<stdio.h>; <stdlib.h>; \"usb.h\"',
        "Meta Macros": "NA",
        "Meta Arrays": "int buf_data[16]; /* declared local array, not statically initialized */",
        "Speed": "Full-Speed",
        "Mode": "Device",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Full-Speed Device mode Isochronous Transfer operation. The test performs: NIC programming and GIC IRQ enable, clearing of Buffer_PointerLO and event_trb_addr memory regions, USB controller soft reset, PHY configuration, event ring setup, endpoint configuration, enumeration sequence, then Isochronous OUT transfer TRB setup on physical endpoint 6 (offset 0x60) with Buffer_PointerLO_1, size 0x3ff, TRB control 0x869, and DEPCMD 0x20506, followed by Isochronous IN transfer TRB setup on physical endpoint 7 (offset 0x70) with DEPCMD 0x40506, with multiple DSTS frame number polling loops.",
        "Test Description": "Verify USB Full-Speed Device mode Isochronous Transfer by performing controller soft reset, PHY and event ring configuration, endpoint configuration for control and isochronous endpoints, device enumeration (GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION), SET_ADDRESS, and executing isochronous OUT and isochronous IN transfers with interrupt-driven completion, multiple frame number validations via DSTS polling, and handshake synchronization.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.\n3-81. [Full procedure as specified in test case]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Clear buffer pointer and event TRB address memory regions (20 DWORDs each).\n3. Perform USB controller soft reset via DCTL and poll for reset completion.\n4. Configure USB2 PHY via GUSB2PHYCFG.\n5. Set up event ring by configuring event address low, event address high, event size, and event count registers.\n6. Configure global controller register GCTL for port direction.\n7. Configure device configuration register DCFG.\n8. Enable device events via DEVTEN register.\n9. Configure GUCTL register.\n10. Issue START_NEW_CONFIGURATION command and configure 8 endpoints using endpoint command registers with parameter registers.\n11. Configure TX resources for 8 endpoints by issuing transfer resource commands and polling for completion.\n12. Enable physical endpoints 0 and 1 via DALEPENA.\n13. Start device controller run via DCTL.\n14. Enable system register interrupt.\n15. Wait for link state connect/reset events via interrupt-driven waits.\n16. Re-configure DCFG and wait for additional events if needed.\n17. Write enumeration marker, read DCFG and DSTS, reconfigure DCFG and DCTL, enable all endpoints via DALEPENA.\n18. Execute setup stage for initial control transfer and wait for completion.\n19. Reconfigure USB2 PHY and wait for events.\n20. Set device address in DCFG, prepare status TRB, issue endpoint command on EP1 IN, and poll for completion.\n21. Wait for interrupt events and poll handshake register for synchronization.\n22. Execute full enumeration sequence: GET_DEVICE_DESCRIPTOR with 18-byte device descriptor data response, GET_CONFIGURATION_DESCRIPTOR short read (9 bytes), SET_CONFIGURATION with status data stage, GET_CONFIGURATION_DESCRIPTOR full read (60 bytes) with complete configuration descriptor data, each followed by status stage.\n23. Poll second handshake register for synchronization.\n24. Poll DSTS register for frame number field (bits [11:3]) to become non-zero.\n25. Prepare and execute Isochronous OUT transfer on physical endpoint 6 with 1023-byte TRB (isochronous type) using endpoint command with frame number parameter, wait for interrupt completion.\n26. Poll DSTS register for frame number bits [5:3] to equal 0x2.\n27. Wait for interrupt and poll DSTS register for frame number bits [5:3] to equal 0x3.\n28. Prepare and execute Isochronous IN transfer on physical endpoint 7 with 1023-byte TRB (isochronous type) using endpoint command with frame number parameter, wait for interrupt completion.\n29. Poll DSTS register for frame number bits [5:3] to equal 0x4.\n30. Write final completion marker and wait for final interrupt.\n31. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; 0xa0243ff4; 0xa0243ff8; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Meta Validation / Acceptance Criteria": "1. Soft reset completion: read_reg(MIZAR_USB_DCTL) must return 0xf00000.\n2-24. [Full validation criteria as specified]",
        "Validation / Acceptance Criteria": "1. USB controller soft reset completes successfully as indicated by DCTL register clearing the soft reset bit.\n2. All endpoint configuration commands complete successfully by polling the endpoint command register until the command active bit clears.\n3. All endpoint TX resource configuration commands complete for all 8 endpoints.\n4. Link state connect and reset events are received via interrupt-driven completion.\n5. Handshake synchronization registers return non-zero values indicating host-side readiness.\n6. Device enumeration completes successfully including GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short and full reads), and SET_CONFIGURATION control transfers.\n7. Each control transfer stage (setup, data, status) completes with endpoint command completion and interrupt-driven event notification.\n8. DSTS register frame number field (bits [11:3]) becomes non-zero indicating valid frame number after enumeration.\n9. Isochronous OUT transfer on physical endpoint 6 completes with 1023-byte isochronous TRB and interrupt notification.\n10. DSTS frame number bits [5:3] reach value 0x2 after isochronous OUT transfer.\n11. DSTS frame number bits [5:3] reach value 0x3 after second frame wait.\n12. Isochronous IN transfer on physical endpoint 7 completes with 1023-byte isochronous TRB and interrupt notification.\n13. DSTS frame number bits [5:3] reach value 0x4 after isochronous IN transfer.\n14. IRQ handler correctly reads and acknowledges event count, clears system register interrupt, and clears GIC IRQ.\n15. Test completes with finish(0) indicating successful pass.",
        "Remarks": "The testcase uses interrupt-driven synchronization with int_pend flag cleared by Default_IRQHandler. Isochronous OUT transfer uses TRB control value 0x869 (isochronous type) on physical endpoint 6 (offset 0x60) with transfer size 0x3ff (1023 bytes) and DEPCMD value 0x20506. Isochronous IN transfer uses TRB control value 0x869 on physical endpoint 7 (offset 0x70) with DEPCMD value 0x40506. Four separate DSTS frame number polling loops are used."
    },
    {
        "Index": "4",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_fs",
        "Feature": "USB Host Full-Speed Enumeration",
        "Meta Headers": '<stdio.h>; <stdlib.h>; \"usb.h\"',
        "Meta Macros": "NA",
        "Meta Arrays": "int data_in[512]; /* declared local array, not statically initialized */\nint data_out[512]; /* declared local array, not statically initialized */",
        "Speed": "Full-Speed",
        "Mode": "Host",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Host mode Full-Speed enumeration. The test performs: NIC programming and GIC IRQ enable, global controller configuration via read-modify-write of GCTL, GFLADJ, GUCTL, PIPE control register, PHY control register, reading host capability parameters, configuring port wake enables, event ring and command ring setup, scratchpad buffer and DCBAA configuration, port connect detection, port reset, Enable Slot command, Address Device commands (BSR=1 then BSR=0), Configure Endpoint command, and full enumeration sequence.",
        "Test Description": "Verify USB Host mode Full-Speed enumeration by performing xHCI controller initialization, port configuration with wake enables, event ring and command ring setup, scratchpad buffer and device context base address array configuration, port connect detection, port reset, Enable Slot command, Address Device commands (BSR=1 then BSR=0), Configure Endpoint command, and full enumeration sequence including GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short and full reads), and SET_CONFIGURATION.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.\n3-77. [Full procedure as specified in test case]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Configure global controller registers GCTL, GFLADJ, and GUCTL via read-modify-write.\n3. Configure PIPE control register and PHY control register.\n4. Read host capability parameters: HCSPARAMS1, Supported Protocol USB2 and USB3 capability registers, and PORTSC.\n5. Enable port wake on connect, disconnect, and over-current events via PORTSC.\n6. Read Doorbell Offset register.\n7. Set up Event Ring Segment Table with event ring base address and size 0x20.\n8. Read HCSPARAMS2 and PAGESIZE for scratchpad buffer requirements.\n9. Configure Scratchpad Buffer Array with two scratchpad buffer entries.\n10. Load Device Context Base Address Array with scratchpad pointer and two device context entries.\n11. Configure Command Ring via CRCR_LO/HI with ring cycle state bit.\n12. Configure CONFIG register with MaxSlotsEn and CIE.\n13. Configure DCBAAP_LO/HI with DCBAA base address.\n14. Configure Event Ring registers: ERSTSZ, ERDP_LO/HI, ERSTBA_LO/HI.\n15. Configure IMOD and IMAN interrupter registers.\n16. Enable USBCMD interrupt enable, then enable sysreg interrupt, then set USBCMD run/stop to run.\n17. Wait for first interrupt (port status change event).\n18. Read and clear USBSTS, update IMAN, read port connect status from PORTSC.\n19. Advance Event Ring Dequeue Pointer and issue port reset via PORTSC.\n20. Clear USBSTS, update IMAN, write PORTSC, read PORTSC status.\n21. Wait for second interrupt (port reset complete).\n22. Execute set_address: issue Enable Slot command via command ring, ring doorbell, wait for interrupt, configure Input Context for Address Device with BSR=1, issue Address Device command, ring doorbell, wait for interrupt, reconfigure Input Context for Address Device with BSR=0, issue Address Device command, ring doorbell, wait for interrupt.\n23. Configure Input Context for SET_CONFIGURATION with endpoint context including slot context and EP0 context.\n24. Issue Configure Endpoint command via command ring, ring doorbell, wait for interrupt.\n25. Execute enumeration: set up GET_DEVICE_DESCRIPTOR TRBs (setup/data/status), GET_CONFIGURATION_DESCRIPTOR short read TRBs (setup/data/status with 9-byte length), SET_CONFIGURATION TRBs (setup/status), GET_CONFIGURATION_DESCRIPTOR full read TRBs (setup/data/status with 60-byte length), clear USBSTS/IMAN, advance ERDP, ring endpoint doorbell, wait for interrupt, poll event ring for all events completion.\n26. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "MIZAR_USB_GCTL; MIZAR_USB_GFLADJ; MIZAR_USB_GUCTL; MIZAR_USB_BASE; MIZAR_USB_HCSPARAMS1; MIZAR_USB_SUPTPRT2_DW2; MIZAR_USB_SUPTPRT3_DW2; MIZAR_USB_PORTSC_20; MIZAR_USB_DBOFF; Event_Ring_Segment_Table; MIZAR_USB_HCSPARAMS2; MIZAR_USB_PAGESIZE; Scratchpad_Buffer_Array; Device_Context_Base_Address_Array; MIZAR_USB_CRCR_LO; MIZAR_USB_CRCR_HI; MIZAR_USB_CONFIG; MIZAR_USB_DCBAAP_LO; MIZAR_USB_DCBAAP_HI; MIZAR_USB_ERSTSZ; MIZAR_USB_ERDP_LO; MIZAR_USB_ERDP_HI; MIZAR_USB_ERSTBA_LO; MIZAR_USB_ERSTBA_HI; MIZAR_USB_IMOD; MIZAR_USB_IMAN; MIZAR_USB_USBCMD; MIZAR_LSS_SYSREG_INTR_EN0; MIZAR_USB_USBSTS; Default_Input_Context; Default_Command_Ring; MIZAR_USB_DB; Default_Event_Ring_Array; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0; EP0_TR_Dequeue_Pointer",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB",
        "Meta Validation / Acceptance Criteria": "1. First interrupt wait: int_pend must be cleared to 0 by Default_IRQHandler() after port status change event.\n2-14. [Full validation criteria as specified]",
        "Validation / Acceptance Criteria": "1. Port status change event is received via interrupt after enabling controller run.\n2. Port reset completes successfully as indicated by interrupt event.\n3. Enable Slot command completes successfully via command completion event.\n4. Address Device command with BSR=1 completes successfully.\n5. Address Device command with BSR=0 completes successfully.\n6. Configure Endpoint command completes successfully.\n7. All enumeration transfers complete: GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short and full reads), and SET_CONFIGURATION.\n8. Event ring completion entry at offset 0xf0 becomes non-zero indicating all enumeration events completed.\n9. IRQ handler correctly clears IMAN, clears system register interrupt, and clears GIC IRQ.\n10. Test completes with finish(0) indicating successful pass.",
        "Remarks": "This is a USB Host mode xHCI enumeration test for Full-Speed device. The testcase uses interrupt-driven synchronization with int_pend flag cleared by Default_IRQHandler. The IRQ handler condition uses logical AND (&&) instead of bitwise AND (&) for bit 31 check on MIZAR_LSS_SYSREG_RAW_STCR0, which may be a source code bug. The set_address function issues two Address Device commands: first with BSR=1 then with BSR=0. The GET_QUALIFIER section in enumeration() is commented out."
    },
    {
        "Index": "5",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_hs",
        "Feature": "USB Host High-Speed Enumeration",
        "Meta Headers": '<stdio.h>; <stdlib.h>; \"usb.h\"',
        "Meta Macros": "NA",
        "Meta Arrays": "int data_in[512]; /* declared local array, not statically initialized */\nint data_out[512]; /* declared local array, not statically initialized */",
        "Speed": "High-Speed",
        "Mode": "Host",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Host mode High-Speed enumeration using the xHCI host controller. The test performs: NIC programming and GIC IRQ enable, global controller configuration, port configuration, event ring and command ring setup, scratchpad buffer and DCBAA configuration, port connect detection, port reset, Enable Slot command, Address Device commands with High-Speed slot context 0x08300000, Configure Endpoint command, and full High-Speed enumeration sequence including GET_QUALIFIER.",
        "Test Description": "Verify USB Host mode High-Speed enumeration by performing xHCI controller initialization, port configuration with wake enables, event ring and command ring setup, scratchpad buffer and device context base address array configuration, port connect detection, port reset, Enable Slot command, Address Device commands (BSR=1 then BSR=0) with High-Speed slot context, Configure Endpoint command, and full High-Speed enumeration sequence including GET_DEVICE_DESCRIPTOR, GET_QUALIFIER, GET_CONFIGURATION_DESCRIPTOR (two short reads and one full read), and SET_CONFIGURATION.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.\n3-77. [Full procedure as specified in test case]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Configure global controller registers GCTL, GFLADJ, and GUCTL via read-modify-write.\n3. Configure PIPE control register and PHY control register.\n4. Read host capability parameters: HCSPARAMS1, Supported Protocol USB2 and USB3 capability registers, and PORTSC.\n5. Enable port wake on connect, disconnect, and over-current events via PORTSC.\n6. Read Doorbell Offset register.\n7. Set up Event Ring Segment Table with event ring base address and size 0x20.\n8. Read HCSPARAMS2 and PAGESIZE for scratchpad buffer requirements.\n9. Configure Scratchpad Buffer Array with two scratchpad buffer entries.\n10. Load Device Context Base Address Array with scratchpad pointer and two device context entries.\n11. Configure Command Ring via CRCR_LO/HI with ring cycle state bit.\n12. Configure CONFIG register with MaxSlotsEn and CIE.\n13. Configure DCBAAP_LO/HI with DCBAA base address.\n14. Configure Event Ring registers: ERSTSZ, ERDP_LO/HI, ERSTBA_LO/HI.\n15. Configure IMOD and IMAN interrupter registers.\n16. Enable USBCMD interrupt enable, then enable sysreg interrupt, then set USBCMD run/stop to run.\n17. Wait for first interrupt (port status change event).\n18. Read and clear USBSTS, update IMAN, read port connect status from PORTSC.\n19. Advance Event Ring Dequeue Pointer and issue port reset via PORTSC.\n20. Clear USBSTS, update IMAN, write PORTSC, read PORTSC status.\n21. Wait for second interrupt (port reset complete).\n22. Execute set_address: issue Enable Slot command via command ring, ring doorbell, wait for interrupt, configure Input Context for Address Device with BSR=1 (High-Speed slot context), issue Address Device command, ring doorbell, wait for interrupt, reconfigure Input Context for Address Device with BSR=0, issue Address Device command, ring doorbell, wait for interrupt.\n23. Configure Input Context for SET_CONFIGURATION with High-Speed slot context and endpoint context.\n24. Issue Configure Endpoint command via command ring, ring doorbell, wait for interrupt.\n25. Execute enumeration: set up GET_DEVICE_DESCRIPTOR TRBs (setup/data/status with 18-byte length), GET_QUALIFIER TRBs (setup/data/status with 10-byte length), GET_CONFIGURATION_DESCRIPTOR short read TRBs (setup/data/status with 9-byte length), SET_CONFIGURATION TRBs (setup/status), GET_CONFIGURATION2 short read TRBs (setup/data/status with 9-byte length), GET_CONFIGURATION3 full read TRBs (setup/data/status with 60-byte length), clear USBSTS/IMAN, advance ERDP, ring endpoint doorbell, wait for interrupt, poll event ring for all events completion.\n26. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "MIZAR_USB_GCTL; MIZAR_USB_GFLADJ; MIZAR_USB_GUCTL; MIZAR_USB_BASE; MIZAR_USB_HCSPARAMS1; MIZAR_USB_SUPTPRT2_DW2; MIZAR_USB_SUPTPRT3_DW2; MIZAR_USB_PORTSC_20; MIZAR_USB_DBOFF; Event_Ring_Segment_Table; MIZAR_USB_HCSPARAMS2; MIZAR_USB_PAGESIZE; Scratchpad_Buffer_Array; Device_Context_Base_Address_Array; MIZAR_USB_CRCR_LO; MIZAR_USB_CRCR_HI; MIZAR_USB_CONFIG; MIZAR_USB_DCBAAP_LO; MIZAR_USB_DCBAAP_HI; MIZAR_USB_ERSTSZ; MIZAR_USB_ERDP_LO; MIZAR_USB_ERDP_HI; MIZAR_USB_ERSTBA_LO; MIZAR_USB_ERSTBA_HI; MIZAR_USB_IMOD; MIZAR_USB_IMAN; MIZAR_USB_USBCMD; MIZAR_LSS_SYSREG_INTR_EN0; MIZAR_USB_USBSTS; Default_Input_Context; Default_Command_Ring; MIZAR_USB_DB; Default_Event_Ring_Array; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0; EP0_TR_Dequeue_Pointer",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB",
        "Meta Validation / Acceptance Criteria": "1. First interrupt wait: int_pend must be cleared to 0 by Default_IRQHandler() after port status change event.\n2-14. [Full validation criteria as specified]",
        "Validation / Acceptance Criteria": "1. Port status change event is received via interrupt after enabling controller run.\n2. Port reset completes successfully as indicated by interrupt event.\n3. Enable Slot command completes successfully via command completion event.\n4. Address Device command with BSR=1 completes successfully.\n5. Address Device command with BSR=0 completes successfully.\n6. Configure Endpoint command completes successfully.\n7. All enumeration transfers complete: GET_DEVICE_DESCRIPTOR, GET_QUALIFIER, GET_CONFIGURATION_DESCRIPTOR (two short reads and one full read), and SET_CONFIGURATION.\n8. Event ring completion entry at offset 0x150 becomes non-zero indicating all enumeration events completed.\n9. IRQ handler correctly clears IMAN, clears system register interrupt, and clears GIC IRQ.\n10. Test completes with finish(0) indicating successful pass.",
        "Remarks": "This is a USB Host mode xHCI enumeration test for High-Speed device. The slot context value 0x08300000 in set_address and 0x38300000 in SET_CONFIGURATION indicate High-Speed device speed (speed field = 3 for High-Speed). The GFLADJ value is 0xa07f000 which differs from the Full-Speed variant (0xa87f000). The enumeration function includes GET_QUALIFIER (descriptor type 0x06, 10-byte length) which is specific to High-Speed devices. Event ring polling occurs at offset 0x150 (vs 0xf0 in the FS variant) due to the additional transfer groups."
    },
    {
        "Index": "6",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_ls",
        "Feature": "USB Host Low-Speed Enumeration",
        "Meta Headers": '<stdio.h>; <stdlib.h>; \"usb.h\"',
        "Meta Macros": "NA",
        "Meta Arrays": "int data_in[512]; /* declared local array, not statically initialized */\nint data_out[512]; /* declared local array, not statically initialized */",
        "Speed": "Low-Speed",
        "Mode": "Host",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Host mode Low-Speed enumeration using the xHCI host controller. The test performs: NIC programming and GIC IRQ enable, global controller configuration, port configuration, event ring and command ring setup, scratchpad buffer and DCBAA configuration, port connect detection, port reset, Enable Slot command inline, Address Device commands with Low-Speed slot context 0x08200000 and EP0 max packet size 8, and enumeration sequence with GET_CONFIGURATION2 (8-byte) and GET_CONFIGURATION3 (24-byte).",
        "Test Description": "Verify USB Host mode Low-Speed enumeration by performing xHCI controller initialization, port configuration with wake enables, event ring and command ring setup, scratchpad buffer and device context base address array configuration, port connect detection, port reset, Enable Slot command, Address Device commands (BSR=1 then BSR=0) with Low-Speed slot context and 8-byte max packet size, and enumeration sequence including GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short read), SET_CONFIGURATION, GET_CONFIGURATION2 (short read with 8-byte length), and GET_CONFIGURATION3 (24-byte read).",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.\n3-108. [Full procedure as specified in test case]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Configure global controller registers GCTL, GFLADJ, and GUCTL via read-modify-write.\n3. Configure PIPE control register and PHY control register.\n4. Read host capability parameters: HCSPARAMS1, Supported Protocol USB2 and USB3 capability registers, and PORTSC.\n5. Enable port wake on connect, disconnect, and over-current events via PORTSC.\n6. Read Doorbell Offset register.\n7. Set up Event Ring Segment Table with event ring base address and size 0x30.\n8. Read HCSPARAMS2 and PAGESIZE for scratchpad buffer requirements.\n9. Configure Scratchpad Buffer Array with two scratchpad buffer entries.\n10. Load Device Context Base Address Array with scratchpad pointer and two device context entries.\n11. Configure Command Ring via CRCR_LO/HI with ring cycle state bit.\n12. Configure CONFIG register with MaxSlotsEn and CIE.\n13. Configure DCBAAP_LO/HI with DCBAA base address.\n14. Configure Event Ring registers: ERSTSZ, ERDP_LO/HI, ERSTBA_LO/HI.\n15. Configure IMOD and IMAN interrupter registers.\n16. Enable USBCMD interrupt enable, then enable sysreg interrupt, then set USBCMD run/stop to run.\n17. Wait for first interrupt (port status change event).\n18. Read and clear USBSTS, update IMAN, read port connect status from PORTSC.\n19. Advance Event Ring Dequeue Pointer and issue port reset via PORTSC.\n20. Clear USBSTS, update IMAN, write PORTSC, read PORTSC status.\n21. Wait for second interrupt (port reset complete).\n22. Issue Enable Slot command via command ring TRB, clear USBSTS/IMAN, advance ERDP, write PORTSC, ring Host Controller doorbell.\n23. Wait for third interrupt (Enable Slot completion).\n24. Configure Input Context for Address Device with BSR=1 using Low-Speed slot context (speed=2) and EP0 max packet size 8.\n25. Issue Address Device command with BSR=1, clear USBSTS/IMAN, advance ERDP, ring doorbell.\n26. Wait for fourth interrupt (Address Device BSR=1 completion), read event completion.\n27. Reconfigure Input Context for Address Device with BSR=0 using same Low-Speed parameters.\n28. Issue Address Device command with BSR=0, clear USBSTS/IMAN, advance ERDP, ring doorbell.\n29. Wait for fifth interrupt (Address Device BSR=0 completion), read event completion.\n30. Execute enumeration: set up GET_DEVICE_DESCRIPTOR TRBs (setup/data/status with 18-byte length), GET_CONFIGURATION_DESCRIPTOR short read TRBs (setup/data/status with 9-byte length), SET_CONFIGURATION TRBs (setup/status), GET_CONFIGURATION2 short read TRBs (setup/data/status with 8-byte length), GET_CONFIGURATION3 read TRBs (setup/data/status with 24-byte length), ring endpoint doorbell, poll event ring for all events completion.\n31. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "MIZAR_USB_GCTL; MIZAR_USB_GFLADJ; MIZAR_USB_GUCTL; MIZAR_USB_BASE; MIZAR_USB_HCSPARAMS1; MIZAR_USB_SUPTPRT2_DW2; MIZAR_USB_SUPTPRT3_DW2; MIZAR_USB_PORTSC_20; MIZAR_USB_DBOFF; Event_Ring_Segment_Table; MIZAR_USB_HCSPARAMS2; MIZAR_USB_PAGESIZE; Scratchpad_Buffer_Array; Device_Context_Base_Address_Array; MIZAR_USB_CRCR_LO; MIZAR_USB_CRCR_HI; MIZAR_USB_CONFIG; MIZAR_USB_DCBAAP_LO; MIZAR_USB_DCBAAP_HI; MIZAR_USB_ERSTSZ; MIZAR_USB_ERDP_LO; MIZAR_USB_ERDP_HI; MIZAR_USB_ERSTBA_LO; MIZAR_USB_ERSTBA_HI; MIZAR_USB_IMOD; MIZAR_USB_IMAN; MIZAR_USB_USBCMD; MIZAR_LSS_SYSREG_INTR_EN0; MIZAR_USB_USBSTS; Default_Command_Ring; MIZAR_USB_DB; Default_Input_Context; Default_Event_Ring_Array; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0; EP0_TR_Dequeue_Pointer",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB",
        "Meta Validation / Acceptance Criteria": "1. First interrupt wait: int_pend must be cleared to 0 by Default_IRQHandler() after port status change event.\n2-11. [Full validation criteria as specified]",
        "Validation / Acceptance Criteria": "1. Port status change event is received via interrupt after enabling controller run.\n2. Port reset completes successfully as indicated by interrupt event.\n3. Enable Slot command completes successfully via command completion event.\n4. Address Device command with BSR=1 completes successfully.\n5. Address Device command with BSR=0 completes successfully.\n6. All enumeration transfers complete: GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short read), SET_CONFIGURATION, GET_CONFIGURATION2 (short read with 8-byte length), and GET_CONFIGURATION3 (24-byte read).\n7. Event ring completion entry at offset 0x110 becomes non-zero indicating all enumeration events completed.\n8. IRQ handler correctly clears IMAN, clears system register interrupt, and clears GIC IRQ.\n9. Test completes with finish(0) indicating successful pass.",
        "Remarks": "This is a USB Host mode xHCI enumeration test for Low-Speed device. The slot context value 0x08200000 indicates Low-Speed device speed (speed field = 2 for Low-Speed). The EP0 max packet size is 8 bytes (0x00080020) which is the maximum allowed for Low-Speed devices. The Event Ring Segment Table size is 0x30 (differs from 0x20 in FS/HS variants). The Enable Slot command, Address Device BSR=1, and Address Device BSR=0 are implemented inline rather than in a separate set_address() function. The SET_CONFIGURATION Input Context and Configure Endpoint command block is entirely commented out. GET_CONFIGURATION2 uses a data transfer length of 0x8 (8 bytes). GET_CONFIGURATION3 uses wLength of 24 bytes."
    }
]

# TestPlan sheet columns
tp_columns = [
    'Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
    'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
    'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
    'Code Generation'
]

# MetaData sheet columns
md_columns = [
    'Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
    'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
    'Meta Headers', 'Meta Macros', 'Meta Arrays'
]

# Create workbook
wb = openpyxl.Workbook()

# Rename default sheet to TestPlan
ws_tp = wb.active
ws_tp.title = 'TestPlan'

# Create MetaData sheet
ws_md = wb.create_sheet('MetaData')

# Formatting
header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap_alignment = Alignment(wrap_text=True, vertical='top')
header_alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

def populate_sheet(ws, columns, data, key_map=None):
    # Write headers
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Write data rows
    for row_idx, tc in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            key = col_name
            if key_map and col_name in key_map:
                key = key_map[col_name]
            value = tc.get(key, '')
            if value is None:
                value = ''
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value))
            cell.alignment = wrap_alignment
            cell.border = thin_border
    
    # Freeze first row
    ws.freeze_panes = 'A2'
    
    # Auto-size columns
    for col_idx, col_name in enumerate(columns, 1):
        max_length = len(col_name)
        for row in range(2, len(data) + 2):
            cell_value = str(ws.cell(row=row, column=col_idx).value or '')
            lines = cell_value.split('\n')
            for line in lines:
                if len(line) > max_length:
                    max_length = len(line)
        # Cap width at 60 for readability
        adjusted_width = min(max_length + 2, 60)
        if adjusted_width < 12:
            adjusted_width = 12
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

# Populate TestPlan sheet
populate_sheet(ws_tp, tp_columns, test_cases)

# Populate MetaData sheet
populate_sheet(ws_md, md_columns, test_cases)

# Set MetaData sheet to veryHidden
ws_md.sheet_state = 'veryHidden'

# Save workbook
output_path = filename
wb.save(output_path)

# Verify
assert os.path.exists(output_path), f'File {output_path} does not exist'
assert os.path.getsize(output_path) > 0, f'File {output_path} is empty'

# Verify can reopen
wb_verify = openpyxl.load_workbook(output_path)
assert 'TestPlan' in wb_verify.sheetnames, 'TestPlan sheet missing'
assert 'MetaData' in wb_verify.sheetnames, 'MetaData sheet missing'
assert wb_verify['MetaData'].sheet_state == 'veryHidden', 'MetaData not veryHidden'

print(f'SUCCESS: {output_path}')
print(f'File size: {os.path.getsize(output_path)} bytes')
print(f'Sheets: {wb_verify.sheetnames}')
print(f'TestPlan rows: {ws_tp.max_row - 1}')
print(f'MetaData rows: {ws_md.max_row - 1}')
print(f'MetaData state: {wb_verify["MetaData"].sheet_state}')

# Output base64 for GitHub upload
with open(output_path, 'rb') as f:
    b64_content = base64.b64encode(f.read()).decode('utf-8')

# Write base64 to a text file for retrieval
with open('xlsx_base64.txt', 'w') as f:
    f.write(b64_content)

print(f'Base64 length: {len(b64_content)}')
print(f'Filename: {filename}')
