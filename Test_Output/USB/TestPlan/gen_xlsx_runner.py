#!/usr/bin/env python3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import os, sys, json, base64

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"USB_TestPlan_{timestamp}.xlsx"
filepath = os.path.join("/tmp", filename)

json_data = [{"Index":"1","SS / Module":"USB","Test Case Name":"usb_host_enumeration_hs","Feature":"USB Host Enumeration - High Speed","Speed":"High-Speed (HS)","Mode":"Host","Memory Start Offset":"NA","Memory End Offset":"NA","Test Description":"Verify USB host enumeration of a high-speed device.","Remarks":"Test remarks here.","Test Steps / Procedure":"Step 1\nStep 2","Impacted Registers":"REG1, REG2","Validation / Acceptance Criteria":"Criteria 1","Meta Test Description":"Meta desc 1","Meta Test Steps / Procedure":"Meta steps 1","Meta Impacted Registers":"Meta regs 1","Meta Validation / Acceptance Criteria":"Meta criteria 1","Meta Headers":"#include <stdio.h>","Meta Macros":"MACRO1","Meta Arrays":"int arr[10];"}]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "TestPlan"
ws.cell(row=1, column=1, value="Index")
ws.cell(row=2, column=1, value="1")
wb.save(filepath)
wb.close()

print(f"GENERATED: {filepath}")
print(f"SIZE: {os.path.getsize(filepath)}")

with open(filepath, "rb") as f:
    b64 = base64.b64encode(f.read()).decode()
    print(f"BASE64_START")
    print(b64)
    print(f"BASE64_END")
