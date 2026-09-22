#!/usr/bin/env python3
"""USB TestPlan Excel Generator - Agent 7
Generates USB_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx with TestPlan and MetaData sheets.
Run: python3 generate_usb_testplan_v2.py
"""
import json
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    exit(1)

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"USB_TestPlan_{timestamp}.xlsx"

# Input JSON data
json_data = [
    {
        "Index": "1",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Bulk_Transfer_test",
        "Feature": "Full-Speed Device Bulk Transfer",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"usb.h\"",
        "Meta Macros": "NA",
        "Meta Arrays": "buf_data[16]",
        "Speed": "Full-Speed",
        "Mode": "Device Mode",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Full-Speed Device mode Bulk Transfer operation. It begins by calling nic_programming() and GIC_EnableAllIRQ(). It clears 20 DWORDs at Buffer_PointerLO and event_trb_addr. It performs a soft reset by writing 0x40f00000 to MIZAR_USB_DCTL and polling MIZAR_USB_DCTL until the value equals 0xf00000. It configures MIZAR_USB_GUSB2PHYCFG with 0x40002407, sets up the event ring by writing Default_Event_Ring_Array to MIZAR_USB_GEVNTADRLO, 0x0 to MIZAR_USB_GEVNTADRHI, 0x30 to MIZAR_USB_GEVNTSIZ, and 0x0 to MIZAR_USB_GEVNTCOUNT. It reads MIZAR_USB_GCTL, then writes 0x30c12214 for port direction. It reads MIZAR_USB_DCFG and writes 0x480801. It writes 0x1f to MIZAR_USB_DEVTEN. It reads MIZAR_USB_GUCTL and writes 0xa400010. It calls set_configuration() 9 times to configure endpoints using MIZAR_USB_DEPCMDPAR1, MIZAR_USB_DEPCMDPAR0, and MIZAR_USB_DEPCMD with various parameters, polling MIZAR_USB_DEPCMD until the command completes. It configures endpoint TX resources by looping 8 times, writing 0x1 to MIZAR_USB_DEPCMDPAR0+(i*0x10) and 0x402 to MIZAR_USB_DEPCMD+(i*0x10), polling until completion. It enables physical endpoints 0 and 1 by writing 0x3 to MIZAR_USB_DALEPENA, then writes 0x80f00000 to MIZAR_USB_DCTL. It enables interrupt at sysreg by writing 0x80000000 to MIZAR_LSS_SYSREG_INTR_EN0. It waits for link state connect/reset events via interrupt pending flags. It writes enumeration marker 0xdeadbee0 to 0xA0243ffc. It reads MIZAR_USB_DCFG and MIZAR_USB_DSTS, reconfigures MIZAR_USB_DCFG, writes 0x80f00a00 to MIZAR_USB_DCTL, enables all endpoints via MIZAR_USB_DALEPENA with 0xff. It calls setup_stage() which prepares TRBs at event_trb_addr and Buffer_PointerLO, issues DEPCMD start transfer command 0x506, and waits for interrupt. It reconfigures MIZAR_USB_GUSB2PHYCFG with 0x40002547. It sets device address by writing 0x480809 to MIZAR_USB_DCFG. It prepares TRBs and issues start transfer on endpoint 1. It polls 0xa0243ff4 until non-zero (handshake). It calls enumeration() which performs GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION sequences. It polls 0xa0243ff8 until non-zero. It then performs two Bulk Transfer operations. The Default_IRQHandler reads MIZAR_LSS_SYSREG_MSK_STS0 and MIZAR_LSS_SYSREG_RAW_STCR0, reads MIZAR_USB_GEVNTCOUNT and writes back the count to clear, clears sysreg interrupt, and clears GIC IRQ 84. The test ends with finish(0).",
        "Test Description": "This test validates USB Full-Speed Device mode Bulk Transfer functionality. The test performs a controller soft reset via DCTL, configures the USB 2.0 PHY, sets up the event ring buffer, configures the global controller register for device-mode port direction, configures device speed and event enables via DCFG and DEVTEN, and sets up the utility control register GUCTL. It then configures all endpoints using endpoint command registers and enables them via DALEPENA. The test enables system-level interrupts and waits for link state connect and reset events. It performs USB enumeration including device descriptor, configuration descriptor retrieval, set address, and set configuration sequences. After enumeration, it executes two Bulk Transfer operations on bulk endpoints, issuing start transfer commands and waiting for transfer completion via interrupts. An interrupt handler reads system interrupt status, processes event counts via the event count register, clears interrupt status, and clears the GIC interrupt. The test completes successfully by calling finish.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization. 2. Call GIC_EnableAllIRQ() to enable all IRQs. 3. Clear 20 DWORDs at Buffer_PointerLO and event_trb_addr. 4. Write 0x40f00000 to MIZAR_USB_DCTL to initiate soft reset. 5. Poll MIZAR_USB_DCTL until value equals 0xf00000. 6-36. Full enumeration and bulk transfer sequence. IRQ HANDLER: clear int_pend, read sysreg status, process event count, clear interrupts.",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable all GIC IRQs. 2. Clear transfer buffer and event TRB memory regions (20 DWORDs each). 3. Perform USB controller soft reset via DCTL and poll until reset completes. 4. Configure USB 2.0 PHY settings. 5. Set up event ring by writing event address low, event address high, event size, and event count registers. 6. Read and configure GCTL for device-mode port direction. 7. Read and configure DCFG for device speed settings. 8. Enable device events via DEVTEN. 9. Read and configure GUCTL utility control register. 10. Issue Start New Configuration command followed by 8 Set Endpoint Configuration commands. 11. Configure transfer resources for 8 endpoints. 12. Enable physical endpoints 0 and 1 via DALEPENA. 13. Set DCTL to run the controller and enable system-level interrupt. 14-22. Wait for events, perform enumeration. 23-24. Execute two Bulk Transfers. 25. Verify test completion via finish call.",
        "Meta Impacted Registers": "Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; 0xa0243ff4; 0xa0243ff8; Buffer_PointerLO_1; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Meta Validation / Acceptance Criteria": "1. Soft reset validation: poll MIZAR_USB_DCTL until rd_data equals 0xf00000. 2. Endpoint command completion. 3. Start transfer command completion. 4. Interrupt-driven event processing. 5. Event counter check. 6. Handshake polling. 7. IRQ handler validation. 8. Test passes by reaching finish(0).",
        "Validation / Acceptance Criteria": "1. Controller soft reset completes successfully. 2. All endpoint configuration commands complete. 3. All start transfer commands complete. 4. Device events are received and processed via the interrupt handler. 5. Event counter reaches expected thresholds. 6. Host handshake signals are received. 7. Interrupt handler correctly processes events. 8. Two bulk transfers of 64 bytes each complete successfully. 9. Test passes by reaching the finish call.",
        "Remarks": "Test operates in USB Full-Speed Device mode with interrupt-driven event handling using GIC IRQ 84. Enumeration includes Get Device Descriptor, Get Configuration Descriptor (short and full), and Set Configuration sequences. Two bulk transfers are performed on separate bulk endpoints after enumeration."
    },
    {
        "Index": "2",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Interrupt_Transfer_test",
        "Feature": "Full-Speed Device Interrupt Transfer",
        "Speed": "Full-Speed",
        "Mode": "Device Mode",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Test Description": "This test validates USB Full-Speed Device mode Interrupt Transfer functionality. The test performs a controller soft reset via DCTL, configures the USB 2.0 PHY, sets up the event ring buffer, configures GCTL for device-mode port direction, configures device speed and event enables via DCFG and DEVTEN, and sets up GUCTL. It configures all endpoints and enables them via DALEPENA. After enumeration, it executes two Interrupt Transfer operations on interrupt endpoints with different TRB control types. It then polls DSTS for a valid frame number.",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable all GIC IRQs. 2-22. Same as bulk transfer test. 23. Execute Interrupt Transfer 1 with interrupt-type control on endpoint offset 2. 24. Execute Interrupt Transfer 2 with normal-type control on endpoint offset 3. 25. Poll DSTS for valid frame number. 26. Verify test completion via finish call.",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Validation / Acceptance Criteria": "1. Controller soft reset completes. 2. All endpoint configuration commands complete. 3. All start transfer commands complete. 4. Device events received. 5-6. Handshake signals received. 7. Two interrupt transfers of 64 bytes each complete. 8. DSTS frame number becomes non-zero. 9-10. Test passes.",
        "Remarks": "Test operates in USB Full-Speed Device mode. First interrupt transfer uses TRB control 0x815, second uses 0x813. Additionally polls DSTS for valid frame number."
    },
    {
        "Index": "3",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Isochronous_Transfer_test",
        "Feature": "Full-Speed Device Isochronous Transfer",
        "Speed": "Full-Speed",
        "Mode": "Device Mode",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Test Description": "This test validates USB Full-Speed Device mode Isochronous Transfer functionality. After enumeration, it polls DSTS for a valid frame number, then executes two Isochronous OUT Transfer operations with 1023-byte transfer size and isochronous TRB control type. It polls DSTS for specific frame number values (2, 3, and 4) to synchronize transfers.",
        "Test Steps / Procedure": "1-22. Same as bulk transfer test. 23. Poll DSTS for valid frame number. 24. Execute Isochronous OUT Transfer 1 with 1023-byte size on endpoint offset 6. 25. Poll DSTS for frame numbers 2, 3. 26. Execute Isochronous OUT Transfer 2 on endpoint offset 7. 27. Poll DSTS for frame number 4. 28. Verify test completion.",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Validation / Acceptance Criteria": "1. Controller soft reset completes. 2. All endpoint configuration commands complete. 3. All start transfer commands complete for isochronous endpoints. 4. Device events received. 5. Host handshake signals received. 6. DSTS frame number becomes non-zero. 7. Two isochronous OUT transfers of 1023 bytes each complete. 8. DSTS frame number reaches values 2, 3, and 4. 9. Test passes.",
        "Remarks": "Test uses isochronous TRB control value 0x869 and 1023-byte transfer size. Start transfer commands include frame-number parameters. Extensive frame-number polling on DSTS."
    },
    {
        "Index": "4",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_fs",
        "Feature": "Host Enumeration Full-Speed",
        "Speed": "Full-Speed",
        "Mode": "Host Mode",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Test Description": "This test validates USB Host mode Full-Speed device enumeration. The test configures GCTL, GFLADJ, and GUCTL. It reads HCSPARAMS1, HCSPARAMS2, SUPTPRT2_DW2, SUPTPRT3_DW2, PAGESIZE, and DBOFF. It configures PORTSC_20, sets up xHCI data structures, command ring via CRCR, DCBAAP, CONFIG, interrupter registers, and starts host via USBCMD. It performs Enable Slot, Address Device, Configure Endpoint, then enumeration transfers.",
        "Test Steps / Procedure": "1. Configure GCTL, GFLADJ, GUCTL. 2. Configure PHY and pipe control. 3. Read HCSPARAMS1, HCSPARAMS2, PAGESIZE, SUPTPRT2_DW2, SUPTPRT3_DW2. 4. Configure PORTSC_20. 5. Set up event ring, scratchpad, DCBA, command ring. 6. Configure CRCR, CONFIG, DCBAAP, interrupter registers. 7. Start host via USBCMD. 8. Wait for port connect, perform port reset. 9. Enable Slot, Address Device, Configure Endpoint. 10. Execute enumeration transfers. 11. Poll event ring for completion.",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB; CAPLENGTH",
        "Validation / Acceptance Criteria": "1. Port connection detected. 2. Port reset completes. 3. Enable Slot, Address Device, Configure Endpoint commands complete. 4. All enumeration transfers complete. 5. Test passes by reaching finish.",
        "Remarks": "Host mode Full-Speed enumeration using xHCI. Slot context 0x38100000. GFLADJ 0xa87f000."
    },
    {
        "Index": "5",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_hs",
        "Feature": "Host Enumeration High-Speed",
        "Speed": "High-Speed",
        "Mode": "Host Mode",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Test Description": "This test validates USB Host mode High-Speed device enumeration. Same xHCI flow as Full-Speed but with High-Speed specific parameters. Includes Get Device Qualifier request.",
        "Test Steps / Procedure": "Same as Full-Speed host enumeration with additional Get Device Qualifier and second Get Configuration Descriptor short read.",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB; CAPLENGTH",
        "Validation / Acceptance Criteria": "Same as Full-Speed plus Get Device Qualifier transfer completes.",
        "Remarks": "Slot context 0x38300000 (High-Speed). GFLADJ 0xa07f000. Additional Get Device Qualifier request. Event ring poll at offset 0x150."
    },
    {
        "Index": "6",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_ls",
        "Feature": "Host Enumeration Low-Speed",
        "Speed": "Low-Speed",
        "Mode": "Host Mode",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Test Description": "This test validates USB Host mode Low-Speed device enumeration. Same xHCI flow but with Low-Speed parameters. No Configure Endpoint command (commented out). Third Get Configuration Descriptor requests 24 bytes.",
        "Test Steps / Procedure": "Same as Full-Speed host enumeration without Configure Endpoint. Third Get Config Descriptor is 24 bytes.",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB; CAPLENGTH",
        "Validation / Acceptance Criteria": "Same as Full-Speed. All enumeration transfers complete. Event ring poll at offset 0x110.",
        "Remarks": "Slot context 0x08200000 (Low-Speed). GFLADJ 0xa87f000. Configure Endpoint commented out. Third Get Config Descriptor 24 bytes. Event ring poll at 0x110."
    }
]

# TestPlan columns
tp_cols = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

# MetaData columns
md_cols = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

# Create workbook
wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = "TestPlan"

# Header formatting
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_align = Alignment(wrap_text=True, vertical="top")

# Write headers
for col_idx, col_name in enumerate(tp_cols, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

# Write data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(tp_cols, 1):
        value = row_data.get(col_name, "")
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value if value else "")
        cell.alignment = wrap_align

# Freeze first row
ws_tp.freeze_panes = "A2"

# Auto-size columns
for col_idx, col_name in enumerate(tp_cols, 1):
    max_len = len(col_name)
    for row in range(2, len(json_data) + 2):
        val = ws_tp.cell(row=row, column=col_idx).value
        if val:
            max_len = max(max_len, min(len(str(val)), 80))
    width = min(max_len + 2, 60)
    ws_tp.column_dimensions[get_column_letter(col_idx)].width = width

# --- MetaData Sheet ---
ws_md = wb.create_sheet("MetaData")

# Write headers
for col_idx, col_name in enumerate(md_cols, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

# Write data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(md_cols, 1):
        value = row_data.get(col_name, "")
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value if value else "")
        cell.alignment = wrap_align

# Freeze first row
ws_md.freeze_panes = "A2"

# Auto-size columns
for col_idx, col_name in enumerate(md_cols, 1):
    max_len = len(col_name)
    for row in range(2, len(json_data) + 2):
        val = ws_md.cell(row=row, column=col_idx).value
        if val:
            max_len = max(max_len, min(len(str(val)), 80))
    width = min(max_len + 2, 60)
    ws_md.column_dimensions[get_column_letter(col_idx)].width = width

# Set MetaData sheet to veryHidden
ws_md.sheet_state = "veryHidden"

# Save workbook
wb.save(filename)
print(f"SUCCESS: Workbook saved as {filename}")
print(f"TestPlan rows: {len(json_data)}")
print(f"MetaData rows: {len(json_data)}")

# Validate
try:
    wb2 = load_workbook(filename)
    assert "TestPlan" in wb2.sheetnames
    assert "MetaData" in wb2.sheetnames
    print("VALIDATION: PASSED")
except Exception as e:
    print(f"VALIDATION: FAILED - {e}")

import os
print(f"File size: {os.path.getsize(filename)} bytes")
print(f"Full path: {os.path.abspath(filename)}")
