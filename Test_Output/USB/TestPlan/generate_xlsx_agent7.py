#!/usr/bin/env python3
"""
Agent 7 - Excel Generator Script
Generates USB_TestPlan_YYYYMMDD_HHMMSS.xlsx with TestPlan and MetaData sheets.
Uses IST (GMT+05:30) for timestamp.
"""
import json
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"USB_TestPlan_{timestamp}.xlsx"

# JSON data - 6 test cases
json_data = [
    {
        "Index": "1",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Bulk_Transfer_test",
        "Feature": "USB Full-Speed Device Bulk Transfer",
        "Meta Headers": '<stdio.h>; <stdlib.h>; "usb.h"',
        "Meta Macros": "NA",
        "Meta Arrays": "int buf_data[16]; /* declared local array, not statically initialized, used as local variable */",
        "Speed": "Full-Speed",
        "Mode": "Device",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Full-Speed Device mode Bulk Transfer operation. The test performs the following high-level sequence: NIC programming and GIC IRQ enable, clearing of Buffer_PointerLO and event_trb_addr memory regions (20 DWORDs each), USB controller soft reset via MIZAR_USB_DCTL and polling for reset completion, USB2 PHY configuration via MIZAR_USB_GUSB2PHYCFG, event ring setup (GEVNTADRLO, GEVNTADRHI, GEVNTSIZ, GEVNTCOUNT), global controller configuration via MIZAR_USB_GCTL, device configuration via MIZAR_USB_DCFG, device event enable via MIZAR_USB_DEVTEN, GUCTL configuration, endpoint configuration using set_configuration() for 9 endpoints (START_NEW_CONFIGURATION + 8 endpoint configs), endpoint TX resource configuration loop (8 iterations writing DEPCMDPAR0 and DEPCMD per endpoint with polling), enabling physical endpoints 0 and 1 via DALEPENA, starting device controller run via DCTL, enabling sysreg interrupt, waiting for link state connect/reset events via interrupt waits, enumeration sequence including: setup_stage for control transfers, SET_ADDRESS phase with TRB and DEPCMD operations, handshake polling at 0xa0243ff4, full enumeration() call which performs GET_DEVICE_DESCRIPTOR (18 bytes device descriptor data), GET_CONFIGURATION_DESCRIPTOR (9 bytes short read then 60 bytes full read with complete configuration descriptor data), SET_CONFIGURATION with status stage, handshake polling at 0xa0243ff8, then bulk OUT transfer TRB setup on physical endpoint 4 (offset 0x40) with Buffer_PointerLO_1 and size 0x40, followed by bulk IN transfer TRB setup on physical endpoint 5 (offset 0x50) with Buffer_PointerLO_1 and size 0x40, with interrupt waits after each transfer, and final marker write 0xdeadbee7 to 0xA0243ffc before finish(0).",
        "Test Description": "Verify USB Full-Speed Device mode Bulk Transfer by performing controller soft reset, PHY and event ring configuration, endpoint configuration for control and bulk endpoints, device enumeration (GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION), SET_ADDRESS, and executing bulk OUT and bulk IN transfers with interrupt-driven completion and handshake synchronization.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.\n3. Loop j=0 to 19: write_reg(Buffer_PointerLO + j*DWORD, 0x0); write_reg(event_trb_addr + j*DWORD, 0x0);\n4. Soft reset via MIZAR_USB_DCTL.\n5-74. [Full procedure as specified]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Clear buffer pointer and event TRB address memory regions (20 DWORDs each).\n3. Perform USB controller soft reset via DCTL and poll for reset completion.\n4. Configure USB2 PHY via GUSB2PHYCFG.\n5. Set up event ring by configuring event address low, event address high, event size, and event count registers.\n6. Configure global controller register GCTL for port direction.\n7. Configure device configuration register DCFG.\n8. Enable device events via DEVTEN register.\n9. Configure GUCTL register.\n10. Issue START_NEW_CONFIGURATION command and configure 8 endpoints using endpoint command registers with parameter registers.\n11. Configure TX resources for 8 endpoints by issuing transfer resource commands and polling for completion.\n12. Enable physical endpoints 0 and 1 via DALEPENA.\n13. Start device controller run via DCTL.\n14. Enable system register interrupt.\n15. Wait for link state connect/reset events via interrupt-driven waits.\n16. Re-configure DCFG and wait for additional events if needed.\n17. Write enumeration marker, read DCFG and DSTS, reconfigure DCFG and DCTL, enable all endpoints via DALEPENA.\n18. Execute setup stage for initial control transfer and wait for completion.\n19. Reconfigure USB2 PHY and wait for events.\n20. Set device address in DCFG, prepare status TRB, issue endpoint command on EP1 IN, and poll for completion.\n21. Wait for interrupt events and poll handshake register for synchronization.\n22. Execute full enumeration sequence: GET_DEVICE_DESCRIPTOR with 18-byte device descriptor data response, GET_CONFIGURATION_DESCRIPTOR short read (9 bytes) and full read (60 bytes) with complete configuration descriptor data, SET_CONFIGURATION with status data stage.\n23. Poll second handshake register for synchronization.\n24. Prepare and execute Bulk OUT transfer on physical endpoint 4 with 64-byte TRB using endpoint command, wait for interrupt completion.\n25. Prepare and execute Bulk IN transfer on physical endpoint 5 with 64-byte TRB using endpoint command, wait for interrupt completion.\n26. Write final completion marker and wait for final interrupt.\n27. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; 0xa0243ff4; 0xa0243ff8; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Meta Validation / Acceptance Criteria": "1. Soft reset completion: read_reg(MIZAR_USB_DCTL) must return 0xf00000.\n2-20. [Full criteria as specified]",
        "Validation / Acceptance Criteria": "1. USB controller soft reset completes successfully as indicated by DCTL register clearing the soft reset bit.\n2. All endpoint configuration commands complete successfully by polling the endpoint command register until the command active bit clears.\n3. All endpoint TX resource configuration commands complete for all 8 endpoints.\n4. Link state connect and reset events are received via interrupt-driven completion.\n5. Handshake synchronization registers return non-zero values indicating host-side readiness.\n6. Device enumeration completes successfully including GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short and full reads), and SET_CONFIGURATION control transfers.\n7. Each control transfer stage (setup, data, status) completes with endpoint command completion and interrupt-driven event notification.\n8. Bulk OUT transfer on physical endpoint 4 completes with 64-byte TRB and interrupt notification.\n9. Bulk IN transfer on physical endpoint 5 completes with 64-byte TRB and interrupt notification.\n10. IRQ handler correctly reads and acknowledges event count, clears system register interrupt, and clears GIC IRQ.\n11. Test completes with finish(0) indicating successful pass.",
        "Remarks": "The testcase uses interrupt-driven synchronization with int_pend flag cleared by Default_IRQHandler. Multiple conditional interrupt waits depend on event_counter <= 0x4. Handshake polling at addresses 0xa0243ff4 and 0xa0243ff8 provides host-side synchronization."
    },
    {
        "Index": "2",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Interrupt_Transfer_test",
        "Feature": "USB Full-Speed Device Interrupt Transfer",
        "Meta Headers": '<stdio.h>; <stdlib.h>; "usb.h"',
        "Meta Macros": "NA",
        "Meta Arrays": "int buf_data[16]; /* declared local array, not statically initialized */",
        "Speed": "Full-Speed",
        "Mode": "Device",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Full-Speed Device mode Interrupt Transfer operation.",
        "Test Description": "Verify USB Full-Speed Device mode Interrupt Transfer by performing controller soft reset, PHY and event ring configuration, endpoint configuration for control and interrupt endpoints, device enumeration (GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION), SET_ADDRESS, and executing interrupt OUT and interrupt IN transfers with interrupt-driven completion, frame number validation, and handshake synchronization.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2-74. [Full procedure as specified]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Clear buffer pointer and event TRB address memory regions (20 DWORDs each).\n3. Perform USB controller soft reset via DCTL and poll for reset completion.\n4. Configure USB2 PHY via GUSB2PHYCFG.\n5. Set up event ring.\n6. Configure GCTL.\n7. Configure DCFG.\n8. Enable device events via DEVTEN.\n9. Configure GUCTL.\n10. Issue START_NEW_CONFIGURATION and configure 8 endpoints.\n11. Configure TX resources for 8 endpoints.\n12. Enable physical endpoints 0 and 1 via DALEPENA.\n13. Start device controller run via DCTL.\n14. Enable system register interrupt.\n15. Wait for link state connect/reset events.\n16. Re-configure DCFG and wait for additional events.\n17. Write enumeration marker, reconfigure DCFG and DCTL, enable all endpoints.\n18. Execute setup stage for initial control transfer.\n19. Reconfigure USB2 PHY and wait for events.\n20. Set device address in DCFG.\n21. Wait for interrupt events and poll handshake register.\n22. Execute full enumeration sequence.\n23. Poll second handshake register.\n24. Prepare and execute Interrupt OUT transfer on physical endpoint 2.\n25. Prepare and execute Interrupt IN transfer on physical endpoint 3.\n26. Poll DSTS register for frame number to become non-zero.\n27. Write final completion marker and wait for final interrupt.\n28. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; 0xa0243ff4; 0xa0243ff8; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Meta Validation / Acceptance Criteria": "1. Soft reset completion: read_reg(MIZAR_USB_DCTL) must return 0xf00000.\n2-21. [Full criteria as specified]",
        "Validation / Acceptance Criteria": "1. USB controller soft reset completes successfully.\n2. All endpoint configuration commands complete successfully.\n3. All endpoint TX resource configuration commands complete for all 8 endpoints.\n4. Link state connect and reset events are received.\n5. Handshake synchronization registers return non-zero values.\n6. Device enumeration completes successfully.\n7. Each control transfer stage completes.\n8. Interrupt OUT transfer on physical endpoint 2 completes with 64-byte TRB.\n9. Interrupt IN transfer on physical endpoint 3 completes with 64-byte TRB.\n10. DSTS register frame number field becomes non-zero.\n11. IRQ handler correctly reads and acknowledges event count.\n12. Test completes with finish(0) indicating successful pass.",
        "Remarks": "The testcase uses interrupt-driven synchronization with int_pend flag cleared by Default_IRQHandler. Interrupt OUT transfer uses TRB control value 0x815 on physical endpoint 2, while Interrupt IN transfer uses TRB control value 0x813 on physical endpoint 3."
    },
    {
        "Index": "3",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Isochronous_Transfer_test",
        "Feature": "USB Full-Speed Device Isochronous Transfer",
        "Meta Headers": '<stdio.h>; <stdlib.h>; "usb.h"',
        "Meta Macros": "NA",
        "Meta Arrays": "int buf_data[16]; /* declared local array, not statically initialized */",
        "Speed": "Full-Speed",
        "Mode": "Device",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Full-Speed Device mode Isochronous Transfer operation.",
        "Test Description": "Verify USB Full-Speed Device mode Isochronous Transfer by performing controller soft reset, PHY and event ring configuration, endpoint configuration for control and isochronous endpoints, device enumeration (GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION), SET_ADDRESS, and executing isochronous OUT and isochronous IN transfers with interrupt-driven completion, multiple frame number validations via DSTS polling, and handshake synchronization.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2-81. [Full procedure as specified]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Clear buffer pointer and event TRB address memory regions.\n3. Perform USB controller soft reset via DCTL.\n4. Configure USB2 PHY.\n5. Set up event ring.\n6. Configure GCTL.\n7. Configure DCFG.\n8. Enable device events via DEVTEN.\n9. Configure GUCTL.\n10. Issue START_NEW_CONFIGURATION and configure 8 endpoints.\n11. Configure TX resources for 8 endpoints.\n12. Enable physical endpoints 0 and 1.\n13. Start device controller run.\n14. Enable system register interrupt.\n15. Wait for link state connect/reset events.\n16. Re-configure DCFG.\n17. Write enumeration marker, reconfigure.\n18. Execute setup stage.\n19. Reconfigure USB2 PHY.\n20. Set device address.\n21. Wait for interrupt events and poll handshake.\n22. Execute full enumeration sequence.\n23. Poll second handshake register.\n24. Poll DSTS register for frame number.\n25. Prepare and execute Isochronous OUT transfer on physical endpoint 6.\n26. Poll DSTS for frame number bits [5:3] to equal 0x2.\n27. Wait for interrupt and poll DSTS for frame number bits [5:3] to equal 0x3.\n28. Prepare and execute Isochronous IN transfer on physical endpoint 7.\n29. Poll DSTS for frame number bits [5:3] to equal 0x4.\n30. Write final completion marker.\n31. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; 0xa0243ff4; 0xa0243ff8; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0",
        "Impacted Registers": "DCTL; GCTL; DCFG; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Meta Validation / Acceptance Criteria": "1. Soft reset completion.\n2-24. [Full criteria as specified]",
        "Validation / Acceptance Criteria": "1. USB controller soft reset completes successfully.\n2. All endpoint configuration commands complete.\n3. All endpoint TX resource configuration commands complete.\n4. Link state connect and reset events are received.\n5. Handshake synchronization registers return non-zero values.\n6. Device enumeration completes successfully.\n7. Each control transfer stage completes.\n8. DSTS register frame number field becomes non-zero.\n9. Isochronous OUT transfer on physical endpoint 6 completes with 1023-byte isochronous TRB.\n10. DSTS frame number bits [5:3] reach value 0x2.\n11. DSTS frame number bits [5:3] reach value 0x3.\n12. Isochronous IN transfer on physical endpoint 7 completes.\n13. DSTS frame number bits [5:3] reach value 0x4.\n14. IRQ handler correctly reads and acknowledges event count.\n15. Test completes with finish(0) indicating successful pass.",
        "Remarks": "Isochronous OUT transfer uses TRB control value 0x869 on physical endpoint 6 with transfer size 0x3ff (1023 bytes) and DEPCMD value 0x20506. Isochronous IN transfer uses TRB control value 0x869 on physical endpoint 7 with DEPCMD value 0x40506."
    },
    {
        "Index": "4",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_fs",
        "Feature": "USB Host Full-Speed Enumeration",
        "Meta Headers": '<stdio.h>; <stdlib.h>; "usb.h"',
        "Meta Macros": "NA",
        "Meta Arrays": "int data_in[512]; /* declared local array */\nint data_out[512]; /* declared local array */",
        "Speed": "Full-Speed",
        "Mode": "Host",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Host mode Full-Speed enumeration.",
        "Test Description": "Verify USB Host mode Full-Speed enumeration by performing xHCI controller initialization, port configuration with wake enables, event ring and command ring setup, scratchpad buffer and device context base address array configuration, port connect detection, port reset, Enable Slot command, Address Device commands (BSR=1 then BSR=0), Configure Endpoint command, and full enumeration sequence including GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short and full reads), and SET_CONFIGURATION.",
        "Meta Test Steps / Procedure": "1. Call nic_programming().\n2-77. [Full procedure as specified]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Configure global controller registers GCTL, GFLADJ, and GUCTL.\n3. Configure PIPE control register and PHY control register.\n4. Read host capability parameters.\n5. Enable port wake on connect, disconnect, and over-current events.\n6. Read Doorbell Offset register.\n7. Set up Event Ring Segment Table.\n8. Read HCSPARAMS2 and PAGESIZE.\n9. Configure Scratchpad Buffer Array.\n10. Load Device Context Base Address Array.\n11. Configure Command Ring.\n12. Configure CONFIG register.\n13. Configure DCBAAP.\n14. Configure Event Ring registers.\n15. Configure IMOD and IMAN.\n16. Enable USBCMD interrupt enable, sysreg interrupt, run.\n17. Wait for first interrupt (port status change).\n18. Read and clear USBSTS, read port connect status.\n19. Advance ERDP and issue port reset.\n20. Clear USBSTS, update IMAN, write PORTSC.\n21. Wait for second interrupt (port reset complete).\n22. Execute set_address: Enable Slot, Address Device BSR=1, Address Device BSR=0.\n23. Configure Input Context for SET_CONFIGURATION.\n24. Issue Configure Endpoint command.\n25. Execute enumeration: GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION.\n26. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "MIZAR_USB_GCTL; MIZAR_USB_GFLADJ; MIZAR_USB_GUCTL; MIZAR_USB_BASE; MIZAR_USB_HCSPARAMS1; MIZAR_USB_SUPTPRT2_DW2; MIZAR_USB_SUPTPRT3_DW2; MIZAR_USB_PORTSC_20; MIZAR_USB_DBOFF; Event_Ring_Segment_Table; MIZAR_USB_HCSPARAMS2; MIZAR_USB_PAGESIZE; Scratchpad_Buffer_Array; Device_Context_Base_Address_Array; MIZAR_USB_CRCR_LO; MIZAR_USB_CRCR_HI; MIZAR_USB_CONFIG; MIZAR_USB_DCBAAP_LO; MIZAR_USB_DCBAAP_HI; MIZAR_USB_ERSTSZ; MIZAR_USB_ERDP_LO; MIZAR_USB_ERDP_HI; MIZAR_USB_ERSTBA_LO; MIZAR_USB_ERSTBA_HI; MIZAR_USB_IMOD; MIZAR_USB_IMAN; MIZAR_USB_USBCMD; MIZAR_LSS_SYSREG_INTR_EN0; MIZAR_USB_USBSTS; Default_Input_Context; Default_Command_Ring; MIZAR_USB_DB; Default_Event_Ring_Array; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0; EP0_TR_Dequeue_Pointer",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB",
        "Meta Validation / Acceptance Criteria": "1. First interrupt wait: int_pend must be cleared.\n2-14. [Full criteria as specified]",
        "Validation / Acceptance Criteria": "1. Port status change event is received via interrupt.\n2. Port reset completes successfully.\n3. Enable Slot command completes successfully.\n4. Address Device command with BSR=1 completes.\n5. Address Device command with BSR=0 completes.\n6. Configure Endpoint command completes.\n7. All enumeration transfers complete.\n8. Event ring completion entry at offset 0xf0 becomes non-zero.\n9. IRQ handler correctly clears IMAN, clears system register interrupt, and clears GIC IRQ.\n10. Test completes with finish(0) indicating successful pass.",
        "Remarks": "This is a USB Host mode xHCI enumeration test for Full-Speed device. The slot context value 0x08100000 in set_address and 0x38100000 in SET_CONFIGURATION indicate Full-Speed."
    },
    {
        "Index": "5",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_hs",
        "Feature": "USB Host High-Speed Enumeration",
        "Meta Headers": '<stdio.h>; <stdlib.h>; "usb.h"',
        "Meta Macros": "NA",
        "Meta Arrays": "int data_in[512]; /* declared local array */\nint data_out[512]; /* declared local array */",
        "Speed": "High-Speed",
        "Mode": "Host",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Host mode High-Speed enumeration using the xHCI host controller.",
        "Test Description": "Verify USB Host mode High-Speed enumeration by performing xHCI controller initialization, port configuration with wake enables, event ring and command ring setup, scratchpad buffer and device context base address array configuration, port connect detection, port reset, Enable Slot command, Address Device commands (BSR=1 then BSR=0) with High-Speed slot context, Configure Endpoint command, and full High-Speed enumeration sequence including GET_DEVICE_DESCRIPTOR, GET_QUALIFIER, GET_CONFIGURATION_DESCRIPTOR (two short reads and one full read), and SET_CONFIGURATION.",
        "Meta Test Steps / Procedure": "1. Call nic_programming().\n2-77. [Full procedure as specified]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Configure global controller registers GCTL, GFLADJ, and GUCTL.\n3. Configure PIPE control register and PHY control register.\n4. Read host capability parameters.\n5. Enable port wake on connect, disconnect, and over-current events.\n6. Read Doorbell Offset register.\n7. Set up Event Ring Segment Table.\n8. Read HCSPARAMS2 and PAGESIZE.\n9. Configure Scratchpad Buffer Array.\n10. Load Device Context Base Address Array.\n11. Configure Command Ring.\n12. Configure CONFIG register.\n13. Configure DCBAAP.\n14. Configure Event Ring registers.\n15. Configure IMOD and IMAN.\n16. Enable USBCMD interrupt enable, sysreg interrupt, run.\n17. Wait for first interrupt.\n18. Read and clear USBSTS, read port connect status.\n19. Advance ERDP and issue port reset.\n20. Clear USBSTS, update IMAN, write PORTSC.\n21. Wait for second interrupt.\n22. Execute set_address with High-Speed slot context.\n23. Configure Input Context for SET_CONFIGURATION with High-Speed slot context.\n24. Issue Configure Endpoint command.\n25. Execute enumeration: GET_DEVICE_DESCRIPTOR, GET_QUALIFIER, GET_CONFIGURATION_DESCRIPTOR (two short reads and one full read), SET_CONFIGURATION.\n26. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "MIZAR_USB_GCTL; MIZAR_USB_GFLADJ; MIZAR_USB_GUCTL; MIZAR_USB_BASE; MIZAR_USB_HCSPARAMS1; MIZAR_USB_SUPTPRT2_DW2; MIZAR_USB_SUPTPRT3_DW2; MIZAR_USB_PORTSC_20; MIZAR_USB_DBOFF; Event_Ring_Segment_Table; MIZAR_USB_HCSPARAMS2; MIZAR_USB_PAGESIZE; Scratchpad_Buffer_Array; Device_Context_Base_Address_Array; MIZAR_USB_CRCR_LO; MIZAR_USB_CRCR_HI; MIZAR_USB_CONFIG; MIZAR_USB_DCBAAP_LO; MIZAR_USB_DCBAAP_HI; MIZAR_USB_ERSTSZ; MIZAR_USB_ERDP_LO; MIZAR_USB_ERDP_HI; MIZAR_USB_ERSTBA_LO; MIZAR_USB_ERSTBA_HI; MIZAR_USB_IMOD; MIZAR_USB_IMAN; MIZAR_USB_USBCMD; MIZAR_LSS_SYSREG_INTR_EN0; MIZAR_USB_USBSTS; Default_Input_Context; Default_Command_Ring; MIZAR_USB_DB; Default_Event_Ring_Array; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0; EP0_TR_Dequeue_Pointer",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB",
        "Meta Validation / Acceptance Criteria": "1. First interrupt wait: int_pend must be cleared.\n2-14. [Full criteria as specified]",
        "Validation / Acceptance Criteria": "1. Port status change event is received via interrupt.\n2. Port reset completes successfully.\n3. Enable Slot command completes.\n4. Address Device command with BSR=1 completes.\n5. Address Device command with BSR=0 completes.\n6. Configure Endpoint command completes.\n7. All enumeration transfers complete: GET_DEVICE_DESCRIPTOR, GET_QUALIFIER, GET_CONFIGURATION_DESCRIPTOR (two short reads and one full read), and SET_CONFIGURATION.\n8. Event ring completion entry at offset 0x150 becomes non-zero.\n9. IRQ handler correctly clears IMAN, clears system register interrupt, and clears GIC IRQ.\n10. Test completes with finish(0) indicating successful pass.",
        "Remarks": "This is a USB Host mode xHCI enumeration test for High-Speed device. The slot context value 0x08300000 in set_address and 0x38300000 in SET_CONFIGURATION indicate High-Speed device speed (speed field = 3). The GFLADJ value is 0xa07f000."
    },
    {
        "Index": "6",
        "SS / Module": "USB",
        "Test Case Name": "usb_host_enumeration_ls",
        "Feature": "USB Host Low-Speed Enumeration",
        "Meta Headers": '<stdio.h>; <stdlib.h>; "usb.h"',
        "Meta Macros": "NA",
        "Meta Arrays": "int data_in[512]; /* declared local array */\nint data_out[512]; /* declared local array */",
        "Speed": "Low-Speed",
        "Mode": "Host",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Host mode Low-Speed enumeration using the xHCI host controller.",
        "Test Description": "Verify USB Host mode Low-Speed enumeration by performing xHCI controller initialization, port configuration with wake enables, event ring and command ring setup, scratchpad buffer and device context base address array configuration, port connect detection, port reset, Enable Slot command, Address Device commands (BSR=1 then BSR=0) with Low-Speed slot context and 8-byte max packet size, and enumeration sequence including GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR (short read), SET_CONFIGURATION, GET_CONFIGURATION2 (short read with 8-byte length), and GET_CONFIGURATION3 (24-byte read).",
        "Meta Test Steps / Procedure": "1. Call nic_programming().\n2-108. [Full procedure as specified]",
        "Test Steps / Procedure": "1. Initialize NIC programming and enable GIC interrupts.\n2. Configure global controller registers GCTL, GFLADJ, and GUCTL.\n3. Configure PIPE control register and PHY control register.\n4. Read host capability parameters.\n5. Enable port wake on connect, disconnect, and over-current events.\n6. Read Doorbell Offset register.\n7. Set up Event Ring Segment Table with size 0x30.\n8. Read HCSPARAMS2 and PAGESIZE.\n9. Configure Scratchpad Buffer Array.\n10. Load Device Context Base Address Array.\n11. Configure Command Ring.\n12. Configure CONFIG register.\n13. Configure DCBAAP.\n14. Configure Event Ring registers.\n15. Configure IMOD and IMAN.\n16. Enable USBCMD interrupt enable, sysreg interrupt, run.\n17. Wait for first interrupt.\n18. Read and clear USBSTS, read port connect status.\n19. Advance ERDP and issue port reset.\n20. Clear USBSTS, update IMAN, write PORTSC.\n21. Wait for second interrupt.\n22. Issue Enable Slot command inline.\n23. Wait for third interrupt.\n24. Configure Input Context for Address Device BSR=1 with Low-Speed slot context.\n25. Issue Address Device BSR=1, ring doorbell.\n26. Wait for fourth interrupt, read event completion.\n27. Reconfigure Input Context for Address Device BSR=0.\n28. Issue Address Device BSR=0, ring doorbell.\n29. Wait for fifth interrupt, read event completion.\n30. Execute enumeration: GET_DEVICE_DESCRIPTOR, GET_CONFIGURATION_DESCRIPTOR, SET_CONFIGURATION, GET_CONFIGURATION2, GET_CONFIGURATION3.\n31. Call finish(0) to complete the test.",
        "Meta Impacted Registers": "MIZAR_USB_GCTL; MIZAR_USB_GFLADJ; MIZAR_USB_GUCTL; MIZAR_USB_BASE; MIZAR_USB_HCSPARAMS1; MIZAR_USB_SUPTPRT2_DW2; MIZAR_USB_SUPTPRT3_DW2; MIZAR_USB_PORTSC_20; MIZAR_USB_DBOFF; Event_Ring_Segment_Table; MIZAR_USB_HCSPARAMS2; MIZAR_USB_PAGESIZE; Scratchpad_Buffer_Array; Device_Context_Base_Address_Array; MIZAR_USB_CRCR_LO; MIZAR_USB_CRCR_HI; MIZAR_USB_CONFIG; MIZAR_USB_DCBAAP_LO; MIZAR_USB_DCBAAP_HI; MIZAR_USB_ERSTSZ; MIZAR_USB_ERDP_LO; MIZAR_USB_ERDP_HI; MIZAR_USB_ERSTBA_LO; MIZAR_USB_ERSTBA_HI; MIZAR_USB_IMOD; MIZAR_USB_IMAN; MIZAR_USB_USBCMD; MIZAR_LSS_SYSREG_INTR_EN0; MIZAR_USB_USBSTS; Default_Command_Ring; MIZAR_USB_DB; Default_Input_Context; Default_Event_Ring_Array; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0; EP0_TR_Dequeue_Pointer",
        "Impacted Registers": "GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB",
        "Meta Validation / Acceptance Criteria": "1. First interrupt wait: int_pend must be cleared.\n2-11. [Full criteria as specified]",
        "Validation / Acceptance Criteria": "1. Port status change event is received via interrupt.\n2. Port reset completes successfully.\n3. Enable Slot command completes.\n4. Address Device command with BSR=1 completes.\n5. Address Device command with BSR=0 completes.\n6. All enumeration transfers complete.\n7. Event ring completion entry at offset 0x110 becomes non-zero.\n8. IRQ handler correctly clears IMAN, clears system register interrupt, and clears GIC IRQ.\n9. Test completes with finish(0) indicating successful pass.",
        "Remarks": "This is a USB Host mode xHCI enumeration test for Low-Speed device. The slot context value 0x08200000 indicates Low-Speed device speed (speed field = 2). The EP0 max packet size is 8 bytes. The SET_CONFIGURATION Input Context and Configure Endpoint command block is entirely commented out."
    }
]

# TestPlan columns
tp_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers", "Validation / Acceptance Criteria",
    "Code Generation"
]

# MetaData columns
md_columns = [
    "Index", "Test Case Name", "Meta Test Description", "Meta Test Steps / Procedure",
    "Meta Impacted Registers", "Meta Validation / Acceptance Criteria",
    "Meta Headers", "Meta Macros", "Meta Arrays"
]

# Create workbook
wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = "TestPlan"

# Header formatting
header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap_alignment = Alignment(wrap_text=True, vertical='top')

# Write TestPlan headers
for col_idx, col_name in enumerate(tp_columns, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# Write TestPlan data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(tp_columns, 1):
        value = row_data.get(col_name, "")
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# Freeze first row
ws_tp.freeze_panes = 'A2'

# Auto-size columns with max width
for col_idx in range(1, len(tp_columns) + 1):
    max_length = len(str(ws_tp.cell(row=1, column=col_idx).value or ""))
    for row_idx in range(2, len(json_data) + 2):
        cell_value = str(ws_tp.cell(row=row_idx, column=col_idx).value or "")
        # For multi-line, use the longest line
        for line in cell_value.split('\n'):
            max_length = max(max_length, len(line))
    # Cap at reasonable width
    adjusted_width = min(max_length + 2, 80)
    adjusted_width = max(adjusted_width, 12)
    ws_tp.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

# --- MetaData Sheet ---
ws_md = wb.create_sheet("MetaData")

# Write MetaData headers
for col_idx, col_name in enumerate(md_columns, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_alignment

# Write MetaData data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(md_columns, 1):
        value = row_data.get(col_name, "")
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_alignment

# Freeze first row
ws_md.freeze_panes = 'A2'

# Auto-size columns with max width
for col_idx in range(1, len(md_columns) + 1):
    max_length = len(str(ws_md.cell(row=1, column=col_idx).value or ""))
    for row_idx in range(2, len(json_data) + 2):
        cell_value = str(ws_md.cell(row=row_idx, column=col_idx).value or "")
        for line in cell_value.split('\n'):
            max_length = max(max_length, len(line))
    adjusted_width = min(max_length + 2, 80)
    adjusted_width = max(adjusted_width, 12)
    ws_md.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

# Set MetaData sheet to veryHidden
ws_md.sheet_state = 'veryHidden'

# Save workbook
output_path = filename
wb.save(output_path)

# Validate
assert os.path.exists(output_path), "File does not exist"
assert os.path.getsize(output_path) > 0, "File is empty"

# Reopen to validate
wb_check = load_workbook(output_path)
assert "TestPlan" in wb_check.sheetnames, "TestPlan sheet missing"
assert "MetaData" in wb_check.sheetnames, "MetaData sheet missing"
assert wb_check["MetaData"].sheet_state == 'veryHidden', "MetaData not veryHidden"
tp_rows = wb_check["TestPlan"].max_row - 1  # minus header
md_rows = wb_check["MetaData"].max_row - 1
wb_check.close()

print(f"SUCCESS: {output_path}")
print(f"File size: {os.path.getsize(output_path)} bytes")
print(f"TestPlan rows: {tp_rows}")
print(f"MetaData rows: {md_rows}")
print(f"Validation: PASSED")

# Output base64 for upload
import base64
with open(output_path, 'rb') as f:
    b64_content = base64.b64encode(f.read()).decode('utf-8')
print(f"BASE64_LENGTH: {len(b64_content)}")

# Write base64 to file for retrieval
with open(output_path + '.b64', 'w') as f:
    f.write(b64_content)
print(f"Base64 file written: {output_path}.b64")
