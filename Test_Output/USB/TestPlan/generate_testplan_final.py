#!/usr/bin/env python3
"""USB TestPlan Excel Generator - Agent 7 Direct Generation"""
import json
import os
import sys
from datetime import datetime, timezone, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

# ============================================================
# INPUT DATA
# ============================================================
json_data = [
    {"Index":"1","SS / Module":"USB","Test Case Name":"USB_FS_Device_Bulk_Transfer_test","Feature":"Full-Speed Device Bulk Transfer","Meta Headers":"<stdio.h>; <stdlib.h>; \"usb.h\"","Meta Macros":"Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0; Buffer_PointerLO_1; Default_Event_Ring_Array; DWORD","Meta Arrays":"buf_data[16]","Speed":"Full-Speed","Mode":"Device Mode","Memory Start Offset":"NA","Memory End Offset":"NA","Meta Test Description":"This testcase validates USB Full-Speed Device mode Bulk Transfer operation.","Test Description":"This test validates USB Full-Speed Device mode Bulk Transfer functionality. The test performs a controller soft reset via the DCTL register and polls for completion. It configures the USB2 PHY via the GUSB2PHYCFG register, sets up the event buffer using GEVNTADRLO, GEVNTADRHI, GEVNTSIZ, and GEVNTCOUNT registers. The controller is placed in device mode by configuring the GCTL register port capability direction field. Device configuration is set via the DCFG register, device events are enabled via the DEVTEN register, and the GUCTL register is configured for device timeout parameters. Eight endpoints are configured using endpoint command parameter registers (DEPCMDPAR0, DEPCMDPAR1) and the endpoint command register (DEPCMD). After enumeration, bulk OUT and bulk IN transfers of 64 bytes each are initiated.","Meta Test Steps / Procedure":"Steps 1-44 as generated","Test Steps / Procedure":"1. Initialize the NIC subsystem and enable all GIC interrupts. 2. Clear buffers. 3. Perform soft reset via DCTL. 4. Configure USB2 PHY. 5. Set up event buffer. 6. Configure GCTL for device mode. 7. Configure DCFG. 8. Enable device events via DEVTEN. 9. Configure GUCTL. 10. Configure 8 endpoints. 11. Allocate transfer resources. 12. Enable endpoints via DALEPENA. 13. Set run state via DCTL. 14. Enable interrupts. 15-22. Enumeration sequence. 23-25. Bulk transfers. 26-28. Validation and finish.","Meta Impacted Registers":"Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0; 0xa0243ff4; 0xa0243ff8","Impacted Registers":"DCTL; GUSB2PHYCFG; GEVNTADRLO; GEVNTADRHI; GEVNTSIZ; GEVNTCOUNT; GCTL; DCFG; DEVTEN; GUCTL; DEPCMDPAR0; DEPCMD; DALEPENA; DSTS; DEPCMDPAR1","Meta Validation / Acceptance Criteria":"Soft reset, endpoint commands, interrupts, handshake, bulk transfers validated","Validation / Acceptance Criteria":"1. Soft reset completes. 2. Endpoint commands complete. 3. Start Transfer commands complete. 4. Events received via interrupts. 5. Enumeration completes. 6. Bulk transfers complete. 7. Test passes.","Remarks":"Full-Speed Device mode, GIC IRQ 84, bulk TRBs on endpoints 0x40/0x50."},
    {"Index":"2","SS / Module":"USB","Test Case Name":"USB_FS_Device_Interrupt_Transfer_test","Feature":"Full-Speed Device Interrupt Transfer","Speed":"Full-Speed","Mode":"Device Mode","Memory Start Offset":"","Memory End Offset":"","Impacted Registers":"DCTL; GUSB2PHYCFG; GEVNTADRLO; GEVNTADRHI; GEVNTSIZ; GEVNTCOUNT; GCTL; DCFG; DEVTEN; GUCTL; DEPCMDPAR0; DEPCMD; DALEPENA; DSTS; DEPCMDPAR1","Test Description":"This test validates USB Full-Speed Device mode Interrupt Transfer functionality. Similar to bulk test but with interrupt-type endpoints and TRB types 0x815/0x813. DSTS polled for frame number.","Test Steps / Procedure":"Same as Folder 1 with interrupt endpoints at offsets 0x20/0x30 and DSTS frame number polling.","Validation / Acceptance Criteria":"Interrupt transfers of 64 bytes complete. DSTS frame number validated.","Remarks":"Interrupt TRB types 0x815 (OUT) and 0x813 (IN) on endpoints 0x20/0x30."},
    {"Index":"3","SS / Module":"USB","Test Case Name":"USB_FS_Device_Isochronous_Transfer_test","Feature":"Full-Speed Device Isochronous Transfer","Speed":"Full-Speed","Mode":"Device Mode","Memory Start Offset":"","Memory End Offset":"","Impacted Registers":"DCTL; GUSB2PHYCFG; GEVNTADRLO; GEVNTADRHI; GEVNTSIZ; GEVNTCOUNT; GCTL; DCFG; DEVTEN; GUCTL; DEPCMDPAR0; DEPCMD; DALEPENA; DSTS; DEPCMDPAR1","Test Description":"This test validates USB Full-Speed Device mode Isochronous Transfer functionality. Uses isochronous TRB type 0x869 with 1023-byte transfers and frame-number-qualified Start Transfer commands.","Test Steps / Procedure":"Same init as Folder 1, isochronous transfers on endpoints 0x60/0x70 with frame number synchronization via DSTS polling.","Validation / Acceptance Criteria":"Isochronous transfers of 1023 bytes complete. Frame numbers validated via DSTS.","Remarks":"Isochronous TRB type 0x869, frame qualifiers 0x20506/0x40506, endpoints 0x60/0x70."},
    {"Index":"4","SS / Module":"USB","Test Case Name":"usb_host_enumeration_fs","Feature":"Host Enumeration Full-Speed","Speed":"Full-Speed","Mode":"Host Mode","Memory Start Offset":"","Memory End Offset":"","Impacted Registers":"GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB","Test Description":"Host mode Full-Speed enumeration via xHCI. Enable Slot, Address Device (BSR=1/0), Configure Endpoint, enumeration with GET DEVICE DESCRIPTOR, GET CONFIG DESCRIPTOR, SET CONFIGURATION.","Test Steps / Procedure":"xHCI init, port detect, reset, Enable Slot, Address Device, Configure Endpoint, enumeration TRBs, poll Event Ring.","Validation / Acceptance Criteria":"All xHCI commands complete. Enumeration transfers complete. Event Ring polled successfully.","Remarks":"xHCI host mode, FS device, slot context 0x08100000."},
    {"Index":"5","SS / Module":"USB","Test Case Name":"usb_host_enumeration_hs","Feature":"Host Enumeration High-Speed","Speed":"High-Speed","Mode":"Host Mode","Memory Start Offset":"","Memory End Offset":"","Impacted Registers":"GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB","Test Description":"Host mode High-Speed enumeration via xHCI. Includes GET DEVICE QUALIFIER DESCRIPTOR. Slot context 0x08300000 for HS.","Test Steps / Procedure":"Same as FS host with additional GET QUALIFIER DESCRIPTOR. Final Event Ring poll at +0x150.","Validation / Acceptance Criteria":"All xHCI commands and enumeration transfers including qualifier descriptor complete.","Remarks":"xHCI host mode, HS device, slot context 0x08300000, includes qualifier descriptor."},
    {"Index":"6","SS / Module":"USB","Test Case Name":"usb_host_enumeration_ls","Feature":"Host Enumeration Low-Speed","Speed":"Low-Speed","Mode":"Host Mode","Memory Start Offset":"","Memory End Offset":"","Impacted Registers":"GCTL; GFLADJ; GUCTL; HCSPARAMS1; SUPTPRT2_DW2; SUPTPRT3_DW2; PORTSC_20; DBOFF; HCSPARAMS2; PAGESIZE; CRCR_LO; CRCR_HI; CONFIG; DCBAAP_LO; DCBAAP_HI; ERSTSZ; ERDP_LO; ERDP_HI; ERSTBA_LO; ERSTBA_HI; IMOD; IMAN; USBCMD; USBSTS; DB","Test Description":"Host mode Low-Speed enumeration via xHCI. MaxPacketSize=8. No qualifier descriptor. Config descriptor 24 bytes.","Test Steps / Procedure":"Same as FS host with LS-specific parameters. MaxPacketSize=8, config descriptor 24 bytes.","Validation / Acceptance Criteria":"All xHCI commands and enumeration transfers complete for LS device.","Remarks":"xHCI host mode, LS device, slot context 0x08200000, MaxPacketSize=8, no qualifier descriptor."}
]

# ============================================================
# CONFIGURATION
# ============================================================
IP_NAME = "USB"
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"{IP_NAME}_TestPlan_{timestamp}.xlsx"
output_dir = os.path.dirname(os.path.abspath(__file__))
output_path = os.path.join(output_dir, filename)

# ============================================================
# COLUMN DEFINITIONS
# ============================================================
TESTPLAN_COLS = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers",
    "Validation / Acceptance Criteria", "Code Generation"
]

METADATA_COLS = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

# ============================================================
# WORKBOOK CREATION
# ============================================================
print(f"Creating workbook: {filename}")
wb = Workbook()

# --- TestPlan Sheet ---
ws_tp = wb.active
ws_tp.title = "TestPlan"

# Header formatting
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_align = Alignment(vertical="top", wrap_text=True)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

# Write TestPlan headers
for col_idx, col_name in enumerate(TESTPLAN_COLS, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Write TestPlan data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(TESTPLAN_COLS, 1):
        value = row_data.get(col_name, "")
        if value is None:
            value = ""
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=str(value))
        cell.alignment = cell_align
        cell.border = thin_border

# Freeze first row
ws_tp.freeze_panes = "A2"

# --- MetaData Sheet ---
ws_md = wb.create_sheet(title="MetaData")

# Write MetaData headers
for col_idx, col_name in enumerate(METADATA_COLS, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Write MetaData data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(METADATA_COLS, 1):
        value = row_data.get(col_name, "")
        if value is None:
            value = ""
        cell = ws_md.cell(row=row_idx, column=col_idx, value=str(value))
        cell.alignment = cell_align
        cell.border = thin_border

# Freeze first row
ws_md.freeze_panes = "A2"

# Set MetaData sheet to veryHidden
ws_md.sheet_state = "veryHidden"

# ============================================================
# AUTO-SIZE COLUMNS
# ============================================================
def auto_size_columns(ws, max_width=60):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        adjusted = min(max(max_len + 2, 12), max_width)
        ws.column_dimensions[col_letter].width = adjusted

auto_size_columns(ws_tp)
auto_size_columns(ws_md)

# ============================================================
# SAVE WORKBOOK
# ============================================================
wb.save(output_path)
print(f"Workbook saved: {output_path}")

# ============================================================
# VALIDATION
# ============================================================
if os.path.exists(output_path) and os.path.getsize(output_path) > 0:
    wb_check = load_workbook(output_path)
    sheets = wb_check.sheetnames
    assert "TestPlan" in sheets, "TestPlan sheet missing"
    assert "MetaData" in sheets, "MetaData sheet missing"
    tp_rows = wb_check["TestPlan"].max_row - 1
    md_rows = wb_check["MetaData"].max_row - 1
    print(f"VALIDATION PASSED")
    print(f"TestPlan rows: {tp_rows}")
    print(f"MetaData rows: {md_rows}")
    print(f"File size: {os.path.getsize(output_path)} bytes")
    print(f"FILENAME={filename}")
else:
    print("VALIDATION FAILED")
    sys.exit(1)
