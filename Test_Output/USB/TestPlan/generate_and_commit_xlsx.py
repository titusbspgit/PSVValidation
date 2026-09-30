#!/usr/bin/env python3
"""USB TestPlan Excel Generator - Agent 7 Direct Generation
Generates USB_TestPlan_<YYYYMMDD>_<HHMMSS>.xlsx with TestPlan and MetaData sheets.
Run: python generate_and_commit_xlsx.py
"""
import datetime
import os
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

# ============================================================
# DATA
# ============================================================
json_data = [
    {
        "Index": "1",
        "SS / Module": "USB",
        "Test Case Name": "USB_FS_Device_Bulk_Transfer_test",
        "Feature": "USB FS Device Bulk Transfer",
        "Meta Headers": '<stdio.h>; <stdlib.h>; "usb.h"',
        "Meta Macros": "NA",
        "Meta Arrays": "NA",
        "Speed": "Full-Speed",
        "Mode": "Device",
        "Memory Start Offset": "NA",
        "Memory End Offset": "NA",
        "Meta Test Description": "This testcase validates USB Full-Speed Device mode Bulk Transfer operation. The test performs a soft reset of the USB controller via MIZAR_USB_DCTL, configures USB2 PHY via MIZAR_USB_GUSB2PHYCFG, sets up the event ring (GEVNTADRLO, GEVNTADRHI, GEVNTSIZ, GEVNTCOUNT), configures the global controller (GCTL), device configuration (DCFG), device event enable (DEVTEN), and GUCTL registers. It then issues endpoint configuration commands via set_configuration() for 9 endpoints (START NEW CONFIGURATION plus 8 endpoint configurations). A transfer resource configuration loop runs for 8 endpoints writing DEPCMDPAR0 and DEPCMD with polling. Physical endpoints 0 and 1 are enabled via DALEPENA, then DCTL run/stop is set. The system-level interrupt is enabled via MIZAR_LSS_SYSREG_INTR_EN0. The test then waits for link state connect/reset events via interrupt waits. After enumeration speed detection (reading DSTS), the test performs USB enumeration consisting of: setup_stage for GET_DESCRIPTOR (Device Descriptor), data stage with 5 DWORDs of device descriptor data, status_stage; setup_stage for GET_DESCRIPTOR (Configuration Descriptor short), data stage with 4 DWORDs, status_stage; setup_stage for SET_CONFIGURATION, data stage with 1 DWORD, status_stage; setup_stage for GET_DESCRIPTOR (Full Configuration Descriptor), data stage with 15 DWORDs of full configuration descriptor data, status_stage. After enumeration, the test performs two Bulk OUT transfers on physical endpoint 4 (offset 0x40) and physical endpoint 5 (offset 0x50) using Buffer_PointerLO_1 with transfer size 0x40 and TRB control 0x813, each followed by interrupt waits. The test writes a final marker 0xdeadbee7 and waits for a final interrupt before calling finish(0). An IRQ handler (Default_IRQHandler) reads MIZAR_LSS_SYSREG_MSK_STS0, MIZAR_LSS_SYSREG_RAW_STCR0, reads and writes back GEVNTCOUNT, clears the sysreg interrupt, and clears GIC IRQ 84.",
        "Test Description": "Verify USB Full-Speed Device mode Bulk Transfer operation by performing controller soft reset, PHY configuration, event ring setup, global and device configuration, endpoint configuration for 8 endpoints, link state event handling, full USB enumeration (Device Descriptor, Configuration Descriptor short, SET_CONFIGURATION, Full Configuration Descriptor), and two Bulk OUT data transfers on bulk endpoints, with interrupt-driven synchronization throughout.",
        "Meta Test Steps / Procedure": "1. Call nic_programming() for NIC initialization.\n2. Call GIC_EnableAllIRQ() to enable all GIC interrupts.\n3. Loop j=0 to 19: write_reg(Buffer_PointerLO + jDWORD, 0x0); write_reg(event_trb_addr + jDWORD, 0x0); \u2014 clears 20 DWORDs at Buffer_PointerLO and 20 DWORDs at event_trb_addr.\n4. wr_data = 0x40f00000; write_reg(MIZAR_USB_DCTL, 0x40f00000) \u2014 initiate soft reset.\n5. rd_data = read_reg(MIZAR_USB_DCTL); poll while(rd_data != 0xf00000) { wait_on(5); rd_data = read_reg(MIZAR_USB_DCTL); } \u2014 wait for soft reset completion.\n6. printf(\"Soft Rst is done\\n\");\n7. wr_data = 0x40002407; write_reg(MIZAR_USB_GUSB2PHYCFG, 0x40002407) \u2014 configure USB2 PHY.\n8. write_reg(MIZAR_USB_GEVNTADRLO, Default_Event_Ring_Array) \u2014 set event ring base address low.\n9. write_reg(MIZAR_USB_GEVNTADRHI, 0x0) \u2014 set event ring base address high.\n10. write_reg(MIZAR_USB_GEVNTSIZ, 0x30) \u2014 set event ring size.\n11. write_reg(MIZAR_USB_GEVNTCOUNT, 0x0) \u2014 clear event count.\n12. rd_data = read_reg(MIZAR_USB_GCTL) \u2014 read global control.\n13. wr_data = 0x30c12214; write_reg(MIZAR_USB_GCTL, 0x30c12214) \u2014 configure global control (port direction).\n14. rd_data = read_reg(MIZAR_USB_DCFG) \u2014 read device configuration.\n15. write_reg(MIZAR_USB_DCFG, 0x480801) \u2014 set device configuration.\n16. write_reg(MIZAR_USB_DEVTEN, 0x1f) \u2014 enable device events (bits 0-4).\n17. rd_data = read_reg(MIZAR_USB_GUCTL) \u2014 read GUCTL.\n18. write_reg(MIZAR_USB_GUCTL, 0xa400010) \u2014 configure GUCTL.\n19-27. Configure 9 endpoints via set_configuration calls.\n28. Loop i=0 to 7: TX resource configuration for 8 endpoints.\n29-31. Enable endpoints, set run/stop, enable system interrupt.\n32-43. Link state handling, speed detection, endpoint enable.\n44-58. Setup stage, SET ADDRESS, handshake polling.\n59. Full enumeration sequence (GET_DESCRIPTOR Device, Config short, SET_CONFIGURATION, Config full).\n60-74. Handshake, bulk transfers on EP4 and EP5, final marker, finish(0).\n75-82. IRQ Handler: clear pending, read status, clear events, clear GIC IRQ 84.",
        "Test Steps / Procedure": "1. Initialize NIC and enable all GIC interrupts.\n2. Clear 20 DWORDs at the buffer pointer and event TRB address regions.\n3. Perform USB controller soft reset by writing to DCTL and polling until reset completes.\n4. Configure USB2 PHY settings.\n5. Set up the event ring by writing event ring base address (low and high), event ring size, and clearing event count.\n6. Configure global controller register for port direction.\n7. Set device configuration, enable device events (disconnect, reset, connection done, link state change, wakeup), and configure GUCTL.\n8. Issue START NEW CONFIGURATION command on endpoint 0.\n9. Configure 8 physical endpoints (EP0-EP7) with their respective parameters using set_configuration calls.\n10. Issue Transfer Resource Configuration command for all 8 endpoints in a loop, polling each for completion.\n11. Enable physical endpoints 0 and 1 via DALEPENA, set DCTL run/stop, and enable system-level interrupt.\n12. Wait for link state connect/reset event interrupts with conditional additional waits based on event_counter.\n13. Re-write device configuration after link events.\n14. Read device status for speed detection, update device configuration and DCTL, enable all 8 physical endpoints.\n15. Execute setup_stage for initial control transfer setup, then update USB2 PHY configuration.\n16. Wait for conditional interrupts, then set device address in DCFG.\n17. Prepare and issue a Start Transfer command on EP1 for SET ADDRESS status, poll for completion.\n18. Wait for interrupts, then poll handshake register until non-zero.\n19. Execute full enumeration sequence: GET_DESCRIPTOR (Device Descriptor) with 5 DWORDs of descriptor data and status stage; GET_DESCRIPTOR (Configuration Descriptor short) with 4 DWORDs and status stage; SET_CONFIGURATION with zero data and status stage; GET_DESCRIPTOR (Full Configuration Descriptor) with 15 DWORDs and status stage.\n20. Poll second handshake register until non-zero.\n21. Write enumeration complete marker, wait for interrupt.\n22. Execute Bulk OUT transfer on physical endpoint 4 with 64-byte transfer size, poll for command completion, wait for two interrupts.\n23. Execute Bulk OUT transfer on physical endpoint 5 with 64-byte transfer size, poll for command completion, wait for interrupt.\n24. Write final marker, wait for final interrupt, and call finish to complete the test with pass status.",
        "Meta Impacted Registers": "Buffer_PointerLO; event_trb_addr; MIZAR_USB_DCTL; MIZAR_USB_GUSB2PHYCFG; MIZAR_USB_GEVNTADRLO; MIZAR_USB_GEVNTADRHI; MIZAR_USB_GEVNTSIZ; MIZAR_USB_GEVNTCOUNT; MIZAR_USB_GCTL; MIZAR_USB_DCFG; MIZAR_USB_DEVTEN; MIZAR_USB_GUCTL; MIZAR_USB_DEPCMDPAR0; MIZAR_USB_DEPCMD; MIZAR_USB_DALEPENA; MIZAR_LSS_SYSREG_INTR_EN0; 0xA0243ffc; MIZAR_USB_DSTS; MIZAR_USB_DEPCMDPAR1; 0xa0243ff4; 0xa0243ff8; MIZAR_LSS_SYSREG_MSK_STS0; MIZAR_LSS_SYSREG_RAW_STCR0",
        "Impacted Registers": "DCTL; GEVNTSIZ; GCTL; DEVTEN; GUCTL; DALEPENA; DSTS",
        "Meta Validation / Acceptance Criteria": "1. Soft reset completion: poll read_reg(MIZAR_USB_DCTL) until value equals 0xf00000 (bit 30 cleared).\n2. Endpoint configuration command completion: for each set_configuration call, poll read_reg(MIZAR_USB_DEPCMD+trb_address) until value != cmd (command active bit cleared). Applies to 9 set_configuration calls with cmd values 0x409 and 0x401.\n3. Transfer resource configuration: for i=0..7, poll read_reg(MIZAR_USB_DEPCMD+(i*0x10)) until value != 0x402.\n4. Interrupt-driven synchronization: int_pend flag set to 1 before each wait, cleared to 0 by Default_IRQHandler. Multiple interrupt waits throughout test flow must each complete (int_pend transitions from 1 to 0).\n5. Event counter checks: conditional interrupt waits gated by (event_counter <= 0x4).\n6. Start Transfer command completion: poll read_reg(MIZAR_USB_DEPCMD+offset) until value != 0x506 for offsets 0x0, 0x10, 0x40, 0x50 at various points.\n7. Handshake polling at 0xa0243ff4: poll until non-zero value read.\n8. Handshake polling at 0xa0243ff8: poll until non-zero value read.\n9. IRQ handler validation: read_reg(MIZAR_LSS_SYSREG_MSK_STS0) and read_reg(MIZAR_LSS_SYSREG_RAW_STCR0) for interrupt status; if bit 31 set in rd_data, write 0x80000000 to MIZAR_LSS_SYSREG_RAW_STCR0 to clear; event_count read from MIZAR_USB_GEVNTCOUNT and written back to clear.\n10. Test completion: finish(0) called indicating pass (argument 0).\n11. Debug markers written to 0xA0243ffc: 0xdeadbee0, 0xdeadbee1, 0xdeadbee3, 0xdeadbee4, 0xdeadbee5, 0xdeadbee6, 0xdeadbee7 at various enumeration stages.",
        "Validation / Acceptance Criteria": "1. USB controller soft reset completes successfully (DCTL reset bit clears).\n2. All endpoint configuration commands complete (command active bit clears in endpoint command registers).\n3. All transfer resource configuration commands complete for 8 endpoints.\n4. All interrupt-driven synchronization points complete (IRQ handler fires and clears pending flag).\n5. Start Transfer commands complete on control and bulk endpoints.\n6. Handshake registers return non-zero values indicating host-side readiness.\n7. Full USB enumeration completes: Device Descriptor, Configuration Descriptor (short and full), and SET_CONFIGURATION all succeed.\n8. Two Bulk OUT transfers of 64 bytes each complete on bulk endpoints.\n9. IRQ handler correctly reads event count, clears event count register, clears system interrupt status, and clears GIC IRQ.\n10. Test completes with pass status (finish(0)).",
        "Remarks": "The test uses interrupt-driven synchronization with int_pend flag and Default_IRQHandler throughout. Conditional interrupt waits are gated by event_counter <= 0x4. The IRQ handler has a logical bug: uses && (logical AND) instead of & (bitwise AND) for checking bit 31 of rd_data. Buffer_PointerLO, Buffer_PointerLO_1, event_trb_addr, Default_Event_Ring_Array, and DWORD are external symbols defined in usb.h (not accessible in testcase folder). Several Agent 4 register mappings are unresolved including MIZAR_USB_GUSB2PHYCFG, MIZAR_USB_GEVNTADRLO, MIZAR_USB_GEVNTADRHI, MIZAR_USB_GEVNTCOUNT, MIZAR_USB_DCFG, MIZAR_USB_DEPCMDPAR0, MIZAR_USB_DEPCMD, MIZAR_USB_DEPCMDPAR1, and system register macros. Hardcoded addresses 0xA0243ffc, 0xa0243ff4, 0xa0243ff8 are used for debug markers and handshake polling. The status_stage() function contains two separate interrupt waits with a wait_on(5) delay between them.",
        "Code Generation": ""
    }
]

# ============================================================
# COLUMN DEFINITIONS
# ============================================================
testplan_columns = [
    "Index", "SS / Module", "Feature", "Test Case Name", "Test Description",
    "Speed", "Mode", "Memory Start Offset", "Memory End Offset", "Remarks",
    "Test Steps / Procedure", "Impacted Registers",
    "Validation / Acceptance Criteria", "Code Generation"
]

metadata_columns = [
    "Index", "Test Case Name", "Meta Test Description",
    "Meta Test Steps / Procedure", "Meta Impacted Registers",
    "Meta Validation / Acceptance Criteria", "Meta Headers",
    "Meta Macros", "Meta Arrays"
]

# ============================================================
# STYLES
# ============================================================
header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_alignment = Alignment(wrap_text=True, vertical="top")
header_alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

# ============================================================
# WORKBOOK CREATION
# ============================================================
def create_workbook():
    # IST timestamp
    ist = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
    now_ist = datetime.datetime.now(ist)
    timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
    filename = f"USB_TestPlan_{timestamp}.xlsx"

    wb = Workbook()

    # --- TestPlan Sheet ---
    ws_tp = wb.active
    ws_tp.title = "TestPlan"

    # Headers
    for col_idx, col_name in enumerate(testplan_columns, 1):
        cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Data rows
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(testplan_columns, 1):
            value = row_data.get(col_name, "")
            cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment

    # Freeze first row
    ws_tp.freeze_panes = "A2"

    # Auto-size columns
    for col_idx, col_name in enumerate(testplan_columns, 1):
        max_len = len(col_name)
        for row in range(2, len(json_data) + 2):
            val = str(ws_tp.cell(row=row, column=col_idx).value or "")
            lines = val.split("\n")
            for line in lines:
                if len(line) > max_len:
                    max_len = len(line)
        width = min(max_len + 4, 80)
        if width < 12:
            width = 12
        ws_tp.column_dimensions[get_column_letter(col_idx)].width = width

    # --- MetaData Sheet ---
    ws_md = wb.create_sheet(title="MetaData")

    # Headers
    for col_idx, col_name in enumerate(metadata_columns, 1):
        cell = ws_md.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Data rows
    for row_idx, row_data in enumerate(json_data, 2):
        for col_idx, col_name in enumerate(metadata_columns, 1):
            value = row_data.get(col_name, "")
            cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_alignment

    # Freeze first row
    ws_md.freeze_panes = "A2"

    # Auto-size columns
    for col_idx, col_name in enumerate(metadata_columns, 1):
        max_len = len(col_name)
        for row in range(2, len(json_data) + 2):
            val = str(ws_md.cell(row=row, column=col_idx).value or "")
            lines = val.split("\n")
            for line in lines:
                if len(line) > max_len:
                    max_len = len(line)
        width = min(max_len + 4, 80)
        if width < 12:
            width = 12
        ws_md.column_dimensions[get_column_letter(col_idx)].width = width

    # Set MetaData sheet to veryHidden
    ws_md.sheet_state = "veryHidden"

    # Save
    script_dir = os.path.dirname(os.path.abspath(__file__))
    filepath = os.path.join(script_dir, filename)
    wb.save(filepath)
    print(f"Workbook saved: {filepath}")

    # Validate
    assert os.path.exists(filepath), "File does not exist!"
    assert os.path.getsize(filepath) > 0, "File is empty!"
    wb2 = load_workbook(filepath)
    assert "TestPlan" in wb2.sheetnames, "TestPlan sheet missing!"
    assert "MetaData" in wb2.sheetnames, "MetaData sheet missing!"
    print(f"Validation PASSED. Sheets: {wb2.sheetnames}")
    print(f"TestPlan rows: {ws_tp.max_row - 1}")
    print(f"MetaData rows: {ws_md.max_row - 1}")
    print(f"Filename: {filename}")
    return filepath, filename

if __name__ == "__main__":
    create_workbook()
