#!/usr/bin/env python3
"""Build xlsx locally, print base64 for upload"""
import openpyxl, base64, sys, os
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
wb = openpyxl.Workbook()
hf = Font(bold=True, color="FFFFFF")
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap = Alignment(wrap_text=True, vertical="top")
ws = wb.active; ws.title = "TestPlan"
cols = ["Index","SS / Module","Feature","Test Case Name","Speed","Mode","Memory Start Offset","Memory End Offset","Test Description","Test Steps / Procedure","Impacted Registers","Validation / Acceptance Criteria","Remarks","Code Generation (Required / Not)"]
for c,h in enumerate(cols,1):
    cell=ws.cell(row=1,column=c,value=h); cell.font=hf; cell.fill=hfill; cell.alignment=wrap
ws.cell(row=2,column=1,value="1").alignment=wrap
ws.cell(row=2,column=2,value="USB").alignment=wrap
ws.cell(row=2,column=3,value="Full-Speed Device Bulk Transfer").alignment=wrap
ws.cell(row=2,column=4,value="USB_FS_Device_Bulk_Transfer_test").alignment=wrap
ws.cell(row=2,column=5,value="Full-Speed").alignment=wrap
ws.cell(row=2,column=6,value="Device").alignment=wrap
ws.cell(row=2,column=7,value="NA").alignment=wrap
ws.cell(row=2,column=8,value="NA").alignment=wrap
ws.cell(row=2,column=9,value="Test Description placeholder").alignment=wrap
ws.cell(row=2,column=10,value="Test Steps placeholder").alignment=wrap
ws.cell(row=2,column=11,value="GCTL, GUCTL, DALEPENA, DCFG, DCTL, DEVTEN, DSTS, GUSB2PHYCFG, GEVNTADRLO, GEVNTADRHI, GEVNTSIZ, GEVNTCOUNT, DEPCMDPAR2, DEPCMDPAR1, DEPCMDPAR0, DEPCMD").alignment=wrap
ws.cell(row=2,column=12,value="Validation placeholder").alignment=wrap
ws.cell(row=2,column=13,value="Remarks placeholder").alignment=wrap
dv=DataValidation(type="list",formula1='"Required,Not Required"',allow_blank=True)
dv.sqref="N2:N2"; ws.add_data_validation(dv)
ws.freeze_panes="A2"
md=wb.create_sheet("MetaData"); md.sheet_state="veryHidden"
mcols=["Test Case Name","Meta Headers","Meta Macros","Meta Arrays","Meta Types","Meta Test Description","Meta Test Steps / Procedure","Meta Impacted Registers","Meta Validation / Acceptance Criteria"]
for c,h in enumerate(mcols,1):
    cell=md.cell(row=1,column=c,value=h); cell.font=hf; cell.fill=hfill; cell.alignment=wrap
md.cell(row=2,column=1,value="USB_FS_Device_Bulk_Transfer_test").alignment=wrap
md.freeze_panes="A2"
p="/tmp/usb_tp.xlsx"; wb.save(p)
with open(p,"rb") as f: print(base64.b64encode(f.read()).decode())
