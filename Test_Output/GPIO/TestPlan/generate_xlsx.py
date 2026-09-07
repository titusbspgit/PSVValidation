#!/usr/bin/env python3
"""Auto-generated script to create GPIO TestPlan XLSX workbook.
Run: python3 generate_xlsx.py
Requires: pip install openpyxl
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone, timedelta
import os
import json

# IST timezone
IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp_str = now_ist.strftime('%Y%m%d_%H%M%S')
filename = f'GPIO_TestPlan_{timestamp_str}.xlsx'

# JSON data
json_data = [
    {
        "Index": "1",
        "SS / Module": "GPIO",
        "Test Case Name": "gpio_reg_wr_rd_test",
        "Feature": "Register Read/Write Verification",
        "Test Description": "This test verifies GPIO register default values after reset and performs write-read verification using multiple data patterns. In the first phase, all readable GPIO registers (gp0_gpio_8, gp0_gpio_9, gp0_gpio_10) are read and their values compared against expected default values. In the second phase, six distinct data patterns are written to all writable GPIO registers and then read back. The read-back value is compared against an expected value computed using the write mask and read mask for each register. The test passes only if all default value checks and all write-read comparisons succeed across all registers and all patterns.",
        "Speed": "NA",
        "Mode": "NA",
        "Memory Start Offset": "0x0",
        "Memory End Offset": "0x8",
        "Remarks": "The addr_array is declared with size 49 but only 3 register addresses are explicitly populated (remaining entries are zero-initialized). The skip_array and skip_rst_array contain flags to skip certain register indices during write-read and default-value checks respectively, indicating some registers (indices 32, 37-48) are VRRW type and excluded from standard write-read testing. The default value comparison applies a 0xfffffffe mask to exclude bit-0 (data_in pin level) which may fluctuate. The soft_reset_chk function is disabled via #ifdef 0 and is not executed.",
        "Test Steps / Procedure": "1. Initialize the test environment and load the register address table, default value table, read mask table, and write mask table for all GPIO registers under test.\n2. Perform default value verification: Read each readable GPIO register (gp0_gpio_8, gp0_gpio_9, gp0_gpio_10), apply a bit mask to the read value, and compare against the expected default value from the specification. Skip registers flagged in the reset-skip list or those that are not readable.\n3. Perform write-read verification: For each of six data patterns (all-ones, alternating-bit patterns, and half-word pattern), write the pattern (masked by the register's write mask) to each writable GPIO register.\n4. After writing each pattern, read back each register (masked by the register's read mask) and compare the result against an expected value computed from the written data, write mask, read mask, and default value. Skip registers flagged in the skip list or those that are not writable or readable.\n5. Repeat steps 3-4 for all six data patterns.\n6. Evaluate the overall result: if any default value mismatch or write-read mismatch was detected, report test failure. Otherwise, report test pass.",
        "Impacted Registers": "gp0_gpio_8; gp0_gpio_9; gp0_gpio_10",
        "Validation / Acceptance Criteria": "1. All readable GPIO registers (gp0_gpio_8, gp0_gpio_9, gp0_gpio_10) must return their specification-defined default values after reset, with bit-0 masked out during comparison.\n2. For each of the six write-read data patterns, every writable register must accept the written value (masked by its write mask) and the subsequent read-back (masked by its read mask) must match the expected value derived from the written data, write mask, read mask, and default value.\n3. The test passes only when zero default-value mismatches and zero write-read mismatches are detected across all registers and all patterns. Any single mismatch causes the test to fail.",
        "Meta Headers": "<stdio.h>; <stdlib.h>; \"test_common.h\"; \"test_define.c\"; <gpio/gpio_def.h>; <gpio/gpio_offset.h>",
        "Meta Macros": "CNT",
        "Meta Arrays": "addr_array[49]; default_value_array[49]; read_mask_array[49]; write_mask_array[49]; skip_array[49]; skip_rst_array[49]; chk_val[6]",
        "Meta Test Description": "This testcase performs two phases of GPIO register verification. Phase 1 (chk_rst_val): Iterates over addr_array[0..CNT-1] (CNT=49, only 3 addresses populated: MIZAR_GPIO_GP0_GPIO_8, MIZAR_GPIO_GP0_GPIO_9, MIZAR_GPIO_GP0_GPIO_10). For each register, if skip_rst_array[i]==1 or read_mask_array[i]==0x00000000, the register is skipped. Otherwise, read_reg(addr) is called, the result is masked with 0xfffffffe, and compared against default_value_array[i]. Mismatches increment def_fail_cnt. Phase 2 (chk_rd_wr): Iterates over 6 data patterns in chk_val[6]={0xffffffff, 0xaaaaaaaa, 0x55555555, 0xf5f5f5f5, 0xA5A5A5A5, 0xffff0000}. For each pattern, iterates addr_array. Registers with skip_array[i]==1 or write_mask_array[i]==0x00000000 are skipped for writing. write_reg(addr, data_wr & write_mask_array[i]) is called. Then for reading, registers with skip_array[i]==1, write_mask_array[i]==0x00000000, or read_mask_array[i]==0x00000000 are skipped. data_rd = read_reg(addr) & read_mask_array[i]. Expected value is computed as exp_val = ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])) where wr_n = write_mask_array[i] ^ 0xffffffff. Mismatches increment wr_fail_cnt. test_case() calls chk_rst_val() then chk_rd_wr(), and calls finish(1) if def_fail_cnt > 0 or wr_fail_cnt > 0, else finish(0). soft_reset_chk() is disabled inside #ifdef 0.",
        "Meta Test Steps / Procedure": "1. test_case() is called. 2. chk_rst_val() is invoked: for i=0 to CNT-1, addr=addr_array[i]. If skip_rst_array[i]==1, skip. If read_mask_array[i]==0x00000000, skip. Otherwise data_rd=read_reg(addr), data=(data_rd & 0xfffffffe), compare data against default_value_array[i]. On mismatch, increment def_fail_cnt. 3. chk_rd_wr() is invoked: chk_val[6]={0xffffffff, 0xaaaaaaaa, 0x55555555, 0xf5f5f5f5, 0xA5A5A5A5, 0xffff0000}. For j=0 to 5: data_wr=chk_val[j]. Write pass: for i=0 to CNT-1, addr=addr_array[i]. If skip_array[i]==1, skip. If write_mask_array[i]==0x00000000, skip. Else write_reg(addr, data_wr & write_mask_array[i]). Read pass: for i=0 to CNT-1, addr=addr_array[i]. If skip_array[i]==1, skip. If write_mask_array[i]==0x00000000, skip. If read_mask_array[i]==0x00000000, skip. Else data_rd=read_reg(addr) & read_mask_array[i]. wr_n=write_mask_array[i] ^ 0xffffffff. exp_val=((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])). Compare data_rd to exp_val. On mismatch, increment wr_fail_cnt. 4. If def_fail_cnt > 0 or wr_fail_cnt > 0, call finish(1) (fail). Else call finish(0) (pass).",
        "Meta Impacted Registers": "MIZAR_GPIO_GP0_GPIO_8; MIZAR_GPIO_GP0_GPIO_9; MIZAR_GPIO_GP0_GPIO_10",
        "Meta Validation / Acceptance Criteria": "Phase 1 (chk_rst_val): For each register where skip_rst_array[i]!=1 and read_mask_array[i]!=0x00000000, the value (read_reg(addr) & 0xfffffffe) must equal default_value_array[i]. Any mismatch increments def_fail_cnt. Phase 2 (chk_rd_wr): For each of 6 patterns and each register where skip_array[i]!=1, write_mask_array[i]!=0x00000000, and read_mask_array[i]!=0x00000000, the value (read_reg(addr) & read_mask_array[i]) must equal ((data_wr & read_mask_array[i] & write_mask_array[i]) | ((write_mask_array[i] ^ 0xffffffff) & read_mask_array[i] & default_value_array[i])). Any mismatch increments wr_fail_cnt. Final: test passes (finish(0)) only if both def_fail_cnt == 0 and wr_fail_cnt == 0. Otherwise finish(1) is called."
    },
    {
        "Index": "2",
        "SS / Module": "GPIO",
        "Test Case Name": "test_gpio_level_sel_intr_en",
        "Feature": "Level-Triggered Interrupt Enable",
        "Test Description": "This test verifies GPIO level-triggered interrupt generation and clearing for all 32 GPIO pins (pins 8 through 39). The test runs in two phases. In the first phase, each GPIO pin is configured in input mode with active-high level interrupt selection. The group interrupt is enabled for the pin under test, and an external stimulus is applied via an SRAM write. The interrupt service routine verifies that the per-pin raw interrupt status is asserted in gp0_gpio_8 through gp0_gpio_39, the group interrupt status is asserted in gp0_intr1_intr_sts1, then clears the interrupt and verifies it is cleared at both the per-pin and group levels, as well as at the system register level. In the second phase, the same procedure is repeated with active-low level interrupt selection, using an inverted stimulus pattern. The test passes only if all 32 pins pass both active-high and active-low interrupt verification with no errors.",
        "Speed": "NA",
        "Mode": "Interrupt Mode",
        "Memory Start Offset": "0x0",
        "Memory End Offset": "0x88",
        "Remarks": "GPIO pins 8-39 are accessed via computed addressing from the base register gp0_gpio_8 using the expression MIZAR_GPIO_GP0_GPIO_8 + (i*4) where i ranges from 0 to 31, covering per-pin registers gp0_gpio_8 through gp0_gpio_39. Two external system register macros (MIZAR_LSS_SYSREG_INTR_EN1 and MIZAR_LSS_SYSREG_RAW_STCR1) are used for system-level interrupt enable and status clearing but could not be resolved to GPIO register spec entries as they belong to the LSS_SYSREG block. The SRAM address 0xA0243ffc is used as an interrupt trigger stimulus and is not a GPIO register. The test uses GIC interrupt line 87 for GPIO0 context. The int_pend flag is polled in a while loop with wait_on(10) to synchronize with the interrupt handler. The active-high phase writes 0x00180000 (io_ctrl=1, level_sel=1) and the active-low phase writes 0x00100000 (io_ctrl=1, level_sel=0) to each per-pin register.",
        "Test Steps / Procedure": "1. Enable the GIC interrupt line for the GPIO block and enable the system-level interrupt for GPIO in the system register.\n2. For each GPIO pin (pins 8 through 39), configure the pin in input mode with active-high level interrupt selection by writing to the per-pin GPIO register (gp0_gpio_8 through gp0_gpio_39 via computed addressing).\n3. Enable the group interrupt for the current pin by writing the corresponding bit to gp0_intr1_intr_en1.\n4. Apply an interrupt stimulus by writing to the SRAM trigger location and wait for the interrupt service routine to execute.\n5. In the interrupt handler, verify that the per-pin raw interrupt status bit is asserted in the per-pin GPIO register.\n6. Verify that the group interrupt status bit for the current pin is asserted in gp0_intr1_intr_sts1.\n7. Clear the per-pin interrupt by writing the interrupt clear bit and verify the register reads back the expected cleared value.\n8. Disable the group interrupt and verify gp0_intr1_intr_sts1 reads zero.\n9. Clear the system register interrupt status and verify it is cleared. Clear the GIC interrupt.\n10. Repeat steps 2-9 for all 32 pins with active-high level selection.\n11. Repeat the entire sequence (steps 2-9) for all 32 pins with active-low level interrupt selection, using an inverted stimulus pattern.\n12. Evaluate the overall result: if any interrupt assertion, clearing, or status check failed across any pin or phase, report test failure. Otherwise, report test pass.",
        "Impacted Registers": "gp0_gpio_8; gp0_intr1_intr_en1; gp0_intr1_intr_sts1",
        "Validation / Acceptance Criteria": "1. For each GPIO pin (8-39), the per-pin raw interrupt status bit must be asserted in the corresponding per-pin GPIO register (gp0_gpio_8 through gp0_gpio_39) after the level-triggered interrupt stimulus is applied.\n2. The group interrupt status bit for the active pin must be asserted in gp0_intr1_intr_sts1.\n3. After clearing the per-pin interrupt, the per-pin register must read back the expected cleared value confirming the interrupt is no longer active.\n4. After disabling the group interrupt, gp0_intr1_intr_sts1 must read zero.\n5. The system register interrupt status must be successfully cleared after each interrupt service.\n6. All 32 GPIO pins must pass all validation checks in both the active-high and active-low level interrupt phases.\n7. The test passes only when the total error count is zero across all pins and both phases.",
        "Meta Headers": "<lss_sysreg.h>; <stdio.h>; <test_define.c>; <test_common.h>; <gpio/gpio_def.h>; <gpio/gpio_offset.h>",
        "Meta Macros": "CNT",
        "Meta Arrays": "addr_array[49]; default_value_array[49]; read_mask_array[49]; write_mask_array[49]; skip_array[49]",
        "Meta Test Description": "This testcase verifies GPIO level-triggered interrupt functionality for all 32 GPIO pins (GPIO 8 through GPIO 39). The test operates in two phases. Phase 1 (Active-High Level Interrupt): For each pin i (0 to 31), write_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4), 0x00180000) configures the pin in input mode (io_ctrl bit 20 = 1) with level_sel = 1 (bit 19 = 1) for active-high level interrupt. Then write_reg(MIZAR_GPIO_GP0_INTR1_INTR_EN1, 1<<i) enables the group interrupt for that pin. write_reg(0xA0243ffc, 0xffffffff) writes to an SRAM location to trigger the interrupt stimulus. The test sets int_pend=1 and polls while(int_pend==1) with wait_on(10) until the ISR clears int_pend. Phase 2 (Active-Low Level Interrupt): For each pin i, write_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4), 0x00100000) configures input mode with level_sel = 0 (bit 19 = 0) for active-low level interrupt. Group interrupt is enabled similarly. write_reg(0xA0243ffc, ~(1<<i)) writes inverted pattern to SRAM to trigger the low-level interrupt. Polling loop waits for ISR. In Default_IRQHandler(): rdata = read_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4)) reads the per-pin register. Checks (rdata & 0x2) != 0 to verify intr_raw_sts bit is set. Reads rdata_grp = read_reg(MIZAR_GPIO_GP0_INTR1_INTR_STS1) and checks (rdata_grp & (1<<i)) != 0 for group interrupt status. Clears interrupt by write_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4), 0x00110000) setting intr_clr bit 16. Reads back and verifies rdata == 0x100001 (interrupt cleared, io_ctrl=1, data_in=1). Disables group interrupt by write_reg(MIZAR_GPIO_GP0_INTR1_INTR_EN1, 0x00000000). Reads MIZAR_GPIO_GP0_INTR1_INTR_STS1 and verifies == 0x0. Clears sysreg status by write_reg(MIZAR_LSS_SYSREG_RAW_STCR1, LSS_SYSREG_RAW_STCR1_GPIO0_INTR) and verifies cleared via read_reg(MIZAR_LSS_SYSREG_RAW_STCR1). Clears GIC IRQ. On any failure, test_err is incremented. finish(test_err) is called at the end. Prior to the loops, GIC_EnableIRQ(87) is called for GPIO0 context, and write_reg(MIZAR_LSS_SYSREG_INTR_EN1, LSS_SYSREG_INTR_EN1_GPIO0_INTR) enables the sysreg-level interrupt.",
        "Meta Test Steps / Procedure": "1. GIC_EnableIRQ(87) is called (GPIO0 context). 2. write_reg(MIZAR_LSS_SYSREG_INTR_EN1, LSS_SYSREG_INTR_EN1_GPIO0_INTR) enables sysreg interrupt for GPIO0. 3. Active-high loop: for i=0 to 31: a. wr_val = 1<<i. b. write_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4), 0x00180000) sets io_ctrl=1, level_sel=1. c. wait_on(50). d. write_reg(MIZAR_GPIO_GP0_INTR1_INTR_EN1, wr_val) enables group interrupt for pin i. e. wait_on(10). f. write_reg(0xA0243ffc, 0xffffffff) SRAM stimulus. g. int_pend=1; while(int_pend==1) { printf; wait_on(10); }. 4. wait_on(100). 5. Active-low loop: for i=0 to 31: a. wr_val = 1<<i. b. write_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4), 0x00100000) sets io_ctrl=1, level_sel=0. c. wait_on(50). d. write_reg(MIZAR_GPIO_GP0_INTR1_INTR_EN1, wr_val) enables group interrupt for pin i. e. wait_on(10). f. write_reg(0xA0243ffc, ~(wr_val)) inverted SRAM stimulus. g. int_pend=1; while(int_pend==1) { printf; wait_on(10); }. 6. finish(test_err). 7. Default_IRQHandler(): a. wr_val=1<<i; int_pend=0. b. write_reg(0xA0243ffc, 0xffffffff) reset SRAM. c. rdata = read_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4)). d. Check (rdata & 0x2) != 0 (intr_raw_sts set). e. rdata_grp = read_reg(MIZAR_GPIO_GP0_INTR1_INTR_STS1). f. Check (rdata_grp & (1<<i)) != 0 (group interrupt set). g. write_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4), 0x00110000) clear interrupt (intr_clr bit). h. wait_on(20). i. rdata = read_reg(MIZAR_GPIO_GP0_GPIO_8 + (i*4)). j. Check rdata == 0x100001 (interrupt cleared). k. write_reg(MIZAR_GPIO_GP0_INTR1_INTR_EN1, 0x00000000) disable group interrupt. l. rdata_grp = read_reg(MIZAR_GPIO_GP0_INTR1_INTR_STS1). m. Check rdata_grp == 0x0 (group interrupt cleared). n. write_reg(MIZAR_LSS_SYSREG_RAW_STCR1, LSS_SYSREG_RAW_STCR1_GPIO0_INTR) clear sysreg status. o. rdata = read_reg(MIZAR_LSS_SYSREG_RAW_STCR1). p. Check (rdata & LSS_SYSREG_RAW_STCR1_GPIO0_INTR) == 0 (sysreg cleared). q. GIC_ClearIRQ(87). r. On any failure, test_err++.",
        "Meta Impacted Registers": "MIZAR_GPIO_GP0_GPIO_8; MIZAR_GPIO_GP0_INTR1_INTR_EN1; MIZAR_GPIO_GP0_INTR1_INTR_STS1; MIZAR_LSS_SYSREG_INTR_EN1; MIZAR_LSS_SYSREG_RAW_STCR1; 0xA0243ffc",
        "Meta Validation / Acceptance Criteria": "In Default_IRQHandler: 1. (rdata & 0x2) != 0x0 per-pin intr_raw_sts bit must be set after level interrupt trigger. 2. (rdata_grp & (1<<i)) != 0 group interrupt status bit for pin i must be set in MIZAR_GPIO_GP0_INTR1_INTR_STS1. 3. After writing 0x00110000 to clear interrupt: rdata == 0x100001 per-pin register must show interrupt cleared (io_ctrl=1, data_in=1, all interrupt bits cleared). 4. After disabling group interrupt (write 0x00000000 to MIZAR_GPIO_GP0_INTR1_INTR_EN1): rdata_grp == 0x0 group interrupt status must be zero. 5. After clearing sysreg status via write_reg(MIZAR_LSS_SYSREG_RAW_STCR1, LSS_SYSREG_RAW_STCR1_GPIO0_INTR): (rdata & LSS_SYSREG_RAW_STCR1_GPIO0_INTR) == 0 sysreg interrupt status must be cleared. 6. All 32 pins must pass all checks in both active-high and active-low phases. 7. finish(test_err) is called; test passes only if test_err == 0."
    }
]

# TestPlan columns
tp_cols = ['Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
           'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
           'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
           'Code Generation']

# MetaData columns
md_cols = ['Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
           'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
           'Meta Headers', 'Meta Macros', 'Meta Arrays']

# Create workbook
wb = openpyxl.Workbook()

# TestPlan sheet
ws_tp = wb.active
ws_tp.title = 'TestPlan'

# MetaData sheet
ws_md = wb.create_sheet('MetaData')

# Formatting
header_font = Font(bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
wrap_align = Alignment(wrap_text=True, vertical='top')

# Write TestPlan headers
for col_idx, col_name in enumerate(tp_cols, 1):
    cell = ws_tp.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

# Write TestPlan data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(tp_cols, 1):
        value = row_data.get(col_name, '')
        cell = ws_tp.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_align

# Write MetaData headers
for col_idx, col_name in enumerate(md_cols, 1):
    cell = ws_md.cell(row=1, column=col_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = wrap_align

# Write MetaData data
for row_idx, row_data in enumerate(json_data, 2):
    for col_idx, col_name in enumerate(md_cols, 1):
        value = row_data.get(col_name, '')
        cell = ws_md.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_align

# Auto-size columns
def auto_size(ws, columns, max_width=60):
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(col_name)
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        max_len = max(max_len, len(line))
        adjusted = min(max_len + 2, max_width)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = adjusted

auto_size(ws_tp, tp_cols)
auto_size(ws_md, md_cols)

# Freeze first row
ws_tp.freeze_panes = 'A2'
ws_md.freeze_panes = 'A2'

# Set MetaData sheet to veryHidden
ws_md.sheet_state = 'veryHidden'

# Save
script_dir = os.path.dirname(os.path.abspath(__file__))
output_path = os.path.join(script_dir, filename)
wb.save(output_path)
print(f'Workbook saved: {output_path}')
print(f'Filename: {filename}')
print(f'TestPlan rows: {len(json_data)}')
print(f'MetaData rows: {len(json_data)}')
