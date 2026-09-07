#!/usr/bin/env python3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timezone, timedelta
import os, sys, base64

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"GPIO_TestPlan_{timestamp}.xlsx"

json_data = [
    {"Index":"1","SS / Module":"GPIO","Test Case Name":"gpio_reg_wr_rd_test","Feature":"Register Write-Read Verification","Test Description":"This test validates the GPIO GP0 registers (gp0_gpio_8, gp0_gpio_9, gp0_gpio_10) in two phases. Phase 1 verifies that each register's default (reset) value matches the expected default value after applying a read mask. Phase 2 writes six distinct test patterns to each register (respecting write masks), reads back the values (respecting read masks), computes the expected result considering writable and non-writable bit fields, and verifies correctness. The test passes only if all default value checks and all write-read checks succeed across all registers and all patterns.","Speed":"NA","Mode":"NA","Memory Start Offset":"NA","Memory End Offset":"NA","Remarks":"The test uses a skip mechanism via skip_array and skip_rst_array to selectively bypass certain register indices during write-read and default-value checks respectively. The read data in the default value check phase is masked with a fixed mask that clears bit 0 before comparison. Six distinct bit patterns are used to exercise various bit combinations across writable fields. The soft_reset_chk function is disabled (inside #ifdef 0 block) and is not executed.","Test Steps / Procedure":"1. Read the default (reset) value of each GPIO GP0 register (gp0_gpio_8, gp0_gpio_9, gp0_gpio_10), skipping registers flagged in the reset-skip list or those that are not readable.\n2. Mask each read value and compare against the expected default value for that register. Record any mismatches.\n3. For each of six predefined test patterns, write the pattern to each register after applying the write mask, skipping registers flagged in the skip list or those that are not writable.\n4. After writing each pattern, read back each register value, applying the read mask.\n5. Compute the expected read-back value by combining the written data (masked by both read and write masks) with the default value for non-writable bits.\n6. Compare the actual read-back value against the computed expected value. Record any mismatches.\n7. Repeat steps 3 through 6 for all six test patterns.\n8. Determine overall test result: pass if no default-value mismatches and no write-read mismatches occurred; fail otherwise.","Impacted Registers":"gp0_gpio_8; gp0_gpio_9; gp0_gpio_10","Validation / Acceptance Criteria":"1. All GPIO GP0 registers (gp0_gpio_8, gp0_gpio_9, gp0_gpio_10) must return their expected default values when read after reset, after applying the read mask.\n2. For each of the six test patterns written to each register, the read-back value (masked by the read mask) must match the expected value computed from the written data, write mask, read mask, and default values for non-writable bits.\n3. The test passes (finish with success) only if zero default-value mismatches and zero write-read mismatches are recorded across all registers and all test patterns.\n4. Any single mismatch in either phase causes the test to fail (finish with failure).","Meta Headers":"<stdio.h>; <stdlib.h>; \"test_common.h\"; \"test_define.c\"; <gpio/gpio_def.h>; <gpio/gpio_offset.h>","Meta Macros":"SOFT_RST_REG_DATA; CNT","Meta Arrays":"addr_array[49]; default_value_array[49]; read_mask_array[49]; write_mask_array[49]; skip_array[49]; skip_rst_array[49]; chk_val[6]","Meta Test Description":"This testcase performs two phases of register validation on GPIO GP0 registers.","Meta Test Steps / Procedure":"1. Enter test_case(). 2. Call chk_rst_val(). 3. Call chk_rd_wr(). 4. Check results.","Meta Impacted Registers":"MIZAR_GPIO_GP0_GPIO_8; MIZAR_GPIO_GP0_GPIO_9; MIZAR_GPIO_GP0_GPIO_10","Meta Validation / Acceptance Criteria":"Phase 1 and Phase 2 checks as described."},
    {"Index":"2","SS / Module":"GPIO","Test Case Name":"test_gpio_level_sel_intr_en","Feature":"GPIO Level-Select Interrupt Enable Verification","Test Description":"This test validates the GPIO level-sensitive interrupt enable functionality for 32 GPIO pads. It runs in two phases.","Speed":"NA","Mode":"NA","Memory Start Offset":"NA","Memory End Offset":"NA","Remarks":"The test exercises both active-high and active-low level-sensitive interrupt modes across 32 GPIO pads.","Test Steps / Procedure":"1. Enable the GIC interrupt line for the GPIO controller.\n2. Enable the system-level interrupt.\n3-7. Configure, trigger, verify interrupts for all 32 pads in both phases.","Impacted Registers":"gp0_gpio_8; gp0_intr1_intr_en1; gp0_intr1_intr_sts1","Validation / Acceptance Criteria":"1-6. All validations must pass for all 32 GPIO pads in both phases.","Meta Headers":"<lss_sysreg.h>; <stdio.h>; <test_define.c>; <test_common.h>; <gpio/gpio_def.h>; <gpio/gpio_offset.h>","Meta Macros":"CNT","Meta Arrays":"addr_array[49]; default_value_array[49]; read_mask_array[49]; write_mask_array[49]; skip_array[49]","Meta Test Description":"This testcase validates GPIO level-sensitive interrupt enable functionality.","Meta Test Steps / Procedure":"1-8. Full procedure as described.","Meta Impacted Registers":"MIZAR_LSS_SYSREG_INTR_EN1; MIZAR_GPIO_GP0_GPIO_8; MIZAR_GPIO_GP0_INTR1_INTR_EN1; 0xA0243ffc; MIZAR_GPIO_GP0_INTR1_INTR_STS1; MIZAR_LSS_SYSREG_RAW_STCR1","Meta Validation / Acceptance Criteria":"IRQ handler checks as described."}
]

wb = openpyxl.Workbook()
ws_tp = wb.active
ws_tp.title = "TestPlan"
hf = Font(bold=True, color="FFFFFF", size=11)
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wa = Alignment(wrap_text=True, vertical="top")

tp_cols = ["Index","SS / Module","Feature","Test Case Name","Test Description","Speed","Mode","Memory Start Offset","Memory End Offset","Remarks","Test Steps / Procedure","Impacted Registers","Validation / Acceptance Criteria","Code Generation"]
for ci, cn in enumerate(tp_cols, 1):
    c = ws_tp.cell(row=1, column=ci, value=cn); c.font=hf; c.fill=hfill; c.alignment=wa

for ri, rd in enumerate(json_data, 2):
    for ci, cn in enumerate(tp_cols, 1):
        v = rd.get(cn, "")
        ws_tp.cell(row=ri, column=ci, value=v).alignment = wa

ws_tp.freeze_panes = "A2"

ws_md = wb.create_sheet("MetaData")
md_cols = ["Index","Test Case Name","Meta Test Description","Meta Test Steps / Procedure","Meta Impacted Registers","Meta Validation / Acceptance Criteria","Meta Headers","Meta Macros","Meta Arrays"]
for ci, cn in enumerate(md_cols, 1):
    c = ws_md.cell(row=1, column=ci, value=cn); c.font=hf; c.fill=hfill; c.alignment=wa

for ri, rd in enumerate(json_data, 2):
    for ci, cn in enumerate(md_cols, 1):
        v = rd.get(cn, "")
        ws_md.cell(row=ri, column=ci, value=v).alignment = wa

ws_md.freeze_panes = "A2"
ws_md.sheet_state = "veryHidden"

for ci in range(1, len(tp_cols)+1):
    ws_tp.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = 30
for ci in range(1, len(md_cols)+1):
    ws_md.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = 40

wb.save(filename)
with open(filename, "rb") as f:
    b64 = base64.b64encode(f.read()).decode()
print(f"FILENAME={filename}")
print(f"B64START")
print(b64)
print(f"B64END")
