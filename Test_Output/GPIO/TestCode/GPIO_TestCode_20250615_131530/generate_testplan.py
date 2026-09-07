#!/usr/bin/env python3
"""Auto-generated TestPlan Excel Generator for GPIO"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os
import datetime

def generate():
    wb = openpyxl.Workbook()
    
    # ===== TestPlan Sheet =====
    ws_tp = wb.active
    ws_tp.title = 'TestPlan'
    
    tp_headers = ['Index', 'SS / Module', 'Feature', 'Test Case Name', 'Test Description',
                  'Speed', 'Mode', 'Memory Start Offset', 'Memory End Offset', 'Remarks',
                  'Test Steps / Procedure', 'Impacted Registers', 'Validation / Acceptance Criteria',
                  'Code Generation']
    
    tp_data = {
        'Index': '1',
        'SS / Module': 'gpio',
        'Feature': 'Register Write Read Test',
        'Test Case Name': 'gpio_reg_wr_rd_test',
        'Test Description': 'This testcase validates GPIO register default values and write-read functionality across 49 GPIO registers including gp0_gpio_8, gp0_gpio_9, and gp0_gpio_10. In the first phase, each register is read and its value (with the least significant bit masked off) is compared against the expected default value. Registers that are not readable or are marked for skipping are excluded. In the second phase, six distinct data patterns are written to each writable register and then read back. The read-back value is compared against an expected value computed using the write mask, read mask, and default value for each register. The test passes only if all default value checks and all write-read checks succeed with zero mismatches.',
        'Speed': 'NA',
        'Mode': 'NA',
        'Memory Start Offset': 'NA',
        'Memory End Offset': 'NA',
        'Remarks': 'The addr_array contains 49 register entries but only 3 register macros are visible in test_define.c (the remaining entries are truncated in the source). Some registers are skipped for write-read testing via skip_array and for reset value checking via skip_rst_array. A comment in test_define.c notes that the din bit value may become 1 automatically if not forced, causing the bit-level select to go high and potentially causing default value mismatches. The soft_reset_chk function is defined but compiled out via #ifdef 0. The least significant bit is masked off during default value comparison (mask 0xfffffffe).',
        'Test Steps / Procedure': '1. Read all 49 GPIO registers and verify that each readable register contains its expected default reset value (with the least significant bit masked off). Skip registers marked as non-readable or excluded from reset checking.\n2. For each of six test data patterns, write the pattern to all writable GPIO registers, applying the appropriate write mask for each register. Skip registers marked as non-writable or excluded.\n3. After each write pass, read back all writable and readable GPIO registers and compare the read value against the expected value, which accounts for the write mask, read mask, and default value of each register.\n4. Verify that all default value checks and all write-read checks pass with zero mismatches. If any mismatch is detected, the test fails; otherwise the test passes.',
        'Impacted Registers': 'gp0_gpio_8; gp0_gpio_9; gp0_gpio_10',
        'Validation / Acceptance Criteria': '1. Each readable GPIO register must return its expected default reset value after masking the least significant bit.\n2. For each of six test patterns, after writing to each writable register, the read-back value must match the expected value computed from the write mask, read mask, and default value.\n3. The test passes only if zero mismatches are detected across both the default value check phase and the write-read check phase. Any mismatch causes test failure.',
        'Code Generation': ''
    }
    
    # ===== MetaData Sheet =====
    ws_md = wb.create_sheet('MetaData')
    
    md_headers = ['Index', 'Test Case Name', 'Meta Test Description', 'Meta Test Steps / Procedure',
                  'Meta Impacted Registers', 'Meta Validation / Acceptance Criteria',
                  'Meta Headers', 'Meta Macros', 'Meta Arrays']
    
    md_data = {
        'Index': '1',
        'Test Case Name': 'gpio_reg_wr_rd_test',
        'Meta Test Description': 'This testcase performs two phases of GPIO register validation. Phase 1 (chk_rst_val): Iterates over 49 register addresses in addr_array. For each register, it checks skip_rst_array to determine if the register should be skipped for reset value checking. It also checks read_mask_array; if the read mask is 0x00000000, the register is not readable and is skipped. For readable registers, it reads the register value using read_reg(addr), masks the read data with 0xfffffffe, and compares against default_value_array[i]. Mismatches increment def_fail_cnt. Phase 2 (chk_rd_wr): Iterates over 6 test patterns in chk_val[6] = {0xffffffff, 0xaaaaaaaa, 0x55555555, 0xf5f5f5f5, 0xA5A5A5A5, 0xffff0000}. For each pattern, it writes the pattern (masked with write_mask_array[i]) to each register address in addr_array, skipping entries where skip_array[i]==1 or write_mask_array[i]==0x00000000. Then it reads back each register, masking with read_mask_array[i], and computes the expected value as ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])) where wr_n = write_mask_array[i] ^ 0xffffffff. Mismatches increment wr_fail_cnt. The test passes (finish(0)) if both def_fail_cnt and wr_fail_cnt are zero; otherwise it fails (finish(1)). The addr_array contains register address macros including MIZAR_GPIO_GP0_GPIO_8, MIZAR_GPIO_GP0_GPIO_9, MIZAR_GPIO_GP0_GPIO_10. A soft_reset_chk function is defined but conditionally compiled out (#ifdef 0).',
        'Meta Test Steps / Procedure': '1. Call chk_rst_val(): Loop i from 0 to CNT-1 (49 registers). For each i: get addr = addr_array[i]. If skip_rst_array[i]==1, skip. If read_mask_array[i]==0x00000000, skip (not readable). Otherwise, data_rd = read_reg(addr); data = data_rd & 0xfffffffe; compare data with default_value_array[i]. If mismatch, increment def_fail_cnt. 2. Call chk_rd_wr(): Loop j from 0 to 5 over chk_val[6]={0xffffffff, 0xaaaaaaaa, 0x55555555, 0xf5f5f5f5, 0xA5A5A5A5, 0xffff0000}. For each pattern: (a) Write phase: loop i from 0 to CNT-1. If skip_array[i]==1, skip. If write_mask_array[i]==0x00000000, skip. Otherwise write_reg(addr, data_wr & write_mask_array[i]). (b) Read phase: loop i from 0 to CNT-1. If skip_array[i]==1, skip. If write_mask_array[i]==0x00000000, skip. If read_mask_array[i]==0x00000000, skip. Otherwise data_rd = read_reg(addr) & read_mask_array[i]; compute wr_n = write_mask_array[i] ^ 0xffffffff; exp_val = (data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i]); compare data_rd with exp_val. If mismatch, increment wr_fail_cnt. 3. Check if def_fail_cnt > 0 or wr_fail_cnt > 0: if yes, finish(1) (fail); else finish(0) (pass).',
        'Meta Impacted Registers': 'MIZAR_GPIO_GP0_GPIO_8; MIZAR_GPIO_GP0_GPIO_9; MIZAR_GPIO_GP0_GPIO_10',
        'Meta Validation / Acceptance Criteria': 'Phase 1 (default value check): For each readable, non-skipped register, (read_reg(addr) & 0xfffffffe) must equal default_value_array[i]. Any mismatch increments def_fail_cnt. Phase 2 (write-read check): For each writable and readable, non-skipped register, after writing (data_wr & write_mask_array[i]), the read-back value (read_reg(addr) & read_mask_array[i]) must equal ((data_wr & read_mask_array[i] & write_mask_array[i]) | (wr_n & read_mask_array[i] & default_value_array[i])). Any mismatch increments wr_fail_cnt. Final pass condition: def_fail_cnt == 0 AND wr_fail_cnt == 0 results in finish(0) (pass). Otherwise finish(1) (fail).',
        'Meta Headers': '<stdio.h>; <stdlib.h>; "test_common.h"; "test_define.c"; <gpio/gpio_def.h>; <gpio/gpio_offset.h>',
        'Meta Macros': 'SOFT_RST_REG_ADDRESS; SOFT_RST_REG_DATA; CNT; DEBUG_DISPLAY',
        'Meta Arrays': 'addr_array[49]; default_value_array[49]; read_mask_array[49]; write_mask_array[49]; skip_array[49]; skip_rst_array[49]; chk_val[6]'
    }
    
    # Format and populate
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    wrap = Alignment(wrap_text=True, vertical='top')
    
    for sheet, headers, data in [(ws_tp, tp_headers, tp_data), (ws_md, md_headers, md_data)]:
        for col_idx, h in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap
        for col_idx, h in enumerate(headers, 1):
            cell = sheet.cell(row=2, column=col_idx, value=data.get(h, ''))
            cell.alignment = wrap
        # Auto-size
        for col_idx, h in enumerate(headers, 1):
            max_len = len(str(h))
            val = str(data.get(h, ''))
            content_len = max(len(line) for line in val.split('\n')) if val else 0
            col_width = min(max(max_len, content_len) + 2, 60)
            sheet.column_dimensions[get_column_letter(col_idx)].width = col_width
        sheet.freeze_panes = 'A2'
    
    ws_md.sheet_state = 'veryHidden'
    
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'GPIO_TestPlan_20250615_131530.xlsx')
    wb.save(output_path)
    print(f'Workbook saved to: {output_path}')
    
    # Validate
    wb2 = openpyxl.load_workbook(output_path)
    assert 'TestPlan' in wb2.sheetnames
    assert 'MetaData' in wb2.sheetnames
    print('Validation PASSED')
    return output_path

if __name__ == '__main__':
    generate()
