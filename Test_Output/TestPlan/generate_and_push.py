#!/usr/bin/env python3
"""
Temporary helper - generates LPDDR4 TestPlan XLSX.
Run once, then delete this file.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import os

def generate():
    ist = timezone(timedelta(hours=5, minutes=30))
    now_ist = datetime.now(ist)
    ts = now_ist.strftime("%Y%m%d_%H%M%S")
    fname = f"LPDDR4_TestPlan_{ts}.xlsx"
    fpath = os.path.join(os.path.dirname(os.path.abspath(__file__)), fname)

    cols = [
        "Index", "SS / Module", "Test Case Name", "Feature",
        "Headers", "Macros", "Arrays", "Speed", "Mode",
        "Memory Start Offset", "Memory End Offset",
        "Test Description", "Test Steps / Procedure",
        "Impacted Registers", "Validation / Acceptance Criteria", "Remarks"
    ]

    widths = [8, 14, 40, 22, 35, 45, 10, 30, 30, 20, 20, 60, 65, 18, 65, 60]

    data = {
        "Index": "1",
        "SS / Module": "LPDDR4",
        "Test Case Name": "lpddr4_mem_half_data_bus_width_test",
        "Feature": "Half Data Bus Width",
        "Headers": '#include <stdio.h>\n#include <stdlib.h>\n#include "test_common.h"\n#include "lpddr4.h"',
        "Macros": (
            "APS_DRAM - Preprocessor conditional compilation flag for APS DRAM configuration (used in #if defined(APS_DRAM)); no #define value; externally defined.\n"
            "MPS_DRAM - Preprocessor conditional compilation flag for MPS DRAM configuration (used in #else path of APS_DRAM); no #define value; externally defined.\n"
            "SG2667 - Preprocessor conditional compilation flag for Speed Grade 2667 selection (used in #if defined(SG2667)); no #define value; externally defined.\n"
            "SG2133 - Preprocessor conditional compilation flag for Speed Grade 2133 selection (used in #elif defined(SG2133)); no #define value; externally defined.\n"
            "A53_INITIATOR - Preprocessor conditional compilation flag for A53 initiator selection (used in #if defined(A53_INITIATOR)); no #define value; externally defined.\n"
            "AI_INITIATOR - Preprocessor conditional compilation flag for AI initiator selection (used in #if defined(AI_INITIATOR)); no #define value; externally defined.\n"
            "DSP_INITIATOR - Preprocessor conditional compilation flag for DSP initiator selection (used in #if defined(DSP_INITIATOR)); no #define value; externally defined."
        ),
        "Arrays": "NA",
        "Speed": "SG3200 (default), SG2667 (when SG2667 defined), SG2133 (when SG2133 defined)",
        "Mode": "Half Data Bus Width (bus_width = 1)",
        "Memory Start Offset": "0x100",
        "Memory End Offset": "0x108",
        "Test Description": (
            "Verify LPDDR4 memory data integrity in half data bus width mode after training. "
            "The test configures the memory controller for half bus width operation, performs LPDDR4 training, "
            "waits for stabilization, then reads back data from both memory ports at specific offsets and verifies "
            "the data matches expected patterns. The test supports both APS_DRAM and MPS_DRAM configurations "
            "with multiple speed grades (3200, 2667, 2133)."
        ),
        "Test Steps / Procedure": (
            "1. Perform GPV programming to set up the bus interconnect.\n"
            "2. Initialize error counter to 0.\n"
            "3. Configure port addresses and controller/PHY base addresses based on DRAM type (APS_DRAM or MPS_DRAM).\n"
            "4. Set bus width to Half Bus mode (bus_width = 1).\n"
            "5. Select speed grade: SG3200 (default), SG2667, or SG2133.\n"
            "6. Disable Data Bus Inversion (DBI = 0).\n"
            "7. Disable Data Mask (DM_EN = 0).\n"
            "8. Execute LPDDR4 memory training sequence.\n"
            "9. Confirm training completion.\n"
            "10. Wait 1000 units for stabilization.\n"
            "11. Conditionally read 64-bit data from Port0 at offsets 0x100 and 0x108 and verify against expected patterns (condition based on initiator/DRAM combination).\n"
            "12. Read 64-bit data from Port1 at offsets 0x100 and 0x108 and verify against expected patterns.\n"
            "13. Report test result based on accumulated error count."
        ),
        "Impacted Registers": "NA",
        "Validation / Acceptance Criteria": (
            "1. Port0 (conditional): Data read from offset 0x100 must equal 0x3333333333333333 and data read from offset 0x108 must equal 0x2222222222222222.\n"
            "2. Port1 (unconditional): Data read from offset 0x100 must equal 0x3333333333333333 and data read from offset 0x108 must equal 0x2222222222222222.\n"
            "3. The test passes when the error counter (err0) is 0, indicating all read-back data matched the expected patterns.\n"
            "4. The test fails when err0 is greater than 0, indicating one or more data mismatches were detected.\n"
            "5. Note: The error condition uses logical AND (&&), so both values at offsets 0x100 and 0x108 must mismatch simultaneously to trigger an error."
        ),
        "Remarks": (
            "All 7 macros (APS_DRAM, MPS_DRAM, SG2667, SG2133, A53_INITIATOR, AI_INITIATOR, DSP_INITIATOR) are compile-time configuration flags "
            "expected to be defined externally via compiler command-line options. Port0 read-back verification is conditionally compiled using an "
            "XOR-based preprocessor expression involving initiator and DRAM type combinations. Port1 read-back verification is always compiled. "
            "The variables ctl_base, phy_base, bus_width, SG, DBI, DM_EN are declared in included headers (lpddr4.h or test_common.h) and assigned "
            "within test_case(). The local variables index and offset are declared but their usage is not visible in the available source context."
        )
    }

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestPlan"

    hf = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    hfill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cf = Font(name="Calibri", size=10)
    ca = Alignment(horizontal="left", vertical="top", wrap_text=True)
    bd = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = ha
        cell.border = bd

    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=2, column=i, value=data.get(c, ""))
        cell.font = cf
        cell.alignment = ca
        cell.border = bd

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 300
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}2"

    wb.save(fpath)
    sz = os.path.getsize(fpath)
    print(f"SUCCESS|{fname}|{fpath}|{sz}")
    return fname, fpath, sz

if __name__ == "__main__":
    generate()
