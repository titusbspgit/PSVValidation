#!/usr/bin/env python3
"""
Excel TestPlan Generator for LPDDR4
Generates LPDDR4_TestPlan_YYYYMMDD_HHMMSS.xlsx
Run: python3 generate_excel.py
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta
import os, sys

IST = timezone(timedelta(hours=5, minutes=30))
now_ist = datetime.now(IST)
timestamp = now_ist.strftime("%Y%m%d_%H%M%S")
filename = f"LPDDR4_TestPlan_{timestamp}.xlsx"

data = [
  {
    "Index": "1",
    "SS / Module": "LPDDR4",
    "Test Case Name": "lpddr4_mem_half_data_bus_width_test",
    "Feature": "Half Data Bus Width",
    "Meta Headers": "#include <stdio.h>\n#include <stdlib.h>\n#include \"test_common.h\"\n#include \"lpddr4.h\"",
    "Meta Macros": "APS_DRAM - Preprocessor conditional compilation flag for APS DRAM configuration (used in #if defined(APS_DRAM)); no #define value; externally defined.\nMPS_DRAM - Preprocessor conditional compilation flag for MPS DRAM configuration (used in #else path of APS_DRAM); no #define value; externally defined.\nSG2667 - Preprocessor conditional compilation flag for Speed Grade 2667 selection (used in #if defined(SG2667)); no #define value; externally defined.\nSG2133 - Preprocessor conditional compilation flag for Speed Grade 2133 selection (used in #elif defined(SG2133)); no #define value; externally defined.\nA53_INITIATOR - Preprocessor conditional compilation flag for A53 initiator selection (used in #if defined(A53_INITIATOR)); no #define value; externally defined.\nAI_INITIATOR - Preprocessor conditional compilation flag for AI initiator selection (used in #if defined(AI_INITIATOR)); no #define value; externally <NRP> conditional compilation flag for DSP initiator selection (used in #if defined(DSP_INITIATOR)); no #define value; externally defined.",
    "Meta Arrays": "NA",
    "Speed": "SG3200 (default), SG2667 (when SG2667 defined), SG2133 (when SG2133 defined)",
    "Mode": "Half Data Bus Width (bus_width = 1)",
    "Memory Start Offset": "0x100",
    "Memory End Offset": "0x108",
    "Meta Test Description": "This testcase verifies LPDDR4 memory data integrity when operating in half data bus width mode (bus_width = 1). The test configures the LPDDR4 controller and PHY base addresses based on the DRAM type (APS_DRAM or MPS_DRAM), sets bus_width to 1 (Half Bus), selects the speed grade (SG3200 default, SG2667, or SG2133 via preprocessor defines), disables Data Bus Inversion (<NRP> = 0) and <PERSON> (DM_EN = 0), performs LPDDR4 memory training via <PERSON>(), confirms training completion via training_done(), waits 1000 units for stabilization via wait_on(1000), then performs 64-bit read-back verification on <PERSON> (conditionally compiled based on XOR of initiator/DRAM combinations) and <PERSON> (unconditionally) at offsets 0x100 and 0x108. Read data is compared against expected patterns 0x3333333333333333 and 0x2222222222222222. Errors are accumulated in err0 and the test concludes with finish(err0).",
    "Test Description": "Verify LPDDR4 memory data integrity in half data bus width mode after training. The test configures the memory controller for half bus width operation, performs LPDDR4 training, waits for stabilization, then reads back data from both memory ports at specific offsets and verifies the data matches expected patterns. The test supports both APS_DRAM and MPS_DRAM configurations with multiple speed grades (3200, 2667, 2133).",
    "Meta Test Steps / Procedure": "1. File-scope declarations:\n   unsigned long int port0_addr;\n   unsigned long int port1_addr;\n \n2. Function: int test_case()\n \n3. Local variable declarations:\n   unsigned long int index;\n   unsigned long int offset;\n   unsigned long int read_data0;\n   unsigned long int read_data1;\n   unsigned long int read_data2;\n   unsigned long int read_data3;\n   unsigned int err0;\n \n4. Call gpv_programming();\n \n5. err0 = 0;\n \n6. Address configuration (preprocessor conditional):\n   #if defined(APS_DRAM)\n     port0_addr = 0;\n     port1_addr = 0x15A0000000;\n     ctl_base = 0x9EE03000;\n     phy_base = 0x9F000000;\n   #else\n     port0_addr = 0;\n     port1_addr = 0x11A0000000;\n     ctl_base = 0x11D003000;\n     phy_base = 0x11D500000;\n   #endif\n \n7. bus_width = 1; //0 -> Full Bus, 1 -> Half Bus, 2 -> Quarter bus\n \n8. Speed grade selection (preprocessor conditional):\n   #if defined(SG2667)\n     SG = 2667;\n   #elif defined(SG2133)\n     SG = 2133;\n   #else\n     SG = 3200;\n   #endif\n \n9. <NRP> = 0;\n \n10. DM_EN = 0;\n \n11. Call <PERSON>();\n \n12. Call training_done();\n \n13. Call wait_on(1000);\n \n14. Port0 read-back verification (preprocessor conditional):\n   #if !((defined(APS_DRAM) && defined(A53_INITIATOR)) ^(defined(MPS_DRAM) && defined(AI_INITIATOR)) ^(defined(MPS_DRAM) && defined(DSP_INITIATOR)))\n     read_reg64(port0_addr + 0x100, &read_data0);\n     read_reg64(port0_addr + 0x108, &read_data1);\n     if((read_data0 != 0x3333333333333333) && (read_data1 != 0x2222222222222222))\n     {\n       printf(\"ERROR: port0 read_data0 = %lx, read_data1 = %lx\\n\", read_data0, read_data1);\n       err0++;\n     }\n   #endif\n \n15. Port1 read-back verification (unconditional):\n   read_reg64(port1_addr + 0x100, &read_data2);\n   read_reg64(port1_addr + 0x108, &read_data3);\n   if((read_data2 != 0x3333333333333333) && (read_data3 != 0x2222222222222222))\n   {\n     printf(\"ERROR: port1 read_data2 = %lx, read_data3 = %lx\\n\", read_data2, read_data3);\n     err0++;\n   }\n \n16. Call finish(err0);",
    "Test Steps / Procedure": "1. Perform GPV programming to set up the bus interconnect.\n2. Initialize error counter to 0.\n3. Configure port addresses and controller/PHY base addresses based on DRAM type (APS_DRAM or MPS_DRAM).\n4. Set bus width to Half Bus mode (bus_width = 1).\n5. Select speed grade: SG3200 (default), SG2667, or SG2133.\n6. Disable Data Bus Inversion (<NRP> = 0).\n7. Disable Data Mask (DM_EN = 0).\n8. Execute LPDDR4 memory training sequence.\n9. Confirm training completion.\n10. Wait 1000 units for stabilization.\n11. Conditionally read 64-bit data from <LOCATION> at offsets 0x100 and 0x108 and verify against expected patterns (condition based on initiator/DRAM combination).\n12. Read 64-bit data from <LOCATION> at offsets 0x100 and 0x108 and verify against expected <LOCATION>. Report test result based on accumulated error count.",
    "Meta Impacted Registers": "NA",
    "Impacted Registers": "NA",
    "Meta Validation / Acceptance Criteria": "Port0 validation (conditionally compiled based on: #if !((defined(APS_DRAM) && defined(A53_INITIATOR)) ^(defined(MPS_DRAM) && defined(AI_INITIATOR)) ^(defined(MPS_DRAM) && defined(DSP_INITIATOR)))):\n  read_reg64(port0_addr + 0x100, &read_data0): expected read_data0 == 0x3333333333333333\n  read_reg64(port0_addr + 0x108, &read_data1): expected read_data1 == 0x2222222222222222\n  Fail condition: if((read_data0 != 0x3333333333333333) && (read_data1 != 0x2222222222222222)) then err0++ and printf error message\n \nPort1 validation (unconditional):\n  read_reg64(port1_addr + 0x100, &read_data2): expected read_data2 == <NRP>  read_reg64(port1_addr + 0x108, &read_data3): expected read_data3 == 0x2222222222222222\n  Fail condition: if((read_data2 != 0x3333333333333333) && (read_data3 != 0x2222222222222222)) then err0++ and printf error message\n \nPass condition: err0 == 0 (no mismatches detected)\nFail condition: err0 > 0 (one or more read-back mismatches)\nFinal call: finish(err0)\n \nNote: The comparison uses && (logical AND), meaning both values at offset 0x100 and 0x108 must mismatch simultaneously to trigger an error increment.",
    "Validation / Acceptance Criteria": "1. <LOCATION> (conditional): Data read from offset 0x100 must equal 0x3333333333333333 and data read from offset 0x108 must equal 0x2222222222222222.\n2. <LOCATION> (unconditional): Data read from offset 0x100 must equal 0x3333333333333333 and data read from offset 0x108 must equal 0x2222222222222222.\n3. The test passes when the error counter (err0) is 0, indicating all read-back data matched the expected patterns.\n4. The test fails when err0 is greater than 0, indicating one or more data mismatches were detected.\n5. Note: The error condition uses logical AND (&&), so both values at offsets 0x100 and 0x108 must mismatch simultaneously to trigger an error.",
    "Remarks": "All 7 macros (APS_DRAM, MPS_DRAM, SG2667, SG2133, A53_INITIATOR, AI_INITIATOR, DSP_INITIATOR) are compile-time configuration flags expected to be defined externally via compiler command-line options. Port0 read-back verification is conditionally compiled using an XOR-based preprocessor expression involving initiator and DRAM type combinations. Port1 read-back verification is always compiled. The variables ctl_base, phy_base, bus_width, SG, <NRP>, DM_EN are declared in included headers (lpddr4.h or test_common.h) and assigned within test_case(). The local variables index and offset are declared but their usage is not visible in the available source context.",
    "Code Generation": ""
  }
]

wb = openpyxl.Workbook()

# TestPlan sheet
ws_tp = wb.active
ws_tp.title = "TestPlan"
tp_columns = ["Index","SS / Module","Feature","Test Case Name","Test Description","Speed","Mode","Memory Start Offset","Memory End Offset","Remarks","Test Steps / Procedure","Impacted Registers","Validation / Acceptance Criteria","Code Generation"]
for ci, cn in enumerate(tp_columns, 1):
    ws_tp.cell(row=1, column=ci, value=cn)
for ri, rd in enumerate(data, 2):
    for ci, cn in enumerate(tp_columns, 1):
        ws_tp.cell(row=ri, column=ci, value=rd.get(cn, ""))

# MetaData sheet
ws_md = wb.create_sheet("MetaData")
md_columns = ["Index","Test Case Name","Meta Test Description","Meta Test Steps / Procedure","Meta Impacted Registers","Meta Validation / Acceptance Criteria","Meta Headers","Meta Macros","Meta Arrays"]
for ci, cn in enumerate(md_columns, 1):
    ws_md.cell(row=1, column=ci, value=cn)
for ri, rd in enumerate(data, 2):
    for ci, cn in enumerate(md_columns, 1):
        ws_md.cell(row=ri, column=ci, value=rd.get(cn, ""))

# Formatting
hf = Font(bold=True, color="FFFFFF")
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wa = Alignment(wrap_text=True, vertical="top")
ha = Alignment(wrap_text=True, vertical="top", horizontal="center")

for ws in [ws_tp, ws_md]:
    for ci in range(1, ws.max_column+1):
        c = ws.cell(row=1, column=ci)
        c.font = hf; c.fill = hfill; c.alignment = ha
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.alignment = wa
    for ci in range(1, ws.max_column+1):
        ml = 0
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=ci, max_col=ci):
            for cell in row:
                if cell.value:
                    for line in str(cell.value).split('\n'):
                        ml = max(ml, len(line))
        aw = min(ml+2, 60)
        if aw < 12: aw = 12
        ws.column_dimensions[get_column_letter(ci)].width = aw
    ws.freeze_panes = "A2"

ws_md.sheet_state = "veryHidden"

script_dir = os.path.dirname(os.path.abspath(__file__))
save_path = os.path.join(script_dir, filename)
wb.save(save_path)

# Verify
fs = os.path.getsize(save_path)
wb2 = openpyxl.load_workbook(save_path)
print(f"OK: {filename} ({fs} bytes) sheets={wb2.sheetnames}")
